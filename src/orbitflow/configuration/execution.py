"""Approved predefined execution: preparation/confirmation independent of Excel.

Workbook ingestion is an adapter. Confirmation records are local replay guards,
not authentication. No writes occur until confirm(..., apply=True).
Fresh target/ownership/prerequisite checks do not replan or replace approved CLI.
Independent rows share stage observations; verification and EXEC-only save remain
separate from command delivery.
"""
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from contextlib import contextmanager
import hashlib
import math
import os
from pathlib import Path
import re
import sqlite3
from time import monotonic
from uuid import uuid4

from orbitflow.progress import Progress
from orbitflow.timing import profiling, correlation, stage, measured, timed_call, timed_context, command_category
from orbitflow.configuration.plans import canonical, decode, identity, require
from orbitflow.configuration.v2.catalogue import BY_ID
from orbitflow.configuration.v2.planning import BulkPlan, evaluate_decisions
from orbitflow.config import ExecutionConfig
from orbitflow.logging import module_logger, sanitize_text, capability_metadata
from orbitflow.result_spool import ResultSpool, ResultJournal
from orbitflow.execution import execute_devices
from orbitflow.capabilities.save import command_rejected, save_configuration
from orbitflow.capabilities.switchport import SwitchportService
from orbitflow.capabilities.interfaces import InterfaceService
from orbitflow.capabilities.vlans import VlanService
from orbitflow.capabilities.configuration import ConfigurationService
from orbitflow.configuration_backup import BackupWriter
from orbitflow.configuration.v2.observation import snapshot
from orbitflow.configuration import execution_observation as observations
from orbitflow.inventory import DeviceInventoryResolver, JsonInventoryStore
from orbitflow.transport import connect_device
from orbitflow.vendors.common import DeviceCLI
from orbitflow.vendors.configuration_facts import _safe_source, interface_settings
from orbitflow.vendors.aggregation import for_observation, relationship_uncertain
from orbitflow.vendors.interface_names import canonical_interface_name
from orbitflow.vendors.vlan_types import parse_vlan_list
from orbitflow.vendors.cisco import configuration_execution as cisco_execution
from orbitflow.vendors.cisco import evc_execution
from orbitflow.vendors.ubiquiti import configuration_execution as edge_execution


def _now():
    return datetime.now(timezone.utc).isoformat()


class ExecutionBlocked(Exception):
    """A fixed execution gate code, never device text or arbitrary exception data."""
    def __init__(self, code):
        self.code = code
        super().__init__(code)


def _gate(condition, code):
    if not condition:
        raise ExecutionBlocked(code)


@dataclass(frozen=True)
class ApplySettings:
    """Separate apply limit; all timeouts are finite seconds."""
    max_concurrent_devices: int = 1
    command_timeout: float = 30.0
    save_timeout: float = 300.0
    connection_start_interval: float = 1.0
    connection_retry_delay: float = 5.0

    def __post_init__(self):
        ExecutionConfig(self.max_concurrent_devices, self.connection_start_interval, self.connection_retry_delay)
        for value in (self.command_timeout, self.save_timeout):
            require(type(value) in (int, float) and math.isfinite(value) and value > 0, 'invalid apply timeout')


def verification_commands(request, result):
    t, p = BY_ID[request['template_id']], request['inputs']
    target = result.get('resolved_target', p.get('interface', ''))
    if t.family not in {'C3750X', 'C3850X', 'EdgeSwitch', 'ME3600X', 'ASR920'}:
        return []
    adapter = edge_execution if t.family == 'EdgeSwitch' else cisco_execution
    return adapter.verification_commands(t.function, p, target)


class ExecutionService:
    """Local transactional confirmation store. Keep authority directory protected."""
    def __init__(self, authority_dir='data/configuration/execution'):
        self.path = Path(authority_dir)
        self.path.mkdir(parents=True, exist_ok=True)
        with self._db() as db:
            db.execute('CREATE TABLE IF NOT EXISTS confirmations (reference TEXT PRIMARY KEY, digest TEXT UNIQUE, payload TEXT, state TEXT, run_path TEXT)')

    @contextmanager
    def _db(self):
        db = sqlite3.connect(self.path / 'confirmations.sqlite3', timeout=30)
        try:
            with db:
                yield db
        finally:
            db.close()

    def prepare(self, original, decisions, *, expected_digest, settings=ApplySettings()):
        """No connection. Return exact summary/reference; decisions is row-ID mapping."""
        original = BulkPlan(original.content)
        require(original.digest == expected_digest, 'source digest mismatch')
        settings = ApplySettings(**asdict(settings))
        data = original.to_dict()
        from orbitflow.configuration.v2.planning import UNVERIFIED_SOURCE
        require(not data['source'].startswith(UNVERIFIED_SOURCE),
                'unverified historical recovery cannot authorize execution')
        intent = evaluate_decisions(original, decisions, expected_digest=expected_digest,
                                    now=datetime.fromisoformat(data['evaluated_at'])).to_dict()
        selected, skipped = [], []
        results = {r['row_id']: r for r in intent['results']}
        for request in data['requests']:
            result = results[request['row_id']]
            t = BY_ID[request['template_id']]
            reason = ('UNSUPPORTED_TEMPLATE' if t.family not in {'C3750X', 'C3850X', 'EdgeSwitch', 'ME3600X', 'ASR920'} else
                      result['decision'] if result['decision'] != 'Approve' else
                      'BLOCKED' if not result['deployment_ready'] else '')
            if not reason and t.function in {'Access Port', 'Trunk Port', 'VLAN Port', 'EVC'}:
                observation = data['observations'].get(request['inputs']['management_ip'], {})
                platform = observation.get('identity', {}).get('platform', '')
                aggregation = for_observation(observation, platform)
                source = request['inputs']['interface']
                if aggregation.owner(source):
                    reason = 'NOT_ALLOW_ON_CHILD_PORT_CHANNEL'
                elif relationship_uncertain(aggregation, source,
                        [c for n in observation.get('facts', [])
                         if n['kind'] in {'interface', 'l2_interface'}
                         and canonical_interface_name(platform, n['value'], normalize_aggregation=True)
                         == canonical_interface_name(platform, source, normalize_aggregation=True)
                         for c in n['children']]) or aggregation.capture_incomplete:
                    reason = 'OWNERSHIP_REQUIRES_REVIEW'
            entry = dict(request=request, result=result,
                         identity=data['observations'].get(request['inputs']['management_ip'], {}).get('identity'),
                         verification_commands=verification_commands(request, result))
            if reason:
                skipped.append(dict(row_id=request['row_id'], decision=result['decision'], reason=reason, request=request, identity=entry['identity'], blockers=result.get('blockers', result['findings'])))
            else:
                selected.append(entry)
        # Dependency closure is evaluated against executable rows, in order.
        ready = set()
        executable = []
        for entry in selected:
            if not set(entry['result']['dependencies']) <= ready:
                skipped.append(dict(row_id=entry['request']['row_id'], decision=entry['result']['decision'], reason='DEPENDENCY_NOT_EXECUTABLE', request=entry['request'], identity=entry['identity']))
            else:
                executable.append(entry)
                ready.add(entry['request']['row_id'])
        summary = dict(source_digest=original.digest, selected=executable, skipped=skipped,
                       settings=asdict(settings),
                       save_behavior='Save entire running configuration once after attempted writes, including partial changes. EdgeSwitch may interrupt management; recognized save question is answered y.')
        digest = hashlib.sha256(canonical(summary).encode()).hexdigest()
        payload = canonical(dict(summary=summary, plan=data, decisions=intent['decisions']))
        reference = uuid4().hex
        # Added presentation metadata must not make a consumed pre-revision
        # confirmation executable again. Reuse its exact original bound payload.
        legacy = dict(summary, skipped=[{k: v for k, v in e.items() if k not in {'request', 'identity'}} for e in skipped])
        legacy_digest = hashlib.sha256(canonical(legacy).encode()).hexdigest()
        with self._db() as db:
            row = db.execute("SELECT reference, state, run_path, digest, payload FROM confirmations WHERE digest IN (?, ?) ORDER BY CASE WHEN state='awaiting_confirmation' THEN 1 ELSE 0 END", (digest, legacy_digest)).fetchone()
            if row is None:
                db.execute('INSERT OR IGNORE INTO confirmations VALUES (?, ?, ?, ?, ?)', (reference, digest, payload, 'awaiting_confirmation', ''))
                row = db.execute('SELECT reference, state, run_path, digest, payload FROM confirmations WHERE digest=?', (digest,)).fetchone()
        digest = row[3]
        summary = decode(row[4])['summary']
        with module_logger('configuration', 'execution') as (logger, _):
            logger.info(f'Prepare; digest={original.digest}; confirmation={row[0]}; summary={digest}')
        return dict(status=row[1], reference=row[0], summary_digest=digest, summary=summary, run_path=row[2])

    def prepare_workbook(self, original, workbook, *, expected_digest, settings=ApplySettings()):
        from orbitflow.configuration.v2.workbook import consume_decisions
        intent = consume_decisions(original, workbook, expected_digest=expected_digest,
                                   now=datetime.fromisoformat(original.to_dict()['evaluated_at']))
        return self.prepare(original, intent.to_dict()['decisions'], expected_digest=expected_digest, settings=settings)

    def confirm(self, reference, *, summary_digest, apply=False, credentials=None,
                transport_config=None, inventory_path='data/inventory/inventory.json',
                run_root='outputs/runs/configuration_execution', backups_dir=None, progress=None):
        """Consume confirmation before connecting; replay never starts work.

        backups_dir=None stores captures in the run; explicit overrides allocate
        an external ROOT/run-id directory and retain relative references.
        """
        require(apply is True, 'explicit apply intent required')
        with self._db() as db:
            db.execute('BEGIN IMMEDIATE')
            record = db.execute('SELECT digest, payload, state, run_path FROM confirmations WHERE reference=?', (reference,)).fetchone()
            require(record is not None and record[0] == summary_digest, 'confirmation summary mismatch')
            require(record[2] == 'awaiting_confirmation', 'confirmation already consumed')
            payload = decode(record[1])
            require(hashlib.sha256(canonical(payload['summary']).encode()).hexdigest() == summary_digest, 'corrupt confirmation')
            # Revalidate immutable content before claiming; authority state is local, not authentication.
            original = BulkPlan(canonical(payload['plan']))
            require(original.digest == payload['summary']['source_digest'], 'confirmation source mismatch')
            db.execute("UPDATE confirmations SET state='claimed' WHERE reference=?", (reference,))
        summary = payload['summary']
        try:
            def started(path):
                with self._db() as db:
                    db.execute('UPDATE confirmations SET run_path=? WHERE reference=?', (str(path), reference))
            spool = _run(summary, credentials or {}, transport_config, inventory_path=inventory_path,
                         run_root=run_root, backups_dir=backups_dir, on_started=started, progress=progress)
            with self._db() as db:
                db.execute("UPDATE confirmations SET state='completed', run_path=? WHERE reference=?", (str(spool.path), reference))
            return spool
        except BaseException:
            with self._db() as db:
                db.execute("UPDATE confirmations SET state='interrupted' WHERE reference=?", (reference,))
            raise


@measured('full_observation')
def _observe(session, context, cli, entries, timeout):
    interfaces = timed_call('interface_collection', InterfaceService(timeout=timeout).collect, session, context, cli=cli)
    vlans = timed_call('vlan_config_collection', VlanService(timeout=timeout).collect, session, context, cli=cli, preserve_full_evidence=True)
    obs = timed_call('normalized_snapshot', observations.project, session, context, cli, interfaces, vlans, snapshot)
    vids = set()
    for entry in entries:
        p = entry['request']['inputs']
        vids.update(p[k] for k in ('vlan_id', 'pvid', 'native_vlan') if k in p)
        vids.update(p.get('vlans', []))
        vids.update(parse_vlan_list(p.get('tagged_vlans', '')))
    obs['vlan_lookups'] = {str(v): timed_call('vlan_lookup', VlanService(timeout=timeout).lookup, session, context, v, cli=cli) for v in sorted(vids)}
    if evc_execution.supported_context(context):
        evcs = [e for e in entries if BY_ID[e['request']['template_id']].function == 'EVC']
        if context.device_family == 'ME3600X':
            for vid in sorted({e['request']['inputs']['bridge_domain'] for e in evcs}):
                obs['vlan_lookups'][str(vid)] = timed_call('vlan_lookup', VlanService(timeout=timeout).lookup, session, context, vid, cli=cli)
        obs['service_details'] = {}
        for e in evcs:
            p = e['request']['inputs']
            target = e['result'].get('resolved_target') or p['interface']
            obs['service_details'][target + ':' + str(p['service_instance_id'])] = timed_call(
                'service_detail', VlanService(timeout=timeout).service_detail, session, context, target, p['service_instance_id'], cli=cli)
    return obs


def _set(value):
    if value == 'ALL':
        return set(range(1, 4095))
    if value == 'NONE':
        return set()
    return set(value) if isinstance(value, list) else None


@measured('pre_row_target_checks')
def _check_target(entry, observed, context):
    """Validate actual interface identity/ownership without planning or rendering.

    The immutable approval supplies commands and expected results. Fresh evidence
    supplies target safety only; warnings and state differences are not apply gates.
    """
    request, approved = entry['request'], entry['result']
    function = BY_ID[request['template_id']].function
    if function == 'VLAN Database':
        return
    source = request['inputs']['interface']
    target = approved.get('resolved_target') or source
    key = lambda name: canonical_interface_name(context.platform, name, normalize_aggregation=True)
    for name in {key(source), key(target)}:
        _gate(sum(key(i['port_name']) == name for i in observed['interfaces']) == 1,
              'INTERFACE_ABSENT_OR_AMBIGUOUS')
        _gate(len(interface_settings(observed['facts'], context.platform, name)) == 1,
              'INTERFACE_DECLARATION_ABSENT_OR_AMBIGUOUS')
    adapter = edge_execution if context.platform == 'ubiquiti_edgeswitch' else cisco_execution
    code = (evc_execution.target_blocker(entry, observed, context) if function == 'EVC' else
            adapter.target_blocker(function, source, target, observed, context.platform))
    _gate(not code, code)


@measured('row_verification', failure='Unknown')
def verify_row(session, context, cli, entry, before, *, timeout=30.0, correlation=None):
    """Read-only verification returning Verified/Failed/Unknown and scoped evidence.

    Access-only missing operational VLAN may use reliable explicit exact-target
    normalized config; preserve parser unknowns and disclose the fallback source.
    EVC verification independently compares configured child intent, operational
    detail and preserved parent/sibling settings on ME3600X/ASR920 profiles.
    """
    return _verify_observed(session, context, cli, entry, before, timeout=timeout, correlation=correlation)


def _verify_observed(session, context, cli, entry, before, *, timeout=30.0, correlation=None, after=None):
    """Verify a target against a fresh stage observation; never infer intermediate effects."""
    from orbitflow.configuration.v2.evidence import scoped_evidence
    request, result = entry['request'], entry['result']
    t, p = BY_ID[request['template_id']], request['inputs']
    if t.function == 'VLAN Database':
        lookup = after['vlan_lookups'][str(p['vlan_id'])] if after is not None else timed_call('vlan_lookup', VlanService(timeout=timeout).lookup, session, context, p['vlan_id'], cli=cli)
        status = ('Verified' if lookup.get('name') == p['name'] else 'Failed') if lookup['status'] == 'present' else 'Failed' if lookup['status'] == 'absent' else 'Unknown'
        return dict(status=status, lookup=lookup, warnings=['VLAN_INACTIVE'] if lookup.get('operational_status') in {'suspended', 'inactive', 'act/lshut'} else [])
    if after is None:
        after = _observe(session, context, cli, [entry], timeout)
    if t.function == 'EVC':
        return evc_execution.verify(entry, before, after, context)
    target = result.get('resolved_target', p['interface'])
    key = lambda value: canonical_interface_name(context.platform, value, normalize_aggregation=True)
    actual = [i for i in after['interfaces'] if key(i['port_name']) == key(target)]
    previous = [i for i in before['interfaces'] if key(i['port_name']) == key(target)]
    if len(actual) != 1 or len(previous) != 1:
        return dict(status='Unknown', error_category='InterfaceEvidenceUnavailable')
    description = p.get('description', previous[0]['port_description'])
    checks = [actual[0]['port_description'] == description]
    # Retain scoped config and prove admin setting preservation independently of link state.
    def settings(obs):
        nodes = interface_settings(obs['facts'], context.platform, target)
        return nodes[0]['children'] if len(nodes) == 1 else None
    old, new = settings(before), settings(after)
    if old is None or new is None:
        return dict(status='Unknown', error_category='ConfigurationEvidenceUnavailable')
    modified = {'description'} if t.function == 'Interface Description' else (
        {'description', 'mode', 'access', 'routed'} if t.function == 'Access Port' else
        {'description', 'mode', 'allowed', 'allowed_add', 'allowed_remove', 'allowed_except', 'native'} if t.function == 'Trunk Port' else
        {'description', 'include', 'exclude', 'tagged', 'untagged_vlans', 'pvid'})
    if t.function == 'Trunk Port' and context.device_family == 'C3750X' and 'switchport trunk encapsulation dot1q' in result['commands']:
        modified.add('trunk_encapsulation')
        checks.append([n['value'] for n in new if n['kind'] == 'trunk_encapsulation'] == ['dot1q'])
    # Compare stable fact values, not capture line numbers. Omitted fields remain
    # protected separately below; unrelated settings must survive every template.
    def signature(node):
        return (node['kind'], node['value'], [signature(child) for child in node.get('children', [])])
    def preserved(nodes):
        return [signature(n) for n in nodes if n['kind'] not in modified]
    checks.append(preserved(old) == preserved(new))
    evidence = dict(configuration=timed_call('scoped_evidence', scoped_evidence, request, after), description=actual[0], preservation='Matched' if all(checks) else 'Description or unrelated setting mismatch')
    if t.function != 'Interface Description':
        try:
            switchport = SwitchportService(timeout=timeout).collect(session, context, target, cli=cli, correlation=correlation).to_dict()
        except Exception as exc:
            evidence['warnings'] = ['Operational switchport evidence unavailable: ' + type(exc).__name__]
            return dict(status='Unknown', error_category=type(exc).__name__, evidence=evidence)
        evidence['switchport'] = switchport
        f = switchport['facts']
        if t.function == 'Access Port':
            access = f['access_vlan']
            if access is None:
                missing_access_field = 'Access Mode VLAN' not in switchport.get('fields', {})
                evidence['missing_fields' if missing_access_field else 'invalid_fields'] = ['Access Mode VLAN']
                explicit = [n['value'] for n in new if n['kind'] == 'access']
                if missing_access_field and after.get('configuration_reliable') is True and f['enabled'] is True and f['administrative_mode'] == 'static access' and len(explicit) == 1 and isinstance(explicit[0], list) and len(explicit[0]) == 1 and type(explicit[0][0]) is int and 1 <= explicit[0][0] <= 4094:
                    access = explicit[0][0]
                    evidence['access_vlan_source'] = 'fresh exact-interface explicit configuration'
                    evidence['warnings'] = ['ACCESS_VLAN_CONFIG_FALLBACK: Access Mode VLAN not reported']
                else:
                    evidence['warnings'] = [('Access Mode VLAN not reported' if missing_access_field else 'Access Mode VLAN malformed') + '; unambiguous explicit configuration fallback unavailable']
                with module_logger('configuration', 'execution') as (logger, _):
                    logger.warning('Access Mode VLAN ' + ('not reported' if missing_access_field else 'malformed') + '; verification source=' + evidence.get('access_vlan_source', 'unavailable'),
                                   extra=capability_metadata(context, correlation, interface=target, stage='verification', outcome='Fallback' if access is not None else 'Unknown'))
            evidence['access_vlan_check'] = dict(actual=access, requested=p['vlan_id'])
            required = [f['enabled'], f['administrative_mode'], access]
            if any(v is None for v in required):
                return dict(status='Failed' if not all(checks) or f['enabled'] is False or f['administrative_mode'] not in (None, 'static access') else 'Unknown', evidence=evidence)
            checks += [f['enabled'] is True, f['administrative_mode'] == 'static access', access == p['vlan_id']]
            evidence['vlan_lookup'] = after['vlan_lookups'].get(str(p['vlan_id']))
            lookup = evidence['vlan_lookup'] or {}
            evidence.setdefault('warnings', []).extend(['GLOBAL_VLAN_NOT_ACTIVE'] if lookup.get('status') == 'absent' or lookup.get('operational_status') in {'inactive', 'suspended', 'act/lshut'} else [])
        elif t.function == 'Trunk Port':
            desired = result['desired_allowed_state']
            values = set(desired['values'])
            allowed = set(range(1, 4095)) - values if desired['all'] else values
            native = result['projected_state']['native']
            if any(f[k] is None for k in ('enabled', 'administrative_mode', 'allowed_vlans', 'native_vlan')) or len(native) != 1:
                evidence['warnings'] = ['Missing operational fields: ' + ', '.join(k for k in ('enabled', 'administrative_mode', 'allowed_vlans', 'native_vlan') if f[k] is None)]
                return dict(status='Unknown', evidence=evidence)
            checks += [f['enabled'] is True, f['administrative_mode'] == 'trunk', _set(f['allowed_vlans']) == allowed, f['native_vlan'] == native[0]]
        else:
            desired = result['projected_state']
            if any(f[k] is None for k in ('pvid', 'tagged_vlans', 'untagged_vlans')):
                evidence['warnings'] = ['Missing operational fields: ' + ', '.join(k for k in ('pvid', 'tagged_vlans', 'untagged_vlans') if f[k] is None)]
                return dict(status='Unknown', evidence=evidence)
            checks += [f['pvid'] == desired['pvid'], _set(f['tagged_vlans']) == set(desired['tagged'])]
            # PVID untagged handling is checked independently from participation.
            if desired['pvid'] is not None:
                checks.append((desired['pvid'] in _set(f['untagged_vlans'])) == (desired['pvid'] not in set(desired['tagged'])))
            evidence['lookups'] = after['vlan_lookups']
            wanted = set(parse_vlan_list(p.get('tagged_vlans', '')))
            for vid in wanted | ({p['pvid']} if 'pvid' in p else set()):
                lookup = after['vlan_lookups'][str(vid)]
                ports = [v for v in lookup.get('ports', []) if key(v['interface']) == key(target)]
                if lookup['status'] != 'present' or not lookup.get('ports_complete') or len(ports) != 1:
                    return dict(status='Unknown', evidence=evidence)
                port = ports[0]
                remove = p['action'] == 'Remove' and vid in wanted
                checks.append(port['configured'].lower() == ('exclude' if remove else 'include'))
                if vid in wanted:
                    checks.append(port['tagging'].lower() == ('untagged' if remove else 'tagged'))
    return dict(status='Verified' if all(checks) else 'Failed', evidence=evidence)


def _template_stages(entries):
    """Preserve selection order; split dependencies and potentially overlapping targets.

    Full fresh collection is retained at stage boundaries. Distinct targets may
    share it; repeated targets require independently observed intermediate effects.
    """
    pending, targets, ids = [], set(), set()
    for entry in entries:
        r, result = entry['request'], entry['result']
        p = r['inputs']
        if BY_ID[r['template_id']].function == 'VLAN Database':
            touched = {('vlan', p['vlan_id'])}
        else:
            # Both source and resolved owner protect physical/bundle overlaps.
            touched = {('interface', canonical_interface_name(entry['identity']['platform'], name, normalize_aggregation=True))
                       for name in (p['interface'], result.get('resolved_target') or p['interface'])}
        if pending and (r['template_id'] != pending[0]['request']['template_id']
                        or targets & touched or ids & set(result['dependencies'])):
            yield pending
            pending, targets, ids = [], set(), set()
        pending.append(entry)
        targets.update(touched)
        ids.add(r['row_id'])
    if pending:
        yield pending


def _run(summary, credentials, transport_config, *, inventory_path, run_root, backups_dir, on_started=None, progress=None):
    with Progress(progress) as reporter:
        reporter.emit('-', 'confirmed_start', rows_total=len(summary['selected']),
                      devices_total=len({e['request']['inputs']['management_ip'] for e in summary['selected']}))
        return _run_with_progress(summary, credentials, transport_config, inventory_path=inventory_path,
                                  run_root=run_root, backups_dir=backups_dir, on_started=on_started, reporter=reporter)


def _run_with_progress(summary, credentials, transport_config, *, inventory_path, run_root, backups_dir, on_started, reporter):
    run_started = monotonic()
    settings = ApplySettings(**summary['settings'])
    groups = {}
    for entry in summary['selected']:
        groups.setdefault(entry['request']['inputs']['management_ip'], []).append(entry)
    secrets = [v for credential in credentials.values() for v in (credential.username, credential.password, credential.secret) if v]
    capture_secrets = {v for credential in credentials.values() for v in (credential.password, credential.secret) if v}
    pattern = re.compile('|'.join(re.escape(v) for v in sorted(set(secrets), key=len, reverse=True))) if secrets else None
    def clean(text):
        text = sanitize_text(text)
        return pattern.sub('[REDACTED]', text) if pattern else text
    spool = ResultSpool.create(run_root, 'configuration_execution', len(groups), clean=clean)
    reporter.correlation.update(run_id=spool.path.name, digest=summary['source_digest'])
    if on_started:
        on_started(spool.path)
    (spool.path / 'summary.json').write_text(canonical(summary), encoding='utf-8')
    journal = ResultJournal(spool.path / 'events.jsonl', clean=clean)
    backup = BackupWriter(Path(backups_dir) / spool.path.name if backups_dir else spool.path / 'backups', datetime.now(timezone.utc), exact_directory=True)
    from orbitflow.terminal_progress import location
    location('Execution run', clean(str(spool.path)))
    location('Protected backup folder (capture pending)', clean(str(backup.path)))
    store = JsonInventoryStore(inventory_path)
    with module_logger('configuration', 'execution') as (logger, log_path):
        locations = {e['request']['row_id']: f"{BY_ID[e['request']['template_id']].sheet} row {e['request'].get('row', '?')}" for e in summary['selected']}
        hostnames = {ip: clean((entries[0].get('identity') or {}).get('hostname', '')) for ip, entries in groups.items()}
        def event(ip, stage, *, _display=True, **values):
            item = dict(run_id=spool.path.name, digest=summary['source_digest'], management_ip=ip,
                        stage=stage, timestamp=_now(), log_reference=str(log_path), **values)
            journal.append(item)
            display = {k: clean(str(values[k])) for k in ('status', 'command', 'command_index', 'command_total', 'error_category') if k in values}
            if values.get('row_id'):
                display['source'] = locations.get(values['row_id'], 'retained row')
            timeout = settings.save_timeout if stage == 'save_intent' else settings.command_timeout if stage in {'command_intent', 'mode_intent'} else None
            if stage in {'preflight', 'row_preflight', 'backup', 'verification_intent'}:
                display['command_timeout_seconds'] = settings.command_timeout
            if _display:
                reporter.emit(ip, stage, hostname=hostnames.get(ip, ''), timeout=timeout, **display)
            logger.info(clean(f"run={spool.path.name}; digest={summary['source_digest']}; ip={ip}; stage={stage}; row={values.get('row_id', '-')}; command={values.get('command_index', '-')}; status={values.get('status', '-')}; category={values.get('error_category', '-') }; backup={values.get('backup_path', '-')}; warning={'; '.join(short_findings([values['warning']])) if values.get('warning') else '-'}"))
        @contextmanager
        def owned_cli(session, ip):
            with timed_context('shell_setup', lambda: DeviceCLI(session)) as cli:
                try:
                    with observations.observation_scope(session, cli):
                        yield cli
                finally:
                    event(ip, 'cleanup_start')

        def worker(ip):
            with profiling(logger, clean, run_id=spool.path.name, digest=summary['source_digest'], management_ip=ip,
                           device_family=groups[ip][0]['identity']['device_family']):
                with stage('device_total', failure='Unknown') as timing:
                    try:
                        value = device_worker(ip)
                        save = value.get('save', {})
                        if save.get('status') == 'Unknown' or any(r['status'] == 'Unknown' for r in value['rows']):
                            timing['outcome'] = 'Unknown'
                        elif value.get('error_category') or save.get('explicit_error') or any(r['status'] != 'Verified' for r in value['rows']):
                            timing['outcome'] = 'Failed'
                        else:
                            timing['outcome'] = 'Completed'
                        return value
                    finally:
                        event(ip, 'cleanup', status='Completed')

        def device_worker(ip):
            entries = groups[ip]
            attempted, rows, usable = False, [], True
            save = dict(status='Not required')
            event(ip, 'connect')
            try:
                require(ip in credentials, 'credentials unavailable')
                with timed_context('connection', lambda: connect_device(ip, credentials[ip], transport_config)) as session:
                    with owned_cli(session, ip) as cli:
                        event(ip, 'identify')
                        context = timed_call('identification', DeviceInventoryResolver(store, sanitize_fact=clean).resolve, session, management_ip=ip, cli=cli)
                        _gate(all(identity(context) == e['identity'] for e in entries), 'FRESH_IDENTITY_MISMATCH')
                        cli._reuse_session_setup = True
                        event(ip, 'preflight')
                        before = timed_call('initial_preflight', _observe, session, context, cli, entries, settings.command_timeout)
                        exec_prompt = cli.prompt
                        adapter = edge_execution if context.platform == 'ubiquiti_edgeswitch' else cisco_execution
                        require(adapter.is_exec_prompt(exec_prompt), 'privileged EXEC required')
                        event(ip, 'backup')
                        content = timed_call('protected_backup_collection', ConfigurationService(timeout=settings.command_timeout).collect, session, context, cli=cli)
                        for secret in capture_secrets:
                            content = content.replace(secret, '[REDACTED]')
                        path = timed_call('protected_backup_write', backup.write, context.hostname, context.platform, content)
                        try:
                            backup_reference = Path(os.path.relpath(path, spool.path)).as_posix()
                        except ValueError:
                            # Windows external overrides may be on another drive.
                            backup_reference = path.resolve().as_posix()
                        event(ip, 'backup_complete', backup_reference=backup_reference, backup_path=str(path.resolve()))
                        completed = set()
                        try:
                            for stage_index, stage_entries in enumerate(_template_stages(entries), 1):
                                stage_id = f'stage-{stage_index}'
                                template_id = stage_entries[0]['request']['template_id']
                                stage_rows = []
                                with correlation(template_id=template_id, stage_id=stage_id), stage('template_stage', failure='Unknown') as stage_timing:
                                    # All candidate prechecks and durable evidence precede any stage write.
                                    pre_error = None
                                    try:
                                        _gate(usable, 'SESSION_UNUSABLE')
                                        cli.require_session(session)
                                        _gate(adapter.is_exec_prompt(cli.prompt), 'SESSION_UNUSABLE')
                                        exec_prompt = cli.prompt
                                        if stage_index != 1:
                                            before = timed_call('stage_pre_observation', _observe, session, context, cli, stage_entries, settings.command_timeout)
                                        _gate(before.get('configuration_reliable') is True, 'STAGE_EVIDENCE_UNAVAILABLE')
                                        event(ip, 'stage_precheck', stage_id=stage_id, template_id=template_id,
                                              row_ids=[e['request']['row_id'] for e in stage_entries], observation=before)
                                    except Exception as exc:
                                        pre_error = exc
                                        observations.invalidate()
                                    for entry in stage_entries:
                                        r, approved = entry['request'], entry['result']
                                        rid = r['row_id']
                                        row = dict(row_id=rid, decision=approved['decision'], status='Not executed', stage_id=stage_id, _started=monotonic())
                                        rows.append(row)
                                        stage_rows.append((entry, row))
                                        try:
                                            if pre_error is not None:
                                                raise pre_error
                                            _gate(set(approved['dependencies']) <= completed, 'DEPENDENCY_FAILED')
                                            event(ip, 'row_preflight', row_id=rid)
                                            from orbitflow.configuration.v2.evidence import scoped_evidence
                                            row['pre_config'] = timed_call('scoped_evidence', scoped_evidence, r, before)
                                            event(ip, 'pre_row_evidence', row_id=rid, stage_id=stage_id, evidence=row['pre_config'])
                                            _check_target(entry, before, context)
                                            if BY_ID[r['template_id']].function == 'VLAN Port' and r['inputs']['action'] != 'Remove':
                                                row['missing_prerequisite_vlans'] = [str(v) for v in (set(parse_vlan_list(r['inputs'].get('tagged_vlans', ''))) | ({r['inputs']['pvid']} if 'pvid' in r['inputs'] else set())) if before['vlan_lookups'].get(str(v), {}).get('status') != 'present']
                                                _gate(not row['missing_prerequisite_vlans'], 'VLAN_PREREQUISITE_UNAVAILABLE')
                                            row['_ready'] = True
                                        except Exception as exc:
                                            row['error_category'] = type(exc).__name__
                                            if isinstance(exc, ExecutionBlocked):
                                                row['blocker'] = exc.code
                                    for entry, row in stage_rows:
                                        r, approved = entry['request'], entry['result']
                                        rid = r['row_id']
                                        if not row.pop('_ready', False):
                                            continue
                                        with correlation(row_id=rid), stage('row_apply', failure='Unknown') as row_timing:
                                            try:
                                                _gate(usable, 'SESSION_UNUSABLE')
                                                rejected = False
                                                for index, command in enumerate(approved['commands'], 1):
                                                    cli.execution_mode(exec_prompt, context.platform)
                                                    observations.invalidate()
                                                    event(ip, 'command_intent', row_id=rid, command_index=index, command_total=len(approved['commands']), command=command)
                                                    attempted = True
                                                    row['status'] = 'Unknown'
                                                    started = monotonic()
                                                    with correlation(command_index=str(index)), stage('approved_command', command_category=command_category(command), failure='Unknown') as command_timing:
                                                        try:
                                                            response = cli.run_command(command, timeout=settings.command_timeout)
                                                        except BaseException:
                                                            event(ip, 'command_unknown', row_id=rid, command_index=index,
                                                                  response=_safe_source(getattr(cli, 'last_response', ''), context.platform), status='Unknown')
                                                            raise
                                                        rejected = command_rejected(response)
                                                        event(ip, 'command_result', row_id=rid, command_index=index,
                                                              response=_safe_source(response, context.platform), status='Rejected' if rejected else 'Returned', elapsed_seconds=monotonic() - started)
                                                        command_timing['outcome'] = 'Rejected' if rejected else 'Returned'
                                                    if rejected:
                                                        row['status'] = 'Failed'
                                                        break
                                                if not adapter.is_exec_prompt(cli.prompt):
                                                    # Recovery stays on the synchronized owned session.
                                                    for index in range(2):
                                                        cli.execution_mode(exec_prompt, context.platform)
                                                        command = 'exit' if context.platform == 'ubiquiti_edgeswitch' else 'end'
                                                        event(ip, 'mode_intent', row_id=rid, command=command)
                                                        response = timed_call('mode_recovery_command', cli.run_command, command, timeout=settings.command_timeout)
                                                        event(ip, 'mode_result', row_id=rid, response=_safe_source(response, context.platform))
                                                        if command_rejected(response):
                                                            usable = False
                                                            require(False, 'EXEC recovery rejected')
                                                        if adapter.is_exec_prompt(cli.prompt):
                                                            break
                                                    require(adapter.is_exec_prompt(cli.prompt), 'EXEC recovery failed')
                                                # Rebind presentation only after an echo-associated completed
                                                # approved sequence/recovery reaches privileged EXEC.
                                                exec_prompt = cli.prompt
                                                row['_rejected'] = rejected
                                                row['_verify'] = True
                                            except Exception as exc:
                                                observations.invalidate()
                                                row['error_category'] = type(exc).__name__
                                                if isinstance(exc, ExecutionBlocked):
                                                    row['blocker'] = exc.code
                                                try:
                                                    cli.require_session(session)
                                                    usable = usable and cli.prompt == exec_prompt
                                                except Exception:
                                                    usable = False
                                            row_timing['outcome'] = row['status']
                                    # A single fresh full snapshot proves final effects on independent
                                    # targets. Overlaps/dependencies were split before any send.
                                    after, post_error = None, None
                                    candidates = [(e, row) for e, row in stage_rows if row.get('_verify')]
                                    try:
                                        if candidates:
                                            _gate(usable, 'SESSION_UNUSABLE')
                                            cli.require_session(session)
                                            _gate(adapter.is_exec_prompt(cli.prompt), 'SESSION_UNUSABLE')
                                            exec_prompt = cli.prompt
                                            if BY_ID[template_id].function == 'VLAN Database':
                                                after = dict(configuration_reliable=True, vlan_lookups={
                                                    str(e['request']['inputs']['vlan_id']): timed_call('vlan_lookup', VlanService(timeout=settings.command_timeout).lookup,
                                                        session, context, e['request']['inputs']['vlan_id'], cli=cli)
                                                    for e, _ in candidates})
                                            else:
                                                after = timed_call('stage_post_observation', _observe, session, context, cli, stage_entries, settings.command_timeout)
                                            _gate(after.get('configuration_reliable') is True, 'STAGE_EVIDENCE_UNAVAILABLE')
                                            event(ip, 'stage_post_observation', stage_id=stage_id, observation=after)
                                    except Exception as exc:
                                        observations.invalidate()
                                        post_error = exc
                                    for entry, row in stage_rows:
                                        rid = entry['request']['row_id']
                                        if row.pop('_verify', False):
                                            rejected = row.pop('_rejected', False)
                                            row['status'] = 'Failed' if rejected else 'Unknown'
                                            try:
                                                if post_error is not None:
                                                    raise post_error
                                                event(ip, 'verification_intent', row_id=rid, commands=entry['verification_commands'])
                                                started = monotonic()
                                                with correlation(row_id=rid):
                                                    verification = timed_call('post_row_verification', _verify_observed,
                                                        session, context, cli, entry, before, timeout=settings.command_timeout,
                                                        correlation={'run_id': spool.path.name, 'digest': summary['source_digest'], 'row_id': rid, 'stage_id': stage_id}, after=after)
                                                verification['stage_id'] = stage_id
                                                verification['observation_scope'] = 'stage final state; independent target'
                                                row.update(status='Failed' if rejected else verification['status'], verification=verification)
                                                row['verification_seconds'] = monotonic() - started
                                                for warning in verification.get('evidence', {}).get('warnings', []) if isinstance(verification.get('evidence'), dict) else []:
                                                    event(ip, 'verification_warning', row_id=rid, warning=warning)
                                                if row['status'] == 'Verified':
                                                    completed.add(rid)
                                            except Exception as exc:
                                                if not rejected:
                                                    row['status'] = 'Unknown'
                                                row['error_category'] = type(exc).__name__
                                        if row['status'] != 'Verified':
                                            observations.invalidate()
                                        row['elapsed_seconds'] = max(0.0, monotonic() - row.pop('_started'))
                                        logger.info('Execution timing', extra=dict(run_id=spool.path.name, digest=summary['source_digest'],
                                            management_ip=ip, device_family=context.device_family, template_id=template_id,
                                            row_id=rid, stage_id=stage_id, stage='row_total', outcome=row['status'], elapsed_seconds=row['elapsed_seconds']))
                                        event(ip, 'row_result', **row)
                                    event(ip, 'stage_result', stage_id=stage_id,
                                          rows=[dict(row_id=row['row_id'], status=row['status']) for _, row in stage_rows])
                                    stage_timing['outcome'] = ('Unknown' if any(row['status'] == 'Unknown' for _, row in stage_rows)
                                                              else 'Verified' if all(row['status'] == 'Verified' for _, row in stage_rows) else 'Failed')
                        finally:
                            try:
                                cli.require_session(session)
                                usable = usable and adapter.is_exec_prompt(cli.prompt)
                                if usable:
                                    exec_prompt = cli.prompt
                            except Exception:
                                usable = False
                            if attempted and usable:
                                observations.invalidate()
                                event(ip, 'save_intent', command='write memory')
                                save = dict(status='Unknown', error_category='SaveCapabilityIncomplete')
                                try:
                                    with reporter.save_scope(ip, clean(context.hostname), settings.save_timeout,
                                                             audit=lambda stage: event(ip, stage, _display=False)):
                                        save = timed_call('device_save', save_configuration, session, context, cli=cli, exec_prompt=exec_prompt, timeout=settings.save_timeout, correlation={'run_id': spool.path.name, 'digest': summary['source_digest']})
                                except BaseException:
                                    event(ip, 'save_result', status='Unknown', error_category='SaveCapabilityIncomplete',
                                          response=_safe_source(getattr(cli, 'last_response', ''), context.platform))
                                    raise
                            elif attempted:
                                save = dict(status='Unknown', error_category='SessionUnusable')
                            event(ip, 'save_result', **save)
                return dict(rows=rows, save=save)
            except Exception as exc:
                event(ip, 'device_failure', error_category=type(exc).__name__)
                done = {r['row_id'] for r in rows}
                for entry in entries:
                    rid = entry['request']['row_id']
                    if rid not in done:
                        row = dict(row_id=rid, status='Not executed', error_category=type(exc).__name__)
                        rows.append(row)
                        event(ip, 'row_result', **row)
                return dict(rows=rows, save=save, error_category=type(exc).__name__)
        def persist(outcome):
            ip = list(groups)[outcome.position - 1]
            value = outcome.value or {}
            failed = (outcome.error_category or value.get('error_category')
                      or value.get('save', {}).get('status') == 'Unknown'
                      or value.get('save', {}).get('explicit_error')
                      or any(r['status'] != 'Verified' for r in value.get('rows', [])))
            spool.append(outcome, target=ip, failed=bool(failed))
            event(ip, 'device_complete', status='Failed' if failed else 'Completed')
        event('-', 'confirm', status='Claimed')
        reporter.active.pop('-', None)
        run_outcome = 'Unknown'
        try:
            with spool.collection():
                execute_devices(list(groups), worker, config=ExecutionConfig(settings.max_concurrent_devices, settings.connection_start_interval, settings.connection_retry_delay), on_outcome=persist)
            event('-', 'recording', status='Completed')
            run_outcome = 'Completed'
        finally:
            logger.info('Execution timing', extra=dict(run_id=spool.path.name, digest=summary['source_digest'],
                        stage='run_total', outcome=run_outcome, elapsed_seconds=max(0.0, monotonic() - run_started)))
    return spool


def compact(value):
    """Compact scalar/list presentation; integer lists use inclusive ranges."""
    if isinstance(value, list) and value and all(type(v) is int for v in value):
        values = sorted(set(value))
        groups = []
        start = end = values[0]
        for v in values[1:]:
            if v == end + 1:
                end = v
            else:
                groups.append(str(start) if start == end else f'{start}-{end}')
                start = end = v
        groups.append(str(start) if start == end else f'{start}-{end}')
        return ','.join(groups)
    if isinstance(value, dict):
        return '; '.join(f'{k}={compact(v)}' for k, v in value.items())
    if isinstance(value, list):
        return ', '.join(compact(v) for v in value)
    return readable(value)


def short_findings(findings):
    """Execution presentation only: retain plan findings unchanged in bound data."""
    messages = []
    def project(value):
        if isinstance(value, list):
            for item in value:
                project(item)
        elif isinstance(value, dict):
            code = value.get('code') or value.get('category')
            if code == 'TAGGED_VLAN_NOT_IN_MEMBERSHIP':
                return
            if code:
                detail = {k: v for k, v in value.items() if k in {'message', 'vlan', 'vlan_id', 'vlans', 'interface', 'reason'}}
                messages.append(str(code) + (': ' + compact(detail) if detail else ''))
            else:
                messages.append('Existing VLAN finding (see retained detail)')
        else:
            text = str(value)
            category, sep, payload = text.partition(': ')
            if category == 'TAGGED_VLAN_NOT_IN_MEMBERSHIP':
                return
            if sep and payload[:1] in '{[':
                try:
                    before = len(messages)
                    project(decode(payload))
                    if len(messages) > before and category != 'EXISTING_VLAN_WARNING':
                        messages[before] = category + ': ' + messages[before]
                    return
                except (ValueError, TypeError):
                    text = category + ' (see retained detail)'
            messages.append(text)
    for finding in findings:
        project(finding)
    return list(dict.fromkeys(messages))


def operator_summary(prepared):
    """Project bound preparation for concise operator confirmation."""
    source = prepared['summary']
    devices = {}
    changes = no_ops = 0
    reasons = []
    categories = {}
    def reason(source, findings):
        for message in short_findings(findings):
            category, _, detail = message.partition(': ')
            group = categories.setdefault(category, dict(sources=[], details=[]))
            group['sources'].append(source)
            if detail and detail not in group['details']:
                group['details'].append(detail)
    for entry in source['selected']:
        r, result = entry['request'], entry['result']
        ip = r['inputs']['management_ip']
        t = BY_ID[r['template_id']]
        device = devices.setdefault(ip, dict(ip=ip, hostname=(entry.get('identity') or {}).get('hostname', ''), family=t.family, actions={}))
        action = t.function + ': ' + r['inputs'].get('action', 'Set')
        device['actions'][action] = device['actions'].get(action, 0) + 1
        no_ops += bool(result['no_op'])
        changes += not result['no_op']
        reason(f"{t.sheet} row {r.get('row', '?')} / {ip}", result.get('findings', []))
    for entry in source['skipped']:
        r = entry.get('request', {})
        location = f"{BY_ID[r['template_id']].sheet} row {r.get('row', '?')} / {r['inputs']['management_ip']}" if r else 'retained row'
        reason(location, [entry['reason'], *entry.get('blockers', [])])
    reasons = [dict(reason=category + (': ' + '; '.join(group['details']) if group['details'] else ''),
                    count=len(group['sources']), sources=list(dict.fromkeys(group['sources'])))
               for category, group in categories.items()]
    return dict(reference=prepared['reference'], source_digest=source['source_digest'], summary_digest=prepared['summary_digest'],
                device_count=len(devices), changes=changes, no_ops=no_ops, skipped=len(source['skipped']),
                concurrency=source['settings']['max_concurrent_devices'], devices=list(devices.values()), reasons=reasons,
                save_behavior=source['save_behavior'], fresh_checks='Offline selection only. Fresh identity, state, VLAN prerequisites and protected backup are checked after confirmation.')


def render_summary(summary):
    """Concise text for CLI or other operator surfaces, without nested evidence."""
    lines = [f"Devices: {summary['device_count']} | Changes: {summary['changes']} | No-op: {summary['no_ops']} | Skipped: {summary['skipped']} | Concurrency: {summary['concurrency']}"]
    for index, device in enumerate(summary['devices'], 1):
        lines.append(f"{index}. {device['ip']} {device['hostname']} | selected rows: {sum(device['actions'].values())}")
    if summary['reasons']:
        lines.append('Warnings/blockers or skipped selections: see preparation findings and configuration logs for details.')
    lines.append('Fresh checks + backup before writes; execute approved rows; save entire running config once when usable after attempted writes (including partial changes). EdgeSwitch may interrupt management; save answer: y.')
    if any('EVC: ' in action for device in summary['devices'] for action in device['actions']):
        lines.append('EVC rows remove/recreate the entire selected service instance, even when matching. Forwarding may be interrupted; omitted child settings are removed.')
    return '\n'.join(lines)


def readable(value):
    """Readable retained facts/evidence; no JSON encoding in operator cells."""
    if value is None:
        return 'Unknown'
    if isinstance(value, list):
        return '\n'.join(readable(v) for v in value)
    if isinstance(value, dict):
        if 'excerpt' in value:
            return f"{value.get('source_filename', '')}:{value.get('line', '')}: {value['excerpt']}"
        return '\n'.join(f"{k}: {readable(v)}" for k, v in value.items())
    return str(value)


def raw_config(value):
    if isinstance(value, list):
        return '\n'.join(str(v.get('excerpt', '')) if isinstance(v, dict) else str(v) for v in value)
    return readable(value)


def verification_text(verification, result, request):
    if not verification:
        return 'Not performed'
    evidence = verification.get('evidence', {})
    lines = [verification.get('status', 'Unknown')]
    if not isinstance(evidence, dict) or 'verification_scope' not in evidence:
        lines.append('Expected: ' + compact(result.get('projected_state') or request.get('inputs', {})))
    if isinstance(evidence, dict):
        for name in ('verification_scope', 'service_detail', 'access_vlan_check', 'access_vlan_source', 'preservation', 'missing_fields', 'invalid_fields'):
            if name in evidence:
                lines.append(name + ': ' + compact(evidence[name]))
        facts = evidence.get('switchport', {}).get('facts', {})
        if facts:
            lines.append('Actual: ' + compact(facts))
        if evidence.get('description'):
            lines.append('Description: ' + compact(evidence['description'].get('port_description')))
        target = result.get('resolved_target', request.get('inputs', {}).get('interface'))
        for vid, lookup in evidence.get('lookups', {}).items():
            for port in lookup.get('ports', []):
                if port.get('interface') == target:
                    lines.append(f'VLAN {vid}: ' + compact(port))
    lookup = verification.get('lookup')
    if lookup:
        lines.append('Actual: ' + compact({k: lookup.get(k) for k in ('status', 'vlan_id', 'name', 'operational_status')}))
    return '\n'.join(lines)


def export_results(run_path, output):
    """Export retained execution evidence offline, with optional CLI presentation.

    No collection state is changed and Unknown outcomes are never reclassified.
    Library calls retain output ownership; the returned path is unchanged.
    """
    from orbitflow.terminal_progress import workflow_progress, current_display, location
    with workflow_progress('Execution report stages', quiet=current_display.get() is None) as display:
        try:
            result = _export_results(run_path, output)
        except BaseException:
            display.finish(False)
            location('Export failed/interrupted; retained execution run', run_path)
            raise
        display.finish(True)
        location('Execution report saved', result)
        location('Retained execution run (historical outcomes unchanged)', run_path)
        return result


def _export_results(run_path, output):
    """Offline literal report grouped by original template, from retained journal only."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    require(not Path(output).exists(), 'output already exists')
    run_path = Path(run_path)
    source = decode((run_path / 'summary.json').read_text(encoding='utf-8'))
    manifest_path = run_path / 'manifest.json'
    run_id = decode(manifest_path.read_text(encoding='utf-8')).get('run_id', run_path.name) if manifest_path.exists() else run_path.name
    book = Workbook(write_only=True)
    summary = book.create_sheet('Run Summary')
    details = book.create_sheet('Evidence Details')
    details.column_dimensions['A'].width = 28
    details.column_dimensions['B'].width = 28
    details.column_dimensions['C'].width = 8
    details.column_dimensions['D'].width = 110
    details.append(['Row / Event', 'Field', 'Part', 'Complete readable text'])
    sheet_rows = {details.title: 1}
    def literal_row(sheet, values, reference=''):
        cells = []
        for column, value in enumerate(values, 1):
            text = ILLEGAL_CHARACTERS_RE.sub('', sanitize_text(str(value)))
            if sheet is details and column == 4 and len(text) > 2400:
                for offset in range(0, len(text), 2400):
                    literal_row(details, [values[0], values[1], offset // 2400 + 1, text[offset:offset + 2400]])
                return
            if len(text) > 30000:
                for offset in range(0, len(text), 30000):
                    literal_row(details, [reference, column, offset // 30000 + 1, text[offset:offset + 30000]])
                text = text[:29000] + f"\nContinued completely in Evidence Details: {reference}, field {column}"
            cell = WriteOnlyCell(sheet, value=text)
            cell.data_type = 's'
            if text.startswith('Evidence Details!A'):
                cell.hyperlink = "#'Evidence Details'!" + text.split('!')[1]
                cell.style = 'Hyperlink'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cells.append(cell)
        sheet_rows[sheet.title] = sheet_rows.get(sheet.title, 0) + 1
        sheet.row_dimensions[sheet_rows[sheet.title]].height = (300 if sheet is details else 90 if reference else 30)
        sheet.append(cells)
    literal_row(summary, ['Run', run_id])
    literal_row(summary, ['Plan digest', source['source_digest']])
    literal_row(summary, ['Save policy', source.get('save_behavior', '')])
    saves, backups, row_offsets, detail_refs = {}, {}, {}, {}
    row_counts, device_counts = {}, {}
    for index, event in enumerate(ResultJournal(run_path / 'events.jsonl').records(), 1):
        stage, ip, rid = event['stage'], event.get('management_ip', ''), event.get('row_id')
        detail_refs[index] = sheet_rows.get(details.title, 1) + 1
        literal_row(details, [f'Event {index}', stage, 1, readable(event)], f'Event {index}')
        if stage in {'row_result', 'save_result'}:
            counts = row_counts if stage == 'row_result' else device_counts
            status = event.get('status', 'Unknown')
            counts[status] = counts.get(status, 0) + 1
        if stage == 'save_result':
            saves[ip] = {k: event[k] for k in ('status', 'explicit_error', 'error_category') if k in event}
        elif stage == 'backup_complete':
            backups[ip] = event.get('backup_reference', '')
        elif rid:
            # Only journal positions are retained across the run; evidence is read per row.
            row_offsets.setdefault(rid, []).append(index)
    for ip in dict.fromkeys([e['request']['inputs']['management_ip'] for e in source['selected']]):
        literal_row(summary, [ip, 'Save: ' + readable(saves.get(ip, {'status': 'Not recorded'})), 'Backup: ' + backups.get(ip, 'Not recorded')])
    literal_row(summary, ['Selected / skipped rows', f"{len(source['selected'])} / {len(source['skipped'])}"])
    literal_row(summary, ['Rows by verification outcome', compact(row_counts)])
    literal_row(summary, ['Devices by save outcome', compact(device_counts)])
    sheets = {}
    headers = ['Run ID', 'IP', 'Row ID', 'Source row', 'Hostname', 'Family', 'Original Inputs', 'Decision', 'No-op', 'Pre-config', 'Config Sent', 'Config Evidence (post-row)', 'Apply Status', 'Status Verification', 'Error', 'Warning', 'Save Status', 'Backup Reference', 'Log Reference']
    entries = [(entry, False) for entry in source['selected']] + [(entry, True) for entry in source['skipped']]
    entries.sort(key=lambda pair: (pair[0].get('request', {}).get('sheet_order', 999), pair[0].get('request', {}).get('row', 999)))
    for entry, skipped in entries:
        r, result = entry.get('request', {}), entry.get('result', {})
        t = BY_ID.get(r.get('template_id'))
        title = t.sheet if t else 'Legacy Skipped'
        if title not in sheets:
            sheet = book.create_sheet(title)
            sheet.freeze_panes = 'D2'
            from openpyxl.utils import get_column_letter
            input_columns = list(t.columns) if t else []
            template_headers = headers[:9] + input_columns + headers[9:] + ['Full detail']
            for col, name in enumerate(template_headers, 1):
                sheet.column_dimensions[get_column_letter(col)].width = 48 if name in {'Pre-config', 'Config Sent', 'Config Evidence (post-row)', 'Status Verification', 'Error', 'Warning'} else 24
            sheet.column_dimensions['G'].hidden = bool(t)
            sheet.column_dimensions['C'].hidden = True
            literal_row(sheet, template_headers)
            sheets[title] = sheet
        rid = entry.get('row_id', r.get('row_id'))
        ip = r.get('inputs', {}).get('management_ip', '')
        event, pre_config, sent = {}, 'Not captured', []
        positions = set(row_offsets.get(rid, []))
        for index, item in enumerate(ResultJournal(run_path / 'events.jsonl').records(), 1):
            if index not in positions:
                continue
            stage = item['stage']
            if stage == 'pre_row_evidence':
                pre_config = item['evidence']
            elif stage == 'command_intent':
                sent.append(dict(command=item['command'], status='Pending / uncertain'))
                event = dict(item, status='Unknown')
            elif stage in {'command_result', 'command_unknown'} and sent:
                sent[-1].update(status=item.get('status', 'Unknown'), response=item.get('response', ''))
            elif stage == 'row_result':
                event = item
                pre_config = item.get('pre_config', pre_config)
        verification = event.get('verification', {})
        evidence = verification.get('evidence', verification.get('lookup', 'Not captured'))
        errors = {k: event[k] for k in ('error_category', 'blocker', 'missing_prerequisite_vlans') if event.get(k)}
        if event.get('blocker'):
            errors['Preflight findings'] = event.get('preflight_findings', [])
        if verification.get('error_category'):
            errors['Verification error'] = verification['error_category']
        if skipped:
            errors.update(reason=entry['reason'], details=entry.get('blockers', []))
        for command in sent:
            if command['status'] == 'Rejected':
                errors['Rejected command'] = command['command'] + ' (see response detail)'
            elif command['status'] in {'Unknown', 'Pending / uncertain'}:
                errors['Uncertain command'] = command['command'] + ' (outcome unknown; see detail)'
        if verification.get('status') == 'Failed':
            errors['Verification'] = 'Failed; compare expected/actual checks and linked detail'
        warnings = list(evidence.get('warnings', [])) if isinstance(evidence, dict) else []
        warnings.extend(verification.get('warnings', []))
        if not event.get('blocker'):
            warnings.extend(event.get('preflight_findings', []))
        warnings = short_findings(warnings)
        errors = {k: short_findings(v) if k in {'Preflight findings', 'details'} else v for k, v in errors.items()}
        apply_status = ('Skipped' if skipped else 'Unknown' if any(c['status'] in {'Unknown', 'Pending / uncertain'} for c in sent) else 'Rejected' if any(c['status'] == 'Rejected' for c in sent) else 'Applied' if sent else 'No-op' if result.get('no_op') and verification else 'Not executed')
        values = [run_id, ip, rid, r.get('row', ''), (entry.get('identity') or {}).get('hostname', ''), t.family if t else '',
            readable(r.get('source_inputs') or r.get('inputs', {})), entry.get('decision', result.get('decision', '')), result.get('no_op', ''),
            raw_config(pre_config), '\n'.join('command: ' + c['command'] + ' [' + c['status'] + ']' for c in sent) if sent else 'None (no-op)' if result.get('no_op') else 'None attempted', raw_config(evidence.get('configuration', evidence.get('evidence', []))) if isinstance(evidence, dict) else readable(evidence),
            apply_status, verification_text(verification, result, r), compact(errors), '\n'.join(warnings),
            saves.get(ip, {}).get('status', 'Not recorded'), backups.get(ip, ''), event.get('log_reference', '')]
        if not positions:
            detail_start = sheet_rows.get(details.title, 1) + 1
            literal_row(details, [rid, 'Retained source / review', 1, readable(entry)], rid)
        else:
            detail_start = detail_refs[min(positions)]
        input_values = [compact((r.get('source_inputs') or r.get('inputs', {})).get(k, '')) for k in (t.columns if t else [])]
        literal_row(sheets[title], values[:9] + input_values + values[9:] + [f"Evidence Details!A{detail_start}"], rid)
    with Path(output).open('xb') as stream:
        book.save(stream)
    return Path(output)
