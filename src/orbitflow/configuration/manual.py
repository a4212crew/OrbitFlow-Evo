"""Excel manual blocks: fresh shells, durable evidence and success-only row save.

Public API: create_template, load_rows, summary, execute, regenerate_report.
No legacy plan conversion, automatic replay, rollback or independent verification.
"""
from copy import deepcopy
from contextlib import contextmanager
from dataclasses import dataclass, asdict
from threading import Event, current_thread, main_thread
import signal
from datetime import datetime, timezone
from pathlib import Path
import hashlib
import ipaddress
import math
import os
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from orbitflow.config import ExecutionConfig
from orbitflow.configuration.manual_auth import authorize
from orbitflow.configuration.plans import canonical, safe_text
from orbitflow.configuration_backup import BackupWriter, filename_component
from orbitflow.execution import execute_devices
from orbitflow.result_spool import ResultSpool, ResultJournal
from orbitflow.progress import Progress
from orbitflow.logging import sanitize_text, module_logger
from orbitflow.transport import connect_device
from orbitflow.inventory import DeviceInventoryResolver, JsonInventoryStore
from orbitflow.capabilities.configuration import ConfigurationService
from orbitflow.capabilities.save import command_rejected, save_configuration
from orbitflow.vendors.common import DeviceCLI, normalize_output
from orbitflow.vendors.configuration_facts import _safe_source
from orbitflow.vendors import manual_configuration as vendor

INPUT_ERROR = 'Invalid input: IP, command, credential or persistence restriction'
SCOPE = 'Transcript only — configuration not independently verified'
# Fail closed on abbreviations, execution wrappers and programmable/interactive
# constructs that could bypass per-row persistence or require automatic answers.
FORBIDDEN = ('write', 'copy', 'save', 'commit', 'reload', 'reboot', 'erase',
             'delete', 'format', 'execute', 'do', 'run', 'alias', 'macro', 'event',
             'kron', 'tclsh', 'python', 'shell', 'bash', 'cli', 'banner', 'more',
             'terminal', 'screen-length', 'enable', 'disable', 'quit', 'logout')
CONFIG_ROOTS = frozenset(('interface description vlan switchport shutdown no default hostname ip ipv6 router network spanning-tree service encapsulation bridge-domain rewrite split-horizon mtu speed duplex negotiation exit end name address exit-address-family access-list route-map match set class-map policy-map bandwidth priority logging snmp-server ntp clock line transport location contact').split())
INTERACTION = re.compile(r'(?i)(?:--\s*more\s*--|\[more\]|password\s*:|username\s*:|\(y/n\)|\[y/n\]|\[yes/no\]|\[confirm\]|press (?:any|enter|return)|continue\?|are you sure|destination filename|overwrite\?|\[(?:yes|no|y|n)(?:/[^\]]+)?\]|(?:enter|input|confirm|proceed)[^\r\n]{0,120}:|\?\s*$)')


@contextmanager
def interrupt_guard(cancelled):
    """Notify active workers before the scheduler waits for owned cleanup."""
    previous = None
    if current_thread() is main_thread():
        previous = signal.getsignal(signal.SIGINT)
        def interrupted(signum, frame):
            cancelled.set()
            if callable(previous):
                previous(signum, frame)
            else:
                raise KeyboardInterrupt()
        signal.signal(signal.SIGINT, interrupted)
    try:
        yield
    except BaseException:
        cancelled.set()
        raise
    finally:
        if previous is not None:
            signal.signal(signal.SIGINT, previous)


def now():
    return datetime.now(timezone.utc).isoformat()


@dataclass(frozen=True)
class CaptureSettings:
    first_response: float = 10.0
    inactivity: float = 1.0
    maximum: float = 30.0
    setup_timeout: float = 30.0
    save_timeout: float = 60.0

    def __post_init__(self):
        for value in asdict(self).values():
            if type(value) not in (int, float) or not math.isfinite(value) or not 0 < value <= 300:
                raise ValueError('Capture deadlines must be finite seconds in (0, 300]')
        if max(self.first_response, self.inactivity) >= self.maximum:
            raise ValueError('Capture overall deadline must exceed first response and inactivity')


def commands(text):
    if not isinstance(text, str) or not text or len(text) > 32767:
        raise ValueError('Invalid multiline command cell')
    lines = text.split('\n')
    if len(lines) > 1000:
        raise ValueError('Too many commands')
    for line in lines:
        # Validate without modifying literal whitespace or authored command text.
        safe_text(line.strip())
        if any(ord(c) < 32 or ord(c) == 127 for c in line):
            raise ValueError('Control characters forbidden')
        tokens = line.strip().lower().split()
        head = tokens[0]
        if (any(word.startswith(head) for word in FORBIDDEN)
                or head in {'configure', 'conf'}
                or any(c in line for c in ';|?\\')
                or any(token in {'startup-config', 'nvram:', 'flash:', 'bootflash:'} for token in tokens)):
            raise ValueError('Persistence, interaction or execution wrapper forbidden')
        if head not in CONFIG_ROOTS:
            raise ValueError('Ambiguous command or alias is not supported')
        if head in {'no', 'default'} and (len(tokens) < 2 or tokens[1] not in CONFIG_ROOTS):
            raise ValueError('Ambiguous negated/default command')
        # Credential commands rejected through shared protections, including
        # vendor labels beyond the v1 string policy.
        if any(_safe_source(line, platform) != line for platform in vendor.PROFILES):
            raise ValueError('Credential-bearing command forbidden')
    return lines


def create_template(directory='data/input/generated_template'):
    """Allocate an exclusive daily template; command cells accept Alt+Enter."""
    root = Path(directory)
    root.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = 'Manual Configuration'
    sheet.append(['management_ip', 'command'])
    sheet.column_dimensions['A'].width = 24
    sheet.column_dimensions['B'].width = 90
    sheet['B2'].alignment = Alignment(wrap_text=True, vertical='top')
    from openpyxl.comments import Comment
    sheet['B1'].comment = Comment('Literal newline-separated commands (Alt+Enter). Starts in global configuration mode. No credentials, save/commit, interactive commands or execution wrappers.', 'OrbitFlow')
    for sequence in range(1, 100000):
        path = root / f'manual_config_{datetime.now():%Y%m%d}_{sequence:03d}.xlsx'
        try:
            stream = path.open('xb')
        except FileExistsError:
            continue
        try:
            with stream:
                book.save(stream)
        except BaseException:
            path.unlink(missing_ok=True)
            raise
        return path
    raise ValueError('Template allocation exhausted')


def load_rows(path):
    """Retain every nonempty input row and fixed validation errors, never secrets."""
    book = load_workbook(path, read_only=True, data_only=False)
    rows = []
    try:
        sheet = book.active
        iterator = sheet.iter_rows()
        headers = next(iterator, ())
        if [cell.value for cell in headers] != ['management_ip', 'command']:
            raise ValueError('Manual workbook requires exactly management_ip and command columns')
        for number, cells in enumerate(iterator, 2):
            values = [cell.value for cell in cells]
            if all(value is None for value in values):
                continue
            item = dict(source_row=number, management_ip='', commands=[], error='')
            try:
                item['management_ip'] = str(ipaddress.ip_address(str(values[0]).strip()))
                if any(cell.data_type == 'f' for cell in cells):
                    raise ValueError()
                item['commands'] = commands(values[1])
            except Exception:
                item['error'] = INPUT_ERROR
                item['commands'] = []
            rows.append(item)
            if len(rows) > 10000:
                raise ValueError('Manual workbook exceeds 10000 rows')
    finally:
        book.close()
    if not rows:
        raise ValueError('Manual workbook contains no rows')
    return rows


def summary(rows, execution_config=ExecutionConfig(1), settings=CaptureSettings()):
    """Digest binds confirmation to exact ordered validated rows and settings."""
    if not 1 <= execution_config.max_concurrent_devices <= 5:
        raise ValueError('Manual concurrency must be 1–5')
    body = dict(rows=rows, execution=asdict(execution_config), capture=asdict(settings))
    return dict(digest=hashlib.sha256(canonical(body).encode()).hexdigest(),
                rows=len(rows), devices=len({r['management_ip'] for r in rows if r['management_ip']}),
                commands=sum(len(r['commands']) for r in rows),
                invalid=sum(bool(r['error']) for r in rows), concurrency=execution_config.max_concurrent_devices)


def execute(rows, *, profile, password, confirmation_digest, credentials, transport_config,
            execution_config=ExecutionConfig(1), settings=CaptureSettings(),
            inventory_path='data/inventory/inventory.json',
            run_root='outputs/runs/manual_configuration',
            backups_dir='outputs/backups/configuration', progress=None, on_started=None):
    """Apply only after local password and summary-bound explicit confirmation.

    credentials maps IP to existing DeviceCredentials; no command-workbook secrets.
    Returns a retained run path. Report regeneration is a separate offline operation.
    """
    authorize(profile, password)  # Before allocation, connection, identity or backup.
    rows = deepcopy(rows)
    # API callers receive the same validation as XLSX input; do not trust supplied
    # row dictionaries, commands or digest as execution authorization by themselves.
    if not 1 <= len(rows) <= 10000:
        raise ValueError('Manual input requires 1–10000 rows')
    for row in rows:
        if row.get('error') not in {'', INPUT_ERROR}:
            raise ValueError('Invalid manual input diagnostic')
        if row.get('error') and row.get('commands') != []:
            raise ValueError('Invalid rows cannot retain command cells')
        if row.get('management_ip'):
            ipaddress.ip_address(row['management_ip'])
        if set(row) != {'source_row', 'management_ip', 'commands', 'error'}:
            raise ValueError('Invalid manual row schema')
        if type(row['source_row']) is not int or row['source_row'] < 2:
            raise ValueError('Invalid source row')
        if not row['error']:
            ipaddress.ip_address(row['management_ip'])
            if commands('\n'.join(row['commands'])) != row['commands']:
                raise ValueError('Invalid commands')
    if len({r['source_row'] for r in rows}) != len(rows):
        raise ValueError('Duplicate source row identity')
    rows = sorted(rows, key=lambda r: r['source_row'])
    bound = summary(rows, execution_config, settings)
    if confirmation_digest != bound['digest']:
        raise ValueError('Explicit manual summary confirmation required')
    for row in rows:
        for command in row['commands']:
            if any(value and value in command for credential in credentials.values()
                   for value in (credential.password, credential.secret)):
                raise ValueError('Credential-bearing command forbidden')
    secret_values = sorted({v for c in credentials.values() for v in (c.username, c.password, c.secret) if v}, key=len, reverse=True)
    def clean(text):
        text = normalize_output(str(text))
        for value in secret_values:
            text = text.replace(value, '[REDACTED]')
        # Whole lines retained; shared credential-label policy also covers
        # unanticipated secrets. Terminal escape sequences are removed first.
        lines = []
        for line in sanitize_text(text).split('\n'):
            for platform in vendor.PROFILES:
                line = _safe_source(line, platform)
            lines.append(line)
        return '\n'.join(lines)
    def safe_record(value):
        # Field names are fixed application schema, never device input. Redact
        # values before the shared writer so short credentials cannot corrupt
        # keys such as source_row/status and make offline evidence unreadable.
        if isinstance(value, dict):
            return {key: item if key in {'kind', 'status', 'save', 'error', 'verification_scope'}
                    else safe_record(item) for key, item in value.items()}
        if isinstance(value, (list, tuple)):
            return [safe_record(item) for item in value]
        return clean(value) if isinstance(value, str) else value
    groups = {}
    for row in rows:
        groups.setdefault(row['management_ip'], []).append(row)
    spool = ResultSpool.create(run_root, 'manual_configuration', len(groups), clean=sanitize_text)
    journal = ResultJournal(spool.path / 'events.jsonl', clean=sanitize_text)
    backup = BackupWriter(Path(backups_dir) / spool.path.name, datetime.now(timezone.utc), exact_directory=True)
    store = JsonInventoryStore(inventory_path)
    # Recovery needs source identity/counts, never a replayable command workbook.
    input_rows = [dict(source_row=r['source_row'], management_ip=r['management_ip'],
                       total=len(r['commands']), error=r['error']) for r in rows]
    journal.append(safe_record(dict(kind='input', rows=input_rows, summary=bound, settings=asdict(settings))))
    if on_started:
        on_started(spool.path)
    cancelled = Event()
    def checkpoint():
        if cancelled.is_set():
            raise KeyboardInterrupt()
    with interrupt_guard(cancelled), Progress(progress) as reporter, module_logger('manual_configuration', 'execution') as (logger, log_path):
        def event(ip, kind, **values):
            journal.append(safe_record(dict(timestamp=now(), management_ip=ip, kind=kind, log_reference=str(log_path), **values)))
            reporter.emit(ip, kind, **{k: v for k, v in values.items() if k in {'source_row', 'command_index', 'command_total', 'status'}})
            logger.info('Manual configuration stage', extra=dict(management_ip=ip, stage=kind, run_id=spool.path.name))

        def worker(ip):
            stopped = False
            session = None
            transcript = None
            path = None
            try:
                for source in groups[ip]:
                    row = dict(source_row=source['source_row'], management_ip=ip,
                               status='Incomplete', reason='', attempted=0, total=len(source['commands']),
                               last_command=None, save='Not initiated', verification_scope=SCOPE,
                               transcript=str(path.relative_to(spool.path)) if path else '')
                    phase = 'setup'
                    save_sent = False
                    pending = ''
                    try:
                        # Allocate the IP fallback before identification so even
                        # login/setup failures retain a device transcript.
                        if transcript is None:
                            path = spool.path / (filename_component(ip or 'invalid-input') + '.txt')
                            transcript = path.open('x', encoding='utf-8', newline='\r\n')
                        row['transcript'] = path.name
                        def record(label, text):
                            transcript.write(f'{now()} row={source["source_row"]} {label}\n{clean(text)}\n')
                            transcript.flush()
                            os.fsync(transcript.fileno())
                        record('ROW', 'Fresh shell required; ' + SCOPE)
                        checkpoint()
                        if stopped:
                            row['reason'] = 'Not attempted: earlier device row failed or unknown'
                        elif source['error']:
                            row['reason'] = source['error']
                            stopped = True
                        else:
                            checkpoint()
                            if session is None:
                                if ip not in credentials:
                                    raise ValueError('Credentials unavailable')
                                event(ip, 'connect', source_row=source['source_row'])
                                session = connect_device(ip, credentials[ip], transport_config)
                            # This context owns a NEW shell every row. No literal
                            # exit is used for cleanup, even from unknown submodes.
                            with DeviceCLI(session, timeout=settings.setup_timeout) as cli:
                                context = DeviceInventoryResolver(store, sanitize_fact=clean).resolve(session, management_ip=ip, cli=cli)
                                vendor.adapter(context)  # Before any user configuration writes.
                                phase = 'backup'
                                content = ConfigurationService(timeout=settings.setup_timeout).collect(session, context, cli=cli)
                                for value in secret_values:
                                    content = content.replace(value, '[REDACTED]')
                                saved_backup = backup.write(context.hostname or ip, context.platform, content)
                                event(ip, 'backup', source_row=source['source_row'], backup=str(saved_backup))
                                phase = 'setup'
                                checkpoint()
                                vendor.setup(cli, context, settings.setup_timeout)
                                pending = ''
                                recent = ''
                                received = 0
                                def receive(text):
                                    nonlocal pending, recent, received
                                    received += len(text.encode('utf-8'))
                                    if received > 8 * 1024 * 1024:
                                        raise ValueError('Capture byte limit exceeded')
                                    pending += text
                                    recent = (recent + text)[-65536:]
                                    while '\n' in pending:
                                        line, pending = pending.split('\n', 1)
                                        record('RECEIVE', line)
                                    if len(pending) > 65536:
                                        raise ValueError('Capture line limit exceeded')
                                    if command_rejected(normalize_output(recent)) or INTERACTION.search(normalize_output(recent)):
                                        if pending:
                                            record('RECEIVE', pending)
                                            pending = ''
                                    if command_rejected(normalize_output(recent)):
                                        raise ValueError('Vendor rejected command')
                                    if INTERACTION.search(normalize_output(recent)):
                                        raise ValueError('Unsupported interactive exchange')
                                    checkpoint()
                                phase = 'command capture'
                                for index, command in enumerate(source['commands'], 1):
                                    checkpoint()
                                    row.update(attempted=index, last_command=index)
                                    event(ip, 'command_intent', source_row=source['source_row'], command_index=index, command_total=row['total'])
                                    record('SEND', command)  # Durable intent before transmission.
                                    cli.capture_literal(command, first_response=settings.first_response,
                                                        inactivity=settings.inactivity, maximum=settings.maximum, receive=receive)
                                if pending:
                                    record('RECEIVE', pending)
                                    pending = ''
                                checkpoint()
                                phase = 'finalization'
                                anchor = vendor.finalize(cli, context, settings.setup_timeout)
                                # Mode response may contain delayed block output.
                                record('FINALIZATION RECEIVE', cli.last_response)
                                if command_rejected(cli.last_response) or INTERACTION.search(cli.last_response):
                                    raise ValueError('Finalization rejected or interactive')
                                record('BLOCK CAPTURED', 'No detected error; inactivity is not proof of completion')
                                event(ip, 'save_intent', source_row=source['source_row'])
                                checkpoint()
                                phase = 'save'
                                row['save'] = 'Unknown'
                                save_sent = True  # Conservatively Unknown even if sendall fails partially.
                                result = save_configuration(session, context, cli=cli, exec_prompt=anchor, timeout=settings.save_timeout, manual=True)
                                record('SAVE RECEIVE', result['response'])
                                checkpoint()
                                row['save'] = 'Failed' if result.get('explicit_error') else result['status']
                                if result['status'] != 'Saved' or result.get('explicit_error'):
                                    raise ValueError('Save not confirmed or explicitly rejected')
                                phase = 'recording'
                                row.update(status='Completed', reason='')
                    except BaseException as exc:
                        if pending and transcript is not None:
                            try:
                                record('RECEIVE (interrupted)', pending)
                            except BaseException:
                                phase = 'recording'
                            pending = ''
                        stopped = True
                        row['status'] = 'Incomplete'
                        # Never include exception text or raw receive data.
                        detail = str(exc) if str(exc) in {
                            'Credentials unavailable', 'Vendor rejected command', 'Unsupported interactive exchange',
                            'Capture byte limit exceeded', 'Capture line limit exceeded',
                            'Manual entry/save unsupported for identified platform', 'Privileged EXEC required',
                            'Global configuration mode not established', 'Save mode synchronization failed',
                            'Save not confirmed or explicitly rejected', 'Finalization rejected or interactive',
                            'Literal command overall deadline', 'Literal command first response deadline',
                            'Literal command disconnected'} else type(exc).__name__
                        row['reason'] = f'{phase}: {detail}'
                        if save_sent and row['save'] not in {'Saved', 'Failed'}:
                            row['save'] = 'Unknown'
                        if not isinstance(exc, Exception):
                            row['reason'] = f'{phase}: Interrupted'
                    try:
                        event(ip, 'row_result', **{k: v for k, v in row.items() if k != 'management_ip'})
                    except BaseException:
                        # A recording failure must never initiate save, continue
                        # rows or report a successful row. Recovery uses intent.
                        raise
            finally:
                if transcript is not None:
                    transcript.close()
                if session is not None:
                    session.close()
            return dict(management_ip=ip, failed=stopped)

        def persist(outcome):
            try:
                spool.append(outcome, payload=safe_record(outcome.value), target=clean(list(groups)[outcome.position - 1]),
                             failed=bool(outcome.error_category or (outcome.value or {}).get('failed')))
            except BaseException:
                cancelled.set()
                raise
        with spool.collection():
            execute_devices(list(groups), worker, config=execution_config, on_outcome=persist)
    return spool.path


def regenerate_report(run_path, output=None):
    """Offline recovery from durable input/intents/results; never connects/replays.

    Unterminated journal lines are ignored. A save intent without confirmed result
    remains Unknown. Missing rows remain Incomplete, even if collection returned.
    """
    run_path = Path(run_path)
    journal = ResultJournal(run_path / 'events.jsonl')
    # Metadata aggregation is bounded by the 10000-row input limit;
    # command/response data stay in per-device transcripts on disk.
    inputs, results, attempts, saves = [], {}, {}, set()
    for event in journal.records():
        if event['kind'] == 'input':
            inputs = event['rows']
        elif event['kind'] == 'row_result':
            results[event['source_row']] = event
        elif event['kind'] == 'command_intent':
            attempts[event['source_row']] = event['command_index']
        elif event['kind'] == 'save_intent':
            saves.add(event['source_row'])
    if not inputs:
        raise ValueError('Manual run has no durable input snapshot')
    if output is None:
        root = Path('outputs/reports/manual_configuration')
        root.mkdir(parents=True, exist_ok=True)
        output = root / (run_path.name + '-' + datetime.now().strftime('%Y%m%dT%H%M%S%f') + '.xlsx')
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    # Claim output before constructing streaming XML writers; allocation failure
    # cannot leave unclosed openpyxl generators.
    stream = output.open('xb')
    book = None
    try:
        with stream:
            book = Workbook(write_only=True)
            sheet = book.create_sheet('Manual Configuration')
            fields = ['source_row', 'management_ip', 'status', 'reason', 'attempted', 'total', 'last_command', 'save', 'verification_scope', 'transcript']
            sheet.append(fields)
            from openpyxl.cell import WriteOnlyCell
            for source in inputs:
                number = source['source_row']
                result = results.get(number) or dict(source_row=number, management_ip=source['management_ip'], status='Incomplete',
                            reason='Interrupted/recording incomplete' if number in attempts or number in saves else 'Not attempted: run incomplete',
                            attempted=attempts.get(number, 0), total=source['total'], last_command=attempts.get(number),
                            save='Unknown' if number in saves else 'Not initiated', verification_scope=SCOPE, transcript='')
                cells = []
                for name in fields:
                    cell = WriteOnlyCell(sheet, str(result.get(name, '') if result.get(name) is not None else ''))
                    cell.data_type = 's'  # Formula injection protection for all persisted text.
                    cells.append(cell)
                sheet.append(cells)
            sheet.close()
            book.save(stream)
    except BaseException:
        if book is not None:
            for sheet in book.worksheets:
                try:
                    if not sheet.closed:
                        sheet.close()
                    if sheet._writer is not None and Path(sheet._writer.out).exists():
                        sheet._writer.cleanup()
                except Exception:
                    pass  # Preserve the primary export failure and retained run.
        output.unlink(missing_ok=True)
        raise
    return output
