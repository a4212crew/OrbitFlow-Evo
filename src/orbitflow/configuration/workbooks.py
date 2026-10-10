"""Literal Excel input/output surfaces for offline configuration jobs."""

from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from .plans import PlanError, require


def read(path):
    try:
        book = load_workbook(path, data_only=False, read_only=True, keep_links=False)
    except (BadZipFile, InvalidFileException):
        raise PlanError("invalid workbook file") from None
    tables = {}
    try:
        for sheet in book:
            require(sheet.max_row <= 20000 and sheet.max_column <= 40, "workbook size limit exceeded")
            rows = list(sheet.iter_rows())
            if not rows:
                tables[sheet.title] = []
                continue
            headers = [c.value for c in rows[0]]
            require(all(type(h) is str and h for h in headers) and len(set(headers)) == len(headers),
                    "invalid or duplicate workbook headers")
            records = []
            for index, cells in enumerate(rows[1:], 2):
                if not any(c.value is not None for c in cells):
                    continue
                record = {h: (c.value if c.value is not None else "") for h, c in zip(headers, cells)}
                record["_line"] = index
                record["_formula"] = any(c.data_type in ("f", "e") for c in cells)
                records.append(record)
            tables[sheet.title] = records
    finally:
        book.close()
    return tables


def write(path, sheets):
    """All strings are literal, including Excel formula prefixes. Never truncate."""
    book = Workbook()
    book.remove(book.active)
    for name, (headers, rows) in sheets.items():
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            require(len(row) == len(headers), "invalid output row")
            sheet.append(row)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    require(len(cell.value) <= 32767, "Excel cell limit exceeded")
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            width = min(65, max(18, max(len(str(c.value or "").split("\n")[0]) for c in col) + 2))
            sheet.column_dimensions[col[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            lines = max(max(len(str(c.value or "")) // 55 + 1, str(c.value or "").count("\n") + 1) for c in row)
            sheet.row_dimensions[row[0].row].height = min(240, max(30, lines * 15))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("xb") as stream:
        book.save(stream)
    book.close()
