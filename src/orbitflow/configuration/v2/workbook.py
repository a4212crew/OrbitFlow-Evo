"""Thin XLSX adapter. Only the decision cells may change during decision consumption."""
from pathlib import Path
import hashlib
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

from orbitflow.configuration.plans import require, PlanError, canonical
from orbitflow.vendors.configuration_facts import _safe_source
from .catalogue import CATALOGUE, BY_ID, BY_SHEET, normalize_display_inputs
from .planning import request, evaluate_decisions, MAX_ROWS
from .evidence import scoped_evidence

FULL_EXTRA = ("Row ID", "Existing Evidence", "Proposed CLI", "Live Validation", "Consistency",
         "Dependencies", "Observed At", "Findings", "No-op", "Validation Status", "Error Details", "Resolved Target", "Warnings", "Blockers", "Ordered Overlap Details", "Projected State", "Service Instance Differences", "Engineer Decision")

OMITTED = frozenset({"Row ID", "Observed At", "Findings", "Resolved Target",
                     "Ordered Overlap Details", "Projected State", "Service Instance Differences"})
EXTRA = tuple(name for name in FULL_EXTRA if name not in OMITTED)


def output_columns(template):
    return EXTRA if template in CATALOGUE[:16] else FULL_EXTRA


class WorkbookRequests(list):
    """Adapter metadata retains sanitized invalid rows alongside valid requests."""
    def __init__(self, layout):
        super().__init__()
        self.layout = layout
        self.errors = []
        self.source_digest = ""


def literal(cell, value):
    text = str(value) if value is not None else ""
    if len(text) > 32767:
        # Excel/openpyxl would otherwise truncate silently. Never shorten the
        # immutable plan: replace only its XLSX presentation, deterministically.
        digest = cell.parent.parent["Introduction"]["B1"].value
        require(digest, "oversized Excel value requires an immutable plan")
        text = (f"Complete value in immutable plan.json (plan digest {digest}). "
                f"Field: {cell.parent.title}!{cell.coordinate}; "
                f"characters: {len(text)}; value SHA-256: "
                f"{hashlib.sha256(text.encode('utf-8')).hexdigest()}. "
                "Use the sheet/row and field header to locate the original input/result; "
                "Evidence Details references identify the request and source line in observations. "
                "This cell is a reference, not the complete value.")
    cell.value = text
    cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _base():
    book = Workbook()
    intro = book.active
    intro.title = "Introduction"
    intro.append(["OrbitFlow bulk configuration planning v2"])
    intro.append(["Order per device: worksheets left-to-right, then rows top-to-bottom. Review never authorizes device execution."])
    intro.append(["Predefined templates collect live read-only evidence. Manual CLI never connects; impact remains unknown."])
    intro.append(["Fill rows from row 3; use header comments and sheet A1 comments for field rules. Never enter credentials."])
    intro.append(["Review the generated plan and set Pending/Approve/Reject; decisions bind to immutable plan.json."])
    intro.append(["Reference: docs/operations/configuration-template-reference.md. Templates 17–22 are provisional; hardware validation remains pending."])
    intro.append(["Platform", "Template / function", "Purpose", "Open template"])
    for column, width in {"A": 22, "B": 28, "C": 68, "D": 32}.items():
        intro.column_dimensions[column].width = width
    for row in intro.iter_rows(min_row=1, max_row=6, max_col=1):
        row[0].alignment = Alignment(vertical="top", wrap_text=True)
        if row[0].row != 1:  # B1 is reserved for the immutable output plan digest.
            intro.merge_cells(start_row=row[0].row, start_column=1, end_row=row[0].row, end_column=4)
        intro.row_dimensions[row[0].row].height = 30
    for cell in intro[7]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    purposes = {
        "VLAN Database": "Create or rename a VLAN in the device database.",
        "Access Port": "Assign an interface to one access VLAN.",
        "Trunk Port": "Add, remove or replace trunk VLAN membership.",
        "Interface Description": "Update an existing interface description.",
        "EVC": "Create or recreate an Ethernet service instance.",
        "VLAN Port": "Change port VLAN membership, tagging or PVID.",
        "L2VPN": "Plan a Layer 2 subinterface and bridge-domain attachment.",
        "Subinterface VSI": "Plan a VLAN subinterface attached to an existing VSI.",
        "Manual CLI": "Supply literal offline commands; impact is unknown.",
    }
    for t in CATALOGUE:
        intro.append([t.family, t.function, purposes[t.function], t.sheet])
        link = intro.cell(intro.max_row, 4)
        link.hyperlink = f"#'{t.sheet}'!A1"
        link.font = Font(color="0000FF", underline="single")
        sheet = book.create_sheet(t.sheet)
        sheet.append(["Back to Introduction"])
        sheet["A1"].hyperlink = "#'Introduction'!A1"
        sheet.append(t.columns)
        sheet.freeze_panes = "B3"
        for index, col in enumerate(t.columns, 1):
            cell = sheet.cell(2, index)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.comment = Comment(("Required" if col in ("management_ip", *t.required) else "Optional") +
                                   ("; comma-separated VLAN IDs or ranges, e.g. 445,545,2400-2444" if col == "vlans" else ""), "OrbitFlow")
            if t.function == 'EVC' and col in {'outer_vlan', 'inner_vlan', 'push_vlan'}:
                cell.comment.text = ('Conditional: ' + {'outer_vlan': 'required for Dot1q/QinQ; forbidden for Untagged; inclusive ID/list/range 1-4094',
                    'inner_vlan': 'required for QinQ only; inclusive ID/list/range 1-4094 or any',
                    'push_vlan': 'required for either Push only; single integer 1-4094'}[col])
            if t.function == 'VLAN Port' and col == 'tagged_vlans':
                cell.comment.text = 'Conditional: required for Remove; Add/Change requires tagged_vlans and/or pvid; IDs 2-4093'
            examples = {"management_ip": "192.0.2.10", "vlan_id": "445", "name": "Customer",
                        "description": "Customer Site Alpha", "service_instance_id": "10 (1-4000), independent of VLAN and BD",
                        "bridge_domain": "445; independent of ingress tag", "outer_vlan": "445",
                        "inner_vlan": "545 (QinQ only)", "pop_tags": "Legacy Pop1/Pop2; use rewrite instead",
                        "rewrite": "None omits child rewrite; Pop 1 symmetric; Pop 2 symmetric; Push dot1q; Push dot1q symmetric. Complete child recreation.",
                        "push_vlan": "Required only for Push",
                        "replace_untagged": "No preserves unrelated untagged membership; Yes explicitly replaces it",
                        "pvid": "Independent PVID; PVID-only uses Add/Change and blank tagged_vlans; may also be tagged; Remove+pvid invalid",
                        "tagged_vlans": "445,545,2400-2444; Add/Change includes then tags; Remove untags then excludes; no Replace",
                        "split_horizon_group": "Optional; explicit 0 valid; blank omits setting on recreation",
                        "interface_description": "Optional parent description; blank preserves parent text",
                        "subinterface": "10", "vsi": "existing-VSI", "bridge_group": "CUSTOMER",
                        "scope": "Affected resource metadata only; no CLI wrapper. Literal commands must include required hierarchy", "commands": "Literal multiline CLI; no credentials",
                        "interface": "Existing physical/aggregate interface; descriptions also permit existing subinterfaces"}
            if col in examples:
                cell.comment.text += "; " + examples[col]
            sheet.column_dimensions[get_column_letter(index)].width = 25 if col != "commands" else 65
            choices = {"action": "Add,Remove,Replace", "mode": "Access,Tagged" if t.family == "EdgeSwitch" else "configuration,interface,service",
                       "encapsulation": "dot1q,QinQ,untagged" if t.function in {"EVC", "L2VPN"} else "dot1q termination,vlan-type dot1q",
                       "rewrite": "None,Pop1,Pop2,Push", "replace_untagged": "No,Yes",
                       "platform": "cisco_ios,cisco_xe,cisco_xr,huawei_vrp,ubiquiti_edgeswitch"}
            if t.function == 'EVC':
                choices.update(encapsulation='Untagged,Dot1q,QinQ', rewrite='None,Pop 1 symmetric,Pop 2 symmetric,Push dot1q,Push dot1q symmetric')
            if t.function == 'VLAN Port':
                choices['action'] = 'Add/Change,Remove'
            if col in choices:
                dv = DataValidation(type="list", formula1='"' + choices[col] + '"')
                dv.errorTitle, dv.error, dv.showErrorMessage = "Invalid choice", "Select a listed value", True
                sheet.add_data_validation(dv)
                dv.add(f"{get_column_letter(index)}3:{get_column_letter(index)}5002")
        help_text = template_help(t)
        if help_text:
            sheet['A1'].comment = Comment(help_text, 'OrbitFlow', width=800, height=500)
    return book


def template_help(t):
    if t.family not in {'C3750X', 'C3850X', 'ME3600X', 'ASR920', 'EdgeSwitch'}:
        return ('Catalogue/provisional, not owner-approved or live-validated. Required: ' +
                ', '.join(('management_ip', *t.required)) + '. Optional: ' +
                (', '.join(t.optional) or 'none') +
                '. See docs/operations/configuration-template-reference.md for current input grammar, examples, CLI and limitations. ' +
                ('Manual CLI is literal/offline only; no device connection, wrappers or semantic validation. Impact unknown.'
                 if t.function == 'Manual CLI' else 'Existing shared read-only facts are required; unavailable/ambiguous evidence blocks.'))
    required = 'Required: ' + ', '.join(('management_ip', *t.required)) + '. '
    if t.function == 'EVC':
        return required + (
            'Complete desired child: always no service instance <ID>, then service instance <ID> ethernet and full recreation. '
            'Valid encapsulation mappings: Untagged / blank / blank -> encapsulation untagged; '
            'Dot1q / 445 / blank -> encapsulation dot1q 445; Dot1q / 100-200 / blank -> encapsulation dot1q 100-200; '
            'Dot1q / 2-444,446-544 / blank -> encapsulation dot1q 2-444,446-544; '
            'QinQ / 300 / 200 -> encapsulation dot1q 300 second-dot1q 200; '
            'QinQ / 100 / 1-4094 -> encapsulation dot1q 100 second-dot1q 1-4094. any inner normalizes to 1-4094. '
            'Valid rewrites: Untagged+None or either Push; single Dot1q+Pop 1 or either Push; '
            'Dot1q range/list+None or either Push; QinQ fixed inner+Pop 1/Pop 2/Push symmetric; '
            'QinQ ranged inner+None/Pop 1/Push symmetric. Push needs independent push_vlan. '
            'CLI: None omits rewrite; rewrite ingress tag pop 1 symmetric; rewrite ingress tag pop 2 symmetric; '
            'rewrite ingress tag push dot1q 100; rewrite ingress tag push dot1q 1402 symmetric; bridge-domain 2000 split-horizon group 0. '
            'Unverified (blocked): single Dot1q+None; fixed QinQ+None; Untagged+Pop. '
            'Lab rejected (blocked): single Dot1q+Pop 2; ranged Dot1q+Pop; '
            'Invalid: Untagged with tags; '
            'QinQ+non-symmetric Push; ranged QinQ+Pop 2; missing push_vlan; reversed range 200-100. '
            'Blank child description/rewrite/group omits old setting; blank interface_description preserves parent. '
            'ME3600X needs trunk + allowed vlan none; missing global BD VLAN warns and is never auto-created; ASR920 uses inline BD. Empty compatible parent may initialize. '
            'No implicit no shutdown, MTU or policy changes. Live hardware/firmware validation pending.')
    if t.function == 'VLAN Port':
        return required + (
            'Add/Change requires tagged_vlans and/or pvid. Valid: tagged_vlans=445,545-546; pvid-only=445; tagged_vlans=445 and pvid=445. '
            'CLI: vlan participation include 445,545-546; vlan tagging 445,545-546; vlan pvid 445. '
            'Remove requires tagged_vlans and blank pvid: no vlan tagging 545; vlan participation exclude 545. '
            'Invalid: Replace, Remove+pvid, VLAN 1/4094, 200-100, empty Add/Change. Explicit ordered Remove then Add/Change for replacement. '
            'Global VLAN absence is informational. Preserve unrelated configuration; Current/Configured/Tagging remain independent.')
    if t.function == 'VLAN Database':
        return required + ('Valid: 445 / Customer; create/rename/no-op. Invalid: multiple VLAN IDs, reversed range, names longer than 32. '
                           'EdgeSwitch IDs 2-4093, names may contain spaces: vlan database; vlan 901; vlan name 901 "VLAN 901"; exit. '
                           'Cisco IDs 1-4094 excluding 1002-1005, ASCII names without spaces: configure terminal; vlan 445; name Customer; end. '
                           'Operational lookup preferred; complete reliable configuration fallback annotated; unknown blocks.')
    if t.function == 'Interface Description':
        return required + ('Valid: explicit physical or supported aggregate target, description=Customer Site Alpha. '
                           'CLI: interface <target>; description Customer Site Alpha (EdgeSwitch quoted); end/exit. '
                           'Invalid: missing/ambiguous interface, secret-bearing text, unreliable evidence. Exact target; physical member is not redirected. '
                           'Preserve VLAN/EVC/BD, MTU, admin state and all unrelated commands; matching description is No-op.')
    if t.function == 'Access Port':
        return required + ('Optional description; full canonical interface, VLAN 1-4094 excluding 1002-1005. '
            'Valid: GigabitEthernet1/0/1, vlan_id=445 -> switchport; switchport mode access; switchport access vlan 445. '
            'Exact explicit match is No-op. Invalid: Gi1/0/1, VLAN 1002, empty VLAN. Missing global VLAN warns; '
            'configured trunk/routed conversion and physical Port-channel members block. No implicit no shutdown. '
            'See docs/operations/configuration-template-reference.md.')
    if t.function == 'Trunk Port':
        return required + ('Optional native_vlan/description; blank preserves. Add/Remove/Replace only; '
            'vlans=445,545-546 expands inclusively; no ALL/NONE input or reserved 1002-1005. '
            'CLI: Add -> switchport trunk allowed vlan add 445,545,546; Remove -> switchport trunk allowed vlan remove 445; '
            'Replace -> switchport trunk allowed vlan 445,545,546. Add requires explicit baseline; Remove supports observed default ALL. '
            'Invalid: Add without explicit baseline, reversed range, abbreviated interface. Full match No-op. '
            'Trunks resolve evidenced Port-channel owners; access/routed conversion blocks. '
            'C3750X new trunk adds switchport trunk encapsulation dot1q; C3850X never does. '
            'Compatible Add then Add / Replace then Remove warn and project; other overlapping pairs block. '
            'See docs/operations/configuration-template-reference.md.')
    return required + 'See docs/operations/configuration-template-reference.md; unknown fields and credentials are invalid.'


def create_input(path):
    require(not Path(path).exists(), "output already exists")
    with Path(path).open("xb") as stream:
        _base().save(stream)


def read_input(path):
    book = load_workbook(path, read_only=True, data_only=False)
    result = WorkbookRequests(book.sheetnames)
    with Path(path).open("rb") as stream:
        result.source_digest = hashlib.file_digest(stream, "sha256").hexdigest()
    try:
        require(set(book.sheetnames) == {"Introduction", *BY_SHEET}, "unexpected workbook sheets")
        for order, sheet in enumerate(book):
            if sheet.title == "Introduction":
                continue
            t = BY_SHEET[sheet.title]
            require(sheet.max_row <= 5002, "worksheet row limit exceeded")
            rows = sheet.iter_rows()
            next(rows)
            header = tuple(c.value for c in next(rows))
            require(header == t.columns, "input columns changed")
            for number, cells in enumerate(rows, 3):
                values = [c.value for c in cells]
                if all(v in (None, "") for v in values):
                    continue
                try:
                    require(not any(c.data_type == "f" for c in cells), "formulas forbidden in input")
                    source_inputs = dict(zip(t.columns, values))
                    result.append(request(t.id, normalize_display_inputs(t, source_inputs), source_inputs=source_inputs,
                                          sheet_order=order, row=number))
                except (PlanError, ValueError, TypeError) as error:
                    source_inputs = dict(zip(t.columns, values))
                    original = {k: _safe_source(str(v)) if v is not None else "" for k, v in zip(t.columns, values)}
                    details = []
                    # Reuse the exact schema for field-local diagnostics, retaining
                    # cross-field validation as a row error when needed.
                    defaults = dict(management_ip="192.0.2.1", interface="GigabitEthernet1/0/1", vlan_id=445,
                                    name="Customer", action="Replace", vlans="445", description="Customer")
                    if t.function == 'EVC':
                        defaults.update(service_instance_id=10, encapsulation='Dot1q', bridge_domain=445,
                                        outer_vlan='445', rewrite='Pop 1 symmetric')
                    elif t.function == 'VLAN Port':
                        defaults.update(interface='0/1', action='Add/Change', tagged_vlans='445')
                    for field in t.columns:
                        try:
                            if cells[t.columns.index(field)].data_type == "f":
                                raise PlanError("formulas forbidden")
                            if t.family in {"C3750X", "C3850X", "ME3600X", "ASR920", "EdgeSwitch"}:
                                probe = {k: defaults[k] for k in t.required}
                                probe['management_ip'] = defaults['management_ip']
                                if t.function == 'EVC':
                                    probe.update(outer_vlan='445', rewrite='Pop 1 symmetric')
                                elif t.function == 'VLAN Port':
                                    probe['tagged_vlans'] = '445'
                                probe[field] = source_inputs.get(field)
                                normalize_display_inputs(t, probe)
                        except (PlanError, ValueError, TypeError) as exc:
                            detail = str(exc) if isinstance(exc, PlanError) else "invalid value"
                            details.append(field + ": " + detail)
                    result.errors.append(dict(sheet=t.sheet, row=number, reason="INVALID_INPUT_ROW",
                                              source_inputs=original, details=details or ["row: " + (str(error) if isinstance(error, PlanError) else "invalid input combination")]))
                require(len(result) + len(result.errors) <= MAX_ROWS, "workbook row limit exceeded")
        return result
    finally:
        book.close()


def export_plan(plan, path):
    require(not Path(path).exists(), "output already exists")
    book = _base()
    data = plan.to_dict()
    book["Introduction"]["B1"] = plan.digest
    results = {r["row_id"]: r for r in data["results"]}
    names = data["layout"]
    book._sheets = [book[s] for s in names]
    for t in CATALOGUE:
        sheet = book[t.sheet]
        for i, name in enumerate(output_columns(t), len(t.columns) + 1):
            literal(sheet.cell(2, i), name)
            sheet.column_dimensions[get_column_letter(i)].width = 40
        dv = DataValidation(type="list", formula1='"Pending,Approve,Reject"')
        sheet.add_data_validation(dv)
        col = get_column_letter(len(t.columns) + len(output_columns(t)))
        dv.add(f"{col}3:{col}5002")
    details = book.create_sheet("Evidence Details")
    details.append(["Reference", "Source", "Line", "Exact sanitized excerpt"])
    summary = book.create_sheet("Summary")
    summary.append(["Sheet", "Row", "Result"])
    for error in data["input_errors"]:
        summary.append([error["sheet"], error["row"], error["reason"]])
        t = BY_SHEET[error['sheet']]
        sheet = book[t.sheet]
        for index, field in enumerate(t.columns, 1):
            literal(sheet.cell(error['row'], index), error.get('source_inputs', {}).get(field, ''))
        for field, value in (("Validation Status", "Invalid"), ("Error Details", "; ".join(error.get('details', [error['reason']]))), ("Engineer Decision", "Pending")):
            literal(sheet.cell(error['row'], len(t.columns) + output_columns(t).index(field) + 1), value)
    for r in data["requests"]:
        t, p, result = BY_ID[r["template_id"]], r["inputs"], results[r["row_id"]]
        sheet = book[t.sheet]
        summary.append([t.sheet, r["row"], "No-op" if result["no_op"] else result["decision"]])
        obs = data["observations"].get(p["management_ip"], {}) if t.family != "All" else {}
        excerpts = scoped_evidence(r, obs)
        for e in excerpts:
            row = details.max_row + 1
            for col, v in enumerate((r["row_id"], e["source_filename"], e["line"], e["excerpt"]), 1):
                literal(details.cell(row, col), v)
        evidence = "\n".join(e["excerpt"] for e in excerpts)
        if len(evidence) > 30000:
            evidence = "Evidence Details: " + r["row_id"]
        commands = result["commands"]
        cli = commands if isinstance(commands, str) else "\n".join(commands)
        original = r["source_inputs"]
        values = [(",".join(map(str, original[k])) if isinstance(original.get(k), list) else original.get(k, "")) for k in t.columns]
        extras = [r["row_id"], evidence, cli,
                   result["live_validation"], result["consistency"], ",".join(result["dependencies"]), obs.get("observed_at", ""),
                   ",".join(result["findings"]), str(result["no_op"]), "Valid", "; ".join(f"conflicts with {q['sheet']} row {q['row']}" for q in result.get("conflicting_rows", [])) + ("; " + result["prerequisite_details"] if result.get("prerequisite_details") else ""),
                   result.get("resolved_target", p.get("interface", "")), ",".join(result.get("warnings", [])), ",".join(result.get("blockers", [])), canonical(result.get("overlaps", [])), canonical(result.get("projected_state")), canonical(result.get("service_instance_changes")), result["decision"]]
        fields = dict(zip(FULL_EXTRA, extras))
        if result.get('ingress_conflicts'):
            fields['Error Details'] = '; '.join(filter(None, [fields['Error Details']] +
                [item['explanation'] for item in result['ingress_conflicts']]))
        values += [fields[name] for name in output_columns(t)]
        for i, v in enumerate(values, 1):
            literal(sheet.cell(r["row"], i), v)
        sheet.row_dimensions[r["row"]].height = 90
    for c in ("A", "B", "C", "D"):
        details.column_dimensions[c].width = 70 if c == "D" else 28
    with Path(path).open("xb") as stream:
        book.save(stream)


def consume_decisions(original, path, *, expected_digest, now=None):
    """Future preflight adapter: caller pins a trusted plan digest; no authority.

    Validates protected values, then evaluates untrusted decision intent in memory.
    Never persists, re-signs or modifies the source plan.
    """
    require(original.digest == expected_digest, "source plan digest mismatch")
    # Regenerate the exact expected workbook in memory; compare every populated
    # cell and physical sheet order except ordinary engineer decision cells.
    from tempfile import TemporaryDirectory
    book = load_workbook(path, data_only=False)
    decisions = {}
    with TemporaryDirectory() as directory:
        expected_path = Path(directory) / "expected.xlsx"
        export_plan(original, expected_path)
        expected = load_workbook(expected_path, data_only=False)
        try:
            require(book.sheetnames == expected.sheetnames, "review worksheet order changed")
            decision_cells = {}
            for r in original.to_dict()["requests"]:
                t = BY_ID[r["template_id"]]
                decision_cells[(t.sheet, r["row"], len(t.columns) + len(output_columns(t)))] = r["row_id"]
            for sheet in expected:
                actual = book[sheet.title]
                require(actual.max_row == sheet.max_row and actual.max_column == sheet.max_column, "review dimensions changed")
                for row in sheet:
                    for cell in row:
                        other = actual.cell(cell.row, cell.column)
                        require(other.data_type != "f", "review formulas forbidden")
                        identity = decision_cells.get((sheet.title, cell.row, cell.column))
                        if identity:
                            decisions[identity] = other.value
                        else:
                            require(other.value == cell.value, "review content changed; re-plan required")
        finally:
            expected.close()
            book.close()
    return evaluate_decisions(original, decisions, expected_digest=expected_digest, now=now)
