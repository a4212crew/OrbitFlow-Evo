"""Interactive operator adapter. Run with python -m orbitflow.menu."""
import argparse
import math
from ipaddress import ip_address
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from time import monotonic
from uuid import uuid4

from orbitflow.environment import PROFILE, PATHS, ROUTING, load_profile, save_profile, validate_session, validate_paths
from orbitflow.config import ExecutionConfig
from orbitflow.progress import Progress


class Cancelled(Exception):
    """Operator returned without starting a workflow."""


class OperatorError(ValueError):
    """Fixed, actionable console diagnostic; never contains underlying output."""


@contextmanager
def diagnostic(message):
    try:
        yield
    except (Cancelled, PermissionError, FileNotFoundError):
        raise
    except Exception:
        raise OperatorError(message) from None


class Console:
    """Injectable terminal navigation over existing single-purpose services."""

    def __init__(self, profile=PROFILE, *, read=input, write=print, session_check=validate_session):
        self.profile = Path(profile)
        self.read, self.write, self.session_check = read, write, session_check

    def notify(self, text):
        """Flush real terminal output; injected writers remain usable in tests."""
        if self.write is print:
            print(text, flush=True)
        else:
            self.write(text)

    @contextmanager
    def activity(self, label, total=None, *, execution=False):
        """Render presentation-only activity; always clear it before prompts."""
        from orbitflow.terminal_progress import WorkflowProgress, ExecutionProgress, current_display
        renderer = ExecutionProgress if execution else WorkflowProgress
        display = renderer(label, write=self.write, total=total)
        token = current_display.set(display) if not execution else None
        try:
            with self._activity_timer(display):
                yield display
        finally:
            if token is not None:
                current_display.reset(token)

    @contextmanager
    def _activity_timer(self, display):
        with Progress(lambda event: display.tick()) as timer:
            timer.emit('', 'preflight')
            display.start()
            try:
                yield display
            except BaseException:
                display.finish(False)
                raise
            else:
                display.finish(True)
            finally:
                from orbitflow.terminal_progress import location
                for path in sorted(getattr(display, 'resources', ())):
                    location('Log', path)

    def observation_progress(self, ips, display):
        """Render stages only; ResultSpool checkpoints own positional counts.

        Targets may repeat an IP. Neither identity nor worker status text is
        evidence of a committed result; the active display receives that from
        the shared persistence owner.
        """
        def subscriber(message):
            tail = message.partition(': ')[2]
            from orbitflow.terminal_progress import worker_stage
            stage = tail.split(';', 1)[0]
            stage = {'interface collection': 'interface', 'VLAN/config collection': 'vlan',
                     'audit/snapshot': 'evaluate'}.get(stage, stage)
            worker_stage(stage)
        return subscriber

    def choose_plan(self, directory):
        """List original plan.json artifacts, newest first; validate again on use."""
        from orbitflow.configuration.v2.planning import load
        root = Path(directory)
        candidates = ([root / 'plan.json'] + list(root.glob('*/plan.json'))) if root.is_dir() else []
        files = []
        for path in candidates:
            if not path.is_file():
                continue
            try:
                load(path)
            except Exception:
                self.write('Skipped invalid plan candidate; inspect original artifacts locally.')
                continue
            files.append(path)
        files.sort(key=lambda p: (-p.stat().st_mtime_ns, str(p)))
        for index, path in enumerate(files, 1):
            self.write(f'{index}. {path.parent} / {path.name}')
        if not files:
            self.write('No valid immutable plans found.')
        self.write(f'{len(files)+1}. Enter explicit plan JSON path')
        while True:
            value = self.ask('Immutable plan number (b cancels): ')
            if value == str(len(files)+1):
                return self.choose_file('Immutable plan JSON', '')
            if value.isdecimal() and 1 <= int(value) <= len(files):
                return files[int(value)-1]
            # Explicit paths retain the earlier console interaction contract.
            if Path(value).is_file() or Path(value).suffix.lower() == '.json':
                if not Path(value).is_file():
                    raise FileNotFoundError('Missing input file')
                return Path(value)
            self.write('Choose a listed plan number or the explicit path option.')

    def ask(self, prompt, default=None):
        value = self.read(prompt).strip()
        if value.casefold() in ('b', 'back', 'cancel'):
            raise Cancelled()
        return value or default or ''

    def concurrency(self):
        while True:
            value = self.ask('Concurrent devices [1-5, default 1]: ', '1')
            if value in ('1', '2', '3', '4', '5'):
                return int(value)
            self.write('Enter a whole integer from 1 to 5.')

    def setup(self):
        try:
            old = load_profile(self.profile)
        except (OSError, ValueError):
            from orbitflow.config import load_execution_config
            # Carry forward legacy pacing, never its stored concurrency.
            legacy = load_execution_config() if not self.profile.exists() else ExecutionConfig()
            old = {'teleport': {}, 'execution': dict(connection_start_interval=legacy.connection_start_interval,
                    connection_retry_delay=legacy.connection_retry_delay), 'paths': PATHS.copy()}
        self.write('Non-secret environment settings only; b cancels without saving.')
        routing = {key: self.ask(f'{key}: ', old['teleport'].get(key)) for key in ROUTING}
        for key in ('teleport_key_path', 'teleport_cert_path'):
            value = self.ask(f'{key} (optional file path): ', old['teleport'].get(key))
            if value:
                routing[key] = value
        execution = {}
        for key, default in (('connection_start_interval', 1.0), ('connection_retry_delay', 5.0)):
            while True:
                value = self.ask(f'{key} [{old["execution"].get(key, default)}]: ',
                                 str(old['execution'].get(key, default)))
                try:
                    number = float(value)
                    if not math.isfinite(number) or number < 0:
                        raise ValueError()
                except ValueError:
                    self.write('Enter finite nonnegative seconds, for example 1.0; b cancels.')
                    continue
                execution[key] = number
                break
        paths = {key: self.ask(f'{key} [{old["paths"][key]}]: ', old['paths'][key]) for key in PATHS}
        with diagnostic('Invalid environment profile. Check routing fields and destination paths; nothing was saved.'):
            save_profile(dict(teleport=routing, execution=execution, paths=paths,
                              **({'manual_configuration_auth': old['manual_configuration_auth']}
                                 if 'manual_configuration_auth' in old else {})), self.profile)
        self.write(f'Environment saved: {self.profile}')

    def manual_password(self):
        from getpass import getpass
        from orbitflow.configuration.manual_auth import authorize, hash_password
        settings = self.settings()
        if 'manual_configuration_auth' in settings:
            authorize(settings, getpass('Current manual authorization password: '))
        password = getpass('New manual authorization password (12–1024 characters): ')
        if password != getpass('Repeat new password: '):
            raise OperatorError('Passwords differ; nothing saved')
        settings['manual_configuration_auth'] = {'password_hash': hash_password(password)}
        save_profile(settings, self.profile)
        self.write('Manual authorization password updated.')

    def environment_menu(self):
        self.submenu('Environment Setting', [('Edit environment profile', self.setup),
                     ('Set/change manual authorization password', self.manual_password)])

    def manual_template(self):
        from orbitflow.configuration.manual import create_template
        self.write(str(create_template()))

    def manual_execute(self):
        from getpass import getpass
        from orbitflow.configuration.manual import load_rows, summary, execute
        from orbitflow.configuration.manual_auth import authorize
        from orbitflow.targets import load_device_credentials
        from orbitflow.transport import TransportConfig
        settings = self.settings()
        source = self.choose_workbook(settings['paths']['templates'], prefix='manual_')
        rows = load_rows(source)
        devices = self.choose_file('Device credentials workbook', settings['paths']['devices'])
        execution = ExecutionConfig(self.concurrency(), **settings['execution'])
        bound = summary(rows, execution)
        self.write(f'Manual configuration: rows={bound["rows"]}; devices={bound["devices"]}; commands={bound["commands"]}; invalid={bound["invalid"]}; concurrency={bound["concurrency"]}')
        password = getpass('Manual authorization password: ')
        authorize(settings, password)
        if self.ask('Confirm manual execution? [y/N]: ').lower() != 'y':
            raise Cancelled()
        self.session_check(settings)
        validate_paths(settings)
        credentials = load_device_credentials(devices, [r['management_ip'] for r in rows if r['management_ip']])
        with self.activity('Manual configuration', total=bound['devices'], execution=True) as display:
            path = execute(rows, profile=settings, password=password, confirmation_digest=bound['digest'],
                           credentials=credentials, transport_config=TransportConfig(**settings['teleport']),
                           execution_config=execution, inventory_path=settings['paths']['inventory'],
                           progress=display.event, on_started=lambda path: self.write(f'Retained manual run: {path}'),
                           on_output=lambda output: self.write(f'Updated consolidated report: {output}\nCapture folder: {output.parent}'))

    def manual_recover(self):
        from orbitflow.configuration.manual import regenerate_report
        path = self.ask('Retained manual run directory: ')
        self.write(str(regenerate_report(path)))

    def manual_menu(self):
        self.submenu('Manual Configuration', [('Generate Manual Input Template', self.manual_template),
                     ('Execute Manual Configuration', self.manual_execute),
                     ('Regenerate Manual Report (offline)', self.manual_recover)])

    def settings(self):
        if not self.profile.exists():
            self.setup()
        with diagnostic("Invalid environment profile. Open Environment Setting and correct routing, pacing and paths."):
            return load_profile(self.profile)

    def choose_file(self, label, default):
        path = Path(self.ask(f'{label} [{default}] (b cancels): ', str(default)))
        if not path.is_file():
            raise FileNotFoundError('Missing input file')
        return path

    def choose_workbook(self, directory, *, prefix=''):
        """Choose a numbered .xlsx input, newest mtime first, or an explicit path."""
        root = Path(directory)
        files = sorted((p for p in root.iterdir() if p.is_file()
                        and p.suffix.lower() == '.xlsx' and p.name.startswith(prefix) and not p.name.startswith('~$')),
                       key=lambda p: (-p.stat().st_mtime_ns, p.name)) if root.exists() else []
        for index, path in enumerate(files, 1):
            self.write(f'{index}. {path.name}')
        if not files:
            self.write('No configuration workbooks found in the templates directory.')
        self.write(f'{len(files) + 1}. Enter explicit file path')
        while True:
            value = self.ask('Configuration workbook number (b cancels): ')
            if value == str(len(files) + 1):
                try:
                    path = self.choose_file('Configuration workbook', '')
                except FileNotFoundError:
                    self.write('Missing input file; choose again.')
                    continue
                if path.suffix.lower() == '.xlsx' and path.name.startswith(prefix) and not path.name.startswith('~$'):
                    return path
                self.write('Choose an .xlsx configuration workbook.')
            elif value.isdecimal() and 1 <= int(value) <= len(files):
                return files[int(value) - 1]
            else:
                self.write('Choose a listed workbook number or the explicit path option.')

    def select_targets(self, targets):
        """Select loaded row positions without flooding the console with addresses."""
        self.write(f'Validated device rows: {len(targets)} (numbered 1 through {len(targets)}).')
        while True:
            value = self.ask('Device rows [all, or comma-separated row numbers; default all]: ', 'all')
            self._selected_all = value.lower() == 'all'
            if self._selected_all:
                return list(range(len(targets)))
            try:
                positions = [int(part.strip()) - 1 for part in value.split(',')]
                if not positions or len(set(positions)) != len(positions) or any(i < 0 or i >= len(targets) for i in positions):
                    raise ValueError()
                return sorted(positions)
            except ValueError:
                self.write('Choose all or unique device row numbers in the displayed range.')

    @staticmethod
    def target_address(target):
        """Render only canonical IP literals; never echo arbitrary workbook text."""
        value = target.get('management_ip', '')
        try:
            # IPv6 scope identifiers permit arbitrary text, including controls.
            if not isinstance(value, str) or '%' in value:
                raise ValueError()
            return str(ip_address(value))
        except ValueError:
            return '[missing or invalid IP address]'

    def network(self, settings, devices, label):
        from orbitflow.targets import load_targets
        with diagnostic('Invalid device workbook. Check required management_ip, username and password columns.'):
            targets = load_targets(devices, isolate_invalid=True)
        self.target_positions = self.select_targets(targets)
        targets = [targets[i] for i in self.target_positions]
        limit = self.concurrency()
        self.write(f'{label}; input={devices}; target rows={len(targets)}; concurrency={limit}; inventory={settings["paths"]["inventory"]}')
        if self._selected_all:
            self.write(f'Selected targets: all {len(targets)} device rows.')
        else:
            self.write('Selected targets: ' + ', '.join(
                f'Row {position + 1}: {self.target_address(target)}'
                for position, target in zip(self.target_positions[:10], targets[:10])))
            if len(targets) > 10:
                self.write(f'{len(targets) - 10} additional selected rows omitted; total {len(targets)}.')
        if self.ask('Continue network operation? [y/N]: ').lower() != 'y':
            raise Cancelled()
        with diagnostic('Teleport session unavailable, expired or routing mismatched. Run tsh login externally; check proxy, cluster and identity files.'):
            self.session_check(settings)
        validate_paths(settings)
        from orbitflow.transport import TransportConfig
        return targets, TransportConfig(**settings['teleport']), ExecutionConfig(limit, **settings['execution'])

    def template(self):
        from orbitflow.configuration.v2.workbook import create_input
        root = Path(self.settings()['paths']['templates'])
        root.mkdir(parents=True, exist_ok=True)
        day = datetime.now().strftime('%Y%m%d')
        sequence = 1
        while True:
            path = root / f'config_template_{day}_{sequence:03d}.xlsx'
            lock = path.with_suffix('.lock')
            try:
                with lock.open('xb'):
                    pass
            except FileExistsError:
                sequence += 1
                continue
            try:
                if path.exists():
                    sequence += 1
                    continue
                create_input(path)  # Existing service also opens exclusively.
                break
            finally:
                lock.unlink(missing_ok=True)
        self.write(str(path))

    def planning(self):
        from orbitflow.configuration.v2.workbook import read_input, export_plan
        from orbitflow.configuration.v2.catalogue import BY_ID
        from orbitflow.configuration.v2.planning import plan, save
        settings = self.settings()
        workbook = self.choose_workbook(settings['paths']['templates'])
        with diagnostic('Invalid configuration workbook. Check template sheets and required input columns.'):
            requests = read_input(workbook)
        ips = {r['inputs']['management_ip'] for r in requests if BY_ID[r['template_id']].family != 'All'}
        observations = {}
        if ips:
            devices = self.choose_file('Device workbook', settings['paths']['devices'])
            selected, transport, execution = self.network(settings, devices, f'Read-only planning; selected targets={len(ips)}; output={settings["paths"]["plans"]}')
            selected_ips = {row['management_ip'] for row in selected}
            requests[:] = [r for r in requests if BY_ID[r['template_id']].family == 'All' or r['inputs']['management_ip'] in selected_ips]
            ips &= selected_ips
            from orbitflow.configuration.v2.observation import collect, bounded_observations
            from orbitflow.targets import load_device_credentials
            # Only fixed stage labels cross this presentation boundary; no raw event text.
            with self.activity('Processing', len(ips)) as display:
                observation_progress = self.observation_progress(ips, display)
                spool = collect(requests, load_device_credentials(devices, sorted(ips)), transport,
                                inventory_path=settings['paths']['inventory'], execution_config=execution,
                                progress=observation_progress)
            with self.activity('Loading persisted observations'):
                observations = bounded_observations(spool)
        with self.activity('Planning'):
            result = plan(requests, observations, source='workbook-sha256:' + requests.source_digest)
        from orbitflow.configuration.v2.cli import _default_output
        output = _default_output(requests, root=Path(settings['paths']['plans']))
        with self.activity('Saving immutable plan'):
            save(result, output)
        with self.activity('Exporting review workbook'):
            export_plan(result, output.with_suffix('.xlsx'))
        self.notify('Plan and review export complete.')
        self.write(f'Plan: {output}\nReview: {output.with_suffix(".xlsx")}\nDigest: {result.digest}')

    def execute(self):
        from orbitflow.configuration.execution import ApplySettings, ExecutionService, export_results, operator_summary, render_summary
        from orbitflow.configuration.v2.planning import load
        from orbitflow.targets import load_device_credentials, load_targets
        from orbitflow.transport import TransportConfig
        settings = self.settings()
        source = self.choose_plan(settings['paths']['plans'])
        workbook = self.choose_file('Engineer-reviewed workbook', source.with_suffix('.xlsx'))
        with diagnostic('Invalid immutable plan. Select the original intact plan JSON and matching reviewed workbook.'):
            original = load(source)
        digest = original.digest
        self.write(f'Immutable plan digest: {digest}')
        devices = self.choose_file('Device workbook', settings['paths']['devices'])
        with diagnostic('Invalid device workbook. Check required management_ip, username and password columns.'):
            load_targets(devices, isolate_invalid=True)
        self.write('Apply concurrency is separate from read-only execution settings.')
        limit = self.concurrency()
        service = ExecutionService('data/configuration/execution')
        with self.activity('Preparing/validating plan'), diagnostic('Plan/workbook invalid or confirmation consumed. Check matching original artifacts and decisions; review retained runs before any fresh plan.'):
            prepared = service.prepare_workbook(original, workbook, expected_digest=digest,
                                               settings=ApplySettings(max_concurrent_devices=limit, **settings['execution']))
        summary_text = render_summary(operator_summary(prepared))
        self.write('\n'.join(line for line in summary_text.splitlines() if not line.startswith(
            ('Warnings/blockers or skipped selections', 'Fresh checks + backup before writes'))))
        self.write('Review findings before confirming. Whole-config save may include partial changes; EdgeSwitch may interrupt management.')
        if prepared['status'] != 'awaiting_confirmation':
            raise OperatorError('Confirmation consumed. Review retained execution outcomes; do not replay.')
        if self.ask('Confirm execution? [y/N]: ').lower() != 'y':
            raise Cancelled()
        with diagnostic('Teleport session unavailable, expired or routing mismatched. Run tsh login externally; check proxy, cluster and identity files.'):
            self.session_check(settings)
        validate_paths(settings)
        ips = sorted({row['request']['inputs']['management_ip'] for row in prepared['summary']['selected']})
        started = monotonic()
        with self.activity('Processing', total=len(ips), execution=True) as display:
            spool = service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True,
                                    credentials=load_device_credentials(devices, ips),
                                    transport_config=TransportConfig(**settings['teleport']),
                                    inventory_path=settings['paths']['inventory'], progress=display.event)
        from orbitflow.configuration.execution_cli import outcome_summary
        destination = self.output(settings['paths']['execution_reports'], 'execution')
        with self.activity('Exporting execution report'):
            report = export_results(spool.path, destination)
        self.write(outcome_summary(spool, monotonic() - started, skipped=len(prepared['summary'].get('skipped', []))))
        self.write(str(report))

    @staticmethod
    def output(root, prefix):
        root = Path(root)
        root.mkdir(parents=True, exist_ok=True)
        return root / f'{prefix}_{datetime.now():%Y%m%d_%H%M%S}_{uuid4().hex}.xlsx'

    def run_workflow(self, kind, methodology=False):
        settings = self.settings()
        devices = self.choose_file('Device workbook', settings['paths']['devices'])
        destinations = {'inventory': 'outputs/reports/inventory', 'report': 'outputs/reports/interface_vlan',
                        'backup': 'outputs/backups/configuration', 'compliance': 'outputs/reports/vlan_compliance'}
        label = f'Read-only {kind}; output={destinations[kind]}'
        if kind == 'backup':
            label += '; sensitive configuration capture: restrict backup file access'
        if kind == 'compliance':
            from orbitflow.vlan_compliance import DEFAULT_POLICY
            label += f'; policy={DEFAULT_POLICY}; methodology={methodology}'
        targets, transport, execution = self.network(settings, devices, label)
        common = dict(inventory_path=settings['paths']['inventory'], execution_config=execution, log_root='outputs/logs')
        from orbitflow.terminal_progress import QuietOutput
        quiet = QuietOutput()
        with self.activity('Processing'):
            if kind == 'inventory':
                from orbitflow.inventory_refresh import refresh_inventory_from_excel
                result = refresh_inventory_from_excel(devices, transport, export_path=self.output('outputs/reports/inventory', 'inventory'),
                                                      spool_root='outputs/runs/inventory_refresh', target_positions=self.target_positions, **common)
            elif kind == 'report':
                from orbitflow.reporting import run_report
                result = run_report(targets, transport, spool_root='outputs/runs/interface_vlan_report', output=quiet, **common)
            elif kind == 'backup':
                from orbitflow.configuration_backup import run_backup
                result = run_backup(targets, transport, output=quiet, **common)
            else:
                from orbitflow.vlan_compliance import run_compliance, DEFAULT_POLICY
                from orbitflow.compliance import JsonPolicyProvider
                result = run_compliance(targets, transport, policy_provider=JsonPolicyProvider(DEFAULT_POLICY),
                                        spool_root='outputs/runs/vlan_compliance', cleanup=False,
                                        methodology_report=methodology, output=quiet, **common)
        # Some service returns contain sensitive records; never print arbitrary results.
        if isinstance(result, (str, Path)):
            self.write(f'Output: {result}')
        self.write('Run ended; review per-device outcomes and retained run logs.')

    def recover(self, kind):
        # Offline paths deliberately do not require a profile or a Teleport session.
        run = Path(self.ask('Retained run directory (b cancels): '))
        if kind == 'execution':
            from orbitflow.configuration.execution import export_results
            exporter = export_results
            root = 'outputs/reports/configuration_execution'
            if self.profile.exists():
                try:
                    root = load_profile(self.profile)['paths']['execution_reports']
                except (OSError, ValueError):
                    pass  # Broken environment must not prevent offline recovery.
            kwargs = {}
        else:
            root = f'outputs/reports/{kind}'
            if kind == 'vlan_compliance':
                from orbitflow.compliance_report import export_compliance_spool
                from orbitflow.methodology_report import export_methodology_spool
                exporter = export_methodology_spool if self.ask('Methodology report? [y/N]: ').lower() == 'y' else export_compliance_spool
            elif kind == 'inventory':
                from orbitflow.inventory_refresh import export_inventory_spool
                exporter = export_inventory_spool
            else:
                from orbitflow.reporting import export_report_spool
                exporter = export_report_spool
            kwargs = {'cleanup': False, 'allow_partial': self.ask('Export partial run? [y/N]: ').lower() == 'y'}
        with self.activity('Exporting report'), diagnostic('Invalid export or retained spool. Check run type and completeness; retain originals and select partial export only when supported.'):
            result = exporter(run, self.output(root, kind), **kwargs)
        self.write(str(result))

    def mac_lookup(self, search=False):
        """Two user operations over the shared read-only MAC application."""
        from orbitflow.mac_lookup import normalize_mac, validate_vlan, resolve_target, run_lookup, TargetResolutionError
        from orbitflow.targets import load_targets
        from orbitflow.transport import TransportConfig
        mac, vlan, address = None, None, None
        if search:
            value = self.ask('MAC (aabb.ccdd.eeff or aa:bb:cc:dd:ee:ff; b cancels): ')
            with diagnostic('Invalid MAC. Use dotted or colon notation.'):
                if '-' in value:
                    raise ValueError()
                mac = normalize_mac(value)
            with diagnostic('Invalid device IP address.'):
                address = str(ip_address(self.ask('Device IP (b cancels): ')))
        else:
            with diagnostic('Invalid VLAN. Enter a number from 1 through 4094.'):
                vlan = validate_vlan(self.ask('VLAN (b cancels): '))
        settings = self.settings()
        devices = self.choose_file('Device workbook', settings['paths']['devices'])
        if search:
            with diagnostic('Invalid device workbook. Check required device and credential columns.'):
                loaded = load_targets(devices, isolate_invalid=True)
            try:
                target = resolve_target(loaded, address)
            except TargetResolutionError as error:
                raise OperatorError(str(error)) from None
            self.write(f'Read-only MAC search; target={address}; output=outputs/reports/mac_table')
            if self.ask('Continue network operation? [y/N]: ').lower() != 'y':
                raise Cancelled()
            with diagnostic('Teleport session unavailable or expired. Run tsh login externally.'):
                self.session_check(settings)
            validate_paths(settings)
            targets, transport, execution = [target], TransportConfig(**settings['teleport']), ExecutionConfig(1, **settings['execution'])
        else:
            targets, transport, execution = self.network(settings, devices, 'Read-only MAC VLAN listing; output=outputs/reports/mac_table')
        with self.activity('Collecting MAC addresses'):
            result = run_lookup(targets, transport, mac=mac, vlan=vlan,
                                inventory_path=settings['paths']['inventory'], execution_config=execution)
        self.write(f'Output: {result}; inspect Details, Totals and Run Errors for mapping and collection limits.')
        if search:
            from openpyxl import load_workbook
            from orbitflow.logging import sanitize_text
            book = load_workbook(result, read_only=True)
            try:
                matches = 0
                for row in book['Details'].iter_rows(min_row=2, values_only=True):
                    if row[2] == mac and row[10].endswith('; selected'):
                        self.write(sanitize_text(f'Observation {row[0]}; MAC {row[2]}; attachment {row[3]}; '
                                                 f'domain {row[4] or "unknown"}; mapping {row[10]}; '
                                                 f'Physical interface: {row[13] or "unresolved/nonphysical"}; '
                                                 f'Port Type: {row[14]}; classification {row[11]}'))
                        matches += 1
                errors = list(book['Run Errors'].iter_rows(min_row=2, values_only=True))
                if any(row[1] == 'collection' for row in errors):
                    self.write('MAC collection failed; no successful no-match result is available.')
                elif not matches:
                    self.write('No matching MAC observations.')
                self.write(f'Matches: {matches}; collection/enrichment errors: {len(errors)}.')
            finally:
                book.close()

    def review(self):
        root = Path('outputs/reports/configuration_execution')
        if self.profile.exists():
            with diagnostic('Invalid environment profile. Correct execution report path in Environment Setting.'):
                root = Path(load_profile(self.profile)['paths']['execution_reports'])
        self.write('Execution workbooks (open for complete/partial/Unknown outcomes; no commands are sent):')
        skipped = 0
        for index, path in enumerate(sorted(root.glob('*.xlsx')), 1):
            book = None
            try:
                from openpyxl import load_workbook
                book = load_workbook(path, read_only=True, data_only=True)
                rows = list(book['Run Summary'].iter_rows(values_only=True))
                self.write(str(path))
                for row in rows:
                    if len(row) >= 2 and row[0] in ('Run', 'Selected / skipped rows', 'Rows by verification outcome', 'Devices by save outcome'):
                        self.write(f'{row[0]}: {row[1]}')
            except Exception:
                skipped += 1
                if skipped <= 10:
                    self.write(f'Skipped workbook #{index}: unreadable or missing Run Summary; check report file locally.')
            finally:
                if book is not None:
                    book.close()
        if skipped > 10:
            self.write(f'{skipped - 10} additional invalid workbooks skipped.')

    def submenu(self, title, actions):
        while True:
            self.write('')
            self.write(title + '\n' + '\n'.join(f'{i}. {label}' for i, (label, _) in enumerate(actions, 1)) + f'\n{len(actions)+1}. Back')
            choice = self.ask('Selection: ')
            if choice == str(len(actions)+1):
                return
            if choice.isdigit() and 1 <= int(choice) <= len(actions):
                self.safe(actions[int(choice)-1][1])
            else:
                self.write('Invalid selection.')

    def safe(self, action):
        try:
            action()
        except Cancelled:
            self.write('Cancelled.')
        except SystemExit:
            self.write('Workflow ended; inspect its module logs.')
        except OperatorError as error:
            self.write(str(error))
        except PermissionError:
            self.write('Permission denied. Check access to input, profile and output directories; keep backup access restricted.')
        except FileNotFoundError:
            self.write('Required input or executable missing. Check selected files and that tsh is installed on PATH.')
        except Exception:
            self.write('Operation failed. Check profile, input files, module logs and retained runs. For Teleport expiry, run tsh login outside the console.')

    def run(self):
        configuration = [('Generate Input Template', self.template), ('Generate Configuration Plan', self.planning),
                         ('Execute Approved Plan', self.execute), ('Review Executed Plan', self.review),
                         ('Regenerate Execution Report', lambda: self.recover('execution')),
                         ('Manual Configuration', self.manual_menu)]
        compliance = [('Run VLAN Compliance', lambda: self.run_workflow('compliance')),
                      ('Run VLAN Compliance + Methodology Report', lambda: self.run_workflow('compliance', True)),
                      ('Regenerate Previous Report (offline)', lambda: self.recover('vlan_compliance'))]
        def workflow(kind):
            actions = [('Run', lambda: self.run_workflow(kind))]
            if kind != 'backup':
                actions.append(('Regenerate Previous Report (offline)', lambda: self.recover('inventory' if kind == 'inventory' else 'interface_vlan')))
            self.submenu(kind, actions)
        actions = [('Configuration Management', lambda: self.submenu('Configuration Management', configuration)),
                   ('Device Inventory', lambda: workflow('inventory')), ('Interface & VLAN Reporting', lambda: workflow('report')),
                   ('VLAN Compliance', lambda: self.submenu('VLAN Compliance', compliance)),
                   ('Configuration Backup', lambda: workflow('backup')), ('Environment Setting', self.environment_menu),
                   ('MAC Address Lookup', lambda: self.submenu('MAC Address Lookup', [
                       ('Search MAC Address', lambda: self.mac_lookup(True)),
                       ('List All MAC Addresses', self.mac_lookup)]))]
        try:
            while True:
                self.write('')
                self.write('\n'.join(f'{i}. {label}' for i, (label, _) in enumerate(actions, 1)) + f'\n{len(actions)+1}. Exit')
                choice = self.ask('Selection: ')
                if choice == str(len(actions)+1):
                    return
                if choice.isdigit() and 1 <= int(choice) <= len(actions):
                    self.safe(actions[int(choice)-1][1])
                else:
                    self.write('Invalid selection.')
        except (EOFError, KeyboardInterrupt, Cancelled):
            self.write('Console exited; inspect retained execution runs after an interruption.')


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--profile', type=Path, default=PROFILE, help='Non-secret TOML environment profile')
    args = parser.parse_args(argv)
    Console(args.profile).run()


if __name__ == '__main__':
    main()
