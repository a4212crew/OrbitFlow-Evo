"""Read-only batch Interface/VLAN reporting over shared device capabilities."""

from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
import re
import sys

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from orbitflow.capabilities.interfaces import InterfaceService
from orbitflow.capabilities.vlans import VlanService
from orbitflow.inventory import DeviceInventoryResolver, JsonInventoryStore
from orbitflow.logging import module_logger, sanitize_text
from orbitflow.targets import REQUIRED_COLUMNS
from orbitflow.transport import DeviceCredentials, connect_device
from orbitflow.vendors.interface_names import canonical_interface_name

INTERFACE_COLUMNS = (
    "Device Name", "Device IP", "Platform", "Device Family", "Interface",
    "Description", "Admin Status", "Oper Status", "Mode", "Attached VLANs",
    "Access VLAN", "Native VLAN", "PVID", "Allowed VLANs", "Tagged VLANs",
    "Untagged VLANs", "Service VLAN", "Outer VLAN", "Inner VLAN",
    "Service Binding Type", "Service Binding Name", "Collection Time",
    "Control VLAN", "Excluded VLANs", "VLAN Source",
)
DATABASE_COLUMNS = (
    "Device Name", "Device IP", "Platform", "Object Type", "Object ID",
    "VLAN ID(s)", "Name", "Collection Time",
)
ERROR_COLUMNS = ("Device IP", "Device Name", "Stage", "Error Category", "Time")
_VLAN_FIELDS = (
    "access_vlan", "native_vlan", "pvid", "allowed_vlans", "tagged_vlans",
    "untagged_vlans", "service_vlan", "outer_vlan", "inner_vlan", "control_vlan",
    "referenced_vlans",
)


def _ids(values):
    return ", ".join(str(value) for value in sorted(set(values)))


def _detail(observations, field):
    values = []
    for observation in observations:
        value = getattr(observation, field)
        if isinstance(value, tuple):
            value = _ids(value) if value else ("none" if field == "allowed_vlans" else "")
        values.append("" if value is None else str(value))
    # One line per observation, including placeholders, preserves associations.
    return "\n".join(value or "-" for value in values) if len(values) > 1 else "".join(values)


def build_rows(context, interfaces, vlans):
    """Join within one resolved device; preserve unmatched and logical interfaces."""
    key = lambda name: canonical_interface_name(context.platform, name)
    records = {key(record.port_name): record for record in interfaces}
    observations = defaultdict(list)
    if vlans is not None:
        for observation in vlans.interfaces:
            observations[key(observation.interface_name)].append(observation)
    rows = []
    identity = [context.hostname, context.management_ip, context.platform]
    for name in sorted(records.keys() | observations.keys()):
        record = records.get(name)
        items = sorted(observations[name], key=repr)
        attached = set()
        for item in items:
            for field in _VLAN_FIELDS:
                value = getattr(item, field)
                attached.update(value if isinstance(value, tuple) else (() if value is None else (value,)))
        collected = record.collection_time if record else vlans.collection_time
        rows.append(identity + [
            context.device_family, record.port_name if record else items[0].interface_name,
            record.port_description if record else "",
            record.admin_status if record else "", record.oper_status if record else "",
            _detail(items, "mode"), _ids(attached),
            *[_detail(items, field) for field in (
                "access_vlan", "native_vlan", "pvid", "allowed_vlans", "tagged_vlans",
                "untagged_vlans", "service_vlan", "outer_vlan", "inner_vlan",
                "service_binding_type", "service_binding_name",
            )], collected.isoformat(),
            *[_detail(items, field) for field in ("control_vlan", "excluded_vlans", "vlan_source")],
        ])
    database = [] if vlans is None else [
        identity + [obj.object_type, obj.object_id, _ids(obj.vlan_ids), obj.name,
                    vlans.collection_time.isoformat()]
        for obj in sorted(vlans.objects, key=lambda obj: (obj.object_type, obj.object_id, obj.name, obj.vlan_ids))
    ]
    return rows, database


def write_workbook(path, interfaces, database, errors, *, clean=sanitize_text):
    """Write once using bounded worksheet memory; untrusted values are literal text."""
    workbook = Workbook(write_only=True)
    for title, columns, rows in (
        ("Interfaces", INTERFACE_COLUMNS, interfaces),
        ("VLAN_Database", DATABASE_COLUMNS, database),
        ("Run_Errors", ERROR_COLUMNS, errors),
    ):
        sheet = workbook.create_sheet(title)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
        for index, column in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = (
                42 if column in {"Description", "Attached VLANs", "Allowed VLANs"}
                else 28 if column in {"Collection Time", "Time", "Device Name", "Service Binding Name"}
                else max(18, len(column) + 2)
            )
        for index, row in enumerate(_with_header(columns, rows)):
            cells = []
            for value in row:
                text = ILLEGAL_CHARACTERS_RE.sub("", str(value) if index == 0 else clean(value))
                if len(text) > 32767:
                    raise ValueError("Report cell exceeds Excel text limit")
                cell = WriteOnlyCell(sheet, value=text)
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if index == 0:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="24476A")
                cells.append(cell)
            sheet.append(cells)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def _with_header(columns, rows):
    yield columns
    yield from rows


def run_report(targets, transport_config, *, inventory_path="data/live_validation/inventory.json",
               reports_dir="reports", log_root="logs", output=None, clock=None):
    """Accept the approved list of target dictionaries; isolate failures by stage."""
    output = output if output is not None else sys.stdout
    clock = clock or (lambda: datetime.now(timezone.utc))
    targets = list(targets)
    # Redact known input credentials even if echoed into normalized device text.
    secrets = {value for target in targets for field, value in target.items()
               if field in {"username", "password", "secret", "token", "otp", "private_key"}
               and isinstance(value, str) and value}
    pattern = re.compile("|".join(re.escape(value) for value in sorted(secrets, key=lambda v: (-len(v), v)))) if secrets else None

    def clean(value):
        text = sanitize_text(value)
        return pattern.sub("[REDACTED]", text) if pattern else text

    started = clock()
    path = Path(reports_dir) / f"device_interface_vlan_report_{started.strftime('%Y%m%dT%H%M%S_%fZ')}.xlsx"
    interface_rows, database_rows, errors = [], [], []
    resolver = DeviceInventoryResolver(JsonInventoryStore(inventory_path))
    interface_service, vlan_service = InterfaceService(), VlanService()
    with module_logger("reporting", "interface_vlan_report", log_root=log_root) as (logger, _):
        logger.info("Interface/VLAN batch started")

        def failure(ip, name, stage, exc):
            category = type(exc).__name__
            errors.append([clean(ip), clean(name), stage, category, clock().isoformat()])
            logger.error(f"Report stage failed: {stage}",
                         extra={"management_ip": clean(ip), "error_category": category})

        for target in targets:
            ip, name, stage = target.get("management_ip", ""), "", "input"
            context, interfaces, vlans = None, [], None
            try:
                if not all(isinstance(target.get(field), str) and target[field].strip() for field in REQUIRED_COLUMNS):
                    raise ValueError("Missing required target field")
                stage = "connect"
                with connect_device(ip, DeviceCredentials(target["username"], target["password"]), transport_config) as session:
                    stage = "inventory"
                    context = resolver.resolve(session, management_ip=ip)
                    name = context.hostname
                    for stage, service in (("interfaces", interface_service), ("vlans", vlan_service)):
                        try:
                            result = service.collect(session, context)
                            if stage == "interfaces":
                                interfaces = result
                            else:
                                vlans = result
                        except Exception as exc:
                            failure(ip, name, stage, exc)
                    stage = "disconnect"
            except Exception as exc:
                failure(ip, name, stage, exc)
            if context is not None:
                try:
                    rows, objects = build_rows(context, interfaces, vlans)
                    interface_rows.extend(rows)
                    database_rows.extend(objects)
                except Exception as exc:
                    failure(ip, name, "normalize", exc)
        try:
            write_workbook(path, interface_rows, database_rows, errors, clean=clean)
        except Exception as exc:
            logger.error("Report workbook write failed", extra={"error_category": type(exc).__name__})
            raise
        logger.info("Interface/VLAN batch completed")
    print(f"Devices: {len(targets)}; stage failures: {len(errors)}; report: {path}", file=output)
    return path
