"""Refresh approved Excel targets and export the complete latest inventory."""

from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from orbitflow.execution import execute_devices
from orbitflow.result_spool import ResultSpool
from orbitflow.inventory import DeviceInventoryResolver, JsonInventoryStore
from orbitflow.logging import module_logger, sanitize_text
from orbitflow.models import DeviceContext
from orbitflow.targets import REQUIRED_COLUMNS, load_targets
from orbitflow.transport import DeviceCredentials, TransportConfig, connect_device


REFRESHED = "refreshed"
FAILED = "failed-with-previous-data-retained"
NOT_REQUESTED = "not-requested-this-run"
INVENTORY_COLUMNS = (
    "supplied_in_current_input", "refresh_status", *DeviceContext.__dataclass_fields__,
)
ATTEMPT_COLUMNS = (
    "input_position", "management_ip", "device_ids", "status", "stage",
    "error_category", "attempted_at", "reconciliation_events",
)


def refresh_inventory_from_excel(
    excel_path: str | Path,
    transport_config: TransportConfig,
    *,
    inventory_path: str | Path,
    export_path: str | Path,
    log_root: str | Path = "logs",
    clock=None,
    execution_config=None,
    spool_root=None,
) -> Path:
    """Refresh only input targets, then export all latest-known stable facts.

    The shared loader's required columns are management_ip, username, password.
    Bad rows are isolated. Successful membership follows the resolved physical
    device ID; failed attempts use the store's observed-address matching rule.
    For repeated/alias targets, the last attempt affecting an identity wins.
    Run_Attempts retains every outcome, including unresolved failures and cleanup
    errors after a successful refresh. Run status is not persisted as history.

    Inventory transactions are thread-safe within this process. Export paths
    must have one writer. No live interface or VLAN capabilities are invoked.
    """
    source, inventory, destination = map(Path, (excel_path, inventory_path, export_path))
    # Reject accidental input/inventory destruction before performing any refresh.
    if len({path.resolve() for path in (source, inventory, destination)}) != 3:
        raise ValueError("Input, inventory and export paths must be distinct")
    targets = load_targets(source, isolate_invalid=True)
    clock = clock or (lambda: datetime.now(timezone.utc))
    sensitive = {target[field] for target in targets for field in ("username", "password")
                 if target.get(field)}
    pattern = re.compile("|".join(re.escape(value) for value in
                         sorted(sensitive, key=lambda value: (-len(value), value)))) if sensitive else None

    def clean(value):
        text = sanitize_text(value)
        return pattern.sub("[REDACTED]", text) if pattern else text

    store = JsonInventoryStore(inventory)
    # Fail before connecting if the existing snapshot cannot be read.
    store.contexts()
    spool = ResultSpool.create(spool_root or destination.parent / "runs",
                               "inventory_refresh", len(targets), clean=clean)
    with module_logger("inventory", "inventory_refresh", log_root=log_root) as (logger, _):
        logger.info("Inventory refresh started")
        def collect(item):
            position, target = item
            resolver = DeviceInventoryResolver(store, clock=clock, sanitize_fact=clean)
            statuses = {}
            ip = target["management_ip"]
            attempted_at = clock()
            stage, context, events = "input", None, ()
            affected = ()
            category = ""
            try:
                if not all(target.get(field) for field in REQUIRED_COLUMNS):
                    raise ValueError("Missing required target field")
                # An address containing a credential cannot be safely persisted.
                if clean(ip) != ip:
                    raise ValueError("Unsafe target address")
                stage = "connect"
                with connect_device(ip, DeviceCredentials(target["username"], target["password"]),
                                    transport_config) as session:
                    stage = "inventory"
                    context = resolver.resolve(session, management_ip=ip)
                    events = resolver.last_events
                    affected = (context.device_id,)
                    statuses[context.device_id] = REFRESHED
                    stage = "disconnect"
                stage = "complete"
            except Exception as exc:
                category = type(exc).__name__
                logger.error("Inventory refresh attempt failed", exc_info=True,
                             extra={"error_category": clean(category)})
                if context is None:
                    # Resolver failures already record their own safe metadata.
                    # Input/connect failures need the same retention behaviour.
                    if stage != "inventory" and ip and clean(ip) == ip:
                        store.record_failure(ip, attempted_at, f"inventory refresh failed ({clean(category)})")
                    affected = tuple(item.device_id for item in store.contexts()
                                     if ip and ip in item.observed_management_ips)
                    for device_id in affected:
                        statuses[device_id] = FAILED
            return statuses, (position, clean(ip), affected,
                             "failed" if category else REFRESHED, stage, clean(category),
                             attempted_at, events)

        def persist(outcome):
            target = targets[outcome.position - 1]
            if outcome.error_category:
                logger.error("Inventory worker failed",
                             extra={"error_category": clean(outcome.error_category)})
                payload = ({}, (outcome.position, target["management_ip"], (),
                                "failed", "inventory", outcome.error_category, clock(), ()))
            else:
                payload = outcome.value
            spool.append(outcome, payload=payload, target=target["management_ip"],
                         failed=payload[1][3] == "failed")

        logger.info(f"Inventory run spool: {clean(str(spool.path))}")
        with spool.collection():
            execute_devices(enumerate(targets, 1), collect, config=execution_config, on_outcome=persist)
            # Only per-identity membership is retained; attempt/report rows stay on disk.
            statuses = {}
            for record in spool.records():
                statuses.update(record['payload'][0])
            def inventory_rows():
                for context in sorted(store.contexts(), key=lambda item: (item.management_ip, item.device_id)):
                    state = statuses.get(context.device_id, NOT_REQUESTED)
                    facts = asdict(context)
                    yield (context.device_id in statuses, state,
                           *(facts[field] for field in DeviceContext.__dataclass_fields__))
            # Capture the run's export view for retry even if canonical inventory later changes.
            spool.snapshot(inventory_rows())
        export_inventory_spool(spool.path, destination)
        logger.info("Inventory refresh and export completed")
    return destination


def _write_export(path, inventory_rows, attempts, *, clean):
    """Write literal, sanitized cells and replace the destination only on success."""
    workbook = Workbook(write_only=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        for title, columns, rows in (
            ("Inventory", INVENTORY_COLUMNS, inventory_rows),
            ("Run_Attempts", ATTEMPT_COLUMNS, attempts),
        ):
            sheet = workbook.create_sheet(title)
            sheet.freeze_panes = "A2"
            for index, column in enumerate(columns, 1):
                sheet.column_dimensions[get_column_letter(index)].width = max(22, len(column) + 2)
            headers = []
            for column in columns:
                cell = WriteOnlyCell(sheet, value=column)
                cell.font = Font(bold=True)
                headers.append(cell)
            sheet.append(headers)
            row_number = 1
            for row_number, row in enumerate(rows, 2):
                cells = []
                for value in row:
                    if isinstance(value, datetime):
                        value = value.isoformat()
                    elif isinstance(value, (tuple, list)):
                        value = "; ".join(value)
                    text = ILLEGAL_CHARACTERS_RE.sub("", clean(value))
                    if len(text) > 32767:
                        raise ValueError("Inventory export cell exceeds Excel text limit")
                    cell = WriteOnlyCell(sheet, value=text)
                    cell.data_type = "s"
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cells.append(cell)
                sheet.append(cells)
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{row_number}"
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        for sheet in workbook:
            if not sheet.closed:
                sheet.close()
            if sheet._writer and Path(sheet._writer.out).exists():
                sheet._writer.cleanup()
        workbook.close()
        if temporary.exists():
            temporary.unlink()


def export_inventory_spool(spool_path, path, *, cleanup=True, allow_partial=False):
    """Retry the captured run export; never reads or mutates canonical inventory."""
    spool = ResultSpool(spool_path)
    if spool.manifest['task'] != 'inventory_refresh':
        raise ValueError('Not an inventory refresh spool')
    def consume(run):
        attempts = (record['payload'][1] for record in run.records())
        rows = run.snapshot_rows() if (run.path / "snapshot.jsonl").exists() else ()
        if not (run.path / "snapshot.jsonl").exists() and not allow_partial:
            raise ValueError("Inventory export snapshot is missing")
        _write_export(Path(path), rows, attempts, clean=sanitize_text)
        return Path(path)
    return spool.consume(consume, cleanup=cleanup, allow_partial=allow_partial)
