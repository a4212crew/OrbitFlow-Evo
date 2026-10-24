"""Streaming, atomic, literal-text workbook output for normalized application data."""

from pathlib import Path
from orbitflow.logging import sanitize_text
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def write_tables(path, tables, *, clean=sanitize_text, comments=None):
    """Atomically stream literal text; reject oversized cells with their location.

    ValueError identifies sheet, one-based row, column letter and character
    count without including cell contents. The existing destination is retained.
    Optional comments(title, one_based_column) supplies sanitized text for the
    current data row; callers own evidence size/projection policy.
    """
    workbook = Workbook(write_only=True)
    path = Path(path)
    temporary = path.with_name(path.name + ".tmp")
    try:
        for title, columns, rows in tables:
            sheet = workbook.create_sheet(title)
            sheet.freeze_panes = "A2"

            for index, column in enumerate(columns, 1):
                sheet.column_dimensions[get_column_letter(index)].width = (
                    42 if column in {"Description", "Service Mappings", "Tagged VLANs", "Bridge Domains",
                                     "Configuration Evidence", "Evidence Source", "Evidence Lines",
                                     "Expected", "Observed", "Missing VLANs", "Missing Objects", "Evidence", "JSON Fragment"}
                    else 28 if column in {"Collection Time", "Time", "Device Name", "Service Binding Name"}
                    else max(18, len(column) + 2)
                )
            for index, row in enumerate(_with_header(columns, rows)):
                cells = []
                for column_index, value in enumerate(row, 1):
                    text = ILLEGAL_CHARACTERS_RE.sub("", str(value) if index == 0 else clean(value))
                    if len(text) > 32767:
                        raise ValueError(
                            f"Report cell exceeds Excel text limit (32767): "
                            f"sheet={title}, row={index + 1}, "
                            f"column={get_column_letter(column_index)}, characters={len(text)}"
                        )
                    cell = WriteOnlyCell(sheet, value=text)
                    cell.data_type = "s"
                    if index and comments:
                        comment = comments(title, column_index)
                        if comment:
                            cell.comment = Comment(clean(comment), "OrbitFlow")
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if index == 0:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill("solid", fgColor="24476A")
                    cells.append(cell)
                sheet.append(cells)
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{index + 1}"
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        for sheet in workbook:
            if not sheet.closed:
                sheet.close()
            if sheet._writer and Path(sheet._writer.out).exists():
                sheet._writer.cleanup()
        workbook.close()
        if temporary.exists():
            temporary.unlink()


def _with_header(columns, rows):
    yield columns
    yield from rows

