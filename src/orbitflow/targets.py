"""Shared approved Excel target input; no device discovery or output logging."""

from pathlib import Path
from openpyxl import load_workbook

REQUIRED_COLUMNS = ("management_ip", "username", "password")
OPTIONAL_COLUMNS = ("secret",)


def load_targets(path: str | Path, *, isolate_invalid: bool = False) -> list[dict[str, str]]:
    """Load required fields and optional secret, ignoring unknown columns.

    By default incomplete rows raise, preserving inventory validation behavior.
    Batch consumers may retain them for per-row error recording instead.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise ValueError("Excel input is empty")
        headers = {
            str(value).strip().casefold(): index
            for index, value in enumerate(header_row)
            if value is not None
        }
        missing = [name for name in REQUIRED_COLUMNS if name not in headers]
        if missing:
            raise ValueError(
                f"Excel input is missing required columns: {', '.join(missing)}"
            )
        targets = []
        for row_number, row in enumerate(rows, start=2):
            values = {
                name: (
                    ""
                    if headers[name] >= len(row) or row[headers[name]] is None
                    else (str(row[headers[name]]) if name == "secret"
                          else str(row[headers[name]]).strip())
                )
                for name in (*REQUIRED_COLUMNS, *OPTIONAL_COLUMNS) if name in headers
            }
            if not any(values.values()):
                continue
            if not all(values[name] for name in REQUIRED_COLUMNS) and not isolate_invalid:
                raise ValueError(f"Excel row {row_number} has an empty required field")
            targets.append(values)
        if not targets:
            raise ValueError("Excel input contains no device rows")
        return targets
    finally:
        workbook.close()



def load_device_credentials(path, management_ips, *, provider=None):
    """Resolve requested IPs with provider precedence and isolated Excel failures.

    The provider returns DeviceCredentials or None. Duplicate Excel targets are
    ambiguous even when their values agree. Never expose input values in errors.
    """
    from orbitflow.transport import DeviceCredentials
    ips = list(dict.fromkeys(management_ips))
    resolved = {}
    pending = []
    for ip in ips:
        try:
            supplied = provider(ip) if provider else None
        except Exception:
            continue  # Provider failure is isolated; never downgrade authority.
        if supplied is not None:
            resolved[ip] = supplied
        else:
            pending.append(ip)
    if not pending:
        return resolved
    try:
        rows = load_targets(path, isolate_invalid=True)
    except Exception:
        return resolved
    by_ip = {}
    for row in rows:
        by_ip.setdefault(row.get("management_ip", ""), []).append(row)
    for ip in pending:
        matches = by_ip.get(ip, [])
        if len(matches) == 1 and all(matches[0].get(k) for k in REQUIRED_COLUMNS):
            row = matches[0]
            resolved[ip] = DeviceCredentials(row["username"], row["password"], secret=row.get("secret") or None)
    return resolved
