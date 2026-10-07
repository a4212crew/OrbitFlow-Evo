from contextlib import contextmanager
from copy import deepcopy
from dataclasses import replace
from datetime import datetime, timezone
from importlib.util import module_from_spec, spec_from_file_location
from io import StringIO
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

from openpyxl import load_workbook
import pytest

from orbitflow import vlan_compliance as app
from orbitflow import compliance_report as report
from orbitflow.compliance import JsonPolicyProvider, evaluate_vlan_compliance, parse_policy
from orbitflow.config import ExecutionConfig
from orbitflow.execution import DeviceOutcome
from orbitflow.models import DeviceContext, InterfaceRecord, InterfaceVlanObservation, VlanObject, VlanState
from orbitflow.result_spool import ResultSpool
from orbitflow.vendors.configuration_facts import observe_configuration


NOW = datetime(2026, 10, 3, tzinfo=timezone.utc)
POLICY_DATA = json.loads(app.DEFAULT_POLICY.read_text(encoding="utf-8"))
POLICY = parse_policy(POLICY_DATA)


def context(ip="192.0.2.1"):
    return DeviceContext("device-" + ip, ip, (ip,), "switch", "Cisco", "cisco_ios",
                         "C3750X", "", "", (), "", "", "", NOW, NOW)


def interface(name="Gi0/1"):
    return InterfaceRecord("switch", "192.0.2.1", "cisco_ios", name, "", "up", "up", NOW)


def state(tags=(445, 545, 2449), port_type="trunk", objects=None):
    if objects is None:
        objects = tuple(VlanObject("vlan", str(v)) for v in POLICY.required_domains)
    database = ",".join(o.domain_id for o in objects if o.object_type == "vlan")
    allowed = ",".join(map(str, tags)) if isinstance(tags, tuple) else tags.lower()
    config = (f"vlan {database}\n!\n" if database else "") + (
        f"interface GigabitEthernet0/1\n switchport trunk encapsulation dot1q\n"
        f" switchport mode {port_type}\n switchport trunk allowed vlan {allowed or 'none'}\n!")
    return VlanState("switch", "192.0.2.1", "cisco_ios",
                     (InterfaceVlanObservation("GigabitEthernet0/1", port_type=port_type, tagged_vlans=tags),),
                     objects, NOW, observe_configuration(config, "cisco_ios"))


def evaluate(vlans=None, *, policy=POLICY):
    return evaluate_vlan_compliance(context(), [interface()], vlans or state(), policy)


def test_default_policy_and_provider(tmp_path):
    assert POLICY.required_vlans == (*range(2400, 2445), 2449, 4001)
    assert set(map(int, POLICY.required_domains)) == {445, 545, *range(2400, 2445), 2449, 4001}
    assert JsonPolicyProvider(app.DEFAULT_POLICY).load() == POLICY
    path = tmp_path / "invalid.json"
    path.write_text('{"schema_version": 1, "schema_version": 2}')
    with pytest.raises(ValueError, match="Invalid compliance policy JSON"):
        JsonPolicyProvider(path).load()


@pytest.mark.parametrize("section,key,value", [
    (None, "schema_version", True), (None, "schema_version", 1),
    (None, "extra", "typo"), (None, "policy_id", "password=do-not-echo"),
    ("database", "required_domains", []), ("database", "object_types", ["unknown"]),
    ("database", "required_domains", ["named-service"]),
    ("interface", "match_all", [True]), ("interface", "match_any", [0]),
    ("database", "required_domains", [1]), ("interface", "match_all", ["1-10"]),
    ("interface", "required_vlans", [4002]), ("interface", "required_vlans", ["2444-2400"]),
    ("interface", "required_vlans", ["ALL"]), ("interface", "match_any", "2449"),
    ("interface", "rule_id", "required-forwarding-domains"),
    (None, "service_rules", []),
])
def test_invalid_policy_rejected_without_echo(section, key, value):
    data = deepcopy(POLICY_DATA)
    (data if section is None else data[section])[key] = value
    with pytest.raises(ValueError) as caught:
        parse_policy(data)
    assert "do-not-echo" not in str(caught.value)


def test_policy_change_changes_both_rules_without_code():
    data = deepcopy(POLICY_DATA)
    data["database"].update(required_domains=[10, 20, 30])
    data["interface"].update(match_all=[10], match_any=[20, 30], required_vlans=[10, 20, 30])
    findings = evaluate(state((10, 20), objects=tuple(VlanObject("vlan", str(v)) for v in (10, 20))), policy=parse_policy(data))
    assert findings[0]["missing_vlans"] == [30]
    assert findings[1]["missing_vlans"] == [30]


@pytest.mark.parametrize("tags,status,missing", [
    ((445, 545, 2449), "non_compliant", [*range(2400, 2445), 4001]),
    ((445, 545, 4001), "non_compliant", [*range(2400, 2445), 2449]),
    ((445, 2449, 4001), "not_applicable", []),
    ((545, 2449, 4001), "not_applicable", []),
    ((445, 545), "not_applicable", []),
    ((445, 545, *POLICY.required_vlans), "compliant", []),
    ("ALL", "compliant", []), ("NONE", "not_applicable", []),
    ((), "not_applicable", []),
])
def test_trunk_signature_and_membership(tags, status, missing):
    finding = evaluate(state(tags))[1]
    assert finding["interface"] == "Gi0/1"
    assert finding["status"] == status
    assert finding["missing_vlans"] == missing


def test_evidence_missing_is_not_silently_assessed():
    findings = evaluate(replace(state(), configuration=None))
    assert all(f["status"] == "unable_to_assess" for f in findings)


def install_fakes(monkeypatch, fail=None):
    events, closed = [], []

    @contextmanager
    def connect(ip, credentials, config):
        assert credentials.password == "synthetic-password"
        events.append((ip, "connect"))
        if ip.endswith(".1") and fail == "connect":
            raise RuntimeError("synthetic-password token=never-output")
        session = SimpleNamespace(ip=ip)
        try:
            yield session
        finally:
            closed.append(ip)
            if ip.endswith(".1") and fail == "disconnect":
                raise RuntimeError("synthetic-password")

    @contextmanager
    def cli(session):
        session.cli = object()
        yield session.cli

    def resolve(session, *, management_ip, cli):
        assert cli is session.cli
        events.append((session.ip, "inventory"))
        if session.ip.endswith(".1") and fail == "inventory":
            raise RuntimeError("synthetic-password")
        session.context = replace(context(session.ip), hostname="=synthetic-password", serial_number="secret-output")
        return session.context

    def collect(stage):
        def run(session, ctx, *, cli):
            assert cli is session.cli and ctx is session.context
            events.append((session.ip, stage))
            if session.ip.endswith(".1") and fail == stage:
                raise RuntimeError("synthetic-password token=never-output")
            if stage == "interfaces":
                return [replace(interface(), port_description="secret-output")]
            return state((445, 545, *POLICY.required_vlans))
        return run

    monkeypatch.setattr(app, "connect_device", connect)
    monkeypatch.setattr(app, "DeviceCLI", cli)
    monkeypatch.setattr(app, "DeviceInventoryResolver", Mock(return_value=Mock(resolve=resolve)))
    monkeypatch.setattr(app, "InterfaceService", Mock(return_value=Mock(collect=collect("interfaces"))))
    monkeypatch.setattr(app, "VlanService", Mock(return_value=Mock(collect=collect("vlans"))))
    return events, closed


def targets(count=2):
    return [{"management_ip": f"192.0.2.{i}", "username": "synthetic-user", "password": "synthetic-password"}
            for i in range(1, count + 1)]


def options(tmp_path):
    return dict(inventory_path=tmp_path / "inventory.json", spool_root=tmp_path / "runs",
                log_root=tmp_path / "logs", output=StringIO(), execution_config=ExecutionConfig(2, 0, 0))


@pytest.mark.parametrize("stage", [None, "input", "connect", "inventory", "interfaces", "vlans", "disconnect"])
def test_application_failure_isolation_and_safe_reusable_results(tmp_path, monkeypatch, stage):
    events, closed = install_fakes(monkeypatch, stage)
    inputs = targets()
    if stage == "input":
        inputs[0]["password"] = ""
    opts = options(tmp_path)
    run = app.collect_compliance(inputs, object(), **opts)
    records = list(ResultSpool(run.spool_path).records())
    assert len(records) == 2
    assert run.failed_devices == (0 if stage is None else 1)
    assert all(f["status"] == "compliant" for f in records[1]["payload"]["findings"])
    assert [event for ip, event in events if ip.endswith(".2")] == ["connect", "inventory", "interfaces", "vlans"]
    assert "192.0.2.2" in closed
    if stage == "interfaces":
        assert [f["status"] for f in records[0]["payload"]["findings"]] == ["compliant", "compliant"]
    if stage == "vlans":
        assert [f["status"] for f in records[0]["payload"]["findings"]] == ["unable_to_assess", "unable_to_assess"]
    output_text = json.dumps(records) + opts["output"].getvalue()
    output_text += "".join(p.read_text() for p in (tmp_path / "logs").rglob("*.log"))
    for secret in ("synthetic-password", "synthetic-user", "secret-output", "never-output"):
        assert secret not in output_text
    path = report.export_compliance_spool(run.spool_path, tmp_path / "compliance.xlsx")
    assert not run.spool_path.exists()
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["Findings", "Run_Errors", "Details"]
        sheet = workbook["Findings"]
        assert sheet.freeze_panes == "A2"
        assert sheet["C4"].value == "=[REDACTED]"
        assert sheet["C4"].data_type == "s"
        assert sheet.auto_filter.ref == "A1:AA5"
    finally:
        workbook.close()


def test_report_failure_recovery_never_recollects(tmp_path, monkeypatch):
    events, _ = install_fakes(monkeypatch)
    writer = report.write_tables
    monkeypatch.setattr(report, "write_tables", Mock(side_effect=OSError("synthetic-password")))
    with pytest.raises(OSError):
        app.run_compliance(targets(), object(), reports_dir=tmp_path, **options(tmp_path))
    spool_path, = (tmp_path / "runs").iterdir()
    assert ResultSpool(spool_path).manifest["status"] == "output_failed"
    previous = list(events)
    monkeypatch.setattr(report, "write_tables", writer)
    report.export_compliance_spool(spool_path, tmp_path / "recovered.xlsx")
    assert events == previous
    assert not spool_path.exists()


def test_worker_failure_and_policy_fail_before_connect(tmp_path, monkeypatch):
    install_fakes(monkeypatch)
    evaluate = app.evaluate_vlan_compliance
    def fail(context, *args, **kwargs):
        if context and context.management_ip.endswith(".1"):
            raise RuntimeError("synthetic-password")
        return evaluate(context, *args, **kwargs)
    monkeypatch.setattr(app, "evaluate_vlan_compliance", fail)
    run = app.collect_compliance(targets(), None, **options(tmp_path))
    records = list(ResultSpool(run.spool_path).records())
    assert records[0]["payload"]["errors"][0]["stage"] == "worker"
    assert records[1]["status"] == "success"
    monkeypatch.setattr(app, "connect_device", Mock(side_effect=AssertionError("must not connect")))
    with pytest.raises(ValueError):
        app.collect_compliance(targets(), None, policy_provider=Mock(load=Mock(side_effect=ValueError())), **options(tmp_path))
    app.connect_device.assert_not_called()


def test_large_normalized_evidence_is_preserved_in_excel_details(tmp_path):
    objects = tuple(VlanObject("vlan", str(v)) for v in range(2, 4095))
    findings = evaluate(state(objects=objects))
    spool = ResultSpool.create(tmp_path / "runs", "vlan_compliance", 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={"findings": findings, "errors": []})
    path = report.export_compliance_spool(spool.path, tmp_path / "large.xlsx", cleanup=False)
    workbook = load_workbook(path)
    try:
        assert workbook["Findings"]["J2"].value == "See Details: finding 1, observed"
        parts = [row[4] for row in list(workbook["Details"].values)[1:]]
        assert json.loads("".join(parts)) == findings[0]["observed"]
    finally:
        workbook.close()


def test_partial_spool_requires_opt_in_and_stays_retained(tmp_path):
    spool = ResultSpool.create(tmp_path / "runs", "vlan_compliance", 2)
    with pytest.raises(ValueError):
        with spool.collection():
            spool.append(DeviceOutcome(1), payload={"findings": evaluate(), "errors": []})
    with pytest.raises(ValueError, match="Incomplete"):
        report.export_compliance_spool(spool.path, tmp_path / "partial.xlsx")
    report.export_compliance_spool(spool.path, tmp_path / "partial.xlsx", allow_partial=True)
    assert ResultSpool(spool.path).manifest["status"] == "partial_consumed"


def test_1500_devices_use_bounded_workers_and_complete_spool(tmp_path, monkeypatch):
    from threading import Barrier, Lock, get_ident

    events, closed = install_fakes(monkeypatch)
    connect = app.connect_device
    barrier, lock = Barrier(5), Lock()
    sessions = []
    active = peak = 0
    caller = get_ident()

    @contextmanager
    def overlapping(*args):
        nonlocal active, peak
        assert get_ident() != caller
        with lock:
            active += 1
            peak = max(peak, active)
        try:
            with connect(*args) as session:
                with lock:
                    sessions.append(session)
                    ordinal = len(sessions)
                if ordinal <= 5:
                    barrier.wait(timeout=10)
                yield session
        finally:
            with lock:
                active -= 1

    monkeypatch.setattr(app, "connect_device", overlapping)
    opts = options(tmp_path)
    opts["execution_config"] = ExecutionConfig(5, 0, 0)
    run = app.collect_compliance(targets(1500), None, **opts)
    assert peak == 5 and active == 0
    assert len(closed) == 1500
    assert len({id(s.cli) for s in sessions}) == len({id(s.context) for s in sessions}) == 1500
    assert len(events) == 1500 * 4
    assert run.finding_counts["compliant"] == 3000
    assert run.failed_devices == 0
    spool = ResultSpool(run.spool_path)
    assert spool.manifest["completed_count"] == 1500
    assert [record["input_position"] for record in spool.records()] == list(range(1, 1501))
    writer = report.write_tables
    def on_caller(*args, **kwargs):
        assert get_ident() == caller
        return writer(*args, **kwargs)
    monkeypatch.setattr(report, "write_tables", on_caller)
    path = report.export_compliance_spool(run.spool_path, tmp_path / "scale.xlsx")
    workbook = load_workbook(path, read_only=True)
    try:
        assert sum(1 for _ in workbook["Findings"].rows) == 3001
    finally:
        workbook.close()


def test_cli_run_and_recovery(tmp_path, monkeypatch):
    spec = spec_from_file_location("compliance_cli", Path(__file__).parents[1] / "scripts/device_vlan_compliance.py")
    cli = module_from_spec(spec)
    spec.loader.exec_module(cli)
    monkeypatch.setattr(cli, "load_targets", Mock(return_value=targets()))
    monkeypatch.setattr(cli, "load_execution_config", Mock(return_value=ExecutionConfig(1, 0, 0)))
    monkeypatch.setattr(cli, "run_compliance", Mock(return_value=tmp_path / "output.xlsx"))
    cli.main(["run", "devices.xlsx", "--proxy", "proxy", "--cluster", "cluster",
              "--bastion-host", "bastion", "--bastion-user", "operator", "--keep-spool"])
    kwargs = cli.run_compliance.call_args.kwargs
    assert kwargs["spool_root"] == Path("outputs/runs/vlan_compliance")
    assert kwargs["reports_dir"] == Path("outputs/reports/vlan_compliance")
    assert kwargs["cleanup"] is False
    monkeypatch.setattr(cli, "export_compliance_spool", Mock(return_value=tmp_path / "output.xlsx"))
    cli.main(["export", "retained-run", "report.xlsx", "--allow-partial"])
    assert cli.export_compliance_spool.call_args.kwargs["allow_partial"] is True
    cli.export_compliance_spool.side_effect = RuntimeError("password=never-output")
    with pytest.raises(SystemExit) as caught:
        cli.main(["export", "retained-run", "report.xlsx"])
    assert "never-output" not in str(caught.value)
