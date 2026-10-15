"""Milestone 1 signed-off fixtures, exclusively synthetic read-only evidence."""
from dataclasses import asdict
import pytest
from openpyxl import load_workbook
from test_bulk_configuration_v2 import NOW, IP, template, observed, inputs
from orbitflow.configuration.v2.planning import request, plan, BulkPlan, load, save
from orbitflow.configuration.v2.workbook import create_input, read_input, export_plan, EXTRA, consume_decisions
from orbitflow.configuration.v2.evidence import scoped_evidence
from orbitflow.configuration.plans import PlanError, canonical
from orbitflow.vendors.configuration_facts import observe_configuration
from orbitflow.vendors.cisco.vlans import parse_ios_running_config, parse_vlan_lookup
from orbitflow.compliance.resolvers import reconstruct_allowed

NAME = 'GigabitEthernet1/0/1'
FAMILIES = ['C3750X', 'C3850X']


def observation(family, config, lookups=None):
    t = template(family, 'Access Port')
    obs = observed(t)
    platform = obs['identity']['platform']
    nodes = observe_configuration(config, platform, preserve_full_evidence=True)
    interfaces, objects = parse_ios_running_config(config)
    obs.update(facts=[asdict(n) for n in nodes], vlans=[asdict(i) for i in interfaces],
               objects=[asdict(o) for o in objects])
    from orbitflow.models import DeviceContext, VlanState
    from orbitflow.compliance.resolvers import resolver_for, evidence
    profile = 'c3750x_switching' if family == 'C3750X' else 'c3850_switching'
    ctx = DeviceContext('one', IP, (IP,), 'example', 'Cisco', platform, family if family == 'C3750X' else 'C3850', family,
                        profile, (), 'EXAMPLE', '1', '', NOW, NOW)
    resolver = resolver_for(ctx)(ctx, VlanState('example', IP, platform, interfaces, objects, NOW, nodes), planning=True)
    resolver.resolve()
    obs.update(audit=list(resolver.rows.values()), evidence=evidence(nodes), problems=resolver.problems,
               methodologies=[])
    if lookups is not None:
        obs['vlan_lookups'] = lookups
    from orbitflow.compliance.vlan import safe_data
    from orbitflow.vendors.configuration_facts import _safe_source
    return safe_data(obs, lambda text: _safe_source(text, platform))


def result(family, function, config, params, lookups=None):
    t = template(family, function)
    r = request(t.id, dict(management_ip=IP, **params))
    obs = observation(family, config, lookups)
    return plan([r], {IP: obs}, now=NOW, decisions={r['row_id']: 'Approve'}).to_dict()['results'][0], r, obs


@pytest.mark.parametrize('family', FAMILIES)
@pytest.mark.parametrize('lookup', [None, {'status': 'error', 'name': '', 'operational_status': ''}, {'status': 'unknown', 'name': '', 'operational_status': ''}])
def test_range_lookup_fallback_rename_evidence(family, lookup):
    config = 'vlan 1035-1109\n!\n'
    row, r, obs = result(family, 'VLAN Database', config, dict(vlan_id=1040, name='Customer'), None if lookup is None else {'1040': lookup})
    assert row['commands'] == ['configure terminal', 'vlan 1040', 'name Customer', 'end']
    assert {'VLAN_LOOKUP_FALLBACK', 'VLAN_RENAME_IMPACT'} <= set(row['warnings'])
    assert not row['blockers']
    assert scoped_evidence(r, obs)[0]['excerpt'] == 'vlan 1035-1109'


@pytest.mark.parametrize('family', FAMILIES)
def test_lookup_inactive_noop_absence_and_unknown(family):
    lookup = {'445': dict(status='present', name='Customer', operational_status='suspended')}
    row, _, _ = result(family, 'VLAN Database', 'hostname example\n', dict(vlan_id=445, name='Customer'), lookup)
    assert row['no_op'] and not row['commands'] and 'VLAN_INACTIVE' in row['warnings']
    row, _, _ = result(family, 'VLAN Database', 'hostname example\n', dict(vlan_id=445, name='Customer'), {'445': dict(status='absent', name='', operational_status='')})
    assert not row['no_op'] and not row['blockers']
    t = template(family, 'VLAN Database')
    r = request(t.id, dict(management_ip=IP, vlan_id=445, name='Customer'))
    row = plan([r], {IP: {'status': 'evidence_unavailable'}}, now=NOW).to_dict()['results'][0]
    assert 'VLAN_STATE_UNKNOWN' in row['blockers']
    with pytest.raises(PlanError):
        request(t.id, dict(management_ip=IP, vlan_id=445, name='Customer Site'))


@pytest.mark.parametrize('output,status', [
    ('VLAN Name Status Ports\n---- -------- -------- -----\n445 PPPoE active Te1/0/1,\n                         Te1/0/2, Gi1/0/3\n', 'present'),
    ('% VLAN id 445 not found in current VLAN database', 'absent'),
    ('VLAN 445 does not exist', 'absent'),
    ('% Invalid input detected at marker', 'error'),
    ('% Error fetching VLAN', 'error'), ('garbage', 'unknown'),
])
def test_tabular_lookup(output, status):
    lookup = parse_vlan_lookup(output, 445)
    assert lookup['status'] == status
    if status == 'present':
        assert lookup['name'] == 'PPPoE' and lookup['operational_status'] == 'active'


@pytest.mark.parametrize('family', FAMILIES)
@pytest.mark.parametrize('mode,warning,blocker', [
    ('switchport mode access\n switchport trunk allowed vlan 130', 'RESIDUAL_TRUNK_CONFIGURATION', None),
    ('switchport mode dynamic auto', 'UNSPECIFIED_OR_DYNAMIC_ACCESS_MODE', None),
    ('switchport mode dynamic desirable', 'UNSPECIFIED_OR_DYNAMIC_ACCESS_MODE', None),
    ('', 'UNSPECIFIED_OR_DYNAMIC_ACCESS_MODE', None),
    ('switchport mode trunk', None, 'ACCESS_CONVERSION_BLOCKED'),
    ('no switchport', None, 'ACCESS_CONVERSION_BLOCKED'),
])
def test_access_modes_full_cli_shutdown(family, mode, warning, blocker):
    config = f'interface {NAME}\n {mode}\n description Before\n switchport access vlan 130\n shutdown\n spanning-tree portfast\n authentication password DO_NOT_KEEP\n!\n'
    row, r, obs = result(family, 'Access Port', config, dict(interface=NAME, vlan_id=445, description='After'))
    assert row['commands'] == ['configure terminal', f'interface {NAME}', 'description After', 'switchport', 'switchport mode access', 'switchport access vlan 445', 'end']
    assert 'GLOBAL_VLAN_REQUIRED' in row['warnings']
    assert 'no shutdown' not in row['commands']
    if warning:
        assert warning in row['warnings']
    if blocker:
        assert blocker in row['blockers'] and not row['deployment_ready']
    excerpts = '\n'.join(e['excerpt'] for e in scoped_evidence(r, obs))
    assert ' shutdown' in excerpts and ' spanning-tree portfast' in excerpts and excerpts.endswith('!')
    assert 'DO_NOT_KEEP' not in excerpts


@pytest.mark.parametrize('family', FAMILIES)
def test_access_exact_noop_requires_explicit_vlan(family):
    for vlan, noop in [(' switchport access vlan 445\n', True), ('', False)]:
        row, _, _ = result(family, 'Access Port', f'interface {NAME}\n switchport mode access\n{vlan}!\n', dict(interface=NAME, vlan_id=445))
        assert row['no_op'] is noop
        assert bool(row['commands']) != noop


@pytest.mark.parametrize('family', FAMILIES)
def test_lag_access_block_trunk_redirect_description_exact(family):
    config = f'interface {NAME}\n channel-group 1 mode active\n description Physical\n!\ninterface Port-channel1\n switchport mode trunk\n switchport trunk encapsulation dot1q\n switchport trunk allowed vlan 130-135\n description Aggregate\n!\n'
    row, r, obs = result(family, 'Access Port', config, dict(interface=NAME, vlan_id=445))
    assert 'PHYSICAL_PORT_CHANNEL_MEMBER' in row['blockers'] and f'interface {NAME}' in row['commands']
    assert 'interface Port-channel1' in '\n'.join(e['excerpt'] for e in scoped_evidence(r, obs))
    row, _, _ = result(family, 'Trunk Port', config, dict(interface=NAME, action='Add', vlans=[130, 445]))
    assert row['resolved_target'] == 'Port-channel1'
    assert 'interface Port-channel1' in row['commands'] and 'switchport trunk allowed vlan add 130,445' in row['commands']
    row, _, _ = result(family, 'Interface Description', config, dict(interface=NAME, description='After'))
    assert row['commands'] == ['configure terminal', f'interface {NAME}', 'description After', 'end']
    assert not row['blockers']
    row, _, _ = result(family, 'Trunk Port', config.split('interface Port-channel1')[0], dict(interface=NAME, action='Replace', vlans=[445]))
    assert 'PORT_CHANNEL_OWNER_UNRESOLVED' in row['blockers']
    t = template(family, 'Trunk Port')
    a = request(t.id, dict(management_ip=IP, interface=NAME, action='Replace', vlans=[445]))
    b = request(t.id, dict(management_ip=IP, interface='Port-channel1', action='Replace', vlans=[545]), row=4)
    rows = plan([a,b], {IP: obs}, now=NOW).to_dict()['results']
    assert all('ORDERED_RESOURCE_OVERLAP' in row['findings'] and row['conflicting_rows'] for row in rows)


@pytest.mark.parametrize('family', FAMILIES)
@pytest.mark.parametrize('action,baseline,vlans,noop,blocker', [
    ('Add', ' switchport trunk allowed vlan 130-135\n switchport trunk allowed vlan remove 133\n switchport trunk allowed vlan add 445\n', [130,445], True, None),
    ('Add', ' switchport trunk allowed vlan 130-135\n switchport trunk allowed vlan remove 133\n', [130,133], False, None),
    ('Add', '', [445], True, 'ADD_REQUIRES_EXPLICIT_BASELINE'),
    ('Remove', '', [445], False, None),
    ('Remove', ' switchport trunk allowed vlan remove 445\n', [445], True, None),
    ('Remove', ' switchport trunk allowed vlan none\n', [445], True, None),
    ('Replace', '', [445], False, None),
])
def test_trunk_ordered_all_actions(family, action, baseline, vlans, noop, blocker):
    config = f'interface {NAME}\n switchport mode trunk\n switchport trunk encapsulation dot1q\n{baseline}!\n'
    row, _, _ = result(family, 'Trunk Port', config, dict(interface=NAME, action=action, vlans=vlans))
    assert row['no_op'] is noop
    if blocker:
        assert blocker in row['blockers']
    else:
        assert not row['blockers']
    if not noop:
        fragment = '' if action == 'Replace' else action.lower()+' '
        assert 'switchport trunk allowed vlan '+fragment+','.join(map(str,vlans)) in row['commands']
    if action in {'Replace','Remove'} and not noop:
        assert 'SERVICE_IMPACT_REVIEW' in row['warnings']


@pytest.mark.parametrize('family', FAMILIES)
@pytest.mark.parametrize('mode,blocked', [('', False), ('switchport mode access', True), ('no switchport', True)])
def test_trunk_conversion_encapsulation(family, mode, blocked):
    row, _, _ = result(family, 'Trunk Port', f'interface {NAME}\n {mode}\n shutdown\n!\n', dict(interface=NAME, action='Replace', vlans=[445]))
    assert ('TRUNK_CONVERSION_BLOCKED' in row['blockers']) is blocked
    assert ('switchport trunk encapsulation dot1q' in row['commands']) == (family == 'C3750X')
    assert 'no shutdown' not in row['commands']


@pytest.mark.parametrize('family', FAMILIES)
@pytest.mark.parametrize('target', [NAME, 'Port-channel1', 'Vlan1035', NAME+'.100'])
def test_description_exact_target_noop_unrelated(family, target):
    config = f'interface {target}\n switchport mode access\n switchport access vlan 445\n switchport trunk allowed vlan 545\n shutdown\n!\n'
    row, _, _ = result(family, 'Interface Description', config, dict(interface=target, description='After'))
    assert row['commands'] == ['configure terminal', f'interface {target}', 'description After', 'end']
    assert not row['blockers']
    row, _, _ = result(family, 'Interface Description', config.replace(' shutdown', ' description After\n shutdown'), dict(interface=target, description='After'))
    assert row['no_op'] and not row['commands']


def test_symbolic_all_operations():
    assert reconstruct_allowed([('allowed_remove', [445]), ('allowed_add', [445,545])]) == dict(all=True, values=[], explicit=True)
    assert reconstruct_allowed([('allowed_except', [130,131]), ('allowed_remove', [132])]) == dict(all=True, values=[130,131,132], explicit=True)


@pytest.mark.parametrize('family', FAMILIES)
def test_invalid_rows_retained_fields_safe_continuation(family, tmp_path):
    t = template(family, 'Access Port')
    source = tmp_path/'input.xlsx'
    create_input(source)
    book = load_workbook(source)
    sheet = book[t.sheet]
    for n, values in enumerate([(IP, NAME, 445, None), (IP, 'Gi1/0/1', 1002, 'password DO_NOT_KEEP'), (IP, NAME, 545, None)], 3):
        for col,value in enumerate(values,1):
            sheet.cell(n,col,value)
    book.save(source)
    rows = read_input(source)
    assert len(rows) == 2 and len(rows.errors) == 1
    assert {'interface', 'vlan_id', 'description'} <= {s.split(':')[0] for s in rows.errors[0]['details']}
    original = plan(rows, {IP: observation(family, f'interface {NAME}\n!\n')}, now=NOW)
    dest = tmp_path/'plan.xlsx'
    export_plan(original,dest)
    book = load_workbook(dest)
    sheet = book[t.sheet]
    extra = lambda label,row: sheet.cell(row,len(t.columns)+EXTRA.index(label)+1).value
    assert sheet.cell(4,2).value == 'Gi1/0/1' and sheet.cell(4,3).value == '1002'
    assert extra('Validation Status',4) == 'Invalid' and extra('Proposed CLI',4) is None
    assert 'row 5' in extra('Error Details',3) and 'row 3' in extra('Error Details',5)
    assert 'DO_NOT_KEEP' not in canonical(original.to_dict())
    save(original,tmp_path/'plan.json')
    assert load(tmp_path/'plan.json').digest == original.digest
    assert consume_decisions(original,dest,expected_digest=original.digest,now=NOW).digest == original.digest
    with pytest.raises(PlanError):
        export_plan(original,dest)

@pytest.mark.parametrize('family', FAMILIES)
def test_optional_trunk_changes_do_not_repeat_membership(family):
    config = f'interface {NAME}\n switchport mode trunk\n switchport trunk encapsulation dot1q\n switchport trunk allowed vlan 445\n description Before\n!\n'
    row, _, _ = result(family, 'Trunk Port', config, dict(interface=NAME, action='Add', vlans=[445], description='After', native_vlan=545))
    assert row['commands'] == ['configure terminal', f'interface {NAME}', 'description After', 'switchport trunk native vlan 545', 'end']


@pytest.mark.parametrize('family', FAMILIES)
def test_malformed_operations_and_lag_conflicts_block(family):
    config = f'interface {NAME}\n switchport mode trunk\n switchport trunk allowed vlan frobnicate 445\n!\n'
    row, _, _ = result(family, 'Trunk Port', config, dict(interface=NAME, action='Replace', vlans=[445]))
    assert 'UNSUPPORTED_ALLOWED_VLAN_OPERATION' in row['blockers'] and row['commands']
    config = f'interface {NAME}\n channel-group 1 mode active\n switchport mode access\n switchport access vlan 445\n!\ninterface Port-channel1\n switchport mode trunk\n switchport trunk encapsulation dot1q\n switchport trunk allowed vlan 545\n!\n'
    row, _, _ = result(family, 'Trunk Port', config, dict(interface=NAME, action='Replace', vlans=[445]))
    assert 'PORT_CHANNEL_CONFIG_CONFLICT' in row['blockers']


def test_all_invalid_workbook_exports_without_observation(tmp_path):
    t = template('C3750X', 'VLAN Database')
    source = tmp_path/'input.xlsx'
    create_input(source)
    book = load_workbook(source)
    book[t.sheet].append([IP, 1002, 'Invalid Name'])
    book.save(source)
    rows = read_input(source)
    assert not rows and rows.errors
    original = plan(rows, now=NOW)
    export_plan(original, tmp_path/'plan.xlsx')
    assert not original.to_dict()['observations']


def test_lookup_service_borrows_cli_and_uses_only_show(monkeypatch):
    from orbitflow.capabilities.vlans import VlanService, VlanCapabilityError
    from orbitflow.models import DeviceContext
    ctx = DeviceContext('one', IP, (IP,), 'example', 'Cisco', 'cisco_xe', 'C3850', 'C3850',
                        'c3850_switching', (), 'EXAMPLE', '1', '', NOW, NOW)
    session = object()
    calls = []
    class CLI:
        def require_session(self, supplied):
            assert supplied is session
        def configure(self, **kwargs):
            assert kwargs['paging_command'] == 'terminal length 0'
        def run_command(self, command, **kwargs):
            calls.append(command)
            assert command == 'show vlan id 445'
            return 'VLAN Name Status Ports\n445 Customer active Gi1/0/1,\n                         Gi1/0/2'
    assert VlanService().lookup(session, ctx, 445, cli=CLI())['name'] == 'Customer'
    assert calls == ['show vlan id 445']
    with pytest.raises(VlanCapabilityError):
        VlanService().lookup(session, ctx, '445', cli=CLI())
    for stage in ('setup', 'command'):
        class FailedCLI(CLI):
            def configure(self, **kwargs):
                if stage == 'setup':
                    raise RuntimeError('password MUST_NOT_PERSIST')
                super().configure(**kwargs)
            def run_command(self, command, **kwargs):
                raise RuntimeError('password MUST_NOT_PERSIST')
        lookup = VlanService().lookup(session, ctx, 445, cli=FailedCLI())
        assert lookup == dict(status='error', name='', operational_status='')

@pytest.mark.parametrize('family', FAMILIES)
def test_explicit_unreliable_configuration_cannot_fallback(family):
    t = template(family, 'VLAN Database')
    r = request(t.id, dict(management_ip=IP, vlan_id=445, name='Customer'))
    obs = observation(family, 'vlan 445\n name Customer\n!\n', {'445': dict(status='error', name='', operational_status='')})
    obs['configuration_reliable'] = False
    row = plan([r], {IP: obs}, now=NOW).to_dict()['results'][0]
    assert 'VLAN_STATE_UNKNOWN' in row['blockers'] and not row['no_op']

@pytest.mark.parametrize('family', ['C3750X', 'C3850X'])
@pytest.mark.parametrize('number', [1, 2, 3, 4])
def test_ordered_access_description_pairs(family, number, tmp_path):
    access, desc = template(family, 'Access Port'), template(family, 'Interface Description')
    interface = f'GigabitEthernet1/0/{number}'
    config = f'vlan 445\n name Customer\n!\ninterface {interface}\n description Original\n switchport mode access\n switchport access vlan 445\n shutdown\n!'
    obs = observation(family, config)
    a = request(access.id, dict(management_ip=IP, interface=interface, vlan_id=445, description='Access text'), sheet_order=1, row=number+2)
    b = request(desc.id, dict(management_ip=IP, interface=interface, description='Final text'), sheet_order=2, row=number+6)
    original = plan([b, a], {IP: obs}, now=NOW)
    rows = original.to_dict()['results']
    assert 'description Access text' in rows[0]['commands']
    assert 'description Final text' in rows[1]['commands']
    for row in rows:
        assert row['consistency'] == 'consistent' and not row['blockers']
        assert 'ORDERED_RESOURCE_OVERLAP_WARNING' in row['warnings']
        detail = row['overlaps'][0]
        assert detail['earlier'] == dict(sheet=access.sheet, row=number+2)
        assert detail['later'] == dict(sheet=desc.sheet, row=number+6)
        assert detail['effective_result']['description'] == 'Final text'
    assert original.digest == plan([a, b], {IP: obs}, now=NOW).digest
    save(original, tmp_path/'plan.json')
    assert load(tmp_path/'plan.json').digest == original.digest
    export_plan(original, tmp_path/'plan.xlsx')
    book = load_workbook(tmp_path/'plan.xlsx')
    for t, r in ((access, a), (desc, b)):
        sheet = book[t.sheet]
        assert 'Ordered Overlap Details' not in [c.value for c in sheet[2]]
        assert 'ORDERED_RESOURCE_OVERLAP_WARNING' in sheet.cell(r['row'], len(t.columns) + EXTRA.index('Warnings') + 1).value
        assert 'shutdown' in sheet.cell(r['row'], len(t.columns)+EXTRA.index('Existing Evidence')+1).value
    assert consume_decisions(original, tmp_path/'plan.xlsx', expected_digest=original.digest, now=NOW).digest == original.digest


@pytest.mark.parametrize('family', ['C3750X', 'C3850X'])
@pytest.mark.parametrize('actions', [('Add', 'Add'), ('Replace', 'Remove')])
def test_ordered_trunk_projection(family, actions):
    t = template(family, 'Trunk Port')
    obs = observation(family, 'interface GigabitEthernet1/0/1\n switchport mode trunk\n switchport trunk allowed vlan 100-102\n!')
    a = request(t.id, dict(management_ip=IP, interface='GigabitEthernet1/0/1', action=actions[0], vlans=[200,201]), row=3)
    b = request(t.id, dict(management_ip=IP, interface='GigabitEthernet1/0/1', action=actions[1], vlans=[201,202] if actions[1]=='Add' else [201]), row=4)
    rows = plan([a,b], {IP:obs}, now=NOW).to_dict()['results']
    assert rows[1]['allowed_state'] == rows[0]['desired_allowed_state']
    assert rows[1]['desired_allowed_state']['values'] == ([100,101,102,200,201,202] if actions[1]=='Add' else [200])
    assert all(r['consistency']=='consistent' for r in rows)
    assert all('ORDERED_RESOURCE_OVERLAP_WARNING' in r['warnings'] for r in rows)
    if actions[1]=='Remove':
        assert all('SERVICE_IMPACT_REVIEW' in r['warnings'] for r in rows)


@pytest.mark.parametrize('family', ['C3750X', 'C3850X'])
@pytest.mark.parametrize('case', ['missing', 'baseline', 'conversion', 'description'])
def test_overlap_safety_blockers(family, case, tmp_path):
    function = 'Interface Description' if case=='description' else 'Access Port' if case=='conversion' else 'Trunk Port'
    t = template(family, function)
    p = inputs(t)
    if function=='Trunk Port': p.update(vlans=[445])
    obs = observation(family, 'interface GigabitEthernet1/0/1\n switchport mode trunk\n!')
    if case=='missing': obs = dict(status='evidence_unavailable')
    a = request(t.id,p,row=3)
    u = template(family,'Trunk Port') if case=='conversion' else t
    v = inputs(u)
    if case=='description': v['description']='Different description'
    b = request(u.id,v,row=4)
    original = plan([a,b], {IP:obs}, now=NOW)
    rows = original.to_dict()['results']
    assert all('ORDERED_RESOURCE_OVERLAP' in r['blockers'] for r in rows)
    assert all(r['consistency']=='unresolved' for r in rows)
    if case=='baseline':
        assert rows[0]['no_op']
        assert 'ADD_REQUIRES_EXPLICIT_BASELINE' in rows[0]['blockers']
        assert rows[0]['commands']==[]
        export_plan(original, tmp_path/'blocked.xlsx')
        book = load_workbook(tmp_path/'blocked.xlsx')
        assert 'does not waive' in book[t.sheet].cell(3,len(t.columns)+EXTRA.index('Error Details')+1).value

@pytest.mark.parametrize('family', FAMILIES)
@pytest.mark.parametrize('baseline', [' switchport trunk allowed vlan 100-102\n', ''])
def test_sequential_port_channel_and_all(family, baseline, tmp_path):
    t = template(family, 'Trunk Port')
    obs = observation(family, f'interface {NAME}\n channel-group 1 mode active\n!\ninterface Port-channel1\n switchport mode trunk\n{baseline}!\n')
    a = request(t.id, dict(management_ip=IP, interface=NAME, action='Replace', vlans=[200,201]), row=3)
    b = request(t.id, dict(management_ip=IP, interface='Port-channel1', action='Remove', vlans=[201]), row=4)
    original = plan([a,b], {IP:obs}, now=NOW)
    rows = original.to_dict()['results']
    assert all(r['resolved_target']=='Port-channel1' and not r['blockers'] for r in rows)
    assert rows[1]['desired_allowed_state']['values']==[200]
    assert rows[1]['observed_allowed_state']['all'] == (not baseline)
    assert not rows[1]['allowed_state']['all']
    assert 'SERVICE_IMPACT_REVIEW' in rows[1]['warnings']
    export_plan(original,tmp_path/'plan.xlsx')
    book=load_workbook(tmp_path/'plan.xlsx')
    assert 'ORDERED_RESOURCE_OVERLAP_WARNING' in book[t.sheet].cell(3,len(t.columns)+EXTRA.index('Warnings')+1).value
    rejected=consume_decisions(original,tmp_path/'plan.xlsx',expected_digest=original.digest,now=NOW)
    assert rejected.digest==original.digest
    # An absent logical owner remains a blocker on both requests.
    bad=observation(family,f'interface {NAME}\n channel-group 1 mode active\n!\n')
    failed=plan([a,b],{IP:bad},now=NOW).to_dict()['results']
    assert 'PORT_CHANNEL_OWNER_UNRESOLVED' in failed[0]['blockers']
    assert all(r['blockers'] for r in failed)


@pytest.mark.parametrize('family', FAMILIES)
def test_projected_description_noop_and_stale_evidence(family):
    from datetime import timedelta
    from orbitflow.configuration.v2.planning import evaluate_decisions
    access,desc=template(family,'Access Port'),template(family,'Interface Description')
    obs=observation(family,f'interface {NAME}\n switchport mode access\n switchport access vlan 445\n description Original\n!')
    a=request(access.id,dict(management_ip=IP,interface=NAME,vlan_id=445,description='Final'),sheet_order=1)
    b=request(desc.id,dict(management_ip=IP,interface=NAME,description='Final'),sheet_order=2)
    original=plan([a,b],{IP:obs},now=NOW)
    assert original.to_dict()['results'][1]['no_op']
    assert original.to_dict()['results'][1]['commands']==[]
    stale=plan([a,b],{IP:obs},now=NOW+timedelta(hours=2)).to_dict()['results']
    assert all('STALE_EVIDENCE' in r['blockers'] for r in stale)
    assert all('ORDERED_RESOURCE_OVERLAP_WARNING' in r['warnings'] for r in stale)
    rejected=evaluate_decisions(original,{a['row_id']:'Reject'},expected_digest=original.digest,now=NOW)
    assert not rejected.to_dict()['results'][1]['no_op']
    assert rejected.digest==original.digest


@pytest.mark.parametrize('decision', ['Pending', 'Approve', 'Reject'])
def test_owner_pending_vlan_access_description_and_trunk(tmp_path, decision):
    db = template('C3750X', 'VLAN Database')
    access = template('C3750X', 'Access Port')
    desc = template('C3750X', 'Interface Description')
    trunk = template('C3750X', 'Trunk Port')
    names = ['GigabitEthernet1/1/1', 'GigabitEthernet1/1/2', 'GigabitEthernet1/1/3']
    obs = observation('C3750X', ''.join(f'interface {n}\n switchport mode access\n!\n' for n in names + ['GigabitEthernet1/0/13']))
    prerequisites = [request(db.id, dict(management_ip=IP, vlan_id=v, name=f'Customer{v}'), sheet_order=1, row=i+3) for i,v in enumerate([2432,2433,2434])]
    ports = [request(access.id, dict(management_ip=IP, interface=n, vlan_id=v), sheet_order=2, row=i+3) for i,(n,v) in enumerate(zip(names,[2432,2433,2434]))]
    conversion = request(trunk.id, dict(management_ip=IP, interface='GigabitEthernet1/0/13', action='Add', vlans=[2432]), sheet_order=3)
    description = request(desc.id, dict(management_ip=IP, interface=names[0], description='Customer'), sheet_order=4)
    original = plan(prerequisites + ports + [conversion, description], {IP:obs}, now=NOW)
    rows = original.to_dict()['results']
    for prerequisite, row in zip(prerequisites, rows[3:6]):
        assert row['dependencies'] == [prerequisite['row_id']]
        assert 'DEPENDENCY_NOT_APPROVED' in row['warnings']
        assert not row['blockers'] and row['consistency'] == 'consistent'
        assert not row['deployment_ready']
    assert 'ORDERED_RESOURCE_OVERLAP_WARNING' in rows[3]['warnings']
    assert {'ADD_REQUIRES_EXPLICIT_BASELINE', 'TRUNK_CONVERSION_BLOCKED'} <= set(rows[6]['blockers'])
    save(original, tmp_path/'plan.json')
    assert load(tmp_path/'plan.json').digest == original.digest
    export_plan(original, tmp_path/'plan.xlsx')
    book = load_workbook(tmp_path/'plan.xlsx')
    for r in prerequisites + ports + [description]:
        t = template('C3750X', 'VLAN Database') if r in prerequisites else access if r in ports else desc
        book[t.sheet].cell(r['row'], len(t.columns)+len(EXTRA), decision if r in prerequisites else 'Approve')
    book.save(tmp_path/'plan.xlsx')
    reviewed = consume_decisions(original, tmp_path/'plan.xlsx', expected_digest=original.digest, now=NOW)
    assert reviewed.digest == original.digest
    evaluated = reviewed.to_dict()['results']
    assert all(r['deployment_ready'] == (decision == 'Approve') for r in evaluated[3:6])
    assert all(not r['execution_authorized'] for r in evaluated)


@pytest.mark.parametrize('family', FAMILIES)
@pytest.mark.parametrize('description_first', [False, True])
def test_access_description_cli_follows_both_workbook_orders(family, description_first):
    access, desc = template(family, 'Access Port'), template(family, 'Interface Description')
    operation = (access, dict(management_ip=IP, interface=NAME, vlan_id=445, description='Port description'))
    description = (desc, dict(management_ip=IP, interface=NAME, description='Separate description'))
    ordered = [description, operation] if description_first else [operation, description]
    requests = [request(t.id, p, sheet_order=i+1) for i,(t,p) in enumerate(ordered)]
    obs = observation(family, f'vlan 445\n name Customer\n!\ninterface {NAME}\n switchport mode access\n switchport access vlan 445\n description Original\n!')
    rows = plan(requests, {IP:obs}, now=NOW, decisions={r['row_id']:'Approve' for r in requests}).to_dict()['results']
    assert all(not r['blockers'] and r['deployment_ready'] for r in rows)
    assert all('ORDERED_RESOURCE_OVERLAP_WARNING' in r['warnings'] for r in rows)
    for row, (_,p) in zip(rows,ordered):
        assert 'description ' + p['description'] in row['commands']
    assert rows[-1]['projected_state']['description'] == ordered[-1][1]['description']
