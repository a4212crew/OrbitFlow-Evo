from dataclasses import replace
import json
import pytest
from openpyxl import load_workbook
from orbitflow.compliance.methodology import resolve_methodologies
from orbitflow.compliance import evaluate_vlan_compliance
from orbitflow.models import VlanState
from orbitflow.vendors.configuration_facts import observe_configuration
from orbitflow.execution import DeviceOutcome
from orbitflow.result_spool import ResultSpool
from orbitflow.methodology_report import export_methodology_spool
from test_vlan_configuration_audit import FAMILIES
from test_vlan_compliance import context, interface, NOW, POLICY


def snapshot(config, family):
    ctx = replace(context(), platform=FAMILIES[family], device_family=family)
    state = VlanState('switch', ctx.management_ip, ctx.platform, (), (), NOW,
                      observe_configuration(config, ctx.platform, source_filename='fixture.cfg'))
    return ctx, state


def resolve(config, family='C3850', observed=()):
    ctx, state = snapshot(config, family)
    return resolve_methodologies(ctx, [interface(n) for n in observed], state)


@pytest.mark.parametrize('family,config,method', [
    ('C3850', 'interface Gi0/1\n switchport mode access\n switchport access vlan 445', 'M01'),
    ('ASR920', 'interface Gi0/1\n service instance 7 ethernet\n  encapsulation dot1q 445\n!\nbridge-domain 445\n member Gi0/1 service-instance 7', 'M02'),
    ('ME3600X', 'interface Gi0/1\n service instance 7 ethernet\n  encapsulation dot1q 445\n  bridge-domain 445', 'M03'),
    ('NCS540', 'interface TenGigE0/0/0/1\n!\ninterface TenGigE0/0/0/1.999 l2transport\n encapsulation dot1q 445 second-dot1q 17\n!\nl2vpn\n bridge group G\n  bridge-domain B\n   interface TenGigE0/0/0/1.999', 'M04'),
    ('EdgeSwitch', 'vlan database\n vlan 445\nexit\ninterface 0/1\n vlan participation include 445\n vlan pvid 445\nexit', 'M05'),
    ('NE05E', 'interface GE0/1\n port link-type trunk\n port trunk allow-pass vlan 445', 'M06'),
    ('NE05E', 'vsi svc\n#\ninterface GE0/1\n#\ninterface GE0/1.999\n dot1q termination vid 445\n l2 binding vsi svc', 'M07'),
])
def test_all_methods_and_compliance_nonmutation(family, config, method):
    ctx, state = snapshot(config, family)
    before = evaluate_vlan_compliance(ctx, [], state, POLICY)
    records = resolve_methodologies(ctx, [], state)
    assert any(method in r['methodology'] for r in records)
    assert evaluate_vlan_compliance(ctx, [], state, POLICY) == before
    for r in records:
        assert r['configuration_facts_digest']
        for item in r['evidence']:
            assert item['excerpt'] == config.splitlines()[item['line'] - 1]


@pytest.mark.parametrize('global_bd,status', [('445', 'mixed_equivalent'), ('545', 'conflict')])
def test_mixed_service_bindings(global_bd, status):
    config = ('interface Gi0/1\n service instance 7 ethernet\n  encapsulation dot1q 445 second-dot1q 17\n'
              '  bridge-domain 445\n service instance 8 ethernet\n  encapsulation untagged\n  bridge-domain 545\n'
              '!\nbridge-domain ' + global_bd + '\n member Gi0/1 service-instance 7')
    rows = resolve(config, 'ASR920')
    one, two = rows
    assert one['methodology'] == ['M02', 'M03']
    assert one['status'] == status and one['review_needed']
    assert one['mapping']['inner_vlan'] == [17]
    assert one['mapping']['outer_vlan'] == [445]
    assert two['methodology'] == ['M03']
    assert two['service_instance_id'] == '8'


def test_child_identity_qinq_and_no_inference():
    config = ('interface TenGigE0/0/0/1\n!\ninterface TenGigE0/0/0/1.999 l2transport\n'
              ' encapsulation dot1q 445 second-dot1q 17\n!\nl2vpn\n bridge group G\n'
              '  bridge-domain 4001\n   interface TenGigE0/0/0/1.999')
    rows = resolve(config, 'NCS540', ('TenGigE0/0/0/1.999',))
    child = next(r for r in rows if r['methodology'] == ['M04'])
    assert child['parent_interface'] == 'TenGigE0/0/0/1'
    assert child['interface_match_status'] == 'matched'
    assert child['mapping']['inner_vlan'] == [17]
    assert child['mapping']['audit_vlan'] == [445]
    assert child['mapping']['bridge_domains'] == [['G', '4001']]
    assert len(rows) == 1


def test_unknown_partial_secret_and_missing_reference():
    config = ('interface Gi0/1\n switchport mode trunk\n encapsulation mystery token topsecret\n'
              '!\nbridge-domain 445\n member Gi0/99 service-instance 7')
    rows = resolve(config, 'ME3600X')
    assert rows[0]['review_needed']
    assert 'UNSUPPORTED_FORWARDING_SYNTAX' in json.dumps(rows)
    assert 'topsecret' not in json.dumps(rows)
    assert not any(r.get('config_interface_name') == 'Gi0/99' for r in rows)
    assert any(r['record_kind'] == 'reference_exception' for r in rows)
    assert resolve('interface Gi0/1\n service instance 1 ethernet', 'ASR920')[0]['status'] != 'resolved'
    assert resolve('interface Gi0/1', 'C3850')[0]['methodology'] == []


def test_unknown_marker_does_not_change_compliance():
    base = 'interface Gi0/1\n switchport mode trunk\n'
    ctx, clean = snapshot(base, 'C3850')
    _, unknown = snapshot(base + ' encapsulation unsupported secret value', 'C3850')
    assert evaluate_vlan_compliance(ctx, [], clean, POLICY) == evaluate_vlan_compliance(ctx, [], unknown, POLICY)


def test_workbook_recovery_literal_cells_legacy_and_long_evidence(tmp_path):
    rows = resolve('interface Gi0/1\n description =1+1\n switchport mode access\n switchport access vlan 445')
    rows[0]['mapping']['long'] = 'x' * 40000
    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 2)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={'methodologies': rows, 'findings': [], 'errors': []})
        spool.append(DeviceOutcome(2), payload={'findings': [], 'errors': []}, target='192.0.2.2')
    path = export_methodology_spool(spool.path, tmp_path / 'review.xlsx')
    assert spool.path.exists()
    workbook = load_workbook(path)
    assert workbook.sheetnames == ['Methodology Resolution', 'Run Errors', 'Details']
    sheet = workbook.worksheets[0]
    assert sheet.freeze_panes == 'A2' and sheet.auto_filter.ref
    headers = [c.value for c in sheet[1]]
    result = dict(zip(headers, [c.value for c in sheet[2]]))
    assert result['Methodology'] == 'M01' and result['Mapping'] == 'See Details'
    assert sheet.cell(3, headers.index('Status') + 1).value == 'unable_to_assess'
    for tab in workbook:
        assert all(c.data_type != 'f' for row in tab for c in row)
    first = [[list(row) for row in tab.values] for tab in workbook]
    workbook.close()
    export_methodology_spool(spool.path, path)
    second = load_workbook(path)
    assert first == [[list(row) for row in tab.values] for tab in second]
    second.close()

@pytest.mark.parametrize('family,config,method,code', [
    ('NE05E', 'vsi a\n#\nvsi b\n#\ninterface GE0/1\n#\ninterface GE0/1.123\n dot1q termination vid 445\n l2 binding vsi a\n l2 binding vsi b', 'M07', 'CONFLICTING_VSI_BINDING'),
    ('NCS540', 'interface TenGigE0/0/0/1.123 l2transport\n encapsulation untagged', 'M04', 'UNBOUND_L2TRANSPORT_SUBINTERFACE'),
    ('C3850', 'interface Gi0/1\n switchport trunk allowed vlan 445', 'M01', 'UNRESOLVED_SWITCHPORT_MODE'),
    ('EdgeSwitch', 'interface 0/1\n vlan pvid 445\nexit', 'M05', 'PVID_NOT_IN_MEMBERSHIP'),
])
def test_source_backed_review_exceptions(family, config, method, code):
    record = next(r for r in resolve(config, family) if method in r['methodology'])
    assert record['review_needed'] and record['status'] != 'resolved'
    assert code in json.dumps(record['configuration_findings'])
    assert record['evidence']


def test_unknown_family_and_missing_snapshot():
    ctx, state = snapshot('interface Gi0/1\n switchport mode trunk', 'C3850')
    for context_, state_ in [(replace(ctx, device_family='unknown'), state), (ctx, None)]:
        record = resolve_methodologies(context_, [], state_)[0]
        assert record['methodology'] == [] and record['status'] == 'unable_to_assess'


def test_mixed_switching_and_service_keeps_membership_separate():
    records = resolve('interface Gi0/1\n switchport mode trunk\n switchport trunk allowed vlan 445\n'
                      ' service instance 7 ethernet\n  encapsulation dot1q 17\n  bridge-domain 545', 'ME3600X')
    conventional = next(r for r in records if r['methodology'] == ['M01'])
    service = next(r for r in records if r['methodology'] == ['M03'])
    assert conventional['review_needed']
    assert 'tagged' not in conventional['mapping']
    assert conventional['mapping']['configured_switching'][1]['value'] == [445]
    assert service['mapping']['outer_vlan'] == [17]
    assert service['mapping']['resolved_bridge_domain'] == ['545']


def test_sanitized_workbook_and_supplied_secrets(tmp_path):
    ctx, state = snapshot('interface Gi0/1\n description password hiddenvalue\n switchport mode access\n switchport access vlan 445', 'C3850')
    records = resolve_methodologies(ctx, [], state, clean=lambda s: s.replace('suppliedvalue', '[REDACTED]'))
    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={'methodologies': records, 'errors': []})
    path = export_methodology_spool(spool.path, tmp_path / 'safe.xlsx')
    import zipfile
    with zipfile.ZipFile(path) as archive:
        content = b''.join(archive.read(n) for n in archive.namelist())
    assert b'hiddenvalue' not in content
    assert b'password [REDACTED]' in content


def test_run_exports_both_reports_before_cleanup(tmp_path, monkeypatch):
    from orbitflow import vlan_compliance as app
    from types import SimpleNamespace
    records = resolve('interface Gi0/1\n switchport mode access\n switchport access vlan 445')
    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={'methodologies': records, 'errors': [], 'findings': []})
    monkeypatch.setattr(app, 'collect_compliance', lambda *a, **kw: SimpleNamespace(spool_path=spool.path))
    path = app.run_compliance([], None, reports_dir=tmp_path, methodology_report=True)
    assert path.exists()
    assert (tmp_path / ('methodology_resolution_' + spool.path.name + '.xlsx')).exists()
    assert not spool.path.exists()


@pytest.mark.parametrize('statement,expected', [
    ('encapsulation dot1q 445 exact', ' encapsulation dot1q 445 exact'),
    ('encapsulation mystery arbitrary-sensitive-value', ' encapsulation [REDACTED] [REDACTED]'),
    ('encapsulation dot1q 445 token confidential', '[unsupported forwarding statement omitted: unsafe or oversized evidence]'),
    ('encapsulation dot1q 445 community confidential', '[unsupported forwarding statement omitted: unsafe or oversized evidence]'),
    ('encapsulation =HYPERLINK("sensitive")', ' encapsulation [REDACTED]'),
    ('encapsulation ' + 'z' * 600, '[unsupported forwarding statement omitted: unsafe or oversized evidence]'),
])
def test_unknown_disclosure_workbook_and_policy(tmp_path, statement, expected):
    base = 'interface Gi0/1\n switchport mode trunk\n'
    ctx, state = snapshot(base + ' ' + statement, 'C3850')
    _, clean = snapshot(base, 'C3850')
    assert evaluate_vlan_compliance(ctx, [], state, POLICY) == evaluate_vlan_compliance(ctx, [], clean, POLICY)
    rows = resolve_methodologies(ctx, [], state)
    proof = rows[0]['configuration_findings'][-1]['evidence'][0]
    assert proof == dict(source_filename='fixture.cfg', line=3, excerpt=expected)
    assert rows[0]['review_needed']
    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={'methodologies': rows})
    book = load_workbook(export_methodology_spool(spool.path, tmp_path / 'review.xlsx'))
    summary = str(list(book.worksheets[0].values))
    details = ''.join(str(row[-1]) for row in list(book['Details'].values)[1:])
    assert expected in summary
    assert expected == json.loads(details)['configuration_findings'][-1]['evidence'][0]['excerpt']
    assert 'confidential' not in details and 'sensitive' not in details
    assert all(cell.data_type != 'f' for tab in book for row in tab for cell in row)
    book.close()


def test_unknown_service_context_and_bounded_capture():
    config = ('interface Gi0/1\n service instance 7 ethernet\n  encapsulation dot1q 445 exact\n'
              '  bridge-domain 445\n service instance 8 ethernet\n  encapsulation dot1q 545\n  bridge-domain 545')
    rows = resolve(config, 'ASR920')
    assert 'UNSUPPORTED_FORWARDING_SYNTAX' in json.dumps(rows[0])
    assert 'UNSUPPORTED_FORWARDING_SYNTAX' not in json.dumps(rows[1])
    rows = resolve('interface Gi0/1\n' + ' encapsulation dot1q 445 exact\n' * 1000)
    assert len(rows) == 1
    assert len(rows[0]["configuration_findings"]) == 257
    assert rows[-1]['config_interface_name'] == 'Gi0/1'
    assert 'storage limit' in json.dumps(rows[-1])


def test_out_of_scope_noise_and_xr_attachment_identity():
    config = ('interface TenGigE0/0/0/1\n description transit\n!\n'
              'interface BVI445\n ipv4 address 192.0.2.1/24\n!\n'
              'interface PW-Ether7\n!\n'
              'interface TenGigE0/0/0/1.445 l2transport\n encapsulation dot1q 445\n!\n'
              'l2vpn\n bridge group G\n  bridge-domain B\n'
              '   routed interface BVI445\n   interface PW-Ether7\n   interface TenGigE0/0/0/1.445')
    rows = resolve(config, 'NCS540')
    assert len(rows) == 3
    by_name = {r['config_interface_name']: r for r in rows}
    for name, subtype in [('BVI445', 'routed_attachment'), ('PW-Ether7', 'pseudowire_attachment')]:
        assert by_name[name]['subtype'] == subtype
        assert by_name[name]['methodology'] == []
        assert by_name[name]['record_kind'] == 'service_context'
        assert by_name[name]['parent_interface'] == ''
    assert by_name['TenGigE0/0/0/1.445']['methodology'] == ['M04']
    noise = ('member arbitrary-global-value\ninterface Gi0/1\n switchport nonegotiate\n'
             ' switchport port-security maximum 2\n ip address 192.0.2.1 255.255.255.0')
    assert 'UNSUPPORTED_FORWARDING_SYNTAX' not in json.dumps(resolve(noise))
