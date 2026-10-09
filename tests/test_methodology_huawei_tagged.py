"""NE05E operator examples: review projection only, no device side effects."""

from dataclasses import replace
import json
import zipfile

import pytest
from openpyxl import load_workbook

from orbitflow.compliance import evaluate_vlan_compliance
from orbitflow.compliance.resolvers import resolver_for
from orbitflow.execution import DeviceOutcome
from orbitflow.methodology_report import COLUMNS, export_methodology_spool
from orbitflow.result_spool import ResultSpool
from test_methodology_resolution import resolve, snapshot
from test_vlan_compliance import POLICY


@pytest.mark.parametrize('name,tag', [
    ('GigabitEthernet0/2/11.1382', 1382), ('GigabitEthernet0/2/11.1383', 1383),
    ('GE0/2/5.303', 303), ('GE0/2/6.302', 302),
    ('GE0/2/0.1376', 1376), ('GE0/2/0.1377', 1377),
    ('GE0/2/11.1370', 1370), ('GE0/2/11.1371', 1371),
    ('GE0/2/11.999', 1382),  # Never use the suffix as a VLAN.
    ('GE0/2/11.1', 1), ('GE0/2/11.4094', 4094),  # Syntax range, not audit range.
])
def test_real_subinterfaces_preserve_identity_evidence_and_compliance(name, tag):
    parent = name.rsplit('.', 1)[0]
    config = (f'interface {parent}\n!\ninterface {name}\n vlan-type dot1q {tag}\n'
              ' description Uplink to DC-VIC-CORE-RTR3\n!')
    ctx, state = snapshot(config, 'NE05E')
    before = evaluate_vlan_compliance(ctx, [], state, POLICY)
    facts_before = resolver_for(ctx)(ctx, state).result()
    record, = resolve(config, 'NE05E', (name,))
    assert record['methodology'] == ['M08']
    assert record['subtype'] == 'vlan_tagged_subinterface'
    assert record['status'] == 'resolved' and not record['review_needed']
    assert record['config_interface_name'] == record['interface_name'] == name
    assert record['parent_interface'] == parent
    assert record['configuration_owner'] == name
    assert record['interface_match_status'] == 'matched'
    assert record['mapping'] == dict(outer_vlan=[tag], inner_vlan=[],
                                     encapsulation_type='vlan-type-dot1q', role='not_determined')
    assert not record['configuration_findings']
    for proof in record['evidence']:
        assert proof['excerpt'] == config.splitlines()[proof['line'] - 1]
        assert proof['source_filename'] == 'fixture.cfg'
    assert evaluate_vlan_compliance(ctx, [], state, POLICY) == before
    assert resolver_for(ctx)(ctx, state).result() == facts_before
    assert before[0]['observed']['valid_database_vlans'] == []


@pytest.mark.parametrize('header,command', [
    ('interface GE0/2/11.1382', ' VLAN-TYPE DOT1Q 1382'),
    ('interface GE0/2/11.1382', '\tvlan-type\t dot1q  \t1382  '),
    ('INTERFACE\tGE0/2/11.1382', ' VlAn-TyPe DoT1q 1382'),
])
def test_case_whitespace_are_review_only(header, command):
    config = f'interface GE0/2/11\n!\n{header}\n{command}\n!'
    record, = resolve(config, 'NE05E')
    assert record['methodology'] == ['M08'] and record['status'] == 'resolved'
    assert record['mapping']['outer_vlan'] == [1382]
    assert any(p['line'] == 4 and p['excerpt'] == command for p in record['evidence'])
    ctx, state = snapshot(config, 'NE05E')
    # Tolerant spellings must not enter the established audit facts or policy.
    def strip(nodes):
        return tuple(replace(n, children=strip(n.children)) for n in nodes
                     if n.kind not in {'methodology_interface', 'methodology_encapsulation', 'methodology_unknown'})
    baseline = replace(state, configuration=strip(state.configuration))
    assert resolver_for(ctx)(ctx, state).result() == resolver_for(ctx)(ctx, baseline).result()
    assert evaluate_vlan_compliance(ctx, [], state, POLICY) == evaluate_vlan_compliance(ctx, [], baseline, POLICY)


@pytest.mark.parametrize('tag', ['0', '4095', '99999'])
def test_invalid_numeric_tags_preserve_observer_error_contract(tag):
    with pytest.raises(ValueError, match='invalid VLAN range'):
        resolve(f'interface GE0/2/11.1382\n vlan-type dot1q {tag}', 'NE05E')


@pytest.mark.parametrize('command', ['', 'vlan-type', 'vlan-type dot1q',
    'vlan-type dot1q nope', 'vlan-type dot1q -1', 'vlan-type dot1q 1382 trailing',
    'VLAN-TYPE DOT1Q 4095', 'vlan-type dot1q 1382 second-dot1q 17'])
def test_absent_or_malformed_encapsulation_never_claims_m08(command):
    rows = resolve(f'interface GE0/2/11.1382\n {command}', 'NE05E')
    assert not any('M08' in r['methodology'] for r in rows)
    assert all(r['review_needed'] for r in rows)
    if command:
        assert 'UNSUPPORTED_FORWARDING_SYNTAX' in json.dumps(rows)


def test_m06_m07_separation_and_physical_scope():
    config = ('vsi svc\n!\ninterface GE0/2/11\n port link-type trunk\n'
              ' port trunk allow-pass vlan 445\n!\ninterface GE0/2/11.1382\n'
              ' vlan-type dot1q 1382\n!\ninterface GE0/2/11.445\n'
              ' dot1q termination vid 445\n l2 binding vsi svc')
    rows = resolve(config, 'NE05E')
    assert {tuple(r['methodology']) for r in rows} == {('M06',), ('M07',), ('M08',)}
    tagged = next(r for r in rows if r['methodology'] == ['M08'])
    assert not tagged['review_needed']
    assert 'vsi' not in tagged['mapping'] and 'audit_vlan' not in tagged['mapping']
    physical = resolve('interface GE0/2/11\n vlan-type dot1q 1382', 'NE05E')
    assert all(r['methodology'] != ['M08'] for r in physical)


@pytest.mark.parametrize('extra,code', [
    (' vlan-type dot1q 1383', 'CONFLICTING_VLAN_TAGS'),
    (' dot1q termination vid 1382', 'MIXED_INTERFACE_CONSTRUCTS'),
    (' l2 binding vsi absent', 'VSI_REFERENCE_NOT_FOUND'),
    (' port link-type trunk', 'MIXED_INTERFACE_CONSTRUCTS'),
])
def test_genuine_ambiguity_has_source_evidence(extra, code):
    rows = resolve('interface GE0/2/11\n!\ninterface GE0/2/11.1382\n'
                   ' vlan-type dot1q 1382\n' + extra, 'NE05E')
    record = next(r for r in rows if r['methodology'] == ['M08'])
    assert record['review_needed']
    assert next(p for p in record['configuration_findings'] if p['code'] == code)['evidence']
    assert record['mapping']['role'] == 'not_determined'


def test_missing_parent_and_identical_duplicates():
    record, = resolve('interface GE0/2/11.1382\n vlan-type dot1q 1382', 'NE05E')
    assert record['methodology'] == ['M08']
    assert [p['code'] for p in record['configuration_findings']] == ['PARENT_INTERFACE_NOT_FOUND']
    assert record['review_needed']
    record, = resolve('interface GE0/2/11\n!\ninterface GE0/2/11.1382\n'
                      ' vlan-type dot1q 1382\n vlan-type dot1q 1382', 'NE05E')
    assert record['status'] == 'resolved' and record['mapping']['outer_vlan'] == [1382]


def test_m08_workbook_mapping_determinism_and_no_secrets(tmp_path):
    rows = resolve('interface GE0/2/11\n!\ninterface GE0/2/11.1382\n'
                   ' vlan-type dot1q 1382\n description password hiddenvalue', 'NE05E')
    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={'methodologies': rows})
    path = export_methodology_spool(spool.path, tmp_path / 'review.xlsx')
    book = load_workbook(path)
    first = [list(tab.values) for tab in book]
    assert tuple(first[0][0]) == COLUMNS
    row = dict(zip(COLUMNS, first[0][1]))
    assert row['Methodology'] == 'M08' and row['Subtype'] == 'vlan_tagged_subinterface'
    assert json.loads(row['Mapping'])['role'] == 'not_determined'
    assert json.loads(row['Mapping'])['outer_vlan'] == [1382]
    assert row['Status'] == 'resolved' and row['Review Needed'] == 'False'
    assert 'vlan-type dot1q 1382' in row['Configuration Evidence']
    assert all(c.data_type != 'f' for tab in book for row_ in tab for c in row_)
    book.close()
    export_methodology_spool(spool.path, path)
    book = load_workbook(path)
    assert first == [list(tab.values) for tab in book]
    book.close()
    with zipfile.ZipFile(path) as archive:
        assert b'hiddenvalue' not in b''.join(archive.read(n) for n in archive.namelist())
