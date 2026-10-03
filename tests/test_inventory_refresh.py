"""Refresh/export contracts using the real loader, resolver and JSON store."""

from contextlib import contextmanager
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from zipfile import ZipFile
from pathlib import Path
from functools import partial
import json

from openpyxl import Workbook, load_workbook
import pytest

from orbitflow import inventory_refresh as refresh
from orbitflow.config import ExecutionConfig
from orbitflow.inventory import DeviceInventoryResolver, JsonInventoryStore
from orbitflow.models import DeviceContext
from orbitflow.transport import TransportConfig
from scripts import device_inventory_refresh as command


BEFORE = datetime(2026, 9, 28, tzinfo=timezone.utc)
NOW = BEFORE + timedelta(days=1)
CONFIG = TransportConfig("proxy", "cluster", "bastion", "user")
COMMAND_ARGS = ["--proxy", "proxy", "--cluster", "cluster",
                "--bastion-host", "bastion", "--bastion-user", "user"]


def test_operator_command_defaults_overrides_and_unique_names(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    calls = []
    def capture(source, config, **kwargs):
        calls.append((source, config, kwargs))
        return kwargs["export_path"]
    monkeypatch.setattr(command, "refresh_inventory_from_excel", capture)
    class FrozenClock:
        @staticmethod
        def now(tz):
            return NOW
    monkeypatch.setattr(command, "datetime", FrozenClock)
    first = command.main(["targets.xlsx", *COMMAND_ARGS])
    second = command.main(["targets.xlsx", *COMMAND_ARGS])
    assert first != second
    assert first.parent == Path("outputs/reports/inventory")
    assert first.name.startswith("inventory_20260929T000000_000000Z_")
    source, transport, options = calls[0]
    assert source == "targets.xlsx"
    assert transport == CONFIG
    assert options["inventory_path"] == Path("data/inventory/inventory.json")
    assert options["spool_root"] == Path("outputs/runs/inventory_refresh")
    assert options["log_root"] == Path("outputs/logs")
    assert options["execution_config"] == ExecutionConfig()
    Path("custom.toml").write_text("[execution]\nmax_concurrent_devices = 2\nconnection_start_interval = 0\n")
    command.main(["targets.xlsx", *COMMAND_ARGS, "--config", "custom.toml",
                  "--inventory-path", "custom/state.json", "--reports-dir", "custom/reports",
                  "--spool-root", "custom/runs", "--log-root", "custom/logs",
                  "--teleport-key-path", "identity", "--teleport-cert-path", "identity-cert.pub"])
    _, transport, options = calls[-1]
    assert transport.teleport_key_path == Path("identity")
    assert transport.teleport_cert_path == Path("identity-cert.pub")
    assert options["inventory_path"] == Path("custom/state.json")
    assert options["export_path"].parent == Path("custom/reports")
    assert options["spool_root"] == Path("custom/runs")
    assert options["log_root"] == Path("custom/logs")
    assert options["execution_config"] == ExecutionConfig(2, 0)


def test_operator_command_full_export_and_recoverable_failure(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(command, "TransportConfig", lambda **kwargs: CONFIG)
    monkeypatch.setattr(command, "refresh_inventory_from_excel",
                        partial(refresh.refresh_inventory_from_excel, clock=lambda: NOW))
    store = JsonInventoryStore("data/inventory/inventory.json")
    failed = seed(store, "192.0.2.1", "SERIAL1")
    omitted = seed(store, "192.0.2.3", "SERIAL3")
    connected, _, _ = setup_network(monkeypatch, {
        "192.0.2.1": "connect-failure", "192.0.2.2": ("SERIAL2", "new-router"),
    })
    source = input_file(tmp_path, [target("192.0.2.1"), target("192.0.2.2"),
                                   ["192.0.2.4", "", "runtime-password"]])
    book = load_workbook(source)
    book.active.cell(1, 4, "Secret")
    for row in range(2, 5):
        book.active.cell(row, 4, "runtime-enable")
    book.save(source)
    book.close()
    first = command.main([str(source), *COMMAND_ARGS])
    original = first.read_bytes()
    rows = sheets(first)
    by_id = {row["device_id"]: row for row in rows["Inventory"]}
    assert by_id[failed.device_id]["refresh_status"] == refresh.FAILED
    assert by_id[failed.device_id]["hostname"] == failed.hostname
    assert by_id[failed.device_id]["last_successful_collection"] == BEFORE.isoformat()
    assert by_id[omitted.device_id]["refresh_status"] == refresh.NOT_REQUESTED
    assert len(rows["Inventory"]) == 3
    assert [row["status"] for row in rows["Run_Attempts"]] == ["failed", "refreshed", "failed"]
    assert sorted(connected) == ["192.0.2.1", "192.0.2.2"]
    second = command.main([str(source), *COMMAND_ARGS])
    assert second != first and second.is_file()
    assert first.read_bytes() == original
    assert list(Path("outputs/runs/inventory_refresh").iterdir()) == []
    logs = list(Path("outputs/logs/inventory").glob("????-??-??/inventory_refresh.log"))
    assert logs
    writer = refresh._write_export
    def fail(*args, **kwargs):
        raise OSError("runtime-login runtime-password runtime-enable token=private-test-token")
    monkeypatch.setattr(refresh, "_write_export", fail)
    with pytest.raises(SystemExit) as error:
        command.main([str(source), *COMMAND_ARGS])
    assert error.value.code != 0
    assert error.value.__suppress_context__
    assert first.read_bytes() == original
    spool = next(Path("outputs/runs/inventory_refresh").iterdir())
    manifest = json.loads((spool / "manifest.json").read_text())
    assert manifest["collection_complete"]
    assert manifest["completed_count"] == 3
    assert manifest["failed_count"] == 2
    persisted = "".join(path.read_text() for path in spool.glob("*.json*"))
    persisted += store.path.read_text() + "".join(path.read_text() for path in logs)
    output = capsys.readouterr()
    persisted += str(error.value) + output.out + output.err
    with ZipFile(first) as archive:
        persisted += "".join(archive.read(name).decode() for name in archive.namelist() if name.endswith(".xml"))
    for secret in ("runtime-login", "runtime-password", "runtime-enable", "private-test-token"):
        assert secret not in persisted
    monkeypatch.setattr(refresh, "_write_export", writer)
    def forbidden(*args, **kwargs):
        raise AssertionError("Recovery must not connect")
    monkeypatch.setattr(refresh, "connect_device", forbidden)
    recovered = refresh.export_inventory_spool(spool, "outputs/reports/inventory/recovered.xlsx")
    assert sheets(recovered) == rows
    assert not spool.exists()


@pytest.mark.parametrize("stage", ["config", "input"])
def test_operator_command_suppresses_raw_setup_errors(tmp_path, monkeypatch, capsys, stage):
    monkeypatch.chdir(tmp_path)
    def fail(*args, **kwargs):
        raise ValueError("runtime-password token=hidden")
    monkeypatch.setattr(command, "load_execution_config" if stage == "config"
                        else "refresh_inventory_from_excel", fail)
    with pytest.raises(SystemExit) as error:
        command.main(["targets.xlsx", *COMMAND_ARGS])
    output = capsys.readouterr()
    assert "runtime-password" not in str(error.value) + output.out + output.err
    assert "hidden" not in str(error.value) + output.out + output.err
    assert error.value.__suppress_context__


def seed(store, ip, serial, hostname="old-router"):
    return store.reconcile(DeviceContext(
        device_id="", management_ip=ip, observed_management_ips=(ip,),
        hostname=hostname, vendor="Cisco", platform="cisco_ios",
        device_family="C3750X", hardware_model="WS-C3750X", capability_profile="c3750x_switching",
        capability_flags=("switchport",), serial_number=serial, software_version="15.1",
        uptime="1 day", last_successful_collection=BEFORE, last_collection_attempt=BEFORE,
    ))[0]


def input_file(tmp_path, rows):
    path = tmp_path / "targets.xlsx"
    book = Workbook()
    book.active.append(["management_ip", "username", "password"])
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()
    return path


def setup_network(monkeypatch, outcomes):
    connected, closed, commands = [], [], []

    @contextmanager
    def connect(ip, credentials, config):
        assert config is CONFIG
        assert credentials.username == "runtime-login"
        assert credentials.password == "runtime-password"
        connected.append(ip)
        if outcomes[ip] == "connect-failure":
            raise OSError("runtime-login runtime-password token=do-not-export")
        try:
            yield ip
        finally:
            closed.append(ip)
        if outcomes[ip] == "disconnect-failure":
            raise OSError("runtime-password")

    class Runner:
        prompt = "router#"

        def __init__(self, ip):
            self.ip = ip

        def close(self):
            pass

        def run_command(self, command, timeout=10):
            commands.append(command)
            result = outcomes[self.ip]
            if result == "probe-failure":
                raise RuntimeError("runtime-password")
            serial, hostname = ("SERIAL1", "router") if result == "disconnect-failure" else result
            if command == "show version":
                return f"Cisco IOS Software, Version 15.2\n{hostname} uptime is 4 days\nWS-C3750X"
            assert command == "show inventory"
            return f"NAME: chassis\nPID: WS-C3750X, VID: V02\nSN: {serial}"

    monkeypatch.setattr(refresh, "connect_device", connect)
    monkeypatch.setattr(DeviceInventoryResolver, "_new_runner", lambda self, session: Runner(session))
    return connected, closed, commands


def run(tmp_path, rows):
    return refresh.refresh_inventory_from_excel(
        input_file(tmp_path, rows), CONFIG, inventory_path=tmp_path / "inventory.json",
        export_path=tmp_path / "inventory.xlsx", log_root=tmp_path / "logs", clock=lambda: NOW, execution_config=ExecutionConfig(1),
    )


def sheets(path):
    book = load_workbook(path)
    try:
        return {sheet.title: [dict(zip(next(sheet.values), row))
                              for row in list(sheet.values)[1:]] for sheet in book}
    finally:
        book.close()


def target(ip):
    return [ip, "runtime-login", "runtime-password"]


def test_full_export_refreshes_only_input_and_preserves_previous_facts(tmp_path, monkeypatch):
    store = JsonInventoryStore(tmp_path / "inventory.json")
    good = seed(store, "192.0.2.1", "SERIAL1")
    failed = seed(store, "192.0.2.2", "SERIAL2")
    omitted = seed(store, "192.0.2.3", "SERIAL3")
    connected, closed, commands = setup_network(monkeypatch, {
        good.management_ip: ("SERIAL1", "new-router"), failed.management_ip: "connect-failure",
        "192.0.2.4": "probe-failure", "192.0.2.5": ("SERIAL5", "new-device"),
    })
    path = run(tmp_path, [target(good.management_ip), target(failed.management_ip),
                          target("192.0.2.4"), target("192.0.2.5")])
    data = sheets(path)
    assert len(data["Inventory"]) == 4
    rows = {row["device_id"]: row for row in data["Inventory"]}
    assert rows[good.device_id]["refresh_status"] == refresh.REFRESHED
    assert rows[good.device_id]["hostname"] == "new-router"
    assert rows[failed.device_id]["refresh_status"] == refresh.FAILED
    assert rows[failed.device_id]["supplied_in_current_input"] == "True"
    assert rows[failed.device_id]["last_successful_collection"] == BEFORE.isoformat()
    assert rows[failed.device_id]["last_collection_attempt"] == NOW.isoformat()
    assert rows[omitted.device_id]["refresh_status"] == refresh.NOT_REQUESTED
    assert rows[omitted.device_id]["supplied_in_current_input"] == "False"
    contexts = {item.device_id: item for item in store.contexts()}
    assert contexts[omitted.device_id] == omitted
    retained = contexts[failed.device_id]
    assert replace(retained, last_collection_attempt=BEFORE, collection_status="success", collection_error="") == failed
    assert [row["status"] for row in data["Run_Attempts"]] == ["refreshed", "failed", "failed", "refreshed"]
    assert data["Run_Attempts"][2]["device_ids"] is None
    assert connected == ["192.0.2.1", "192.0.2.2", "192.0.2.4", "192.0.2.5"]
    assert closed == ["192.0.2.1", "192.0.2.4", "192.0.2.5"]
    assert set(commands) == {"show version", "show inventory"}
    with ZipFile(path) as archive:
        exported = "".join(archive.read(name).decode() for name in archive.namelist() if name.endswith(".xml"))
    persisted = store.path.read_text() + exported + "".join(p.read_text() for p in (tmp_path / "logs").rglob("*.log"))
    for secret in ("runtime-login", "runtime-password", "do-not-export"):
        assert secret not in persisted


def test_serial_move_replacement_and_hostname_collision_keep_existing_semantics(tmp_path, monkeypatch):
    store = JsonInventoryStore(tmp_path / "inventory.json")
    original = seed(store, "192.0.2.1", "SERIAL1", "router")
    setup_network(monkeypatch, {
        "192.0.2.2": ("SERIAL1", "router"),
        "192.0.2.1": ("SERIAL2", "router"),
    })
    data = sheets(run(tmp_path, [target("192.0.2.2"), target("192.0.2.1")]))
    assert len(data["Inventory"]) == 2
    contexts = {item.serial_number: item for item in store.contexts()}
    assert contexts["SERIAL1"].device_id == original.device_id
    assert contexts["SERIAL1"].observed_management_ips == ("192.0.2.1", "192.0.2.2")
    assert contexts["SERIAL2"].device_id != original.device_id
    assert data["Run_Attempts"][0]["reconciliation_events"] == "management_ip_changed"
    assert "likely_replacement_or_ip_reassignment" in data["Run_Attempts"][1]["reconciliation_events"]
    assert "hostname_collision" in data["Run_Attempts"][1]["reconciliation_events"]


def test_replacement_does_not_label_old_identity_refreshed(tmp_path, monkeypatch):
    store = JsonInventoryStore(tmp_path / "inventory.json")
    old = seed(store, "192.0.2.1", "SERIAL1")
    setup_network(monkeypatch, {"192.0.2.1": ("SERIAL2", "router")})
    rows = sheets(run(tmp_path, [target("192.0.2.1")]))["Inventory"]
    assert next(row for row in rows if row["device_id"] == old.device_id)["refresh_status"] == refresh.NOT_REQUESTED
    assert next(item for item in store.contexts() if item.device_id == old.device_id) == old


@pytest.mark.parametrize("outcome", ["connect-failure", "probe-failure"])
def test_repeated_alias_failure_is_visible_and_retains_latest_success(tmp_path, monkeypatch, outcome):
    store = JsonInventoryStore(tmp_path / "inventory.json")
    old = seed(store, "192.0.2.1", "SERIAL1")
    setup_network(monkeypatch, {"192.0.2.2": ("SERIAL1", "updated"), "192.0.2.1": outcome})
    rows = sheets(run(tmp_path, [target("192.0.2.2"), target("192.0.2.1")]))["Inventory"]
    assert len(rows) == 1
    assert rows[0]["device_id"] == old.device_id
    assert rows[0]["refresh_status"] == refresh.FAILED
    assert rows[0]["hostname"] == "updated"
    assert rows[0]["last_successful_collection"] == NOW.isoformat()


def test_invalid_row_isolated_and_missing_serial_does_not_merge(tmp_path, monkeypatch):
    store = JsonInventoryStore(tmp_path / "inventory.json")
    old = seed(store, "192.0.2.1", "")
    connected, _, _ = setup_network(monkeypatch, {"192.0.2.2": ("", "old-router")})
    data = sheets(run(tmp_path, [["192.0.2.1", "runtime-login", ""], target("192.0.2.2")]))
    assert connected == ["192.0.2.2"]
    assert len(data["Inventory"]) == 2
    assert data["Run_Attempts"][0]["stage"] == "input"
    assert data["Run_Attempts"][0]["device_ids"] == old.device_id


def test_secrets_in_facts_redacted_and_excel_formulas_are_literal(tmp_path, monkeypatch):
    store = JsonInventoryStore(tmp_path / "inventory.json")
    seed(store, "192.0.2.2", "SERIAL2", "=1+1")
    setup_network(monkeypatch, {"192.0.2.1": ("SERIAL1", "runtime-password")})
    path = run(tmp_path, [target("192.0.2.1")])
    book = load_workbook(path)
    try:
        sheet = book["Inventory"]
        headers = [cell.value for cell in sheet[1]]
        cell = sheet.cell(2, headers.index("hostname") + 1)
        assert cell.value == "[REDACTED]"
        assert cell.data_type == "s"
        formula = sheet.cell(3, headers.index("hostname") + 1)
        assert formula.value == "=1+1"
        assert formula.data_type == "s"
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref
    finally:
        book.close()
    assert "runtime-password" not in (tmp_path / "inventory.json").read_text()


def test_disconnect_failure_does_not_erase_successful_refresh(tmp_path, monkeypatch):
    setup_network(monkeypatch, {"192.0.2.1": "disconnect-failure"})
    data = sheets(run(tmp_path, [target("192.0.2.1")]))
    assert data["Inventory"][0]["refresh_status"] == refresh.REFRESHED
    assert data["Run_Attempts"][0]["status"] == "failed"
    assert data["Run_Attempts"][0]["stage"] == "disconnect"


def test_sensitive_serial_is_rejected_instead_of_becoming_shared_identity(tmp_path, monkeypatch):
    setup_network(monkeypatch, {"192.0.2.1": ("runtime-password", "router")})
    data = sheets(run(tmp_path, [target("192.0.2.1")]))
    assert data["Inventory"] == []
    assert data["Run_Attempts"][0]["status"] == "failed"
    assert "runtime-password" not in (tmp_path / "inventory.json").read_text()


def test_all_unknown_failures_export_empty_inventory_and_attempts(tmp_path, monkeypatch):
    setup_network(monkeypatch, {"192.0.2.1": "connect-failure"})
    data = sheets(run(tmp_path, [target("192.0.2.1")]))
    assert data["Inventory"] == []
    assert data["Run_Attempts"][0]["status"] == "failed"


def test_export_cannot_overwrite_input(tmp_path):
    source = input_file(tmp_path, [target("192.0.2.1")])
    original = source.read_bytes()
    with pytest.raises(ValueError, match="distinct"):
        refresh.refresh_inventory_from_excel(source, CONFIG, inventory_path=tmp_path / "inventory.json", export_path=source)
    assert source.read_bytes() == original


def test_next_run_resets_membership_without_clearing_stored_failure(tmp_path, monkeypatch):
    store = JsonInventoryStore(tmp_path / "inventory.json")
    seed(store, "192.0.2.1", "SERIAL1")
    setup_network(monkeypatch, {"192.0.2.1": "probe-failure", "192.0.2.2": ("SERIAL2", "router")})
    run(tmp_path, [target("192.0.2.1")])
    rows = sheets(run(tmp_path, [target("192.0.2.2")]))["Inventory"]
    old = next(row for row in rows if row["management_ip"] == "192.0.2.1")
    assert old["refresh_status"] == refresh.NOT_REQUESTED
    assert old["supplied_in_current_input"] == "False"
    assert old["collection_status"] == "failed"
    assert old["hostname"] == "old-router"


def test_later_success_clears_failed_attempt_and_keeps_identity(tmp_path, monkeypatch):
    store = JsonInventoryStore(tmp_path / "inventory.json")
    old = seed(store, "192.0.2.1", "SERIAL1")
    setup_network(monkeypatch, {"192.0.2.1": "probe-failure"})
    run(tmp_path, [target("192.0.2.1")])
    setup_network(monkeypatch, {"192.0.2.1": ("SERIAL1", "recovered")})
    row = sheets(run(tmp_path, [target("192.0.2.1")]))["Inventory"][0]
    assert row["device_id"] == old.device_id
    assert row["refresh_status"] == refresh.REFRESHED
    assert row["collection_status"] == "success"
    assert row["collection_error"] is None


def test_export_write_failure_preserves_prior_workbook(tmp_path, monkeypatch):
    setup_network(monkeypatch, {"192.0.2.1": ("SERIAL1", "router")})
    path = run(tmp_path, [target("192.0.2.1")])
    previous = path.read_bytes()
    source = input_file(tmp_path, [target("192.0.2.1")])

    def fail_save(self, filename):
        raise OSError("Disk write failed")

    monkeypatch.setattr(Workbook, "save", fail_save)
    with pytest.raises(OSError):
        refresh.refresh_inventory_from_excel(source, CONFIG, inventory_path=tmp_path / "inventory.json",
                                             export_path=path, log_root=tmp_path / "logs")
    assert path.read_bytes() == previous
    assert not path.with_name(path.name + ".tmp").exists()


def test_concurrent_store_instances_reconcile_aliases_and_failures(tmp_path):
    from threading import Barrier
    from orbitflow.execution import execute_devices
    path = tmp_path / "inventory.json"
    store = JsonInventoryStore(path)
    old = seed(store, "old", "OLD")
    barrier = Barrier(8)
    def worker(index):
        # Separate instances and path spellings must share the transaction lock.
        local = JsonInventoryStore(path.parent / "." / path.name)
        barrier.wait(timeout=10)
        if index == 0:
            local.record_failure("old", NOW, "safe failure")
        else:
            seed(local, f"address-{index}", "SHARED" if index < 4 else f"SERIAL-{index}")
    outcomes = execute_devices(range(8), worker, config=ExecutionConfig(8))
    assert all(not item.error_category for item in outcomes)
    contexts = store.contexts()
    assert len(contexts) == 6
    shared = next(item for item in contexts if item.serial_number == "SHARED")
    assert set(shared.observed_management_ips) == {"address-1", "address-2", "address-3"}
    retained = next(item for item in contexts if item.device_id == old.device_id)
    assert retained.hostname == old.hostname
    assert retained.collection_status == "failed"
    assert retained.last_collection_attempt == NOW


def test_refresh_workers_overlap_and_export_after_cleanup(tmp_path, monkeypatch):
    from threading import Barrier, get_ident
    barrier = Barrier(2)
    caller = get_ident()
    connected, closed, _ = setup_network(monkeypatch, {
        "192.0.2.1": ("SERIAL1", "one"), "192.0.2.2": ("SERIAL2", "two"),
    })
    original_connect = refresh.connect_device
    @contextmanager
    def connect(*args):
        assert get_ident() != caller
        with original_connect(*args) as session:
            barrier.wait(timeout=10)
            yield session
    monkeypatch.setattr(refresh, "connect_device", connect)
    writer = refresh._write_export
    def write(*args, **kwargs):
        assert get_ident() == caller
        assert len(closed) == 2
        return writer(*args, **kwargs)
    monkeypatch.setattr(refresh, "_write_export", write)
    path = refresh.refresh_inventory_from_excel(
        input_file(tmp_path, [target("192.0.2.1"), target("192.0.2.2")]), CONFIG,
        inventory_path=tmp_path / "inventory.json", export_path=tmp_path / "out.xlsx",
        log_root=tmp_path / "logs", execution_config=ExecutionConfig(2),
    )
    data = sheets(path)
    assert len(connected) == len(data["Inventory"]) == 2
    assert [row["management_ip"] for row in data["Run_Attempts"]] == ["192.0.2.1", "192.0.2.2"]
    assert all(row["status"] == refresh.REFRESHED for row in data["Run_Attempts"])


def test_inventory_spool_retry_uses_original_export_view(tmp_path, monkeypatch):
    setup_network(monkeypatch, {'192.0.2.1': ('SERIAL1', 'original')})
    writer = refresh._write_export
    def fail(*args, **kwargs):
        raise OSError('runtime-password')
    monkeypatch.setattr(refresh, '_write_export', fail)
    with pytest.raises(OSError):
        run(tmp_path, [target('192.0.2.1')])
    spool = next((tmp_path / 'runs').iterdir())
    for file in spool.glob('*.json*'):
        assert 'runtime-password' not in file.read_text()
    store = JsonInventoryStore(tmp_path / 'inventory.json')
    seed(store, '192.0.2.1', 'SERIAL1', 'changed-after-run')
    canonical = store.path.read_bytes()
    monkeypatch.setattr(refresh, '_write_export', writer)
    def forbidden(*args, **kwargs):
        raise AssertionError('must not reconnect')
    monkeypatch.setattr(refresh, 'connect_device', forbidden)
    path = refresh.export_inventory_spool(spool, tmp_path / 'retry.xlsx')
    data = sheets(path)
    assert data['Inventory'][0]['hostname'] == 'original'
    assert data['Run_Attempts'][0]['status'] == refresh.REFRESHED
    assert canonical == store.path.read_bytes()
    assert not spool.exists()


def test_1500_inventory_attempts_stream_at_concurrency_five(tmp_path, monkeypatch):
    # Repeated targets also require one outcome per input position. Real resolver
    # and canonical reconciliation must retain one physical device.
    connected, closed, _ = setup_network(monkeypatch, {'192.0.2.1': ('SERIAL1', 'router')})
    writer = refresh._write_export
    def streaming_writer(path, rows, attempts, **kwargs):
        assert iter(rows) is rows
        assert iter(attempts) is attempts
        return writer(path, rows, attempts, **kwargs)
    monkeypatch.setattr(refresh, '_write_export', streaming_writer)
    path = refresh.refresh_inventory_from_excel(
        input_file(tmp_path, [target('192.0.2.1') for _ in range(1500)]), CONFIG,
        inventory_path=tmp_path / 'inventory.json', export_path=tmp_path / 'inventory.xlsx',
        log_root=tmp_path / 'logs', clock=lambda: NOW, execution_config=ExecutionConfig(5))
    data = sheets(path)
    assert len(data['Inventory']) == 1
    assert [int(row['input_position']) for row in data['Run_Attempts']] == list(range(1, 1501))
    assert len(connected) == len(closed) == 1500
    assert len(JsonInventoryStore(tmp_path / 'inventory.json').contexts()) == 1
    assert list((tmp_path / 'runs').iterdir()) == []
