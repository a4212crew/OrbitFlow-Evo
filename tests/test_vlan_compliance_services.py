"""Deterministic IOS-XR policy analysis, with no transport or vendor CLI parsing."""

from copy import deepcopy
from dataclasses import replace
import json

import pytest
from openpyxl import load_workbook

from orbitflow.compliance import evaluate_vlan_compliance, parse_policy
from orbitflow.models import InterfaceVlanObservation, VlanObject
from orbitflow.execution import DeviceOutcome
from orbitflow.result_spool import ResultSpool
from orbitflow.compliance_report import export_compliance_spool
from test_vlan_compliance import POLICY_DATA, context, interface, state


PARENT = "TenGigE0/0/0/18"
OTHER = "TenGigE0/0/0/19"


def policy_data():
    data = deepcopy(POLICY_DATA)
    rule = data["service_rules"][0]
    # Small custom baseline keeps generic evaluator tests independent of defaults.
    rule["required"] = deepcopy(rule["signature"][:2]) + [
        [{"field": "object_id", "value": "LEAPTEL-PPPoE/LEAPTEL-PPPoE"}],
        [{"field": "object_id", "value": "URL-PPPOE/URL-PPPOE"}],
    ]
    rule["required"].append([{"field": "object_id", "value": "TEST/RSVD-RSP0"}])
    rule["baseline_complete"] = True
    return data


def snapshot(lbb="LBB-PPPoE", parents=None):
    domains = (lbb, "BD_VLAN545", "LEAPTEL-PPPoE", "URL-PPPOE", "RSVD-RSP0")
    identities = ("LBB/" + lbb, "VLAN545/BD_VLAN545", "LEAPTEL-PPPoE/LEAPTEL-PPPoE",
                  "URL-PPPOE/URL-PPPOE", "TEST/RSVD-RSP0")
    objects = tuple(VlanObject("bridge_domain", identity, domain_id=domain)
                    for identity, domain in zip(identities, domains))
    records, profiles = [], []
    for parent, selected in (parents or {PARENT: domains}).items():
        records.append(replace(interface(parent), platform="cisco_xr"))
        profiles.append(InterfaceVlanObservation(parent, port_type="routed"))
        for index, domain in enumerate(selected, 1):
            name = f"{parent}.{index}"
            records.append(replace(interface(name), platform="cisco_xr"))
            profiles.append(InterfaceVlanObservation(name, port_type="service",
                            bridge_domains=(domain,), service_mappings=(f"{index} -> {domain}",)))
    return records, replace(state(), platform="cisco_xr", objects=objects, interfaces=tuple(profiles))


def evaluate(records, vlans, data=None):
    return evaluate_vlan_compliance(replace(context(), platform="cisco_xr"),
                                    records, vlans, parse_policy(data or policy_data()))


@pytest.mark.parametrize("lbb", ["LBB-PPPoE", "LBB-PPPOE"])
def test_services_consolidate_and_case_alternatives(lbb):
    records, vlans = snapshot(lbb)
    before = vlans
    findings = evaluate(records, vlans)
    assert [f["status"] for f in findings] == ["compliant", "compliant"]
    assert findings[1]["interface"] == PARENT
    assert len(findings[1]["observed"]["members"]) == 5
    assert vlans == before
    assert all(p.port_type != "trunk" for p in vlans.interfaces)


@pytest.mark.parametrize("provider", ["LEAPTEL-PPPoE", "URL-PPPOE"])
def test_parent_scope_and_missing_named_baseline(provider):
    records, vlans = snapshot(parents={PARENT: ("LBB-PPPoE", "BD_VLAN545", provider),
                                      OTHER: ("RSVD-RSP0",)})
    findings = evaluate(records, vlans)
    assert findings[1]["status"] == "non_compliant"
    assert [{"field": "object_id", "value": "TEST/RSVD-RSP0"}] in findings[1]["missing_objects"]
    assert findings[2]["interface"] == OTHER
    assert findings[2]["status"] == "not_applicable"
    assert findings[0]["status"] == "compliant"


def test_database_full_identity_not_numeric_tags_or_domain_name():
    records, vlans = snapshot()
    objects = tuple(replace(o, object_id="WRONG/RSVD-RSP0", vlan_ids=(2400,))
                    if o.domain_id == "RSVD-RSP0" else o for o in vlans.objects)
    findings = evaluate(records, replace(vlans, objects=objects))
    assert all(f["status"] == "non_compliant" for f in findings)
    assert findings[0]["missing_objects"] == [[{"field": "object_id", "value": "TEST/RSVD-RSP0"}]]


@pytest.mark.parametrize("change", ["ambiguous_domain", "missing_object", "missing_parent", "missing_profile", "duplicate_profile"])
def test_incomplete_or_ambiguous_normalized_evidence(change):
    records, vlans = snapshot()
    if change == "ambiguous_domain":
        vlans = replace(vlans, objects=vlans.objects + (VlanObject("bridge_domain", "OTHER/BD_VLAN545", domain_id="BD_VLAN545"),))
    elif change == "missing_object":
        vlans = replace(vlans, objects=vlans.objects[1:])
    elif change == "missing_parent":
        records = records[1:]
    elif change == "missing_profile":
        vlans = replace(vlans, interfaces=vlans.interfaces[:-1])
    else:
        vlans = replace(vlans, interfaces=vlans.interfaces + (vlans.interfaces[-1],))
    assert evaluate(records, vlans)[1]["status"] == "unable_to_assess"


def test_unobserved_service_reference_does_not_satisfy_parent():
    records, vlans = snapshot()
    records = records[:-1]
    finding = evaluate(records, vlans)[1]
    assert finding["status"] == "non_compliant"
    assert len(finding["observed"]["members"]) == 4


def test_exact_signature_and_policy_changes():
    records, vlans = snapshot(lbb="lbb-pppoe")
    assert evaluate(records, vlans)[1]["status"] == "not_applicable"
    data = policy_data()
    data["service_rules"][0]["signature"][0] = [{"field": "domain_id", "value": "lbb-pppoe"}]
    data["service_rules"][0]["required"][0] = data["service_rules"][0]["signature"][0]
    assert evaluate(records, vlans, data)[1]["status"] == "compliant"


def test_order_does_not_change_findings_and_routed_units_do_not_contribute():
    records, vlans = snapshot()
    assert evaluate(records, vlans) == evaluate(list(reversed(records)), replace(
        vlans, objects=tuple(reversed(vlans.objects)), interfaces=tuple(reversed(vlans.interfaces))))
    profiles = tuple(replace(p, port_type="routed") if p.interface_name.endswith(".5") else p
                     for p in vlans.interfaces)
    assert evaluate(records, replace(vlans, interfaces=profiles))[1]["status"] == "non_compliant"


def test_old_policy_and_incomplete_custom_baseline_are_explicit():
    records, vlans = snapshot()
    old = policy_data()
    del old["service_rules"]
    assert all(f["reason"] == "service_policy_unavailable" for f in evaluate(records, vlans, old))
    incomplete = policy_data()
    incomplete["service_rules"][0]["baseline_complete"] = False
    assert all(f["reason"] == "service_baseline_incomplete" for f in evaluate(records, vlans, incomplete))
    assert evaluate(records, None)[0]["reason"] == "vlan_collection_unavailable"
    assert evaluate(None, vlans)[1]["status"] == "unable_to_assess"


@pytest.mark.parametrize("field,value", [
    ("platform", "cisco_ios"), ("signature", []), ("required", [[]]),
    ("baseline_complete", "true"), ("database_rule", "trunk-required-vlans"),
    ("required", [[{"field": "name", "value": "secret"}]]),
    ("required", [[{"field": "object_id", "value": "password=do-not-echo"}]]),
])
def test_service_policy_validation(field, value):
    data = policy_data()
    data["service_rules"][0][field] = value
    with pytest.raises(ValueError) as caught:
        parse_policy(data)
    assert "do-not-echo" not in str(caught.value)


def test_service_findings_spool_excel_recovery_and_secret_exclusion(tmp_path, monkeypatch):
    import orbitflow.compliance_report as report
    records, vlans = snapshot()
    vlans = replace(vlans, interfaces=tuple(replace(p, description="secret-description",
                    service_details=(InterfaceVlanObservation("secret-raw-output"),)) for p in vlans.interfaces))
    findings = evaluate(records, vlans)
    assert "secret-" not in json.dumps(findings)
    spool = ResultSpool.create(tmp_path / "runs", "vlan_compliance", 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={"findings": findings, "errors": []})
    writer = report.write_tables
    def fail(*args, **kwargs):
        raise OSError("synthetic write failure")
    monkeypatch.setattr(report, "write_tables", fail)
    with pytest.raises(OSError):
        export_compliance_spool(spool.path, tmp_path / "report.xlsx")
    monkeypatch.setattr(report, "write_tables", writer)
    path = export_compliance_spool(spool.path, tmp_path / "report.xlsx")
    workbook = load_workbook(path)
    try:
        assert workbook["Findings"].max_row == 3
        assert PARENT in str(list(workbook["Findings"].values))
    finally:
        workbook.close()

# Independent representation of the supplied observation, not derived from policy.
NCS_NAMES = (
    "IQNET-PPPOE", "LBB-PPPoE", "LEAPTEL-PPPoE",
    *(f"RSVD-RSP{i}" for i in range(38)), "SPIRIT-PPPOE",
    *(f"SUPERLOOP{i}-PPPOE" for i in range(1, 7)), "URL-PPPOE",
)
NCS_IDENTITIES = tuple(f"{name}/{name}" for name in NCS_NAMES) + (
    "VLAN545/BD_VLAN545", "VLAN745/BD_VLAN745",
)


def ncs_snapshot(*, omitted=None, lbb="LBB-PPPoE", database_missing=False):
    identities = tuple(identity.replace("LBB-PPPoE", lbb) for identity in NCS_IDENTITIES)
    objects = tuple(VlanObject("bridge_domain", identity, domain_id=identity.split("/")[1])
                    for identity in identities if not (database_missing and identity == omitted))
    records, profiles = [], []
    for parent, selected in ((PARENT, [o for o in objects if o.object_id != omitted]),
                             (OTHER, [o for o in objects if o.domain_id == "IQNET-PPPOE"])):
        records.append(replace(interface(parent), platform="cisco_xr"))
        profiles.append(InterfaceVlanObservation(parent, port_type="routed"))
        for index, obj in enumerate(selected, 1):
            name = f"{parent}.{index}"
            records.append(replace(interface(name), platform="cisco_xr"))
            profiles.append(InterfaceVlanObservation(name, port_type="service",
                                                     bridge_domains=(obj.domain_id,)))
    return records, replace(state(), platform="cisco_xr", objects=objects, interfaces=tuple(profiles))


@pytest.mark.parametrize("lbb", ["LBB-PPPoE", "LBB-PPPOE"])
def test_complete_supplied_default_ncs_baseline(lbb):
    records, vlans = ncs_snapshot(lbb=lbb)
    rule = parse_policy(POLICY_DATA).service_rules[0]
    assert rule.baseline_complete
    assert len(rule.required) == len(NCS_IDENTITIES) == 51
    assert {group[0].field for group in rule.required} == {"object_id"}
    assert {selector.value for group in rule.required for selector in group} == (
        set(NCS_IDENTITIES) | {"LBB-PPPOE/LBB-PPPOE"})
    findings = evaluate(records, vlans, POLICY_DATA)
    assert [f["status"] for f in findings] == ["compliant", "compliant", "not_applicable"]
    assert all(not f["missing_objects"] for f in findings)


@pytest.mark.parametrize("omitted", NCS_IDENTITIES)
@pytest.mark.parametrize("database_missing", [False, True])
def test_each_default_ncs_identity_is_required(omitted, database_missing):
    records, vlans = ncs_snapshot(omitted=omitted, database_missing=database_missing)
    findings = evaluate(records, vlans, POLICY_DATA)
    group = next(group for group in parse_policy(POLICY_DATA).service_rules[0].required
                 if any(selector.value == omitted for selector in group))
    missing = [[{"field": selector.field, "value": selector.value} for selector in group]]
    assert findings[0]["status"] == ("non_compliant" if database_missing else "compliant")
    assert findings[0]["missing_objects"] == (missing if database_missing else [])
    # Losing a mandatory trigger makes this parent out of scope, as before.
    triggered = omitted not in ("LBB-PPPoE/LBB-PPPoE", "VLAN545/BD_VLAN545")
    assert findings[1]["status"] == ("non_compliant" if triggered else "not_applicable")
    assert findings[1]["missing_objects"] == (missing if triggered else [])
    assert findings[2]["interface"] == OTHER
    assert findings[2]["status"] == "not_applicable"
    assert findings == evaluate(list(reversed(records)), replace(
        vlans, objects=tuple(reversed(vlans.objects)),
        interfaces=tuple(reversed(vlans.interfaces))), POLICY_DATA)
