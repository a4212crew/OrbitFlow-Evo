"""Synthetic C1-C4 regressions. Network and subprocess access are forbidden."""

from dataclasses import replace
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path
import runpy
import socket
import subprocess

from openpyxl import load_workbook
import pytest

from orbitflow.configuration import ApprovalStore, ChangePlan, PlanError, load_plan
from orbitflow.configuration.catalogue import CATALOGUE, parameters
from orbitflow.configuration.guided_schema import digest
from orbitflow.configuration.job_cli import main
from orbitflow.configuration.jobs import MANUAL, REQUEST, load_manifest, plan, prepare
from orbitflow.configuration.plans import canonical
from orbitflow.configuration.workbooks import write
from orbitflow.inventory.store import JsonInventoryStore
from orbitflow.models import DeviceContext

NOW = datetime(2026, 10, 10, tzinfo=timezone.utc)


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("offline planning attempted network or subprocess side effect")
    monkeypatch.setattr(socket, "socket", forbidden)
    monkeypatch.setattr(socket, "create_connection", forbidden)
    monkeypatch.setattr(subprocess, "Popen", forbidden)


@pytest.fixture
def contexts():
    base = DeviceContext("switch", "192.0.2.1", ("192.0.2.1",), "switch", "Cisco", "cisco_xe", "C3850",
                         "WS-C3850", "c3850_switching", ("classic_switchport",), "SYNTH1", "16.12", "", NOW, NOW)
    return [base,
            replace(base, device_id="evc", hostname="evc", management_ip="192.0.2.2", observed_management_ips=("192.0.2.2",),
                    device_family="ASR920", capability_profile="asr920_evc", capability_flags=("evc",), serial_number="SYNTH2"),
            replace(base, device_id="huawei", hostname="huawei", vendor="Huawei", platform="huawei_vrp", management_ip="192.0.2.3",
                    observed_management_ips=("192.0.2.3",), device_family="NE05E", capability_profile="ne05e",
                    capability_flags=("dot1q_subinterface",), serial_number="SYNTH3")]


def request(tmp_path, rows=None, manual=None):
    path = tmp_path / "request.xlsx"
    write(path, {"Requests": (REQUEST, rows or [["r1", "switch", "GigabitEthernet1/0/1", "trunk VLAN add", ""]]),
                 "Manual": (MANUAL, manual or [])})
    return path


def edit(path, sheet, column, value, match=lambda row: True):
    book = load_workbook(path)
    ws = book[sheet]
    headers = [c.value for c in ws[1]]
    for cells in ws.iter_rows(min_row=2):
        row = {k: c.value for k, c in zip(headers, cells)}
        if match(row):
            cells[headers.index(column)].value = value
    book.save(path)
    book.close()


def completed(tmp_path, contexts, rows=None, manual=None):
    src = request(tmp_path, rows, manual)
    out = tmp_path / "guided.xlsx"
    prepare(src, contexts, out)
    edit(out, "Parameters", "value", "3500,3501", lambda r: r["name"] == "vlans")
    edit(out, "Parameters", "value", 3500, lambda r: r["name"] == "vlan")
    return out


def run(tmp_path, path, contexts):
    manifest = plan(path, contexts, tmp_path / "batch", "TEST")
    plans = [load_plan(tmp_path / "batch" / m["path"]) for m in manifest["members"]]
    assert load_manifest(tmp_path / "batch/manifest.json") == manifest
    return manifest, plans


def test_mixed_vendor_order_provenance_and_advanced_review(tmp_path, contexts):
    rows = [["r1", "switch", "GigabitEthernet1/0/1", "trunk VLAN add", ""],
            ["r2", "switch", "GigabitEthernet1/0/1", "manual CLI", "m1"],
            ["r3", "huawei", "GigabitEthernet0/1/1.3500", "tagged subinterface", ""]]
    path = completed(tmp_path, contexts, rows, [["m1", 2, "description Reviewed"], ["m1", 1, "interface GigabitEthernet1/0/1"]])
    edit(path, "Requests", "depends_on", "r1", lambda r: r["row_id"] == "r2")
    manifest, plans = run(tmp_path, path, contexts)
    assert len(plans) == 2 and len(manifest["results"]) == 3
    switch = next(p for p in plans if p.to_dict()["identity"]["inventory_id"] == "switch")
    data = switch.to_dict()
    assert data["status"] == "review_required"
    assert [o["row_id"] for o in data["operations"]] == ["r1", "r2"]
    assert data["operations"][1]["commands"] == ["interface GigabitEthernet1/0/1", "description Reviewed"]
    assert data["execution_policy"]["execution_authorized"] is False
    copy = switch.to_dict()
    copy["operations"].clear()
    assert switch.to_dict() == data
    assert data["source"]["completed_digest"] == manifest["completed_digest"]
    book = load_workbook(tmp_path / "batch/configuration_preview_TEST.xlsx")
    assert book["CLI Preview"].max_row > 8
    assert {"Summary", "Inputs", "Operations", "Findings", "Checks"} <= set(book.sheetnames)
    book.close()


@pytest.mark.parametrize("selection", ["", "cisco.evc.local", "cisco.evc.global"])
def test_evc_ambiguity_requires_engineer_selection(tmp_path, contexts, selection):
    path = completed(tmp_path, contexts, [["r1", "evc", "GigabitEthernet0/0/1", "EVC", ""]])
    if selection:
        edit(path, "Requests", "template_id", selection)
        edit(path, "Requests", "template_version", "1")
        edit(path, "Parameters", "value", 3500,
             lambda r: r["template_id"] == selection and r["required"] is True)
    manifest, plans = run(tmp_path, path, contexts)
    if not selection:
        assert not plans and manifest["results"][0]["status"] == "blocked"
    else:
        assert plans[0].to_dict()["status"] == "offline_validated"
        op = plans[0].to_dict()["operations"][0]
        assert op["method"] == ("M03" if selection.endswith("local") else "M02")
        assert "member GigabitEthernet0/0/1 service-instance 3500" in op["commands"] if selection.endswith("global") else True


@pytest.mark.parametrize("action", ["access port", "trunk VLAN replace"])
def test_cross_operation_conflict(tmp_path, contexts, action):
    path = completed(tmp_path, contexts, [["r1", "switch", "GigabitEthernet1/0/1", "trunk VLAN add", ""],
                                          ["r2", "switch", "GigabitEthernet1/0/1", action, ""]])
    _, plans = run(tmp_path, path, contexts)
    assert plans[0].to_dict()["status"] == "blocked"
    assert "conflicting_interface_operations" in {f["code"] for f in plans[0].to_dict()["findings"]}
    with pytest.raises(PlanError, match="blocked"):
        ApprovalStore(tmp_path / "authority").decide(plans[0], "approved", "engineer", now=NOW, expires_at=NOW + timedelta(days=1))


@pytest.mark.parametrize("dependency,order", [("r2", 1), ("missing", 2), ("r1", 1)])
def test_order_and_dependencies_block(tmp_path, contexts, dependency, order):
    path = completed(tmp_path, contexts, [["r1", "switch", "GigabitEthernet1/0/1", "trunk VLAN add", ""],
                                          ["r2", "switch", "GigabitEthernet1/0/2", "trunk VLAN add", ""]])
    edit(path, "Requests", "depends_on", dependency, lambda r: r["row_id"] == "r1")
    edit(path, "Requests", "order", order, lambda r: r["row_id"] == "r2")
    _, plans = run(tmp_path, path, contexts)
    assert plans[0].to_dict()["status"] == "blocked"


def test_known_manual_remove_conflict(tmp_path, contexts):
    path = completed(tmp_path, contexts, [["r1", "switch", "GigabitEthernet1/0/1", "trunk VLAN add", ""],
                                          ["r2", "switch", "GigabitEthernet1/0/1", "manual CLI", "m"]],
                     [["m", 1, "switchport trunk allowed vlan remove 3500"]])
    _, plans = run(tmp_path, path, contexts)
    assert "manual_removes_requested_vlan" in {f["code"] for f in plans[0].to_dict()["findings"]}


def test_local_approval_never_authorizes_v2(tmp_path, contexts):
    path = completed(tmp_path, contexts, [["r1", "switch", "GigabitEthernet1/0/1", "manual CLI", "m"]],
                     [["m", 1, "mystery forwarding-mode"]])
    _, plans = run(tmp_path, path, contexts)
    store = ApprovalStore(tmp_path / "authority")
    store.decide(plans[0], "approved", "claimed-admin", now=NOW, expires_at=NOW + timedelta(days=1))
    assert store.verify(plans[0], now=NOW, offline_review=True)["status"] == "approved"
    with pytest.raises(PlanError, match="execution is forbidden"):
        store.verify(plans[0], now=NOW)
    store.decide(plans[0], "rejected", "engineer", now=NOW)
    with pytest.raises(PlanError, match="not approved"):
        store.verify(plans[0], now=NOW, offline_review=True)


@pytest.mark.parametrize("field,value", [("interface", "GigabitEthernet1/0/9"), ("inventory_id", "huawei"),
                                         ("row_id", "changed"), ("action", "access port")])
def test_identity_edits_never_retarget(tmp_path, contexts, field, value):
    path = completed(tmp_path, contexts)
    edit(path, "Requests", field, value)
    manifest, plans = run(tmp_path, path, contexts)
    assert not plans and all(r["status"] == "blocked" for r in manifest["results"])


def test_snapshot_change_blocks(tmp_path, contexts):
    path = completed(tmp_path, contexts)
    contexts[0] = replace(contexts[0], serial_number="REPLACED")
    manifest, plans = run(tmp_path, path, contexts)
    assert not plans and "snapshot" in manifest["results"][0]["reason"]


@pytest.mark.parametrize("change", [dict(device_family="unknown"), dict(capability_flags=()),
                                     dict(collection_status="failed")])
def test_unsupported_profiles_and_failed_snapshot_visible(tmp_path, contexts, change):
    contexts[0] = replace(contexts[0], **change)
    path = completed(tmp_path, contexts)
    manifest, plans = run(tmp_path, path, contexts)
    assert not plans and manifest["results"][0]["status"] == "blocked"


def test_duplicate_inventory_identity_visible(tmp_path, contexts):
    contexts.append(replace(contexts[0], device_id="duplicate"))
    path = completed(tmp_path, contexts)
    manifest, plans = run(tmp_path, path, contexts)
    assert not plans and "ambiguous" in manifest["results"][0]["reason"]


@pytest.mark.parametrize("value", ["", 0, 4095, 1002, True, "3500\nend", "3500;reload"])
def test_typed_parameters_rejected(tmp_path, contexts, value):
    path = completed(tmp_path, contexts)
    edit(path, "Parameters", "value", value, lambda r: r["name"] == "vlans")
    manifest, plans = run(tmp_path, path, contexts)
    assert not plans and manifest["results"][0]["status"] == "blocked"


@pytest.mark.parametrize("secret", ["password superSensitive", "snmp-server community superSensitive",
                                  "username admin privilege 15 superSensitive", "crypto key superSensitive"])
def test_secrets_never_persist(tmp_path, contexts, secret):
    path = completed(tmp_path, contexts, [["r1", "switch", "GigabitEthernet1/0/1", "manual CLI", "m"]],
                     [["m", 1, secret]])
    manifest, plans = run(tmp_path, path, contexts)
    assert not plans
    assert "superSensitive" not in canonical(manifest)
    assert "superSensitive" not in path.with_suffix(".json").read_text()
    book = load_workbook(path)
    assert all("superSensitive" not in str(c.value) for sheet in book for row in sheet for c in row)
    book.close()


@pytest.mark.parametrize("decision", ["skip", "reject"])
def test_exclusions_and_missing_rows_remain_visible(tmp_path, contexts, decision):
    path = completed(tmp_path, contexts, [["r1", "switch", "GigabitEthernet1/0/1", "trunk VLAN add", ""],
                                          ["r2", "switch", "GigabitEthernet1/0/2", "trunk VLAN add", ""]])
    edit(path, "Requests", "decision", decision, lambda r: r["row_id"] == "r2")
    manifest, plans = run(tmp_path, path, contexts)
    assert len(manifest["results"]) == 2 and plans[0].to_dict()["status"] == "blocked"
    assert manifest["results"][1]["status"] == ("skipped" if decision == "skip" else "rejected")


@pytest.mark.parametrize("mutation", ["commands", "policy", "findings", "source", "template", "unknown"])
def test_v2_semantic_tampering_rejected(tmp_path, contexts, mutation):
    path = completed(tmp_path, contexts)
    _, plans = run(tmp_path, path, contexts)
    data = plans[0].to_dict()
    if mutation == "commands":
        data["operations"][0]["commands"].append("reload")
    elif mutation == "policy":
        data["execution_policy"]["execution_authorized"] = True
    elif mutation == "findings":
        data["status"] = "approved"
    elif mutation == "source":
        data["source"]["catalogue_version"] = "999"
    elif mutation == "template":
        data["operations"][0]["template_version"] = "999"
    else:
        data["extra"] = True
    with pytest.raises(PlanError):
        ChangePlan(canonical(data))


def test_digest_and_bound_manifest_tampering(tmp_path, contexts):
    path = completed(tmp_path, contexts)
    manifest, plans = run(tmp_path, path, contexts)
    target = tmp_path / "batch" / manifest["members"][0]["path"]
    envelope = json.loads(target.read_text())
    envelope["plan"]["change_id"] = "CHANGED"
    target.write_text(canonical(envelope))
    with pytest.raises(PlanError, match="digest"):
        load_manifest(tmp_path / "batch/manifest.json")
    envelope["digest"] = digest(envelope["plan"])
    target.write_text(canonical(envelope))
    with pytest.raises(PlanError, match="membership"):
        load_manifest(tmp_path / "batch/manifest.json")


def test_cli_and_exclusive_artifacts(tmp_path, contexts):
    inventory = tmp_path / "inventory.json"
    for c in contexts:
        JsonInventoryStore(inventory, id_factory=lambda c=c: c.device_id).reconcile(c)
    path = completed(tmp_path, contexts)
    assert main(["plan", str(path), "--inventory", str(inventory), "--output", str(tmp_path / "batch"), "--batch-id", "TEST"]) == 0
    assert main(["validate-batch", str(tmp_path / "batch/manifest.json")]) == 0
    assert main(["plan", str(path), "--inventory", str(inventory), "--output", str(tmp_path / "batch")]) == 2
    with pytest.raises(PlanError, match="already exists"):
        prepare(tmp_path / "request.xlsx", contexts, path)


def test_catalogue_optional_parameters_and_no_implicit_si_defaults():
    evc = next(t for t in CATALOGUE if t.template_id == "cisco.evc.local")
    with pytest.raises(PlanError, match="required"):
        parameters(evc, {"outer_vlan": 3500})
    assert parameters(evc, {"outer_vlan": 3500, "service_instance": 7, "bridge_domain": 9}) == {
        "outer_vlan": 3500, "service_instance": 7, "bridge_domain": 9}


def test_manual_fields_generated_without_reference(tmp_path, contexts):
    path = completed(tmp_path, contexts, [["r1", "switch", "GigabitEthernet1/0/1", "manual CLI", ""]])
    edit(path, "Manual", "command", "description Literal engineer input")
    _, plans = run(tmp_path, path, contexts)
    assert plans[0].to_dict()["status"] == "review_required"


def test_missing_row_and_invalid_sibling_block_device_not_other_device(tmp_path, contexts):
    path = completed(tmp_path, contexts, [["r1", "switch", "GigabitEthernet1/0/1", "trunk VLAN add", ""],
                                          ["r2", "switch", "GigabitEthernet1/0/2", "trunk VLAN add", ""],
                                          ["r3", "huawei", "GigabitEthernet0/1/1.3500", "tagged subinterface", ""]])
    book = load_workbook(path)
    book["Requests"].delete_rows(3)
    book.save(path)
    book.close()
    manifest, plans = run(tmp_path, path, contexts)
    assert len(manifest["results"]) == 3
    statuses = {p.to_dict()["identity"]["inventory_id"]: p.to_dict()["status"] for p in plans}
    assert statuses == {"switch": "blocked", "huawei": "offline_validated"}


@pytest.mark.parametrize("value", ["Gi1/0/1", "GigabitEthernet01/0/1"])
def test_interface_aliases_fail_closed(tmp_path, contexts, value):
    path = completed(tmp_path, contexts, [["r1", "switch", value, "trunk VLAN add", ""]])
    manifest, plans = run(tmp_path, path, contexts)
    assert not plans and manifest["results"][0]["status"] == "blocked"


def test_duplicate_service_and_encapsulation_conflicts(tmp_path, contexts):
    path = completed(tmp_path, contexts, [["r1", "evc", "GigabitEthernet0/0/1", "EVC", ""],
                                          ["r2", "evc", "GigabitEthernet0/0/1", "EVC", ""]])
    edit(path, "Requests", "template_id", "cisco.evc.local")
    edit(path, "Requests", "template_version", "1")
    edit(path, "Parameters", "value", 3500, lambda r: r["template_id"] == "cisco.evc.local" and r["required"] is True)
    _, plans = run(tmp_path, path, contexts)
    codes = {f["code"] for f in plans[0].to_dict()["findings"]}
    assert {"duplicate_service_instance", "overlapping_encapsulation"} <= codes


def test_independent_trunk_adds_preserve_explicit_order(tmp_path, contexts):
    path = completed(tmp_path, contexts, [["r1", "switch", "GigabitEthernet1/0/1", "trunk VLAN add", ""],
                                          ["r2", "switch", "GigabitEthernet1/0/1", "trunk VLAN add", ""]])
    edit(path, "Parameters", "value", "3502", lambda r: r["row_id"] == "r2")
    edit(path, "Requests", "order", 10, lambda r: r["row_id"] == "r1")
    _, plans = run(tmp_path, path, contexts)
    assert plans[0].to_dict()["status"] == "offline_validated"
    assert [o["row_id"] for o in plans[0].to_dict()["operations"]] == ["r2", "r1"]


def test_formula_input_blocked_and_literal_output(tmp_path, contexts):
    path = completed(tmp_path, contexts)
    edit(path, "Parameters", "value", "=3500+1")
    manifest, plans = run(tmp_path, path, contexts)
    assert not plans and "formula" in manifest["results"][0]["reason"]
    output = tmp_path / "literal.xlsx"
    write(output, {"Literal": (["value"], [["=1+1"], ["@SUM(1)"]])})
    book = load_workbook(output)
    assert book["Literal"]["A2"].data_type == "s"
    book.close()


def test_completed_secret_parameters_omitted(tmp_path, contexts):
    path = completed(tmp_path, contexts, [["r1", "switch", "GigabitEthernet1/0/1", "access port", ""]])
    edit(path, "Parameters", "value", "password NeverPersist", lambda r: r["name"] == "description")
    manifest, plans = run(tmp_path, path, contexts)
    assert not plans and "NeverPersist" not in canonical(manifest)


def test_repeated_plan_determinism_and_changes_invalidate_approval(tmp_path, contexts):
    path = completed(tmp_path, contexts)
    _, plans = run(tmp_path, path, contexts)
    other = plan(path, contexts, tmp_path / "repeat", "TEST")
    assert other["members"][0]["digest"] == plans[0].digest
    store = ApprovalStore(tmp_path / "authority")
    store.decide(plans[0], "approved", "engineer", now=NOW, expires_at=NOW + timedelta(days=1))
    edit(path, "Parameters", "value", "3502")
    changed = plan(path, contexts, tmp_path / "changed", "TEST")
    changed_plan = load_plan(tmp_path / "changed" / changed["members"][0]["path"])
    assert changed_plan.digest != plans[0].digest
    with pytest.raises(PlanError, match="unapproved"):
        store.verify(changed_plan, now=NOW, offline_review=True)


def test_manifest_rehashed_missing_outcomes_rejected(tmp_path, contexts):
    path = completed(tmp_path, contexts)
    run(tmp_path, path, contexts)
    target = tmp_path / "batch/manifest.json"
    envelope = json.loads(target.read_text())
    envelope["manifest"]["results"].clear()
    envelope["digest"] = digest(envelope["manifest"])
    target.write_text(canonical(envelope))
    with pytest.raises(PlanError, match="missing from batch"):
        load_manifest(target)


def test_duplicate_rows_and_unknown_parameters_are_visible(tmp_path, contexts):
    path = completed(tmp_path, contexts)
    book = load_workbook(path)
    book["Parameters"].append(["orphan", "cisco.trunk.add", "vlans", "vlans", True, "", "3500", "3500"])
    book["Parameters"].append(["orphan2", "cisco.trunk.add", "vlans", "vlans", True, "", "3500", "3500"])
    book.save(path)
    book.close()
    manifest, plans = run(tmp_path, path, contexts)
    assert len(manifest["results"]) == 3 and plans[0].to_dict()["status"] == "blocked"


def test_corrupt_workbook_fails_without_traceback(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_text("not a workbook")
    assert main(["prepare", str(path), "--output", str(tmp_path / "guided.xlsx")]) == 2


def test_generated_operator_samples_end_to_end(tmp_path):
    generate = runpy.run_path(str(Path(__file__).parents[1] / "scripts/configuration_samples.py"))["generate"]
    sample = generate(tmp_path / "samples")
    contexts = JsonInventoryStore(sample / "inventory.json").contexts()
    manifest = plan(sample / "completed.xlsx", contexts, tmp_path / "sample-batch", "SAMPLE")
    assert len(manifest["members"]) == 3 and len(manifest["results"]) == 6
    assert load_manifest(tmp_path / "sample-batch/manifest.json") == manifest
    plans = [load_plan(tmp_path / "sample-batch" / m["path"]).to_dict() for m in manifest["members"]]
    assert {p["identity"]["inventory_id"]: p["status"] for p in plans} == {
        "sample-switch": "review_required", "sample-evc": "offline_validated", "sample-huawei": "offline_validated"}
    prepare(sample / "conflicts.xlsx", contexts, sample / "conflict-guided.xlsx")
    edit(sample / "conflict-guided.xlsx", "Parameters", "value", "3500,3501")
    negative = plan(sample / "conflict-guided.xlsx", contexts, tmp_path / "negative", "NEGATIVE")
    assert len(negative["results"]) == 4 and all(r["status"] == "blocked" for r in negative["results"])
    assert load_manifest(tmp_path / "negative/manifest.json") == negative


@pytest.mark.parametrize("template", CATALOGUE, ids=lambda t: t.template_id)
def test_all_catalogue_profiles_render_closed_v2_operations(template):
    from orbitflow.configuration.guided_schema import operation
    for platform, family, profile in template.profiles:
        ident = dict(platform=platform, device_family=family, capability_profile=profile)
        values = {p.name: ("3500,3501" if p.kind == "vlans" else 3500) for p in template.parameters if p.required}
        op = operation(dict(row_id="example", source_row=2, interface="GigabitEthernet0/1" +
                            (".3500" if platform == "huawei_vrp" else ""), action=template.action,
                            template_id=template.template_id, template_version=template.version, order=1,
                            depends_on=[], parameters=values, manual_cli=[]), ident, [template.capability])
        assert op["commands"] and op["method"] == template.method
        assert not {"commit", "save", "write memory"}.intersection(op["commands"])


def test_generated_ids_do_not_collide_and_duplicates_retain_rows(tmp_path, contexts):
    path = completed(tmp_path, contexts, [["row-000002", "switch", "GigabitEthernet1/0/1", "trunk VLAN add", ""],
                                          ["", "switch", "GigabitEthernet1/0/2", "trunk VLAN add", ""],
                                          ["duplicate", "switch", "GigabitEthernet1/0/3", "trunk VLAN add", ""],
                                          ["duplicate", "switch", "GigabitEthernet1/0/4", "trunk VLAN add", ""]])
    manifest, _ = run(tmp_path, path, contexts)
    assert len(manifest["results"]) == 4
    assert len({r["row_id"] for r in manifest["results"]}) == 4
    assert len([r for r in manifest["results"] if "duplicate source" in r["reason"]]) == 2


def test_equivalent_ipv6_addresses_are_ambiguous(tmp_path, contexts):
    contexts[0] = replace(contexts[0], management_ip="2001:db8::1", observed_management_ips=("2001:db8::1",))
    contexts[1] = replace(contexts[1], management_ip="2001:0db8::1", observed_management_ips=("2001:0db8::1",))
    path = completed(tmp_path, contexts)
    manifest, plans = run(tmp_path, path, contexts)
    assert not plans and "ambiguous" in manifest["results"][0]["reason"]
