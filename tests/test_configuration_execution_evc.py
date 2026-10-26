"""Issue 154 synthetic execution; no device connections."""
from dataclasses import replace
from types import SimpleNamespace
import socket

import pytest

from orbitflow.configuration import execution as ex
from orbitflow.configuration.v2.planning import request, plan
from orbitflow.configuration.plans import PlanError
from orbitflow.vendors.cisco import evc_execution as evc
from orbitflow.vendors.cisco.evc import child_commands
from orbitflow.capabilities.save import save_configuration
from orbitflow.transport import DeviceCredentials
from orbitflow.result_spool import ResultJournal
from test_configuration_v2_approved_services import observation, template, evc_inputs, evc_config, IP, NOW
from test_configuration_execution import context, fake_runtime, WorkerCLI, cli_with


def fixture(family, config=None, inputs=None, function='EVC'):
    ctx = replace(context(family, IP), platform='cisco_ios' if family == 'ME3600X' else 'cisco_xe',
                  capability_profile='me3600x_evc' if family == 'ME3600X' else 'asr920_evc')
    config = config or evc_config(family, parent=' mtu 9000\n shutdown\n')
    obs = observation(family, config)
    # Use the same canonical identity as the synthetic live resolver.
    obs['identity'] = ex.identity(ctx)
    obs['vlan_lookups'] = {'445': dict(status='present', name='Customer', operational_status='active')}
    p = inputs or evc_inputs()
    r = request(template(family, function).id, p)
    original = plan([r], {IP: obs}, now=NOW)
    result = ex.evaluate_decisions(original, {r['row_id']: 'Approve'}, expected_digest=original.digest, now=NOW).to_dict()['results'][0]
    return ctx, obs, r, original, dict(request=r, result=result)


def presence(p, state='Up'):
    from orbitflow.vendors.cisco.vlans import parse_service_detail
    return parse_service_detail(f"Id: {p['service_instance_id']}\nAssociated Interface: {p['interface']}\nEncapsulation: dot1q 445\nState: {state}\n", p['interface'], p['service_instance_id'])


def attach_presence(obs, p):
    obs.setdefault('service_details', {})[p['interface'] + ':' + str(p['service_instance_id'])] = presence(p)
    return obs


def post(family, p, **kw):
    config = evc_config(family, ' service instance ' + str(p['service_instance_id']) + ' ethernet\n' +
                        ''.join('  ' + line + '\n' for line in child_commands(p)),
                        ' mtu 9000\n shutdown\n' + (' description ' + p['interface_description'] + '\n' if p.get('interface_description') else ''))
    obs = observation(family, config.replace(kw.get('remove', '\0'), ''))
    obs['vlan_lookups'] = {'445': dict(status='present', name='Customer', operational_status='active')}
    return attach_presence(obs, p)


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
@pytest.mark.parametrize('existing', [False, True])
def test_create_recreate_fresh_checks_verifier_and_summary(tmp_path, family, existing):
    p = evc_inputs(split_horizon_group=0, description='CHILD', interface_description='PARENT')
    config = evc_config(family, (' service instance 3000 ethernet\n  encapsulation dot1q 445\n  rewrite ingress tag pop 1 symmetric\n  bridge-domain 445\n' if existing else ''), ' mtu 9000\n shutdown\n')
    ctx, before, r, original, entry = fixture(family, config, p)
    assert not evc.target_blocker(entry, before, ctx)
    after = post(family, p)
    assert evc.verify(entry, before, after, ctx)['status'] == 'Verified'
    assert evc.verify(entry, before, post(family, p, remove=' shutdown\n'), ctx)['status'] == 'Verified'
    prepared = ex.ExecutionService(tmp_path).prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    assert len(prepared['summary']['selected']) == 1
    assert 'Forwarding may be interrupted' in ex.render_summary(ex.operator_summary(prepared))
    assert 'switchport' not in ' '.join(ex.verification_commands(r, entry['result']))


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
@pytest.mark.parametrize('change,code', [
    (' service instance 20 ethernet\n  encapsulation dot1q 445\n  bridge-domain 800\n', 'INGRESS_TAG_OVERLAP'),
    (' service instance 3000 ethernet\n  encapsulation dot1q 446\n  bridge-domain 446\n', 'SERVICE_INSTANCE_DRIFT'),
    (' ip address 192.0.2.1 255.255.255.0\n', 'INCOMPATIBLE_EVC_PARENT'),
    (' channel-group 1 mode active\n', 'OWNERSHIP_REQUIRES_REVIEW'),
])
def test_fresh_negative_gates(family, change, code):
    ctx, before, _, _, entry = fixture(family)
    fresh = observation(family, evc_config(family, change))
    fresh['vlan_lookups'] = before['vlan_lookups']
    assert evc.target_blocker(entry, fresh, ctx) == code


def test_distinct_profiles_and_me_database_prerequisite():
    ctx, before, _, _, entry = fixture('ME3600X')
    before['vlan_lookups'] = {}
    assert evc.target_blocker(entry, before, ctx) == 'VLAN_PREREQUISITE_UNAVAILABLE'
    assert not evc.supported_context(replace(ctx, platform='cisco_xe'))
    ctx, before, _, _, entry = fixture('ASR920')
    before['vlan_lookups'] = {}
    assert not evc.target_blocker(entry, before, ctx)


def test_me_database_conflict_and_unknown_verification():
    ctx, before, r, _, entry = fixture('ME3600X')
    after = post('ME3600X', r['inputs'])
    after['vlan_lookups'] = {}
    after['service_details'] = {}
    assert evc.verify(entry, before, after, ctx)['status'] == 'Unknown'
    before['objects'] = []
    assert evc.target_blocker(entry, before, ctx) == 'VLAN_PREREQUISITE_UNAVAILABLE'


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
@pytest.mark.parametrize('decision', ['Pending', 'Reject'])
def test_nonapproved_never_selected(tmp_path, monkeypatch, family, decision):
    _, _, r, original, _ = fixture(family)
    monkeypatch.setattr(socket, 'create_connection', lambda *a, **kw: pytest.fail('offline preparation connected'))
    prepared = ex.ExecutionService(tmp_path).prepare(original, {r['row_id']: decision}, expected_digest=original.digest)
    assert not prepared['summary']['selected']


@pytest.mark.parametrize('family,function', [('ME3600X', 'VLAN Database'), ('ME3600X', 'Interface Description'), ('ASR920', 'Interface Description')])
def test_non_evc_templates_selected_and_verified(tmp_path, monkeypatch, family, function):
    p = dict(management_ip=IP, vlan_id=445, name='Customer') if function == 'VLAN Database' else dict(management_ip=IP, interface='GigabitEthernet0/1', description='NEW')
    ctx, before, r, original, entry = fixture(family, inputs=p, function=function)
    prepared = ex.ExecutionService(tmp_path).prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    assert len(prepared['summary']['selected']) == 1
    after = observation(family, evc_config(family, parent=' description NEW\n mtu 9000\n shutdown\n'))
    monkeypatch.setattr(ex, 'VlanService', lambda **kw: SimpleNamespace(lookup=lambda *a, **k: dict(status='present', name='Customer', operational_status='active')))
    result = ex._verify_observed(None, ctx, None, entry, before, after=after if function != 'VLAN Database' else None)
    # InterfaceService supplies independent description evidence in real collections.
    if function != 'VLAN Database':
        after['interfaces'][0]['port_description'] = 'NEW'
        result = ex._verify_observed(None, ctx, None, entry, before, after=after)
    assert result['status'] == 'Verified'


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
@pytest.mark.parametrize('failure', ['', 'reject', 'timeout'])
def test_real_lifecycle_exact_order_partial_save_recovery(tmp_path, monkeypatch, family, failure):
    ctx, before, r, original, entry = fixture(family)
    verifier = ex._verify_observed
    saves = fake_runtime(monkeypatch, ctx, before)
    monkeypatch.setattr(ex, '_verify_observed', verifier)
    after = post(family, r['inputs'])
    calls = iter([before, after])
    monkeypatch.setattr(ex, '_observe', lambda *a: next(calls))
    send = WorkerCLI.run_command
    def command(self, cmd, **kw):
        if cmd.startswith('bridge-domain') and failure:
            self.sent.append(cmd)
            if failure == 'timeout':
                self.dead = True
                raise TimeoutError()
            return '% Invalid input'
        return send(self, cmd, **kw)
    monkeypatch.setattr(WorkerCLI, 'run_command', command)
    service = ex.ExecutionService(tmp_path / 'authority')
    prepared = service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    with pytest.raises(PlanError):
        service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'])
    spool = service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True,
                            credentials={IP: DeviceCredentials('test-user', 'hidden')}, run_root=tmp_path / 'runs')
    row = next(spool.records())['payload']['rows'][0]
    assert row['status'] == {'': 'Verified', 'reject': 'Failed', 'timeout': 'Unknown'}[failure]
    assert len(saves) == (0 if failure == 'timeout' else 1)
    commands = entry['result']['commands']
    sent = WorkerCLI.instances[0].sent
    if not failure:
        assert sent == commands
    else:
        index = commands.index('bridge-domain 445')
        assert sent[:index + 1] == commands[:index + 1]
        assert sent[index + 1:] == ([] if failure == 'timeout' else ['end'])
    events = list(ResultJournal(spool.path / 'events.jsonl').records())
    assert any(e['stage'] == 'pre_row_evidence' for e in events)
    assert ex.export_results(spool.path, tmp_path / 'execution.xlsx').exists()
    with pytest.raises(PlanError):
        service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True)


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
def test_service_submode_and_independent_save(family):
    ctx, *_ = fixture(family)
    cli, channel = cli_with([b'configure terminal\nfixture(config)#',
        b'interface GigabitEthernet0/1\nfixture(config-if)#',
        b'service instance 3000 ethernet\nfixture(config-if-srv)#',
        b'end\nfixture#', b'write memory\n% Error saving configuration\nfixture#'])
    for cmd in ('configure terminal', 'interface GigabitEthernet0/1', 'service instance 3000 ethernet', 'end'):
        cli.execution_mode('fixture#', ctx.platform)
        cli.run_command(cmd)
    result = save_configuration(cli.session, ctx, cli=cli, exec_prompt='fixture#', timeout=10)
    assert result['status'] == 'Saved' and result['explicit_error']


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
@pytest.mark.parametrize('enc,outer,inner,rewrite', [
    ('Untagged', None, None, 'Push dot1q'),
    ('Untagged', None, None, 'None'),
    ('Dot1q', '445', None, 'Pop 1 symmetric'),
    ('Dot1q', '2-444,446-544', None, 'Push dot1q symmetric'),
    ('QinQ', '300', '200', 'Pop 2 symmetric'),
    ('QinQ', '100', '1-4094', 'Pop 1 symmetric'),
])
def test_presence_does_not_claim_configured_intent_matrix(family, enc, outer, inner, rewrite):
    p = evc_inputs()
    p.update(encapsulation=enc, rewrite=rewrite, split_horizon_group=0)
    p.pop('outer_vlan')
    if outer:
        p['outer_vlan'] = outer
    if inner:
        p['inner_vlan'] = inner
    if rewrite.startswith('Push'):
        p['push_vlan'] = 1402
    ctx, before, _, _, entry = fixture(family, inputs=p)
    after = post(family, p)
    assert evc.verify(entry, before, after, ctx)['status'] == 'Verified'
    altered = dict(p, split_horizon_group=1)
    assert evc.verify(entry, before, post(family, altered), ctx)['status'] == 'Verified'


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
def test_operational_conflict_and_unreliable_evidence(family):
    ctx, before, r, _, entry = fixture(family)
    key = r['inputs']['interface'] + ':3000'
    before['service_details'] = {key: dict(status='present', facts=[])}
    assert evc.target_blocker(entry, before, ctx) == 'SERVICE_OBSERVATION_CONFLICT'
    after = post(family, r['inputs'])
    after['service_details'] = {key: dict(status='absent')}
    assert evc.verify(entry, before, after, ctx)['status'] == 'Failed'
    after['configuration_reliable'] = False
    assert evc.verify(entry, before, after, ctx)['status'] == 'Failed'
    after = post(family, r['inputs'])
    after['interfaces'] = []
    assert evc.verify(entry, before, after, ctx)['status'] == 'Verified'


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
def test_presence_does_not_claim_policy_preservation(family):
    child = ' service instance 3000 ethernet\n  description OLD\n  encapsulation dot1q 445\n  rewrite ingress tag pop 1 symmetric\n  bridge-domain 445 split-horizon group 0\n  service-policy input CHILD\n'
    sibling = ' service instance 20 ethernet\n  encapsulation dot1q 800\n  bridge-domain 800\n  service-policy input OTHER\n'
    ctx, before, r, _, entry = fixture(family, evc_config(family, child + sibling, ' mtu 9000\n shutdown\n service-policy input PARENT\n'))
    # Rebuild actual full configuration with unrelated settings still present.
    text = evc_config(family, ' service instance 3000 ethernet\n' + ''.join('  ' + c + '\n' for c in child_commands(r['inputs'])) + sibling,
                      ' mtu 9000\n shutdown\n service-policy input PARENT\n')
    after = observation(family, text)
    after['vlan_lookups'] = before['vlan_lookups']
    attach_presence(after, r['inputs'])
    assert evc.verify(entry, before, after, ctx)['status'] == 'Verified'
    damaged = observation(family, text.replace('  service-policy input OTHER\n', ''))
    damaged['vlan_lookups'] = before['vlan_lookups']
    attach_presence(damaged, r['inputs'])
    assert evc.verify(entry, before, damaged, ctx)['status'] == 'Verified'


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
@pytest.mark.parametrize('timeout', [False, True])
def test_independent_row_continuation_and_unknown_gate(tmp_path, monkeypatch, family, timeout):
    ctx, before, r, _, _ = fixture(family)
    p2 = dict(evc_inputs(), service_instance_id=3001, outer_vlan='446')
    r2 = request(template(family).id, p2, row=4)
    original = plan([r, r2], {IP: before}, now=NOW)
    verifier = ex._verify_observed
    saves = fake_runtime(monkeypatch, ctx, before)
    monkeypatch.setattr(ex, '_verify_observed', verifier)
    p1 = r['inputs']
    def child(p, bd=True):
        return ' service instance ' + str(p['service_instance_id']) + ' ethernet\n' + ''.join(
            '  ' + c + '\n' for c in child_commands(p) if bd or not c.startswith('bridge-domain'))
    middle = observation(family, evc_config(family, child(p1, False), ' mtu 9000\n shutdown\n'))
    final = observation(family, evc_config(family, child(p1, False) + child(p2), ' mtu 9000\n shutdown\n'))
    for obs in (middle, final):
        obs['vlan_lookups'] = before['vlan_lookups']
    attach_presence(middle, p1)
    attach_presence(final, p2)
    observations = iter([before, middle, middle, final])
    monkeypatch.setattr(ex, '_observe', lambda *a: next(observations))
    send = WorkerCLI.run_command
    failed = []
    def command(self, cmd, **kw):
        if cmd.startswith('bridge-domain') and not failed:
            failed.append(True)
            self.sent.append(cmd)
            if timeout:
                self.dead = True
                raise TimeoutError()
            return '% Invalid input'
        return send(self, cmd, **kw)
    monkeypatch.setattr(WorkerCLI, 'run_command', command)
    service = ex.ExecutionService(tmp_path / 'authority')
    prepared = service.prepare(original, {r['row_id']: 'Approve', r2['row_id']: 'Approve'}, expected_digest=original.digest)
    assert len(prepared['summary']['selected']) == 2
    spool = service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True,
                            credentials={IP: DeviceCredentials('test-user', 'hidden')}, run_root=tmp_path / 'runs')
    rows = next(spool.records())['payload']['rows']
    assert [r['status'] for r in rows] == (['Unknown', 'Not executed'] if timeout else ['Failed', 'Verified'])
    assert len(saves) == (0 if timeout else 1)


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
def test_fresh_collector_uses_service_detail_not_switchport(monkeypatch, family):
    ctx, before, _, _, entry = fixture(family)
    calls = []
    class Vlans:
        def collect(self, *a, **kw):
            return object()
        def lookup(self, session, context, vid, **kw):
            calls.append(('vlan', vid))
            return dict(status='present')
        def service_detail(self, session, context, target, sid, **kw):
            calls.append(('service', target, sid))
            return dict(status='absent', evidence=['scoped'])
    monkeypatch.setattr(ex, 'InterfaceService', lambda **kw: SimpleNamespace(collect=lambda *a, **k: []))
    monkeypatch.setattr(ex, 'VlanService', lambda **kw: Vlans())
    monkeypatch.setattr(ex.observations, 'project', lambda *a: dict(before))
    observed = ex._observe(None, ctx, None, [entry], 1)
    assert calls == ([('vlan', 445)] if family == 'ME3600X' else []) + [('service', 'GigabitEthernet0/1', 3000)]
    assert observed['service_details']['GigabitEthernet0/1:3000']['status'] == 'absent'


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
def test_aggregate_conflict_blocks(family):
    ctx, before, _, _, entry = fixture(family)
    before['audit'][0]['configuration_findings'].append(dict(code='PORT_CHANNEL_CONFIG_CONFLICT'))
    assert evc.target_blocker(entry, before, ctx) == 'PORT_CHANNEL_CONFIG_CONFLICT'

@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
@pytest.mark.parametrize('state', ['Up', 'Down'])
@pytest.mark.parametrize('sid', [445, 545])
def test_operational_presence_hardware_shape_nested_delimiters(family, state, sid):
    from orbitflow.vendors.cisco.vlans import parse_service_detail
    p = dict(evc_inputs(), interface='GigabitEthernet0/7', service_instance_id=sid)
    ctx, before, _, _, entry = fixture(family, inputs=p)
    # Sanitized ME3600X field shape; nested ! must not control success.
    text = f'Id: {sid}\nAssociated Interface: GigabitEthernet0/7\nEncapsulation: dot1q {sid} vlan protocol type 0x8100\nRewrite: ingress tag pop 1 symmetric\nState: {state}\n'
    after = observation(family, f'interface GigabitEthernet0/7\n service instance {sid} ethernet\n  encapsulation dot1q {sid}\n !\n service instance 900 ethernet\n  encapsulation dot1q 900\n !\n!\nend\n')
    after['service_details'] = {p['interface'] + ':' + str(sid): parse_service_detail(text, p['interface'], sid)}
    result = evc.verify(entry, before, after, ctx)
    assert result['status'] == 'Verified'
    assert result['evidence']['verification_scope'] == 'Presence verified'
    assert result['evidence']['service_detail']['operational_status'] == state.lower()
    assert ('SERVICE_INSTANCE_DOWN' in result['evidence']['warnings']) == (state == 'Down')
    rendered = ex.verification_text(result, entry['result'], entry['request'])
    assert 'Presence verified' in rendered


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
@pytest.mark.parametrize('text,expected', [
    ('Service Instance 3000 on interface Gi0/1 does not exist', 'Failed'),
    ('Id: 3001\nInterface: Gi0/1\nEncapsulation: dot1q 445\nState: Up', 'Failed'),
    ('Id: 3000\nInterface: Gi0/2\nEncapsulation: dot1q 445\nState: Up', 'Failed'),
    ('% Invalid input detected', 'Unknown'),
    ('% Error: request timed out', 'Unknown'),
    ('arbitrary nonempty output', 'Unknown'),
    ('Id: 3000\nState: Up', 'Unknown'),
    ('Id: 3000\nInterface: Gi0/1\nState: Up', 'Unknown'),
    ('Id: 3000\nId: 3000\nInterface: Gi0/1\nEncapsulation: dot1q 445\nState: Up', 'Unknown'),
    ('Id: 3000\nInterface: Gi0/1\nEncapsulation: dot1q 445\nState: Up\nState: Down', 'Unknown'),
])
def test_operational_presence_negative_detail(family, text, expected):
    from orbitflow.vendors.cisco.vlans import parse_service_detail
    ctx, before, r, _, entry = fixture(family)
    after = post(family, r['inputs'])
    after['service_details'] = {r['inputs']['interface'] + ':3000': parse_service_detail(text, r['inputs']['interface'], 3000)}
    assert evc.verify(entry, before, after, ctx)['status'] == expected
    after['service_details'] = {}
    assert evc.verify(entry, before, after, ctx)['status'] == 'Unknown'


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
@pytest.mark.parametrize('fault', [TimeoutError, OSError])
def test_service_detail_exchange_failure_is_unknown(monkeypatch, family, fault):
    from contextlib import contextmanager
    from orbitflow.capabilities.vlans import VlanService
    from orbitflow.vendors.cisco import vlans
    ctx, before, r, _, entry = fixture(family)
    class BrokenCLI:
        def run_command(self, command, **kw):
            assert command == 'show ethernet service instance id 3000 interface GigabitEthernet0/1 detail'
            raise fault()
    @contextmanager
    def shell(*args, **kwargs):
        yield BrokenCLI()
    monkeypatch.setattr(vlans, 'open_cisco_cli', shell)
    after = post(family, r['inputs'])
    after['service_details'] = {'GigabitEthernet0/1:3000': VlanService().service_detail(None, ctx, 'GigabitEthernet0/1', 3000)}
    assert evc.verify(entry, before, after, ctx)['status'] == 'Unknown'


@pytest.mark.parametrize('family,initialized,first,failure', [
    (family, initialized, first, '')
    for family in ['ME3600X', 'ASR920']
    for initialized in [False, True]
    for first in ['Approve', 'Pending', 'Reject', 'Blocked', 'BlockedEmpty']
] + [('ME3600X', False, 'Approve', failure) for failure in ['reject', 'timeout']]
  + [('ME3600X', False, first, 'reject') for first in ['Pending', 'Reject']])
def test_ordered_parent_initialization_lifecycle(tmp_path, monkeypatch, family, initialized, first, failure):
    """Four disjoint owner-shaped rows, real preparation/worker/report lifecycle."""
    ctx, _, _, _, _ = fixture(family)
    setup = ' switchport mode trunk\n switchport trunk allowed vlan none\n'
    prefix = 'vlan 445\n name Customer\n!\n' if family == 'ME3600X' else ''
    def observed(params, ready):
        text = prefix + 'interface GigabitEthernet0/6\n shutdown\n'
        if ready and family == 'ME3600X':
            text += setup
        for p in params:
            text += f" service instance {p['service_instance_id']} ethernet\n"
            text += ''.join('  ' + c + '\n' for c in child_commands(p))
        obs = observation(family, text + '!\nend\n')
        obs['identity'] = ex.identity(ctx)
        obs['vlan_lookups'] = {'445': dict(status='present', name='Customer', operational_status='active')}
        for p in params:
            attach_presence(obs, p)
        return obs
    params = [dict(evc_inputs(), interface='GigabitEthernet0/6', service_instance_id=si,
                   outer_vlan=str(si)) for si in [445, 30, 31, 32]]
    if first == 'Blocked':
        params[0]['outer_vlan'] = '900'
    sibling = dict(evc_inputs(), interface='GigabitEthernet0/6', service_instance_id=900, outer_vlan='900')
    # A blocked first row uses an existing initialized parent's ingress conflict.
    initial_children = [sibling] if first == 'Blocked' else []
    ready = initialized or first == 'Blocked'
    before = observed(initial_children, ready)
    if first == 'BlockedEmpty':
        attach_presence(before, params[0])  # Operational/configuration contradiction blocks only this row.
    requests = [request(template(family).id, p, row=i+5) for i, p in enumerate(params)]
    decisions = {r['row_id']: ('Approve' if i or first in {'Blocked', 'BlockedEmpty'} else first) for i, r in enumerate(requests)}
    original = plan(requests, {IP: before}, now=NOW)
    from orbitflow.configuration.v2.workbook import export_plan, output_columns
    from openpyxl import load_workbook
    snapshot, digest = original.content, original.digest
    workbook = tmp_path / 'plan.xlsx'
    export_plan(original, workbook)
    book = load_workbook(workbook)
    t = template(family)
    sheet = book[t.sheet]
    for r in requests:
        cli = sheet.cell(r['row'], len(t.columns) + output_columns(t).index('Proposed CLI') + 1).value
        assert cli == '\n'.join(original.to_dict()['results'][requests.index(r)]['commands'])
        sheet.cell(r['row'], len(t.columns) + len(output_columns(t)), decisions[r['row_id']])
    book.save(workbook)
    book.close()
    service = ex.ExecutionService(tmp_path / 'authority')
    prepared = service.prepare_workbook(original, workbook, expected_digest=original.digest)
    assert original.content == snapshot and original.digest == digest
    summary = ex.operator_summary(prepared)
    effective = [e['result']['commands'] for e in prepared['summary']['selected']]
    assert 'effective_cli' not in summary
    evaluated = ex.evaluate_decisions(original, decisions, expected_digest=original.digest, now=NOW).to_dict()['results']
    selected = [i for i, result in enumerate(evaluated) if result['deployment_ready']]
    assert selected == ([0, 1, 2, 3] if first == 'Approve' else [1, 2, 3])
    assert sum('switchport mode trunk' in evaluated[i]['commands'] for i in selected) == (family == 'ME3600X' and not ready and first in {'Approve', 'BlockedEmpty'})
    verifier = ex._verify_observed
    saves = fake_runtime(monkeypatch, ctx, before)
    monkeypatch.setattr(ex, '_verify_observed', verifier)
    captures = []
    children_now = list(initial_children)
    for index in selected:
        captures.append(observed(children_now, ready))
        children_now.append(params[index])
        ready = True
        captures.append(observed(children_now, ready))
    if failure:
        captures = [before] * 8
    observations = iter(captures)
    monkeypatch.setattr(ex, '_observe', lambda *a: next(observations))
    send = WorkerCLI.run_command
    mode_sent = [initialized or first == 'Blocked']
    def command(self, cmd, **kw):
        if failure == 'reject' and cmd.startswith('service instance ') and not mode_sent[0]:
            self.sent.append(cmd)
            return '% Invalid input: parent not initialized'
        if cmd == 'switchport mode trunk':
            if failure:
                self.sent.append(cmd)
                if failure == 'timeout':
                    self.dead = True
                    raise TimeoutError()
                return '% Invalid input'
            if mode_sent[0]:
                self.sent.append(cmd)
                return '% Command rejected: interface already in trunk mode'
            mode_sent[0] = True
        return send(self, cmd, **kw)
    monkeypatch.setattr(WorkerCLI, 'run_command', command)
    spool = service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True,
        credentials={IP: DeviceCredentials('test-user', 'hidden')}, run_root=tmp_path / 'runs')
    rows = next(spool.records())['payload']['rows']
    if failure:
        assert rows[0]['status'] == ('Unknown' if failure == 'timeout' else 'Failed')
        if failure == 'timeout':
            assert all(row['status'] != 'Verified' for row in rows[1:])
            assert not any(c.startswith('service instance ') for c in WorkerCLI.instances[0].sent)
        assert WorkerCLI.instances[0].sent.count('switchport mode trunk') == (first == 'Approve')
        assert len(saves) == (0 if failure == 'timeout' else 1)
        return
    assert all(row['status'] == 'Verified' for row in rows)
    assert WorkerCLI.instances[0].sent == [c for i in selected for c in evaluated[i]['commands']]
    assert len(saves) == 1
    assert ex.export_results(spool.path, tmp_path / 'result.xlsx').exists()


def test_parent_readiness_is_not_an_execution_gate():
    ctx, initialized, r, _, entry = fixture('ME3600X')
    empty = observation('ME3600X', 'vlan 445\n name Customer\n!\ninterface GigabitEthernet0/1\n shutdown\n!\n')
    empty['vlan_lookups'] = initialized['vlan_lookups']
    assert not evc.target_blocker(entry, empty, ctx)
    entry['result']['commands'][2:2] = ['switchport mode trunk', 'switchport trunk allowed vlan none']
    assert not evc.target_blocker(entry, empty, ctx)
    assert not evc.target_blocker(entry, initialized, ctx)
    partial = observation('ME3600X', 'vlan 445\n name Customer\n!\ninterface GigabitEthernet0/1\n switchport mode trunk\n!\n')
    partial['vlan_lookups'] = initialized['vlan_lookups']
    assert not evc.target_blocker(entry, partial, ctx)


def test_parent_initialization_decisions_order_and_digest():
    from orbitflow.configuration.v2.planning import BulkPlan
    from orbitflow.configuration.plans import canonical
    obs = observation('ME3600X', 'vlan 445\n name Customer\n!\ninterface Gi0/6\n!\n')
    requests = [request(template('ME3600X').id, dict(evc_inputs(), interface='Gi0/6',
                service_instance_id=si, outer_vlan=str(si)), row=i+3) for i, si in enumerate([445, 30, 31, 32])]
    original = plan(requests, {IP: obs}, now=NOW)
    approvals = {r['row_id']: 'Approve' for r in requests}
    all_rows = ex.evaluate_decisions(original, approvals, expected_digest=original.digest, now=NOW).to_dict()['results']
    later = ex.evaluate_decisions(original, {**approvals, requests[0]['row_id']: 'Reject'},
                                  expected_digest=original.digest, now=NOW).to_dict()['results']
    assert ['switchport mode trunk' in r['commands'] for r in all_rows] == [True, False, False, False]
    assert [r['commands'] for r in later] == [r['commands'] for r in original.to_dict()['results']]
    changed = original.to_dict()
    changed['results'][0]['commands'].remove('switchport mode trunk')
    with pytest.raises(PlanError, match='tampered planning result'):
        BulkPlan(canonical(changed))
    with pytest.raises(PlanError, match='digest mismatch'):
        ex.evaluate_decisions(original, approvals, expected_digest='0' * 64, now=NOW)
    reversed_requests = [request(template('ME3600X').id, r['inputs'], row=i+3) for i, r in enumerate(reversed(requests))]
    reversed_plan = plan(reversed_requests, {IP: obs}, now=NOW,
                         decisions={r['row_id']: 'Approve' for r in reversed_requests})
    assert reversed_plan.digest != original.digest
    assert ['switchport mode trunk' in r['commands'] for r in reversed_plan.to_dict()['results']] == [True, False, False, False]


def test_pending_plan_mixed_decisions_preserve_original_cli(tmp_path):
    from orbitflow.configuration.v2.workbook import export_plan, consume_decisions, output_columns
    from openpyxl import load_workbook
    obs = observation('ME3600X', 'vlan 445\n name Customer\n!\ninterface Gi0/6\n!\n')
    requests = [request(template('ME3600X').id, dict(evc_inputs(), interface='Gi0/6', service_instance_id=si,
                               outer_vlan=str(si)), row=i+3) for i, si in enumerate([445, 30, 31, 32])]
    original = plan(requests, {IP: obs}, now=NOW)
    assert all(r['decision'] == 'Pending' for r in original.to_dict()['results'])
    assert ['switchport mode trunk' in r['commands'] for r in original.to_dict()['results']] == [True, False, False, False]
    path = tmp_path / 'plan.xlsx'
    export_plan(original, path)
    book = load_workbook(path)
    t = template('ME3600X')
    for r, decision in zip(requests, ['Pending', 'Approve', 'Reject', 'Approve']):
        book[t.sheet].cell(r['row'], len(t.columns) + len(output_columns(t)), decision)
    book.save(path)
    book.close()
    reevaluated = consume_decisions(original, path, expected_digest=original.digest, now=NOW)
    prepared = ex.ExecutionService(tmp_path / 'authority').prepare_workbook(original, path, expected_digest=original.digest)
    selected = prepared['summary']['selected']
    assert [e['request']['row'] for e in selected] == [4, 6]
    assert all(not e['result']['dependencies'] for e in selected)
    assert [e['result']['commands'] for e in selected] == [r['commands'] for r in reevaluated.to_dict()['results'] if r['deployment_ready']]
    text = ex.render_summary(ex.operator_summary(prepared))
    assert 'switchport mode trunk' not in text
    assert 'switchport trunk allowed vlan none' not in text
    assert all('switchport mode trunk' not in e['result']['commands'] for e in selected)
    changed = ex.ExecutionService(tmp_path / 'other').prepare(original,
        {r['row_id']: 'Approve' for r in requests}, expected_digest=original.digest)
    assert changed['summary_digest'] != prepared['summary_digest']


def test_pending_parent_setup_is_scoped_by_device_and_canonical_parent(tmp_path):
    other_ip = '192.0.2.92'
    obs = observation('ME3600X', 'vlan 445\n name Customer\n!\ninterface Gi0/6\n!\ninterface Gi0/7\n!\n')
    requests = [request(template('ME3600X').id, dict(evc_inputs(), management_ip=ip,
        interface=parent, service_instance_id=si, outer_vlan=str(si)), row=row)
        for row, ip, parent, si in [(5, IP, 'Gi0/6', 445),
            (6, IP, 'GigabitEthernet0/6', 30), (7, IP, 'Gi0/7', 31),
            (8, other_ip, 'Gi0/6', 32)]]
    original = plan(requests, {IP: obs, other_ip: obs}, now=NOW)
    rows = original.to_dict()['results']
    assert ['switchport mode trunk' in r['commands'] for r in rows] == [True, False, True, True]
    assert 'PROPOSED_PARENT_SETUP_PREREQUISITE' in rows[1]['warnings']
    for first in ['Pending', 'Reject', 'Approve']:
        evaluated = ex.evaluate_decisions(original, {requests[0]['row_id']: first,
            requests[1]['row_id']: 'Approve'}, expected_digest=original.digest, now=NOW)
        assert [r['commands'] for r in evaluated.to_dict()['results']] == [r['commands'] for r in rows]
