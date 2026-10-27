"""Synthetic retained-run pipeline; no private captures or device connections."""
from copy import deepcopy
from io import StringIO
import json
import socket
import tracemalloc

import pytest
from openpyxl import load_workbook

from test_bulk_configuration_v2 import NOW, template, inputs, observed
from orbitflow.configuration.plans import canonical, PlanError
from orbitflow.configuration.v2.observation import request_digest
from orbitflow.configuration.v2.planning import load, evaluate_decisions
from orbitflow.configuration.v2.recovery import recover
from orbitflow.configuration.v2.workbook import _base, read_input, consume_decisions
from orbitflow.configuration.execution import ExecutionService
from orbitflow.execution import DeviceOutcome
from orbitflow.result_spool import ResultSpool
from orbitflow.indexed_json import IndexedJSON, chunks, digest, read_json


def test_indexed_json_canonical_roundtrip_and_isolation():
    source = {'z': {'x': ['é', 1, False, None]}, 'a': {}}
    values = IndexedJSON(source.items())
    try:
        data = {'observations': values, 'requests': []}
        text = ''.join(chunks(data))
        assert text == canonical({'observations': source, 'requests': []})
        assert digest(data) == __import__('hashlib').sha256(text.encode()).hexdigest()
        copy = values['z']
        copy['x'].append('changed')
        assert values['z'] == source['z']
        recovered = read_json(StringIO(text))
        assert dict(recovered['observations']) == source
        recovered['observations'].close()
    finally:
        values.close()


@pytest.mark.parametrize('text', ['{"a":1,"a":2}', '{"a":NaN}', '[1,]', '"unterminated', '{} {}'])
def test_streaming_reader_rejects_invalid_json(text):
    with pytest.raises(ValueError):
        read_json(StringIO(text))


@pytest.mark.parametrize('count,padding', [(88, 900000), (500, 72000), (1500, 24000)])
def test_full_recovery_review_preparation_pipeline(tmp_path, monkeypatch, count, padding, record_property):
    monkeypatch.setattr(socket, 'create_connection', lambda *a, **k: pytest.fail('network forbidden'))
    t = template('C3750X', 'VLAN Database')
    book = _base()
    ips = [f'198.18.{i // 250}.{i % 250 + 1}' for i in range(count)]
    for ip in ips:
        p = dict(inputs(t), management_ip=ip)
        book[t.sheet].append([p.get(k, '') for k in t.columns])
    workbook = tmp_path / 'input.xlsx'
    book.save(workbook)
    book.close()
    requests = read_input(workbook)
    spool = ResultSpool.create(tmp_path / 'runs', 'configuration_observation_v2', count)
    spool.bind(dict(request_digest=request_digest(requests), workbook_digest=requests.source_digest))
    with spool.collection():
        for i, ip in enumerate(ips, 1):
            obs = observed(t)
            obs['identity'].update(management_ip=ip, inventory_id=f'device-{i}',
                                   hostname=f'device-{i}', serial_number=f'SYNTHETIC{i}')
            obs['evidence'] = [{'source_filename': 'synthetic-config', 'line': 1,
                                'excerpt': '! synthetic ' + 'x' * padding}]
            spool.append(DeviceOutcome(i, obs), target=ip)
    output = tmp_path / 'plan.json'
    tracemalloc.start()
    try:
        result = recover(workbook, spool.path, output, now=NOW)
        assert output.stat().st_size > 32 * 1024 * 1024
        restored = load(output)
        assert restored.digest == result.digest
        data = restored.view()
        assert data['requests'] == requests
        assert list(data['observations']) == sorted(ips)
        for r in data['results']:
            assert r['commands'] == ['configure terminal', 'vlan 4002', 'name Customer', 'end']
            assert r['decision'] == 'Pending'
            assert not r['execution_authorized']
        intent = consume_decisions(restored, output.with_suffix('.xlsx'),
                                   expected_digest=restored.digest, now=NOW)
        assert len(intent.to_dict()['results']) == count
        decisions = {r['row_id']: 'Approve' for r in requests}
        approved = evaluate_decisions(restored, decisions, expected_digest=restored.digest, now=NOW)
        assert all(r['deployment_ready'] for r in approved.to_dict()['results'])
        service = ExecutionService(tmp_path / 'authority')
        prepared = service.prepare(restored, decisions, expected_digest=restored.digest)
        assert len(prepared['summary']['selected']) == count
        assert service.prepare(restored, decisions, expected_digest=restored.digest)['reference'] == prepared['reference']
        assert len(list((tmp_path / 'authority').glob('confirmation-*.json'))) == 1
        # Exercise the confirmed consumer without transport or configuration.
        monkeypatch.setattr('orbitflow.configuration.execution._run', lambda *a, **k: spool)
        assert service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True) is spool
        with pytest.raises(PlanError, match='consumed'):
            service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    record_property('full_pipeline_peak_python_bytes', peak)
    assert peak < 96 * 1024 * 1024
    assert spool.manifest['collection_complete'] is True
    assert not list(spool.path.glob('read-*'))
    review = load_workbook(output.with_suffix('.xlsx'), read_only=True)
    try:
        assert review['Introduction']['B1'].value == result.digest
        assert review.sheetnames[-2:] == ['Evidence Details', 'Summary']
        assert sum(1 for _ in review['Summary'].iter_rows()) == count + 1
    finally:
        review.close()


def test_indexed_transaction_failure_cleans_scratch(tmp_path, monkeypatch):
    import orbitflow.indexed_json as module
    from tempfile import TemporaryDirectory
    monkeypatch.setattr(module, 'TemporaryDirectory', lambda **kw: TemporaryDirectory(dir=tmp_path, **kw))
    def interrupted():
        yield 'committed-only-if-complete', {'safe': True}
        raise OSError('synthetic disk failure')
    with pytest.raises(OSError):
        IndexedJSON(interrupted())
    assert list(tmp_path.iterdir()) == []
    with pytest.raises(Exception):
        IndexedJSON([('same', 1), ('same', 2)])
    assert list(tmp_path.iterdir()) == []


def test_indexed_value_and_disk_bounds_leave_no_scratch(tmp_path, monkeypatch):
    import orbitflow.indexed_json as module
    from tempfile import TemporaryDirectory
    monkeypatch.setattr(module, 'TemporaryDirectory', lambda **kw: TemporaryDirectory(dir=tmp_path, **kw))
    monkeypatch.setattr(module, 'MAX_WORKING_VALUE_BYTES', 32)
    with pytest.raises(ValueError, match='value size'):
        IndexedJSON([('large', 'x' * 33)])
    monkeypatch.setattr(module, 'MAX_TOTAL_BYTES', 32)
    with pytest.raises(ValueError, match='disk limit'):
        IndexedJSON((str(i), 'x' * 20) for i in range(2))
    assert list(tmp_path.iterdir()) == []


def test_unsafe_observation_rejected_before_sqlite_persistence(tmp_path, monkeypatch):
    from orbitflow.configuration.v2 import planning
    t = template('C3750X', 'VLAN Database')
    obs = observed(t)
    obs['evidence'] = [{'source_filename': 'synthetic', 'line': 1, 'excerpt': 'username operator secret synthetic-only'}]
    with pytest.raises(PlanError, match='unsafe'):
        planning.plan([planning.request(t.id, inputs(t))], {'192.0.2.10': obs}, now=NOW)


def test_large_scalar_chunk_boundary_and_container_bounds(monkeypatch):
    import orbitflow.indexed_json as module
    value = {'observations': {'device': {'text': 'x' * 70000}}, 'number': 123}
    parsed = read_json(StringIO(canonical(value)))
    assert dict(parsed['observations']) == value['observations']
    assert parsed['number'] == 123
    parsed['observations'].close()
    monkeypatch.setattr(module, 'MAX_WORKING_VALUE_BYTES', 32)
    with pytest.raises(ValueError, match='size limit'):
        read_json(StringIO(canonical({'observations': {'device': {'text': 'x' * 100}}})))


def test_control_character_sanitizer_equivalence():
    from orbitflow.vendors.configuration_facts import _safe_source
    for code in range(128):
        source = 'synthetic' + chr(code) + 'evidence'
        omitted = code < 32 and chr(code) not in '\t\n\r' or code == 127
        assert (_safe_source(source) == '[configuration evidence omitted: unsafe control characters]') == omitted


def test_immutable_index_concurrent_readers():
    from concurrent.futures import ThreadPoolExecutor
    values = IndexedJSON((str(i), {'position': i}) for i in range(100))
    try:
        with ThreadPoolExecutor(max_workers=5) as executor:
            assert list(executor.map(lambda i: values[str(i)], range(100))) == [{'position': i} for i in range(100)]
    finally:
        values.close()


def test_atomic_plan_save_failure_retains_destination_absence(tmp_path, monkeypatch):
    from orbitflow.configuration.v2 import planning
    import os
    t = template('C3750X', 'VLAN Database')
    original = planning.plan([planning.request(t.id, inputs(t))], now=NOW)
    output = tmp_path / 'plan.json'
    monkeypatch.setattr(os, 'link', lambda *a: (_ for _ in ()).throw(OSError('synthetic disk full')))
    with pytest.raises(OSError):
        planning.save(original, output)
    assert list(tmp_path.iterdir()) == []


def test_untrusted_result_strings_never_enter_scratch_sqlite(tmp_path, monkeypatch):
    import sqlite3
    from orbitflow.configuration.v2 import planning
    import orbitflow.indexed_json as module
    t = template('C3750X', 'VLAN Database')
    original = planning.plan([planning.request(t.id, inputs(t))], now=NOW)
    data = original.to_dict()
    data['results'][0]['commands'] = ['username operator secret synthetic-untrusted-value']
    output = tmp_path / 'tampered.json'
    output.write_text(canonical({'digest': original.digest, 'plan': data}))
    connect = sqlite3.connect
    class AuditedConnection(sqlite3.Connection):
        def execute(self, sql, parameters=()):
            assert 'synthetic-untrusted-value' not in repr(parameters)
            return super().execute(sql, parameters)
    monkeypatch.setattr(module.sqlite3, 'connect', lambda *a, **k: connect(*a, factory=AuditedConnection, **k))
    with pytest.raises(PlanError, match='tampered'):
        load(output)


@pytest.mark.parametrize('width', [1, 2, 7, 17])
def test_streamed_numeric_and_escaped_string_boundaries(width):
    class SplitStream(StringIO):
        def read(self, size=-1):
            return super().read(min(size, width))
    data = {'values': [1.2, -23.4e-25, 1e20, True, False, None, 'escaped"\\\ntext']}
    parsed = read_json(SplitStream(canonical(data)))
    assert parsed == data


@pytest.mark.parametrize('status', ['live_read_only', 'evidence_unavailable'])
def test_invalid_observation_schema_never_reaches_scratch_values(tmp_path, monkeypatch, status):
    import sqlite3
    import orbitflow.indexed_json as module
    from orbitflow.configuration.v2 import planning
    t = template('C3750X', 'VLAN Database')
    obs = observed(t)
    obs['identity']['platform'] = 'unsupported-synthetic'
    obs['status'] = status
    obs['evidence'] = ['snmp-server synthetic-untrusted-value']
    connect = sqlite3.connect
    class AuditedConnection(sqlite3.Connection):
        def execute(self, sql, parameters=()):
            assert 'synthetic-untrusted-value' not in repr(parameters)
            return super().execute(sql, parameters)
    monkeypatch.setattr(module.sqlite3, 'connect', lambda *a, **k: connect(*a, factory=AuditedConnection, **k))
    with pytest.raises(PlanError):
        planning.plan([planning.request(t.id, inputs(t))], {'192.0.2.10': obs}, now=NOW)
