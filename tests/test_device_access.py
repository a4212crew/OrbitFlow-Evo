"""Deterministic access tests; all credentials and devices are synthetic."""
from io import StringIO
from contextlib import closing
from threading import Barrier, Lock
from dataclasses import replace
from pathlib import Path
from unittest.mock import MagicMock
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from paramiko import AuthenticationException, SSHException
import pytest

from orbitflow import configuration_backup as backup
from orbitflow.config import ExecutionConfig
from orbitflow.execution import execute_devices
from orbitflow.targets import load_targets
from orbitflow.transport import DeviceCredentials, DeviceSession, TransportConfig, connect_device
from orbitflow.transport.exceptions import (
    ConnectionRetryExhausted, DeviceConnectionError, TransportConfigurationError,
    UnsupportedPlatformError, TunnelTimeout,
    ConnectionCleanupError,
)
from orbitflow.transport.policy import ConnectionStartPacer, connection_pacer
from orbitflow.vendors.common import DeviceCLI, InteractiveCLITimeout
from orbitflow.vendors.cisco.ios import CiscoIOSCLI
from orbitflow.vendors.privilege import MissingEnableSecret, EnableAuthenticationFailed
from orbitflow.vendors.privilege import ensure_privileged
from orbitflow.vendors.ubiquiti.prompts import extract_edgeswitch_hostname
from test_capability_context import DEVICES
from test_configuration_backup import device_session, CONFIG


SECRET = 'synthetic-enable-only'
TRANSPORT = TransportConfig('proxy', 'cluster', 'bastion', 'user')


@pytest.mark.parametrize('header', [None, 'Secret', 'secret'])
@pytest.mark.parametrize('value', [None, SECRET])
def test_optional_excel_secret(tmp_path, header, value):
    book = Workbook()
    book.active.append(['management_ip', 'username', 'password'] + ([header] if header else []))
    book.active.append(['192.0.2.1', 'synthetic-user', 'synthetic-password'] + ([value] if header else []))
    path = tmp_path / 'input.xlsx'
    book.save(path)
    book.close()
    target, = load_targets(path)
    assert target.get('secret', '') == (value or '' if header else '')
    assert SECRET not in repr(DeviceCredentials('user', 'password', secret=SECRET))
    assert 'password' not in repr(DeviceCredentials('user', 'password', secret=SECRET))


@pytest.mark.parametrize('end', ['>', '#'])
@pytest.mark.parametrize('annotation', ['', ' (arbitrary annotation)', ' (x (nested) y)', ' (changed # > text)'])
def test_structural_prompt(end, annotation):
    assert extract_edgeswitch_hostname(f'(HOSTNAME{annotation}) {end}') == 'HOSTNAME'
    assert extract_edgeswitch_hostname(f'HOSTNAME{end}') == 'HOSTNAME'


@pytest.mark.parametrize('prompt', ['(host (x) #', '(host x) #', '(host (x)))>',
    '(host (x) (y)) #', '((x)) #', '(host (x)junk) #', '(host (x\ny)) #', '(host) junk#'])
def test_malformed_prompt(prompt):
    with pytest.raises(ValueError, match='unrecognized EdgeSwitch prompt'):
        extract_edgeswitch_hostname(prompt)


def enable_session(device, *, secret=SECRET, response='success', enabled_prompt=None):
    session, client, channel = device_session(device)
    channel.prompt = '(sw (arbitrary (nested) annotation)) >' if device[1] == 'ubiquiti_edgeswitch' else 'sw>'
    session.secret = secret
    original = channel.sendall
    def send(data):
        command = data.decode().strip()
        if command == 'enable':
            channel.sent.append(command)
            channel.pending = [b'enable\r\nPass', b'word: ']
        elif command == SECRET:
            channel.sent.append(command)
            if response == 'success':
                channel.prompt = enabled_prompt if enabled_prompt is not None else channel.prompt[:-1] + '#'
                channel.pending = [(SECRET + '\r\n' + channel.prompt).encode()]
                # Later command echoes use the learned prompt without surrounding whitespace.
                channel.prompt = channel.prompt.strip()
            elif response == 'repeat':
                channel.pending = [b'Password: ']
            elif response == 'timeout':
                channel.pending = [TimeoutError(SECRET)]
            elif response == 'closed':
                channel.pending = [b'']
            else:
                channel.pending = [(SECRET + '\r\n% Access denied\r\n' + channel.prompt).encode()]
        else:
            original(data)
    channel.sendall = send
    channel.outputs['show running-config'] = CONFIG
    return session, client, channel


ACCESS_DEVICES = [d for d in DEVICES if d[1] in {'cisco_ios', 'cisco_xe', 'ubiquiti_edgeswitch'}]


@pytest.mark.parametrize('cli_type', [DeviceCLI, CiscoIOSCLI])
@pytest.mark.parametrize('device', ACCESS_DEVICES, ids=lambda d: d[2])
@pytest.mark.parametrize('fragmented', [False, True])
@pytest.mark.parametrize('scenario', [
    'echo', 'stale_before', 'stale_after', 'changed_before', 'changed_after',
    'denied', 'repeat', 'malformed', 'timeout', 'closed',
])
def test_enable_interactive_state_machine(cli_type, device, fragmented, scenario):
    session, client, channel = enable_session(device)
    channel.prompt = channel.prompt.replace('annotation', 'annotati\u00f6n')
    user_prompt = channel.prompt
    privileged_prompt = user_prompt[:-1] + '#'
    pre = f'{user_prompt}enable\r\n'
    if scenario == 'stale_before':
        pre = f'{user_prompt}\r\n{user_prompt}\r\n' + pre + f'{user_prompt}\r\n'
    if scenario == 'changed_before':
        pre += user_prompt.replace('sw', 'other') + '\r\n'
    pre += 'Password: '
    post = SECRET + '\r\n'
    if scenario == 'stale_after':
        post += f'{user_prompt}\r\n{user_prompt}\r\n'
    if scenario == 'changed_after':
        post += user_prompt.replace('sw', 'other') + '\r\n'
    if scenario == 'denied':
        post += 'Access denied\r\n'
    if scenario == 'repeat':
        post += 'Password: \r\n'
    elif scenario == 'malformed':
        post += '(sw (broken) #\r\n'
    post += privileged_prompt

    def chunks(text):
        data = text.encode()
        return [data[i:i+1] for i in range(len(data))] if fragmented else [data]

    original = channel.sendall
    def send(data):
        if data == b'enable\n':
            channel.sent.append('enable')
            channel.pending = chunks(pre)
        elif data == (SECRET + '\n').encode():
            channel.sent.append(SECRET)
            channel.prompt = privileged_prompt
            channel.pending = ([TimeoutError(SECRET)] if scenario == 'timeout' else
                               [b''] if scenario == 'closed' else chunks(post))
        else:
            original(data)
    channel.sendall = send
    if scenario in {'echo', 'stale_before', 'stale_after'}:
        with session, closing(cli_type(session)) as cli:
            assert cli.prompt == privileged_prompt
            assert SECRET not in cli.run_command('show running-config')
        assert channel.sent.count(SECRET) == 1
    else:
        with pytest.raises(EnableAuthenticationFailed) as caught, session:
            cli_type(session)
        assert str(caught.value) == 'Enable authentication failed.'
        assert caught.value.__cause__ is None
        assert 'terminal length 0' not in channel.sent
        assert channel.sent.count(SECRET) == (0 if scenario == 'changed_before' else 1)
        if fragmented and scenario not in {'timeout', 'closed'}:
            assert channel.pending, 'failure must precede the later privileged prompt'
    assert channel.sent.count('enable') == 1
    assert client.closes == channel.close_calls == 1
    assert session.secret is None


@pytest.mark.parametrize('prompt', ['sw>', '(sw (arbitrary (nested) annotation)) >'])
@pytest.mark.parametrize('answered', [False, True])
def test_enable_progress_does_not_reset_deadline(monkeypatch, prompt, answered):
    channel = MagicMock()
    channel.recv.side_effect = [b'Password:' if answered else prompt.encode(), prompt.encode()]
    clock = iter([0, 1, 2, 3])
    monkeypatch.setattr('orbitflow.vendors.privilege.time.monotonic', lambda: next(clock))
    with pytest.raises(EnableAuthenticationFailed, match='^Enable authentication failed\\.$'):
        ensure_privileged(channel, prompt, SECRET, timeout=3)
    assert channel.recv.call_count == 2
    assert [call.args[0] for call in channel.sendall.call_args_list] == (
        [b'enable\n', (SECRET + '\n').encode()] if answered else [b'enable\n'])


@pytest.mark.parametrize('cli_type', [DeviceCLI, CiscoIOSCLI])
@pytest.mark.parametrize('enabled_prompt', ['sw#', '  sw#\t'])
def test_enable_cisco_prompt_formatting(cli_type, enabled_prompt):
    session, client, channel = enable_session(ACCESS_DEVICES[0], enabled_prompt=enabled_prompt)
    with session:
        cli = cli_type(session)
        assert cli.prompt == 'sw#'
        cli.close()
    assert channel.sent[:4] == ['', 'enable', SECRET, 'terminal length 0']
    assert client.closes == channel.close_calls == 1


@pytest.mark.parametrize('enabled_prompt', [
    '(sw (different annotation))#', '(sw (new (nested) annotation))\t#',
    '(sw)#', 'sw#', '  (sw (changed # > text))  #\t',
])
def test_enable_edgeswitch_prompt_presentation_change(enabled_prompt):
    edge = next(d for d in ACCESS_DEVICES if d[1] == 'ubiquiti_edgeswitch')
    session, client, channel = enable_session(edge, enabled_prompt=enabled_prompt)
    with session, DeviceCLI(session) as cli:
        assert cli.prompt == enabled_prompt.strip()
        assert extract_edgeswitch_hostname(cli.prompt) == 'sw'
        assert SECRET not in cli.run_command('show running-config')
    assert channel.sent[:4] == ['', 'enable', SECRET, 'terminal length 0']
    assert channel.sent.count('enable') == channel.sent.count(SECRET) == 1
    assert client.closes == channel.close_calls == 1


@pytest.mark.parametrize('cli_type', [DeviceCLI, CiscoIOSCLI])
@pytest.mark.parametrize('device', ACCESS_DEVICES, ids=lambda d: d[2])
@pytest.mark.parametrize('enabled_prompt', [
    'other#', '(other (arbitrary annotation)) #',
    '(sw (unbalanced) #', '(sw annotation) #', 'sw(config)#',
    '(sw (changed annotation)) >', '% Access denied\r\nsw#',
])
def test_enable_rejects_changed_identity_or_invalid_exec(cli_type, device, enabled_prompt):
    session, client, channel = enable_session(device, enabled_prompt=enabled_prompt)
    with pytest.raises(EnableAuthenticationFailed) as caught, session:
        cli_type(session)
    assert str(caught.value) == 'Enable authentication failed.'
    assert caught.value.__cause__ is None
    assert 'terminal length 0' not in channel.sent
    assert channel.sent.count('enable') == channel.sent.count(SECRET) == 1
    assert client.closes == channel.close_calls == 1


@pytest.mark.parametrize('device', ACCESS_DEVICES, ids=lambda d: d[2])
def test_enable_backup_end_to_end(tmp_path, monkeypatch, device):
    session, client, channel = enable_session(device)
    channel.outputs['show running-config'] += SECRET + '\nsynthetic-login-password\n'
    def connect(ip, credentials, config):
        assert credentials.secret == SECRET
        return session
    monkeypatch.setattr(backup, 'connect_device', connect)
    output = StringIO()
    folder = backup.run_backup([
        dict(management_ip='192.0.2.1', username='synthetic-login-user',
             password='synthetic-login-password', secret=SECRET)
    ], TRANSPORT, backups_dir=tmp_path/'backups', inventory_path=tmp_path/'inventory.json',
        log_root=tmp_path/'logs', output=output, execution_config=ExecutionConfig(1))
    captures = list(folder.glob('*.txt'))
    assert len(captures) == 1
    assert channel.sent[:4] == ['', 'enable', SECRET, 'terminal length 0']
    assert channel.sent.count('enable') == 1
    assert client.closes == channel.close_calls == 1
    assert session.secret is None
    safe = output.getvalue() + ''.join(p.read_text() for p in tmp_path.rglob('*')
                                    if p.suffix in {'.json', '.log', '.txt', '.jsonl'})
    with ZipFile(folder/'failed_devices.xlsx') as archive:
        safe += ''.join(archive.read(name).decode() for name in archive.namelist())
    assert SECRET not in safe and 'synthetic-login-password' not in safe
    assert not list(tmp_path.rglob('results.jsonl'))


@pytest.mark.parametrize('cli_type', [DeviceCLI, CiscoIOSCLI])
@pytest.mark.parametrize('response', ['success', 'wrong', 'repeat', 'timeout', 'closed', 'missing'])
def test_enable_outcomes_cleanup(cli_type, response):
    session, client, channel = enable_session(ACCESS_DEVICES[0],
        secret=None if response == 'missing' else SECRET, response=response)
    if response == 'success':
        with session:
            cli = cli_type(session)
            assert cli.prompt.endswith('#')
            cli.close()
    else:
        expected = MissingEnableSecret if response == 'missing' else EnableAuthenticationFailed
        with pytest.raises(expected) as caught, session:
            cli_type(session)
        assert SECRET not in str(caught.value)
        assert caught.value.__cause__ is None
        assert 'terminal length 0' not in channel.sent
    assert client.closes == channel.close_calls == 1


@pytest.mark.parametrize('response', ['wrong', 'missing'])
def test_failed_workbook_privilege_reason(tmp_path, monkeypatch, response):
    session, client, channel = enable_session(ACCESS_DEVICES[0],
        secret=None if response == 'missing' else SECRET, response=response)
    monkeypatch.setattr(backup, 'connect_device', lambda *args: session)
    folder = backup.run_backup([dict(management_ip='192.0.2.1', username='user',
        password='synthetic-password', secret=SECRET)], TRANSPORT,
        backups_dir=tmp_path/'backups', inventory_path=tmp_path/'inventory.json',
        log_root=tmp_path/'logs', output=StringIO(), execution_config=ExecutionConfig(1))
    book = load_workbook(folder/'failed_devices.xlsx')
    rows = list(book.active.values)
    book.close()
    assert rows[1][-1] == ('Enable secret is required for privileged EXEC.'
                          if response == 'missing' else 'Enable authentication failed.')
    assert SECRET not in repr(rows)
    assert client.closes == channel.close_calls == 1


@pytest.fixture(autouse=True)
def retry_sleeps(monkeypatch):
    sleeps = []
    monkeypatch.setattr('orbitflow.transport.policy.sleep', sleeps.append)
    return sleeps


@pytest.mark.parametrize('system', ['windows', 'linux'])
@pytest.mark.parametrize('failure', [TimeoutError, ConnectionResetError, ConnectionRefusedError, TunnelTimeout])
@pytest.mark.parametrize('recover', [True, False])
def test_bounded_transient_retry(monkeypatch, system, failure, recover, retry_sleeps):
    import importlib
    module = importlib.import_module('orbitflow.transport.' + system)
    attempts, starts = [], []
    session = DeviceSession(object(), lambda: None)
    def connect(*args):
        attempts.append(1)
        if len(attempts) == 1 or not recover:
            raise failure('synthetic-password')
        return session
    monkeypatch.setattr(module, 'connect_' + system, connect)
    monkeypatch.setattr('orbitflow.transport.wait_for_connection_start', lambda: starts.append(1))
    if recover:
        assert connect_device('192.0.2.1', DeviceCredentials('user', secret=SECRET), TRANSPORT, system=system) is session
        assert session.secret == SECRET
    else:
        with pytest.raises(ConnectionRetryExhausted) as caught:
            connect_device('192.0.2.1', DeviceCredentials('user'), TRANSPORT, system=system)
        assert backup.failure_reason('connect', caught.value) == 'Transient connection failure after two attempts.'
    assert len(attempts) == len(starts) == 2
    assert retry_sleeps == [5.0]


@pytest.mark.parametrize('error', [AuthenticationException(SECRET), SSHException(SECRET),
    MissingEnableSecret(SECRET), EnableAuthenticationFailed(SECRET), ValueError(SECRET),
    TransportConfigurationError(SECRET), UnsupportedPlatformError(SECRET),
    InteractiveCLITimeout(SECRET), backup.ConfigurationCommandRejected(SECRET)])
def test_no_retry_of_nontransient_errors(monkeypatch, error, retry_sleeps):
    calls = []
    def connect(*args):
        calls.append(1)
        raise error
    monkeypatch.setattr('orbitflow.transport.windows.connect_windows', connect)
    monkeypatch.setattr('orbitflow.transport.wait_for_connection_start', lambda: None)
    with pytest.raises(type(error)):
        connect_device('192.0.2.1', DeviceCredentials('user'), TRANSPORT, system='windows')
    assert calls == [1]
    assert retry_sleeps == []


@pytest.mark.parametrize('interval', [-1, True, float('nan'), float('inf'), '1'])
def test_invalid_pacing(interval):
    with pytest.raises(ValueError):
        ExecutionConfig(connection_start_interval=interval)


def test_pacing_spaces_starts_without_serializing_active_devices(monkeypatch):
    now, starts = [0.0], []
    def sleep(delay):
        now[0] += delay
    recording_lock = Lock()
    class RecordingPacer(ConnectionStartPacer):
        def wait(self):
            with recording_lock:
                super().wait()
                starts.append(now[0])
    pacer = RecordingPacer(1.0, clock=lambda: now[0], sleep=sleep)
    def make_pacer(interval):
        assert interval == 1.0
        return pacer
    monkeypatch.setattr('orbitflow.execution.ConnectionStartPacer', make_pacer)
    lock, barrier = Lock(), Barrier(5)
    active, peak = 0, 0
    def connect(*args):
        nonlocal active, peak
        with lock:
            active += 1
            peak = max(peak, active)
        barrier.wait(timeout=5)
        def cleanup():
            nonlocal active
            with lock:
                active -= 1
        return DeviceSession(object(), cleanup)
    monkeypatch.setattr('orbitflow.transport.windows.connect_windows', connect)
    def worker(target):
        with connect_device(str(target), DeviceCredentials('user'), TRANSPORT, system='windows'):
            return target
    outcomes = execute_devices(range(5), worker, config=ExecutionConfig())
    assert all(not item.error_category for item in outcomes)
    assert starts == [0, 1, 2, 3, 4]
    assert peak == 5 and active == 0
    assert connection_pacer.get() is None


@pytest.mark.parametrize('system', ['windows', 'linux'])
@pytest.mark.parametrize('cleanup_fails', [False, True])
def test_real_transport_retry_cleans_every_resource_first(monkeypatch, system, cleanup_fails, retry_sleeps):
    import importlib
    module = importlib.import_module('orbitflow.transport.' + system)
    monkeypatch.setattr('orbitflow.transport.wait_for_connection_start', lambda: None)
    targets = [MagicMock(), MagicMock()]
    targets[0].connect.side_effect = TimeoutError(SECRET)
    if cleanup_fails:
        targets[0].close.side_effect = OSError(SECRET)
    resources, attempts = [], []
    if system == 'windows':
        processes = [MagicMock(), MagicMock()]
        sockets = [MagicMock(), MagicMock()]
        for process in processes:
            process.poll.return_value = None
        monkeypatch.setattr(module.subprocess, 'Popen', MagicMock(side_effect=processes))
        monkeypatch.setattr(module, '_free_local_port', lambda: 49152)
        monkeypatch.setattr(module, '_open_forwarded_socket', MagicMock(side_effect=sockets))
        clients = iter(targets)
        resources = [sockets[0].close, processes[0].terminate]
    else:
        bastions, proxies = [MagicMock(), MagicMock()], [MagicMock(), MagicMock()]
        monkeypatch.setattr(module.paramiko.PKey, 'from_path', MagicMock())
        monkeypatch.setattr(module.paramiko, 'ProxyCommand', MagicMock(side_effect=proxies))
        clients = iter([bastions[0], targets[0], bastions[1], targets[1]])
        resources = [bastions[0].get_transport.return_value.open_channel.return_value.close,
                     bastions[0].close, proxies[0].close]
    def backoff(delay):
        assert targets[0].close.call_count == 1
        assert all(resource.call_count == 1 for resource in resources)
        retry_sleeps.append(delay)
    monkeypatch.setattr('orbitflow.transport.policy.sleep', backoff)
    def client():
        item = next(clients)
        if item is targets[1]:
            assert targets[0].close.call_count == 1
            assert all(resource.call_count == 1 for resource in resources)
        attempts.append(item)
        return item
    monkeypatch.setattr(module.paramiko, 'SSHClient', client)
    config = replace(TRANSPORT, tsh_path='tsh', teleport_key_path=Path('key'),
                     teleport_cert_path=Path('cert'))
    if cleanup_fails:
        with pytest.raises(ConnectionCleanupError):
            connect_device('192.0.2.1', DeviceCredentials('user'), config, system=system)
        assert all(item is not targets[1] for item in attempts)
    else:
        with connect_device('192.0.2.1', DeviceCredentials('user'), config, system=system):
            pass
        assert targets[1].close.call_count == 1
    assert all(resource.call_count == 1 for resource in resources)
    assert retry_sleeps == ([] if cleanup_fails else [5.0])


@pytest.mark.parametrize('cause', [AuthenticationException(SECRET), ValueError(SECRET),
    InteractiveCLITimeout(SECRET), EnableAuthenticationFailed(SECRET)])
def test_wrapped_nonretryable_failure(monkeypatch, cause, retry_sleeps):
    calls = []
    def connect(*args):
        calls.append(1)
        raise DeviceConnectionError('safe wrapper') from cause
    monkeypatch.setattr('orbitflow.transport.windows.connect_windows', connect)
    monkeypatch.setattr('orbitflow.transport.wait_for_connection_start', lambda: None)
    with pytest.raises(DeviceConnectionError):
        connect_device('192.0.2.1', DeviceCredentials('user'), TRANSPORT, system='windows')
    assert calls == [1]
    assert retry_sleeps == []


def test_direct_inventory_excludes_enable_secret(tmp_path):
    from orbitflow.inventory import DeviceInventoryResolver, JsonInventoryStore
    session, client, channel = device_session(ACCESS_DEVICES[0])
    session.secret = SECRET
    channel.outputs['show version'] = channel.outputs['show version'].replace('uptime is', 'uptime is ' + SECRET)
    path = tmp_path/'inventory.json'
    with session:
        DeviceInventoryResolver(JsonInventoryStore(path)).resolve(session, management_ip='192.0.2.1')
    assert SECRET not in path.read_text()
    assert '[REDACTED]' in path.read_text()


def test_inventory_refresh_secret_not_in_retained_spool(tmp_path, monkeypatch):
    from orbitflow import inventory_refresh as refresh
    from test_inventory_refresh import setup_network
    setup_network(monkeypatch, {'192.0.2.1': ('SERIAL1', SECRET)})
    original = refresh.connect_device
    def connect(ip, credentials, config):
        assert credentials.secret == SECRET
        return original(ip, credentials, config)
    monkeypatch.setattr(refresh, 'connect_device', connect)
    monkeypatch.setattr(refresh, 'export_inventory_spool', lambda *args: None)
    book = Workbook()
    book.active.append(['management_ip', 'username', 'password', 'Secret'])
    book.active.append(['192.0.2.1', 'runtime-login', 'runtime-password', SECRET])
    source = tmp_path/'input.xlsx'
    book.save(source)
    book.close()
    from test_inventory_refresh import CONFIG as config
    refresh.refresh_inventory_from_excel(source, config, inventory_path=tmp_path/'inventory.json',
        export_path=tmp_path/'export.xlsx', log_root=tmp_path/'logs',
        execution_config=ExecutionConfig(1))
    assert list(tmp_path.rglob('results.jsonl'))
    persisted = ''.join(p.read_text() for p in tmp_path.rglob('*')
                       if p.suffix in {'.json', '.jsonl', '.log'})
    assert SECRET not in persisted
    assert 'runtime-password' not in persisted


def test_already_privileged_annotated_edge_does_not_enable():
    edge = next(d for d in ACCESS_DEVICES if d[1] == 'ubiquiti_edgeswitch')
    session, client, channel = device_session(edge)
    channel.prompt = '(sw (anything (nested))) #'
    with session, DeviceCLI(session) as cli:
        assert extract_edgeswitch_hostname(cli.prompt) == 'sw'
        assert 'enable' not in channel.sent
    assert client.closes == channel.close_calls == 1


def test_partial_annotation_at_prompt_character_receive_boundary():
    edge = next(d for d in ACCESS_DEVICES if d[1] == 'ubiquiti_edgeswitch')
    session, client, channel = device_session(edge)
    channel.prompt = '(sw (arbitrary # > annotation)) #'
    send = channel.sendall
    def split(data):
        send(data)
        if data == b'\n':
            channel.pending = [b'(sw (arbitrary #', b' >', b' annotation)) #']
    channel.sendall = split
    with session, DeviceCLI(session) as cli:
        assert cli.prompt == channel.prompt
        assert extract_edgeswitch_hostname(cli.prompt) == 'sw'
    assert client.closes == channel.close_calls == 1


@pytest.mark.parametrize('delay', [0.0, 5.0, 8.0])
@pytest.mark.parametrize('limit', [1, 5])
def test_retry_backoff_then_pacing_and_run_config(monkeypatch, delay, limit):
    from orbitflow.transport.policy import connection_retry_delay
    now, events = [0.0], []

    def pacing_sleep(seconds):
        events.append(('pacing', seconds))
        now[0] += seconds

    pacer = ConnectionStartPacer(1.0, clock=lambda: now[0], sleep=pacing_sleep)
    monkeypatch.setattr('orbitflow.execution.ConnectionStartPacer', lambda interval: pacer)

    def backoff(seconds):
        assert events[-1] == ('cleanup', 0.0)
        events.append(('backoff', seconds))
        now[0] += seconds
        # A peer takes a start slot just as backoff ends. The retry must wait
        # for the next slot, even though it has already waited for backoff.
        pacer.wait()

    monkeypatch.setattr('orbitflow.transport.policy.sleep', backoff)
    attempts = []

    def connect(*args):
        attempts.append(now[0])
        if len(attempts) == 1:
            events.append(('cleanup', now[0]))
            raise TimeoutError()
        return DeviceSession(object(), lambda: None)

    monkeypatch.setattr('orbitflow.transport.windows.connect_windows', connect)

    def worker(target):
        with connect_device(str(target), DeviceCredentials('user'), TRANSPORT, system='windows'):
            return target

    config = ExecutionConfig(limit, connection_retry_delay=delay)
    outcomes = execute_devices([1], worker, config=config)
    assert not outcomes[0].error_category
    assert attempts == [0.0, delay + 1.0]
    assert events == ([('cleanup', 0.0), ('backoff', delay), ('pacing', 1.0)]
                      if delay else [('cleanup', 0.0), ('pacing', 1.0)])
    assert connection_retry_delay.get() == 5.0
    # A successful first attempt must not back off.
    events.clear()
    outcomes = execute_devices([2], worker, config=config)
    assert not outcomes[0].error_category
    assert not any(event[0] == 'backoff' for event in events)
