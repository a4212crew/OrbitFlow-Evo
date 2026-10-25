"""Issue149 synthetic relationship authority and compatibility regressions."""
from dataclasses import asdict, replace
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from orbitflow.compliance import evaluate_vlan_compliance
from orbitflow.compliance.resolvers import resolver_for
from orbitflow.compliance.methodology import resolve_methodologies
from orbitflow.compliance_report import export_compliance_spool, INTERFACE_COLUMNS
from orbitflow.models import InterfaceRecord, VlanState
from orbitflow.vendors.aggregation import parse_aggregation, identity
from orbitflow.vendors.configuration_facts import observe_configuration
from orbitflow.capabilities.vlans import VlanService
from orbitflow.result_spool import ResultSpool
from orbitflow.execution import DeviceOutcome
from test_aggregation import FAMILIES, fixture
from test_vlan_compliance import POLICY, NOW


def saved(spec, text):
    ctx, _ = fixture(*spec)
    snap = parse_aggregation(text, ctx)
    state = VlanState('synthetic', ctx.management_ip, ctx.platform, (), (), NOW,
                      observe_configuration(text, ctx.platform, source_filename='synthetic.cfg'), snap)
    return ctx, state


@pytest.mark.parametrize('spec', FAMILIES)
def test_all_families_share_owner_down_children_zero_parent_and_subinterface(spec):
    _, text = fixture(*spec)
    term = spec[-1]
    text = text.replace('end\n', f'interface {spec[2]}.445\n description logical child\n{term}\nend\n')
    ctx, state = saved(spec, text)
    resolver = resolver_for(ctx)(ctx, state)
    facts = resolver.result()
    member = facts['interfaces'][identity(ctx.platform, spec[3])]
    assert resolver.aggregation is state.aggregation
    assert member['configuration_owner'] == spec[2]
    assert member['shutdown']
    findings = evaluate_vlan_compliance(ctx, [], state, POLICY)
    child = next(f for f in findings if f['interface'] == spec[3])
    assert child['status'] == 'not_applicable'
    assert child['observed']['interface_type'] == 'Child Port-Channel'
    assert child['observed']['configuration_owner'] == spec[2]
    assert child['missing_vlans'] == [] and not child['observed']['trigger_applicable']
    assert 'owned by ' + spec[2] in child['explanation']
    assert not state.aggregation.owner(spec[2] + '.445')
    empty = text.replace(' ' + spec[4], ' description no binding')
    _, empty_state = saved(spec, empty)
    assert empty_state.aggregation.fields(spec[2]) == ['Yes', 'Parent', '']


@pytest.mark.parametrize('spec', FAMILIES)
@pytest.mark.parametrize('mode', ['missing_parent', 'contradiction', 'partial'])
def test_uncertain_relationships_cannot_pass_or_create_parent(spec, mode):
    _, text = fixture(*spec)
    if mode == 'missing_parent':
        text = text[text.index('interface ' + spec[3]):]
    elif mode == 'contradiction':
        other = spec[4].replace('1', '2')
        text = text.replace(' ' + spec[4], ' ' + spec[4] + '\n ' + other)
    else:
        text = text.removesuffix('end\n')
    ctx, state = saved(spec, text)
    findings = evaluate_vlan_compliance(ctx, [], state, POLICY)
    child = next(f for f in findings if f['interface'] == spec[3])
    assert child['status'] == ('not_applicable' if mode == 'partial' else 'unable_to_assess')
    if mode == 'missing_parent':
        assert not state.aggregation.parents
        assert not any(f['interface'] == spec[2] for f in findings)
    if mode == 'partial':
        parent = next(f for f in findings if f['interface'] == spec[2])
        assert parent['status'] == 'unable_to_assess'


def test_me_evc_parent_child_conflict_and_pseudowire_compliance_only():
    spec = FAMILIES[2]
    text = '''vlan 445,545,2449
!
interface Port-channel1
 service instance 7 ethernet
  encapsulation dot1q 100
  bridge-domain 445
 service instance 8 ethernet
  encapsulation dot1q 200
  bridge-domain 545
 service instance 9 ethernet
  encapsulation dot1q 300
  bridge-domain 2449
!
interface Gi0/1
 channel-group 1 mode active
 switchport mode access
 switchport access vlan 445
!
interface pseudowire7
 encapsulation mpls
 neighbor 192.0.2.1 77
!
end
'''
    ctx, state = saved(spec, text)
    result = evaluate_vlan_compliance(ctx, [], state, POLICY)
    assert not any(f['interface'] == 'pseudowire7' for f in result)
    parent = next(f for f in result if f['interface'] == 'Port-channel1')
    assert parent['status'] == 'non_compliant'
    assert parent['observed']['valid_interface_vlans'] == [445, 545, 2449]
    assert len(parent['observed']['numeric_mappings']) == 3
    child = next(f for f in result if f['interface'] == 'Gi0/1')
    assert child['status'] == 'not_applicable'
    assert child['configuration_health'] == 'wrong_configuration'
    assert 'PORT_CHANNEL_CONFIG_CONFLICT' in {p['code'] for p in child['observed']['configuration_findings']}
    assert any(r['config_interface_name'] == 'pseudowire7' for r in resolve_methodologies(ctx, [], state))
    assert 'pseudowire7' in dict(state.aggregation.evidence)


def test_edgeswitch_alias_join_and_offline_compliance_roundtrip(tmp_path):
    spec = FAMILIES[-1]
    _, text = fixture(*spec)
    text = text.replace(' addport 3/1', ' vlan participation include 445\n vlan tagging 445\n addport 3/1')
    text = text.replace('end\n', 'interface 3/1\n description repeated alias\nexit\nend\n')
    ctx, state = saved(spec, text)
    interfaces = [InterfaceRecord('synthetic', ctx.management_ip, ctx.platform,
                                 name, 'observed', 'up', 'down', NOW)
                  for name in ('3/1', '0/12')]
    findings = evaluate_vlan_compliance(ctx, interfaces, state, POLICY)
    parents = [f for f in findings if f.get('observed', {}).get('config_interface_name') in {'lag 1', '3/1'}]
    assert len(parents) == 1
    parent = parents[0]
    assert parent['interface'] == 'lag 1'
    assert parent['observed']['config_interface_name'] == 'lag 1'
    assert parent['observed']['configuration_owner'] == 'lag 1'
    assert parent['observed']['interface_name'] == '3/1'
    assert parent['observed']['interface_match_status'] == 'matched'
    assert parent['observed']['admin_status'] == 'up'
    assert parent['observed']['oper_status'] == 'down'
    child = next(f for f in findings if f['interface'] == '0/12')
    # Shared contract deterministically selects an actual heading for aliases.
    assert child['observed']['configuration_owner'] == 'lag 1'
    assert child['configuration_health'] == 'wrong_configuration'
    assert 'AGGREGATE_CONFIG_CONFLICT' in {p['code'] for p in child['observed']['configuration_findings']}
    # The independent interface/VLAN report keeps its operational display.
    from orbitflow.reporting import build_rows, INTERFACE_COLUMNS as REPORT_COLUMNS
    report_rows, _ = build_rows(ctx, interfaces, state, state.aggregation)
    assert [dict(zip(REPORT_COLUMNS, r))['Interface'] for r in report_rows] == ['0/12', '3/1']
    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload=dict(findings=findings, errors=[]))
    output = tmp_path / 'compliance.xlsx'
    export_compliance_spool(spool.path, output, cleanup=False)
    wb = load_workbook(output)
    sheet = wb['Interface Results']
    assert tuple(c.value for c in sheet[1]) == INTERFACE_COLUMNS
    rows = [dict(zip(INTERFACE_COLUMNS, r)) for r in sheet.iter_rows(min_row=2, values_only=True)]
    logical = [r for r in rows if r['Interface'] in {'lag 1', '3/1'}]
    assert len(logical) == 1
    assert logical[0]['Interface'] == logical[0]['Config Interface'] == logical[0]['Configuration Owner'] == 'lag 1'
    assert (logical[0]['Admin Status'], logical[0]['Oper Status']) == ('up', 'down')
    member = next(r for r in rows if r['Config Interface'] == '0/12')
    assert member['Type'] == 'Child Port-Channel' and member['Status'] == 'not_applicable'
    assert member['Configuration Owner'] == child['observed']['configuration_owner']
    assert 'owned by lag 1' in member['Explanation']
    assert member['Configuration Evidence'] and member['Evidence Lines']
    assert member['Configuration Health'] == 'wrong_configuration'
    assert 'AGGREGATE_CONFIG_CONFLICT' in member['Configuration Findings']
    wb.close()


@pytest.mark.parametrize('heading', ['3/1', None])
def test_edgeswitch_missing_lag_heading_never_fabricates_display(heading):
    spec = FAMILIES[-1]
    _, text = fixture(*spec)
    text = text.replace('interface lag 1\n description parent\nexit\n',
                        f'interface {heading}\nexit\n' if heading else '')
    ctx, state = saved(spec, text)
    interfaces = [InterfaceRecord('synthetic', ctx.management_ip, ctx.platform,
                                 '3/1', '', 'up', 'down', NOW)]
    findings = evaluate_vlan_compliance(ctx, interfaces, state, POLICY)
    assert not any(f['interface'] == 'lag 1' for f in findings)
    if heading:
        parent = next(f for f in findings if f['interface'] == '3/1')
        assert parent['observed']['config_interface_name'] == '3/1'
        assert parent['observed']['oper_status'] == 'down'
    else:
        child = next(f for f in findings if f['interface'] == '0/12')
        assert child['status'] == 'unable_to_assess'


@pytest.mark.parametrize('spec', FAMILIES[:-1])
def test_other_families_keep_observed_parent_display(spec):
    _, text = fixture(*spec)
    ctx, state = saved(spec, text)
    name = identity(ctx.platform, spec[2])
    interfaces = [InterfaceRecord('synthetic', ctx.management_ip, ctx.platform,
                                 name, '', 'up', 'down', NOW)]
    findings = evaluate_vlan_compliance(ctx, interfaces, state, POLICY)
    parent = next(f for f in findings if f['interface'] == name)
    assert parent['observed']['config_interface_name'] == spec[2]
    assert parent['observed']['configuration_owner'] == spec[2]
    assert parent['observed']['admin_status'] == 'up'
    assert parent['observed']['oper_status'] == 'down'


@pytest.mark.parametrize('complete', [True, False])
def test_report_reuses_only_complete_existing_capture(tmp_path, monkeypatch, complete):
    from io import StringIO
    from orbitflow import reporting
    from test_device_reporting import install_fakes, CONFIG, state as vlan_state
    install_fakes(monkeypatch)
    ctx, text = fixture(*FAMILIES[1])
    aggregation = replace(parse_aggregation(text, ctx), complete=complete)
    vlans = replace(vlan_state(), aggregation=aggregation)
    monkeypatch.setattr(reporting, 'VlanService', Mock(return_value=Mock(collect=Mock(return_value=vlans))))
    reporting.run_report([dict(management_ip='192.0.2.1', username='synthetic-user', password='synthetic-password')],
        CONFIG, reports_dir=tmp_path, spool_root=tmp_path / 'runs', log_root=tmp_path / 'logs',
        inventory_path=tmp_path / 'inventory.json', output=StringIO())
    assert reporting.AggregationService.call_count == (0 if complete else 1)


def test_reused_capture_keeps_independently_observed_empty_parent():
    ctx, _ = fixture(*FAMILIES[-1])
    snap = parse_aggregation('hostname synthetic\nend\n', ctx)
    observed = snap.with_observed(['3/1', '0/1', '3/1.445'])
    assert observed.parents == ('3/1',)
    assert observed.bindings == ()
    assert observed.fields('3/1') == ['Yes', 'Parent', '']
    assert not observed.owner('0/1')


def test_unresolved_membership_does_not_redirect_exact_target_description():
    from test_approved_catalyst_templates import result, NAME
    row, _, _ = result('C3750X', 'Interface Description',
        f'interface {NAME}\n channel-group 1 mode active\n!\n',
        dict(interface=NAME, description='new'))
    assert not row['blockers']
    assert row['resolved_target'] == NAME
    assert row['commands'] == ['configure terminal', f'interface {NAME}', 'description new', 'end']


def test_partial_capture_cannot_authorize_forwarding_via_confirmed_parent():
    from test_approved_catalyst_templates import result, NAME, NOW as PLAN_NOW, IP
    from orbitflow.configuration.v2.planning import plan
    config = (f'interface {NAME}\n channel-group 1 mode active\n!\n'
              'interface Port-channel1\n switchport mode trunk\n'
              ' switchport trunk encapsulation dot1q\n switchport trunk allowed vlan 445\n!\n')
    _, request, observed = result('C3750X', 'Trunk Port', config,
        dict(interface=NAME, action='Replace', vlans=[445, 545]))
    ctx, _ = fixture(*FAMILIES[0])
    aggregation = parse_aggregation(config, ctx)
    assert aggregation.owner(NAME) == 'port-channel1' and not aggregation.complete
    from orbitflow.compliance.vlan import safe_data
    observed['aggregation'] = safe_data(asdict(aggregation))
    row = plan([request], {IP: observed}, now=PLAN_NOW).to_dict()['results'][0]
    assert 'PORT_CHANNEL_OWNER_UNRESOLVED' in row['blockers']
    assert not row['deployment_ready']
