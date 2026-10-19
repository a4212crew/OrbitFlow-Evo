"""Offline owner-supplied MAC shapes and independent capability contracts."""
from types import SimpleNamespace as NS
from contextlib import contextmanager
import pytest
from openpyxl import load_workbook
from orbitflow.capabilities.mac import normalize_mac, parse_mac_table, resolve_attachment, MacService
from orbitflow.capabilities.bridge_domains import parse_bridge_domains
from orbitflow.capabilities.mac_forwarding import enrich
from orbitflow.mac_lookup import resolve_target, validate_vlan, run_lookup, export_mac_spool, MAIN_COLUMNS


@pytest.mark.parametrize('value', ['AABB.CCDD.EEFF', 'AA:BB:CC:DD:EE:FF', 'aabb-ccdd-eeff'])
def test_normalize(value):
    assert normalize_mac(value) == 'aa:bb:cc:dd:ee:ff'


@pytest.mark.parametrize('value', ['', 'aabbccddeeff', 'aa-bb-cc-dd-ee-ff', 'aabb.ccdd.ee', 'aa:bb:cc:dd:ee:gg', 'aa:bb:cc-dd:ee:ff'])
def test_invalid(value):
    with pytest.raises(ValueError):
        normalize_mac(value)


IOS = '''Vlan Mac Address Type Ports
All 0100.0ccc.cccc STATIC CPU
1100 d401.c312.96f0 DYNAMIC Gi1/0/46
545 00eb.d543.8472 STATIC Vl545
2437 0ea2.03b1.2437 DYNAMIC Te1/0/12
4001 9a49.254f.1f0f DYNAMIC Gi0/24+Efp3999
1201 e8a5.5acc.fed0 DYNAMIC Po2
Total Mac Addresses for this criterion: 30'''
XR = '''Mac Address Type Learned from/Filtered on LC learned Resync Age/Last Change Mapped to
789a.18d9.0fbc dynamic Te0/0/0/19.2440 N/A 17 Oct 10:36:24 N/A
ae78.d132.d87b EVPN (192.0.2.20, 2445) N/A 17 Oct 10:36:24 N/A
9a49.254f.1f0f dynamic (192.0.2.20, 4001) N/A 17 Oct 10:36:24 N/A
4464.3cd7.5481 routed BD id: 0 N/A N/A N/A'''
BRIDGE = '''Bridge group: VLAN545, bridge-domain: BD_VLAN545, id: 0, state: up, ShgId: 0, MSTi: 0
 List of ACs:
 BV545, state: up, BVI MAC addresses: 1
 Te0/0/0/19.545, state: up, Static MAC addresses: 0
Bridge group: LBB-PPPOE, bridge-domain: LBB-PPPOE, id: 2, state: up, ShgId: 0, MSTi: 0
 List of ACs:
 Te0/0/0/19.445, state: up, Static MAC addresses: 0
 List of Access PWs:
 Neighbor 192.0.2.20 pw-id 2445, state: up, Static MAC addresses: 0
 Neighbor 192.0.2.21 pw-id 2445, state: down, Static MAC addresses: 0, backup'''


def test_ios_totals_and_local_selection():
    table = parse_mac_table(IOS, 'cisco_ios')
    assert table.total == 30
    assert len(table.entries) == 6
    assert table.lookup(domain=545).total == 1
    assert table.lookup(mac='9A49.254F.1F0F').entries[0].attachment == 'Gi0/24+Efp3999'
    assert not table.lookup(interface='Gi0/24', platform='cisco_ios').entries
    assert table.lookup(domain=22).total == 0


def test_asr_and_empty():
    output = 'Nile Mac Address Entries\nBD mac addr type ports\nALL 0100.0ccc.cccc STATIC CPU\n2200 e8a5.5acc.fed0 DYNAMIC Gi0/0/1.Efp2200'
    assert parse_mac_table(output, 'cisco_xe').total == 2
    assert parse_mac_table('Vlan Mac Address Type Ports\n----', 'cisco_ios').total == 0


def test_huawei_and_edge():
    output = '''MAC address table of slot 0:
MAC Address VLAN/ PEVLAN CEVLAN Port Type LSP/LSR-ID
            VSI/SI MAC-Tunnel
0ea2-03b1-2435 SUPERLOOP2-PPPOE-2435 - - GE0/2/5.303 dynamic 0/0
744d-289f-92db 745 - - GE0/2/6 dynamic 0/-
Total matching items on slot 0 displayed = 153'''
    table = parse_mac_table(output, 'huawei_vrp')
    assert table.total == 2 and table.scoped_totals == (('slot 0', 153),)
    assert table.entries[0].domain == 'SUPERLOOP2-PPPOE-2435'
    edge = parse_mac_table('VLAN ID MAC Address Interface IfIndex Status\n445 00:26:DE:13:35:FE 0/8 8 Learned\n545 B4:FB:E4:19:9A:6E 5/1 65 Management', 'ubiquiti_edgeswitch')
    assert edge.total == 2 and edge.entries[0].kind == 'dynamic'
    assert edge.entries[1].kind == 'management' and edge.entries[1].ifindex == '65'


@pytest.mark.parametrize('text', ['', 'garbage', 'Vlan Mac Address Type Ports\nbroken row', IOS + '\n--More--', IOS + '\n% Invalid input'])
def test_reject_incomplete(text):
    with pytest.raises(ValueError):
        parse_mac_table(text, 'cisco_ios')


def test_xr_exact_scope():
    table = parse_mac_table(XR, 'cisco_xr')
    assert table.total == 4 and table.entries[1].attachment == '(192.0.2.20, 2445)'
    bridge = parse_bridge_domains(BRIDGE)
    assert bridge.for_destination('BD id: 0')[0].name == 'BD_VLAN545'
    assert bridge.for_destination('(192.0.2.20, 2445)')[0].name == 'LBB-PPPOE'
    assert bridge.domains[1].attachments[-1].backup
    assert bridge.domains[0].attachments[0].bvi_count == 1
    assert bridge.domains[0].attachments[0].static_count is None
    assert bridge.lookup(interface='TenGigE0/0/0/19.545')[0].internal_id == '0'
    assert bridge.lookup(interface='BVI545')[0].internal_id == '0'
    ambiguous = parse_bridge_domains(BRIDGE + '\n' + BRIDGE)
    assert len(ambiguous.for_destination('(192.0.2.20, 2445)')) == 2


@pytest.mark.parametrize('platform,name,attachment,expected', [
    ('cisco_ios', 'Gi0/24', 'Gi0/24+Efp3999', 'GigabitEthernet0/24'),
    ('cisco_xe', 'Te1/0/12', 'Te1/0/12', 'TenGigabitEthernet1/0/12'),
    ('cisco_xr', 'Te0/0/0/19', 'Te0/0/0/19.2440', 'TenGigE0/0/0/19'),
    ('ubiquiti_edgeswitch', '0/8', '0/8', '0/8')])
def test_physical(platform, name, attachment, expected):
    result = resolve_attachment(attachment, platform, [NS(port_name=name)])
    assert result.physical == expected and result.port_type == 'Unknown'
    assert resolve_attachment(attachment, platform, []).physical == ''


def test_domain_not_vlan():
    table = parse_mac_table('BD mac addr type ports\n4001 9a49.254f.1f0f DYNAMIC Gi0/24+Efp3999', 'cisco_ios')
    context = NS(platform='cisco_ios', device_family='ME3600X')
    detail = NS(interface_name='GigabitEthernet0/24', service_details=(), service_mappings=('445 -> 4001',), port_type='evc')
    from orbitflow.vendors.configuration_facts import observe_configuration
    facts = observe_configuration('interface GigabitEthernet0/24\n service instance 3999 ethernet\n  encapsulation dot1q 445\n  bridge-domain 4001', 'cisco_ios')
    rows = enrich(table, context, [NS(port_name='Gi0/24')], NS(interfaces=[detail], objects=(), configuration=facts))
    assert rows[0].vlans == (445,) and rows[0].resolution.port_type == 'trunk'
    assert enrich(table, context)[0].vlans == ()


def test_target_and_validation():
    target = dict(management_ip='192.0.2.1', username='fixture', password='private-fixture')
    assert resolve_target([target], '192.0.2.1') is target
    for targets in ([], [target, target]):
        with pytest.raises(ValueError):
            resolve_target(targets, '192.0.2.1')
    for vlan in ('0', '4095', 'x', True):
        with pytest.raises(ValueError):
            validate_vlan(vlan)


def test_batch_spool_isolation_and_recovery(tmp_path, monkeypatch):
    import orbitflow.mac_lookup as app
    def collect(target, *args, **kwargs):
        if target['management_ip'] == '192.0.2.2':
            raise ValueError('private-fixture')
        return [[['192.0.2.1', '=formula', 'C3750X', 'GigabitEthernet1/1', 'trunk', 'aa:bb:cc:dd:ee:ff']] * 4, [], [['192.0.2.1', 'table', 30]], []]
    monkeypatch.setattr(app, 'collect_device', collect)
    targets = [dict(management_ip=f'192.0.2.{i}', username='fixture', password='private-fixture') for i in (1, 2)]
    path = run_lookup(targets, None, reports_dir=tmp_path/'reports', spool_root=tmp_path/'runs', inventory_path=tmp_path/'inventory.json', log_root=tmp_path/'logs')
    book = load_workbook(path)
    assert tuple(c.value for c in book['MAC Addresses'][1]) == MAIN_COLUMNS
    assert book['MAC Addresses'].max_row == 5
    assert book['MAC Addresses']['B2'].data_type == 's'
    assert book['MAC Addresses']['G2'].value is None
    assert book['Run Errors'].max_row == 2
    book.close()
    spool = next((tmp_path/'runs').iterdir())
    export_mac_spool(spool, tmp_path/'recovered.xlsx')
    assert (tmp_path/'recovered.xlsx').exists()
    assert all('private-fixture' not in p.read_text(errors='ignore') for p in spool.rglob('*') if p.is_file())


@pytest.mark.parametrize('platform,command,output', [
    ('cisco_ios', 'show mac address-table', IOS),
    ('cisco_xe', 'show mac address-table', 'BD mac addr type ports'),
    ('cisco_xr', 'show l2vpn forwarding bridge-domain mac-address location 0/RP0/CPU0', XR),
    ('huawei_vrp', 'display mac-address', 'MAC Address VLAN/ PEVLAN CEVLAN Port Type\nVSI/SI MAC-Tunnel'),
    ('ubiquiti_edgeswitch', 'show mac-addr-table all', 'VLAN ID MAC Address Interface IfIndex Status')])
def test_independent_borrowed_collection(monkeypatch, platform, command, output):
    import orbitflow.vendors.mac_tables as vendor
    calls = []
    shell = NS(run_command=lambda cmd, **kw: calls.append(cmd) or output)
    @contextmanager
    def opened(session, observed_platform, **kwargs):
        assert kwargs['cli'] is shell and session == 'session' and platform == observed_platform
        yield shell
    monkeypatch.setattr(vendor, 'operational_cli', opened)
    MacService().collect('session', NS(platform=platform), cli=shell)
    assert calls == [command]


def test_exact_service_conflicts_and_wrong_efp():
    from orbitflow.vendors.configuration_facts import observe_configuration
    from orbitflow.capabilities.mac_forwarding import exact_efp_vlans
    facts = observe_configuration('interface GigabitEthernet0/24\n service instance 3999 ethernet\n  encapsulation dot1q 445\n  bridge-domain 4001', 'cisco_ios')
    assert exact_efp_vlans(facts, 'cisco_ios', 'Gi0/24', '3999', '4001') == (445,)
    assert exact_efp_vlans(facts, 'cisco_ios', 'Gi0/24', '4001', '4001') == ()
    assert exact_efp_vlans(facts, 'cisco_ios', 'Gi0/24', '3999', '445') == ()


def test_hybrid_aggregate_nonphysical_and_huawei():
    interfaces = [NS(port_name='Po2'), NS(port_name='GigabitEthernet0/2/6')]
    vlans = NS(interfaces=[NS(interface_name='Port-channel2', port_type='hybrid')])
    result = resolve_attachment('Po2', 'cisco_ios', interfaces, vlans)
    assert result.classification == 'aggregate' and result.physical == 'Port-channel2' and result.port_type == 'hybrid'
    assert resolve_attachment('GE0/2/6', 'huawei_vrp', interfaces).physical == 'GigabitEthernet0/2/6'
    for attachment in ('CPU', 'Vl545', '(192.0.2.20, 2445)', 'BD id: 0'):
        assert not resolve_attachment(attachment, 'cisco_xr', interfaces).physical
    assert not resolve_attachment('5/1', 'ubiquiti_edgeswitch', [NS(port_name='5/1')]).physical


def test_xr_vlan_requires_explicit_relationship_and_unique_join():
    from orbitflow.capabilities.mac_forwarding import lookup_vlan
    table = parse_mac_table(XR, 'cisco_xr')
    bridge = parse_bridge_domains(BRIDGE)
    profile = NS(interface_name='TenGigE0/0/0/19.445', service_details=(), service_mappings=('860 -> LBB-PPPOE',), port_type='service', service_binding_name='LBB-PPPOE/LBB-PPPOE', service_binding_type='bridge-domain')
    context = NS(platform='cisco_xr', device_family='NCS540')
    enriched = enrich(table, context, [], NS(interfaces=[profile], configuration=()), bridge)
    assert [r.entry.mac for r in lookup_vlan(enriched, 860)] == ['ae:78:d1:32:d8:7b']
    assert not lookup_vlan(enriched, 445) and not lookup_vlan(enriched, 2445)
    assert not lookup_vlan(enrich(table, context, [], NS(interfaces=[profile]), parse_bridge_domains(BRIDGE+'\n'+BRIDGE)), 860)


def test_export_failure_retains_spool(tmp_path, monkeypatch):
    import orbitflow.mac_lookup as app
    monkeypatch.setattr(app, 'collect_device', lambda *a, **k: [[], [], [], []])
    real = app.write_tables
    def fail(*args, **kwargs):
        raise OSError('disk unavailable')
    monkeypatch.setattr(app, 'write_tables', fail)
    with pytest.raises(OSError):
        run_lookup([dict(management_ip='192.0.2.1')], None, reports_dir=tmp_path/'reports', spool_root=tmp_path/'runs', inventory_path=tmp_path/'inventory.json', log_root=tmp_path/'logs')
    monkeypatch.setattr(app, 'write_tables', real)
    export_mac_spool(next((tmp_path/'runs').iterdir()), tmp_path/'recovered.xlsx')
    assert (tmp_path/'recovered.xlsx').exists()


def test_one_session_partial_enrichment_and_no_match(monkeypatch):
    import orbitflow.mac_lookup as app
    calls, session, shell = [], object(), object()
    @contextmanager
    def connect(*args, **kwargs):
        calls.append('connect')
        yield session
        calls.append('disconnect')
    @contextmanager
    def cli(observed_session):
        assert observed_session is session
        yield shell
    context = NS(platform='cisco_ios', device_family='C3750X', hostname='sample', hardware_model='WS-C3750X')
    def collect(self, observed_session, observed_context, **kwargs):
        assert observed_session is session and observed_context is context and kwargs['cli'] is shell
        return parse_mac_table(IOS, 'cisco_ios')
    def failed(*args, **kwargs):
        raise ValueError('secret-bearing-device-diagnostic')
    monkeypatch.setattr(app, 'connect_device', connect)
    monkeypatch.setattr(app, 'DeviceCLI', cli)
    monkeypatch.setattr(app.DeviceInventoryResolver, 'resolve', lambda *a, **kw: context)
    monkeypatch.setattr(app.MacService, 'collect', collect)
    monkeypatch.setattr(app.InterfaceService, 'collect', failed)
    monkeypatch.setattr(app.VlanService, 'collect', failed)
    target = dict(management_ip='192.0.2.1', username='fixture', password='private-fixture')
    result = app.collect_device(target, None, None, mac='d401.c312.96f0')
    assert len(result[0]) == 1 and result[0][0][2] == 'WS-C3750X'
    assert result[0][0][3:5] == ['', 'Unknown']
    assert result[0][0][6] == ''
    assert len(result[3]) == 2 and 'secret-bearing' not in repr(result)
    assert calls == ['connect', 'disconnect']
    empty = app.collect_device(target, None, None, mac='aaaa.bbbb.cccc')
    assert empty[0] == [] and empty[2][1][-1] == 0


def test_duplicate_domains_and_parent_lookup():
    table = parse_mac_table('Vlan Mac Address Type Ports\n100 aabb.ccdd.eeff DYNAMIC Gi0/1.100\n200 aabb.ccdd.eeff DYNAMIC Gi0/1.200', 'cisco_ios')
    assert table.lookup(mac='aabb.ccdd.eeff').total == 2
    assert table.lookup(interface='GigabitEthernet0/1', platform='cisco_ios', include_children=True).total == 2
    assert table.lookup(interface='GigabitEthernet0/1.100', platform='cisco_ios').total == 1


def test_cli_offline_export_does_not_load_profile(tmp_path, monkeypatch):
    import orbitflow.mac_cli as cli
    monkeypatch.setattr(cli, 'load_profile', lambda *a: pytest.fail('offline profile'))
    seen = []
    monkeypatch.setattr(cli, 'export_mac_spool', lambda *a, **k: seen.append((a, k)) or tmp_path/'out.xlsx')
    cli.main(['export', 'retained-run', str(tmp_path/'out.xlsx')])
    assert seen[0][1] == {'allow_partial': False}


def test_unresolved_local_attachment_cannot_borrow_sibling_vlan():
    table = parse_mac_table('Mac Address Type Learned from/Filtered on LC learned Resync Age/Last Change Mapped to\naabb.ccdd.eeff dynamic Te0/0/0/19.545 N/A N/A N/A', 'cisco_xr')
    profile = NS(interface_name='TenGigE0/0/0/20.860', service_details=(), service_mappings=('860 -> BD_VLAN545',), port_type='service')
    bridge = parse_bridge_domains(BRIDGE.replace('Te0/0/0/19.545, state:', 'Te0/0/0/20.860, state: up, Static MAC addresses: 0\n Te0/0/0/19.545, state:'))
    rows = enrich(table, NS(platform='cisco_xr', device_family='NCS540'), [], NS(interfaces=[profile]), bridge)
    assert rows[0].vlans == ()
    assert resolve_attachment('Tunnel1', 'cisco_ios', [NS(port_name='Tunnel1')]).physical == ''


def test_xr_preserves_reported_scope_and_rejects_truncated_row():
    table = parse_mac_table('Bridge-domain: supplied-scope, id: 0\n' + XR, 'cisco_xr')
    assert table.entries[0].scope == 'Bridge-domain: supplied-scope, id: 0'
    with pytest.raises(ValueError):
        parse_mac_table(XR.splitlines()[0] + '\naabb.ccdd.eeff dynamic Te0/0/0/19.545 N/A N/A', 'cisco_xr')


def test_actual_configured_parent_is_identity_but_service_reference_is_not():
    from orbitflow.vendors.configuration_facts import observe_configuration
    actual = observe_configuration('interface GigabitEthernet0/24\n service instance 3999 ethernet\n  encapsulation dot1q 445\n  bridge-domain 4001', 'cisco_ios')
    reference = observe_configuration('bridge-domain 4001\n member GigabitEthernet0/24 service-instance 3999', 'cisco_ios')
    assert resolve_attachment('Gi0/24+Efp3999', 'cisco_ios', [], NS(interfaces=[], configuration=actual)).physical == 'GigabitEthernet0/24'
    assert resolve_attachment('Gi0/24+Efp3999', 'cisco_ios', [], NS(interfaces=[], configuration=reference)).physical == ''

@pytest.mark.parametrize('group,expected', [('G1', (860,)), ('G2', ())])
@pytest.mark.parametrize('destination', ['Te0/0/0/19.445', '(192.0.2.20, 2445)'])
def test_xr_group_identity_real_vlan_service(monkeypatch, group, expected, destination):
    from orbitflow.capabilities.vlans import VlanService
    import orbitflow.vendors.cisco.vlans as vendor
    config = f'''interface TenGigE0/0/0/19.445 l2transport
 encapsulation dot1q 860
!
l2vpn
 bridge group {group}
  bridge-domain SAME
   interface TenGigE0/0/0/19.445
!
'''
    @contextmanager
    def opened(*args, **kwargs):
        yield NS(prompt='RP/0/RP0/CPU0:lab#', run_command=lambda *a, **k: config)
    monkeypatch.setattr(vendor, 'open_cisco_cli', opened)
    context = NS(platform='cisco_xr', device_family='NCS540', hostname='lab', management_ip='192.0.2.1')
    state = VlanService().collect(None, context)
    assert state.interfaces[0].service_binding_name == f'{group}/SAME'
    bridges = parse_bridge_domains('''Bridge group: G1, bridge-domain: SAME, id: 0, state: up
 List of ACs:
 Te0/0/0/19.445, state: up
 List of Access PWs:
 Neighbor 192.0.2.20 pw-id 2445, state: up, backup
Bridge group: G2, bridge-domain: SAME, id: 1, state: down
 List of ACs:
 Te0/0/0/20.445, state: down''')
    table = parse_mac_table(XR.splitlines()[0] + f'\naabb.ccdd.eeff dynamic {destination} N/A N/A N/A', 'cisco_xr')
    row = enrich(table, context, (), state, bridges)[0]
    assert row.domains == ('G1/SAME',)
    assert row.vlans == expected
    assert row.mapping == ('resolved' if expected else 'unknown VLAN mapping')
    duplicate = parse_bridge_domains('\n'.join(['Bridge group: G2, bridge-domain: SAME, id: 2, state: up', ' List of ACs:', ' Te0/0/0/19.445, state: up']))
    from orbitflow.capabilities.bridge_domains import BridgeSnapshot
    if not destination.startswith('('):
        assert not enrich(table, context, (), state, BridgeSnapshot(bridges.domains + duplicate.domains))[0].vlans


BRIDGE_SUMMARY = '''Bridge group: G, bridge-domain: D, id: 0, state: up
 ACs: 9 (8 up), VFIs: 0, PWs: 0 (0 up), PBBs: 0 (0 up), VNIs: 0 (0 up)
 List of ACs:
 BV545, state: up, BVI MAC addresses: 1
 Te0/0/0/19.445, state: down, Static MAC addresses: 2
 Te0/0/0/20.445, state: unresolved
 List of Access PWs:
 Neighbor 192.0.2.1 pw-id 445, state: down, backup
 List of VFIs:
 List of Access VFIs:
Bridge group: G, bridge-domain: EMPTY, id: 1, state: down
 ACs: 0 (0 up), VFIs: 0, PWs: 0 (0 up), PBBs: 0 (0 up), VNIs: 0 (0 up)
 List of ACs:
 List of Access PWs:
 List of VFIs:
 List of Access VFIs:
'''


def test_bridge_complete_summary_and_empty_sections():
    domains = parse_bridge_domains(BRIDGE_SUMMARY).domains
    assert dict(domains[0].counts) == {'ACs': 9, 'ACs up': 8, 'VFIs': 0, 'PWs': 0, 'PWs up': 0, 'PBBs': 0, 'PBBs up': 0, 'VNIs': 0, 'VNIs up': 0}
    assert domains[1].attachments == ()
    assert domains[0].attachments[-1].backup
    assert domains[0].attachments[2].state == 'unresolved'


@pytest.mark.parametrize('content', ['nonsense-truncated', 'Te0/0/0/19.445, state:', 'Neighbor 192.0.2.1 pw-id 1, state: up', 'List of VFIs:\n vfi, state: up', 'List of Access VFIs:\n Te0/0/0/19.445, state: up', 'List of VFIs:\n Neighbor 192.0.2.1 pw-id 1, state: up', 'List of Unknowns:'])
def test_bridge_rejects_malformed_or_wrong_section(content):
    with pytest.raises(ValueError):
        parse_bridge_domains('Bridge group: G, bridge-domain: D, id: 0, state: up\n List of ACs:\n ' + content)


def test_ncs_collection_vlan_match_survives_empty_vfi_sections(monkeypatch):
    import orbitflow.mac_lookup as app
    @contextmanager
    def connect(*args, **kwargs):
        yield object()
    @contextmanager
    def cli(*args):
        yield NS(run_command=lambda command, **kwargs: BRIDGE_SUMMARY)
    context = NS(platform='cisco_xr', device_family='NCS540', hostname='lab', hardware_model='NCS540')
    profile = NS(interface_name='TenGigE0/0/0/19.445', service_details=(),
                 service_mappings=('860 -> D',), service_binding_name='G/D', service_binding_type='bridge-domain', port_type='service')
    table = parse_mac_table(XR.splitlines()[0] + '\naabb.ccdd.eeff dynamic Te0/0/0/19.445 N/A N/A N/A', 'cisco_xr')
    monkeypatch.setattr(app, 'connect_device', connect)
    monkeypatch.setattr(app, 'DeviceCLI', cli)
    @contextmanager
    def operational_cli(*args, **kwargs):
        yield kwargs['cli']
    monkeypatch.setattr('orbitflow.vendors.mac_tables.operational_cli', operational_cli)
    monkeypatch.setattr(app.DeviceInventoryResolver, 'resolve', lambda *a, **k: context)
    monkeypatch.setattr(app.MacService, 'collect', lambda *a, **k: table)
    monkeypatch.setattr(app.InterfaceService, 'collect', lambda *a, **k: [NS(port_name='TenGigE0/0/0/19')])
    monkeypatch.setattr(app.VlanService, 'collect', lambda *a, **k: NS(interfaces=[profile]))
    result = app.collect_device(dict(management_ip='192.0.2.1', username='fixture', password='fixture'), None, None, vlan=860)
    assert result[0] == [['192.0.2.1', 'lab', 'NCS540', 'TenGigE0/0/0/19', 'Unknown', 'aa:bb:cc:dd:ee:ff', '']]
    assert result[1][0][4:6] == ['G/D', '860']
    assert result[3] == []
    assert result[4]


def configured_state(config, platform):
    from orbitflow.vendors.configuration_facts import observe_configuration
    return NS(interfaces=(), objects=(), configuration=observe_configuration(config, platform))


@pytest.mark.parametrize('family,platform', [('ME3600X', 'cisco_ios'), ('ASR920', 'cisco_xe')])
@pytest.mark.parametrize('encaps,extra,expected', [
    ('untagged', '', 'access'),
    ('dot1q 445', '', 'trunk'),
    ('untagged', ' service instance 2 ethernet\n  encapsulation dot1q 545\n  bridge-domain 700', 'hybrid'),
    ('untagged', ' service instance 2 ethernet\n  encapsulation untagged\n  bridge-domain 700', 'Unknown'),
    ('dot1q 445', 'bridge-domain 800\n member Gi0/1 service-instance 1', 'Unknown'),
])
def test_family_classification_shared_with_audit(family, platform, encaps, extra, expected):
    state = configured_state(f'interface Gi0/1\n service instance 1 ethernet\n  encapsulation {encaps}\n  bridge-domain 600\n{extra}', platform)
    context = NS(platform=platform, device_family=family)
    table = parse_mac_table('BD mac addr type ports\n600 aabb.ccdd.eeff dynamic Gi0/1.Efp1', platform)
    row = enrich(table, context, (), state)[0]
    assert row.resolution.port_type == expected
    assert 600 not in row.vlans and 1 not in row.vlans
    from orbitflow.compliance.resolvers import resolver_for
    kind = resolver_for(context)(context, state).result()['interfaces']['gigabitethernet0/1']['interface_type']
    assert row.resolution.semantics == kind


@pytest.mark.parametrize('global_binding', ['', 'bridge-domain 600\n member Gi0/1 service-instance 7',
    'bridge-domain 800\n member Gi0/1 service-instance 7'])
def test_exact_efp_local_global_and_duplicate_rows(global_binding):
    state = configured_state('interface Gi0/1\n service instance 7 ethernet\n  encapsulation dot1q 445\n  bridge-domain 600\n'+global_binding, 'cisco_ios')
    table = parse_mac_table('BD mac addr type ports\n'+ '600 aabb.ccdd.eeff dynamic Gi0/1+Efp7\n'*2, 'cisco_ios')
    rows = enrich(table, NS(platform='cisco_ios', device_family='ME3600X'), (), state)
    assert len(rows) == 2
    assert all(r.vlans == (() if '800' in global_binding else (445,)) for r in rows)


@pytest.mark.parametrize('number', [1, 2, 5])
def test_edgeswitch_exact_lag_destinations(number):
    from orbitflow.vendors.ubiquiti.vlans import parse_edgeswitch_config
    config = f'interface 0/8\n addport 3/{number}\nexit\ninterface lag {number}\n switchport mode trunk\nexit'
    state = configured_state(config, 'ubiquiti_edgeswitch')
    state.interfaces = parse_edgeswitch_config(config)[0]
    table = parse_mac_table(f'VLAN ID MAC Address Interface IfIndex Status\n445 aa:bb:cc:dd:ee:ff 3/{number} 100 Learned\n545 aa:bb:cc:dd:ee:ff 5/1 65 Management', 'ubiquiti_edgeswitch')
    rows = enrich(table, NS(platform='ubiquiti_edgeswitch', device_family='EdgeSwitch'), [NS(port_name='0/8')], state)
    assert rows[0].entry.attachment == f'3/{number}' and rows[0].entry.ifindex == '100'
    assert rows[0].resolution.physical == f'lag {number}'
    assert rows[0].resolution.classification == 'aggregate' and rows[0].resolution.port_type == 'trunk'
    assert rows[1].resolution.classification == 'nonphysical' and not rows[1].resolution.physical
    assert not resolve_attachment(f'3/{number}', 'ubiquiti_edgeswitch', [NS(port_name='0/8')]).physical
    assert state.interfaces[0].tagged_vlans == () and not state.interfaces[0].bridge_domains
    from orbitflow.compliance.resolvers import EdgeSwitchAuditResolver
    facts = EdgeSwitchAuditResolver(NS(platform='ubiquiti_edgeswitch', device_family='EdgeSwitch'), state).result()
    assert facts['interfaces'][f'lag {number}']['valid_interface_vlans'] == []
    assert facts['interfaces']['0/8']['configuration_owner'] == f'lag {number}'


@pytest.mark.parametrize('number', [1, 2, 5])
@pytest.mark.parametrize('observed,config,resolved', [
    (['lag {n}'], '', True),
    ([], 'interface lag {n}\nexit', True),
    (['lag {n}'], 'interface lag {n}\nexit', True),
    (['0/8'], '', False),
    ([], 'interface 0/8\n addport 3/{n}\nexit', False),
    (['lag 9'], 'interface lag 8\nexit', False),
    (['lag {n}', 'LAG {n}'], '', False),
    ([], 'interface lag {n}\nexit\ninterface lag {n}\nexit', False),
])
def test_edgeswitch_lag_requires_unique_independent_identity(number, observed, config, resolved):
    state = configured_state(config.format(n=number), 'ubiquiti_edgeswitch')
    interfaces = [NS(port_name=name.format(n=number)) for name in observed]
    table = parse_mac_table(
        f'VLAN ID MAC Address Interface IfIndex Status\n'
        f'445 aa:bb:cc:dd:ee:ff 3/{number} 100 Learned\n'
        '545 aa:bb:cc:dd:ee:ff 5/1 65 Management', 'ubiquiti_edgeswitch')
    rows = enrich(table, NS(platform='ubiquiti_edgeswitch', device_family='EdgeSwitch'), interfaces, state)
    assert len(rows) == 2
    lag, management = rows
    assert lag.entry.attachment == f'3/{number}' and lag.entry.ifindex == '100'
    assert lag.vlans == (445,)
    assert lag.resolution.physical == (f'lag {number}' if resolved else '')
    assert lag.resolution.classification == ('aggregate' if resolved else 'unresolved')
    assert lag.resolution.port_type == 'Unknown'
    assert management.entry.attachment == '5/1' and management.entry.ifindex == '65'
    assert management.resolution.classification == 'nonphysical'
    assert management.resolution.physical == ''


@pytest.mark.parametrize('present', [True, False])
def test_edgeswitch_logical_aggregate_workbook_contract(tmp_path, monkeypatch, present):
    import orbitflow.mac_lookup as app
    resolution = resolve_attachment('3/2', 'ubiquiti_edgeswitch',
                                    [NS(port_name='lag 2')] if present else [])
    def collect(target, *args, **kwargs):
        return [[['192.0.2.1', 'lab', 'EdgeSwitch', resolution.physical,
                  resolution.port_type, 'aa:bb:cc:dd:ee:ff']], [], [], [], []]
    monkeypatch.setattr(app, 'collect_device', collect)
    path = run_lookup([dict(management_ip='192.0.2.1')], None,
                      reports_dir=tmp_path/'reports', spool_root=tmp_path/'runs',
                      inventory_path=tmp_path/'inventory.json', log_root=tmp_path/'logs')
    book = load_workbook(path)
    assert tuple(c.value for c in book['MAC Addresses'][1]) == MAIN_COLUMNS
    assert book['MAC Addresses'].max_column == 7
    assert book['MAC Addresses'].max_row == 2
    assert book['MAC Addresses']['D2'].value == ('lag 2' if present else None)
    assert book['MAC Addresses']['E2'].value == 'Unknown'
    assert set(book.sheetnames) == {'MAC Addresses', 'Details', 'Totals', 'Run Errors', 'Bridge Domains'}
    book.close()


@pytest.mark.parametrize('tags,group,ambiguous,expected', [
    ([], 'G', False, ()), ([445], 'G', False, (445,)),
    ([445, 445], 'G', False, (445,)), ([445, 545], 'G', False, ()),
    ([445], 'OTHER', False, ()), ([445], 'G', True, ())])
def test_xr_untagged_mac_only_peer_association(tags, group, ambiguous, expected):
    config = 'interface TenGigE0/0/0/1.99 l2transport\n encapsulation untagged\n!\n'
    for index, tag in enumerate(tags):
        config += f'interface TenGigE0/0/0/2.{index} l2transport\n encapsulation dot1q {tag}\n!\n'
    config += f'l2vpn\n bridge group {group}\n  bridge-domain D\n   interface TenGigE0/0/0/1.99\n'
    for index in range(len(tags)):
        config += f'   interface TenGigE0/0/0/2.{index}\n'
    if ambiguous:
        config += ' bridge group X\n  bridge-domain D\n   interface TenGigE0/0/0/1.99\n'
    state = configured_state(config, 'cisco_xr')
    bridge = parse_bridge_domains('Bridge group: G, bridge-domain: D, id: 99, state: up\n List of ACs:\n Te0/0/0/1.99, state: up')
    table = parse_mac_table('Mac Address Type Learned from/Filtered on\naabb.ccdd.eeff dynamic Te0/0/0/1.99 N/A N/A N/A', 'cisco_xr')
    row = enrich(table, NS(platform='cisco_xr', device_family='NCS540'), (), state, bridge)[0]
    assert row.vlans == expected and row.entry.attachment == 'Te0/0/0/1.99'
    assert row.resolution.port_type == 'Unknown'


@pytest.mark.parametrize('family,platform', [('C3750X', 'cisco_ios'), ('C3850', 'cisco_xe'), ('ME3600X', 'cisco_ios'), ('ASR920', 'cisco_xe')])
def test_numeric_vlan_source_retained_without_configuration(family, platform):
    table = parse_mac_table('Vlan Mac Address Type Ports\n445 aabb.ccdd.eeff dynamic Po2\n445 aabb.ccdd.eeff static Vl445\nAll aabb.ccdd.eeff static CPU', platform)
    rows = enrich(table, NS(platform=platform, device_family=family))
    assert [r.vlans for r in rows] == [(445,), (445,), ()]
    assert len(rows) == 3 and rows[-1].resolution.classification == 'cpu'


@pytest.mark.parametrize('family,platform', [('ME3600X', 'cisco_ios'), ('ASR920', 'cisco_xe')])
def test_cpu_pw_and_unresolved_efp_remain_distinct(family, platform):
    table = parse_mac_table('BD mac addr type ports\nALL aabb.ccdd.eeff static CPU\n600 aabb.ccdd.eeff dynamic PW700\n600 aabb.ccdd.eeff dynamic Gi0/1.Efp600', platform)
    rows = enrich(table, NS(platform=platform, device_family=family))
    assert len(rows) == 3
    assert [r.resolution.classification for r in rows] == ['cpu', 'remote', 'unresolved']
    assert all(not r.vlans and not r.resolution.physical for r in rows)


@pytest.mark.parametrize('defect', ['missing', 'conflicting_binding', 'conflicting_encapsulation'])
def test_xr_peer_fallback_rejects_uncertain_peer_facts(defect):
    from orbitflow.capabilities.mac_forwarding import xr_untagged_peer_vlan
    config = '''interface TenGigE0/0/0/1.99 l2transport
 encapsulation untagged
!
interface TenGigE0/0/0/2.1 l2transport
 encapsulation dot1q 445
!
'''
    if defect != 'missing':
        config += 'interface TenGigE0/0/0/3.1 l2transport\n encapsulation dot1q 445\n'
        if defect == 'conflicting_encapsulation':
            config += ' encapsulation dot1q 545\n'
        config += '!\n'
    config += '''l2vpn
 bridge group G
  bridge-domain D
   interface TenGigE0/0/0/1.99
   interface TenGigE0/0/0/2.1
   interface TenGigE0/0/0/3.1
'''
    if defect == 'conflicting_binding':
        config += ' bridge group OTHER\n  bridge-domain D\n   interface TenGigE0/0/0/3.1\n'
    state = configured_state(config, 'cisco_xr')
    assert not xr_untagged_peer_vlan(state.configuration, 'cisco_xr', 'Te0/0/0/1.99', 'G/D')


@pytest.mark.parametrize('platform,family,attachment,owner', [
    ('cisco_ios', 'C3750X', 'Gi1/0/46', 'GigabitEthernet1/0/46'),
    ('cisco_xe', 'C3850', 'Po2', 'Port-channel2'),
    ('cisco_ios', 'ME3600X', 'Gi0/24+Efp3999', 'GigabitEthernet0/24'),
    ('cisco_xe', 'ASR920', 'Gi0/0/1.Efp2200', 'GigabitEthernet0/0/1'),
    ('cisco_xr', 'NCS540', 'Te0/0/0/19.2440', 'TenGigE0/0/0/19'),
    ('huawei_vrp', 'NE05', 'GE0/2/5.303', 'GigabitEthernet0/2/5'),
    ('huawei_vrp', 'NE05E', 'GE0/2/5.303', 'GigabitEthernet0/2/5'),
    ('ubiquiti_edgeswitch', 'EdgeSwitch', '3/2', 'lag 2'),
    ('ubiquiti_edgeswitch', 'EdgeSwitch', '0/8', '0/8'),
])
@pytest.mark.parametrize('source', ['observed', 'configured', 'both', 'conflict', 'absent'])
def test_exact_owner_description(platform, family, attachment, owner, source):
    from orbitflow.capabilities.mac import MacSnapshot, MacEntry
    description = '=user controlled uplink'
    interfaces = [NS(port_name=owner, port_description=description if source in {'observed', 'both', 'conflict'} else ''),
                  NS(port_name=attachment, port_description='child only')] if ('.' in attachment or 'Efp' in attachment) else [
                      NS(port_name=owner, port_description=description if source in {'observed', 'both', 'conflict'} else '')]
    # Member and child descriptions cannot override the resolved owner.
    interfaces.append(NS(port_name='0/9', port_description='arbitrary member'))
    profile = NS(interface_name=owner, description='different' if source == 'conflict' else description if source in {'configured', 'both'} else '',
                 port_type='', service_details=(), service_mappings=())
    state = NS(interfaces=[profile], configuration=())
    snapshot = MacSnapshot((MacEntry('aa:bb:cc:dd:ee:ff', '445', attachment, 'DYNAMIC', 'dynamic'),), 1)
    row = enrich(snapshot, NS(platform=platform, device_family=family), interfaces, state)[0]
    assert row.description == ('' if source in {'conflict', 'absent'} else description)
    assert len(enrich(snapshot, NS(platform=platform, device_family=family), interfaces, state)) == 1
    assert row.resolution.physical == owner


@pytest.mark.parametrize('attachment,platform', [('CPU', 'cisco_ios'), ('Drop', 'cisco_ios'),
    ('Vl545', 'cisco_ios'), ('PW123', 'cisco_xe'), ('(192.0.2.1, 123)', 'cisco_xr'),
    ('5/1', 'ubiquiti_edgeswitch'), ('3/2', 'ubiquiti_edgeswitch')])
def test_nonphysical_description_blank(attachment, platform):
    from orbitflow.capabilities.mac_forwarding import interface_description
    interfaces = [NS(port_name=attachment, port_description='must not leak')]
    if attachment == '3/2':
        interfaces = [NS(port_name='lag 2', port_description='one'), NS(port_name='LAG 2', port_description='two')]
    resolution = resolve_attachment(attachment, platform, interfaces)
    assert interface_description(resolution, platform, interfaces) == ''


@pytest.mark.parametrize('vlan_failure', [False, True])
@pytest.mark.parametrize('description', ['=SUM(1,2)', '+formula', '-formula', '@formula', 'line\x00break password-fixture'])
def test_description_workbook_and_recovery(tmp_path, monkeypatch, description, vlan_failure):
    import orbitflow.mac_lookup as app
    from orbitflow.capabilities.mac import MacEntry, MacSnapshot
    @contextmanager
    def connection(*a, **kw):
        yield object()
    monkeypatch.setattr(app, 'connect_device', connection)
    monkeypatch.setattr(app, 'DeviceCLI', connection)
    monkeypatch.setattr(app.DeviceInventoryResolver, 'resolve', lambda *a, **kw:
                        NS(platform='cisco_ios', device_family='C3750X', hostname='lab', hardware_model='C3750X'))
    monkeypatch.setattr(app.MacService, 'collect', lambda *a, **kw:
                        MacSnapshot((MacEntry('aa:bb:cc:dd:ee:ff', '445', 'Gi1/0/1', 'DYNAMIC', 'dynamic'),) * 2, 2))
    monkeypatch.setattr(app.InterfaceService, 'collect', lambda *a, **kw:
                        [NS(port_name='GigabitEthernet1/0/1', port_description=description)])
    def optional_vlan(*a, **kw):
        if vlan_failure:
            raise ValueError('optional failure')
        return None
    monkeypatch.setattr(app.VlanService, 'collect', optional_vlan)
    path = run_lookup([dict(management_ip='192.0.2.1', username='test-user', password='password-fixture')], None,
                      reports_dir=tmp_path/'reports', spool_root=tmp_path/'runs',
                      inventory_path=tmp_path/'inventory.json', log_root=tmp_path/'logs')
    recovered = export_mac_spool(next((tmp_path/'runs').iterdir()), tmp_path/'recovered.xlsx')
    for output in (path, recovered):
        book = load_workbook(output)
        sheet = book['MAC Addresses']
        assert MAIN_COLUMNS == ('Device IP', 'Device Name', 'Equipment Type', 'Physical Interface name', 'Port Type', 'Mac address', 'Interface Description')
        assert sheet.max_row == 3 and sheet.max_column == 7
        assert [c.value for c in sheet[2]][:6] == ['192.0.2.1', 'lab', 'C3750X', 'GigabitEthernet1/0/1', 'Unknown', 'aa:bb:cc:dd:ee:ff']
        assert sheet['G2'].data_type == 's'
        assert sheet['G2'].value == description.replace('\x00', '').replace('password-fixture', '[REDACTED]')
        assert book['Totals']['C2'].value == '2'
        book.close()


def test_description_from_vlan_when_interface_collection_unavailable():
    from orbitflow.capabilities.mac_forwarding import interface_description
    state = configured_state('interface GigabitEthernet0/24\n description configured uplink\n!', 'cisco_ios')
    state.interfaces = [NS(interface_name='GigabitEthernet0/24', description='configured uplink', port_type='')]
    resolution = resolve_attachment('Gi0/24+Efp3999', 'cisco_ios', (), state)
    assert resolution.physical == 'GigabitEthernet0/24'
    assert interface_description(resolution, 'cisco_ios', (), state) == 'configured uplink'
    assert interface_description(resolve_attachment('Gi0/25+Efp3999', 'cisco_ios', (), state), 'cisco_ios', (), state) == ''
