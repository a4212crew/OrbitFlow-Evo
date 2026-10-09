"""Read-only rewrite syntax and family projection; no network access."""

import json
import zipfile

import pytest
from openpyxl import load_workbook

from orbitflow.compliance import evaluate_vlan_compliance
from orbitflow.execution import DeviceOutcome
from orbitflow.methodology_report import export_methodology_spool
from orbitflow.result_spool import ResultSpool
from test_methodology_resolution import resolve, snapshot
from test_vlan_compliance import POLICY


XR = """interface GigabitEthernet0/0/0/6.2441 l2transport
 encapsulation dot1q 2441
 rewrite ingress tag pop 1 symmetric
!
l2vpn
 bridge group RSVD-RSP34
  bridge-domain RSVD-RSP34
   interface GigabitEthernet0/0/0/6.2441
"""
ME = """interface GigabitEthernet0/1
 switchport mode trunk
 switchport trunk allowed vlan none
 service instance 2244 ethernet
  encapsulation dot1q 2244
  rewrite ingress tag pop 1 symmetric
  bridge-domain 2244
 service instance 2245 ethernet
  encapsulation dot1q 2245
  rewrite ingress tag pop 1 symmetric
  bridge-domain 2245
"""


@pytest.mark.parametrize('family,config', [('NCS540', XR), ('ME3600X', ME), ('ASR920', ME)])
def test_operator_rewrites_preserve_policy_and_exact_evidence(family, config):
    ctx, state = snapshot(config, family)
    # Removing a statement shifts downstream line numbers: compare against the
    # original snapshot with only review facts removed, preserving provenance.
    from dataclasses import replace
    def strip(nodes):
        return tuple(replace(n, children=strip(n.children)) for n in nodes if n.kind != 'tag_rewrite')
    baseline = replace(state, configuration=strip(state.configuration))
    assert evaluate_vlan_compliance(ctx, [], state, POLICY) == evaluate_vlan_compliance(ctx, [], baseline, POLICY)
    records = resolve(config, family)
    assert 'UNSUPPORTED_FORWARDING_SYNTAX' not in json.dumps(records)
    assert 'MIXED_INTERFACE_CONSTRUCTS' not in json.dumps(records)
    services = [r for r in records if r['methodology'] in (['M03'], ['M04'])]
    assert len(services) == (1 if family == 'NCS540' else 2)
    for row in services:
        if family == 'NCS540':
            # The supplied excerpt omits the parent stanza; retain that genuine
            # existing finding, without warning about the recognized rewrite.
            assert [p['code'] for p in row['configuration_findings']] == ['PARENT_INTERFACE_NOT_FOUND']
        else:
            assert not row['review_needed']
        profile, = row['mapping']['rewrite_profiles']
        assert profile['direction'] == 'ingress' and profile['operation'] == 'pop'
        assert profile['tag_count'] == 1 and profile['symmetric'] is True
        assert profile['parameters'] == '1'
        assert profile['config_interface_name'] == row['config_interface_name']
        assert profile['service_instance_id'] == row['service_instance_id']
        assert profile['platform_support'] == 'not_assessed'
        proof, = profile['evidence']
        assert proof['excerpt'] == config.splitlines()[proof['line'] - 1]
        assert proof['source_filename'] == 'fixture.cfg'
        assert proof in row['evidence']
    if family == 'NCS540':
        assert services[0]['parent_interface'] == 'GigabitEthernet0/0/0/6'
        assert services[0]['mapping']['outer_vlan'] == [2441]
        assert services[0]['mapping']['bridge_domains'] == [['RSVD-RSP34', 'RSVD-RSP34']]
    if family == 'ME3600X':
        conventional = next(r for r in records if r['methodology'] == ['M01'])
        assert not conventional['review_needed']
        assert 'rewrite_profiles' not in conventional['mapping']
        assert 'tagged' not in conventional['mapping']
        assert conventional['mapping']['configured_switching'][1]['value'] == 'NONE'
        assert [r['mapping']['resolved_bridge_domain'] for r in services] == [['2244'], ['2245']]


def test_xr_complete_parent_context_resolves_and_duplicate_rewrite_is_not_conflict():
    config = 'interface GigabitEthernet0/0/0/6\n!\n' + XR.replace(
        ' rewrite ingress tag pop 1 symmetric', ' rewrite ingress tag pop 1 symmetric\n rewrite ingress tag pop 1 symmetric')
    row, = resolve(config, 'NCS540')
    assert row['status'] == 'resolved'
    assert len(row['mapping']['rewrite_profiles']) == 2


@pytest.mark.parametrize('tail,operation,count,tags,translation,symmetric', [
    ('pop 2 symmetric', 'pop', 2, [], '', True),
    ('push dot1q 17', 'push', 1, [['dot1q', 17]], '', False),
    ('push dot1q 17 second-dot1q 18 symmetric', 'push', 2, [['dot1q', 17], ['second-dot1q', 18]], '', True),
    ('push dot1ad 17 dot1q 18', 'push', 2, [['dot1ad', 17], ['dot1q', 18]], '', False),
    ('translate 1-to-1 dot1q 17 symmetric', 'translate', 1, [['dot1q', 17]], '1-to-1', True),
    ('translate 2-to-1 dot1ad 17', 'translate', 2, [['dot1ad', 17]], '2-to-1', False),
    ('translate 1-to-2 dot1q 17 second-dot1q 18', 'translate', 1, [['dot1q', 17], ['second-dot1q', 18]], '1-to-2', False),
    ('translate 2-to-2 dot1ad 17 dot1q 18 symmetric', 'translate', 2, [['dot1ad', 17], ['dot1q', 18]], '2-to-2', True),
])
@pytest.mark.parametrize('family', ['ME3600X', 'ASR920', 'NCS540'])
def test_unambiguous_variants_are_observation_only(family, tail, operation, count, tags, translation, symmetric):
    config = (XR if family == 'NCS540' else ME).replace('pop 1 symmetric', tail)
    rows = resolve(config, family)
    profiles = [p for r in rows for p in r['mapping'].get('rewrite_profiles', [])]
    assert profiles
    for p in profiles:
        assert (p['operation'], p['tag_count'], p['output_tags'], p['translation'], p['symmetric']) == (
            operation, count, tags, translation, symmetric)
        assert p['parameters'] == tail.removeprefix(operation + ' ').removesuffix(' symmetric')
        assert p['syntax_status'] == 'recognized' and p['platform_support'] == 'not_assessed'


@pytest.mark.parametrize('tail', ['pop 3', 'pop', 'push dot1q 4095', 'push dot1q 17 arbitrary-value',
    'translate 1-to-2 dot1q 17', 'translate 3-to-1 dot1q 17',
    'translate 2-to-1 dot1q 17 second-dot1q 18', 'pop 1 token hiddenvalue'])
def test_unknown_rewrites_review_only_and_sanitized(tail):
    config = ME.replace('rewrite ingress tag pop 1 symmetric', 'rewrite ingress tag ' + tail, 1)
    rows = resolve(config, 'ME3600X')
    first = next(r for r in rows if r['service_instance_id'] == '2244')
    second = next(r for r in rows if r['service_instance_id'] == '2245')
    assert first['review_needed'] and not second['review_needed']
    assert 'UNSUPPORTED_FORWARDING_SYNTAX' in json.dumps(first)
    assert 'hiddenvalue' not in json.dumps(rows) and 'arbitrary-value' not in json.dumps(rows)


@pytest.mark.parametrize('family,config', [
    ('C3750X', ME), ('C3850', ME),
    ('ME3600X', 'interface Gi0/1\n switchport mode trunk\n rewrite ingress tag pop 1 symmetric'),
    ('NCS540', XR.replace(' l2transport', '')),
    ('NCS540', XR.replace('GigabitEthernet0/0/0/6.2441', 'GigabitEthernet0/0/0/6')),
    ('NE05E', 'interface GE0/1\n rewrite ingress tag pop 1 symmetric'),
])
def test_rewrite_family_and_service_scope_requires_review(family, config):
    rows = resolve(config, family)
    assert any(r['review_needed'] and any(p['code'] in {'REWRITE_SCOPE_UNSUPPORTED', 'UNSUPPORTED_FORWARDING_SYNTAX'}
                                        for p in r['configuration_findings']) for r in rows)


@pytest.mark.parametrize('replacement', ['switchport mode access', 'switchport mode trunk\n switchport access vlan 445'])
def test_coexistence_exception_does_not_hide_conflicts(replacement):
    rows = resolve(ME.replace('switchport mode trunk', replacement), 'ME3600X')
    assert 'MIXED_INTERFACE_CONSTRUCTS' in json.dumps(rows)
    assert next(r for r in rows if r['methodology'] == ['M01'])['review_needed']


def test_conflicting_rewrites_and_qinq_remain_distinct():
    config = XR.replace('dot1q 2441', 'dot1q 2441 second-dot1q 17').replace(
        'rewrite ingress tag pop 1 symmetric',
        'rewrite ingress tag pop 1 symmetric\n rewrite ingress tag pop 2 symmetric')
    row, = resolve(config, 'NCS540')
    assert row['mapping']['outer_vlan'] == [2441] and row['mapping']['inner_vlan'] == [17]
    assert row['review_needed'] and 'CONFLICTING_TAG_REWRITES' in json.dumps(row)
    assert len(row['mapping']['rewrite_profiles']) == 2


def test_rewrite_workbook_roundtrip_no_secret_and_determinism(tmp_path):
    rows = resolve(ME + '  rewrite ingress tag pop 1 password hiddenvalue', 'ME3600X')
    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={'methodologies': rows})
    path = export_methodology_spool(spool.path, tmp_path / 'review.xlsx')
    with zipfile.ZipFile(path) as archive:
        assert b'hiddenvalue' not in b''.join(archive.read(n) for n in archive.namelist())
    book = load_workbook(path)
    first = [list(tab.values) for tab in book]
    headers, *values = book['Methodology Resolution'].values
    service = next(dict(zip(headers, row)) for row in values if row[headers.index('Service Instance')] == '2244')
    assert json.loads(service['Mapping'])['rewrite_profiles'][0]['operation'] == 'pop'
    assert 'rewrite ingress tag pop 1 symmetric' in service['Configuration Evidence']
    assert all(c.data_type != 'f' for tab in book for row in tab for c in row)
    book.close()
    export_methodology_spool(spool.path, path)
    book = load_workbook(path)
    assert first == [list(tab.values) for tab in book]
    details = [json.loads(row[-1]) for row in list(book['Details'].values)[1:]]
    assert details == rows
    book.close()
