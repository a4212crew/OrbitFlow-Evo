"""Synthetic Issue #147 fixtures: never collected device logs."""
from dataclasses import asdict, replace
from unittest.mock import Mock
from openpyxl import load_workbook
import pytest
from orbitflow.vendors.aggregation import parse_aggregation, identity
from orbitflow.capabilities.aggregation import AggregationService
from orbitflow.reporting import build_rows, export_report_spool, INTERFACE_COLUMNS
from orbitflow.result_spool import ResultSpool
from orbitflow.execution import DeviceOutcome
from test_device_reporting import context, interface


FAMILIES = [
    ('C3750X', 'cisco_ios', 'Po1', 'Gi1/0/1', 'channel-group 1 mode active', '!'),
    ('C3850', 'cisco_xe', 'Port-channel1', 'Gi1/0/1', 'channel-group 1 mode active', '!'),
    ('ME3600X', 'cisco_ios', 'Po1', 'Gi0/1', 'channel-group 1 mode active', '!'),
    ('ASR920', 'cisco_xe', 'Po1', 'Gi0/0/1', 'channel-group 1 mode active', '!'),
    ('NCS540', 'cisco_xr', 'BE1', 'Te0/0/0/1', 'bundle id 1 mode active', '!'),
    ('NE05E', 'huawei_vrp', 'Eth-Trunk1', 'GE0/2/2', 'eth-trunk 1', '#'),
    ('EdgeSwitch', 'ubiquiti_edgeswitch', 'lag 1', '0/12', 'addport 3/1', 'exit'),
]


def fixture(family, platform, parent, child, binding, terminator):
    ctx = replace(context(platform=platform), device_family=family)
    text = (f'interface {parent}\n description parent\n{terminator}\n'
            f'interface {child}\n description unrelated-name\n shutdown\n {binding}\n'
            f' service-policy input keep-qos\n password synthetic-private\n{terminator}\nend\n')
    return ctx, text


@pytest.mark.parametrize('spec', FAMILIES)
def test_all_seven_families_down_explicit_binding_and_evidence(spec):
    family, platform, parent, child, binding, terminator = spec
    ctx, text = fixture(*spec)
    snapshot = parse_aggregation(text, ctx)
    p, c = identity(platform, parent), identity(platform, child)
    assert snapshot.complete
    assert snapshot.fields(parent) == ['Yes', 'Parent', c]
    assert snapshot.fields(child) == ['Yes', 'Child', p]
    assert snapshot.references(parent) == tuple(sorted((p, c)))
    assert snapshot.references(child) == tuple(sorted((p, c)))
    assert 'synthetic-private' not in repr(snapshot)
    assert 'service-policy input keep-qos' in dict(snapshot.evidence)[c]
    rows, _ = build_rows(ctx, [interface(parent), interface(child)], None, snapshot)
    assert len(rows) == 2
    assert all(dict(zip(INTERFACE_COLUMNS, r))['LAG'] == 'Yes' for r in rows)


@pytest.mark.parametrize('spec', FAMILIES)
def test_configured_empty_partial_unknown_and_observed_default_parent(spec):
    ctx, _ = fixture(*spec)
    parent, term = spec[2], spec[-1]
    full = f'interface {parent}\n{term}\nend\n'
    snapshot = parse_aggregation(full, ctx)
    assert snapshot.fields(parent) == ['Yes', 'Parent', '']
    partial = parse_aggregation(full.removesuffix('end\n'), ctx)
    assert partial.fields(parent) == ['Yes', 'Parent', 'Unknown']
    assert partial.problems
    default = parse_aggregation('hostname synthetic\nend\n', ctx, [parent])
    assert default.fields(parent) == ['Yes', 'Parent', '']


def test_subinterfaces_parent_child_projection_and_nested_evc():
    ctx = context()
    text = '''interface Po1
 description parent
!
interface Po1.445
 encapsulation dot1q 445
!
interface Gi0/1
 channel-group 1 mode active
 service instance 12 ethernet
  encapsulation dot1q 12
  service-policy input keep
 !
 shutdown
!
interface Gi0/2
 channel-group 1 mode active
!
end
'''
    snap = parse_aggregation(text, ctx)
    assert len(snap.bindings) == 2
    assert snap.fields('Port-channel1.445') == ['', '', '']
    assert set(snap.references('Po1')) == {'port-channel1', 'port-channel1.445', 'gigabitethernet0/1', 'gigabitethernet0/2'}
    assert set(snap.references('Gi0/1')) == {'port-channel1', 'gigabitethernet0/1'}
    assert ' shutdown' in dict(snap.evidence)['gigabitethernet0/1']
    assert ' service instance 12 ethernet' in dict(snap.evidence)['gigabitethernet0/1']
    assert snap.references('Po1.445') == ('port-channel1.445',)


def test_conflicting_membership_and_missing_parent_are_nonfatal_unknown():
    text = 'interface Po1\n!\ninterface Gi0/1\n channel-group 1\n channel-group 2\n!\nend\n'
    snap = parse_aggregation(text, context())
    assert not snap.complete and not snap.bindings
    assert 'contradictory member binding' in snap.problems
    snap = parse_aggregation('interface Gi0/1\n channel-group 1\n!\nend\n', context())
    assert not snap.complete and not snap.parents
    assert snap.fields('Gi0/1') == ['Yes', 'Child', 'port-channel1']


def test_operational_counts_descriptions_and_partner_tables_never_bind():
    for spec in FAMILIES:
        ctx, text = fixture(*spec)
        text = text.replace(' ' + spec[4], ' description ' + spec[4])
        text += '\n1Po1(RU)LACP Gi0/1(bndl)\nBE1 0/0/2\nLocal GE0/2/2\nPartner GE0/2/2\n'
        assert parse_aggregation(text, ctx).bindings == ()


def test_private_key_and_vendor_credentials_sanitized_without_losing_following_lines():
    text = '''interface Gi0/1
 authentication key-chain very-private
 uisp-key secret-value
 -----BEGIN RSA PRIVATE KEY-----
 private-key-body
 -----END RSA PRIVATE KEY-----
 description retained
!
end
'''
    snap = parse_aggregation(text, context())
    serialized = repr(asdict(snap))
    for secret in ('very-private', 'secret-value', 'private-key-body'):
        assert secret not in serialized
    assert 'description retained' in serialized


@pytest.mark.parametrize('long', [False, True])
def test_spool_evidence_recovery_atomic_retry_and_lossless_overflow(tmp_path, monkeypatch, long):
    from orbitflow import reporting
    from orbitflow.execution import DeviceOutcome
    text = 'interface Po1\n description ' + ('x' * 65000 if long else 'parent') + '\n!\nend\n'
    snap = parse_aggregation(text, context())
    rows, db = build_rows(context(), [interface('Po1')], None, snap)
    spool = ResultSpool.create(tmp_path / 'spools', 'interface_vlan_report', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1, ()), payload=[rows, db, [], asdict(snap)], target='192.0.2.1')
    path = tmp_path / 'report.xlsx'
    path.write_bytes(b'original')
    writer = reporting.write_tables
    monkeypatch.setattr(reporting, 'write_tables', Mock(side_effect=OSError('synthetic')))
    with pytest.raises(OSError):
        export_report_spool(spool.path, path)
    assert path.read_bytes() == b'original'
    assert spool.path.exists()
    monkeypatch.setattr(reporting, 'write_tables', writer)
    export_report_spool(spool.path, path, cleanup=False)
    book = load_workbook(path)
    assert book['Interfaces']['E2'].comment
    if long:
        evidence = list(book['Configuration_Evidence'].values)[1:]
        assert ''.join(row[-1] for row in evidence) == dict(snap.evidence)['port-channel1']
        assert [int(row[4]) for row in evidence] == [1, 2, 3]
        assert all(len(row[-1]) <= 30000 for row in evidence)
        assert evidence[0][3] in book['Interfaces']['E2'].comment.text
    else:
        assert book.sheetnames == ['Interfaces', 'VLAN_Database', 'Run_Errors']
        assert 'description parent' in book['Interfaces']['E2'].comment.text
    book.close()

@pytest.mark.parametrize('spec', FAMILIES)
def test_service_borrows_same_cli_and_preserves_device_scope(monkeypatch, spec):
    from contextlib import contextmanager
    from orbitflow.capabilities import aggregation
    ctx, text = fixture(*spec)
    session, shell = object(), Mock()
    shell.run_command.return_value = text
    @contextmanager
    def manager(given_session, *, cli, **kwargs):
        assert given_session is session and cli is shell
        yield shell
    monkeypatch.setattr(aggregation, 'open_cisco_cli', manager)
    monkeypatch.setattr(aggregation, 'open_prompt_cli', manager)
    snap = AggregationService().collect(session, ctx, cli=shell, interfaces=[interface(spec[2])])
    assert snap.device_id == ctx.device_id and snap.device_family == ctx.device_family
    shell.run_command.assert_called_once_with('display current-configuration' if spec[1] == 'huawei_vrp' else 'show running-config', timeout=30)
    shell.close.assert_not_called()
    other = parse_aggregation(text, replace(ctx, device_id='other-device'))
    assert other.device_id != snap.device_id


@pytest.mark.parametrize('failed', [False, True])
def test_reporting_stage_integration_and_isolated_failure(tmp_path, monkeypatch, failed):
    from io import StringIO
    from test_device_reporting import install_fakes, CONFIG
    from orbitflow import reporting
    install_fakes(monkeypatch)
    snap = parse_aggregation('interface Po1\n!\ninterface Gi0/1\n channel-group 1\n!\nend\n', context())
    calls = []
    def collect(session, ctx, *, cli, interfaces):
        assert cli is session.cli and ctx is session.context
        calls.append(ctx.management_ip)
        if failed and ctx.management_ip == '192.0.2.1':
            raise ValueError('synthetic-secret')
        return snap
    monkeypatch.setattr(reporting, 'AggregationService', Mock(return_value=Mock(collect=collect)))
    targets = [dict(management_ip=f'192.0.2.{n}', username='synthetic-user', password='synthetic-password') for n in (1, 2)]
    path = reporting.run_report(targets, CONFIG, inventory_path=tmp_path/'inventory.json', reports_dir=tmp_path/'reports', log_root=tmp_path/'logs', output=StringIO())
    book = load_workbook(path)
    rows = list(book['Interfaces'].values)[1:]
    assert len(rows) == 2 and calls == ['192.0.2.1', '192.0.2.2']
    assert rows[1][5:8] == ('Yes', 'Child', 'port-channel1')
    errors = list(book['Run_Errors'].values)[1:]
    assert len(errors) == int(failed)
    if failed:
        assert errors[0][2] == 'aggregation'
        assert rows[0][5:8] == (None, None, None)
    book.close()


def test_normal_workbook_stores_stanzas_once_and_views_reference_them(tmp_path):
    ctx, text = fixture(*FAMILIES[0])
    snap = parse_aggregation(text, ctx)
    records = [interface(FAMILIES[0][2]), interface(FAMILIES[0][3])]
    rows, db = build_rows(ctx, records, None, snap)
    spool = ResultSpool.create(tmp_path/'spools', 'interface_vlan_report', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload=[rows, db, [], asdict(snap)])
    path = tmp_path/'report.xlsx'
    export_report_spool(spool.path, path)
    book = load_workbook(path)
    comments = [book['Interfaces'].cell(i, 5).comment.text for i in (2, 3)]
    assert '\n'.join(comments).count('description parent') == 1
    assert '\n'.join(comments).count('description unrelated-name') == 1
    assert all('port-channel1' in comment and 'gigabitethernet1/0/1' in comment for comment in comments)
    book.close()


def test_edgeswitch_six_observed_default_parents_and_populated_lag():
    ctx = replace(context(platform='ubiquiti_edgeswitch'), device_family='EdgeSwitch')
    observed = [f'3/{n}' for n in range(1, 7)]
    text = 'hostname synthetic\ninterface lag 2\nexit\ninterface 0/12\n addport 3/2\nexit\ninterface 0/16\n shutdown\n addport 3/2\nexit\nend\n'
    snap = parse_aggregation(text, ctx, observed)
    assert len(snap.parents) == 6
    assert snap.fields('lag 2') == ['Yes', 'Parent', '0/12, 0/16']
    assert all(snap.fields(n) == ['Yes', 'Parent', ''] for n in observed if n != '3/2')
    empty = parse_aggregation('hostname synthetic\nend\n', ctx, observed)
    assert all(empty.fields(n) == ['Yes', 'Parent', ''] for n in observed)


def test_xr_two_down_members_subinterface_and_zero_active_counts():
    ctx = replace(context(platform='cisco_xr'), device_family='NCS540')
    text = 'interface Bundle-Ether1\n shutdown\n!\ninterface BE1.445 l2transport\n encapsulation dot1q 445\n!\n'
    for n in (1, 2):
        text += f'interface Te0/0/0/{n}\n shutdown\n bundle id 1 mode active\n!\n'
    snap = parse_aggregation(text + 'end\nBE1 0/0/2\n', ctx)
    assert snap.complete
    assert snap.fields('BE1')[2] == 'tengige0/0/0/1, tengige0/0/0/2'
    assert snap.fields('BE1.445') == ['', '', '']
    assert snap.subinterfaces == (('bundle-ether1.445', 'bundle-ether1'),)


def test_duplicate_alias_stanzas_and_malformed_binding_are_unknown():
    ctx = context()
    text = 'interface Po1\n!\ninterface Gi0/1\n channel-group invalid\n!\nend\n'
    assert 'malformed member binding' in parse_aggregation(text, ctx).problems
    text = 'interface Po1\n!\ninterface Port-channel1\n!\nend\n'
    snap = parse_aggregation(text, ctx)
    assert not snap.complete and 'duplicate interface stanza' in snap.problems


def test_rejects_unavailable_configuration_and_unsupported_family():
    for output in ('', '% Invalid input detected'):
        with pytest.raises(ValueError, match='unavailable|rejected'):
            parse_aggregation(output, context())
    with pytest.raises(ValueError, match='unsupported'):
        AggregationService().collect(object(), replace(context(), device_family='unknown'))


@pytest.mark.parametrize('spec', FAMILIES)
@pytest.mark.parametrize('borrow', [False, True])
def test_real_shared_cli_setup_and_shell_ownership(spec, borrow):
    from test_cli_lifecycle import Channel, session_for
    from orbitflow.vendors.common import DeviceCLI
    ctx, text = fixture(*spec)
    huawei = spec[1] == 'huawei_vrp'
    prompt = '<synthetic>' if huawei else 'synthetic#'
    command = 'display current-configuration' if huawei else 'show running-config'
    channel = Channel(prompt, {'': '', 'screen-length 0 temporary': '',
        'terminal length 0': '', 'terminal width 511': '', command: text})
    session, client = session_for([channel])
    if borrow:
        with DeviceCLI(session) as cli:
            snap = AggregationService().collect(session, ctx, cli=cli)
            assert channel.close_calls == 0
    else:
        snap = AggregationService().collect(session, ctx)
    assert snap.complete and len(snap.bindings) == 1
    assert channel.close_calls == 1
    assert len(client.opened) == 1 and not client.closed


def test_large_parent_reference_list_also_uses_lossless_overflow(tmp_path):
    text = 'interface Po1\n!\n' + ''.join(f'interface Gi1/0/{n}\n channel-group 1\n!\n' for n in range(1, 501)) + 'end\n'
    snap = parse_aggregation(text, context())
    rows, db = build_rows(context(), [interface('Po1')], None, snap)
    spool = ResultSpool.create(tmp_path/'spools', 'interface_vlan_report', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload=[rows, db, [], asdict(snap)])
    path = tmp_path/'report.xlsx'
    export_report_spool(spool.path, path)
    book = load_workbook(path)
    comment = book['Interfaces']['E2'].comment.text
    assert len(comment) <= 32767 and comment.endswith(':references')
    parts = [row for row in list(book['Configuration_Evidence'].values)[1:] if row[3].endswith(':references')]
    references = ''.join(row[-1] for row in parts)
    assert len(references.split(', ')) == 501
    assert references.count(':port-channel1') == 1
    book.close()


def test_unrecognized_or_wrapped_partial_output_never_proves_empty():
    ctx = context()
    snap = parse_aggregation('unexpected output\nend\n', ctx, ['Po1'])
    assert not snap.complete and snap.fields('Po1')[-1] == 'Unknown'
    assert 'unsupported configuration shape' in snap.problems
    snap = parse_aggregation('interface Po1\n!\ninterface Gi0/1\n channel-group\n 1 mode active\n!\nend\n', ctx)
    assert not snap.complete and 'malformed member binding' in snap.problems


def test_edgeswitch_fastpath_outer_configure_exit_completeness():
    ctx = replace(context(platform='ubiquiti_edgeswitch'), device_family='EdgeSwitch')
    text = '!Current Configuration:\nconfigure\nvlan database\n vlan 10\nexit\ninterface lag 2\nexit\ninterface 0/12\n addport 3/2\nexit\nexit\n'
    snap = parse_aggregation(text, ctx, ['3/1', '3/2'])
    assert snap.complete
    assert snap.fields('3/1') == ['Yes', 'Parent', '']
    assert snap.fields('3/2') == ['Yes', 'Parent', '0/12']
    partial = parse_aggregation(text.removesuffix('exit\n'), ctx, ['3/1'])
    assert not partial.complete and partial.fields('3/1')[-1] == 'Unknown'


@pytest.mark.parametrize('prompt', ['synthetic#', '(synthetic) #'])
def test_fastpath_closed_export_trailing_prompt_and_sections(prompt):
    # Constructed firmware-shaped export, not a captured owner transcript.
    ctx = replace(context(platform='ubiquiti_edgeswitch'), device_family='EdgeSwitch')
    text = ('!Current Configuration:\nconfigure\nvlan database\n vlan 10\nexit\n'
            'interface lag 2\nexit\ninterface 0/12\n shutdown\n addport 3/2\nexit\n'
            'interface 0/16\n addport 3/2\nexit\nexit\n! trailing comment\n' + prompt + '\n')
    snap = parse_aggregation(text, ctx, [f'3/{n}' for n in range(1, 7)])
    assert snap.complete
    assert snap.fields('3/2')[-1] == '0/12, 0/16'
    assert all(snap.fields(f'3/{n}')[-1] == '' for n in (1, 3, 4, 5, 6))
    partial = parse_aggregation(text.replace('exit\n! trailing', '! trailing'), ctx, ['3/2'])
    assert not partial.complete
    assert partial.fields('0/12') == ['Yes', 'Child', '3/2']
    assert partial.fields('3/2')[-1] == 'Unknown'
    assert 'addport 3/2' in dict(partial.evidence)['0/12']
    # A section following closure is not a prompt; never silently discard it.
    extra = parse_aggregation(text + 'router unsupported-section\nexit\n', ctx, ['3/2'])
    assert not extra.complete and extra.fields('3/2')[-1] == 'Unknown'


def test_xr_indented_final_end_preserves_down_members_and_projection():
    ctx = replace(context(platform='cisco_xr'), device_family='NCS540')
    text = ('!! IOS XR Configuration\ninterface Bundle-Ether1\n shutdown\n!\n'
            'interface Bundle-Ether1.445 l2transport\n encapsulation dot1q 445\n!\n'
            'interface Te0/0/0/1\n shutdown\n bundle id 1 mode active\n!\n'
            'interface Te0/0/0/2\n shutdown\n bundle id 1 mode active\n end\n')
    snap = parse_aggregation(text, ctx)
    assert snap.complete and len(snap.bindings) == 2
    assert snap.fields('BE1.445') == ['', '', '']
    assert len(snap.references('BE1')) == 4
    assert len(snap.references('Te0/0/0/1')) == 2
    partial = parse_aggregation(text.removesuffix(' end\n'), ctx)
    assert not partial.complete and len(partial.bindings) == 2


def test_safe_configuration_failure_categories():
    from orbitflow.vendors.aggregation import (
        AggregationConfigurationUnavailable, AggregationConfigurationRejected,
        AggregationParserUnsupported)
    with pytest.raises(AggregationConfigurationUnavailable):
        parse_aggregation('', context())
    with pytest.raises(AggregationConfigurationRejected):
        parse_aggregation('% Invalid input detected', context())
    with pytest.raises(AggregationParserUnsupported):
        AggregationService().collect(object(), replace(context(), device_family='unknown'))
