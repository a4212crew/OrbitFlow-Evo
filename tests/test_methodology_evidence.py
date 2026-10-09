"""Synthetic CLI only: complete disclosure is independent of interpretation."""
from dataclasses import asdict
import json
import zipfile

import pytest
from openpyxl import load_workbook

from orbitflow import methodology_report as report
from orbitflow.compliance import evaluate_vlan_compliance
from orbitflow.compliance.methodology import resolve_methodologies
from orbitflow.execution import DeviceOutcome
from orbitflow.result_spool import ResultSpool
from orbitflow.vendors.configuration_facts import observe_configuration
from test_methodology_resolution import snapshot
from test_vlan_compliance import POLICY


@pytest.mark.parametrize('family,config,expected', [
    ('ME3600X', 'interface FutureEthernet0/7.123\n encapsulation mpls\n', ' encapsulation mpls'),
    ('ASR920', 'bridge-domain named-domain\n member FutureEthernet9/2 service-instance 42 split-horizon group 3',
     ' member FutureEthernet9/2 service-instance 42 split-horizon group 3'),
    ('NCS540', 'interface PW-Ether987\n encapsulation mpls\n', ' encapsulation mpls'),
    ('NE05E', 'interface GE0/7.987\n\tqinq future-mode pe-vid 445 ce-vid 545',
     '\tqinq future-mode pe-vid 445 ce-vid 545'),
    ('EdgeSwitch', 'interface 3/77\n vlan participation future-mode named-domain\nexit',
     ' vlan participation future-mode named-domain'),
    ('C3850', 'interface FutureEthernet1/7\n switchport trunk allowed vlan future-operation 445',
     ' switchport trunk allowed vlan future-operation 445'),
    ('ME3600X', 'bridge-domain monkey-domain\n member auth-link/1 service-instance 7 split-horizon',
     ' member auth-link/1 service-instance 7 split-horizon'),
    ('ME3600X', 'pseudowire-class future-pw\n encapsulation mpls', ' encapsulation mpls'),
    ('ASR920', 'interface Gi0/1\n service instance 7 ethernet\n  encapsulation dot1q 445\n'
     '  bridge-domain 445\n   split-horizon group 2', '   split-horizon group 2'),
    ('NE05E', 'vsi future-vsi\n pwsignal ldp\n  peer 192.0.2.1 static', '  peer 192.0.2.1 static'),
])
def test_cross_vendor_original_evidence(family, config, expected):
    ctx, state = snapshot(config, family)
    before = evaluate_vlan_compliance(ctx, [], state, POLICY)
    rows = resolve_methodologies(ctx, [], state)
    proof = [p for r in rows for p in r['evidence'] if p['excerpt'] == expected]
    assert proof
    assert all(p['source_filename'] == 'fixture.cfg' and
               config.splitlines()[p['line'] - 1] == expected for p in proof)
    assert evaluate_vlan_compliance(ctx, [], state, POLICY) == before
    assert '[REDACTED]' not in json.dumps(rows)


@pytest.mark.parametrize('family,tail', [
    ('ME3600X', 'password'), ('ASR920', 'key-string'),
    ('NCS540', 'authentication-key'), ('NE05E', 'cipher'),
    ('NE05E', 'irreversible-cipher'), ('EdgeSwitch', 'community'),
    ('C3850', 'pre-shared'),
])
def test_credential_values_never_enter_snapshot_spool_or_workbook(tmp_path, family, tail, caplog):
    # Deliberately synthetic sentinel, not a credential copied into a fixture.
    sentinel = 'test-only-disclosure-sentinel'
    config = f'interface Gi0/1\n encapsulation experimental {tail} {sentinel}\n'
    ctx, state = snapshot(config, family)
    assert sentinel not in json.dumps([asdict(n) for n in state.configuration])
    rows = resolve_methodologies(ctx, [], state)
    assert 'credential-bearing syntax' in json.dumps(rows)
    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={'methodologies': rows})
    assert sentinel not in (spool.path / 'results.jsonl').read_text()
    book = report.export_methodology_spool(spool.path, tmp_path / 'safe.xlsx')
    with zipfile.ZipFile(book) as archive:
        assert all(sentinel.encode() not in archive.read(n) for n in archive.namelist())
    assert sentinel not in caplog.text


@pytest.mark.parametrize('payload', ['"free form disclosure"', 'value;another-command',
    '\x1b[31mterminal', '\x1cunsafe', 'https://example.invalid/value', 'a' * 64,
    'eyJsynthetic.payload.signature',
    'AbcDef012345678901234567890'])
def test_ambiguous_payloads_have_located_omissions(payload):
    ctx, state = snapshot('interface Gi0/1\n encapsulation future ' + payload, 'C3850')
    rows = resolve_methodologies(ctx, [], state)
    proof = rows[0]['evidence'][-1]
    assert proof['line'] == 2 and proof['source_filename'] == 'fixture.cfg'
    assert 'omitted:' in proof['excerpt'] and payload not in proof['excerpt']


def test_large_evidence_round_trip_recovery_and_sheet_continuations(tmp_path, monkeypatch):
    # Exceeds the old statement cutoff, cell limit and (simulated) sheet limit.
    long_line = ' encapsulation future ' + '445,545,' * 5000 + ' exact'
    config = 'interface Gi0/1\n' + '\n'.join(
        f' encapsulation future-mode domain-{i}' for i in range(600)) + '\n' + long_line
    ctx, state = snapshot(config, 'C3850')
    rows = resolve_methodologies(ctx, [], state)
    expected = {(p['source_filename'], p['line']): p['excerpt'] for p in rows[0]['evidence']}
    assert expected['fixture.cfg', 602] == long_line
    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 3)
    with spool.collection():
        for position in range(1, 4):
            spool.append(DeviceOutcome(position), payload={'methodologies': rows})
    monkeypatch.setattr(report, 'SHEET_ROWS', 200)
    path = report.export_methodology_spool(spool.path, tmp_path / 'large.xlsx')

    def read():
        book = load_workbook(path, read_only=True)
        content = {sheet.title: list(sheet.values) for sheet in book}
        assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
        book.close()
        return content

    first = read()
    assert tuple(first['Methodology Resolution'][0]) == report.COLUMNS
    summary = dict(zip(report.COLUMNS, first['Methodology Resolution'][1]))
    assert 'Input Position 1, Record 1' in summary['Configuration Evidence']
    fragments, cli = {}, {}
    for title, table in first.items():
        assert len(table) <= 200
        for row in table[1:]:
            if title.startswith('Details'):
                position, record, part, text = row
                fragments.setdefault((position, record), []).append((int(part), text))
            if title.startswith('Evidence Details'):
                position, record, source, line, part, text = row
                cli.setdefault((position, source, int(line)), []).append((int(part), text))
    for position in ('1', '2', '3'):
        recovered = json.loads(''.join(text for _, text in sorted(fragments[position, '1'])))
        assert recovered == rows[0]
        assert {(source, line): ''.join(t for _, t in sorted(parts))
                for (pos, source, line), parts in cli.items() if pos == position} == expected
    report.export_methodology_spool(spool.path, path)
    assert read() == first


def test_export_failure_keeps_spool_for_retry(tmp_path, monkeypatch):
    ctx, state = snapshot('interface Gi0/1\n encapsulation mpls', 'ME3600X')
    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={'methodologies': resolve_methodologies(ctx, [], state)})
    writer = report.write_tables
    def fail(*args, **kwargs):
        raise OSError('synthetic export failure')
    monkeypatch.setattr(report, 'write_tables', fail)
    with pytest.raises(OSError):
        report.export_methodology_spool(spool.path, tmp_path / 'retry.xlsx', cleanup=True)
    assert spool.path.exists()
    monkeypatch.setattr(report, 'write_tables', writer)
    assert report.export_methodology_spool(spool.path, tmp_path / 'retry.xlsx').exists()


def test_supporting_evidence_does_not_change_decisions_or_audit():
    config = ('interface Gi0/1\n service instance 7 ethernet\n'
              '  encapsulation dot1q 445\n  bridge-domain 445')
    ctx, base = snapshot(config, 'ASR920')
    _, more = snapshot(config + '\n   split-horizon group 2', 'ASR920')
    before = resolve_methodologies(ctx, [], base)[0]
    after = resolve_methodologies(ctx, [], more)[0]
    for field in ('methodology', 'status', 'configuration_classification', 'finding_severity', 'standard_finding'):
        assert before[field] == after[field]
    assert evaluate_vlan_compliance(ctx, [], base, POLICY) == evaluate_vlan_compliance(ctx, [], more, POLICY)


def test_original_filename_crlf_tabs_and_formula_defense(tmp_path):
    from dataclasses import replace
    config = 'interface OddEthernet-77/2.9\r\n\tencapsulation mpls\r\n'
    ctx, state = snapshot(config, 'ME3600X')
    filename = '=1+1 backup.cfg'
    state = replace(state, configuration=observe_configuration(config, ctx.platform, source_filename=filename))
    rows = resolve_methodologies(ctx, [], state)
    assert rows[0]['evidence'][-1] == dict(source_filename=filename, line=2, excerpt='\tencapsulation mpls')
    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={'methodologies': rows})
    book = load_workbook(report.export_methodology_spool(spool.path, tmp_path / 'source.xlsx'))
    assert book['Evidence Details']['C2'].value == filename
    assert book['Evidence Details']['C2'].data_type == 's'
    book.close()


def test_large_domain_context_is_complete_without_one_record_per_line():
    config = 'bridge-domain named-domain\n' + '\n'.join(
        f' member FutureEthernet0/{i} service-instance 7 split-horizon' for i in range(600))
    ctx, state = snapshot(config, 'ASR920')
    rows = resolve_methodologies(ctx, [], state)
    assert len(rows) == 1
    assert {p['line'] for p in rows[0]['evidence']} == set(range(1, 602))
    assert rows[0]['methodology'] == [] and rows[0]['review_needed']


@pytest.mark.parametrize('family,context,support', [
    ('ME3600X', 'pseudowire-class future-pw', [' encapsulation mpls', ' pw-class other-pw']),
    ('ASR920', 'pseudowire-class future-pw', [' encapsulation mpls']),
    ('NCS540', 'pseudowire-class future-pw', [' encapsulation mpls']),
    ('ASR920', 'bridge-domain 445', [' split-horizon group 2']),
    ('NE05E', 'vsi future-vsi', [' pwsignal ldp', '  peer 192.0.2.1 static']),
])
def test_unmatched_supporting_context_preserves_decisions_and_workbook(tmp_path, family, context, support):
    from collections import Counter

    config = 'interface Gi0/1\n!\n' + context
    ctx, base = snapshot(config, family)
    augmented = config + '\n' + '\n'.join(support)
    _, more = snapshot(augmented, family)
    before = resolve_methodologies(ctx, [], base)
    after = resolve_methodologies(ctx, [], more)
    supporting = [r for r in after if r['subtype'] == 'supporting_evidence']
    assert len(supporting) == 1
    existing = [r for r in after if r not in supporting]
    fields = ('record_kind', 'methodology', 'status', 'configuration_classification',
              'finding_severity', 'standard_finding', 'configuration_findings', 'standard_findings')
    assert [{k: r[k] for k in fields} for r in existing] == [{k: r[k] for k in fields} for r in before]
    for field in ('configuration_classification', 'standard_finding'):
        assert Counter(r[field] for r in existing) == Counter(r[field] for r in before)
    assert sum(r['review_needed'] for r in after) == sum(r['review_needed'] for r in before)
    assert sum(len(r['configuration_findings']) for r in after) == sum(len(r['configuration_findings']) for r in before)
    row = supporting[0]
    assert row['record_kind'] == 'service_context'
    assert row['configuration_classification'] == 'not_applicable'
    assert row['finding_severity'] == 'none'
    assert not row['review_needed'] and not row['configuration_findings']
    assert evaluate_vlan_compliance(ctx, [], base, POLICY) == evaluate_vlan_compliance(ctx, [], more, POLICY)
    expected = {line: text for line, text in enumerate(augmented.splitlines(), 1) if line >= 3}
    assert {p['line']: p['excerpt'] for p in row['evidence']} == expected
    assert all(p['source_filename'] == 'fixture.cfg' for p in row['evidence'])

    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={'methodologies': after})
    path = report.export_methodology_spool(spool.path, tmp_path / 'support.xlsx')
    book = load_workbook(path, read_only=True)
    try:
        assert tuple(next(book['Methodology Resolution'].values)) == report.COLUMNS
        record = str(after.index(row) + 1)
        summary = next(dict(zip(report.COLUMNS, values)) for values in
                       list(book['Methodology Resolution'].values)[1:] if values[1] == record)
        assert summary['Configuration Classification'] == 'not_applicable'
        assert summary['Review Needed'] == 'False'
        assert all(text in summary['Configuration Evidence'] for text in support)
        fragments = [(int(part), text) for pos, rec, part, text in
                     list(book['Details'].values)[1:] if rec == record]
        assert json.loads(''.join(text for _, text in sorted(fragments))) == row
        cli = {int(line): text for pos, rec, source, line, part, text in
               list(book['Evidence Details'].values)[1:] if rec == record and source == 'fixture.cfg'}
        assert cli == expected
    finally:
        book.close()


@pytest.mark.parametrize('tail', [' encapsulation future-mode', ' encapsulation mpls future-mode'])
def test_unknown_pseudowire_syntax_still_requires_review(tail):
    ctx, state = snapshot('pseudowire-class future-pw\n encapsulation mpls\n' + tail, 'ME3600X')
    rows = resolve_methodologies(ctx, [], state)
    problems = [r for r in rows if r['review_needed']]
    assert problems
    assert any(p['code'] == 'UNSUPPORTED_FORWARDING_SYNTAX'
               for r in problems for p in r['configuration_findings'])
    assert any(p['line'] == 3 and p['excerpt'] == tail for r in problems for p in r['evidence'])
