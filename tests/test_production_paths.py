"""Production path contracts, using synthetic sessions and isolated working directories."""
from io import StringIO
from pathlib import Path
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from orbitflow import reporting, configuration_backup as backup
from orbitflow.logging import module_logger
from test_device_reporting import install_fakes, CONFIG
from test_configuration_backup import device_session, DEVICES


@pytest.mark.parametrize('kind', ['report', 'backup'])
@pytest.mark.parametrize('override', [False, True])
def test_command_paths(monkeypatch, kind, override):
    from scripts import device_interface_vlan_report, device_configuration_backup
    command = device_interface_vlan_report if kind == 'report' else device_configuration_backup
    runner = Mock()
    monkeypatch.setattr(command, 'run_report' if kind == 'report' else 'run_backup', runner)
    monkeypatch.setattr(command, 'load_targets', Mock(return_value=[]))
    paths = {'inventory_path': 'data/inventory/inventory.json', 'log_root': 'outputs/logs'}
    paths.update({'reports_dir': 'outputs/reports/interface_vlan', 'spool_root': 'outputs/runs/interface_vlan_report'}
                 if kind == 'report' else {'backups_dir': 'outputs/backups/configuration'})
    args = ['targets.xlsx', '--proxy', 'proxy', '--cluster', 'cluster',
            '--bastion-host', 'host', '--bastion-user', 'user']
    if override:
        paths = {key: 'legacy/' + key for key in paths}
        for key, value in paths.items():
            args.extend(['--' + key.replace('_', '-'), value])
    command.main(args)
    for key, value in paths.items():
        assert runner.call_args.kwargs[key] == Path(value)


def test_default_report_recovery_paths(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    install_fakes(monkeypatch)
    store = Mock(wraps=reporting.JsonInventoryStore)
    monkeypatch.setattr(reporting, 'JsonInventoryStore', store)
    original = reporting.write_workbook
    monkeypatch.setattr(reporting, 'write_workbook', Mock(side_effect=OSError('fixture failure')))
    targets = [dict(management_ip='192.0.2.1', username='synthetic-user', password='synthetic-password')]
    with pytest.raises(OSError):
        reporting.run_report(targets, CONFIG, output=StringIO())
    store.assert_called_once_with('data/inventory/inventory.json')
    root = Path('outputs/runs/interface_vlan_report')
    spool = next(root.iterdir())
    assert 'synthetic-password' not in ''.join(p.read_text() for p in spool.glob('*.json*'))
    monkeypatch.setattr(reporting, 'write_workbook', original)
    monkeypatch.setattr(reporting, 'connect_device', Mock(side_effect=AssertionError('no recollection')))
    path = reporting.export_report_spool(spool, 'outputs/reports/interface_vlan/recovered.xlsx')
    assert path.is_file() and not spool.exists()
    path = reporting.run_report([], CONFIG, output=StringIO())
    assert path.parent == Path('outputs/reports/interface_vlan')
    assert path.name.startswith('device_interface_vlan_report_')
    assert not list(root.iterdir())
    assert list(Path('outputs/logs/reporting').glob('*/interface_vlan_report.log'))
    assert not any(Path(p).exists() for p in ('reports', 'backups', 'logs', 'data/live_validation'))


def test_default_backup_paths_security_and_isolation(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    # Legacy data must not be loaded, modified, or removed implicitly.
    legacy = Path('data/live_validation/inventory.json')
    legacy.parent.mkdir(parents=True)
    legacy.write_text('legacy sentinel')
    session, client, channel = device_session(DEVICES[0])
    channel.outputs['show running-config'] = 'hostname fixture\n! fixture-password\nend\n'
    monkeypatch.setattr(backup, 'connect_device', lambda *args: session)
    targets = [dict(management_ip='192.0.2.1', username='fixture-user', password='fixture-password'), {}]
    folder = backup.run_backup(targets, object(), output=StringIO())
    assert folder.parent == Path('outputs/backups/configuration')
    assert (folder / '.gitignore').read_text() == '*\n'
    captures = list(folder.glob('*.txt'))
    assert len(captures) == 1 and '[REDACTED]' in captures[0].read_text()
    assert 'fixture-password' not in captures[0].read_text()
    assert Path('data/inventory/inventory.json').is_file()
    assert legacy.read_text() == 'legacy sentinel'
    assert client.closes == channel.close_calls == 1
    book = load_workbook(folder / 'failed_devices.xlsx')
    assert book.active.max_row == 2
    book.close()
    logs = list(Path('outputs/logs/backup').glob('*/configuration_backup.log'))
    assert logs and 'fixture-password' not in logs[0].read_text()
    assert not any(Path(p).exists() for p in ('reports', 'backups', 'logs'))
    assert not Path('outputs/runs').exists()


def test_shared_logger_default_and_override(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    with module_logger('inventory') as (logger, path):
        logger.info('Safe event')
        assert path.is_file() and path.parts[:3] == ('outputs', 'logs', 'inventory')
    with module_logger('backup', log_root='legacy-logs') as (_, path):
        assert path.parts[0] == 'legacy-logs'
        with module_logger('transport') as (_, nested):
            assert nested.parts[0] == 'legacy-logs'
    with module_logger('transport') as (_, path):
        assert path.parts[:2] == ('outputs', 'logs')
    assert not Path('logs').exists()
