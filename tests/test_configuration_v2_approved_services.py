"""Signed-off Templates 9–16: synthetic observations, no network writes."""
from dataclasses import asdict
from datetime import datetime, timezone
import socket

import pytest
from openpyxl import load_workbook

from orbitflow.configuration.plans import PlanError, canonical
from orbitflow.configuration.v2.catalogue import BY_SHEET, CATALOGUE
from orbitflow.configuration.v2.planning import request, plan, load, save
from orbitflow.configuration.v2.workbook import create_input, read_input, export_plan, consume_decisions, EXTRA
from orbitflow.configuration.v2.observation import snapshot
from orbitflow.configuration.v2.evidence import scoped_evidence
from orbitflow.models import DeviceContext, InterfaceRecord, VlanState
from orbitflow.vendors.configuration_facts import observe_configuration
from orbitflow.vendors.cisco.vlans import parse_ios_running_config, parse_service_detail
from orbitflow.vendors.ubiquiti.vlans import parse_edgeswitch_config, parse_vlan_lookup

IP = '192.0.2.91'
NOW = datetime(2026, 10, 13, tzinfo=timezone.utc)


def template(family, function='EVC'):
    return BY_SHEET[family + ' ' + ('Description' if function == 'Interface Description' else function)]


def test_approved_template_numbering_and_no_legacy_fields():
    assert [(t.family, t.function) for t in CATALOGUE[8:16]] == [
        ('ME3600X', 'VLAN Database'), ('ME3600X', 'EVC'), ('ME3600X', 'Interface Description'),
        ('ASR920', 'EVC'), ('ASR920', 'Interface Description'), ('EdgeSwitch', 'VLAN Database'),
        ('EdgeSwitch', 'VLAN Port'), ('EdgeSwitch', 'Interface Description')]
    assert not {'pop_tags', 'service_instance'} & set(template('ME3600X').columns)
    assert not {'mode', 'vlans', 'replace_untagged', 'vlan_id', 'untagged_vlans'} & set(template('EdgeSwitch', 'VLAN Port').columns)


@pytest.mark.parametrize('name', ['Vlan445', 'BDI445', 'GigabitEthernet0/1.445'])
def test_evc_requires_supported_physical_or_aggregate_parent(name):
    with pytest.raises(PlanError, match='unsupported EVC parent'):
        request(template('ASR920').id, {**evc_inputs(), 'interface': name})


def observation(family, config):
    platform = 'ubiquiti_edgeswitch' if family == 'EdgeSwitch' else 'cisco_ios' if family == 'ME3600X' else 'cisco_xe'
    profile = '' if family == 'EdgeSwitch' else 'me3600x_evc' if family == 'ME3600X' else 'asr920_evc'
    ctx = DeviceContext('synthetic-91', IP, (IP,), 'fixture', 'fixture', platform, family, family, profile, (), 'FIXTURE', '1', '', NOW, NOW)
    ports, objects = parse_edgeswitch_config(config) if family == 'EdgeSwitch' else parse_ios_running_config(config, evc=True, vlan_database=family != 'ASR920')
    facts = observe_configuration(config, platform, preserve_full_evidence=True)
    interfaces = tuple(InterfaceRecord('fixture', IP, platform, n.value, '', 'up', 'up', NOW) for n in facts if n.kind == 'interface')
    return snapshot(ctx, interfaces, VlanState('fixture', IP, platform, ports, objects, NOW, facts))


def evc_inputs(**changes):
    return dict(management_ip=IP, interface='GigabitEthernet0/1', service_instance_id=3000,
                encapsulation='Dot1q', outer_vlan='445', rewrite='Pop 1 symmetric', bridge_domain=445, **changes)


def evc_config(family, child='', parent=''):
    return ('vlan 445\n name Customer\n!\n' if family == 'ME3600X' else '') + 'interface GigabitEthernet0/1\n' + (
        ' switchport mode trunk\n switchport trunk allowed vlan none\n' if family == 'ME3600X' else '') + parent + child + '!\nend\n'


def result(family, p, config, **extra):
    t = template(family, 'VLAN Port' if family == 'EdgeSwitch' else 'EVC')
    r = request(t.id, p)
    obs = observation(family, config)
    obs.update(extra)
    return plan([r], {IP: obs}, now=NOW, decisions={r['row_id']: 'Approve'}).to_dict()['results'][0]


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
def test_matching_evc_always_recreated_full_evidence_and_omitted_settings(family):
    config = evc_config(family, ' service instance 3000 ethernet\n  description OLD CHILD\n  encapsulation dot1q 445\n  rewrite ingress tag pop 1 symmetric\n  bridge-domain 445 split-horizon group 0\n  service-policy input CUSTOMER\n service instance 20 ethernet\n  encapsulation dot1q 800\n  bridge-domain 800\n', ' description OLD PARENT\n mtu 9000\n shutdown\n')
    p = evc_inputs()
    p['split_horizon_group'] = 0
    row = result(family, p, config)
    assert not row['no_op'] and row['deployment_ready']
    assert row['commands'][2:4] == ['no service instance 3000', 'service instance 3000 ethernet']
    assert 'bridge-domain 445 split-horizon group 0' in row['commands']
    assert not any('20 ethernet' in c or 'shutdown' in c or 'mtu' in c for c in row['commands'])
    assert {'SERVICE_INSTANCE_RECREATE_IMPACT', 'OMITTED_CHILD_SETTINGS'} <= set(row['warnings'])
    assert any('service-policy' in c for c in row['omitted_child_commands'])
    obs = observation(family, config)
    evidence = '\n'.join(e['excerpt'] for e in scoped_evidence(request(template(family).id, p), obs))
    assert 'OLD PARENT' in evidence and 'OLD CHILD' in evidence and '20 ethernet' not in evidence and '\n!' in evidence


@pytest.mark.parametrize('enc,outer,inner', [('Untagged', None, None), ('Dot1q', '445', None), ('Dot1q', '100-200', None),
                                          ('Dot1q', '2-444,446-544', None), ('QinQ', '300', '200'), ('QinQ', '100', '1-4094')])
@pytest.mark.parametrize('rewrite', ['None', 'Pop 1 symmetric', 'Pop 2 symmetric', 'Push dot1q', 'Push dot1q symmetric'])
@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
def test_approved_rewrite_matrix(family, enc, outer, inner, rewrite):
    p = {**evc_inputs(), 'encapsulation': enc, 'rewrite': rewrite}
    p.pop('outer_vlan')
    if outer:
        p['outer_vlan'] = outer
    if inner:
        p['inner_vlan'] = inner
    if rewrite.startswith('Push'):
        p['push_vlan'] = 1402
    allowed = ({'None', 'Push dot1q', 'Push dot1q symmetric'} if enc == 'Untagged' else
               {'Pop 1 symmetric', 'Push dot1q', 'Push dot1q symmetric'} if outer == '445' else
               {'None', 'Push dot1q', 'Push dot1q symmetric'} if enc == 'Dot1q' else
               {'Pop 1 symmetric', 'Pop 2 symmetric', 'Push dot1q symmetric'} if inner == '200' else
               {'None', 'Pop 1 symmetric', 'Push dot1q symmetric'})
    if rewrite not in allowed:
        with pytest.raises(PlanError, match='(?:UNVERIFIED|LAB_REJECTED)_REWRITE_COMBINATION'):
            request(template(family).id, p)
    else:
        row = result(family, p, evc_config(family))
        assert row['deployment_ready']
        if rewrite == 'None':
            assert not any('rewrite' in c for c in row['commands'])
        elif rewrite.startswith('Push'):
            assert f"rewrite ingress tag push dot1q 1402" + (' symmetric' if rewrite.endswith('symmetric') else '') in row['commands']


def test_shared_ranges_any_and_split_horizon_zero():
    config = 'interface GigabitEthernet0/1\n service instance 1 ethernet\n  encapsulation dot1q 2-444,446-544 second-dot1q any\n  bridge-domain 700 split-horizon group 0\n!'
    facts = observe_configuration(config, 'cisco_ios')
    child = facts[0].children[0]
    assert next(n.value for n in child.children if n.kind == 'encapsulation') == tuple([*range(2,445), *range(446,545)])
    assert next(n.value for n in child.children if n.kind == 'inner_vlan') == tuple(range(1,4095))
    bd = next(n for n in child.children if n.kind == 'bridge_domain')
    assert bd.children[0].value == '0'
    ports, _ = parse_ios_running_config(config, evc=True)
    assert 444 in ports[0].tagged_vlans and 445 not in ports[0].tagged_vlans
    assert ports[0].inner_vlan_expression == '1-4094'


@pytest.mark.parametrize('inner,blocked', [('200', True), ('201', False), ('190-210', True), ('any', True)])
def test_qinq_pair_overlap(inner, blocked):
    p = {**evc_inputs(), 'encapsulation': 'QinQ', 'outer_vlan': '300', 'inner_vlan': '200', 'rewrite': 'Pop 2 symmetric'}
    child = f' service instance 20 ethernet\n  encapsulation dot1q 300 second-dot1q {inner}\n  bridge-domain 800\n'
    row = result('ASR920', p, evc_config('ASR920', child))
    assert ('INGRESS_TAG_OVERLAP' in row['blockers']) == blocked


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
@pytest.mark.parametrize('parent', [' ip address 192.0.2.1 255.255.255.0\n', ' channel-group 1 mode active\n', ' switchport mode access\n'])
def test_parent_conversion_and_ownership_block(family, parent):
    config = ('vlan 445\n!\n' if family == 'ME3600X' else '') + 'interface GigabitEthernet0/1\n' + parent + '!\nend\n'
    assert not result(family, evc_inputs(), config)['deployment_ready']


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
def test_empty_initialization_preserves_shutdown(family):
    config = ('vlan 445\n!\n' if family == 'ME3600X' else '') + 'interface GigabitEthernet0/1\n shutdown\n mtu 9000\n!\nend\n'
    row = result(family, evc_inputs(), config)
    assert row['deployment_ready']
    assert ('switchport trunk allowed vlan none' if family == 'ME3600X' else 'no ip address') in row['commands']
    assert 'no shutdown' not in row['commands']


def test_operational_detail_reuses_rewrite_grammar_and_scoped_absence():
    text = 'Id: 3000\nInterface: GigabitEthernet0/1\nEncapsulation: dot1q 2-444,446-544\nRewrite: ingress tag push dot1q 1402 vlan-type 0x8100 symmetric\nBridge-domain: 700 split-horizon group 0\nState: Down\n'
    detail = parse_service_detail(text, 'GigabitEthernet0/1', 3000)
    assert detail['status'] == 'present' and detail['operational_status'] == 'down'
    rewrite = next(n['value'] for n in detail['facts'] if n['kind'] == 'tag_rewrite')
    assert rewrite['symmetric'] and rewrite['output_tags'] == [['dot1q', 1402]]
    assert parse_service_detail('Service instance 3000 interface GigabitEthernet0/1 does not exist', 'GigabitEthernet0/1', 3000)['status'] == 'absent'
    assert parse_service_detail('Service instance 20 interface GigabitEthernet0/2 does not exist', 'GigabitEthernet0/1', 3000)['status'] == 'unknown'
    assert parse_service_detail('garbage', 'GigabitEthernet0/1', 3000)['status'] == 'unknown'


EDGE = 'vlan database\nvlan 445,545,745-746,901,2436-2444\nvlan name 901 "VLAN 901"\nexit\ninterface 0/1\ndescription "Customer Site Alpha"\nvlan pvid 445\nvlan participation include 445,545\nvlan participation exclude 746\nvlan tagging 545,746\nstorm-control broadcast level 10\nmtu 9000\nexit\n'


def edge_inputs(**changes):
    return dict(management_ip=IP, interface='0/1', action='Add/Change', **changes)


def test_edge_config_and_operational_independent_facts_full_evidence():
    ports, objects = parse_edgeswitch_config(EDGE)
    assert ports[0].configured_tagged_vlans == (545,746) and ports[0].excluded_vlans == (746,)
    assert {v.object_id for v in objects} >= {'2436', '2444', '745', '746'}
    assert next(v.name for v in objects if v.object_id == '901') == 'VLAN 901'
    lookup = parse_vlan_lookup('VLAN ID: 746\nVLAN Name: Customer\nVLAN Type: Static\nPort Current Configured Tagging\n0/1 Exclude Include Tagged\n', 746)
    assert lookup['ports'] == [dict(interface='0/1', current='Exclude', configured='Include', tagging='Tagged')]
    assert parse_vlan_lookup('VLAN does not exist.', 746)['status'] == 'absent'
    r = request(template('EdgeSwitch', 'VLAN Port').id, edge_inputs(tagged_vlans='746'))
    evidence = '\n'.join(e['excerpt'] for e in scoped_evidence(r, observation('EdgeSwitch', EDGE)))
    assert 'storm-control' in evidence and evidence.endswith('exit')
    row = result('EdgeSwitch', edge_inputs(tagged_vlans='746'), EDGE)
    assert row['commands'] == ['configure', 'interface 0/1', 'vlan participation include 746', 'exit', 'exit']


def test_edge_pvid_only_tagged_pvid_global_absence_and_noops():
    row = result('EdgeSwitch', edge_inputs(pvid=901), EDGE)
    assert row['commands'][2:-2] == ['vlan participation include 901', 'vlan pvid 901']
    row = result('EdgeSwitch', edge_inputs(pvid=445, tagged_vlans='445'), EDGE)
    assert row['deployment_ready'] and row['commands'][2:-2] == ['vlan tagging 445']
    row = result('EdgeSwitch', edge_inputs(tagged_vlans='902'), EDGE)
    assert row['deployment_ready'] and 'VLAN_NOT_IN_DATABASE' in row['warnings']
    assert not any(c.startswith('vlan 902') for c in row['commands'])
    row = result('EdgeSwitch', edge_inputs(tagged_vlans='545', pvid=445), EDGE)
    assert row['no_op'] and row['commands'] == []


def test_edge_remove_order_no_replace_and_projected_pvid():
    t = template('EdgeSwitch', 'VLAN Port')
    for p in [{**edge_inputs(tagged_vlans='545'), 'action': 'Replace'}, {**edge_inputs(tagged_vlans='545', pvid=901), 'action': 'Remove'}]:
        with pytest.raises(PlanError):
            request(t.id, p)
    row = result('EdgeSwitch', {**edge_inputs(tagged_vlans='746'), 'action': 'Remove'}, EDGE)
    assert row['commands'][2:-2] == ['no vlan tagging 746']
    row = result('EdgeSwitch', {**edge_inputs(tagged_vlans='545'), 'action': 'Remove'}, EDGE)
    assert row['commands'][2:-2] == ['no vlan tagging 545', 'vlan participation exclude 545']
    assert 'PVID_REMOVAL_REQUIRES_REVIEW' in result('EdgeSwitch', {**edge_inputs(tagged_vlans='445'), 'action': 'Remove'}, EDGE)['blockers']
    a = request(t.id, edge_inputs(pvid=901), row=3)
    b = request(t.id, {**edge_inputs(tagged_vlans='445'), 'action': 'Remove'}, row=4)
    rows = plan([b,a], {IP: observation('EdgeSwitch', EDGE)}, now=NOW,
                decisions={a['row_id']: 'Approve', b['row_id']: 'Approve'}).to_dict()['results']
    assert all(r['deployment_ready'] for r in rows)
    assert rows[1]['projected_state']['pvid'] == 901


@pytest.mark.parametrize('family', ['ME3600X', 'EdgeSwitch'])
def test_database_lookup_fallback_create_rename_noop_conflict_unknown(family):
    config = EDGE if family == 'EdgeSwitch' else evc_config(family)
    t = template(family, 'VLAN Database')
    p = dict(management_ip=IP, vlan_id=445, name='New_Name')
    r = request(t.id, p)
    obs = observation(family, config)
    row = plan([r], {IP: obs}, now=NOW).to_dict()['results'][0]
    assert 'VLAN_LOOKUP_FALLBACK' in row['warnings'] and 'VLAN_RENAME_IMPACT' in row['warnings']
    if family == 'EdgeSwitch':
        assert row['commands'] == ['vlan database', 'vlan name 445 "New_Name"', 'exit']
    obs['vlan_lookups'] = {'445': dict(status='absent', name='', operational_status='')}
    assert 'VLAN_OBSERVATION_CONFLICT' in plan([r], {IP: obs}, now=NOW).to_dict()['results'][0]['blockers']
    obs['configuration_reliable'] = False
    obs['vlan_lookups'] = {}
    assert 'VLAN_STATE_UNKNOWN' in plan([r], {IP: obs}, now=NOW).to_dict()['results'][0]['blockers']


@pytest.mark.parametrize('family', ['ME3600X', 'EdgeSwitch'])
def test_reliably_empty_database_is_confirmed_absent(family):
    t = template(family, 'VLAN Database')
    r = request(t.id, dict(management_ip=IP, vlan_id=901, name='Customer'))
    obs = observation(family, 'hostname fixture\nend\n')
    row = plan([r], {IP: obs}, now=NOW, decisions={r['row_id']: 'Approve'}).to_dict()['results'][0]
    assert row['deployment_ready'] and not row['no_op']
    assert 'vlan 901' in row['commands']


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920', 'EdgeSwitch'])
def test_description_only_exact_member_target_and_unchanged_evidence(family):
    t = template(family, 'Interface Description')
    config = EDGE if family == 'EdgeSwitch' else evc_config(family, parent=' channel-group 1 mode active\n description Before\n shutdown\n')
    p = dict(management_ip=IP, interface='0/1' if family == 'EdgeSwitch' else 'GigabitEthernet0/1', description='New Customer Site')
    obs = observation(family, config)
    before = canonical(obs)
    r = request(t.id,p)
    row = plan([r], {IP: obs}, now=NOW, decisions={r['row_id']: 'Approve'}).to_dict()['results'][0]
    assert row['deployment_ready'] and len(row['commands']) == (5 if family == 'EdgeSwitch' else 4)
    assert row['resolved_target'] == p['interface'] and canonical(obs) == before


def test_new_workbook_help_invalid_row_isolation_digest_and_no_network(tmp_path, monkeypatch):
    monkeypatch.setattr(socket, 'socket', lambda *a, **k: pytest.fail('offline workflow connected'))
    source = tmp_path / 'input.xlsx'
    create_input(source)
    book = load_workbook(source)
    for family in ('ME3600X', 'ASR920'):
        sheet = book[template(family).sheet]
        assert '2-444,446-544' in sheet['A1'].comment.text
        assert 'Pop 2' in sheet['A1'].comment.text and 'Invalid' in sheet['A1'].comment.text
    t = template('EdgeSwitch', 'VLAN Port')
    p = edge_inputs(pvid=901)
    for col, field in enumerate(t.columns, 1):
        book[t.sheet].cell(3, col, p.get(field))
        book[t.sheet].cell(4, col, {**p, 'action': 'Remove'}.get(field))
        book[t.sheet].cell(5, col, {**p, 'management_ip': ''}.get(field))
    book.save(source)
    rows = read_input(source)
    assert len(rows) == 1 and len(rows.errors) == 2
    assert any('management_ip' in detail for detail in rows.errors[1]['details'])
    original = plan(rows, {IP: observation('EdgeSwitch', EDGE)}, now=NOW)
    dest = tmp_path / 'plan.xlsx'
    export_plan(original, dest)
    save(original, tmp_path / 'plan.json')
    assert load(tmp_path / 'plan.json').digest == original.digest
    book = load_workbook(dest)
    assert book[t.sheet].cell(4, len(t.columns) + EXTRA.index('Validation Status') + 1).value == 'Invalid'
    book[t.sheet].cell(3, len(t.columns) + len(EXTRA), 'Approve')
    book.save(dest)
    reviewed = consume_decisions(original, dest, expected_digest=original.digest, now=NOW)
    assert reviewed.to_dict()['results'][0]['deployment_ready'] and not reviewed.to_dict()['execution_authorized']
    with pytest.raises(PlanError):
        export_plan(original, dest)


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920', 'EdgeSwitch'])
def test_shared_collection_detail_lookup_failure_isolation_sanitization_and_no_writes(family, tmp_path, monkeypatch):
    from contextlib import contextmanager
    from functools import partial
    from orbitflow.configuration.v2 import observation as collector
    from orbitflow.config import ExecutionConfig
    from orbitflow.transport import DeviceCredentials
    from orbitflow.logging import module_logger
    config = EDGE if family == 'EdgeSwitch' else evc_config(family)
    obs = observation(family, config)
    ident = obs['identity']
    context = DeviceContext(ident['inventory_id'], IP, (IP,), ident['hostname'], 'fixture', ident['platform'], family,
                            family, ident['capability_profile'], (), 'FIXTURE', '1', '', NOW, NOW)
    calls, closed = [], []
    @contextmanager
    def connect(ip, credentials, config):
        try:
            if ip != IP:
                raise ValueError('SYNTHETIC_PASSWORD_MUST_NOT_ESCAPE')
            yield object()
        finally:
            closed.append(ip)
    class CLI:
        def __init__(self, session):
            self.session = session
            self.prompt = '(fixture)#' if family == 'EdgeSwitch' else 'fixture#'
        def __enter__(self):
            return self
        def __exit__(self, *args):
            closed.append('cli')
        def require_session(self, session):
            assert session is self.session
        def configure(self, **options):
            assert options['paging_command'] == 'terminal length 0'
        def run_command(self, command, **kwargs):
            assert command.startswith('show '), 'device write attempted'
            calls.append(command)
            if command == 'show running-config':
                # Known supplied credentials must also disappear from a harmless
                # source description, beyond command-label sanitation.
                return config.replace('Customer Site Alpha', 'SYNTHETIC_PASSWORD_MUST_NOT_ESCAPE')
            if command.startswith('show ethernet'):
                return 'Service instance 3000 interface GigabitEthernet0/1 does not exist'
            if command == 'show vlan 902':
                return 'VLAN does not exist.'
            return '% Invalid input'
    class Resolver:
        def __init__(self, *a, **k):
            pass
        def resolve(self, *a, **k):
            return context
    monkeypatch.setattr(collector, 'connect_device', connect)
    monkeypatch.setattr(collector, 'DeviceCLI', CLI)
    monkeypatch.setattr(collector, 'DeviceInventoryResolver', Resolver)
    monkeypatch.setattr(collector.InterfaceService, 'collect', lambda *a, **k: ())
    monkeypatch.setattr(collector, 'module_logger', partial(module_logger, log_root=tmp_path / 'logs'))
    t = template(family, 'VLAN Port' if family == 'EdgeSwitch' else 'EVC')
    p = edge_inputs(tagged_vlans='902', pvid=901) if family == 'EdgeSwitch' else evc_inputs()
    rows = [request(t.id, p), request(t.id, {**p, 'management_ip': '192.0.2.92'}, row=4)]
    creds = {ip: DeviceCredentials('synthetic-user', 'SYNTHETIC_PASSWORD_MUST_NOT_ESCAPE') for ip in (IP, '192.0.2.92')}
    spool = collector.collect(rows, creds, None, inventory_path=tmp_path / 'inventory.json', spool_root=tmp_path / 'runs', execution_config=ExecutionConfig(1,0,0))
    records = dict(collector.observations(spool))
    assert records[IP]['status'] == 'live_read_only' and records['192.0.2.92']['status'] == 'evidence_unavailable'
    assert calls.count('show running-config') == 1 and closed.count('cli') == 1 and closed.count(IP) == 1
    assert ('show vlan 902' if family == 'EdgeSwitch' else 'show ethernet service instance id 3000 interface GigabitEthernet0/1 detail') in calls
    if family == 'EdgeSwitch':
        assert 'show vlan 901' in calls
    persisted = ''.join(p.read_text(encoding='utf-8') for p in tmp_path.rglob('*') if p.is_file() and p.suffix in {'.json', '.jsonl', '.log'})
    assert 'SYNTHETIC_PASSWORD_MUST_NOT_ESCAPE' not in persisted
    assert 'SYNTHETIC_PASSWORD_MUST_NOT_ESCAPE' not in canonical(records)


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920', 'EdgeSwitch'])
def test_incomplete_running_config_cannot_prove_absence(family):
    from orbitflow.vendors.cisco.vlans import CiscoXEVlanAdapter
    from orbitflow.vendors.ubiquiti.vlans import EdgeSwitchVlanAdapter
    session = object()
    class CLI:
        prompt = 'fixture#'
        def require_session(self, supplied):
            assert supplied is session
        def configure(self, **kwargs):
            pass
        def run_command(self, command, **kwargs):
            assert command == 'show running-config'
            return 'interface 0/1\nvlan pvid 445' if family == 'EdgeSwitch' else 'interface GigabitEthernet0/1\n shutdown\n!'
    adapter = EdgeSwitchVlanAdapter if family == 'EdgeSwitch' else CiscoXEVlanAdapter
    with pytest.raises(ValueError, match='incomplete'):
        adapter(session, cli=CLI(), preserve_full_evidence=True).collect()


def test_earlier_me_database_prerequisite_requires_approval():
    db = template('ME3600X', 'VLAN Database')
    evc = template('ME3600X')
    a = request(db.id, dict(management_ip=IP, vlan_id=445, name='Customer'), sheet_order=1)
    b = request(evc.id, evc_inputs(), sheet_order=2)
    obs = observation('ME3600X', 'interface GigabitEthernet0/1\n!\nend\n')
    for decision in ('Pending', 'Approve', 'Reject'):
        rows = plan([b,a], {IP: obs}, now=NOW, decisions={a['row_id']: decision, b['row_id']: 'Approve'}).to_dict()['results']
        assert rows[1]['dependencies'] == [a['row_id']]
        assert rows[1]['deployment_ready'] == (decision == 'Approve')
        assert rows[1]['consistency'] == 'consistent' and not rows[1]['blockers']
        assert ('DEPENDENCY_NOT_APPROVED' in rows[1]['warnings']) == (decision != 'Approve')


def test_evc_detail_semantic_conflict_down_information_and_source_sanitation():
    family = 'ASR920'
    config = evc_config(family, ' service instance 3000 ethernet\n  encapsulation dot1q 445\n  rewrite ingress tag pop 1 symmetric\n  bridge-domain 445\n')
    text = 'Id: 3000\nInterface: GigabitEthernet0/1\nEncapsulation: dot1q 445\nRewrite: ingress tag pop 1 symmetric\nBridge-domain: 445\nState: Down\npassword SYNTHETIC_SECRET_VALUE\n'
    detail = parse_service_detail(text, 'GigabitEthernet0/1', 3000)
    row = result(family, evc_inputs(), config, service_details={'GigabitEthernet0/1:3000': detail})
    assert row['deployment_ready'] and 'SERVICE_INSTANCE_DOWN' in row['warnings']
    assert 'SYNTHETIC_SECRET_VALUE' not in canonical(row)
    bad = parse_service_detail(text.replace('Bridge-domain: 445', 'Bridge-domain: 700'), 'GigabitEthernet0/1', 3000)
    row = result(family, evc_inputs(), config, service_details={'GigabitEthernet0/1:3000': bad})
    assert 'SERVICE_OBSERVATION_CONFLICT' in row['blockers']


def test_child_omissions_and_separate_parent_description_are_visible():
    p = {**evc_inputs(), 'encapsulation': 'Untagged', 'rewrite': 'None', 'description': 'New Child', 'interface_description': 'New Parent'}
    p.pop('outer_vlan')
    config = evc_config('ASR920', ' service instance 3000 ethernet\n  description Old Child\n  encapsulation untagged\n  rewrite ingress tag push dot1q 1402 symmetric\n  bridge-domain 445 split-horizon group 0\n', ' description Old Parent\n')
    row = result('ASR920', p, config)
    assert row['deployment_ready']
    assert row['commands'].index('description New Parent') < row['commands'].index('service instance 3000 ethernet') < row['commands'].index('description New Child')
    assert not any('rewrite' in c for c in row['commands'])
    assert any('push dot1q' in c for c in row['omitted_child_commands'])
    assert any('split-horizon group 0' in c for c in row['omitted_child_commands'])

@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
@pytest.mark.parametrize('enc,outer,inner,rewrite,reason', [
    ('Dot1q', '445', None, 'None', 'UNVERIFIED'),
    ('QinQ', '300', '200', 'None', 'UNVERIFIED'),
    ('Untagged', None, None, 'Pop 1 symmetric', 'UNVERIFIED'),
    ('Untagged', None, None, 'Pop 2 symmetric', 'UNVERIFIED'),
    ('Dot1q', '445', None, 'Pop 2 symmetric', 'LAB_REJECTED'),
    ('Dot1q', '100-200', None, 'Pop 1 symmetric', 'LAB_REJECTED'),
    ('QinQ', '300', '200', 'Push dot1q', 'LAB_REJECTED'),
    ('QinQ', '100', '1-4094', 'Pop 2 symmetric', 'LAB_REJECTED'),
])
def test_rewrite_rejection_reason_is_explicit(family, enc, outer, inner, rewrite, reason):
    p = {**evc_inputs(), 'encapsulation': enc, 'rewrite': rewrite}
    p.pop('outer_vlan')
    if outer:
        p['outer_vlan'] = outer
    if inner:
        p['inner_vlan'] = inner
    if rewrite.startswith('Push'):
        p['push_vlan'] = 100
    with pytest.raises(PlanError, match=reason + '_REWRITE_COMBINATION: ' + family):
        request(template(family).id, p)


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
@pytest.mark.parametrize('existing', ['absent', 'matching', 'changed', 'unknown', 'conflict'])
def test_evc_recreation_source_and_precise_differences(family, existing):
    from orbitflow.vendors.cisco.evc import child_commands
    p = {**evc_inputs(), 'encapsulation': 'QinQ', 'outer_vlan': '300', 'inner_vlan': '200',
         'rewrite': 'Pop 2 symmetric', 'split_horizon_group': 0}
    desired = child_commands(p)
    child = '' if existing == 'absent' else ' service instance 3000 ethernet\n' + '\n'.join('  ' + c for c in desired) + '\n  !\n'
    if existing == 'changed':
        child = child.replace('  !', '  service-policy input OLD\n  description Old Child\n  !')
    if existing == 'unknown':
        child = child.replace('encapsulation dot1q 300 second-dot1q 200', 'encapsulation unsupported')
    config = evc_config(family, child + ' service instance 20 ethernet\n  encapsulation dot1q 800\n  bridge-domain 800\n  !\n',
                        ' description Parent\n mtu 9000\n shutdown\n')
    extra = {}
    if existing == 'conflict':
        extra['service_details'] = {'GigabitEthernet0/1:3000': parse_service_detail(
            'Service instance 3000 interface GigabitEthernet0/1 does not exist', 'GigabitEthernet0/1', 3000)}
    row = result(family, p, config, **extra)
    assert not row['no_op']
    assert row['commands'][2:4] == ['no service instance 3000', 'service instance 3000 ethernet']
    assert row['commands'][4:-1] == desired
    assert row['deployment_ready'] == (existing not in {'unknown', 'conflict'})
    changes = row['service_instance_changes']
    assert changes['after'] == desired and changes['recreation']
    if existing == 'matching':
        assert changes['removed'] == changes['added'] == row['omitted_child_commands'] == []
        assert 'OMITTED_CHILD_SETTINGS' not in row['warnings']
    if existing == 'changed':
        assert row['omitted_child_commands'] == ['service-policy input OLD', 'description Old Child']
        assert changes['removed'] == row['omitted_child_commands']
    assert ('SERVICE_INSTANCE_RECREATE_IMPACT' in row['warnings']) == (existing != 'absent')
    obs = observation(family, config)
    evidence = scoped_evidence(request(template(family).id, p), obs)
    interface_evidence = [e for e in evidence if e['excerpt'].strip() == '!' or 'interface' in e['excerpt'] or e['excerpt'].startswith(' ')]
    assert len({(e['source_filename'], e['line']) for e in evidence}) == len(evidence)
    assert any(e['excerpt'] == '  !' for e in interface_evidence) == (existing != 'absent')
    node = next(n for n in obs['facts'] if n['kind'] == 'interface')
    service = next(n for n in node['children'] if n['kind'] == 'service_instance' and n['value'] == '20')
    assert service['children'][-1]['value'] == 'stanza_end'
    assert 'service instance 20 ethernet' not in row['commands']
    assert all(c not in row['commands'] for c in ['shutdown', 'mtu 9000', 'description Parent'])


@pytest.mark.parametrize('scenario', ['remove_add', 'pvid_remove', 'repeat', 'unsafe', 'pending', 'reject'])
def test_edge_ordered_workbook_rows_to_plan(tmp_path, scenario):
    source = tmp_path / 'ordered.xlsx'
    create_input(source)
    book = load_workbook(source)
    t = template('EdgeSwitch', 'VLAN Port')
    inputs = {
        'remove_add': [dict(action='Remove', tagged_vlans='545,746'), dict(action='Add/Change', tagged_vlans='545,746', pvid=746)],
        'pvid_remove': [dict(action='Add/Change', pvid=901), dict(action='Remove', tagged_vlans='445')],
        'repeat': [dict(action='Add/Change', tagged_vlans='545', pvid=445)] * 2,
        'unsafe': [dict(action='Add/Change', tagged_vlans='545'), dict(action='Remove', tagged_vlans='445')],
        'pending': [dict(action='Add/Change', pvid=901), dict(action='Remove', tagged_vlans='445')],
        'reject': [dict(action='Add/Change', pvid=901), dict(action='Remove', tagged_vlans='445')],
    }[scenario]
    for number, values in enumerate(inputs, 3):
        p = dict(management_ip=IP, interface='0/1', **values)
        for col, field in enumerate(t.columns, 1):
            book[t.sheet].cell(number, col, p.get(field))
    book.save(source)
    requests = read_input(source)
    assert not requests.errors
    obs = observation('EdgeSwitch', EDGE.replace('mtu 9000', 'poe opmode auto\nmtu 9000'))
    obs['vlan_lookups'] = {'746': parse_vlan_lookup('VLAN ID: 746\nVLAN Name: 746\nVLAN Type: Static\nPort Current Configured Tagging\n0/1 Exclude Include Tagged\n', 746)}
    before = canonical(obs)
    decisions = {r['row_id']: 'Approve' for r in requests}
    if scenario in {'pending', 'reject'}:
        decisions[requests[0]['row_id']] = scenario.title()
    planned = plan(requests, {IP: obs}, now=NOW, decisions=decisions)
    rows = planned.to_dict()['results']
    assert canonical(obs) == before
    assert [r['row'] for r in planned.to_dict()['requests']] == [3, 4]
    if scenario == 'remove_add':
        assert rows[0]['commands'][2:-2] == ['no vlan tagging 545,746', 'vlan participation exclude 545']
        assert rows[1]['commands'][2:-2] == ['vlan participation include 545,746', 'vlan tagging 545,746', 'vlan pvid 746']
        assert rows[1]['projected_state'] == dict(description='Customer Site Alpha', included=[445,545,746], excluded=[], tagged=[545,746], pvid=746)
    elif scenario == 'pvid_remove':
        assert rows[1]['commands'][2:-2] == ['vlan participation exclude 445']
        assert rows[1]['projected_state']['pvid'] == 901
    elif scenario == 'repeat':
        assert all(r['no_op'] and not r['commands'] for r in rows)
    elif scenario == 'pending':
        assert 'DEPENDENCY_NOT_APPROVED' in rows[1]['blockers']
    elif scenario in {'unsafe', 'reject'}:
        assert 'PVID_REMOVAL_REQUIRES_REVIEW' in rows[1]['blockers']
    if scenario in {'remove_add', 'pvid_remove', 'repeat'}:
        assert all(r['deployment_ready'] and 'ORDERED_RESOURCE_OVERLAP_WARNING' in r['warnings'] for r in rows)
    else:
        assert 'ORDERED_RESOURCE_OVERLAP' in rows[1]['blockers']
    assert all(not any('storm-control' in c or 'poe' in c or 'mtu' in c for c in r['commands']) for r in rows)
    for r in requests:
        evidence = '\n'.join(e['excerpt'] for e in scoped_evidence(r, obs))
        assert 'poe opmode auto' in evidence and 'storm-control' in evidence and 'exit' in evidence
    export_plan(planned, tmp_path / 'plan.xlsx')


def test_edge_excluded_tag_lists_ranges_remain_independent():
    config = EDGE.replace('vlan participation exclude 746', 'vlan participation exclude 745-746,2436-2444').replace('vlan tagging 545,746', 'vlan tagging 545,745-746,2436-2444')
    ports, _ = parse_edgeswitch_config(config)
    assert set(ports[0].configured_tagged_vlans) == {545,745,746,*range(2436,2445)}
    assert set(ports[0].excluded_vlans) == {745,746,*range(2436,2445)}
    row = result('EdgeSwitch', edge_inputs(pvid=901), config)
    assert row['projected_state']['tagged'] == list(ports[0].configured_tagged_vlans)
    assert row['projected_state']['excluded'] == list(ports[0].excluded_vlans)
    assert row['commands'][2:-2] == ['vlan participation include 901', 'vlan pvid 901']

@pytest.mark.parametrize('ending', [' !', '  !'])
def test_nested_child_ending_is_preserved_with_child(ending):
    config = evc_config('ASR920', ' service instance 3000 ethernet\n  encapsulation dot1q 445\n  rewrite ingress tag pop 1 symmetric\n  bridge-domain 445 split-horizon group 0\n' + ending + '\n')
    obs = observation('ASR920', config)
    node = next(n for n in obs['facts'] if n['kind'] == 'interface')
    child = next(n for n in node['children'] if n['kind'] == 'service_instance')
    def descendants(n):
        return [n] + [c for sub in n.get('children', []) for c in descendants(sub)]
    assert any(n['excerpt'] == ending and n['value'] == 'stanza_end' for n in descendants(child))
    p = {**evc_inputs(), 'split_horizon_group': 0}
    row = result('ASR920', p, config)
    assert row['omitted_child_commands'] == []


def test_operator_reference_columns_and_workbook_help_agree():
    from pathlib import Path
    from orbitflow.configuration.v2.workbook import _base
    reference = (Path(__file__).resolve().parents[1] / 'docs/operations/configuration-template-reference.md').read_text(encoding='utf-8')
    book = _base()
    for number, t in enumerate(CATALOGUE, 1):
        section = reference.split(f'## Template {number} — {t.sheet}\n', 1)[1].split('\n## ', 1)[0]
        assert all(f'| `{field}` |' in section for field in t.columns)
        assert '| `management_ip` | Required |' in section
        for field in t.required:
            assert f'| `{field}` | Required |' in section
        assert ('Expected' in section or 'complete matrix' in section.lower() or 'scenario table below' in section) and 'evidence' in section.lower()
        help_text = book[t.sheet]['A1'].comment.text
        assert all(field in help_text for field in ('management_ip', *t.required))
        if number >= 17:
            assert 'not owner-approved or live-validated' in help_text
            assert 'catalogue/provisional' in section
    for family in ('ME3600X', 'ASR920'):
        text = book[template(family).sheet]['A1'].comment.text
        assert 'Unverified (blocked): single Dot1q+None; fixed QinQ+None' in text
        assert 'Lab rejected (blocked):' in text
    for value in ('Push dot1q symmetric', 'Pop 2 symmetric', 'Remove', 'Add/Change', 'split_horizon_group', 'push_vlan'):
        assert value in reference
    assert 'Service Instance Differences' not in EXTRA


def test_evc_differences_in_single_digest_bound_workbook(tmp_path):
    t = template('ASR920')
    r = request(t.id, {**evc_inputs(), 'split_horizon_group': 0})
    config = evc_config('ASR920', ' service instance 3000 ethernet\n  encapsulation dot1q 445\n  rewrite ingress tag pop 1 symmetric\n  bridge-domain 445 split-horizon group 0\n  service-policy input OLD\n !\n', ' description Parent\n')
    original = plan([r], {IP: observation('ASR920', config)}, now=NOW)
    path = tmp_path / 'plan.xlsx'
    export_plan(original, path)
    save(original, tmp_path / 'plan.json')
    assert load(tmp_path / 'plan.json').digest == original.digest
    book = load_workbook(path)
    col = len(t.columns) + EXTRA.index('Proposed CLI') + 1
    text = canonical(original.to_dict()['results'][0]['service_instance_changes'])
    assert 'service-policy input OLD' in text and 'description Parent' not in text
    assert 'removed' in text and 'before' in text and 'after' in text
    book[t.sheet].cell(3, len(t.columns) + len(EXTRA), 'Approve')
    book.save(path)
    assert consume_decisions(original, path, expected_digest=original.digest, now=NOW).to_dict()['results'][0]['deployment_ready']
    book[t.sheet].cell(3, col, '{}')
    book.save(path)
    with pytest.raises(PlanError, match='review content changed'):
        consume_decisions(original, path, expected_digest=original.digest, now=NOW)

@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
def test_recreation_differences_are_sanitized_before_persistence(family, tmp_path):
    config = evc_config(family, ' service instance 3000 ethernet\n  encapsulation dot1q 445\n  rewrite ingress tag pop 1 symmetric\n  bridge-domain 445\n  description password SYNTHETIC_CHILD_SECRET\n !\n')
    t = template(family)
    r = request(t.id, evc_inputs())
    original = plan([r], {IP: observation(family, config)}, now=NOW)
    assert 'SYNTHETIC_CHILD_SECRET' not in canonical(original.to_dict())
    assert '[REDACTED]' in canonical(original.to_dict()['results'][0]['service_instance_changes'])
    save(original, tmp_path / 'plan.json')
    export_plan(original, tmp_path / 'plan.xlsx')
    assert 'SYNTHETIC_CHILD_SECRET' not in (tmp_path / 'plan.json').read_text(encoding='utf-8')
    book = load_workbook(tmp_path / 'plan.xlsx')
    assert all('SYNTHETIC_CHILD_SECRET' not in str(cell.value) for sheet in book for row in sheet for cell in row)


@pytest.mark.parametrize('family', ['ME3600X', 'ASR920'])
def test_large_evc_scoping_full_validation_and_bounded_export(family, tmp_path, monkeypatch):
    monkeypatch.setattr(socket, 'socket', lambda *a, **k: pytest.fail('planning connected'))
    children = ''.join(f' service instance {i} ethernet\n  description Neighbor {i}\n  encapsulation dot1q {i}\n  bridge-domain {i}\n  !\n' for i in range(1, 121))
    children += ' service instance 3000 ethernet\n  description Requested\n  encapsulation dot1q 445\n  rewrite ingress tag pop 1 symmetric\n  bridge-domain 445\n  !\n'
    obs = observation(family, evc_config(family, children, ' description Parent\n mtu 9000\n'))
    r = request(template(family).id, evc_inputs())
    original = plan([r], {IP: obs}, now=NOW)
    before = original.to_dict()
    assert len(canonical(before['results'][0]['projected_state'])) > 32767
    selected = scoped_evidence(r, obs)
    assert any(e['excerpt'] == '  !' for e in selected)
    assert any('mtu 9000' in e['excerpt'] for e in selected)
    assert any('Requested' in e['excerpt'] for e in selected)
    assert not any('Neighbor' in e['excerpt'] for e in selected)
    # Exact source fidelity, including line numbers and child/parent endings.
    source = evc_config(family, children, ' description Parent\n mtu 9000\n').splitlines()
    assert all(source[e['line'] - 1] == e['excerpt'] for e in selected)
    neighbor_request = request(template(family).id, {**evc_inputs(), 'outer_vlan': '120'})
    conflict = plan([neighbor_request], {IP: obs}, now=NOW).to_dict()['results'][0]
    assert 'INGRESS_TAG_OVERLAP' in conflict['blockers']
    dest = tmp_path / 'plan.xlsx'
    save(original, tmp_path / 'plan.json')
    export_plan(original, dest)
    book = load_workbook(dest)
    sheet = book[template(family).sheet]
    assert 'Projected State' not in [c.value for c in sheet[2]]
    projected = sheet.cell(3, len(template(family).columns) + EXTRA.index('Proposed CLI') + 1)
    assert all(len(str(c.value or '')) <= 32767 for sh in book for row in sh for c in row)
    assert not any('Neighbor' in str(c.value) for row in book['Evidence Details'] for c in row)
    sheet.cell(3, len(template(family).columns) + len(EXTRA), 'Approve')
    book.save(dest)
    consume_decisions(original, dest, expected_digest=original.digest, now=NOW)
    assert original.to_dict() == before and load(tmp_path / 'plan.json').digest == original.digest
    projected.value += ' altered'
    book.save(dest)
    with pytest.raises(PlanError, match='review content changed'):
        consume_decisions(original, dest, expected_digest=original.digest, now=NOW)


def test_32_rows_invalid_catalyst_names_remain_isolated(tmp_path):
    source = tmp_path / 'input.xlsx'
    create_input(source)
    book = load_workbook(source)
    t = template('C3750X', 'VLAN Database')
    for row in range(3, 35):
        p = dict(management_ip=IP, vlan_id=445 + row, name='Bad Name' if row in (7, 19) else 'Customer')
        for col, field in enumerate(t.columns, 1):
            book[t.sheet].cell(row, col, p.get(field))
    book.save(source)
    requests = read_input(source)
    assert len(requests) == 30 and len(requests.errors) == 2
    original = plan(requests, {}, now=NOW)
    dest = tmp_path / 'plan.xlsx'
    export_plan(original, dest)
    book = load_workbook(dest)
    status_col = len(t.columns) + EXTRA.index('Validation Status') + 1
    assert [book[t.sheet].cell(row, status_col).value for row in (7, 19)] == ['Invalid', 'Invalid']
    assert book['Summary'].max_row == 33
    consume_decisions(original, dest, expected_digest=original.digest, now=NOW)


@pytest.mark.parametrize('length', [32767, 32768, 100000])
def test_excel_literal_limit_and_deterministic_reference(length):
    from openpyxl import Workbook
    from orbitflow.configuration.v2.workbook import literal
    import hashlib
    book = Workbook()
    book.active.title = 'Introduction'
    book.active['B1'] = 'a' * 64
    sheet = book.create_sheet('Evidence Details')
    value = 'x' * length
    literal(sheet['D2'], value)
    if length == 32767:
        assert sheet['D2'].value == value
    else:
        assert len(sheet['D2'].value) <= 32767
        assert hashlib.sha256(value.encode()).hexdigest() in sheet['D2'].value
        assert str(length) in sheet['D2'].value
        literal(sheet['D3'], value)
        assert sheet['D2'].value.replace('D2', 'D3') == sheet['D3'].value

@pytest.mark.parametrize('description_first', [False, True])
@pytest.mark.parametrize('decision', ['Pending', 'Approve', 'Reject'])
def test_edgeswitch_port_description_mixed_sheet_order(tmp_path, description_first, decision):
    port = template('EdgeSwitch', 'VLAN Port')
    desc = template('EdgeSwitch', 'Interface Description')
    source = tmp_path / 'input.xlsx'
    create_input(source)
    book = load_workbook(source)
    if description_first:
        book.move_sheet(desc.sheet, offset=-1)
    values = [(port, dict(management_ip=IP, interface='0/1', action='Add/Change', tagged_vlans='902', pvid=902, description='Port description')),
              (desc, dict(management_ip=IP, interface='0/1', description='New description'))]
    for t, p in values:
        for col, field in enumerate(t.columns, 1):
            book[t.sheet].cell(3, col, p.get(field))
    book.save(source)
    requests = read_input(source)
    obs = observation('EdgeSwitch', 'vlan database\n vlan 901\n exit\ninterface 0/1\n description "Original"\n vlan pvid 901\n vlan participation include 901\n vlan tagging 901\nexit\n')
    original = plan(requests, {IP: obs}, now=NOW)
    for row in original.to_dict()['results']:
        assert not row['blockers'] and row['consistency'] == 'consistent'
        assert 'ORDERED_RESOURCE_OVERLAP_WARNING' in row['warnings']
        assert not row['dependencies']
    expected_description = 'Port description' if description_first else 'New description'
    assert original.to_dict()['results'][-1]['projected_state']['description'] == expected_description
    path = tmp_path / 'plan.xlsx'
    export_plan(original, path)
    book = load_workbook(path)
    for t, _ in values:
        book[t.sheet].cell(3, len(t.columns) + len(EXTRA), decision)
    book.save(path)
    evaluated = consume_decisions(original, path, expected_digest=original.digest, now=NOW)
    assert all(r['decision'] == decision and not r['blockers'] for r in evaluated.to_dict()['results'])
    assert all(r['deployment_ready'] == (decision == 'Approve') for r in evaluated.to_dict()['results'])
    assert original.to_dict()['results'][0]['decision'] == 'Pending'


def test_edgeswitch_description_does_not_waive_pvid_blocker():
    port = request(template('EdgeSwitch', 'VLAN Port').id,
                   dict(management_ip=IP, interface='0/1', action='Remove', tagged_vlans='901'), sheet_order=1)
    desc = request(template('EdgeSwitch', 'Interface Description').id,
                   dict(management_ip=IP, interface='0/1', description='New'), sheet_order=2)
    obs = observation('EdgeSwitch', 'interface 0/1\n vlan pvid 901\n vlan participation include 901\n vlan tagging 901\nexit\n')
    rows = plan([port, desc], {IP: obs}, now=NOW).to_dict()['results']
    assert 'PVID_REMOVAL_REQUIRES_REVIEW' in rows[0]['blockers']
    assert all('ORDERED_RESOURCE_OVERLAP_WARNING' in row['warnings'] for row in rows)
    assert not rows[1]['blockers']


def test_edgeswitch_port_evidence_excludes_database_but_retains_warnings(tmp_path):
    t = template('EdgeSwitch', 'VLAN Port')
    r = request(t.id, dict(management_ip=IP, interface='0/1', action='Add/Change', tagged_vlans='902', pvid=902))
    obs = observation('EdgeSwitch', 'vlan database\n vlan 901\n exit\ninterface 0/1\n vlan pvid 901\n vlan participation include 901\n vlan tagging 901\nexit\n')
    obs['vlan_lookups'] = {'902': dict(status='absent', name='', operational_status='', evidence=[
        dict(source_filename='show vlan 902', line=1, excerpt='VLAN does not exist.')])}
    original = plan([r], {IP: obs}, now=NOW)
    assert 'VLAN_NOT_IN_DATABASE' in original.to_dict()['results'][0]['warnings']
    assert original.to_dict()['observations'][IP]['vlan_lookups'] == obs['vlan_lookups']
    path = tmp_path / 'plan.xlsx'
    export_plan(original, path)
    book = load_workbook(path)
    text = book[t.sheet].cell(3, len(t.columns) + EXTRA.index('Existing Evidence') + 1).value
    assert all(s in text for s in ('vlan pvid 901', 'vlan participation include 901', 'vlan tagging 901'))
    assert 'VLAN does not exist' not in text and 'vlan database' not in text
    assert not any('show vlan' in str(c.value) for row in book['Evidence Details'] for c in row)


@pytest.mark.parametrize('lookup', ['absent', 'error', 'unknown', 'unavailable'])
def test_asr920_lookup_evidence_distinguishes_absence(tmp_path, lookup):
    t = template('ASR920')
    r = request(t.id, evc_inputs())
    text = {'absent': 'Service instance 3000 interface GigabitEthernet0/1 does not exist',
            'error': '% Invalid input detected; service instance 3000 interface GigabitEthernet0/1 does not exist',
            'unknown': 'Service instance 20 interface GigabitEthernet0/1 does not exist',
            'unavailable': ''}[lookup]
    detail = parse_service_detail(text, 'GigabitEthernet0/1', 3000) if lookup != 'unavailable' else dict(status='error', facts=[], evidence=[], operational_status='')
    assert detail['status'] == ('error' if lookup == 'unavailable' else lookup)
    obs = observation('ASR920', evc_config('ASR920', parent=' description Parent\n mtu 9000\n'))
    obs['service_details'] = {'GigabitEthernet0/1:3000': detail}
    original = plan([r], {IP: obs}, now=NOW)
    path = tmp_path / 'plan.xlsx'
    export_plan(original, path)
    book = load_workbook(path)
    evidence = book[t.sheet].cell(3, len(t.columns) + EXTRA.index('Existing Evidence') + 1).value
    assert 'description Parent' in evidence and 'mtu 9000' in evidence
    if lookup == 'absent':
        assert 'Interface GigabitEthernet0/1: service instance 3000 does not exist (reliable lookup)' in evidence
    else:
        assert 'lookup unavailable; absence not established' in evidence
        assert 'does not exist (reliable lookup)' not in evidence
        assert 'SERVICE_DETAIL_FALLBACK' in original.to_dict()['results'][0]['warnings']


def test_me_missing_global_vlan_warns_without_creating_and_keeps_other_blockers():
    config = evc_config('ME3600X').replace('vlan 445\n name Customer\n', '')
    # Remove all global declarations regardless of the fixture name.
    config = config[config.index('interface '):]
    row = result('ME3600X', evc_inputs(), config)
    assert 'GLOBAL_VLAN_REQUIRED' in row['warnings'] and not row['blockers']
    assert row['deployment_ready'] and row['missing_global_vlans'] == [445]
    assert 'vlan 445' not in row['commands']
    bad = result('ME3600X', evc_inputs(), config.replace('switchport mode trunk', 'switchport mode access'))
    assert 'INCOMPATIBLE_EVC_PARENT' in bad['blockers']
    assert 'GLOBAL_VLAN_REQUIRED' in bad['warnings']


def test_all_output_columns_and_hidden_identity_tamper(tmp_path):
    from orbitflow.configuration.v2.workbook import FULL_EXTRA, OMITTED, output_columns
    t = template('EdgeSwitch', 'Interface Description')
    requests = [request(t.id, dict(management_ip=IP, interface='0/1', description=value), row=row)
                for row, value in [(3, 'New'), (4, 'New')]]
    original = plan(requests, {IP: observation('EdgeSwitch', 'interface 0/1\n description "Old"\nexit\n')}, now=NOW)
    path = tmp_path / 'plan.xlsx'
    export_plan(original, path)
    book = load_workbook(path)
    for index, item in enumerate(CATALOGUE):
        headers = tuple(c.value for c in book[item.sheet][2])
        assert headers == item.columns + (EXTRA if index < 16 else FULL_EXTRA)
        if index < 16:
            assert not OMITTED.intersection(headers)
    col = len(t.columns) + len(output_columns(t))
    book[t.sheet].cell(3, col, 'Approve')
    book[t.sheet].cell(4, col, 'Reject')
    book.save(path)
    rows = consume_decisions(original, path, expected_digest=original.digest, now=NOW).to_dict()['results']
    assert [(r['row_id'], r['decision']) for r in rows] == [(requests[0]['row_id'], 'Approve'), (requests[1]['row_id'], 'Reject')]
    # Moving input rows changes protected content even without a visible Row ID.
    book[t.sheet].cell(3, 1, '192.0.2.92')
    book.save(path)
    with pytest.raises(PlanError, match='review content changed'):
        consume_decisions(original, path, expected_digest=original.digest, now=NOW)

@pytest.mark.parametrize('field', ['Existing Evidence', 'Proposed CLI', 'Live Validation', 'Consistency',
                                   'Dependencies', 'No-op', 'Validation Status', 'Error Details', 'Warnings', 'Blockers'])
def test_simplified_workbook_protected_result_cells(tmp_path, field):
    t = template('ASR920')
    r = request(t.id, evc_inputs())
    original = plan([r], {IP: observation('ASR920', evc_config('ASR920'))}, now=NOW)
    path = tmp_path / 'plan.xlsx'
    export_plan(original, path)
    book = load_workbook(path)
    book[t.sheet].cell(3, len(t.columns) + EXTRA.index(field) + 1, 'altered')
    book.save(path)
    with pytest.raises(PlanError, match='review content changed'):
        consume_decisions(original, path, expected_digest=original.digest, now=NOW)


@pytest.mark.parametrize('field', ['overlaps', 'projected_state', 'service_instance_changes', 'resolved_target', 'findings'])
def test_removed_workbook_detail_remains_json_protected(tmp_path, field):
    import json
    t = template('ASR920')
    r = request(t.id, evc_inputs())
    original = plan([r], {IP: observation('ASR920', evc_config('ASR920'))}, now=NOW)
    path = tmp_path / 'plan.json'
    save(original, path)
    payload = json.loads(path.read_text())
    payload['plan']['results'][0][field] = 'altered'
    path.write_text(json.dumps(payload))
    with pytest.raises(PlanError):
        load(path)


def test_me_vlan_dependency_chain_preserves_downstream_gate():
    db = template('ME3600X', 'VLAN Database')
    evc = template('ME3600X')
    a = request(db.id, dict(management_ip=IP, vlan_id=445, name='Customer'), sheet_order=1)
    b = request(evc.id, evc_inputs(), sheet_order=2, row=3)
    c = request(evc.id, {**evc_inputs(), 'service_instance_id':3001, 'outer_vlan':'446'}, sheet_order=2, row=4)
    obs = observation('ME3600X', 'interface GigabitEthernet0/1\n!\nend\n')
    rows = plan([a,b,c], {IP:obs}, now=NOW, decisions={b['row_id']:'Approve', c['row_id']:'Approve'}).to_dict()['results']
    assert rows[1]['dependencies'] == [a['row_id']]
    assert 'DEPENDENCY_NOT_APPROVED' in rows[1]['warnings'] and not rows[1]['blockers']
    assert b['row_id'] in rows[2]['dependencies'] and a['row_id'] in rows[2]['dependencies']
    assert 'DEPENDENCY_NOT_APPROVED' in rows[2]['blockers']
    assert all(not r['deployment_ready'] and not r['execution_authorized'] for r in rows)


def test_me_owner_independent_eleven_row_snapshot():
    config = ('vlan 2200-2201\n!\ninterface TenGigabitEthernet0/1\n'
              ' switchport mode trunk\n switchport trunk allowed vlan none\n'
              ' service instance 10 ethernet\n  encapsulation dot1q 1300\n  bridge-domain 2200\n'
              '!\ninterface GigabitEthernet0/22\n switchport mode trunk\n!\nend\n')
    obs = observation('ME3600X', config)
    obs['vlan_lookups'] = {'2200': dict(status='present', name='VLAN2200', operational_status='active')}
    specs = [('Dot1q', '1300', ''), ('Dot1q', '1234', ''), ('Dot1q', '100', ''),
             ('Dot1q', '101-200', ''), ('QinQ', '2433', '100'), ('QinQ', '2434', '200'),
             ('QinQ', '2438', '100-200'), ('QinQ', '1000', '100'), ('QinQ', '1000', '101'),
             ('Untagged', '', '')]
    requests = []
    for i, (enc, outer, inner) in enumerate(specs):
        p = dict(management_ip=IP, interface='TenGigabitEthernet0/1', service_instance_id=10+i,
                 encapsulation=enc, rewrite='Push dot1q symmetric', push_vlan=3300, bridge_domain=2200)
        if outer:
            p['outer_vlan'] = outer
        if inner:
            p['inner_vlan'] = inner
        requests.append(request(template('ME3600X').id, p, row=i+3))
    requests.append(request(template('ME3600X').id, {**requests[0]['inputs'], 'interface':'GigabitEthernet0/22'}, row=13))
    rows = plan(requests, {IP:obs}, now=NOW).to_dict()['results']
    for row in rows[:10]:
        assert row['consistency'] == 'consistent' and not row['blockers'], row['blockers']
        assert row['dependencies'] == []
        assert {'SERVICE_DETAIL_FALLBACK', 'SERVICE_IMPACT_REVIEW'} <= set(row['warnings'])
        assert not row['deployment_ready'] and not row['execution_authorized']
    assert 'SERVICE_INSTANCE_RECREATE_IMPACT' in rows[0]['warnings']
    assert rows[10]['blockers'] == ['INCOMPATIBLE_EVC_PARENT']
    # Each independent service can be reviewed without approving its siblings.
    approved = plan(requests, {IP:obs}, now=NOW,
                    decisions={requests[9]['row_id']:'Approve'}).to_dict()['results']
    assert approved[9]['deployment_ready']


@pytest.mark.parametrize('change', [dict(service_instance_id=3000, outer_vlan='446'),
                                  dict(service_instance_id=3001, outer_vlan='445'),
                                  dict(service_instance_id=3001, outer_vlan='440-450', rewrite='Push dot1q symmetric', push_vlan=3300)])
def test_me_real_same_parent_conflicts_remain(change):
    a = request(template('ME3600X').id, evc_inputs(), row=3)
    b = request(template('ME3600X').id, {**evc_inputs(), **change}, row=4)
    rows = plan([a,b], {IP:observation('ME3600X', evc_config('ME3600X'))}, now=NOW).to_dict()['results']
    assert all('ORDERED_RESOURCE_OVERLAP' in row['blockers'] for row in rows)


@pytest.mark.parametrize('lookup', [dict(status='present', name='Different'), dict(status='absent')])
def test_me_authoritative_vlan_contradictions_remain(lookup):
    row = result('ME3600X', evc_inputs(), evc_config('ME3600X'), vlan_lookups={'445':{**dict(name='', operational_status=''), **lookup}})
    assert 'VLAN_OBSERVATION_CONFLICT' in row['blockers']


@pytest.mark.parametrize('blocked_first', [True, False])
def test_me_independent_blocker_does_not_spread_to_sibling(blocked_first):
    a = request(template('ME3600X').id, evc_inputs(), row=3)
    b = request(template('ME3600X').id, {**evc_inputs(), 'service_instance_id':3001, 'outer_vlan':'446'}, row=4)
    obs = observation('ME3600X', evc_config('ME3600X'))
    obs['service_details'] = {('GigabitEthernet0/1:3000' if blocked_first else 'GigabitEthernet0/1:3001'):dict(status='present', facts=[], evidence=[], operational_status='unknown')}
    rows = plan([a,b], {IP:obs}, now=NOW, decisions={a['row_id']:'Approve', b['row_id']:'Approve'}).to_dict()['results']
    blocked, independent = rows if blocked_first else rows[::-1]
    assert blocked['blockers'] and 'ORDERED_RESOURCE_OVERLAP' not in blocked['blockers']
    assert not independent['blockers'] and not independent['dependencies'] and independent['deployment_ready']


def test_me_explicit_numeric_vlan_name_is_authoritative():
    config = evc_config('ME3600X').replace('name Customer', 'name 445')
    row = result('ME3600X', evc_inputs(), config,
                 vlan_lookups={'445':dict(status='present', name='VLAN0445', operational_status='active')})
    assert 'VLAN_OBSERVATION_CONFLICT' in row['blockers']


@pytest.mark.parametrize('inner,blocked', [('150-250', True), ('201-300', False)])
def test_me_ordered_qinq_ranges_use_tag_pairs(inner, blocked):
    p = {**evc_inputs(), 'encapsulation':'QinQ', 'outer_vlan':'1000',
         'inner_vlan':'100-200', 'rewrite':'Pop 1 symmetric'}
    a = request(template('ME3600X').id, p, row=3)
    b = request(template('ME3600X').id, {**p, 'service_instance_id':3001, 'inner_vlan':inner}, row=4)
    rows = plan([a,b], {IP:observation('ME3600X', evc_config('ME3600X'))}, now=NOW,
                decisions={b['row_id']:'Approve'}).to_dict()['results']
    assert all(('ORDERED_RESOURCE_OVERLAP' in row['blockers']) == blocked for row in rows)
    if not blocked:
        assert not rows[1]['dependencies'] and rows[1]['deployment_ready']
