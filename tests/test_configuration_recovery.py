"""Offline recovery provenance and fail-closed execution; synthetic fixtures only."""
from datetime import timedelta
import json
import socket

import pytest

from test_bulk_configuration_v2 import NOW, IP, template, inputs, observed
from orbitflow.configuration.plans import PlanError, canonical
from orbitflow.configuration.v2.workbook import _base, read_input, export_plan, consume_decisions
from orbitflow.configuration.v2.observation import request_digest
from orbitflow.configuration.v2.planning import load, evaluate_decisions, BulkPlan
from orbitflow.configuration.v2.recovery import recover
from orbitflow.configuration.execution import ExecutionService
from orbitflow.execution import DeviceOutcome
from orbitflow.result_spool import ResultSpool


def fixture(tmp_path, *, legacy=False, age=0, failed=False):
    workbook = tmp_path / 'input.xlsx'
    t = template('C3750X', 'VLAN Database')
    book = _base()
    p = inputs(t)
    book[t.sheet].append([p.get(k, '') for k in t.columns])
    book.save(workbook)
    requests = read_input(workbook)
    spool = ResultSpool.create(tmp_path / 'spools', 'configuration_observation_v2', 1)
    if not legacy:
        spool.bind(dict(request_digest=request_digest(requests), workbook_digest=requests.source_digest))
    obs = observed(t)
    obs['observed_at'] = (NOW - timedelta(seconds=age)).isoformat()
    with spool.collection():
        spool.append(DeviceOutcome(1, {'status': 'evidence_unavailable'} if failed else obs),
                     target=IP, failed=failed)
    return workbook, spool, requests


def test_bound_offline_recovery_no_network_idempotent_and_immutable(tmp_path, monkeypatch):
    workbook, spool, requests = fixture(tmp_path)
    before = {name: (spool.path / name).read_bytes() for name in ('manifest.json', 'results.jsonl')}
    monkeypatch.setattr(socket, 'create_connection', lambda *a, **k: pytest.fail('network forbidden'))
    first = recover(workbook, spool.path, tmp_path / 'first.json', now=NOW)
    second = recover(workbook, spool.path, tmp_path / 'second.json', now=NOW)
    assert first.digest == second.digest == load(tmp_path / 'first.json').digest
    assert first.to_dict()['requests'] == requests
    assert first.to_dict()['source'] == 'workbook-sha256:' + requests.source_digest
    assert (tmp_path / 'first.xlsx').is_file()
    assert before == {name: (spool.path / name).read_bytes() for name in before}
    with pytest.raises(PlanError, match='already exists'):
        recover(workbook, spool.path, tmp_path / 'first.json', now=NOW)
    assert not list(spool.path.glob('read-*'))


@pytest.mark.parametrize('age', [3601, -1])
def test_bound_recovery_rejects_stale_or_future_evidence(tmp_path, age):
    workbook, spool, _ = fixture(tmp_path, age=age)
    with pytest.raises(PlanError, match='stale'):
        recover(workbook, spool.path, tmp_path / 'plan.json', now=NOW)
    assert not (tmp_path / 'plan.json').exists()


def test_changed_workbook_rejected_even_when_requests_unchanged(tmp_path):
    workbook, spool, _ = fixture(tmp_path)
    from openpyxl import load_workbook
    book = load_workbook(workbook)
    book['Introduction']['A1'] = 'Changed introduction'
    book.save(workbook)
    with pytest.raises(PlanError, match='binding mismatch'):
        recover(workbook, spool.path, tmp_path / 'plan.json', now=NOW,
                allow_unverified_legacy=True)


def test_legacy_diagnostic_retains_age_and_cannot_authorize(tmp_path):
    workbook, spool, requests = fixture(tmp_path, legacy=True, age=7200)
    with pytest.raises(PlanError, match='explicit diagnostic'):
        recover(workbook, spool.path, tmp_path / 'plan.json', now=NOW)
    result = recover(workbook, spool.path, tmp_path / 'plan.json', now=NOW,
                     allow_unverified_legacy=True)
    decisions = {requests[0]['row_id']: 'Approve'}
    reviewed = evaluate_decisions(result, decisions, expected_digest=result.digest, now=NOW).to_dict()
    row = reviewed['results'][0]
    assert 'STALE_EVIDENCE' in row['findings']
    assert 'UNVERIFIED_HISTORICAL_PROVENANCE' in row['findings']
    assert not row['deployment_ready'] and not row['execution_authorized']
    assert reviewed['source_digest'] == result.digest
    with pytest.raises(PlanError, match='cannot authorize'):
        ExecutionService(tmp_path / 'authority').prepare(result, decisions, expected_digest=result.digest)
    data = result.to_dict()
    data['results'][0]['findings'].remove('UNVERIFIED_HISTORICAL_PROVENANCE')
    with pytest.raises(PlanError, match='tampered'):
        BulkPlan(canonical(data))


@pytest.mark.parametrize('field,value', [('task', 'other'), ('collection_complete', False)])
def test_wrong_task_or_incomplete_rejected(tmp_path, field, value):
    workbook, spool, _ = fixture(tmp_path)
    manifest = spool.path / 'manifest.json'
    data = json.loads(manifest.read_text())
    data[field] = value
    manifest.write_text(json.dumps(data))
    with pytest.raises(PlanError):
        recover(workbook, spool.path, tmp_path / 'plan.json', now=NOW)


def test_failed_target_remains_failed_evidence(tmp_path):
    workbook, spool, _ = fixture(tmp_path, failed=True)
    result = recover(workbook, spool.path, tmp_path / 'plan.json', now=NOW)
    assert 'EVIDENCE_UNAVAILABLE' in result.to_dict()['results'][0]['findings']


def test_export_failure_publishes_neither_artifact(tmp_path, monkeypatch):
    workbook, spool, _ = fixture(tmp_path)
    import orbitflow.configuration.v2.recovery as module
    def fail(*args):
        raise OSError('synthetic disk full')
    monkeypatch.setattr(module, 'export_plan', fail)
    with pytest.raises(OSError):
        recover(workbook, spool.path, tmp_path / 'plan.json', now=NOW)
    assert not (tmp_path / 'plan.json').exists()
    assert not (tmp_path / 'plan.xlsx').exists()
    assert not list(tmp_path.glob('.recovery-*'))
    assert ResultSpool(spool.path).manifest['status'] == 'collected'


def test_legacy_workbook_review_keeps_provenance_blocker(tmp_path):
    workbook, spool, _ = fixture(tmp_path, legacy=True)
    original = recover(workbook, spool.path, tmp_path / 'plan.json', now=NOW,
                       allow_unverified_legacy=True)
    from openpyxl import load_workbook
    from orbitflow.configuration.v2.catalogue import BY_ID
    from orbitflow.configuration.v2.workbook import output_columns
    review_path = tmp_path / 'plan.xlsx'
    book = load_workbook(review_path)
    r = original.to_dict()['requests'][0]
    t = BY_ID[r['template_id']]
    book[t.sheet].cell(r['row'], len(t.columns) + len(output_columns(t)), 'Approve')
    book.save(review_path)
    result = consume_decisions(original, review_path, expected_digest=original.digest, now=NOW)
    assert result.to_dict()['results'][0]['decision'] == 'Approve'
    assert not result.to_dict()['results'][0]['deployment_ready']


@pytest.mark.parametrize('mutation', ['target', 'identity', 'outcome', 'request_binding'])
def test_recovery_rejects_inconsistent_metadata(tmp_path, mutation):
    workbook, spool, _ = fixture(tmp_path)
    if mutation == 'request_binding':
        path = spool.path / 'manifest.json'
        data = json.loads(path.read_text())
        data['binding']['request_digest'] = '0' * 64
    else:
        path = spool.path / 'results.jsonl'
        data = json.loads(path.read_text())
        if mutation == 'target':
            data['target'] = '192.0.2.99'
        elif mutation == 'identity':
            data['payload']['identity']['management_ip'] = '192.0.2.99'
        else:
            data['status'] = 'failed'
    path.write_text(json.dumps(data) + '\n')
    with pytest.raises(PlanError):
        recover(workbook, spool.path, tmp_path / 'plan.json', now=NOW)
    assert not (tmp_path / 'plan.json').exists()


def test_second_publication_failure_removes_own_json_only(tmp_path, monkeypatch):
    workbook, spool, _ = fixture(tmp_path)
    import orbitflow.configuration.v2.recovery as module
    link = module.os.link
    def collide(source, target):
        if str(target).endswith('.xlsx'):
            target.write_bytes(b'other operator')
        return link(source, target)
    monkeypatch.setattr(module.os, 'link', collide)
    with pytest.raises(FileExistsError):
        recover(workbook, spool.path, tmp_path / 'plan.json', now=NOW)
    assert not (tmp_path / 'plan.json').exists()
    assert (tmp_path / 'plan.xlsx').read_bytes() == b'other operator'


def test_review_opens_both_workbooks_in_streaming_mode(tmp_path, monkeypatch):
    workbook, spool, _ = fixture(tmp_path)
    original = recover(workbook, spool.path, tmp_path / 'plan.json', now=NOW)
    import orbitflow.configuration.v2.workbook as adapter
    loader = adapter.load_workbook
    calls = []
    def load_streaming(*args, **kwargs):
        calls.append(kwargs.get('read_only'))
        assert kwargs.get('read_only') is True
        return loader(*args, **kwargs)
    monkeypatch.setattr(adapter, 'load_workbook', load_streaming)
    consume_decisions(original, tmp_path / 'plan.xlsx', expected_digest=original.digest, now=NOW)
    assert calls == [True, True]


def test_review_rejects_hidden_cells_beyond_forged_dimension(tmp_path):
    workbook, spool, _ = fixture(tmp_path)
    original = recover(workbook, spool.path, tmp_path / 'plan.json', now=NOW)
    from zipfile import ZipFile, ZIP_DEFLATED
    import re
    review = tmp_path / 'plan.xlsx'
    with ZipFile(review) as archive:
        contents = {name: archive.read(name) for name in archive.namelist()}
    # Keep the original dimension hint, append a real cell beyond that hint.
    name = 'xl/worksheets/sheet1.xml'
    xml = contents[name]
    match = re.search(rb'<dimension ref="([^"]+)"', xml)
    if match is None:
        # Streaming writers omit the optional hint. Supply a lying hint so
        # this still verifies that review scans actual cells beyond it.
        xml = xml.replace(b'<sheetData>', b'<dimension ref="A1:D29"/><sheetData>')
    contents[name] = xml.replace(b'</sheetData>',
        b'<row r="5001"><c r="A5001" t="inlineStr"><is><t>tampered</t></is></c></row></sheetData>')
    with ZipFile(review, 'w', ZIP_DEFLATED) as archive:
        for entry, content in contents.items():
            archive.writestr(entry, content)
    with pytest.raises(PlanError, match='dimensions changed'):
        consume_decisions(original, review, expected_digest=original.digest, now=NOW)


def test_orphan_publication_detected_and_fresh_destination_retry(tmp_path, monkeypatch):
    workbook, spool, _ = fixture(tmp_path)
    monkeypatch.setattr(socket, 'create_connection', lambda *a, **k: pytest.fail('network forbidden'))
    original = recover(workbook, spool.path, tmp_path / 'first.json', now=NOW)
    # The durable state left by process death between the two link operations.
    (tmp_path / 'first.xlsx').unlink()
    before = (tmp_path / 'first.json').read_bytes()
    with pytest.raises(PlanError, match='output already exists'):
        recover(workbook, spool.path, tmp_path / 'first.json', now=NOW)
    retry = recover(workbook, spool.path, tmp_path / 'retry.json', now=NOW)
    assert retry.digest == original.digest
    assert (tmp_path / 'first.json').read_bytes() == before
    assert load(tmp_path / 'retry.json').digest == original.digest
    assert (tmp_path / 'retry.xlsx').is_file()


@pytest.mark.parametrize('mutation', ['digest', 'results'])
def test_failed_plan_load_closes_scratch_with_traceback_retained(tmp_path, monkeypatch, mutation):
    import orbitflow.indexed_json as module
    from tempfile import TemporaryDirectory
    workbook, spool, _ = fixture(tmp_path)
    recover(workbook, spool.path, tmp_path / 'first.json', now=NOW)
    data = json.loads((tmp_path / 'first.json').read_text())
    if mutation == 'digest':
        data['digest'] = '0' * 64
    else:
        data['plan']['results'][0]['commands'] = ['invalid synthetic command']
    (tmp_path / 'bad.json').write_text(json.dumps(data))
    scratch = tmp_path / 'scratch'
    scratch.mkdir()
    monkeypatch.setattr(module, 'TemporaryDirectory', lambda **kw: TemporaryDirectory(dir=scratch, **kw))
    with pytest.raises(PlanError) as retained:
        load(tmp_path / 'bad.json')
    assert retained.traceback is not None
    assert list(scratch.iterdir()) == []


def test_failed_export_closes_scratch_with_traceback_retained(tmp_path, monkeypatch):
    import orbitflow.indexed_json as module
    import orbitflow.configuration.v2.recovery as recovery_module
    from tempfile import TemporaryDirectory
    workbook, spool, _ = fixture(tmp_path)
    scratch = tmp_path / 'scratch'
    scratch.mkdir()
    monkeypatch.setattr(module, 'TemporaryDirectory', lambda **kw: TemporaryDirectory(dir=scratch, **kw))
    def failed_export(*args):
        raise OSError('synthetic disk full')
    monkeypatch.setattr(recovery_module, 'export_plan', failed_export)
    with pytest.raises(OSError) as retained:
        recover(workbook, spool.path, tmp_path / 'first.json', now=NOW)
    assert retained.traceback is not None
    assert list(scratch.iterdir()) == []
    assert not (tmp_path / 'first.json').exists()
    assert not (tmp_path / 'first.xlsx').exists()
