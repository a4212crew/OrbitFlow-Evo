"""Owner-supplied EFP response shapes, exercised without device access."""
from types import SimpleNamespace
from contextlib import nullcontext

import pytest
from openpyxl import load_workbook

from orbitflow.capabilities import VlanService
from orbitflow.configuration.v2.planning import request, plan, save, load
from orbitflow.configuration.v2.workbook import export_plan, consume_decisions, EXTRA
from orbitflow.configuration.plans import PlanError
from orbitflow.vendors.cisco.vlans import parse_service_detail
from test_configuration_v2_approved_services import IP, NOW, observation, template


# Family, requested abbreviation, canonical hardware port, ID, operational
# encapsulation, rewrite, BD, state, description.
SAMPLES = [
    ('ASR920', 'g0/0/6', 'GigabitEthernet0/0/6', 100, '100',
     'pop 1 symmetric', 100, 'Up', ''),
    ('ASR920', 'g0/0/12', 'GigabitEthernet0/0/12', 446,
     '2-444,446-544,546-744,747-1000',
     'push dot1q 446 vlan-type 0x8100 symmetric', 1000, 'Up', 'NEOLOOP'),
    ('ME3600X', 'g0/16', 'GigabitEthernet0/16', 600,
     '600 vlan protocol type 0x8100 second-dot1q 10-20',
     'pop 1 symmetric', None, 'Down', ''),
    ('ME3600X', 'Gi0/23', 'GigabitEthernet0/23', 100,
     '100 vlan protocol type 0x8100 second-dot1q 100-444,446-544,546-836',
     'pop 1 symmetric', 100, 'Up', 'WAP-WIFI-APT'),
]


def response(sample):
    _, _, port, sid, enc, rewrite, bd, state, description = sample
    return (f'Service Instance ID: {sid}\nAssociated Interface: {port}\n'
            + (f'Description: {description}\n' if description else '')
            + f'Encapsulation: dot1q {enc} vlan protocol type 0x8100\n'
            + f'Rewrite: ingress tag {rewrite}\nState: {state}\n'
            + (f'Microblock type: Bridge-domain\nBridge-domain: {bd}\n' if bd else '')
            + 'Microblock type: L2Mcast\nL2Mcast metadata: 123\n'
            + 'Microblock type: dhcp_snoop\nMicroblock type: PPPoE IA UBLOCK\n')


@pytest.mark.parametrize('sample', SAMPLES)
def test_hardware_present_identity_ranges_and_rewrite(sample):
    _, short, port, sid, enc, rewrite, bd, state, description = sample
    text = response(sample)
    detail = parse_service_detail(text, short, sid)
    expanded = parse_service_detail(text, port, sid)
    assert all(detail[k] == expanded[k] for k in ('status', 'facts', 'operational_status'))
    assert detail['status'] == 'present'
    assert detail['operational_status'] == state.lower()
    facts = detail['facts']
    outer = next(n['value'] for n in facts if n['kind'] == 'encapsulation')
    assert outer[0] == int(enc.split()[0].split('-')[0])
    if sid == 446:
        assert 444 in outer and 445 not in outer and 446 in outer and 1000 in outer
    if 'second-dot1q' in enc:
        inner = next(n['value'] for n in facts if n['kind'] == 'inner_vlan')
        assert 20 in inner if sid == 600 else 444 in inner and 445 not in inner and 836 in inner
    tag = next(n['value'] for n in facts if n['kind'] == 'tag_rewrite')
    assert tag['operation'] == rewrite.split()[0] and tag['symmetric']
    if sid == 446:
        assert tag['output_tags'] == [['dot1q', 446]]
    assert [n['value'] for n in facts if n['kind'] == 'bridge_domain'] == ([str(bd)] if bd else [])
    assert [n['value'] for n in facts if n['kind'] == 'description'] == ([description] if description else [])
    assert 'vlan protocol type 0x8100' in '\n'.join(e['excerpt'] for e in detail['evidence'])


@pytest.mark.parametrize('family', ['ASR920', 'ME3600X'])
@pytest.mark.parametrize('short,full', [('g0/16', 'GigabitEthernet0/16'),
                                      ('g0/0/12', 'GigabitEthernet0/0/12'),
                                      ('te0/1', 'TenGigabitEthernet0/1')])
def test_precise_negative_and_mismatched_identity(family, short, full):
    text = f'Service Instance ID 1234 on interface {full} does not exist'
    assert parse_service_detail(text, short, 1234)['status'] == 'absent'
    assert parse_service_detail(text, short, 123)['status'] == 'unknown'
    assert parse_service_detail(text, full + '0', 1234)['status'] == 'unknown'
    assert parse_service_detail(text + '\n' + text, short, 1234)['status'] == 'unknown'


@pytest.mark.parametrize('sample', [SAMPLES[0], SAMPLES[2]])
@pytest.mark.parametrize('change', ['id', 'port', 'partial', 'unsupported', 'rejection', 'authorization', 'contradiction', 'bad_bd', 'states'])
def test_unknown_and_error_are_never_absent(sample, change):
    text = response(sample)
    if change == 'id':
        text = text.replace(f'ID: {sample[3]}', 'ID: 999')
    elif change == 'port':
        text = text.replace(sample[2], 'GigabitEthernet9/9')
    elif change == 'partial':
        text = f'Service Instance ID: {sample[3]}\nAssociated Interface: {sample[2]}\n'
    elif change == 'unsupported':
        text = text.replace('Encapsulation: dot1q', 'Encapsulation: unsupported')
    elif change == 'rejection':
        text += '% Invalid input detected at marker\n'
    elif change == 'authorization':
        text += '% Authorization failed\n'
    elif change == 'bad_bd':
        text += 'Bridge-domain: unsupported\n'
    elif change == 'states':
        text += 'State: Down\n'
    else:
        text += 'Bridge-domain: 100\nBridge-domain: 200\n'
    assert parse_service_detail(text, sample[1], sample[3])['status'] == (
        'error' if change in {'rejection', 'authorization'} else 'unknown')


@pytest.mark.parametrize('sample', SAMPLES)
@pytest.mark.parametrize('failure', [False, True])
def test_public_capability_exact_dispatch_and_execution_failure(sample, failure, monkeypatch):
    commands = []
    class CLI:
        def run_command(self, command, *, timeout):
            commands.append(command)
            assert command == f'show ethernet service instance id {sample[3]} interface {sample[1]} detail'
            assert timeout > 0
            if failure:
                raise TimeoutError('synthetic timeout')
            return response(sample)

    monkeypatch.setattr('orbitflow.vendors.cisco.vlans.open_cisco_cli', lambda *a, **kw: nullcontext(CLI()))
    detail = VlanService().service_detail(None, SimpleNamespace(device_family=sample[0]),
                                         sample[1], sample[3], cli=CLI())
    assert len(commands) == 1
    assert detail['status'] == ('error' if failure else 'present')
    if failure:
        assert detail['facts'] == [] and detail['evidence'] == []


@pytest.mark.parametrize('sample', SAMPLES)
def test_planner_json_and_excel_preserve_detail_and_incomplete_forwarding(tmp_path, sample):
    family, _, port, sid, enc, rewrite, bd, _, description = sample
    normalized_enc = enc.replace(' vlan protocol type 0x8100', '')
    normalized_rewrite = rewrite.replace(' vlan-type 0x8100', '')
    # Config contains the BD even when operational forwarding evidence is absent.
    config = f'interface {port}\n description Parent\n'
    if family == 'ME3600X':
        config += ' switchport mode trunk\n switchport trunk allowed vlan none\n'
    config += (f' service instance {sid} ethernet\n  encapsulation dot1q {normalized_enc}\n'
               f'  rewrite ingress tag {normalized_rewrite}\n  bridge-domain {bd or 600}\n!\nend\n')
    params = dict(management_ip=IP, interface=port, service_instance_id=sid,
                  encapsulation='QinQ' if 'second-dot1q' in enc else 'Dot1q',
                  outer_vlan=enc.split()[0], rewrite='Push dot1q symmetric' if sid == 446 else 'Pop 1 symmetric',
                  bridge_domain=bd or 600)
    if 'second-dot1q' in enc:
        params['inner_vlan'] = enc.split('second-dot1q ')[1]
    if sid == 446:
        params['push_vlan'] = 446
    if description:
        params['description'] = description
    t = template(family)
    r = request(t.id, params)
    obs = observation(family, config)
    detail = parse_service_detail(response(sample), port, sid)
    obs['service_details'] = {f'{port}:{sid}': detail}
    original = plan([r], {IP: obs}, now=NOW)
    json_path = tmp_path / 'plan.json'
    save(original, json_path)
    restored = load(json_path)
    assert restored.to_dict()['observations'][IP]['service_details'] == obs['service_details']
    row = restored.to_dict()['results'][0]
    if bd is None:
        assert 'SERVICE_DETAIL_INCOMPLETE' in row['blockers']
        assert 'SERVICE_INSTANCE_DOWN' in row['warnings'] and not row['deployment_ready']
    else:
        assert not row['blockers'] and 'SERVICE_DETAIL_FALLBACK' not in row['warnings']
    path = tmp_path / 'plan.xlsx'
    export_plan(restored, path)
    book = load_workbook(path)
    evidence = book[t.sheet].cell(3, len(t.columns) + EXTRA.index('Existing Evidence') + 1).value
    assert f'Associated Interface: {port}' in evidence and f'Service Instance ID: {sid}' in evidence
    assert 'description Parent' in evidence and 'vlan protocol type 0x8100' in evidence
    assert ('forwarding evidence incomplete' in evidence) == (bd is None)
    book.close()


@pytest.mark.parametrize('family', ['ASR920', 'ME3600X'])
@pytest.mark.parametrize('status', ['absent', 'unknown', 'error'])
def test_negative_lookup_json_evidence_decisions_and_tampering(tmp_path, family, status):
    port, sid = 'GigabitEthernet0/0/12', 1234
    text = f'Service Instance ID {sid} on interface {port} does not exist'
    if status == 'unknown':
        text = text.replace(str(sid), '123')
    elif status == 'error':
        text = '% Invalid input detected\n' + text
    detail = parse_service_detail(text, 'g0/0/12', sid)
    assert detail['status'] == status
    t = template(family)
    r = request(t.id, dict(management_ip=IP, interface=port, service_instance_id=sid,
                          encapsulation='Dot1q', outer_vlan='600',
                          rewrite='Pop 1 symmetric', bridge_domain=600))
    obs = observation(family, f'interface {port}\n description Parent\n!\nend\n')
    obs['service_details'] = {f'{port}:{sid}': detail}
    original = plan([r], {IP: obs}, now=NOW)
    save(original, tmp_path / 'plan.json')
    restored = load(tmp_path / 'plan.json')
    assert restored.to_dict()['observations'][IP]['service_details'] == obs['service_details']
    path = tmp_path / 'plan.xlsx'
    export_plan(restored, path)
    book = load_workbook(path)
    sheet = book[t.sheet]
    col = len(t.columns) + EXTRA.index('Existing Evidence') + 1
    evidence = sheet.cell(3, col).value
    assert 'description Parent' in evidence
    assert ('does not exist (reliable lookup)' in evidence) == (status == 'absent')
    assert ('lookup unavailable; absence not established' in evidence) == (status != 'absent')
    decision_col = len(t.columns) + EXTRA.index('Engineer Decision') + 1
    for decision in ('Pending', 'Approve', 'Reject'):
        sheet.cell(3, decision_col, decision)
        book.save(path)
        reviewed = consume_decisions(restored, path, expected_digest=restored.digest, now=NOW)
        assert reviewed.to_dict()['results'][0]['decision'] == decision
    sheet.cell(3, col, evidence + '\nAltered evidence')
    book.save(path)
    with pytest.raises(PlanError):
        consume_decisions(restored, path, expected_digest=restored.digest, now=NOW)
    book.close()
