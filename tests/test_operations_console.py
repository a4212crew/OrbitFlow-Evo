"""Deterministic console boundaries; no device connections or tsh authentication."""
from copy import deepcopy
from datetime import datetime, timezone
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from orbitflow.environment import PATHS, load_profile, save_profile, validate, validate_session
from orbitflow.menu import Console, Cancelled, main
from orbitflow.config import load_execution_config, ExecutionConfig
from orbitflow.configuration import execution_cli


@pytest.fixture
def settings(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    values = dict(teleport=dict(proxy='proxy.example:443', cluster='cluster.example', bastion_host='bastion', bastion_user='operator'),
                  execution=dict(connection_start_interval=1.0, connection_retry_delay=5.0), paths=PATHS.copy())
    path = Path('data/sys_config/configuration-profile.toml')
    save_profile(values, path)
    workbook = Workbook()
    workbook.active.append(['management_ip', 'username', 'password'])
    workbook.active.append(['192.0.2.1', 'fixture', 'fixture-password'])
    Path('data/input').mkdir(parents=True, exist_ok=True)
    workbook.save('data/input/devices.xlsx')
    return values


def console(answers, **kwargs):
    responses = iter(answers)
    output = []
    return Console(read=lambda _: next(responses), write=output.append, **kwargs), output


def test_menu_navigation(settings):
    ui, output = console(['bad', '1', '6', '4', '4', '2', '3', '3', '3', '5', '2', '7'])
    ui.run()
    assert any('Invalid selection' in line for line in output)
    assert any('Generate Input Template' in line for line in output)
    assert not any('Review Plans' in line for line in output)


def test_concurrency_repeated_and_cancel():
    ui, _ = console(['0', '6', '-1', '2.0', 'x', '5', '', 'b'])
    assert ui.concurrency() == 5
    assert ui.concurrency() == 1
    with pytest.raises(Cancelled):
        ui.concurrency()


@pytest.mark.parametrize('section,key,value', [('teleport', 'password', 'secret'), ('execution', 'max_concurrent_devices', 5),
                                              ('execution', 'connection_retry_delay', float('nan')), ('paths', 'token', 'secret')])
def test_profile_closed_schema(settings, section, key, value):
    values = deepcopy(settings)
    values[section][key] = value
    with pytest.raises(ValueError):
        validate(values)
    with pytest.raises(ValueError):
        save_profile(values)
    assert load_profile() == settings


def test_profile_precedence_legacy_compatibility(settings):
    Path('orbitflow.toml').write_text('[execution]\nmax_concurrent_devices=3\nconnection_retry_delay=8.0\n')
    assert load_execution_config() == ExecutionConfig(5, 1.0, 5.0)
    alternate = Path('alternate.toml')
    alternate.write_text('[execution]\nmax_concurrent_devices=2\n')
    assert load_execution_config(alternate).max_concurrent_devices == 2
    Path('data/sys_config/configuration-profile.toml').unlink()
    assert load_execution_config().max_concurrent_devices == 3


@pytest.mark.parametrize('expiry,proxy,cluster,code', [('2030-01-01T00:00:00Z', 'proxy.example:443', 'cluster.example', 0),
    ('2020-01-01T00:00:00Z', 'proxy.example:443', 'cluster.example', 0),
    ('2030-01-01T00:00:00Z', 'other', 'cluster.example', 0),
    ('2030-01-01T00:00:00Z', 'proxy.example:443', 'other', 0),
    ('2030-01-01T00:00:00Z', 'proxy.example:443', 'cluster.example', 1)])
@pytest.mark.parametrize('proxy_field', ['proxy_url', 'profile_url'])
def test_session(settings, monkeypatch, expiry, proxy, cluster, code, proxy_field):
    monkeypatch.setattr('orbitflow.environment.platform.system', lambda: 'Windows')
    def runner(args, **kwargs):
        assert args == ['tsh', 'status', '--format=json']
        assert kwargs['capture_output']
        return SimpleNamespace(returncode=code, stdout=json.dumps({'active': dict(valid_until=expiry, cluster=cluster, **{proxy_field: proxy})}))
    if expiry.startswith('2030') and proxy == 'proxy.example:443' and cluster == 'cluster.example' and code == 0:
        validate_session(settings, runner=runner, now=datetime(2026, 1, 1, tzinfo=timezone.utc))
    else:
        with pytest.raises(ValueError):
            validate_session(settings, runner=runner)


def test_first_use_cancel_and_atomic_setup(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    ui, _ = console(['b'])
    with pytest.raises(Cancelled):
        ui.settings()
    assert not ui.profile.exists()
    ui, _ = console(['proxy.example:443', 'cluster', 'bastion', 'user', '', '', '', '', '', '', '', '', ''])
    assert ui.settings()['paths'] == PATHS
    assert 'max_concurrent_devices' not in ui.profile.read_text()


def test_template_collision_and_sheets(settings):
    ui, output = console([])
    ui.template()
    first = Path(output[-1])
    original = first.read_bytes()
    ui.template()
    second = Path(output[-1])
    assert first != second and first.read_bytes() == original
    assert first.name.endswith('_001.xlsx') and second.name.endswith('_002.xlsx')
    book = load_workbook(second, read_only=True)
    assert len(book.sheetnames) == 23
    book.close()


def test_cancel_before_session_or_service(settings):
    ui, _ = console(['', '', 'n'], session_check=lambda _: pytest.fail('session checked after cancel'))
    with pytest.raises(Cancelled):
        ui.network(settings, 'data/input/devices.xlsx', 'Read-only')


@pytest.mark.parametrize('kind,module,function', [('inventory', 'orbitflow.inventory_refresh', 'refresh_inventory_from_excel'),
    ('report', 'orbitflow.reporting', 'run_report'), ('backup', 'orbitflow.configuration_backup', 'run_backup'),
    ('compliance', 'orbitflow.vlan_compliance', 'run_compliance')])
def test_dispatch(settings, monkeypatch, kind, module, function):
    import importlib
    calls = []
    monkeypatch.setattr(importlib.import_module(module), function, lambda *a, **kw: calls.append((a, kw)))
    ui, _ = console(['', '', '2', 'y'], session_check=lambda _: None)
    ui.run_workflow(kind, methodology=True)
    assert calls[0][1]['execution_config'].max_concurrent_devices == 2
    assert calls[0][1]['inventory_path'] == PATHS['inventory']
    if kind == 'compliance':
        assert calls[0][1]['methodology_report'] and not calls[0][1]['cleanup']
    assert 'max_concurrent_devices' not in ui.profile.read_text()


def test_expired_session_stops_dispatch(settings, monkeypatch):
    from orbitflow import reporting
    monkeypatch.setattr(reporting, 'run_report', lambda *a, **k: pytest.fail('network dispatched'))
    def expired(_):
        raise ValueError('expired')
    ui, _ = console(['', '', '', 'y'], session_check=expired)
    with pytest.raises(ValueError):
        ui.run_workflow('report')


@pytest.mark.parametrize('kind,module,function,answers', [('execution', 'orbitflow.configuration.execution', 'export_results', ['run']),
    ('vlan_compliance', 'orbitflow.compliance_report', 'export_compliance_spool', ['run', 'n', 'y']),
    ('inventory', 'orbitflow.inventory_refresh', 'export_inventory_spool', ['run', 'n']),
    ('interface_vlan', 'orbitflow.reporting', 'export_report_spool', ['run', 'n'])])
def test_offline_recovery_without_profile(tmp_path, monkeypatch, kind, module, function, answers):
    import importlib
    monkeypatch.chdir(tmp_path)
    calls = []
    monkeypatch.setattr(importlib.import_module(module), function, lambda *a, **kw: calls.append((a, kw)))
    ui, _ = console(answers, session_check=lambda _: pytest.fail('offline session'))
    ui.recover(kind)
    assert len(calls) == 1 and not ui.profile.exists()
    if kind != 'execution':
        assert calls[0][1]['cleanup'] is False


def test_help(capsys):
    with pytest.raises(SystemExit) as exc:
        main(['--help'])
    assert exc.value.code == 0
    assert '--profile' in capsys.readouterr().out


@pytest.mark.parametrize('confirm', ['n', 'y'])
def test_execution_uses_bound_prepare_confirm(settings, monkeypatch, confirm):
    from orbitflow.configuration import execution as ex
    from orbitflow.configuration.v2 import planning
    Path('plan.json').touch()
    Path('plan.xlsx').touch()
    calls = []
    prepared = dict(status='awaiting_confirmation', reference='bound-reference', summary_digest='bound-digest',
                    summary=dict(selected=[dict(request=dict(inputs=dict(management_ip='192.0.2.1')))]))
    class Service:
        def __init__(self, authority):
            assert authority == 'data/configuration/execution'
        def prepare_workbook(self, plan, workbook, **kwargs):
            calls.append(('prepare', kwargs))
            return prepared
        def confirm(self, reference, **kwargs):
            assert reference == 'bound-reference'
            calls.append(('confirm', kwargs))
            return SimpleNamespace(path=Path('retained-run'))
    monkeypatch.setattr(ex, 'ExecutionService', Service)
    monkeypatch.setattr(ex, 'operator_summary', lambda p: p)
    monkeypatch.setattr(ex, 'render_summary', lambda p: 'summary with safety policy\nWarnings/blockers or skipped selections: verbose\nFresh checks + backup before writes: verbose')
    monkeypatch.setattr(ex, 'export_results', lambda *a: calls.append(('export', a)))
    monkeypatch.setattr(execution_cli, 'outcome_summary', lambda *a, **k: 'Completed')
    monkeypatch.setattr(planning, 'load', lambda p: SimpleNamespace(digest='original-digest'))
    ui, output = console(['plan.json', '', '', '3', confirm], session_check=lambda _: calls.append(('session', {})))
    if confirm == 'n':
        with pytest.raises(Cancelled):
            ui.execute()
        assert [c[0] for c in calls] == ['prepare']
    else:
        ui.execute()
        assert [c[0] for c in calls] == ['prepare', 'session', 'confirm', 'export']
        assert calls[2][1]['summary_digest'] == 'bound-digest'
        assert calls[2][1]['apply'] is True
    assert 'Preparing/validating plan...' in output
    assert 'Preparing/validating plan: finished.' in output
    assert not any('verbose' in text for text in output)
    assert any('Whole-config save' in text for text in output)
    assert calls[0][1]['expected_digest'] == 'original-digest'
    assert calls[0][1]['settings'].max_concurrent_devices == 3


def test_corrupt_plan_fails_before_session(settings):
    Path('corrupt.json').write_text('{}')
    Path('corrupt.xlsx').touch()
    ui, _ = console(['corrupt.json', '', 'digest', '', '1'], session_check=lambda _: pytest.fail('corrupt plan connected'))
    with pytest.raises(Exception):
        ui.execute()


def test_missing_headers_and_missing_file_fail_without_session(settings):
    workbook = Workbook()
    workbook.active.append(['wrong'])
    workbook.save('invalid.xlsx')
    ui, _ = console([], session_check=lambda _: pytest.fail('bad workbook connected'))
    with pytest.raises(ValueError):
        ui.network(settings, 'invalid.xlsx', 'Read-only')
    with pytest.raises(FileNotFoundError):
        ui.network(settings, 'missing.xlsx', 'Read-only')


def test_linux_identity_and_malformed_session(settings, monkeypatch):
    monkeypatch.setattr('orbitflow.environment.platform.system', lambda: 'Linux')
    status = dict(active=dict(proxy_url='proxy.example:443', cluster='cluster.example', valid_until='2030-01-01T00:00:00Z'))
    runner = lambda *a, **k: SimpleNamespace(returncode=0, stdout=json.dumps(status))
    with pytest.raises(ValueError):
        validate_session(settings, runner=runner)
    for key in ('teleport_key_path', 'teleport_cert_path'):
        Path(key).touch()
        settings['teleport'][key] = key
    validate_session(settings, runner=runner)
    status['active']['valid_until'] = 'invalid'
    with pytest.raises(ValueError):
        validate_session(settings, runner=runner)


def test_parallel_template_names(settings):
    from concurrent.futures import ThreadPoolExecutor
    def generate(_):
        ui, output = console([])
        ui.template()
        return output[-1]
    with ThreadPoolExecutor(max_workers=2) as pool:
        paths = list(pool.map(generate, range(2)))
    assert len(set(paths)) == 2
    assert all(Path(path).stat().st_size > 0 for path in paths)


def test_review_offline_summary(settings):
    root = Path(PATHS['execution_reports'])
    root.mkdir(parents=True)
    book = Workbook()
    book.active.title = 'Run Summary'
    book.active.append(['Run', 'fixture-run'])
    book.active.append(['Rows by verification outcome', 'Unknown: 1'])
    book.save(root / 'execution.xlsx')
    ui, output = console([], session_check=lambda _: pytest.fail('review connected'))
    ui.review()
    assert 'Rows by verification outcome: Unknown: 1' in output


@pytest.mark.parametrize('corrupt', [False, True])
def test_recovery_custom_output_or_invalid_profile_fallback(settings, monkeypatch, corrupt):
    from orbitflow.configuration import execution
    settings['paths']['execution_reports'] = 'custom-reports'
    save_profile(settings)
    if corrupt:
        Path('data/sys_config/configuration-profile.toml').write_text('invalid TOML')
    outputs = []
    monkeypatch.setattr(execution, 'export_results', lambda run, output: outputs.append(output))
    ui, _ = console(['retained-run'], session_check=lambda _: pytest.fail('offline session'))
    ui.recover('execution')
    assert outputs[0].parent == Path('outputs/reports/configuration_execution' if corrupt else 'custom-reports')

@pytest.mark.parametrize('bad', ['garbage-secret', 'nan', 'inf', '-1', '1e9999'])
def test_setup_pacing_reprompts_without_echo(settings, bad):
    before = Path('data/sys_config/configuration-profile.toml').read_bytes()
    ui, output = console([''] * 6 + [bad, '2.5', 'b'])
    with pytest.raises(Cancelled):
        ui.setup()
    assert Path('data/sys_config/configuration-profile.toml').read_bytes() == before
    assert any('finite nonnegative' in line for line in output)
    assert not any(bad in line for line in output)
    ui, _ = console([''] * 6 + [bad, '2.5', ''] + [''] * 5)
    ui.setup()
    assert load_profile()['execution']['connection_start_interval'] == 2.5


def test_subset_validation_and_cancel(settings):
    from orbitflow.targets import load_targets
    targets = load_targets('data/input/devices.xlsx') * 3
    ui, output = console(['0', '4', '1,1', 'secret', '3,1'])
    assert ui.select_targets(targets) == [0, 2]
    assert not any('fixture-password' in line for line in output)
    ui, _ = console(['b'])
    with pytest.raises(Cancelled):
        ui.select_targets(targets)


@pytest.mark.parametrize('answers', [['b'], ['4,3,1', '', 'n']])
def test_target_identity_mapping_and_cancel(settings, monkeypatch, answers):
    from orbitflow import targets as target_loader
    rows = [dict(management_ip=address, username='hidden-user', password='hidden-password',
                 secret='hidden-secret', display_hint='hidden-hint\x1b[2J')
            for address in ('192.0.2.1', 'hidden-address\x1b[2J\r\n', '192.0.2.1',
                            '2001:0db8::1', 'fe80::1%hidden-scope\x1b[2J', '')]
    monkeypatch.setattr(target_loader, 'load_targets', lambda *a, **kw: rows)
    ui, output = console(answers, session_check=lambda _: pytest.fail('session checked after cancellation'))
    with pytest.raises(Cancelled):
        ui.network(settings, 'data/input/devices.xlsx', 'Read-only')
    assert all('hidden' not in line and not any(ord(c) < 32 for c in line) for line in output)
    assert rows[1]['management_ip'] == 'hidden-address\x1b[2J\r\n'
    if answers[0] != 'b':
        assert ui.target_positions == [0, 2, 3]
        assert output[-1] == 'Selected targets: Row 1: 192.0.2.1, Row 3: 192.0.2.1, Row 4: 2001:db8::1'


@pytest.mark.parametrize('kind', ['inventory', 'report', 'backup', 'compliance'])
def test_subset_dispatch(settings, monkeypatch, kind):
    import importlib
    book = load_workbook('data/input/devices.xlsx')
    book.active.append(['192.0.2.2', 'fixture', 'fixture-password'])
    book.save('data/input/devices.xlsx')
    book.close()
    module, function = {'inventory': ('inventory_refresh', 'refresh_inventory_from_excel'),
                        'report': ('reporting', 'run_report'), 'backup': ('configuration_backup', 'run_backup'),
                        'compliance': ('vlan_compliance', 'run_compliance')}[kind]
    calls = []
    monkeypatch.setattr(importlib.import_module('orbitflow.' + module), function, lambda *a, **kw: calls.append((a, kw)))
    ui, _ = console(['', '2', '', 'y'], session_check=lambda _: None)
    ui.run_workflow(kind)
    if kind == 'inventory':
        assert calls[0][1]['target_positions'] == [1]
    else:
        assert [r['management_ip'] for r in calls[0][0][0]] == ['192.0.2.2']


def test_safe_diagnostics_are_fixed(settings, monkeypatch):
    from orbitflow.menu import OperatorError
    cases = [(PermissionError('password=hidden'), 'Permission denied'),
             (FileNotFoundError('password=hidden'), 'Required input'),
             (RuntimeError('password=hidden'), 'Operation failed')]
    for error, expected in cases:
        ui, output = console([])
        def fail():
            raise error
        ui.safe(fail)
        assert expected in output[-1] and 'hidden' not in str(output)
    Path('data/sys_config/configuration-profile.toml').write_text('password=hidden')
    ui, output = console([])
    ui.safe(ui.settings)
    assert 'Invalid environment profile' in output[-1] and 'hidden' not in str(output)


def test_bad_session_and_spool_diagnostics(settings):
    def expired(_):
        raise ValueError('token=hidden raw device output')
    ui, output = console(['', '', 'y'], session_check=expired)
    ui.safe(lambda: ui.network(settings, 'data/input/devices.xlsx', 'Read-only'))
    assert 'Teleport session' in output[-1] and 'hidden' not in str(output)
    ui, output = console(['missing-run'])
    ui.safe(lambda: ui.recover('execution'))
    assert 'Invalid export' in output[-1] or 'Required input' in output[-1]


def test_review_skips_invalid_workbooks(settings):
    root = Path(PATHS['execution_reports'])
    root.mkdir(parents=True)
    for i in range(12):
        (root / f'a{i:02}.xlsx').write_bytes(b'password=hidden')
    book = Workbook()
    book.save(root / 'b-not-report.xlsx')
    book.active.title = 'Run Summary'
    book.active.append(['Run', 'good-run'])
    book.save(root / 'z-valid.xlsx')
    book.close()
    ui, output = console([], session_check=lambda _: pytest.fail('offline connected'))
    ui.review()
    assert 'Run: good-run' in output
    assert sum('Skipped workbook #' in line for line in output) == 10
    assert '3 additional invalid workbooks skipped.' in output
    assert 'hidden' not in str(output)


def test_real_offline_planner_and_review_adapter(settings):
    from orbitflow.configuration.v2.catalogue import CATALOGUE
    from orbitflow.configuration.v2.workbook import create_input
    from orbitflow.configuration.v2.planning import load
    t = next(t for t in CATALOGUE if t.family == 'All' and t.function == 'Manual CLI')
    create_input('manual.xlsx')
    book = load_workbook('manual.xlsx')
    values = dict(management_ip='192.0.2.1', platform='cisco_ios', mode='interface', scope='Gi1/0/1', commands=' description Lab\n exit\n')
    for col, name in enumerate(t.columns, 1):
        book[t.sheet].cell(3, col, values[name])
    book.save('manual.xlsx')
    book.close()
    ui, output = console(['1', 'manual.xlsx'], session_check=lambda _: pytest.fail('manual planning connected'))
    ui.planning()
    source = next(Path(PATHS['plans']).rglob('plan.json'))
    original = load(source)
    assert original.digest in output[-1]
    assert source.with_suffix('.xlsx').is_file()
    ui, output = console([str(source), '', '', '1', 'n'], session_check=lambda _: pytest.fail('declined execution connected'))
    with pytest.raises(Cancelled):
        ui.execute()
    assert any(original.digest in line for line in output)
    source.with_suffix('.xlsx').write_bytes(b'corrupt-secret')
    ui, output = console([str(source), '', '', '1'])
    ui.safe(ui.execute)
    assert 'Plan/workbook invalid' in output[-1] and 'corrupt-secret' not in str(output)


def test_real_report_recovery_and_corrupt_spool(settings):
    from orbitflow.result_spool import ResultSpool
    spool = ResultSpool.create('runs', 'interface_vlan_report', 0)
    with spool.collection():
        pass
    ui, output = console([str(spool.path), 'n'], session_check=lambda _: pytest.fail('recovery connected'))
    ui.recover('interface_vlan')
    result = Path(output[-1])
    book = load_workbook(result, read_only=True)
    assert len(book.sheetnames) == 3
    book.close()
    assert (spool.path / 'results.jsonl').exists()
    (spool.path / 'manifest.json').write_text('corrupt token=hidden')
    ui, output = console([str(spool.path), 'n'])
    ui.safe(lambda: ui.recover('interface_vlan'))
    assert 'Invalid export' in output[-1] and 'hidden' not in str(output)


def test_real_planner_subset(settings, monkeypatch):
    from orbitflow.configuration.v2.catalogue import CATALOGUE
    from orbitflow.configuration.v2.workbook import create_input
    from orbitflow.configuration.v2.planning import load
    from orbitflow.configuration.v2 import observation
    t = next(t for t in CATALOGUE if t.family == 'C3750X' and t.function == 'VLAN Database')
    create_input('subset.xlsx')
    book = load_workbook('subset.xlsx')
    for row, ip in enumerate(['192.0.2.1', '192.0.2.2'], 3):
        values = dict(management_ip=ip, vlan_id='100', name='Lab')
        for col, name in enumerate(t.columns, 1):
            book[t.sheet].cell(row, col, values.get(name, ''))
    book.save('subset.xlsx')
    book.close()
    book = load_workbook('data/input/devices.xlsx')
    book.active.append(['192.0.2.2', 'fixture', 'fixture-password'])
    book.save('data/input/devices.xlsx')
    book.close()
    captured = []
    def collect(requests, credentials, *args, **kwargs):
        captured.append((list(requests), credentials))
        kwargs['progress']('192.0.2.2: connect; identity=private')
        kwargs['progress']('192.0.2.2: completed; identity=private')
    monkeypatch.setattr(observation, 'collect', collect)
    monkeypatch.setattr(observation, 'bounded_observations', lambda spool: {})
    ui, output = console(['1', 'subset.xlsx', '', '2', '', 'y'], session_check=lambda _: None)
    ui.planning()
    assert list(captured[0][1]) == ['192.0.2.2']
    assert any('persisted 1/1 (100%)' in line for line in output)
    assert 'Plan and review export complete.' in output
    assert not any('private' in line for line in output)
    original = load(next(Path(PATHS['plans']).rglob('plan.json')))
    assert [r['inputs']['management_ip'] for r in original.to_dict()['requests']] == ['192.0.2.2']


@pytest.mark.parametrize('missing', [False, True])
def test_plan_failure_diagnostic_before_session(settings, missing):
    Path('bad.xlsx').touch()
    if not missing:
        Path('bad.json').write_text('{"token":"hidden"}')
    ui, output = console(['bad.json', 'bad.xlsx'], session_check=lambda _: pytest.fail('bad plan connected'))
    ui.safe(ui.execute)
    assert ('Required input' if missing else 'Invalid immutable plan') in output[-1]
    assert 'hidden' not in str(output)


@pytest.mark.parametrize('fields', [
    {}, {'profile_url': ''}, {'profile_url': None},
    {'profile_url': 'https://other:443'},
    {'proxy_url': 'proxy.example:443', 'profile_url': 'other'},
    {'proxy_url': 'other', 'profile_url': 'proxy.example:443'},
])
def test_session_rejects_missing_or_conflicting_proxy(settings, monkeypatch, fields):
    monkeypatch.setattr('orbitflow.environment.platform.system', lambda: 'Windows')
    status = dict(active=dict(valid_until='2030-01-01T00:00:00Z', cluster='cluster.example', **fields))
    with pytest.raises(ValueError):
        validate_session(settings, runner=lambda *a, **k: SimpleNamespace(returncode=0, stdout=json.dumps(status)))


def test_session_matching_urls(settings, monkeypatch):
    monkeypatch.setattr('orbitflow.environment.platform.system', lambda: 'Windows')
    status = dict(active=dict(valid_until='2030-01-01T00:00:00Z', cluster='cluster.example',
                             proxy_url='proxy.example:443', profile_url='https://proxy.example:443/'))
    validate_session(settings, runner=lambda *a, **k: SimpleNamespace(returncode=0, stdout=json.dumps(status)))


def test_workbook_chooser_order_and_eligibility(tmp_path):
    import os
    for name in ('older.xlsx', 'z.xlsx', 'a.XLSX', '~$lock.xlsx', 'ignored.xls'):
        path = tmp_path / name
        path.touch()
        os.utime(path, ns=(100, 100 if name == 'older.xlsx' else 200))
    (tmp_path / 'folder.xlsx').mkdir()
    ui, output = console(['0', 'bad', '2'])
    assert ui.choose_workbook(tmp_path).name == 'z.xlsx'
    assert output[:4] == ['1. a.XLSX', '2. z.xlsx', '3. older.xlsx', '4. Enter explicit file path']
    ui, _ = console(['4', str(tmp_path / 'older.xlsx')])
    assert ui.choose_workbook(tmp_path).name == 'older.xlsx'
    ui, _ = console(['b'])
    with pytest.raises(Cancelled):
        ui.choose_workbook(tmp_path)


def test_empty_workbook_chooser(tmp_path):
    ui, output = console(['2', 'b'])
    with pytest.raises(Cancelled):
        ui.choose_workbook(tmp_path / 'missing')
    assert any('No configuration workbooks' in line for line in output)
    path = tmp_path / 'explicit.xlsx'
    path.touch()
    ui, _ = console(['1', str(path)])
    assert ui.choose_workbook(tmp_path / 'missing') == path


@pytest.mark.parametrize('selection', ['all', '300', ','.join(str(i) for i in range(1, 301))])
def test_large_target_output_bounded_and_cancel(settings, monkeypatch, selection):
    from orbitflow import targets as target_loader
    rows = [dict(management_ip=f'2001:db8::{i:x}') for i in range(1, 301)]
    monkeypatch.setattr(target_loader, 'load_targets', lambda *a, **kw: rows)
    ui, output = console([selection, '', 'n'], session_check=lambda _: pytest.fail('cancel connected'))
    with pytest.raises(Cancelled):
        ui.network(settings, 'data/input/devices.xlsx', 'Read-only')
    assert len(str(output)) < 1000
    if selection == 'all':
        assert 'Selected targets: all 300 device rows.' in output
        assert '2001:db8' not in str(output)
    elif selection == '300':
        assert 'Selected targets: Row 300: 2001:db8::12c' in output
    else:
        assert '290 additional selected rows omitted; total 300.' in output


def test_navigation_separates_each_menu(settings):
    ui, output = console(['1', '6', '7'])
    ui.run()
    menus = [i for i, text in enumerate(output) if text.startswith(('1. Configuration', 'Configuration Management\n'))]
    assert len(menus) == 3
    assert all(output[i-1] == '' for i in menus)


def test_observation_progress_counts_persisted_and_hides_event_text():
    ui, output = console([])
    callback = ui.observation_progress({'192.0.2.1', '2001:db8::1'})
    callback('192.0.2.1: connect; identity=secret')
    callback('192.0.2.1: completed; identity=secret')
    callback('192.0.2.1: completed; identity=secret')
    callback('2001:db8::1: failed; identity=secret (SecretError)')
    callback('192.0.2.1: token=secret; identity=secret')
    assert len(output) == 4
    assert '0/2 (0%)' in output[0]
    assert '1/2 (50%)' in output[2]
    assert '2/2 (100%)' in output[3]
    assert 'failed' in output[3]
    assert not any('secret' in line or 'SecretError' in line for line in output)
    ui.write = lambda text: (_ for _ in ()).throw(RuntimeError('display'))
    callback('192.0.2.1: connect; identity=secret')


def test_activity_failure_never_claims_finished():
    ui, output = console([])
    with pytest.raises(Cancelled):
        with ui.activity('Export'):
            raise Cancelled()
    assert output == ['Export...', 'Export: stopped; completion not established.']


def test_plan_chooser_validates_orders_and_defaults_review(tmp_path, monkeypatch):
    from orbitflow.configuration.v2 import planning
    import os
    for name in ('old', 'new', 'bad'):
        folder = tmp_path / name
        folder.mkdir()
        (folder / 'plan.json').write_text(name)
    (tmp_path / 'unrelated.json').write_text('{}')
    os.utime(tmp_path / 'old/plan.json', (1, 1))
    os.utime(tmp_path / 'new/plan.json', (2, 2))
    def load(path):
        if path.parent.name == 'bad':
            raise ValueError('secret')
        return SimpleNamespace(digest='digest')
    monkeypatch.setattr(planning, 'load', load)
    ui, output = console(['0', '1'])
    assert ui.choose_plan(tmp_path) == tmp_path / 'new/plan.json'
    assert any('new / plan.json' in text for text in output)
    assert not any('secret' in text or 'unrelated' in text for text in output)
    ui, _ = console(['b'])
    with pytest.raises(Cancelled):
        ui.choose_plan(tmp_path / 'empty')


def test_activity_elapsed_snapshot_and_cleanup(monkeypatch):
    from orbitflow import menu
    from orbitflow.progress import Progress
    clock = [0]
    ui, output = console([])
    progress = Progress(lambda event: None, clock=lambda: clock[0])
    def factory(subscriber):
        progress.subscriber = subscriber
        return progress
    monkeypatch.setattr(menu, 'Progress', factory)
    with ui.activity('Planning'):
        clock[0] = 6
        progress.tick()
    assert output == ['Planning...', 'Planning: still running (6s elapsed).', 'Planning: finished.']
    assert not progress.thread.is_alive()


def test_configuration_evidence_keeps_operational_lookup_provenance(tmp_path):
    from orbitflow.configuration.v2.evidence import scoped_evidence
    from orbitflow.configuration.v2.catalogue import BY_ID
    template = next(t for t in BY_ID.values() if t.family == 'C3750X' and t.function == 'VLAN Database')
    config = dict(source_filename='saved config', line=1, excerpt='vlan 100')
    operational = dict(source_filename='show vlan', line=2, excerpt='100 lab active')
    observation = dict(configuration_reliable=True,
                       facts=[dict(kind='vlan', value=100, children=[], **config)],
                       vlan_lookups={'100': dict(status='present', operational_status='active', evidence=[operational])})
    result = scoped_evidence(dict(template_id=template.id, inputs=dict(vlan_id=100)), observation)
    assert config in result and operational in result
    from orbitflow.configuration.v2.planning import plan, request
    from orbitflow.configuration.v2.workbook import export_plan
    # Exercise workbook presentation independently of live collection.
    artifact = plan([request(template.id, dict(management_ip='192.0.2.1', vlan_id=100, name='Lab'))], {})
    data = artifact.to_dict()
    data['observations']['192.0.2.1'] = observation
    expected = data['results'][0]
    export_plan(SimpleNamespace(digest=artifact.digest, to_dict=lambda: data), tmp_path / 'review.xlsx')
    book = load_workbook(tmp_path / 'review.xlsx')
    try:
        sheet = book[template.sheet]
        cells = dict(zip([c.value for c in sheet[2]], [c.value for c in sheet[3]]))
        assert config['excerpt'] in cells['Existing Evidence']
        assert operational['excerpt'] in cells['Existing Evidence']
        assert cells['Live Validation'] == expected['live_validation']
        assert cells['Blockers'] == ','.join(expected['blockers'])
        assert cells['Engineer Decision'] == expected['decision']
    finally:
        book.close()


@pytest.mark.parametrize('failure', ['save', 'export_plan'])
def test_planning_export_failure_never_claims_complete(settings, monkeypatch, failure):
    from orbitflow.configuration.v2 import planning, workbook
    from orbitflow.configuration.v2.catalogue import CATALOGUE
    template = next(t for t in CATALOGUE if t.family == 'All')
    workbook.create_input('manual.xlsx')
    book = load_workbook('manual.xlsx')
    values = dict(management_ip='192.0.2.1', platform='cisco_ios', mode='interface', scope='Gi1/0/1', commands=' description Lab\n exit\n')
    for col, name in enumerate(template.columns, 1):
        book[template.sheet].cell(3, col, values[name])
    book.save('manual.xlsx')
    book.close()
    def fail(*args):
        raise OSError('private diagnostic')
    monkeypatch.setattr(planning if failure == 'save' else workbook, failure, fail)
    ui, output = console(['1', 'manual.xlsx'], session_check=lambda _: pytest.fail('offline connected'))
    ui.safe(ui.planning)
    assert 'Plan and review export complete.' not in output
    assert not any('private diagnostic' in text for text in output)
    assert any('stopped; completion not established' in text for text in output)
