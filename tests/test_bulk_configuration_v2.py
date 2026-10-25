"""Synthetic, deterministic contracts; no live equipment or credentials."""
from dataclasses import asdict
from datetime import datetime, timedelta, timezone
from contextlib import contextmanager
import socket
from pathlib import Path

import pytest
from openpyxl import load_workbook

from orbitflow.configuration.plans import PlanError, canonical
from orbitflow.configuration.v2.catalogue import CATALOGUE, FAMILIES
from orbitflow.configuration.v2.planning import request, plan, evaluate_decisions as review, BulkPlan, save, load
from orbitflow.configuration.v2.workbook import create_input, read_input, export_plan as export_review, consume_decisions, EXTRA, output_columns
def import_review(original, path, **kwargs):
    return consume_decisions(original, path, expected_digest=original.digest, **kwargs)
from orbitflow.configuration.v2.observation import snapshot
from orbitflow.models import DeviceContext, InterfaceRecord, VlanObject, VlanState, ConfigFact
from orbitflow.compliance.vlan import safe_data

NOW = datetime(2026, 10, 11, tzinfo=timezone.utc)
IP = "192.0.2.10"


def template(family, function):
    return next(t for t in CATALOGUE if (t.family, t.function) == (family, function))


def inputs(t):
    p = {"management_ip": IP}
    examples = {
        "VLAN Database": dict(vlan_id=4002, name="Customer"),
        "Access Port": dict(interface="Gi1/0/1", vlan_id=4002),
        "Trunk Port": dict(interface="Gi1/0/1", action="Add", vlans=[4002]),
        "Interface Description": dict(interface="Gi1/0/1", description="Customer Site Alpha"),
        "EVC": dict(interface="Gi1/0/1", service_instance_id=10, encapsulation="Dot1q", outer_vlan="4002", bridge_domain=4002, rewrite="Pop 1 symmetric"),
        "L2VPN": dict(interface="Gi1/0/1", subinterface=10, encapsulation="dot1q", outer_vlan=4002, bridge_group="CUSTOMER", bridge_domain="BLUE"),
        "VLAN Port": dict(interface="0/1", action="Remove", tagged_vlans="4002"),
        "Subinterface VSI": dict(interface="Gi1/0/1", subinterface=10, vlan_id=4002, encapsulation="vlan-type dot1q", vsi="BLUE"),
        "Manual CLI": dict(platform="cisco_ios", mode="interface", scope="Gi1/0/1", commands=" description Customer Site\n exit\n"),
    }
    values = {**p, **examples[t.function]}
    if t.family in {"C3750X", "C3850X"} and "interface" in values:
        values['interface'] = 'GigabitEthernet1/0/1'
    return values


def observed(t, *, objects=None):
    platform, profile = FAMILIES[t.family]
    family = "C3850" if t.family == "C3850X" else t.family
    context = DeviceContext("device-1", IP, (IP,), "example", "Example", platform, family, family,
                            profile, (), "EXAMPLE", "1", "", NOW, NOW)
    name = inputs(t).get("interface", "Gi1/0/1")
    interface = InterfaceRecord("example", IP, platform, name, "Before", "up", "up", NOW)
    facts = (ConfigFact("interface", name, "running-config", 1, f"interface {name}"),)
    state = VlanState("example", IP, platform, (), tuple(objects or ()), NOW, facts)
    return snapshot(context, (interface,), state)


@pytest.mark.parametrize("t", CATALOGUE, ids=lambda t: t.id)
def test_all_22_schema_and_representative_cli(t):
    assert len(CATALOGUE) == 22
    p = inputs(t)
    r = request(t.id, p)
    result = plan([r], {} if t.family == "All" else {IP: observed(t)}, now=NOW).to_dict()["results"][0]
    cmds = [c.replace("GigabitEthernet1/0/1", "Gi1/0/1") for c in result["commands"]] if isinstance(result["commands"], list) else result["commands"]
    assert not result["execution_authorized"]
    if t.family == "All":
        assert cmds == p["commands"]
        assert result["live_validation"] == "not_live_validated"
    elif t.function == "VLAN Database":
        expected = (["system-view", "vlan 4002", "description Customer", "quit", "quit"] if t.family == "NE05E" else
                    ["vlan database", "vlan 4002", 'vlan name 4002 "Customer"', "exit"] if t.family == "EdgeSwitch" else
                    ["configure terminal", "vlan 4002", "name Customer", "end"])
        assert cmds == expected
    elif t.function == "Interface Description":
        expected = (["system-view", "interface Gi1/0/1", "description Customer Site Alpha", "quit", "quit"] if t.family == "NE05E" else
                    ["configure", "interface Gi1/0/1", 'description "Customer Site Alpha"', "exit", "exit"] if t.family == "EdgeSwitch" else
                    ["configure terminal", "interface Gi1/0/1", "description Customer Site Alpha", "end"])
        assert cmds == expected
    elif t.function == "Access Port":
        assert cmds == ["configure terminal", "interface Gi1/0/1", "switchport", "switchport mode access", "switchport access vlan 4002", "end"]
    elif t.function == "Trunk Port":
        assert cmds == ["configure terminal", "interface Gi1/0/1", "switchport", *(["switchport trunk encapsulation dot1q"] if t.family == "C3750X" else []), "switchport mode trunk", "switchport trunk allowed vlan add 4002", "end"]
    elif t.function == "EVC":
        assert cmds == ["configure terminal", "interface Gi1/0/1", *(["switchport mode trunk", "switchport trunk allowed vlan none"] if t.family == "ME3600X" else ["no ip address"]), "no service instance 10", "service instance 10 ethernet", "encapsulation dot1q 4002", "rewrite ingress tag pop 1 symmetric", "bridge-domain 4002", "end"]
    elif t.function == "L2VPN":
        assert cmds == ["configure terminal", "interface Gi1/0/1.10 l2transport", "encapsulation dot1q 4002", "exit", "l2vpn", "bridge group CUSTOMER", "bridge-domain BLUE", "interface Gi1/0/1.10", "end"]
    elif t.function == "VLAN Port":
        assert cmds == ["configure", "interface 0/1", "vlan participation exclude 4002", "exit", "exit"]
    elif t.function == "Subinterface VSI":
        assert cmds == ["system-view", "interface Gi1/0/1.10", "vlan-type dot1q 4002", "l2 binding vsi BLUE", "quit", "quit"]
    with pytest.raises(PlanError):
        request(t.id, {**p, "password": "FORBIDDEN"})
    for required in t.required:
        bad = dict(p)
        del bad[required]
        with pytest.raises(PlanError):
            request(t.id, bad)


def test_database_independence_order_rejection_and_digest():
    db, access = template("C3750X", "VLAN Database"), template("C3750X", "Access Port")
    a = request(db.id, inputs(db), sheet_order=1)
    b = request(access.id, inputs(access), sheet_order=2)
    original = plan([b, a], {IP: observed(db)}, now=NOW)
    assert original.to_dict()["results"][1]["dependencies"] == []
    approved = review(original, {a["row_id"]: "Approve", b["row_id"]: "Approve"}, expected_digest=original.digest, now=NOW)
    assert all(r["deployment_ready"] for r in approved.to_dict()["results"])
    rejected = review(original, {a["row_id"]: "Reject", b["row_id"]: "Approve"}, expected_digest=original.digest, now=NOW)
    assert rejected.to_dict()["results"][1]["deployment_ready"]
    assert original.digest == approved.digest == rejected.digest
    with pytest.raises(PlanError):
        review(original, {}, expected_digest="0" * 64)


@pytest.mark.parametrize("family", ["C3750X", "C3850X", "ME3600X", "EdgeSwitch", "NE05E"])
def test_database_rename_noop_and_limits(family):
    t = template(family, "VLAN Database")
    r = request(t.id, inputs(t))
    for name, no_op in (("Customer", True), ("Before", False)):
        obs = observed(t, objects=[VlanObject("vlan", "4002", name)])
        result = plan([r], {IP: obs}, now=NOW).to_dict()["results"][0]
        assert result["no_op"] is no_op
        assert bool(result["commands"]) != no_op
        assert ("VLAN_RENAME_IMPACT" in result["findings"]) != no_op
    for invalid in (0, 4095, True, "100"):
        with pytest.raises(PlanError):
            request(t.id, {**inputs(t), "vlan_id": invalid})
    if family == "EdgeSwitch":
        with pytest.raises(PlanError):
            request(t.id, {**inputs(t), "vlan_id": 1})
    else:
        request(t.id, {**inputs(t), "vlan_id": 1})
    if family in {"C3750X", "C3850X", "ME3600X"}:
        with pytest.raises(PlanError):
            request(t.id, {**inputs(t), "vlan_id": 1002})


def test_missing_and_stale_evidence_fail_closed():
    t = template("C3750X", "VLAN Database")
    r = request(t.id, inputs(t))
    for obs in ({}, {IP: observed(t)}):
        result = plan([r], obs, now=NOW + timedelta(hours=2), decisions={r["row_id"]: "Approve"}).to_dict()["results"][0]
        assert result["commands"] and not result["deployment_ready"]


def test_me_bridge_domain_not_outer_tag_and_ne_existing_vsi():
    t = template("ME3600X", "EVC")
    p = {**inputs(t), "outer_vlan": 123, "bridge_domain": 456}
    r = request(t.id, p)
    obs = observed(t, objects=[VlanObject("vlan", "123")])
    assert "GLOBAL_VLAN_REQUIRED" in plan([r], {IP: obs}, now=NOW).to_dict()["results"][0]["findings"]
    t = template("NE05E", "Subinterface VSI")
    r = request(t.id, inputs(t))
    assert "EXISTING_VSI_REQUIRED" in plan([r], {IP: observed(t)}, now=NOW).to_dict()["results"][0]["findings"]


def test_overlap_and_tampered_commands(tmp_path):
    t = template("C3750X", "Interface Description")
    a, b = request(t.id, inputs(t)), request(t.id, {**inputs(t), "description": "Different"}, row=4)
    result = plan([a, b], {IP: observed(t)}, now=NOW)
    assert all("ORDERED_RESOURCE_OVERLAP" in r["findings"] for r in result.to_dict()["results"])
    d = result.to_dict()
    d["results"][0]["commands"] = ["erase startup-config"]
    with pytest.raises(PlanError):
        BulkPlan(canonical(d))
    path = tmp_path / "plan.json"
    save(result, path)
    assert load(path).digest == result.digest
    with pytest.raises(FileExistsError):
        save(result, path)


def test_manual_no_network_workbook_and_signed_review(tmp_path, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("network forbidden")
    monkeypatch.setattr(socket, "socket", forbidden)
    from orbitflow.configuration.v2 import observation
    monkeypatch.setattr(observation, "connect_device", forbidden)
    t = template("All", "Manual CLI")
    path = tmp_path / "input.xlsx"
    create_input(path)
    wb = load_workbook(path)
    for col, name in enumerate(t.columns, 1):
        wb[t.sheet].cell(3, col, inputs(t)[name])
    wb.save(path)
    requests = read_input(path)
    spool = observation.collect(requests, {}, None, inventory_path=tmp_path / "inventory.json", spool_root=tmp_path / "spool")
    assert list(observation.observations(spool)) == []
    original = plan(requests, now=NOW)
    output = tmp_path / "review.xlsx"
    export_review(original, output)
    wb = load_workbook(output)
    wb[t.sheet].cell(3, len(t.columns) + len(output_columns(t)), "Approve")
    wb.save(output)
    approved = import_review(original, output, now=NOW)
    assert approved.to_dict()["results"][0]["commands"] == inputs(t)["commands"]
    assert not approved.to_dict()["results"][0]["deployment_ready"]
    assert approved.digest == original.digest
    assert approved.to_dict()["authenticated"] is False
    assert original.to_dict()["results"][0]["decision"] == "Pending"
    wb[t.sheet].cell(3, 1, "192.0.2.99")
    wb.save(output)
    with pytest.raises(PlanError):
        import_review(original, output, now=NOW)


@pytest.mark.parametrize("text", ["username bob password example", "snmp-server community example", "key-string example", "description secret example"])
def test_manual_secret_rejection(text):
    t = template("All", "Manual CLI")
    with pytest.raises(PlanError):
        request(t.id, {**inputs(t), "commands": text})


def test_shared_live_collection_one_session_and_read_only_commands(tmp_path, monkeypatch, capsys):
    from orbitflow.configuration.v2 import observation
    from orbitflow.transport import DeviceCredentials
    from orbitflow.config import ExecutionConfig
    from orbitflow.execution import Progress
    import sys
    t = template("C3750X", "VLAN Database")
    context = DeviceContext("one", IP, (IP,), "example", "Cisco", "cisco_ios", "C3750X", "C3750X",
                            "c3750x_switching", (), "EXAMPLE", "1", "", NOW, NOW)
    calls, commands = [], []
    session = object()
    @contextmanager
    def connect(ip, credentials, config):
        calls.append(("connect", ip))
        try:
            if ip != IP:
                raise ValueError("must never persist exception text")
            yield session
        finally:
            calls.append(("close", ip))
    class CLI:
        prompt = "example#"
        def __init__(self, supplied):
            assert supplied is session
        def __enter__(self):
            return self
        def __exit__(self, *args):
            calls.append(("cli_closed", IP))
        def require_session(self, supplied):
            assert supplied is session
        def configure(self, **kwargs):
            assert kwargs["paging_command"] == "terminal length 0"
        def run_command(self, command, **kwargs):
            commands.append(command)
            assert command in {"show interfaces description", "show running-config", "show vlan id 445", "show vlan id 4002"}
            return ("Interface  Status  Protocol  Description\nGi1/0/1  up  up  Before" if command == "show interfaces description" else
                    "vlan 445\n name Existing\n!\ninterface Gi1/0/1\n description Before\n!\nusername hidden password MUST_NOT_PERSIST")
    class Resolver:
        def __init__(self, store, **kwargs):
            pass
        def resolve(self, supplied, *, management_ip, cli):
            assert supplied is session and management_ip == IP
            calls.append(("resolve", IP))
            return context
    monkeypatch.setattr(observation, "connect_device", connect)
    monkeypatch.setattr(observation, "DeviceCLI", CLI)
    monkeypatch.setattr(observation, "DeviceInventoryResolver", Resolver)
    from functools import partial
    from orbitflow.logging import module_logger
    monkeypatch.setattr(observation, "module_logger", partial(module_logger, log_root=tmp_path / "logs"))
    rows = [request(t.id, inputs(t)), request(t.id, {**inputs(t), "vlan_id": 445}, row=4),
            request(t.id, {**inputs(t), "management_ip": "192.0.2.11"}, row=5)]
    creds = {ip: DeviceCredentials("synthetic-user", "MUST_NOT_PERSIST") for ip in (IP, "192.0.2.11")}
    spool = observation.collect(rows, creds, None, inventory_path=tmp_path / "inv.json", spool_root=tmp_path / "runs",
                                execution_config=ExecutionConfig(1, 0, 0), progress=Progress(sys.stdout))
    records = dict(observation.observations(spool))
    assert records[IP]["status"] == "live_read_only"
    assert records[IP]["identity"]["inventory_id"] == "one"
    assert records[IP]["evidence"][0]["line"] == 1
    assert records["192.0.2.11"]["status"] == "evidence_unavailable"
    assert commands == ["show interfaces description", "show running-config", "show vlan id 445", "show vlan id 4002"]
    assert calls.count(("connect", IP)) == calls.count(("resolve", IP)) == calls.count(("close", IP)) == 1
    assert "MUST_NOT_PERSIST" not in canonical(records)
    assert "must never persist" not in canonical(records)
    output = capsys.readouterr().out
    stages = ["run start", "connect", "identify", "interface collection", "VLAN/config collection",
              "audit/snapshot", "completed", "failed", "run end"]
    positions = [output.index(stage) for stage in stages]
    assert positions == sorted(positions)
    assert "MUST_NOT_PERSIST" not in output and "must never persist" not in output
    assert "completed=1; failed=1" in output
    import json
    logs = "".join(p.read_text() for p in (tmp_path / "logs").rglob("bulk_planning.log"))
    events = [json.loads(line) for line in logs.splitlines()]
    assert any(e["error_category"] == "ValueError" for e in events)
    assert all("exception_chain" not in e for e in events)
    assert "MUST_NOT_PERSIST" not in logs and "must never persist" not in logs
    assert spool.path.name in logs


def test_workbook_order_and_invalid_row_isolation(tmp_path):
    path = tmp_path / "input.xlsx"
    create_input(path)
    book = load_workbook(path)
    db = template("C3750X", "VLAN Database")
    ac = template("C3750X", "Access Port")
    for t in (db, ac):
        for col, field in enumerate(t.columns, 1):
            book[t.sheet].cell(3, col, inputs(t).get(field))
    book[db.sheet].cell(4, 1, "bad target")
    book.move_sheet(ac.sheet, offset=-1)
    book.save(path)
    rows = read_input(path)
    assert len(rows) == 2 and len(rows.errors) == 1
    assert rows[0]["template_id"] == ac.id
    result = plan(rows, {IP: observed(db)}, now=NOW)
    assert "GLOBAL_VLAN_REQUIRED" in result.to_dict()["results"][0]["findings"]
    output = tmp_path / "review.xlsx"
    export_review(result, output)
    review_book = load_workbook(output)
    assert review_book.sheetnames[:23] == book.sheetnames
    assert review_book[db.sheet]["A4"].value == "bad target"
    assert "bad target" in canonical(result.to_dict())
    assert import_review(result, output, now=NOW).digest == result.digest








def test_secret_bearing_evidence_and_fields_rejected():
    t = template("C3750X", "VLAN Database")
    r, obs = request(t.id, inputs(t)), observed(t)
    obs["evidence"][0]["excerpt"] = "password DO_NOT_PERSIST"
    with pytest.raises(PlanError):
        plan([r], {IP: obs}, now=NOW)
    obs = observed(t)
    obs["identity"]["password"] = "DO_NOT_PERSIST"
    with pytest.raises(PlanError):
        plan([r], {IP: obs}, now=NOW)


def test_planning_numeric_scope_does_not_change_audit():
    from orbitflow.vendors.configuration_facts import observe_configuration
    from orbitflow.compliance.resolvers import resolver_for
    context = DeviceContext("one", IP, (IP,), "example", "Cisco", "cisco_xr", "NCS540", "NCS540",
                            "", (), "EXAMPLE", "1", "", NOW, NOW)
    config = ("interface Gi0/0/0/1\n!\ninterface Gi0/0/0/1.10 l2transport\n encapsulation dot1q 4094\n!\n"
              "l2vpn\n bridge group G\n  bridge-domain arbitrary\n   interface Gi0/0/0/1.10")
    state = VlanState("example", IP, "cisco_xr", (), (), NOW, observe_configuration(config, "cisco_xr"))
    factory = resolver_for(context)
    before = factory(context, state).result()
    planning = factory(context, state, planning=True).result()
    assert planning["valid_database_vlans"] == [4094]
    assert before["valid_database_vlans"] == []
    assert factory(context, state).result() == before


def test_ne_tagged_vsi_planning_relationship_preserves_audit():
    from orbitflow.vendors.configuration_facts import observe_configuration
    from orbitflow.compliance.resolvers import resolver_for
    context = DeviceContext("one", IP, (IP,), "example", "Huawei", "huawei_vrp", "NE05E", "NE05E",
                            "", (), "EXAMPLE", "1", "", NOW, NOW)
    config = "vsi BLUE\n#\ninterface GE0/1\n#\ninterface GE0/1.10\n vlan-type dot1q 4094\n l2 binding vsi BLUE"
    state = VlanState("example", IP, "huawei_vrp", (), (VlanObject("vsi", "BLUE"),), NOW, observe_configuration(config, "huawei_vrp"))
    factory = resolver_for(context)
    before = factory(context, state).result()
    assert factory(context, state, planning=True).result()["valid_database_vlans"] == [4094]
    assert factory(context, state).result() == before


def test_all_template_workbook_roundtrip_preserves_input_values(tmp_path):
    path = tmp_path / "mixed.xlsx"
    create_input(path)
    book = load_workbook(path)
    for t in CATALOGUE:
        p = inputs(t)
        if "vlans" in p:
            p["vlans"] = "4002-4003"
        for col, field in enumerate(t.columns, 1):
            book[t.sheet].cell(3, col, p.get(field))
    book.save(path)
    rows = read_input(path)
    assert len(rows) == 22 and not rows.errors
    result = plan(rows, now=NOW)
    output = tmp_path / "review.xlsx"
    export_review(result, output)
    returned = load_workbook(output)
    t = template("C3750X", "Trunk Port")
    assert returned[t.sheet].cell(3, t.columns.index("vlans") + 1).value == "4002-4003"
    assert import_review(result, output, now=NOW).digest == result.digest


def test_manual_cli_operator_commands_offline(tmp_path, monkeypatch):
    from orbitflow.configuration.v2.cli import main
    def forbidden(*args, **kwargs):
        raise AssertionError("no network or credential prompting")
    monkeypatch.setattr(socket, "socket", forbidden)
    monkeypatch.setattr("orbitflow.targets.load_device_credentials", forbidden)
    source = tmp_path / "source.xlsx"
    main(["input", str(source)])
    book = load_workbook(source)
    t = template("All", "Manual CLI")
    for col, field in enumerate(t.columns, 1):
        book[t.sheet].cell(3, col, inputs(t)[field])
    book.save(source)
    output = tmp_path / "plan.json"
    main(["plan", str(source), str(output)])
    original = load(output)
    assert original.to_dict()["source"].startswith("workbook-sha256:")
    book = load_workbook(output.with_suffix(".xlsx"))
    book[t.sheet].cell(3, len(t.columns) + len(output_columns(t)), "Approve")
    book.save(output.with_suffix(".xlsx"))
    evaluation = import_review(original, output.with_suffix(".xlsx"))
    assert evaluation.to_dict()["execution_authorized"] is False
    assert load(output).digest == original.digest
    with pytest.raises(SystemExit):
        main(["review"])


def test_mixed_device_workbook_evidence_and_manual_status(tmp_path):
    templates = [template("C3750X", "VLAN Database"), template("NE05E", "VLAN Database"), template("All", "Manual CLI")]
    source = tmp_path / "input.xlsx"
    create_input(source)
    book = load_workbook(source)
    observations = {}
    for index, t in enumerate(templates):
        ip = f"192.0.2.{index + 10}"
        p = {**inputs(t), "management_ip": ip}
        for col, field in enumerate(t.columns, 1):
            book[t.sheet].cell(3, col, p.get(field))
        if t.family != "All":
            obs = observed(t, objects=[VlanObject("vlan", "4002", "Before")])
            obs["identity"].update(management_ip=ip, inventory_id=f"device-{index}", hostname=f"device-{index}", serial_number=f"EXAMPLE-{index}")
            observations[ip] = obs
    book.save(source)
    result = plan(read_input(source), observations, now=NOW)
    destination = tmp_path / "review.xlsx"
    export_review(result, destination)
    book = load_workbook(destination)
    for t in templates:
        sheet = book[t.sheet]
        status = sheet.cell(3, len(t.columns) + output_columns(t).index("Live Validation") + 1).value
        evidence = sheet.cell(3, len(t.columns) + output_columns(t).index("Existing Evidence") + 1).value
        assert status == ("not_live_validated" if t.family == "All" else "live_read_only")
        assert not evidence  # fixture has no requested VLAN stanza; unrelated interfaces stay private
    assert len(result.to_dict()["observations"]) == 2


@pytest.mark.parametrize("family", ["C3750X", "C3850X"])
@pytest.mark.parametrize("action,fragment", [("Add", "add "), ("Remove", "remove "), ("Replace", "")])
def test_trunk_actions(family, action, fragment):
    t = template(family, "Trunk Port")
    p = {**inputs(t), "action": action, "native_vlan": 445}
    result = plan([request(t.id, p)], {IP: observed(t)}, now=NOW).to_dict()["results"][0]
    assert f"switchport trunk allowed vlan {fragment}4002" in result["commands"]
    assert "switchport trunk native vlan 445" in result["commands"]
    assert ("ADD_REQUIRES_EXPLICIT_BASELINE" in result["findings"]) == (action == "Add")




def test_vendor_secret_marker_sanitized_before_spooling():
    t = template("NE05E", "VLAN Database")
    obs = observed(t, objects=[VlanObject("vlan", "4002", "cipher UNDISCLOSED")])
    assert "UNDISCLOSED" not in canonical(obs)
    assert "[REDACTED]" in canonical(obs)


def configuration_observation(t, config, *, objects=()):
    from orbitflow.vendors.configuration_facts import observe_configuration
    obs = observed(t, objects=objects)
    obs["facts"] = safe_data([asdict(n) for n in observe_configuration(config, FAMILIES[t.family][0])])
    obs["evidence"][0]["excerpt"] = config
    return obs


@pytest.mark.parametrize("enc,tags,command", [
    ("dot1q", {"outer_vlan": 445}, "encapsulation dot1q 445"),
    ("QinQ", {"outer_vlan": 445, "inner_vlan": 545}, "encapsulation dot1q 445 second-dot1q 545"),
    ("untagged", {}, "encapsulation untagged"),
])
def test_ncs_encapsulation_review_and_noop(enc, tags, command, tmp_path):
    t = template("NCS540", "L2VPN")
    p = {k: v for k, v in inputs(t).items() if k != "outer_vlan"}
    p.update(encapsulation=enc, **tags)
    config = f"interface Gi1/0/1.10 l2transport\n {command}\n!\nl2vpn\n bridge group CUSTOMER\n  bridge-domain BLUE\n   interface Gi1/0/1.10"
    obs = configuration_observation(t, config)
    r = request(t.id, p)
    result = plan([r], {IP: obs}, now=NOW)
    assert result.to_dict()["results"][0]["no_op"]
    assert result.to_dict()["results"][0]["commands"] == []
    obs = observed(t)
    result = plan([r], {IP: obs}, now=NOW)
    assert command in result.to_dict()["results"][0]["commands"]
    review_path = tmp_path / "review.xlsx"
    export_review(result, review_path)
    assert import_review(result, review_path, now=NOW).digest == result.digest
    # A competing subinterface claiming this ingress is reviewable but blocked.
    obs = configuration_observation(t, config.replace(".10", ".11"))
    row = plan([r], {IP: obs}, now=NOW, decisions={r["row_id"]: "Approve"}).to_dict()["results"][0]
    assert "INGRESS_TAG_OVERLAP" in row["findings"] and not row["deployment_ready"]


@pytest.mark.parametrize("enc,tags", [
    ("untagged", {"outer_vlan": 445}), ("dot1q", {}),
    ("dot1q", {"outer_vlan": 445, "inner_vlan": 545}),
    ("QinQ", {"outer_vlan": 445}),
])
def test_ncs_conditional_vlan_validation(enc, tags):
    t = template("NCS540", "L2VPN")
    p = {k: v for k, v in inputs(t).items() if k != "outer_vlan"}
    with pytest.raises(PlanError):
        request(t.id, {**p, "encapsulation": enc, **tags})






def edge_observation():
    t = template("EdgeSwitch", "VLAN Port")
    obs = observed(t, objects=[VlanObject("vlan", str(v)) for v in (445, 545, 645)])
    obs["vlans"] = [dict(interface_name="0/1", mode="hybrid", tagged_vlans=[545], referenced_vlans=[445, 545], pvid=445)]
    obs["audit"][0].update(configured_membership_vlans=[445, 545], configured_tagged_vlans=[545])
    return obs








@pytest.mark.parametrize("case", ["ncs_untagged", "ncs_qinq", "asr_push", "asr_none", "asr_pop2", "edge_pvid", "edge_tagged", "c3850_name"])
def test_revised_options_input_review_roundtrip(case, tmp_path):
    if case.startswith("ncs"):
        t = template("NCS540", "L2VPN")
        p = {k: v for k, v in inputs(t).items() if k != "outer_vlan"}
        p.update(encapsulation="untagged" if case.endswith("untagged") else "QinQ")
        if p["encapsulation"] == "QinQ":
            p.update(outer_vlan=445, inner_vlan=545)
    elif case.startswith("asr"):
        t = template("ASR920", "EVC")
        p = {k: v for k, v in inputs(t).items() if k not in {"outer_vlan", "pop_tags"}}
        if case == "asr_push":
            p.update(encapsulation="Untagged", rewrite="Push dot1q symmetric", push_vlan=445)
        elif case == "asr_pop2":
            p.update(encapsulation="QinQ", outer_vlan="445", inner_vlan="545", rewrite="Pop 2 symmetric")
        else:
            p.update(outer_vlan="445-446", rewrite="None")
    elif case.startswith("edge"):
        t = template("EdgeSwitch", "VLAN Port")
        p = dict(management_ip=IP, interface="0/1", action="Add/Change", pvid=645)
        if case.endswith("tagged"):
            p.update(tagged_vlans="645")
    else:
        t = template("C3850X", "VLAN Database")
        p = {**inputs(t), "name": "Customer_Site_Alpha"}
    source = tmp_path / "input.xlsx"
    create_input(source)
    book = load_workbook(source)
    for col, field in enumerate(t.columns, 1):
        book[t.sheet].cell(3, col, p.get(field))
    book.save(source)
    rows = read_input(source)
    assert len(rows) == 1 and not rows.errors
    obs = edge_observation() if t.family == "EdgeSwitch" else observed(t)
    result = plan(rows, {IP: obs}, now=NOW)
    output = tmp_path / "review.xlsx"
    export_review(result, output)
    book = load_workbook(output)
    sheet = book[t.sheet]
    for col, field in enumerate(t.columns, 1):
        assert sheet.cell(3, col).value == (str(p[field]) if field in p else None)
    assert sheet.cell(3, len(t.columns) + output_columns(t).index("Existing Evidence") + 1).value == (None if case == "c3850_name" else obs["evidence"][0]["excerpt"])
    assert sheet.cell(3, len(t.columns) + output_columns(t).index("Proposed CLI") + 1).value
    sheet.cell(3, len(t.columns) + len(output_columns(t)), "Approve")
    book.save(output)
    approved = import_review(result, output, now=NOW)
    assert approved.digest == result.digest
    assert not approved.to_dict()["results"][0]["execution_authorized"]
    # Approval cannot turn failed evidence into readiness, even for valid inputs.
    unavailable = plan(rows, now=NOW)
    output = tmp_path / "unavailable.xlsx"
    export_review(unavailable, output)
    book = load_workbook(output)
    book[t.sheet].cell(3, len(t.columns) + len(output_columns(t)), "Approve")
    book.save(output)
    assert not import_review(unavailable, output, now=NOW).to_dict()["results"][0]["deployment_ready"]


def test_c3850_valid_name_rename_and_noop():
    t = template("C3850X", "VLAN Database")
    p = {**inputs(t), "name": "Customer_Site_Alpha"}
    r = request(t.id, p)
    obs = observed(t, objects=[VlanObject("vlan", "4002", "Before")])
    row = plan([r], {IP: obs}, now=NOW).to_dict()["results"][0]
    assert "name Customer_Site_Alpha" in row["commands"]
    assert "VLAN_RENAME_IMPACT" in row["findings"]
    obs = observed(t, objects=[VlanObject("vlan", "4002", p["name"])])
    assert plan([r], {IP: obs}, now=NOW).to_dict()["results"][0]["no_op"]






def test_revised_cli_review_and_digest_invalidation(tmp_path, monkeypatch):
    from orbitflow.configuration.v2.cli import main
    monkeypatch.setattr(socket, "socket", lambda *a, **kw: pytest.fail("review connected"))
    t = template("ASR920", "EVC")
    p = {k: v for k, v in inputs(t).items() if k not in {"outer_vlan", "pop_tags"}}
    p.update(encapsulation="Untagged", rewrite="Push dot1q symmetric", push_vlan=445)
    r = request(t.id, p)
    result = plan([r], {IP: observed(t)}, now=NOW)
    source, workbook, output = tmp_path / "plan.json", tmp_path / "review.xlsx", tmp_path / "reviewed.json"
    save(result, source)
    export_review(result, workbook)
    book = load_workbook(workbook)
    book[t.sheet].cell(3, len(t.columns) + len(output_columns(t)), "Reject")
    book.save(workbook)
    revised = import_review(load(source), workbook, now=NOW)
    assert revised.digest == result.digest
    assert not output.exists()
    assert revised.to_dict()["results"][0]["decision"] == "Reject"
    assert "rewrite ingress tag push dot1q 445 symmetric" in revised.to_dict()["results"][0]["commands"]
    assert not revised.to_dict()["results"][0]["deployment_ready"]


@pytest.mark.parametrize("kind", ["interface", "l2_interface"])
@pytest.mark.parametrize("description", ["Before", "Customer Site Alpha"])
def test_xr_description_uses_observed_declaration(kind, description):
    t = template("NCS540", "Interface Description")
    name = "GigabitEthernet0/0/0/1.545"
    obs = observed(t)
    obs["interfaces"] = []
    obs["facts"] = [asdict(ConfigFact(kind, name, "running-config", 7,
        f"interface {name}" + (" l2transport" if kind == "l2_interface" else ""),
        (ConfigFact("description", description, "running-config", 8, f" description {description}"),)))]
    obs = safe_data(obs, lambda value: value)
    r = request(t.id, {**inputs(t), "interface": name})
    result = plan([r], {IP: obs}, now=NOW).to_dict()["results"][0]
    assert "INTERFACE_ABSENT_OR_AMBIGUOUS" not in result["findings"]
    assert result["no_op"] == (description == "Customer Site Alpha")
    if not result["no_op"]:
        assert result["commands"][1] == f"interface {name}" + (" l2transport" if kind == "l2_interface" else "")
    with pytest.raises(PlanError):
        request(t.id, {**inputs(t), "interface": name + " l2transport"})
    for facts in ([], obs["facts"] * 2):
        invalid = {**obs, "facts": facts}
        row = plan([r], {IP: invalid}, now=NOW, decisions={r["row_id"]: "Approve"}).to_dict()["results"][0]
        assert "INTERFACE_DECLARATION_ABSENT_OR_AMBIGUOUS" in row["findings"]
        assert not row["deployment_ready"]


def test_asr_bdi_requires_interface_declaration():
    t = template("ASR920", "Interface Description")
    r = request(t.id, {**inputs(t), "interface": "BDI545"})
    obs = observed(t)
    obs["interfaces"] = []
    obs["facts"] = [asdict(ConfigFact("interface", "BDI545", "running-config", 12, "interface BDI545"))]
    obs = safe_data(obs, lambda value: value)
    assert "INTERFACE_ABSENT_OR_AMBIGUOUS" not in plan([r], {IP: obs}, now=NOW).to_dict()["results"][0]["findings"]
    obs["facts"] = [asdict(ConfigFact("bridge_domain", "545", "running-config", 12, "bridge-domain 545"))]
    obs = safe_data(obs, lambda value: value)
    assert "INTERFACE_ABSENT_OR_AMBIGUOUS" in plan([r], {IP: obs}, now=NOW).to_dict()["results"][0]["findings"]


def test_scoped_evidence_service_dependencies_and_workbook(tmp_path):
    from orbitflow.configuration.v2.evidence import scoped_evidence
    t = template("ME3600X", "EVC")
    r = request(t.id, inputs(t))
    obs = observed(t)
    def fact(kind, value, line, children=()):
        return ConfigFact(kind, value, "running-config", line, f"{kind} {value}", children)
    obs["facts"] = [asdict(fact("interface", "Gi1/0/1", 10, (
        fact("service_instance", "10", 11, (fact("bridge_domain", "4002", 12),)),
        fact("service_instance", "20", 13)))),
        asdict(fact("interface", "Gi1/0/2", 20)),
        asdict(fact("vlan", (4002,), 30)), asdict(fact("vlan", (999,), 40))]
    obs = safe_data(obs, lambda value: value)
    selected = scoped_evidence(r, obs)
    assert [e["line"] for e in selected] == [10, 11, 12, 30]
    result = plan([r], {IP: obs}, now=NOW)
    output = tmp_path / "scoped.xlsx"
    export_review(result, output)
    book = load_workbook(output)
    assert [book["Evidence Details"].cell(i, 3).value for i in range(2, 6)] == ["10", "11", "12", "30"]
    assert import_review(result, output, now=NOW).digest == result.digest
    db = template("ME3600X", "VLAN Database")
    assert [e["line"] for e in scoped_evidence(request(db.id, inputs(db)), obs)] == [30]


def test_shared_excel_per_ip_credentials_and_provider_precedence(tmp_path):
    from openpyxl import Workbook
    from orbitflow.targets import load_device_credentials
    from orbitflow.transport import DeviceCredentials
    book = Workbook()
    book.active.append(["management_ip", "username", "password", "secret"])
    book.active.append([IP, "user-a", "dummy-a", "enable-a"])
    book.active.append(["192.0.2.11", "user-b", "dummy-b", None])
    book.active.append(["192.0.2.12", "dup", "dummy", None])
    book.active.append(["192.0.2.12", "dup", "dummy", None])
    book.active.append(["192.0.2.13", "missing", None, None])
    path = tmp_path / "devices.xlsx"
    book.save(path)
    ips = [IP, "192.0.2.11", "192.0.2.12", "192.0.2.13", "192.0.2.14"]
    credentials = load_device_credentials(path, ips)
    assert set(credentials) == {IP, "192.0.2.11"}
    assert credentials[IP].secret == "enable-a"
    assert credentials["192.0.2.11"].username == "user-b"
    supplied = DeviceCredentials("provider", "provider-value")
    assert load_device_credentials(path, ips, provider=lambda ip: supplied if ip == IP else None)[IP] is supplied



def test_cli_devices_credentials_and_shared_sequential_config(tmp_path, monkeypatch):
    from orbitflow.configuration.v2 import cli, observation
    from openpyxl import Workbook
    source = tmp_path / "mixed.xlsx"
    create_input(source)
    book = load_workbook(source)
    families = ["NCS540", "ASR920", "NE05E"]
    for index, family in enumerate(families):
        t = template(family, "Interface Description")
        p = {**inputs(t), "management_ip": f"192.0.2.{10 + index}"}
        for col, field in enumerate(t.columns, 1):
            book[t.sheet].cell(3, col, p.get(field))
    manual = template("All", "Manual CLI")
    for col, field in enumerate(manual.columns, 1):
        book[manual.sheet].cell(3, col, {**inputs(manual), "management_ip": "192.0.2.99"}.get(field))
    book.save(source)
    devices = Workbook()
    devices.active.append(["management_ip", "username", "password", "secret"])
    devices.active.append([IP, "user-a", "dummy-a", "enable-a"])
    devices.active.append(["192.0.2.11", "user-b", "dummy-b", None])
    credentials_path = tmp_path / "devices.xlsx"
    devices.save(credentials_path)
    captured = {}
    def collect(rows, credentials, transport, **kwargs):
        captured.update(credentials=credentials, config=kwargs["execution_config"])
        return object()
    monkeypatch.setattr(observation, "collect", collect)
    monkeypatch.setattr(observation, "bounded_observations", lambda spool: {})
    output = tmp_path / "plan.json"
    cli.main(["plan", str(source), str(output), "--devices", str(credentials_path),
              "--max-concurrent-devices", "1", "--proxy", "example", "--cluster", "example",
              "--bastion-host", "example", "--bastion-user", "example"])
    assert captured["config"].max_concurrent_devices == 1
    assert set(captured["credentials"]) == {IP, "192.0.2.11"}
    assert captured["credentials"][IP].secret == "enable-a"
    assert captured["credentials"]["192.0.2.11"].password == "dummy-b"
    assert "dummy-a" not in output.read_text() and "enable-a" not in output.read_text()


def test_missing_credentials_never_connect_and_spool_is_safe(tmp_path, monkeypatch):
    from orbitflow.configuration.v2 import observation
    from orbitflow.config import ExecutionConfig
    def forbidden(*args, **kwargs):
        raise AssertionError("no credentials means no connection")
    monkeypatch.setattr(observation, "connect_device", forbidden)
    t = template("ASR920", "Interface Description")
    spool = observation.collect([request(t.id, inputs(t))], {}, None,
        inventory_path=tmp_path / "inventory.json", spool_root=tmp_path / "runs",
        execution_config=ExecutionConfig(1, 0, 0))
    assert dict(observation.observations(spool)) == {IP: {"status": "evidence_unavailable"}}



def test_xr_service_evidence_distinguishes_bridge_groups():
    from orbitflow.configuration.v2.evidence import scoped_evidence
    from orbitflow.vendors.configuration_facts import observe_configuration
    t = template("NCS540", "L2VPN")
    r = request(t.id, inputs(t))
    configuration = """interface Gi1/0/1
 description parent
!
interface Gi1/0/1.10 l2transport
 encapsulation dot1q 4002
!
interface Gi1/0/2.20 l2transport
 encapsulation dot1q 999
!
l2vpn
 bridge group CUSTOMER
  bridge-domain BLUE
   interface Gi1/0/1.10
  !
 !
 bridge group OTHER
  bridge-domain BLUE
   interface Gi1/0/2.20
  !
 !
!
"""
    obs = observed(t)
    obs["facts"] = safe_data([asdict(n) for n in observe_configuration(configuration, "cisco_xr")], lambda value: value)
    excerpts = "\n".join(e["excerpt"] for e in scoped_evidence(r, obs))
    assert "bridge group CUSTOMER" in excerpts and "Gi1/0/1.10" in excerpts
    assert "OTHER" not in excerpts and "Gi1/0/2" not in excerpts




def test_plan_export_failure_logs_safe_category(tmp_path, monkeypatch, capsys):
    from orbitflow.configuration.v2 import cli
    from orbitflow.logging import module_logger
    source = tmp_path / "input.xlsx"
    create_input(source)
    book = load_workbook(source)
    t = template("All", "Manual CLI")
    for k, v in inputs(t).items():
        book[t.sheet].cell(3, t.columns.index(k) + 1, v)
    book.save(source)
    monkeypatch.setattr(cli, "export_plan", lambda *a: (_ for _ in ()).throw(OSError("password=NEVER_ECHO")))
    with module_logger("test", log_root=tmp_path / "logs"):
        with pytest.raises(SystemExit):
            cli.main(["plan", str(source), str(tmp_path / "plan.json")])
    output = capsys.readouterr()
    logs = "".join(p.read_text() for p in (tmp_path / "logs").rglob("*.log"))
    assert "OSError" in logs and "stage=workbook export" in logs and "Export start" in logs
    assert "Planning failed (OSError)" in output.err
    assert "NEVER_ECHO" not in logs + output.out + output.err


def test_real_edgeswitch_audit_warning_description_only():
    t = template("EdgeSwitch", "Interface Description")
    context = DeviceContext("one", IP, (IP,), "example", "Ubiquiti", "ubiquiti_edgeswitch",
                            "EdgeSwitch", "EdgeSwitch", FAMILIES[t.family][1], (), "EXAMPLE", "1", "", NOW, NOW)
    interface = InterfaceRecord("example", IP, context.platform, "Gi1/0/1", "Before", "up", "up", NOW)
    def fact(kind, value, line, text, children=()):
        return ConfigFact(kind, value, "running-config", line, text, children)
    facts = (fact("vlan_database", "", 1, "vlan database", (
                fact("vlan", (445,), 2, "vlan 445"),)),
             fact("interface", interface.port_name, 3, "interface Gi1/0/1", (
                fact("include", (445,), 4, "vlan participation include 445"),
                fact("tagged", (445, 446), 5, "vlan tagging 445,446"),
                fact("methodology_evidence", "stanza_end", 6, "exit"))))
    state = VlanState("example", IP, context.platform, (), (), NOW, facts)
    obs = snapshot(context, (interface,), state)
    assert obs["audit"][0]["configuration_findings"][0]["code"] == "TAGGED_VLAN_NOT_IN_MEMBERSHIP"
    r = request(t.id, inputs(t))
    draft = plan([r], {IP: obs}, now=NOW)
    approved = review(draft, {r["row_id"]: "Approve"}, expected_digest=draft.digest, now=NOW)
    assert approved.to_dict()["results"][0]["deployment_ready"]
    assert "446" in canonical(approved.to_dict()["results"][0]["findings"])
    obs["facts"][1]['children'].append(dict(kind='description', value=inputs(t)['description'],
        source_filename='running-config', line=6, excerpt='description "' + inputs(t)['description'] + '"', children=[]))
    noop = plan([r], {IP: obs}, now=NOW)
    result = review(noop, {r["row_id"]: "Approve"}, expected_digest=noop.digest, now=NOW).to_dict()["results"][0]
    assert result["no_op"] and result["commands"] == [] and result["deployment_ready"]


def test_progress_flushes_each_increment():
    from orbitflow.execution import Progress
    class Output:
        def __init__(self):
            self.flushes = 0
            self.text = ""
        def write(self, value):
            self.text += value
        def flush(self):
            self.flushes += 1
    output = Output()
    progress = Progress(output)
    progress("connect")
    assert output.flushes == 1 and output.text == "connect\n"
    progress("identify")
    assert output.flushes == 2


def test_observation_run_failure_is_logged_and_safe(tmp_path, monkeypatch):
    from functools import partial
    from orbitflow.configuration.v2 import observation
    from orbitflow.logging import module_logger
    monkeypatch.setattr(observation, "module_logger", partial(module_logger, log_root=tmp_path / "logs"))
    def fail(*a, **kw):
        raise OSError("token=DO_NOT_OUTPUT")
    monkeypatch.setattr(observation, "execute_devices", fail)
    t = template("C3750X", "VLAN Database")
    with pytest.raises(OSError):
        observation.collect([request(t.id, inputs(t))], {}, None,
                            inventory_path=tmp_path / "inv.json", spool_root=tmp_path / "spools")
    logs = "".join(p.read_text() for p in (tmp_path / "logs").rglob("*.log"))
    assert "ObservationRunError" in logs and "run end" in logs
    assert "DO_NOT_OUTPUT" not in logs

@pytest.mark.parametrize("operation,category,stage", [
    ("read_input", ValueError, "workbook input"),
    ("plan", PlanError, "planning"),
    ("save", PermissionError, "plan export"),
    ("export_plan", OSError, "workbook export"),
])
def test_cli_failure_stage_and_class_without_exception_text(tmp_path, monkeypatch, capsys, operation, category, stage):
    from functools import partial
    from orbitflow.configuration.v2 import cli
    from orbitflow.logging import module_logger
    source = tmp_path / "input.xlsx"
    create_input(source)
    book = load_workbook(source)
    t = template("All", "Manual CLI")
    for k, v in inputs(t).items():
        book[t.sheet].cell(3, t.columns.index(k) + 1, v)
    book.save(source)
    def fail(*args, **kwargs):
        raise category("password=NEVER_ECHO /private/token=SECRET_PATH raw-device-output")
    monkeypatch.setattr(cli, operation, fail)
    monkeypatch.setattr(cli, "module_logger", partial(module_logger, log_root=tmp_path / "logs"))
    with pytest.raises(SystemExit) as caught:
        cli.main(["plan", str(source), str(tmp_path / "plan.json")])
    assert caught.value.code == 2
    output = capsys.readouterr()
    logs = "".join(p.read_text() for p in (tmp_path / "logs").rglob("*.log"))
    assert f"stage={stage}" in logs and category.__name__ in logs
    assert f"Planning failed ({category.__name__})" in output.err
    for sensitive in ("NEVER_ECHO", "SECRET_PATH", "raw-device-output", "exception_chain"):
        assert sensitive not in logs + output.out + output.err


def test_observation_preserves_distinct_shared_failure_categories(tmp_path, monkeypatch):
    from contextlib import contextmanager
    from functools import partial
    import json
    from orbitflow.configuration.v2 import observation
    from orbitflow.config import ExecutionConfig
    from orbitflow.logging import module_logger
    class DeviceConnectionError(Exception):
        pass
    class InventoryCollectionError(Exception):
        pass
    errors = {IP: DeviceConnectionError, "192.0.2.11": InventoryCollectionError}
    @contextmanager
    def connect(ip, *args):
        raise errors[ip]("password=NEVER_ECHO /private/token=SECRET_PATH raw-device-output")
        yield  # pragma: no cover
    monkeypatch.setattr(observation, "connect_device", connect)
    monkeypatch.setattr(observation, "module_logger", partial(module_logger, log_root=tmp_path / "logs"))
    t = template("C3750X", "VLAN Database")
    rows = [request(t.id, {**inputs(t), "management_ip": ip}, row=3 + i) for i, ip in enumerate(errors)]
    output = []
    from orbitflow.transport import DeviceCredentials
    spool = observation.collect(rows, {ip: DeviceCredentials("test", "synthetic") for ip in errors}, None,
        inventory_path=tmp_path / "inventory.json", spool_root=tmp_path / "runs",
        execution_config=ExecutionConfig(1, 0, 0), progress=output.append)
    records = list(spool.records())
    assert [r["target"] for r in records] == list(errors)
    assert [r["error_category"] for r in records] == [e.__name__ for e in errors.values()]
    assert all(r["payload"]["status"] == "evidence_unavailable" for r in records)
    logs = "".join(p.read_text() for p in (tmp_path / "logs").rglob("*.log"))
    events = [json.loads(line) for line in logs.splitlines()]
    for ip, error in errors.items():
        assert any(e["error_category"] == error.__name__ and e["management_ip"] == ip for e in events)
        assert any(f"{ip}: failed" in line and error.__name__ in line for line in output)
    assert "completed=0; failed=2" in output[-1]
    for sensitive in ("NEVER_ECHO", "SECRET_PATH", "raw-device-output", "exception_chain"):
        assert sensitive not in logs + "".join(output) + canonical(records)


@pytest.mark.parametrize("destination", ["default", "relative", "absolute"])
def test_plan_output_resolution_and_no_overwrite(tmp_path, monkeypatch, destination):
    from orbitflow.configuration.v2.cli import main
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(socket, "socket", lambda *a, **kw: pytest.fail("network forbidden"))
    source = tmp_path / "input.xlsx"
    create_input(source)
    book = load_workbook(source)
    t = template("All", "Manual CLI")
    for col, field in enumerate(t.columns, 1):
        book[t.sheet].cell(3, col, inputs(t)[field])
    book.save(source)
    argv = ["plan", str(source)]
    if destination != "default":
        output = Path("chosen/plan.json") if destination == "relative" else tmp_path / "absolute/plan.json"
        argv.append(str(output))
    main(argv)
    if destination == "default":
        output, = Path("outputs/reports/configuration").glob("*/plan.json")
        assert output.parent.name.endswith("_Manual-only")
    before = output.read_bytes()
    workbook = output.with_suffix(".xlsx")
    assert workbook.is_file()
    with pytest.raises(SystemExit):
        main(["plan", str(source), str(output)])
    assert output.read_bytes() == before
    output.unlink()
    with pytest.raises(SystemExit):
        main(["plan", str(source), str(output)])
    assert not output.exists()  # Companion collision also blocks JSON creation.


@pytest.mark.parametrize("edit", ["digest", "cli", "evidence", "order", "formula", "invalid_decision"])
def test_workbook_decision_integrity(tmp_path, edit):
    t = template("C3750X", "VLAN Database")
    original = plan([request(t.id, inputs(t))], {IP: observed(t)}, now=NOW)
    path = tmp_path / "plan.xlsx"
    export_review(original, path)
    book = load_workbook(path)
    sheet = book[t.sheet]
    if edit == "digest":
        book["Introduction"]["B1"] = "0" * 64
    elif edit == "order":
        book.move_sheet(sheet, offset=-1)
    else:
        field = {"cli": "Proposed CLI", "evidence": "Existing Evidence",
                 "formula": "Engineer Decision", "invalid_decision": "Engineer Decision"}[edit]
        sheet.cell(3, len(t.columns) + output_columns(t).index(field) + 1,
                   "=1+1" if edit == "formula" else "changed")
    book.save(path)
    with pytest.raises(PlanError):
        consume_decisions(original, path, expected_digest=original.digest, now=NOW)
    with pytest.raises(PlanError):
        consume_decisions(original, path, expected_digest="0" * 64, now=NOW)


@pytest.mark.parametrize("family,function", [("C3750X","Access Port"), ("C3850X","Access Port"), ("ME3600X","EVC"), ("ASR920","EVC"), ("EdgeSwitch","VLAN Port"), ("NCS540","L2VPN"), ("NE05E","Subinterface VSI")])
@pytest.mark.parametrize("description_first", [True, False])
def test_description_overlap_all_families_preserves_safety(family, function, description_first):
    operation, desc = template(family, function), template(family, "Interface Description")
    p = inputs(operation)
    d = dict(management_ip=IP, interface=p['interface'], description='Final description')
    ordered = [(desc,d), (operation,p)] if description_first else [(operation,p), (desc,d)]
    requests = [request(t.id, values, sheet_order=i+1) for i,(t,values) in enumerate(ordered)]
    # Unavailable evidence remains a blocker independently of a permitted overlap.
    rows = plan(requests, {IP: dict(status='evidence_unavailable')}, now=NOW).to_dict()['results']
    assert all('ORDERED_RESOURCE_OVERLAP_WARNING' in r['warnings'] for r in rows)
    assert all('EVIDENCE_UNAVAILABLE' in r['findings'] and not r['deployment_ready'] for r in rows)
    assert all('ORDERED_RESOURCE_OVERLAP' not in r['findings'] for r in rows)
    assert all(not r['execution_authorized'] for r in rows)

@pytest.mark.parametrize('families,addresses,descriptor', [
    (['C3750X'], [IP], 'C3750X-192-0-2-10'),
    (['C3750X', 'C3850X'], [IP, '192.0.2.11'], 'multi-device-2'),
    (['C3750X', 'C3750X'], [IP, '192.0.2.11'], 'multi-device-2'),
    (['All'], [IP], 'Manual-only'),
    (['C3750X'], ['../../evil\\path\n:<>|?*'], 'C3750X-evil-path'),
])
def test_readable_output_reservation(tmp_path, families, addresses, descriptor):
    from orbitflow.configuration.v2.cli import _default_output
    rows = [{'template_id': next(t.id for t in CATALOGUE if t.family == f),
             'inputs': {'management_ip': ip}} for f, ip in zip(families, addresses)]
    root = tmp_path / 'reports'
    first = _default_output(rows, now=NOW, root=root)
    assert first == root / ('20261011-000000_' + descriptor) / 'plan.json'
    first.write_text('preserved')
    second = _default_output(rows, now=NOW, root=root)
    assert second.parent.name == first.parent.name + '-01'
    assert first.read_text() == 'preserved'
    assert second.parent.parent == root


@pytest.mark.parametrize('explicit', [False, True])
def test_cli_report_pair_preserves_digest_and_override(tmp_path, monkeypatch, explicit):
    from orbitflow.configuration.v2 import cli
    import logging
    t = template('All', 'Manual CLI')
    class Rows(list):
        source_digest = 'a' * 64
    rows = Rows([request(t.id, inputs(t))])
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(cli, 'read_input', lambda _: rows)
    args = ['plan', 'input.xlsx']
    if explicit:
        args.append('custom/review.json')
    cli._main(args, logging.getLogger('test'), lambda _: None)
    path = tmp_path / 'custom/review.json' if explicit else next(tmp_path.glob('outputs/reports/configuration/*/plan.json'))
    original = load(path)
    assert len(original.digest) == 64
    book = load_workbook(path.with_suffix('.xlsx'))
    assert book['Introduction']['B1'].value == original.digest
    consume_decisions(original, path.with_suffix('.xlsx'), expected_digest=original.digest)
    before = path.read_bytes(), path.with_suffix('.xlsx').read_bytes()
    with pytest.raises((PlanError, FileExistsError)):
        save(original, path)
    with pytest.raises(PlanError):
        export_review(original, path.with_suffix('.xlsx'))
    if explicit:
        with pytest.raises(SystemExit):
            cli._main(args, logging.getLogger('test'), lambda _: None)
    assert before == (path.read_bytes(), path.with_suffix('.xlsx').read_bytes())
    assert not (tmp_path / 'outputs/runs/configuration').exists()


@pytest.mark.parametrize('t', [t for t in CATALOGUE if t.family not in {'All', 'EdgeSwitch'}
    and t.function != 'VLAN Database'], ids=lambda t:t.id)
@pytest.mark.parametrize('decision', ['Pending', 'Approve', 'Reject'])
def test_issue_152_non_ubiquiti_database_rows_do_not_create_dependencies(t, decision):
    # Families without a database worksheet cannot acquire a database dependency.
    databases = [q for q in CATALOGUE if q.family == t.family and q.function == 'VLAN Database']
    target = request(t.id, inputs(t), sheet_order=1)
    requests = [target]
    decisions = {target['row_id']:'Approve'}
    if databases:
        db = request(databases[0].id, inputs(databases[0]))
        requests.insert(0, db)
        decisions[db['row_id']] = decision
    row = plan(requests, {IP:observed(t)}, now=NOW, decisions=decisions).to_dict()['results'][-1]
    assert not row['dependencies']
    assert 'DEPENDENCY_NOT_APPROVED' not in row['findings']
    assert not row['execution_authorized']
    # Unrelated provisional-template/observation blockers are intentionally retained.
