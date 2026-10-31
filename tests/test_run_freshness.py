"""Normal runs collect anew; explicitly retained exports remain historical."""
from datetime import datetime
from io import StringIO

import pytest
from openpyxl import load_workbook

from orbitflow.config import ExecutionConfig
from orbitflow.result_spool import ResultSpool
from orbitflow.targets import load_targets


def isolated(first, second):
    assert first.path != second.path
    assert first.manifest['run_id'] != second.manifest['run_id']
    assert first.manifest['created_at'] < second.manifest['created_at']
    old, new = list(first.records()), list(second.records())
    assert [r['input_position'] for r in old] == [r['input_position'] for r in new]
    assert all(r['run_id'] == second.manifest['run_id'] for r in new)
    assert max(datetime.fromisoformat(r['completed_at']) for r in old) < min(
        datetime.fromisoformat(r['completed_at']) for r in new)
    return old, new


@pytest.mark.parametrize('failed', [False, True])
def test_interface_vlan_excel_runs_collect_again(tmp_path, monkeypatch, failed):
    import test_device_reporting as fixture
    import test_inventory_refresh as excel
    from orbitflow import reporting

    source = excel.input_file(tmp_path, [
        ['192.0.2.1', 'synthetic-user', 'synthetic-password']])
    original = source.read_bytes()
    export = reporting.export_report_spool
    paths = []
    def retain(spool, path):
        paths.append(spool)
        return export(spool, path, cleanup=False)
    monkeypatch.setattr(reporting, 'export_report_spool', retain)
    kwargs = dict(inventory_path=tmp_path / 'inventory.json', reports_dir=tmp_path,
                  spool_root=tmp_path / 'runs', log_root=tmp_path / 'logs',
                  output=StringIO(), execution_config=ExecutionConfig(1, 0, 0))
    first_events, _ = fixture.install_fakes(monkeypatch)
    reporting.run_report(load_targets(source), fixture.CONFIG, **kwargs)
    second_events, _ = fixture.install_fakes(monkeypatch, 'connect' if failed else None)
    report = reporting.run_report(load_targets(source), fixture.CONFIG, **kwargs)
    assert ('connect', '192.0.2.1') in first_events
    assert ('connect', '192.0.2.1') in second_events
    old, new = isolated(*(ResultSpool(p) for p in paths))
    assert old[0]['payload'][0]
    assert bool(new[0]['payload'][0]) is (not failed)
    assert new[0]['status'] == ('failed' if failed else 'success')
    book = load_workbook(report)
    try:
        assert book.sheetnames == ['Interfaces', 'VLAN_Database', 'Run_Errors']
        assert tuple(next(book['Interfaces'].values)) == reporting.INTERFACE_COLUMNS
        assert book['Interfaces'].max_row == (1 if failed else 2)
    finally:
        book.close()
    before = list(second_events)
    export(paths[0], tmp_path / 'historical.xlsx', cleanup=False)
    assert second_events == before
    assert source.read_bytes() == original


@pytest.mark.parametrize('failed', [False, True])
def test_compliance_runs_do_not_borrow_prior_success(tmp_path, monkeypatch, failed):
    import test_vlan_compliance as fixture
    from orbitflow import vlan_compliance as app
    first_events, _ = fixture.install_fakes(monkeypatch)
    first = app.collect_compliance(fixture.targets(1), object(), **fixture.options(tmp_path))
    second_events, _ = fixture.install_fakes(monkeypatch, 'connect' if failed else None)
    second = app.collect_compliance(fixture.targets(1), object(), **fixture.options(tmp_path))
    old, new = isolated(ResultSpool(first.spool_path), ResultSpool(second.spool_path))
    assert first_events[0] == second_events[0] == ('192.0.2.1', 'connect')
    assert all(f['status'] == 'compliant' for f in old[0]['payload']['findings'])
    assert all(f['status'] == ('unable_to_assess' if failed else 'compliant')
               for f in new[0]['payload']['findings'])
    assert second.failed_devices == int(failed)


@pytest.mark.parametrize('failed', [False, True])
def test_mac_runs_and_explicit_export_are_separate(tmp_path, monkeypatch, failed):
    from orbitflow import mac_lookup as app
    calls = []
    def collect(target, *args, **kwargs):
        calls.append(target['management_ip'])
        if failed and len(calls) == 2:
            raise OSError('unreachable')
        return [[['192.0.2.1', 'router', 'C3750X', 'Gi1/1', 'trunk',
                  'aa:bb:cc:dd:ee:ff', 'fresh']], [], [], []]
    monkeypatch.setattr(app, 'collect_device', collect)
    targets = [dict(management_ip='192.0.2.1', username='fixture', password='fixture')]
    kwargs = dict(inventory_path=tmp_path / 'inventory.json', reports_dir=tmp_path,
                  spool_root=tmp_path / 'runs', log_root=tmp_path / 'logs',
                  execution_config=ExecutionConfig(1, 0, 0))
    app.run_lookup(targets, object(), **kwargs)
    first, = (tmp_path / 'runs').iterdir()
    app.run_lookup(targets, object(), **kwargs)
    second, = (p for p in (tmp_path / 'runs').iterdir() if p != first)
    old, new = isolated(ResultSpool(first), ResultSpool(second))
    assert old[0]['payload'][0]
    assert bool(new[0]['payload'][0]) is (not failed)
    assert calls == ['192.0.2.1'] * 2
    app.export_mac_spool(first, tmp_path / 'historical.xlsx')
    assert calls == ['192.0.2.1'] * 2


def test_inventory_failed_refresh_labels_retained_facts(tmp_path, monkeypatch):
    import test_inventory_refresh as fixture
    from orbitflow import inventory_refresh as app
    outcomes = {'192.0.2.1': ('SERIAL1', 'fresh-router')}
    calls, _, _ = fixture.setup_network(monkeypatch, outcomes)
    source = fixture.input_file(tmp_path, [fixture.target('192.0.2.1')])
    kwargs = dict(inventory_path=tmp_path / 'inventory.json', log_root=tmp_path / 'logs',
                  spool_root=tmp_path / 'runs', execution_config=ExecutionConfig(1, 0, 0))
    app.refresh_inventory_from_excel(source, fixture.CONFIG,
        export_path=tmp_path / 'first.xlsx', clock=lambda: fixture.BEFORE, **kwargs)
    outcomes['192.0.2.1'] = 'connect-failure'
    path = app.refresh_inventory_from_excel(source, fixture.CONFIG,
        export_path=tmp_path / 'second.xlsx', clock=lambda: fixture.NOW, **kwargs)
    assert calls == ['192.0.2.1'] * 2
    rows = fixture.sheets(path)
    row, = rows['Inventory']
    assert row['refresh_status'] == app.FAILED
    assert row['last_successful_collection'] == fixture.BEFORE.isoformat()
    assert row['last_collection_attempt'] == fixture.NOW.isoformat()
    assert rows['Run_Attempts'][0]['status'] == 'failed'
