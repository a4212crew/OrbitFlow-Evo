"""Synthetic operator-approved patterns; no claims about operational health."""
import json
import hashlib
from pathlib import Path
import zipfile

import pytest
from openpyxl import load_workbook

from orbitflow.compliance import evaluate_vlan_compliance
from orbitflow.compliance.methodology import resolve_methodologies
from orbitflow.execution import DeviceOutcome
from orbitflow.methodology_report import COLUMNS, export_methodology_spool
from orbitflow.result_spool import ResultSpool
from test_methodology_resolution import resolve, snapshot
from test_vlan_compliance import POLICY


EDGE = 'vlan database\n vlan 445,545\nexit\ninterface 0/1\n'
EVC = ('interface Gi0/1\n switchport mode trunk\n switchport trunk allowed vlan none\n'
       ' service instance 7 ethernet\n  encapsulation dot1q 445\n  bridge-domain 445\n'
       ' service instance 8 ethernet\n  encapsulation dot1q 545\n  bridge-domain 545')
PATTERNS = [
    ('EdgeSwitch', EDGE + ' vlan participation include 445\n vlan pvid 445\n'
     ' vlan participation exclude 545\n vlan tagging 545\nexit',
     'standard_configuration', 'informational', 'EXCLUDED_VLAN_TAGGING_INACTIVE'),
    ('C3850', 'interface Gi0/1\n switchport mode trunk\n switchport trunk allowed vlan 445',
     'standard_configuration', 'informational', 'ALLOWED_VLAN_NOT_IN_DATABASE'),
    ('EdgeSwitch', 'interface 0/1\n switchport\n switchport mode trunk\n switchport trunk allowed vlan 445\nexit',
     'working_non_standard', 'warning', 'EDGESWITCH_SWITCHPORT_STYLE'),
    ('C3850', 'vlan 445\n!\ninterface Gi0/1\n switchport mode trunk\n switchport access vlan 545\n'
     ' switchport trunk allowed vlan 445', 'working_non_standard', 'warning', 'RESIDUAL_SWITCHPORT_SETTINGS'),
    ('ME3600X', EVC, 'standard_configuration', 'none', 'ME3600X_STANDARD_EVC_TRUNK'),
    ('EdgeSwitch', EDGE + ' vlan participation include 445,545\n vlan tagging 445\n vlan pvid 545\n'
     ' addport 3/1\nexit\ninterface lag 1\n vlan participation include 445,545\n'
     ' vlan tagging 545\n vlan pvid 445\nexit',
     'wrong_configuration', 'error', 'LAG_MEMBER_CONFIGURATION_MISMATCH'),
    ('C3850', 'interface Gi0/1\n description unused\n shutdown',
     'not_applicable', 'none', 'NO_RELEVANT_L2_SERVICE'),
]


def test_full_compliance_results_equal_pre_issue81_baseline():
    # Captured using HEAD's untouched parser/resolver in an isolated source copy.
    expected = json.loads((Path(__file__).parent / 'fixtures/methodology_standards_compliance.json').read_text())
    actual = []
    for family, config, *_ in PATTERNS:
        ctx, state = snapshot(config, family)
        result = evaluate_vlan_compliance(ctx, [], state, POLICY)
        actual.append(hashlib.sha256(json.dumps(result, sort_keys=True).encode()).hexdigest())
    assert actual == expected


@pytest.mark.parametrize('family,config,classification,severity,code', PATTERNS)
def test_seven_patterns(family, config, classification, severity, code):
    ctx, state = snapshot(config, family)
    before = evaluate_vlan_compliance(ctx, [], state, POLICY)
    rows = resolve_methodologies(ctx, [], state)
    assert evaluate_vlan_compliance(ctx, [], state, POLICY) == before
    row = rows[0]
    assert (row['configuration_classification'], row['finding_severity'], row['standard_finding']) == (
        classification, severity, code)
    assert all(r['engineering_review_decision'] == 'not_reviewed' for r in rows)
    for row in rows:
        for proof in row['evidence']:
            assert proof['excerpt'] == config.splitlines()[proof['line'] - 1]
            assert proof['source_filename'] == 'fixture.cfg'
    if code == 'ALLOWED_VLAN_NOT_IN_DATABASE':
        assert any(445 in f['missing_vlans'] for f in before)
        assert 'ALLOWED_VLAN_NOT_IN_DATABASE' in json.dumps(before)
        assert any(f['configuration_health'] == 'wrong_configuration' for f in before)
    if code == 'LAG_MEMBER_CONFIGURATION_MISMATCH':
        relation = rows[0]['mapping']['lag_relationship']
        assert relation['member'] == '0/1' and relation['aggregate'] == 'lag 1'
        assert relation['differences'] == {'pvid': {'member': [545], 'aggregate': [445]},
                                           'tagging': {'member': [445], 'aggregate': [545]}}
        assert ' addport 3/1' in [e['excerpt'] for e in relation['evidence']]
    if family == 'ME3600X':
        assert [r['service_instance_id'] for r in rows] == ['', '7', '8']
        assert all(r['standard_finding'] == code for r in rows)


@pytest.mark.parametrize('family', ['C3750X', 'C3850', 'ME3600X'])
@pytest.mark.parametrize('mode', ['access', 'trunk'])
def test_selected_mode_residual_settings(family, mode):
    config = ('vlan 445\n!\ninterface Gi0/1\n switchport mode ' + mode + '\n'
              ' switchport access vlan 545\n switchport trunk allowed vlan 445')
    row = resolve(config, family)[0]
    assert row['configuration_classification'] == 'working_non_standard'
    assert row['standard_finding'] == 'RESIDUAL_SWITCHPORT_SETTINGS'
    for suffix in ('\n no switchport', '\n switchport mode dynamic', '\n switchport trunk allowed vlan arbitrary',
                   '\n switchport mode ' + ('access' if mode == 'trunk' else 'trunk')):
        assert resolve(config + suffix, family)[0]['configuration_classification'] == 'review_needed'


@pytest.mark.parametrize('config,code', [
    (EVC.replace(' switchport mode trunk\n', ''), 'ME3600X_EVC_PREREQUISITES_MISSING'),
    (EVC.replace(' switchport trunk allowed vlan none\n', ''), 'ME3600X_EVC_PREREQUISITES_MISSING'),
    (EVC.replace('allowed vlan none', 'allowed vlan 445'), 'ME3600X_EVC_PREREQUISITES_MISSING'),
    (EVC.replace('encapsulation dot1q 545', 'encapsulation dot1q 445'), 'ME3600X_EVC_SERVICE_OVERLAP'),
    (EVC.replace('  bridge-domain 545', ''), 'ME3600X_EVC_SERVICE_UNRESOLVED'),
])
def test_evc_missing_prerequisites_or_overlap(config, code):
    rows = resolve(config, 'ME3600X')
    assert all(r['configuration_classification'] == 'review_needed' for r in rows)
    assert all(code in {f['code'] for f in r['standard_findings']} for r in rows)


def test_evc_distinct_qinq_and_service_evidence():
    config = EVC.replace('dot1q 445', 'dot1q 445 second-dot1q 17').replace(
        'dot1q 545', 'dot1q 445 second-dot1q 18')
    rows = resolve(config, 'ME3600X')
    assert all(r['configuration_classification'] == 'standard_configuration' for r in rows)
    assert rows[1]['mapping']['inner_vlan'] == [17]
    assert rows[2]['mapping']['inner_vlan'] == [18]
    assert ' switchport mode trunk' in [e['excerpt'] for e in rows[2]['evidence']]


@pytest.mark.parametrize('tail', [
    ' vlan pvid 545',  # Still excluded: real independent conflict.
    ' encapsulation unknown token arbitrary-private-value',
])
def test_excluded_tagging_does_not_hide_other_findings(tail):
    config = PATTERNS[0][1].removesuffix('exit') + tail + '\nexit'
    row = resolve(config, 'EdgeSwitch')[0]
    assert row['configuration_classification'] == 'review_needed'
    assert 'arbitrary-private-value' not in json.dumps(row)


def test_reinclude_cancels_exclusion_and_restores_valid_tagging():
    config = PATTERNS[0][1].removesuffix('exit') + ' vlan participation include 545\nexit'
    row = resolve(config, 'EdgeSwitch')[0]
    assert row['configuration_classification'] == 'standard_configuration'
    assert row['finding_severity'] == 'none'
    assert 'EXCLUDED_VLAN_TAGGING_INACTIVE' not in json.dumps(row['standard_findings'])


@pytest.mark.parametrize('config', [
    'interface 0/1\n switchport mode mystery\nexit',
    'interface 0/1\n switchport trunk allowed vlan 445\nexit',
    'interface 0/1\n switchport mode trunk\n vlan participation include 445\nexit',
    'interface 0/1\n switchport mode trunk\n switchport mode access\nexit',
])
def test_edge_unknown_or_incomplete_is_not_known_nonstandard(config):
    assert resolve(config, 'EdgeSwitch')[0]['configuration_classification'] == 'review_needed'


@pytest.mark.parametrize('reference', ['3/1', 'lag 1'])
def test_lag_configured_diffs_survive_database_filtering(reference):
    # Both sides lose 545 from audit membership; standards still see the difference.
    config = ('vlan database\n vlan 445\nexit\ninterface 0/1\n vlan participation include 445,545\n'
              ' vlan tagging 445,545\n addport ' + reference + '\nexit\ninterface lag 1\n'
              ' vlan participation include 445\n vlan tagging 445\nexit')
    row = resolve(config, 'EdgeSwitch')[0]
    assert row['configuration_classification'] == 'wrong_configuration'
    assert row['mapping']['lag_relationship']['differences']['participation'] == {
        'member': [445, 545], 'aggregate': [445]}
    for replacement in ('3/2', '1', '0/1'):
        missing = resolve(config.replace('addport ' + reference, 'addport ' + replacement), 'EdgeSwitch')[0]
        assert missing['configuration_classification'] == 'review_needed'
        assert 'lag_relationship' not in missing['mapping']


def test_membership_truth_and_absent_member_settings():
    aggregate = 'interface lag 1\n vlan participation include 445\n vlan tagging 445\nexit'
    base = 'vlan database\n vlan 445\nexit\ninterface 0/1\n'
    inherited = resolve(base + ' addport 3/1\nexit\n' + aggregate, 'EdgeSwitch')[0]
    assert inherited['mapping']['lag_relationship']['differences'] == {}
    assert inherited['mapping']['lag_relationship']['unresolved_fields'] == {}
    assert inherited['mapping']['lag_relationship']['inherited_from'] == 'lag 1'
    assert inherited['configuration_classification'] == 'standard_configuration'
    no_relation = resolve(base + ' description lag 1\nexit\n' + aggregate, 'EdgeSwitch')[0]
    assert no_relation['configuration_classification'] == 'not_applicable'
    assert 'lag_relationship' not in no_relation['mapping']


@pytest.mark.parametrize('field,statement,common', [
    ('tagging', ' vlan tagging 545', ' vlan participation include 445,545'),
    ('pvid', ' vlan pvid 445', ' vlan participation include 445,545'),
    ('participation', ' vlan participation include 445,545', ' vlan tagging 545'),
])
@pytest.mark.parametrize('configured_side', ['member', 'aggregate'])
def test_edge_one_sided_lag_fields_are_unknown(field, statement, common, configured_side):
    member = common + ('\n' + statement if configured_side == 'member' else '')
    aggregate = common + ('\n' + statement if configured_side == 'aggregate' else '')
    config = (EDGE + member + '\n addport 3/1\nexit\ninterface lag 1\n' + aggregate + '\nexit')
    ctx, state = snapshot(config, 'EdgeSwitch')
    before = evaluate_vlan_compliance(ctx, [], state, POLICY)
    row = resolve_methodologies(ctx, [], state)[0]
    assert evaluate_vlan_compliance(ctx, [], state, POLICY) == before
    assert row['configuration_classification'] == 'review_needed'
    assert row['standard_finding'] == 'LAG_MEMBER_CONFIGURATION_UNRESOLVED'
    assert row['finding_severity'] == 'warning'
    assert row['engineering_review_decision'] == 'not_reviewed'
    relation = row['mapping']['lag_relationship']
    assert relation['differences'] == {}
    assert relation['inherited_from'] == ''
    missing = 'aggregate' if configured_side == 'member' else 'member'
    unresolved = relation['unresolved_fields'][field]
    assert unresolved[missing] is None
    assert unresolved[configured_side] == {'tagging': [545], 'pvid': [445],
                                         'participation': [445, 545]}[field]
    assert f'{missing} stanza omits {field}' in unresolved['explanation']
    excerpts = [e['excerpt'] for e in relation['evidence']]
    assert all(line in excerpts for line in ('interface 0/1', 'interface lag 1', ' addport 3/1', statement))
    for proof in relation['evidence']:
        assert proof['excerpt'] == config.splitlines()[proof['line'] - 1]
        assert proof['source_filename'] == 'fixture.cfg'


def test_edge_unknown_lag_owner_does_not_imply_equality_or_reverse_inheritance():
    for member in ('', ' vlan participation include 445\n vlan tagging 445\n'):
        config = EDGE + member + ' addport 3/1\nexit\ninterface lag 1\n description trunk\nexit'
        row = resolve(config, 'EdgeSwitch')[0]
        assert row['configuration_classification'] == 'review_needed'
        relation = row['mapping']['lag_relationship']
        assert relation['unresolved_fields']
        assert relation['differences'] == {}
        assert relation['inherited_from'] == ''


def test_edge_confirmed_difference_wins_over_one_sided_unknown():
    config = (EDGE + ' vlan participation include 445,545\n vlan tagging 545\n vlan pvid 445\n'
              ' addport 3/1\nexit\ninterface lag 1\n vlan participation include 445,545\n'
              ' vlan tagging 445\nexit')
    row = resolve(config, 'EdgeSwitch')[0]
    assert row['configuration_classification'] == 'wrong_configuration'
    assert row['standard_finding'] == 'LAG_MEMBER_CONFIGURATION_MISMATCH'
    relation = row['mapping']['lag_relationship']
    assert relation['differences'] == {'tagging': {'member': [545], 'aggregate': [445]}}
    assert relation['unresolved_fields']['pvid']['aggregate'] is None


def test_edge_explicit_empty_tagging_is_a_demonstrable_difference():
    config = (EDGE + ' vlan participation include 445\n vlan tagging 445\n addport 3/1\nexit\n'
              'interface lag 1\n vlan participation include 445\n vlan tagging 445 disable\nexit')
    row = resolve(config, 'EdgeSwitch')[0]
    assert row['configuration_classification'] == 'wrong_configuration'
    assert row['mapping']['lag_relationship']['differences'] == {
        'tagging': {'member': [445], 'aggregate': []}}


def test_edge_explicit_matched_lag_fields_remain_standard():
    settings = ' vlan participation include 445,545\n vlan tagging 545\n vlan pvid 445\n'
    row = resolve(EDGE + settings + ' addport 3/1\nexit\ninterface lag 1\n' + settings + 'exit',
                  'EdgeSwitch')[0]
    assert row['configuration_classification'] == 'standard_configuration'
    relation = row['mapping']['lag_relationship']
    assert relation['differences'] == relation['unresolved_fields'] == {}
    assert relation['inherited_from'] == ''


def test_ambiguous_lag_membership_is_reviewable():
    config = ('interface 0/1\n addport 3/1\n addport 3/2\nexit\n'
              'interface lag 1\nexit\ninterface lag 2\nexit')
    row = resolve(config, 'EdgeSwitch')[0]
    assert row['configuration_classification'] == 'review_needed'
    assert row['standard_finding'] == 'LAG_RELATIONSHIP_UNRESOLVED'


def test_cisco_member_aggregate_vlan_differences():
    config = ('vlan 445,545\n!\ninterface Gi0/1\n switchport mode trunk\n'
              ' switchport trunk allowed vlan 445\n channel-group 1 mode active\n!\n'
              'interface Port-channel1\n switchport mode trunk\n switchport trunk allowed vlan 445,545')
    row = resolve(config)[0]
    assert row['configuration_classification'] == 'wrong_configuration'
    assert row['mapping']['lag_relationship']['differences'] == {
        'participation': {'member': [445], 'aggregate': [445, 545]}}


def test_known_alternate_lag_style_does_not_hide_confirmed_difference():
    config = ('interface 0/1\n switchport mode trunk\n switchport trunk allowed vlan 445\n'
              ' addport 3/1\nexit\ninterface lag 1\n switchport mode trunk\n'
              ' switchport trunk allowed vlan 445,545\nexit')
    row = resolve(config, 'EdgeSwitch')[0]
    assert row['configuration_classification'] == 'wrong_configuration'
    assert row['mapping']['lag_relationship']['differences'] == {
        'participation': {'member': [445], 'aggregate': [445, 545]}}
    assert ' switchport trunk allowed vlan 445,545' in [
        e['excerpt'] for e in row['mapping']['lag_relationship']['evidence']]


def test_inactive_residual_difference_is_not_confirmed_lag_mismatch():
    config = ('vlan 445\n!\ninterface Gi0/1\n switchport mode trunk\n switchport access vlan 17\n'
              ' switchport trunk allowed vlan 445\n channel-group 1 mode active\n!\n'
              'interface Port-channel1\n switchport mode trunk\n switchport access vlan 18\n'
              ' switchport trunk allowed vlan 445')
    row = resolve(config, 'C3750X')[0]
    assert row['mapping']['lag_relationship']['differences'] == {}
    assert row['configuration_classification'] == 'working_non_standard'


@pytest.mark.parametrize('family,config', [
    ('ASR920', 'interface Gi0/1\n service instance 7 ethernet'),
    ('C3850', 'interface Gi0/1\n switchport'),
    ('C3850', 'interface Gi0/1 l2transport'),
    ('NE05E', 'interface GE0/1\n l2 binding vsi missing'),
    ('NCS540', 'interface TenGigE0/0/0/1.7 l2transport'),
])
def test_incomplete_l2_never_not_applicable(family, config):
    assert resolve(config, family)[0]['configuration_classification'] == 'review_needed'


def test_workbook_contract_safe_evidence_and_recovery(tmp_path):
    config = (PATTERNS[2][1].removesuffix('exit') + ' switchport mode token confidential\n'
              ' switchport trunk allowed vlan =HYPERLINK("unsafe")\n'
              ' password hiddenvalue\n description =1+1\nexit')
    rows = resolve(config, 'EdgeSwitch')
    spool = ResultSpool.create(tmp_path / 'runs', 'vlan_compliance', 2)
    error = dict(management_ip='192.0.2.2', stage='methodology', error_category='ValueError')
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={'methodologies': rows})
        spool.append(DeviceOutcome(2), payload={'methodology_errors': [error]})
    path = export_methodology_spool(spool.path, tmp_path / 'review.xlsx')
    book = load_workbook(path)
    headers, *values = book['Methodology Resolution'].values
    old = ('Input Position', 'Record', 'Device ID', 'Device IP', 'Device Name', 'Family',
           'Record Kind', 'Interface', 'Config Interface', 'Interface Match', 'Parent Interface',
           'Configuration Owner', 'Service Instance', 'Methodology', 'Subtype', 'Status', 'Review Needed',
           'Mapping', 'Configuration Findings', 'Configuration Evidence', 'Evidence Source', 'Evidence Lines',
           'Configuration Facts Digest')
    added = {'Configuration Classification', 'Finding Severity', 'Standard Finding', 'Engineering Review Decision'}
    assert set(old) - set(headers) == {'Interface', 'Service Instance', 'Evidence Source'}
    assert set(headers) - set(old) == added
    assert headers == COLUMNS and len(values[0]) == len(COLUMNS)
    displayed = dict(zip(headers, values[0]))
    assert ' switchport mode trunk' in displayed['Configuration Evidence']
    assert displayed['Engineering Review Decision'] == 'not_reviewed'
    detail = json.loads(list(book['Details'].values)[1][-1])
    assert detail['service_instance_id'] == '' and detail['evidence'][0]['source_filename'] == 'fixture.cfg'
    assert detail['configuration_facts_digest'] == rows[0]['configuration_facts_digest']
    fallback = dict(zip(headers, values[-1]))
    assert fallback['Configuration Classification'] == 'review_needed'
    assert fallback['Engineering Review Decision'] == 'not_reviewed'
    assert len(list(book['Run Errors'].values)) == 2
    assert all(c.data_type != 'f' for sheet in book for row in sheet for c in row)
    first = [list(sheet.values) for sheet in book]
    book.close()
    with zipfile.ZipFile(path) as archive:
        content = b''.join(archive.read(n) for n in archive.namelist())
    assert all(value not in content for value in (b'confidential', b'hiddenvalue', b'HYPERLINK', b'unsafe&quot;'))
    export_methodology_spool(spool.path, path)
    book = load_workbook(path)
    assert first == [list(sheet.values) for sheet in book]
    book.close()
