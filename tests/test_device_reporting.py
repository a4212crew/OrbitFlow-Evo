from dataclasses import replace
from datetime import datetime, timezone
from io import StringIO
from importlib.util import module_from_spec, spec_from_file_location
import json
from pathlib import Path
from unittest.mock import Mock

from openpyxl import Workbook, load_workbook
from paramiko import ChannelException
import pytest

from orbitflow import reporting
from orbitflow.capabilities.interfaces import InterfaceCapabilityError
from orbitflow.models import DeviceContext, InterfaceRecord, InterfaceVlanObservation, VlanObject, VlanState
from orbitflow.targets import load_targets
from orbitflow.transport import DeviceSession, TransportConfig
from orbitflow.vendors.interface_names import canonical_interface_name
from test_cli_lifecycle import Channel

NOW = datetime(2026, 9, 25, tzinfo=timezone.utc)
CONFIG = TransportConfig('proxy', 'cluster', 'bastion', 'operator')


def context(ip='192.0.2.1', platform='cisco_xe'):
    return DeviceContext('id-' + ip, ip, (ip,), 'router', 'Cisco', platform, 'ASR920',
                         '', 'asr920_evc', ('evc',), '', '', '', NOW, NOW)


def interface(name='Gi0/1', description='uplink'):
    return InterfaceRecord('router', '192.0.2.1', 'cisco_xe', name, description, 'up', 'down', NOW)


def state(*observations, objects=()):
    return VlanState('router', '192.0.2.1', 'cisco_xe', observations, objects, NOW)


def test_join_service_details_and_vlan_identity():
    vlans = state(
        InterfaceVlanObservation('GigabitEthernet0/1', mode='service', outer_vlan=200, inner_vlan=201,
                                 service_binding_type='bridge_domain', service_binding_name='999'),
        InterfaceVlanObservation('GigabitEthernet0/1', mode='service', outer_vlan=100, inner_vlan=101,
                                 service_binding_type='bridge_domain', service_binding_name='888'),
        InterfaceVlanObservation('GigabitEthernet0/1.42', mode='l2', service_binding_name='777'),
        objects=(VlanObject('bridge_domain', '999'), VlanObject('vlan', '100', 'users', (100,))),
    )
    rows, objects = reporting.build_rows(context(), [interface(), interface('Gi0/2')], vlans)
    data = [dict(zip(reporting.INTERFACE_COLUMNS, row)) for row in rows]
    summary = next(row for row in data if row['Interface'] == 'Gi0/1')
    assert summary['Attached VLANs'] == '100, 101, 200, 201'
    assert summary['Description'] == 'uplink'
    assert (summary['Admin Status'], summary['Oper Status']) == ('up', 'down')
    assert list(zip(summary['Outer VLAN'].splitlines(), summary['Inner VLAN'].splitlines(),
                    summary['Service Binding Name'].splitlines())) == [('100', '101', '888'), ('200', '201', '999')]
    assert len(rows) == 3
    assert next(row for row in data if row['Interface'].endswith('.42'))['Attached VLANs'] == ''
    assert next(row for row in data if row['Interface'] == 'Gi0/2')['Mode'] == ''
    assert objects[0][3:7] == ['bridge_domain', '999', '', '']
    assert objects[1][3:7] == ['vlan', '100', '100', 'users']
    assert reporting.build_rows(context(), [interface('Gi0/2'), interface()], replace(vlans, interfaces=tuple(reversed(vlans.interfaces)))) == (rows, objects)


def test_union_only_observed_ids_and_explicit_none():
    vlans = state(InterfaceVlanObservation('Gi0/1', access_vlan=10, native_vlan=11, pvid=12,
                  allowed_vlans=(), tagged_vlans=(13,), untagged_vlans=(14,), service_vlan=15,
                  control_vlan=16, outer_vlan=17, inner_vlan=18, referenced_vlans=(19,), excluded_vlans=(999,)))
    rows, _ = reporting.build_rows(context(), [interface()], vlans)
    row = dict(zip(reporting.INTERFACE_COLUMNS, rows[0]))
    assert row['Attached VLANs'] == ', '.join(map(str, range(10, 20)))
    assert row['Allowed VLANs'] == 'none'
    assert row['Excluded VLANs'] == '999'
    assert row['Control VLAN'] == '16'
    assert reporting.build_rows(context(), [], None) == ([], [])


@pytest.mark.parametrize('platform,left,right', [
    ('cisco_ios', 'Gi1/0/1', 'GigabitEthernet1/0/1'),
    ('cisco_xe', 'Te0/1', 'TenGigabitEthernet0/1'),
    ('cisco_xr', 'BE10.3', 'Bundle-Ether10.3'),
    ('cisco_xr', 'Te0/0/0/1.2', 'TenGigE0/0/0/1.2'),
    ('huawei_vrp', 'GE0/1/0.2', 'GigabitEthernet0/1/0.2'),
    ('huawei_vrp', 'Eth0/1/0', 'Ethernet0/1/0'),
    ('ubiquiti_edgeswitch', '0/1', '0/1'),
])
def test_canonical_names(platform, left, right):
    assert canonical_interface_name(platform, left) == canonical_interface_name(platform, right)
    assert canonical_interface_name(platform, left + '.99') != canonical_interface_name(platform, right)


def test_workbook_headers_formatting_and_literal_text(tmp_path):
    path = tmp_path / 'report.xlsx'
    rows, objects = reporting.build_rows(context(), [interface(description='=HYPERLINK("bad")')], state())
    reporting.write_workbook(path, rows, objects, [])
    with_workbook = load_workbook(path)
    try:
        assert with_workbook.sheetnames == ['Interfaces', 'VLAN_Database', 'Run_Errors']
        for sheet, columns in zip(with_workbook, [reporting.INTERFACE_COLUMNS, reporting.DATABASE_COLUMNS, reporting.ERROR_COLUMNS]):
            assert tuple(cell.value for cell in sheet[1]) == columns
            assert sheet.freeze_panes == 'A2'
            assert sheet.auto_filter.ref.endswith(str(sheet.max_row))
            assert sheet.column_dimensions['A'].width >= 18
        assert with_workbook['Interfaces']['F2'].data_type == 's'
        assert with_workbook['Interfaces']['F2'].value == '=HYPERLINK("bad")'
    finally:
        with_workbook.close()


def install_fakes(monkeypatch, failing_stage=None, fail_ip='192.0.2.1'):
    events, closed = [], []
    def connect(ip, credentials, config):
        events.append(('connect', ip))
        assert credentials.password == 'synthetic-password'
        assert config is CONFIG
        if failing_stage == 'connect' and ip == fail_ip:
            raise RuntimeError('synthetic-password')
        def close():
            closed.append(ip)
            if failing_stage == 'disconnect' and ip == fail_ip:
                raise RuntimeError('synthetic-password')
        channel = Channel("router#", {"": "", "terminal length 0": ""})
        session = DeviceSession(Mock(invoke_shell=Mock(return_value=channel)), close)
        session.ip = ip
        return session
    def resolve(session, *, management_ip, cli):
        assert cli.session is session
        session.cli = cli
        events.append(('inventory', session))
        if failing_stage == 'inventory' and management_ip == fail_ip:
            raise RuntimeError('synthetic-password')
        session.context = context(management_ip)
        return session.context
    def collect(stage):
        def run(session, ctx, *, cli):
            assert cli is session.cli
            cli.require_session(session)
            assert session.context is ctx
            assert session.ip == ctx.management_ip
            events.append((stage, session))
            if stage == failing_stage and session.ip == fail_ip:
                raise RuntimeError('synthetic-password token=never-output')
            if stage == 'interfaces':
                return [interface(description='synthetic-password token=never-output')]
            return state(InterfaceVlanObservation('Gi0/1', access_vlan=100), objects=(VlanObject('vlan', '100', 'users', (100,)),))
        return run
    monkeypatch.setattr(reporting, 'connect_device', connect)
    monkeypatch.setattr(reporting, 'DeviceInventoryResolver', Mock(return_value=Mock(resolve=resolve)))
    monkeypatch.setattr(reporting, 'InterfaceService', Mock(return_value=Mock(collect=collect('interfaces'))))
    monkeypatch.setattr(reporting, 'VlanService', Mock(return_value=Mock(collect=collect('vlans'))))
    return events, closed


@pytest.mark.parametrize('stage', ['connect', 'inventory', 'interfaces', 'vlans', 'disconnect', None])
def test_batch_isolation_session_context_and_secrets(tmp_path, monkeypatch, stage):
    events, closed = install_fakes(monkeypatch, stage)
    targets = [{'management_ip': f'192.0.2.{i}', 'username': 'synthetic-user', 'password': 'synthetic-password'} for i in (1, 2)]
    writer = Mock(wraps=reporting.write_workbook)
    monkeypatch.setattr(reporting, 'write_workbook', writer)
    output = StringIO()
    path = reporting.run_report(targets, CONFIG, inventory_path=tmp_path / 'inventory.json',
                                reports_dir=tmp_path, log_root=tmp_path / 'logs', output=output, clock=lambda: NOW)
    assert writer.call_count == 1
    workbook = load_workbook(path)
    try:
        assert workbook['Run_Errors'].max_row == (2 if stage else 1)
        if stage:
            assert workbook['Run_Errors']['C2'].value == stage
            assert workbook['Run_Errors']['D2'].value == 'RuntimeError'
        assert any(row[1] == '192.0.2.2' for row in workbook['Interfaces'].iter_rows(min_row=2, values_only=True))
        if stage == 'interfaces':
            assert workbook['Interfaces']['F2'].value is None
        if stage == 'vlans':
            assert workbook['Interfaces']['J2'].value is None
        rendered = str([[row for row in sheet.values] for sheet in workbook])
    finally:
        workbook.close()
    log = next((tmp_path / 'logs' / 'reporting').glob('*/interface_vlan_report.log')).read_text()
    failures = [json.loads(line) for line in log.splitlines() if json.loads(line)['level'] == 'ERROR']
    assert len(failures) == bool(stage)
    if stage:
        assert stage in failures[0]['message']
        assert failures[0]['management_ip'] == '192.0.2.1'
        assert failures[0]['exception_chain'][0]['category'] == 'RuntimeError'
        assert failures[0]['exception_chain'][0]['frames']
    for secret in ['synthetic-user', 'synthetic-password', 'never-output']:
        assert secret not in rendered + log + output.getvalue()
    sessions = [value for event, value in events if event == 'inventory']
    for session in sessions:
        collected = [(event, value) for event, value in events if event in {'interfaces', 'vlans'} and value is session]
        assert [event for event, _ in collected] == ([] if stage == 'inventory' and session.ip == '192.0.2.1' else ['interfaces', 'vlans'])
    assert len(closed) == (1 if stage == 'connect' else 2)


@pytest.mark.parametrize('stage', ['interfaces', 'workbook'])
def test_reporting_failure_logs_safe_exception_chain(tmp_path, monkeypatch, stage):
    install_fakes(monkeypatch)

    def fail(*args, **kwargs):
        local_secret = 'local-only-sensitive-value'
        try:
            raise ChannelException(4, 'raw-device-output synthetic-password')
        except ChannelException as exc:
            raise InterfaceCapabilityError('raw-wrapper-error token=never-output') from exc

    if stage == 'interfaces':
        monkeypatch.setattr(reporting, 'InterfaceService', Mock(return_value=Mock(collect=fail)))
    else:
        monkeypatch.setattr(reporting, 'write_workbook', fail)
    targets = [{'management_ip': '192.0.2.1', 'username': 'synthetic-user', 'password': 'synthetic-password'}]
    output = StringIO()

    def run():
        return reporting.run_report(targets, CONFIG, reports_dir=tmp_path,
                                    log_root=tmp_path / 'logs', output=output, clock=lambda: NOW)

    if stage == 'workbook':
        with pytest.raises(InterfaceCapabilityError):
            run()
        rendered = ''
    else:
        workbook = load_workbook(run())
        try:
            rows = list(workbook['Run_Errors'].values)
            assert rows == [reporting.ERROR_COLUMNS,
                            ('192.0.2.1', 'router', 'interfaces', 'InterfaceCapabilityError', NOW.isoformat())]
            rendered = str(rows)
        finally:
            workbook.close()
    log = next((tmp_path / 'logs' / 'reporting').glob('*/interface_vlan_report.log')).read_text()
    failures = [record for line in log.splitlines() if (record := json.loads(line))['level'] == 'ERROR']
    assert len(failures) == 1
    chain = failures[0]['exception_chain']
    assert [item['category'] for item in chain] == ['InterfaceCapabilityError', 'ChannelException']
    for item in chain:
        assert set(item) == {'category', 'errno', 'frames'}
        assert item['frames']
        for frame in item['frames']:
            assert set(frame) == {'file', 'line', 'function'}
    for sensitive in ('raw-device-output', 'synthetic-password', 'raw-wrapper-error',
                      'never-output', 'local-only-sensitive-value', 'raise ChannelException'):
        assert sensitive not in log + rendered + output.getvalue()


def test_shared_loader_preserves_legacy_validation_and_isolates_bad_row(tmp_path, monkeypatch):
    path = tmp_path / 'targets.xlsx'
    workbook = Workbook()
    workbook.active.append(['management_ip', 'username', 'password', 'unused'])
    workbook.active.append(['192.0.2.1', 'synthetic-user', None, 'ignored'])
    workbook.active.append(['192.0.2.2', 'synthetic-user', 'synthetic-password'])
    workbook.save(path)
    with pytest.raises(ValueError, match='row 2'):
        load_targets(path)
    targets = load_targets(path, isolate_invalid=True)
    events, _ = install_fakes(monkeypatch)
    result = reporting.run_report(targets, CONFIG, reports_dir=tmp_path / 'reports',
                                 log_root=tmp_path / 'logs', output=StringIO(), clock=lambda: NOW)
    workbook = load_workbook(result)
    assert workbook['Run_Errors']['C2'].value == 'input'
    workbook.close()
    assert [value for event, value in events if event == 'connect'] == ['192.0.2.2']


def test_1500_device_batch_writes_once(tmp_path, monkeypatch):
    install_fakes(monkeypatch)
    targets = [{'management_ip': f'device-{i}', 'username': 'synthetic-user', 'password': 'synthetic-password'} for i in range(1500)]
    writer = Mock(wraps=reporting.write_workbook)
    monkeypatch.setattr(reporting, 'write_workbook', writer)
    path = reporting.run_report(targets, CONFIG, reports_dir=tmp_path, log_root=tmp_path / 'logs', output=StringIO(), clock=lambda: NOW)
    assert writer.call_count == 1
    workbook = load_workbook(path, read_only=True)
    assert sum(1 for _ in workbook['Interfaces'].values) == 1501
    assert sum(1 for _ in workbook['VLAN_Database'].values) == 1501
    workbook.close()


def test_command_uses_shared_loader_and_supplied_routing(tmp_path, monkeypatch):
    script = Path(__file__).parents[1] / 'scripts' / 'device_interface_vlan_report.py'
    spec = spec_from_file_location('report_command', script)
    command = module_from_spec(spec)
    spec.loader.exec_module(command)
    targets = [{'management_ip': '192.0.2.1', 'username': 'test-user', 'password': 'test-password'}]
    loader, runner = Mock(return_value=targets), Mock(return_value=tmp_path / 'result.xlsx')
    monkeypatch.setattr(command, 'load_targets', loader)
    monkeypatch.setattr(command, 'run_report', runner)
    command.main(['targets.xlsx', '--proxy', 'proxy:443', '--cluster', 'cluster',
                  '--bastion-host', 'bastion', '--bastion-user', 'operator',
                  '--teleport-key-path', 'identity', '--teleport-cert-path', 'identity-cert.pub',
                  '--reports-dir', str(tmp_path)])
    loader.assert_called_once_with('targets.xlsx', isolate_invalid=True)
    args, kwargs = runner.call_args
    assert args[0] is targets
    assert args[1].proxy == 'proxy:443'
    assert args[1].teleport_key_path == Path('identity')
    assert args[1].teleport_cert_path == Path('identity-cert.pub')
    assert kwargs['reports_dir'] == tmp_path


def test_full_vlan_range_is_not_truncated_by_log_sanitizer(tmp_path):
    rows, _ = reporting.build_rows(context(), [interface()], state(
        InterfaceVlanObservation('Gi0/1', allowed_vlans=tuple(range(1, 4095)))))
    path = tmp_path / 'full-range.xlsx'
    reporting.write_workbook(path, rows, [], [])
    workbook = load_workbook(path)
    assert workbook['Interfaces']['J2'].value == ', '.join(map(str, range(1, 4095)))
    workbook.close()
