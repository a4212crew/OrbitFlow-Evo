"""Offline manual configuration boundary and recovery tests; no live networking."""
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
import json
import socket

import pytest
from openpyxl import Workbook, load_workbook
from orbitflow.configuration import manual as m
from orbitflow.configuration.manual_auth import hash_password, authorize
from orbitflow.transport import DeviceCredentials
from orbitflow.vendors.common import DeviceCLI
from orbitflow.vendors.manual_configuration import finalize as real_manual_finalize


@pytest.fixture(scope='module')
def auth():
    return {'manual_configuration_auth': {'password_hash': hash_password('test authorization password')}}


def row(number, ip='192.0.2.1', text='interface Gi1\n description Customer\n exit'):
    return dict(source_row=number, management_ip=ip, commands=m.commands(text), error='')


def records(path):
    return [e for e in m.ResultJournal(path / 'events.jsonl').records() if e['kind'] == 'row_result']


@pytest.fixture
def fake(monkeypatch, tmp_path):
    calls = []
    fault = {}
    def fail(stage):
        if fault.get('stage') == stage:
            raise fault.get('exception', ValueError)('synthetic failure')
    class Session:
        def close(self):
            calls.append(('close_session',))
    def connect(ip, *args):
        fail('connect')
        calls.append(('connect', ip))
        return Session()
    class CLI:
        serial = 0
        def __init__(self, session, **kwargs):
            fail('shell')
            self.serial = CLI.serial = CLI.serial + 1
            self.session = session
            self.last_response = 'end\nhost#'
            calls.append(('shell', self.serial))
        def __enter__(self):
            return self
        def __exit__(self, *args):
            calls.append(('close_shell', self.serial))
        def capture_literal(self, command, **kwargs):
            calls.append(('command', self.serial, command))
            fail('capture')
            kwargs['receive'](fault.get('receive', command + '\nhost(config-if)#'))
    def resolve(self, session, **kwargs):
        fail('identify')
        return SimpleNamespace(platform=fault.get('platform', 'cisco_ios'), hostname='host')
    def collect(self, *args, **kwargs):
        fail('backup')
        return 'hostname host\n'
    def setup(*args):
        calls.append(('setup',))
        fail('setup')
    def finalize(*args):
        args[0].last_response = 'end\nhost#'
        calls.append(('finalize',))
        fail('finalize')
        return 'host#'
    def save(*args, **kwargs):
        calls.append(('save',))
        fail('save')
        return fault.get('save', dict(status='Saved', response='write memory\nhost#', explicit_error=False))
    monkeypatch.setattr(m, 'connect_device', connect)
    monkeypatch.setattr(m, 'DeviceCLI', CLI)
    monkeypatch.setattr(m.DeviceInventoryResolver, 'resolve', resolve)
    monkeypatch.setattr(m.ConfigurationService, 'collect', collect)
    monkeypatch.setattr(m.vendor, 'setup', setup)
    monkeypatch.setattr(m.vendor, 'finalize', finalize)
    monkeypatch.setattr(m, 'save_configuration', save)
    monkeypatch.chdir(tmp_path)
    return calls, fault


def run(rows, auth, **kwargs):
    config = m.ExecutionConfig(1, 0, 0)
    return m.execute(rows, profile=auth, password='test authorization password',
                     confirmation_digest=m.summary(rows, config)['digest'],
                     credentials={r['management_ip']: DeviceCredentials('operator', 'device-private') for r in rows},
                     transport_config=None, execution_config=config, **kwargs)


def test_duplicate_noncontiguous_ip_fresh_shell_per_row(fake, auth):
    calls, _ = fake
    rows = [row(2), row(3, '192.0.2.2'), row(4)]
    path = run(rows, auth)
    assert [r['source_row'] for r in records(path)] == [2, 4, 3]
    assert all(r['status'] == 'Completed' and r['save'] == 'Saved' for r in records(path))
    assert len([c for c in calls if c[0] == 'save']) == 3
    assert len([c for c in calls if c[0] == 'connect']) == 2
    assert len([c for c in calls if c[0] == 'close_shell']) == 3
    assert [c[2] for c in calls if c[0] == 'command'] == rows[0]['commands'] * 3
    assert len(list(Path('outputs/reports/manual_configuration').rglob('*.txt'))) == 2
    report = m.regenerate_report(path)
    book = load_workbook(report)
    assert [r[0].value for r in list(book.active.rows)[1:]] == ['2', '3', '4']
    with pytest.raises(FileExistsError):
        m.regenerate_report(path, report)


@pytest.mark.parametrize('stage', ['connect', 'shell', 'identify', 'backup', 'setup', 'capture', 'finalize'])
@pytest.mark.parametrize('exception', [ValueError, TimeoutError, KeyboardInterrupt])
def test_every_pre_save_boundary_blocks_device_rows(fake, auth, stage, exception):
    calls, fault = fake
    fault.update(stage=stage, exception=exception)
    path = run([row(2), row(3)], auth)
    assert not any(c[0] == 'save' for c in calls)
    result = records(path)
    assert result[0]['status'] == 'Incomplete'
    assert result[1]['attempted'] == 0
    assert 'earlier device row' in result[1]['reason']
    assert all(r['save'] == 'Not initiated' for r in result)


@pytest.mark.parametrize('response', ['% Invalid input\n', '--More--', 'Password:', 'Are you sure? (y/n)', 'Error: invalid parameter', 'Proceed with defaults?', 'Enter VLAN identifier:', '[yes]'])
def test_rejection_interaction_no_save(fake, auth, response):
    calls, fault = fake
    fault['receive'] = response
    path = run([row(2), row(4)], auth)
    assert not any(c[0] == 'save' for c in calls)
    assert records(path)[1]['attempted'] == 0


@pytest.mark.parametrize('platform', ['cisco_xr', 'huawei_vrp', 'junos'])
def test_unsupported_before_user_writes(fake, auth, platform):
    calls, fault = fake
    fault['platform'] = platform
    path = run([row(2)], auth)
    assert not any(c[0] in {'command', 'save', 'setup'} for c in calls)
    assert records(path)[0]['status'] == 'Incomplete'


@pytest.mark.parametrize('exception', [ValueError, KeyboardInterrupt])
def test_save_transmission_failure_unknown_no_retry(fake, auth, exception):
    calls, fault = fake
    fault.update(stage='save', exception=exception)
    path = run([row(2), row(3)], auth)
    assert len([c for c in calls if c[0] == 'save']) == 1
    assert records(path)[0]['save'] == 'Unknown'
    assert records(path)[1]['attempted'] == 0


@pytest.mark.parametrize('result, expected', [(dict(status='Unknown', response='', explicit_error=False), 'Unknown'),
    (dict(status='Saved', response='% Error', explicit_error=True), 'Failed')])
def test_save_result_error_stops(fake, auth, result, expected):
    calls, fault = fake
    fault['save'] = result
    path = run([row(2), row(3)], auth)
    assert records(path)[0]['save'] == expected
    assert records(path)[0]['status'] == 'Incomplete'
    assert records(path)[1]['attempted'] == 0


def test_gate_before_network_and_allocation(fake, auth):
    calls, _ = fake
    with pytest.raises(ValueError):
        m.execute([row(2)], profile=auth, password='wrong', confirmation_digest='x', credentials={}, transport_config=None)
    with pytest.raises(ValueError):
        m.execute([row(2)], profile={}, password='anything', confirmation_digest='x', credentials={}, transport_config=None)
    with pytest.raises(ValueError):
        m.execute([row(2)], profile=auth, password='test authorization password', confirmation_digest='wrong', credentials={}, transport_config=None)
    assert calls == []
    assert not Path('outputs/runs').exists()


@pytest.mark.parametrize('text', ['wr', 'wri mem', 'cop run start', 'sa', 'comm', 'do wr', 'run save', 'alias exec persist write memory',
    'persist', 'interface Gi1;write memory', 'description password abc', 'username admin secret abc', 'banner motd ^', 'python', 'more'])
def test_unsafe_commands_fail_closed(text):
    with pytest.raises(ValueError):
        m.commands(text)


def test_xlsx_rows_validation_formula_secret_and_source_unchanged(tmp_path):
    path = m.create_template(tmp_path)
    book = load_workbook(path)
    for values in [('192.0.2.1', 'interface Gi1\n description Literal\n exit'),
                   ('192.0.2.2', '=HYPERLINK("bad")'), ('192.0.2.1', 'password private'), ('invalid', 'shutdown')]:
        book.active.append(values)
    book.save(path)
    before = path.read_bytes()
    rows = m.load_rows(path)
    assert len(rows) == 4 and rows[0]['commands'][1] == ' description Literal'
    assert all(r['error'] for r in rows[1:])
    assert 'private' not in json.dumps(rows)
    assert before == path.read_bytes()
    assert m.create_template(tmp_path) != path


@pytest.mark.parametrize('kwargs', [dict(first_response=float('nan')), dict(maximum=float('inf')), dict(inactivity=0), dict(maximum=5), dict(save_timeout=301)])
def test_deadlines_finite_and_bounded(kwargs):
    with pytest.raises(ValueError):
        m.CaptureSettings(**kwargs)


def test_recording_failure_before_save_retains_recoverable_input(fake, auth, monkeypatch):
    calls, _ = fake
    original = m.ResultJournal.append
    def append(self, event):
        if event.get('kind') == 'command_intent':
            raise OSError('disk full')
        return original(self, event)
    monkeypatch.setattr(m.ResultJournal, 'append', append)
    path = run([row(2), row(3)], auth)
    assert not any(c[0] in {'command', 'save'} for c in calls)
    report = m.regenerate_report(path)
    assert report.exists()


def test_save_intent_recovery_unknown_and_partial_line(fake, auth):
    calls, _ = fake
    path = run([row(2)], auth)
    journal = path / 'events.jsonl'
    lines = [line for line in journal.read_text().splitlines(True) if json.loads(line)['kind'] != 'row_result']
    journal.write_text(''.join(lines) + '{"partial":', encoding='utf-8')
    calls.clear()
    report = m.regenerate_report(path)
    values = list(load_workbook(report).active.values)
    assert values[1][2] == 'Incomplete' and values[1][7] == 'Unknown'
    assert calls == []


def test_redaction_ansi_and_formula_literal(fake, auth):
    _, fault = fake
    fault['receive'] = '\x1b[31mdevice-private\x1b[0m\n=1+1\nsecret never-export\n'
    path = run([row(2)], auth)
    text = Path(records(path)[0]['transcript']).read_text()
    assert 'device-private' not in text and 'never-export' not in text and '\x1b' not in text
    assert '=1+1' in text


class Channel:
    def __init__(self, chunks):
        self.chunks = iter(chunks)
        self.sent = []
    def sendall(self, data):
        self.sent.append(data)
    def settimeout(self, timeout):
        assert 0 < timeout <= 30
    def recv(self, maximum):
        chunk = next(self.chunks)
        if isinstance(chunk, BaseException):
            raise chunk
        return chunk


def literal_cli(chunks):
    cli = DeviceCLI.__new__(DeviceCLI)
    cli.session = object()
    cli._channel = Channel(chunks)
    cli._synchronized = True
    cli._accepted_paging = set()
    return cli


def test_prompt_free_fragmented_utf8_capture():
    cli = literal_cli([b'not a prompt\nweird#\n\xe2', b'\x82\xac', socket.timeout()])
    received = []
    cli.capture_literal(' description Literal', first_response=10, inactivity=1, maximum=30, receive=received.append)
    assert ''.join(received) == 'not a prompt\nweird#\n€'
    assert cli._channel.sent == [b' description Literal\n']


@pytest.mark.parametrize('chunks', [[socket.timeout()], [b''], [OSError()]])
def test_literal_failures_invalidate_without_retry(chunks):
    cli = literal_cli(chunks)
    with pytest.raises(Exception):
        cli.capture_literal('shutdown', first_response=10, inactivity=1, maximum=30, receive=lambda _: None)
    assert not cli._synchronized and len(cli._channel.sent) == 1


def test_profile_legacy_hash_roundtrip_and_invalid_hash(tmp_path, auth):
    from orbitflow.environment import validate, save_profile, load_profile, PATHS
    profile = dict(teleport=dict(proxy='proxy', cluster='cluster', bastion_host='bastion', bastion_user='operator'),
                   execution=dict(connection_start_interval=0, connection_retry_delay=0), paths=PATHS.copy())
    validate(profile)
    with pytest.raises(ValueError):
        authorize(profile, 'anything')
    profile.update(auth)
    path = tmp_path / 'profile.toml'
    save_profile(profile, path)
    authorize(load_profile(path), 'test authorization password')
    assert 'test authorization password' not in path.read_text()
    for value in ['plaintext', 'pbkdf2-sha256$1$00$00', 'pbkdf2-sha256$999999999$00$00']:
        profile['manual_configuration_auth'] = {'password_hash': value}
        with pytest.raises(ValueError):
            validate(profile)


@pytest.mark.parametrize('platform, prompt, entry, global_prompt', [
    ('cisco_ios', 'host#', 'configure terminal', 'host(config)#'),
    ('cisco_xe', 'host#', 'configure terminal', 'host(config)#'),
    ('ubiquiti_edgeswitch', '(host) #', 'configure', '(host) (Config)#')])
def test_vendor_setup_reuses_profile_and_global_confirmation(platform, prompt, entry, global_prompt):
    from orbitflow.vendors import manual_configuration as vendor
    calls = []
    class CLI:
        def __init__(self):
            self.prompt = prompt
        def configure(self, **kwargs):
            calls.append(kwargs)
        def execution_mode(self, anchor, selected):
            assert anchor == prompt and selected == platform
        def run_command(self, command, **kwargs):
            assert command == entry
            self.prompt = global_prompt
            return ''
    assert vendor.setup(CLI(), SimpleNamespace(platform=platform), 30) == prompt
    assert calls[0]['paging_command'] == 'terminal length 0'
    if platform.startswith('cisco'):
        assert calls[0]['setup_commands'] == ('terminal width 511',)


def test_independent_device_continues_after_one_device_fails(fake, auth, monkeypatch):
    calls, fault = fake
    original = m.DeviceInventoryResolver.resolve
    def resolve(self, session, **kwargs):
        context = original(self, session, **kwargs)
        if kwargs['management_ip'] == '192.0.2.1':
            return SimpleNamespace(platform='cisco_xr', hostname='unsupported')
        return context
    monkeypatch.setattr(m.DeviceInventoryResolver, 'resolve', resolve)
    rows = [row(2), row(3, '192.0.2.2'), row(4)]
    path = run(rows, auth)
    outcomes = {r['source_row']: r for r in records(path)}
    assert outcomes[2]['status'] == 'Incomplete'
    assert outcomes[4]['attempted'] == 0
    assert outcomes[3]['status'] == 'Completed'
    manifest = json.loads((path / 'manifest.json').read_text())
    assert manifest['failed_count'] == 1


def test_transcript_failure_no_save_or_later_commands(fake, auth, monkeypatch):
    calls, _ = fake
    original = m.os.fsync
    original_open = Path.open
    transcript_streams = []
    def tracked_open(path, *args, **kwargs):
        stream = original_open(path, *args, **kwargs)
        if path.suffix == '.txt' and 'reports' in path.parts:
            transcript_streams.append(stream)
        return stream
    def fail(descriptor):
        if any(not stream.closed and stream.fileno() == descriptor for stream in transcript_streams):
            raise OSError('synthetic storage error')
        return original(descriptor)
    monkeypatch.setattr(Path, 'open', tracked_open)
    monkeypatch.setattr(m.os, 'fsync', fail)
    path = run([row(2), row(3)], auth)
    assert not any(c[0] in {'command', 'save'} for c in calls)
    assert all(r['status'] == 'Incomplete' for r in records(path))


def test_delayed_response_crosses_commands_and_fragmented_rejection(fake, auth, monkeypatch):
    calls, _ = fake
    def capture(self, command, **kwargs):
        calls.append(('command', self.serial, command))
        if command.startswith('interface'):
            kwargs['receive']('host(config-if)#\n% In')
        else:
            kwargs['receive']('valid input\n')
    monkeypatch.setattr(m.DeviceCLI, 'capture_literal', capture)
    path = run([row(2), row(3)], auth)
    assert not any(c[0] == 'save' for c in calls)
    assert records(path)[0]['attempted'] == 2
    assert '% In' in Path(records(path)[0]['transcript']).read_text()


def test_receive_inactivity_and_overall_separate_deadlines(monkeypatch):
    from orbitflow.vendors import common
    ticks = iter([0, 0, 1, 30])
    monkeypatch.setattr(common.time, 'monotonic', lambda: next(ticks))
    cli = literal_cli([b'continuous output'])
    with pytest.raises(common.InteractiveCLITimeout, match='overall deadline'):
        cli.capture_literal('shutdown', first_response=10, inactivity=1, maximum=30, receive=lambda _: None)
    assert len(cli._channel.sent) == 1


def test_sigint_marks_workers_cancelled_before_cleanup():
    import signal
    from threading import Event
    cancelled = Event()
    original = signal.getsignal(signal.SIGINT)
    with pytest.raises(KeyboardInterrupt):
        with m.interrupt_guard(cancelled):
            signal.getsignal(signal.SIGINT)(signal.SIGINT, None)
    assert cancelled.is_set()
    assert signal.getsignal(signal.SIGINT) is original


def test_new_planning_input_omits_legacy_without_weakening_plan_schema(tmp_path):
    from orbitflow.configuration.v2.workbook import create_input, read_input
    from orbitflow.configuration.v2.planning import plan
    from orbitflow.configuration.v2.catalogue import CATALOGUE
    path = tmp_path / 'planning.xlsx'
    create_input(path)
    book = load_workbook(path)
    assert 'Manual CLI' not in book.sheetnames
    assert len(book.sheetnames) == 22 and len(CATALOGUE) == 22
    assert 'Manual CLI' not in str(list(book['Introduction'].values)[7:])
    requests = read_input(path)
    assert requests.layout[-1] == 'Manual CLI'  # Canonical historical v2 shape.
    artifact = plan(requests)
    assert len(artifact.to_dict()['layout']) == 23


def test_console_decline_and_bad_password_before_session(fake, auth, monkeypatch):
    from orbitflow.menu import Console, Cancelled
    from orbitflow.environment import PATHS
    Path('manual_commands.xlsx').touch()
    Path('devices.xlsx').touch()
    profile = dict(teleport={}, execution=dict(connection_start_interval=0, connection_retry_delay=0), paths=PATHS.copy(), **auth)
    monkeypatch.setattr(Console, 'settings', lambda self: profile)
    monkeypatch.setattr(m, 'load_rows', lambda path: [row(2)])
    import getpass
    for password, expected in [('wrong', ValueError), ('test authorization password', Cancelled)]:
        answers = iter(['1', 'manual_commands.xlsx', 'devices.xlsx', '', 'n'])
        shown, prompts = [], []
        def read(prompt):
            prompts.append(prompt)
            return next(answers)
        console = Console(read=read, write=shown.append,
                          session_check=lambda _: pytest.fail('gate connected'))
        def hidden(prompt):
            assert any('Manual configuration: rows=1' in line for line in shown)
            assert not any(word in '\n'.join(shown) for word in ['Fresh shell', 'Transcript only', 'Not saving', 'Whole-config', 'Never replay'])
            return password
        monkeypatch.setattr(getpass, 'getpass', hidden)
        with pytest.raises(expected):
            console.manual_execute()
        if password != 'wrong':
            assert 'Confirm manual execution? [y/N]: ' in prompts
        assert not fake[0]


def test_manual_save_allows_identified_platform_and_checks_errors_before_redaction():
    from orbitflow.capabilities.save import save_configuration
    context = SimpleNamespace(platform='cisco_ios', device_family='Other IOS', capability_profile='',
                              management_ip='192.0.2.1', hostname='host')
    session = object()
    class CLI:
        def require_session(self, owner):
            assert owner is session
        def save_configuration(self, anchor, platform, timeout):
            return dict(status='Saved', response='secret hidden\n% Error saving\nhost#')
    with pytest.raises(ValueError, match='unsupported save family'):
        save_configuration(session, context, cli=CLI(), exec_prompt='host#', timeout=30)
    result = save_configuration(session, context, cli=CLI(), exec_prompt='host#', timeout=30, manual=True)
    assert result['explicit_error']
    assert 'hidden' not in result['response']


def test_manual_concurrency_uses_shared_workers_with_fresh_shells(fake, auth):
    calls, _ = fake
    rows = [row(i + 2, f'192.0.2.{i % 5 + 1}') for i in range(20)]
    config = m.ExecutionConfig(5, 0, 0)
    path = m.execute(rows, profile=auth, password='test authorization password',
                     confirmation_digest=m.summary(rows, config)['digest'],
                     credentials={r['management_ip']: DeviceCredentials('operator', 'device-private') for r in rows},
                     transport_config=None, execution_config=config)
    result = records(path)
    assert len(result) == 20 and all(r['status'] == 'Completed' for r in result)
    for ip in {r['management_ip'] for r in rows}:
        assert [r['source_row'] for r in result if r['management_ip'] == ip] == [r['source_row'] for r in rows if r['management_ip'] == ip]
    assert len([c for c in calls if c[0] == 'shell']) == 20
    assert len([c for c in calls if c[0] == 'close_shell']) == 20
    assert len([c for c in calls if c[0] == 'save']) == 20


def test_password_set_change_and_environment_preservation(fake, auth, monkeypatch):
    from orbitflow.menu import Console, OperatorError
    from orbitflow.environment import PATHS, save_profile, load_profile
    import getpass
    profile = dict(teleport=dict(proxy='proxy', cluster='cluster', bastion_host='bastion', bastion_user='operator'),
                   execution=dict(connection_start_interval=0, connection_retry_delay=0), paths=PATHS.copy())
    save_profile(profile)
    responses = iter(['new test password value', 'new test password value'])
    monkeypatch.setattr(getpass, 'getpass', lambda _: next(responses))
    console = Console(write=lambda _: None)
    console.manual_password()
    authorize(load_profile(), 'new test password value')
    responses = iter(['wrong current password'])
    with pytest.raises(ValueError):
        console.manual_password()
    authorize(load_profile(), 'new test password value')
    answers = iter([''] * 13)
    console.read = lambda _: next(answers)
    console.setup()
    authorize(load_profile(), 'new test password value')


def test_backup_persistence_failure_blocks_writes(fake, auth, monkeypatch):
    calls, _ = fake
    def fail(*args):
        raise OSError('synthetic storage failure')
    monkeypatch.setattr(m.BackupWriter, 'write', fail)
    path = run([row(2), row(3)], auth)
    assert not any(c[0] in {'command', 'save'} for c in calls)
    assert records(path)[1]['attempted'] == 0


def test_shell_cleanup_failure_after_save_blocks_later_rows(fake, auth, monkeypatch):
    calls, _ = fake
    def fail(self, *args):
        calls.append(('close_shell', self.serial))
        raise OSError('synthetic close failure')
    monkeypatch.setattr(m.DeviceCLI, '__exit__', fail)
    path = run([row(2), row(3)], auth)
    assert len([c for c in calls if c[0] == 'save']) == 1
    assert records(path)[0]['status'] == 'Incomplete'
    assert records(path)[0]['save'] == 'Saved'
    assert records(path)[1]['attempted'] == 0
    assert len([c for c in calls if c[0] == 'close_session']) == 1


def test_report_write_failure_retains_offline_recovery(fake, auth, monkeypatch):
    calls, _ = fake
    path = run([row(2)], auth)
    original = m.Workbook.save
    def fail(*args):
        raise OSError('synthetic export failure')
    monkeypatch.setattr(m.Workbook, 'save', fail)
    with pytest.raises(OSError):
        m.regenerate_report(path, 'failed-report.xlsx')
    assert path.exists() and not Path('failed-report.xlsx').exists()
    monkeypatch.setattr(m.Workbook, 'save', original)
    calls.clear()
    assert m.regenerate_report(path, 'recovered-report.xlsx').exists()
    assert calls == []


@pytest.mark.parametrize('command', ['snmp-server host 192.0.2.1 version 2c hidden',
    'description encrypted hidden', 'description cipher hidden', 'description pre-shared hidden'])
def test_vendor_credential_syntax_rejected(command):
    with pytest.raises(ValueError):
        m.commands(command)


@pytest.mark.parametrize('username', ['source_row', 'input', 'row_result'])
def test_short_login_values_cannot_corrupt_persisted_schema(fake, auth, username):
    # A valid login string may happen to equal an internal field name.
    config = m.ExecutionConfig(1, 0, 0)
    rows = [row(2)]
    path = m.execute(rows, profile=auth, password='test authorization password',
                     confirmation_digest=m.summary(rows, config)['digest'],
                     credentials={'192.0.2.1': DeviceCredentials(username, 'private-fixture')},
                     transport_config=None, execution_config=config)
    assert records(path)[0]['source_row'] == 2
    assert m.regenerate_report(path).exists()

@pytest.mark.parametrize('echo', [
    b'host(config-if)#end\r\nhost#',
    b'host(config-subif)#end\r\nhost#',
    b'host(config-if-srv)#end\r\nhost#',
    b'newhost(config-if)#end\r\nnewhost#',
    b'end\r\nnewhost#',
    b'FAILED(config-if)#end\r\nFAILED#',
])
def test_real_reader_manual_finalize_and_save(echo):
    from orbitflow.vendors import manual_configuration as vendor
    cli = literal_cli([echo[:8], echo[8:-2], echo[-2:], socket.timeout()])
    cli.prompt = 'host(config)#'
    cli.execution_mode('host#', 'cisco_ios')
    anchor = vendor.finalize(cli, SimpleNamespace(platform='cisco_ios'), 2)
    assert anchor in {'host#', 'newhost#', 'FAILED#'}
    cli._channel.chunks = iter([(anchor + 'write memory\r\n[OK]\r\n' + anchor).encode()])
    result = cli.save_configuration(anchor, 'cisco_ios', timeout=2)
    assert result['status'] == 'Saved'
    assert cli._channel.sent == [b'end\n', b'write memory\n']


@pytest.mark.parametrize('chunks', [
    [b'host#', socket.timeout()],  # stale prompt without current echo
    [b'host(config-if)#end\n% Invalid input\nhost#', socket.timeout()],
    [b'host(config-if)#end\nhost#', b' trailing stale output', socket.timeout()],
    [b'host(config-if)#end\nhost#', b'\n% Inva', b'lid input\nhost#', socket.timeout()],
    [b'host(config-if)#end\n', b''],
])
def test_real_reader_finalize_failure_never_saves(chunks):
    from orbitflow.vendors import manual_configuration as vendor
    cli = literal_cli(chunks)
    cli.prompt = 'host(config)#'
    cli.execution_mode('host#', 'cisco_ios')
    with pytest.raises(Exception):
        vendor.finalize(cli, SimpleNamespace(platform='cisco_ios'), 2)
    assert cli._channel.sent == [b'end\n']


@pytest.mark.parametrize('text', ['description FAILED LINK', 'description Customer Error Monitor',
                                'hostname failed', 'hostname error'])
def test_real_capture_echo_words_are_not_rejections(text):
    from orbitflow.vendors.manual_configuration import rejected
    cli = literal_cli([(text + '\r\nerror(config)#').encode(), socket.timeout()])
    raw = []
    def receive(chunk):
        raw.append(chunk)
        assert not rejected(m.normalize_output(''.join(raw)))
    cli.capture_literal(text, first_response=10, inactivity=1, maximum=30, receive=receive)
    assert text in ''.join(raw)


@pytest.mark.parametrize('diagnostic', ['% Invalid input detected', 'Error: command rejected',
                                      'Failed to set VLAN tagging.', 'Command not found',
                                      '% Unrecognized command', 'Command authorization failed.', 'Access denied'])
def test_real_capture_fragmented_diagnostics_after_echo(diagnostic):
    from orbitflow.vendors.manual_configuration import rejected
    echo = b'description Customer Error Monitor\r\n'
    cli = literal_cli([echo, diagnostic[:3].encode(), diagnostic[3:].encode(), socket.timeout()])
    raw = []
    def receive(chunk):
        raw.append(chunk)
        if rejected(m.normalize_output(''.join(raw))):
            raise ValueError('Vendor rejected command')
    with pytest.raises(ValueError, match='Vendor rejected'):
        cli.capture_literal('description Customer Error Monitor', first_response=10, inactivity=1,
                            maximum=30, receive=receive)
    assert not cli._synchronized


@pytest.mark.parametrize('failure', ['', 'rejection', 'interaction', 'disconnect', 'stale'])
def test_application_real_reader_finalize_stops_later_rows(fake, auth, monkeypatch, failure):
    """Run application row policy with real capture, synchronization and save readers."""
    from orbitflow.vendors import manual_configuration as vendor
    calls, _ = fake
    shells = []
    class RealCLI(DeviceCLI):
        def __init__(self, session, **kwargs):
            self.session = session
            self._synchronized = True
            self._accepted_paging = set()
            self.prompt = 'host(config)#'
            response = {
                'rejection': b'newhost(config-if)#end\r\n% Invalid input\r\nnewhost#',
                'interaction': b'newhost(config-if)#end\r\nPassword:\r\nnewhost#',
                'stale': b'newhost#',
            }.get(failure, b'newhost(config-if)#end\r\nnewhost#')
            self._channel = Channel([
                b'description FAILED LINK\r\nnewhost(config-if)#', socket.timeout(),
                response[:12], b'' if failure == 'disconnect' else response[12:], socket.timeout(),
                b'newhost#write memory\r\n[OK]\r\nnewhost#'])
            self._channel.close = lambda: calls.append(('real_close',))
            self.execution_mode('host#', 'cisco_ios')
            self.sent = self._channel.sent
            shells.append(self)
    monkeypatch.setattr(m, 'DeviceCLI', RealCLI)
    monkeypatch.setattr(vendor, 'finalize', real_manual_finalize)
    def save(session, context, *, cli, exec_prompt, **kwargs):
        return dict(cli.save_configuration(exec_prompt, context.platform, timeout=2), explicit_error=False)
    monkeypatch.setattr(m, 'save_configuration', save)
    path = run([row(2, text='description FAILED LINK'), row(3, text='description FAILED LINK')], auth)
    results = records(path)
    assert len(shells) == (1 if failure else 2)
    assert [r['status'] for r in results] == (['Incomplete', 'Incomplete'] if failure else ['Completed', 'Completed'])
    for cli in shells:
        assert cli._channel is None
        assert cli.sent.count(b'write memory\n') == (0 if failure else 1)
    assert calls.count(('real_close',)) == len(shells)


def test_real_reader_edgeswitch_manual_finalize_save_question():
    from orbitflow.vendors import manual_configuration as vendor
    cli = literal_cli([b'(newhost) (Interface 0/3)#en', b'd\r\n(newhost)#', socket.timeout(),
                       b'(newhost)#write memory\r\nAre you sure you want to save? (y/n)',
                       b'y\r\n(newhost)#'])
    cli.prompt = '(host) (Config)#'
    cli.execution_mode('(host)#', 'ubiquiti_edgeswitch')
    anchor = vendor.finalize(cli, SimpleNamespace(platform='ubiquiti_edgeswitch'), 2)
    assert cli.save_configuration(anchor, 'ubiquiti_edgeswitch', timeout=2)['status'] == 'Saved'
    assert cli._channel.sent == [b'end\n', b'write memory\n', b'y\n']


def test_blank_separators_and_source_lines(tmp_path):
    text = '\r\n interface Gi1\r\n   \r\n description Customer\r\n\r\n'
    assert m.commands(text) == [' interface Gi1', ' description Customer']
    book = Workbook()
    book.active.append(['management_ip', 'command'])
    book.active.append(['192.0.2.1', text])
    book.active.append(['192.0.2.2', ' \n\n'])
    path = tmp_path / 'manual_input.xlsx'
    book.save(path)
    rows = m.load_rows(path)
    assert rows[0]['command_lines'] == [2, 4]
    assert rows[1]['error'] == m.EMPTY_ERROR
    assert m.summary(rows)['commands'] == 2


def test_manual_chooser_filter_ties_reprompt_and_cancel(tmp_path):
    from orbitflow.menu import Console, Cancelled
    import os
    for name in ['manual_b.xlsx', 'manual_a.xlsx', 'config_template.xlsx', '~$manual_x.xlsx']:
        path = tmp_path / name
        path.touch()
        os.utime(path, ns=(100, 100))
    (tmp_path / 'manual_directory.xlsx').mkdir()
    answers = iter(['99', '1'])
    shown = []
    console = Console(read=lambda _: next(answers), write=shown.append)
    assert console.choose_workbook(tmp_path, prefix='manual_').name == 'manual_a.xlsx'
    assert not any('config_template' in line or '~$' in line for line in shown)
    answers = iter(['3', str(tmp_path / 'config_template.xlsx'), '3', str(tmp_path / 'manual_b.xlsx')])
    assert console.choose_workbook(tmp_path, prefix='manual_').name == 'manual_b.xlsx'
    empty = tmp_path / 'empty'
    empty.mkdir()
    answers = iter(['b'])
    with pytest.raises(Cancelled):
        console.choose_workbook(empty, prefix='manual_')


def test_fast_device_publishes_before_slow_failure(fake, auth, monkeypatch):
    from threading import Event
    from urllib.parse import unquote
    import shutil
    calls, _ = fake
    published = Event()
    original_connect = m.connect_device
    original_capture = m.DeviceCLI.capture_literal
    def connect(ip, *args):
        session = original_connect(ip, *args)
        session.ip = ip
        return session
    def capture(self, command, **kwargs):
        if self.session.ip == '192.0.2.2':
            assert published.wait(5), 'fast report waited for slow device'
            raise KeyboardInterrupt()
        for chunk in ['echo\r', '\nfrag', 'mented device-private\r', '\n']:
            kwargs['receive'](chunk)
    monkeypatch.setattr(m, 'connect_device', connect)
    monkeypatch.setattr(m.DeviceCLI, 'capture_literal', capture)
    outputs = []
    def output(path):
        book = load_workbook(path)
        sheet = book.active
        ip = sheet['B2'].value
        assert path.name.startswith('manual_output_') and len(path.stem) == 29
        capture = path.parent / unquote(sheet['J2'].hyperlink.target)
        assert capture.is_file() and 'reports' in capture.parts
        if ip == '192.0.2.1':
            assert sheet.max_row == 4
            assert sheet['J4'].hyperlink.target == sheet['J2'].hyperlink.target
            assert [sheet.cell(i, 1).value for i in range(2, 5)] == ['2', '3', '4']
            assert capture.name == 'host_192.0.2.1.txt'
            if not outputs:
                assert sheet['C3'].value == 'Incomplete'
                assert book['Devices']['B3'].value != 'Completed'
            text = capture.read_text()
            assert 'echo\nfragmented [REDACTED]\n' in text
            assert not any(label in text for label in ['SEND', 'RECEIVE', 'ROW', 'row=', 'description Customer'])
            relocated = Path('relocated') / path.parent.name
            shutil.copytree(path.parent, relocated, dirs_exist_ok=True)
            assert (relocated / unquote(load_workbook(relocated / path.name).active['J2'].hyperlink.target)).is_file()
            published.set()
        outputs.append(path)
    rows = [row(2), row(3, '192.0.2.2'), row(4)]
    config = m.ExecutionConfig(2, 0, 0)
    path = m.execute(rows, profile=auth, password='test authorization password',
        confirmation_digest=m.summary(rows, config)['digest'],
        credentials={r['management_ip']: DeviceCredentials('operator', 'device-private') for r in rows},
        transport_config=None, execution_config=config, on_output=output)
    assert published.is_set() and len(outputs) == 2
    assert outputs[0] == outputs[1]
    assert len(list(outputs[0].parent.glob('*.xlsx'))) == 1
    assert next(r for r in records(path) if r['source_row'] == 3)['status'] == 'Incomplete'


def test_per_device_export_failure_retains_other_output(fake, auth, monkeypatch):
    original = m.regenerate_report
    refreshes = []
    def export(path, *args, **kwargs):
        if kwargs.get('replace'):
            refreshes.append(path)
        if len(refreshes) == 1:
            raise OSError('synthetic export failure')
        return original(path, *args, **kwargs)
    monkeypatch.setattr(m, 'regenerate_report', export)
    outputs = []
    path = run([row(2), row(3, '192.0.2.2')], auth, on_output=outputs.append)
    assert len(outputs) == 1 and outputs[0].exists()
    events = list(m.ResultJournal(path / 'events.jsonl').records())
    assert any(e['kind'] == 'report_export_failed' for e in events)
    calls, _ = fake
    calls.clear()
    assert original(path).exists() and not calls


def test_legacy_missing_and_real_capture_links(tmp_path):
    path = tmp_path / 'legacy'
    path.mkdir()
    journal = m.ResultJournal(path / 'events.jsonl')
    journal.append(dict(kind='input', rows=[dict(source_row=2, management_ip='192.0.2.1', total=1)]))
    result = dict(kind='row_result', source_row=2, management_ip='192.0.2.1', transcript='legacy.txt')
    journal.append(result)
    report = m.regenerate_report(path, tmp_path / 'missing.xlsx')
    assert load_workbook(report).active['J2'].hyperlink is None
    (path / 'legacy.txt').write_text('device#')
    report = m.regenerate_report(path, tmp_path / 'real.xlsx')
    assert load_workbook(report).active['J2'].hyperlink.target == 'legacy/legacy.txt'


def test_short_collision_names_and_windows_cross_drive_fallback(fake, auth, monkeypatch):
    from datetime import datetime
    from urllib.parse import unquote
    class FixedDate(datetime):
        @classmethod
        def now(cls, *args):
            return cls(2026, 10, 26, 12, 34, 56)
    path = run([row(2)], auth)
    monkeypatch.setattr(m, 'datetime', FixedDate)
    first = m.regenerate_report(path)
    second = m.regenerate_report(path)
    assert first.name == 'manual_output_20261026_123456.xlsx'
    assert second.name == 'manual_output_20261026_123456_001.xlsx'
    def cross_drive(*args):
        raise ValueError('different drives')
    monkeypatch.setattr(m.os.path, 'relpath', cross_drive)
    third = m.regenerate_report(path)
    link = load_workbook(third).active['J2'].hyperlink.target
    assert link.startswith('file:///')
    assert Path(unquote(link[7:])).is_file()


def test_all_empty_block_never_connects(fake, auth, tmp_path):
    book = Workbook()
    book.active.append(['management_ip', 'command'])
    book.active.append(['192.0.2.1', ' \r\n\t\r\n'])
    source = tmp_path / 'manual_empty.xlsx'
    book.save(source)
    rows = m.load_rows(source)
    path = run(rows, auth)
    assert records(path)[0]['reason'] == m.EMPTY_ERROR
    calls, _ = fake
    assert not calls


def test_manual_chooser_newest_first(tmp_path):
    import os
    from orbitflow.menu import Console
    for name, stamp in [('manual_a.xlsx', 100), ('manual_z.xlsx', 200)]:
        path = tmp_path / name
        path.touch()
        os.utime(path, ns=(stamp, stamp))
    console = Console(read=lambda _: '1', write=lambda _: None)
    assert console.choose_workbook(tmp_path, prefix='manual_').name == 'manual_z.xlsx'


def test_atomic_refresh_preserves_existing_workbook(fake, auth, monkeypatch):
    path = run([row(2)], auth)
    report = next(Path('outputs/reports/manual_configuration').rglob('*.xlsx'))
    previous = report.read_bytes()
    def held(*args):
        raise PermissionError('Excel holds destination')
    monkeypatch.setattr(m.os, 'replace', held)
    with pytest.raises(PermissionError):
        m.regenerate_report(path, report, replace=True)
    assert report.read_bytes() == previous
    assert load_workbook(report).active['C2'].value == 'Completed'
    assert not list(report.parent.glob('.manual-*'))


@pytest.mark.parametrize('hostname', ['new/name:unsafe', 'changed-host', ''])
def test_named_capture_and_identity_fallback(fake, auth, monkeypatch, hostname):
    def resolve(*args, **kwargs):
        return SimpleNamespace(platform='cisco_ios', hostname=hostname)
    monkeypatch.setattr(m.DeviceInventoryResolver, 'resolve', resolve)
    path = run([row(2, '2001:db8::1')], auth)
    report = next(Path('outputs/reports/manual_configuration').rglob('*.xlsx'))
    from urllib.parse import unquote
    link = load_workbook(report).active['J2'].hyperlink.target
    capture = report.parent / unquote(link)
    assert capture.name == m.filename_component(hostname or '2001:db8::1') + '_2001_db8__1.txt'
    assert capture.is_file()
    assert len(path.name) == 15
    assert (Path('outputs/backups/configuration') / path.name).is_dir()


def test_capture_rename_failure_retains_fallback(fake, auth, monkeypatch):
    original = Path.replace
    def fail(self, target):
        if self.suffix == '.txt':
            raise PermissionError('held capture')
        return original(self, target)
    monkeypatch.setattr(Path, 'replace', fail)
    path = run([row(2)], auth)
    report = next(Path('outputs/reports/manual_configuration').rglob('*.xlsx'))
    capture = report.parent / load_workbook(report).active['J2'].hyperlink.target
    assert capture.name == '192.0.2.1_192.0.2.1.txt'
    assert capture.is_file() and records(path)[0]['status'] == 'Completed'


def test_short_spool_allocation_is_exclusive(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    def allocate(_):
        return m.ResultSpool.create(tmp_path, 'manual_configuration', 0, short_name=True).path
    with ThreadPoolExecutor(max_workers=4) as pool:
        paths = list(pool.map(allocate, range(8)))
    assert len(set(paths)) == 8
    assert all(__import__('re').fullmatch(r'\d{8}_\d{6}(?:_\d{3})?', p.name) for p in paths)


def test_hostname_changes_update_all_repeated_row_links(fake, auth, monkeypatch):
    names = iter(['first-host', 'second-host'])
    monkeypatch.setattr(m.DeviceInventoryResolver, 'resolve', lambda *a, **k:
        SimpleNamespace(platform='cisco_ios', hostname=next(names)))
    path = run([row(2), row(3)], auth)
    report = next(Path('outputs/reports/manual_configuration').rglob('*.xlsx'))
    sheet = load_workbook(report).active
    assert sheet['J2'].hyperlink.target == sheet['J3'].hyperlink.target == 'second-host_192.0.2.1.txt'
    assert len(list(report.parent.glob('*.txt'))) == 1
    assert m.regenerate_report(path).exists()


def test_identity_failure_links_ip_fallback(fake, auth):
    fake[1]['stage'] = 'identify'
    path = run([row(2)], auth)
    report = next(Path('outputs/reports/manual_configuration').rglob('*.xlsx'))
    sheet = load_workbook(report).active
    assert sheet['J2'].hyperlink.target == '192.0.2.1_192.0.2.1.txt'
    assert records(path)[0]['save'] == 'Not initiated'


def test_capture_name_collision_preserves_existing_file(fake, auth, monkeypatch):
    # Create collision just after the run's report directory is allocated.
    original = m.regenerate_report
    def export(path, *args, **kwargs):
        output = original(path, *args, **kwargs)
        if not kwargs.get('replace'):
            (output.parent / 'host_192.0.2.1.txt').write_text('existing capture')
        return output
    monkeypatch.setattr(m, 'regenerate_report', export)
    path = run([row(2)], auth)
    report = next(Path('outputs/reports/manual_configuration').rglob('*.xlsx'))
    assert (report.parent / 'host_192.0.2.1.txt').read_text() == 'existing capture'
    assert load_workbook(report).active['J2'].hyperlink.target == '192.0.2.1_192.0.2.1.txt'
