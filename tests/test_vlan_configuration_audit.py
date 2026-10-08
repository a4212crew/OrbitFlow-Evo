"""Saved-configuration audit contracts; entirely synthetic, no network access."""

from dataclasses import asdict, replace
import json

import pytest

from orbitflow.compliance import evaluate_vlan_compliance
from orbitflow.vendors.configuration_facts import observe_configuration
from orbitflow.vendors.cisco.vlans import parse_ios_running_config, parse_ios_xr_running_config
from orbitflow.vendors.huawei.vlans import parse_huawei_config
from orbitflow.vendors.ubiquiti.vlans import parse_edgeswitch_config
from orbitflow.models import VlanState
from test_vlan_compliance import context, interface, POLICY, NOW


FAMILIES = {"C3750X": "cisco_ios", "C3850": "cisco_xe", "ME3600X": "cisco_ios",
            "ASR920": "cisco_xe", "NCS540": "cisco_xr", "NE05E": "huawei_vrp",
            "EdgeSwitch": "ubiquiti_edgeswitch"}


def audit(config, family="C3850", observed=(), *, filename="synthetic.cfg"):
    platform = "huawei_vrp" if family == "NE05" else FAMILIES[family]
    ctx = replace(context(), platform=platform, device_family=family)
    if platform in {"cisco_ios", "cisco_xe"}:
        profiles, objects = parse_ios_running_config(config, evc=family in {"ME3600X", "ASR920"}, vlan_database=family != "ASR920")
    else:
        profiles, objects = {"cisco_xr": parse_ios_xr_running_config, "huawei_vrp": parse_huawei_config,
                             "ubiquiti_edgeswitch": parse_edgeswitch_config}[platform](config)
    snapshot = VlanState("switch", ctx.management_ip, platform, profiles, objects, NOW,
                         observe_configuration(config, platform, source_filename=filename))
    records = [replace(interface(name), platform=platform) for name in observed]
    return evaluate_vlan_compliance(ctx, records, snapshot, POLICY)


def row(results, name):
    return next(f for f in results[1:] if f["observed"]["config_interface_name"] == name)


def codes(finding):
    return {p["code"] for p in finding["observed"]["configuration_findings"]}


def test_catalyst_database_evidence_identity_shutdown_and_no_svi_inference():
    config = """vlan 445,545,2449
!
interface Vlan4001
 ip address 192.0.2.1 255.255.255.0
!
interface GigabitEthernet0/1
 description sample
 shutdown
 switchport mode trunk
 switchport trunk allowed vlan 445,545,2449,4001
!
interface Gi0/2
 switchport mode access
 switchport access vlan 445
!
interface GigabitEthernet0/1
 description repeated stanza
!
"""
    results = audit(config, observed=("Gi0/1", "Gi0/99"))
    assert results[0]["observed"]["valid_database_vlans"] == [445, 545, 2449]
    assert results[0]["missing_vlans"] == [*range(2400, 2445), 4001]
    trunk = row(results, "GigabitEthernet0/1")
    assert trunk["interface"] == "Gi0/1"
    assert trunk["observed"]["interface_match_status"] == "matched"
    assert trunk["observed"]["shutdown"] is True
    assert trunk["observed"]["description"] == "repeated stanza"
    assert "ALLOWED_VLAN_NOT_IN_DATABASE" in codes(trunk)
    assert trunk["missing_vlans"] == [*range(2400, 2445), 4001]
    assert row(results, "Gi0/2")["status"] == "not_applicable"
    assert "CONFIG_ONLY_INTERFACE" in codes(row(results, "Gi0/2"))
    assert "Gi0/99" not in json.dumps(results)
    assert "INTERFACE_NOT_IN_CONFIG" not in json.dumps(results)
    for source in trunk["evidence"]["sources"]:
        assert source["source_filename"] == "synthetic.cfg"
        assert config.splitlines()[source["line"] - 1] == source["excerpt"]
    assert len([f for f in results if f["interface"] == "Gi0/1"]) == 1


@pytest.mark.parametrize("family,encapsulation,status", [
    ("C3750X", "", "unable_to_assess"), ("C3750X", " switchport trunk encapsulation dot1q\n", "non_compliant"),
    ("C3850", "", "non_compliant"),
])
def test_all_vlan_uses_database_and_recommends_global_change(family, encapsulation, status):
    result = audit("vlan 445,545,2449\n!\ninterface Gi0/1\n" + encapsulation + " switchport mode trunk\n!", family)
    trunk = row(result, "Gi0/1")
    assert trunk["status"] == status
    assert trunk["observed"]["all_vlan"] is True
    assert "4094" not in json.dumps(trunk)
    if status == "non_compliant":
        assert trunk["missing_vlans"] == [*range(2400, 2445), 4001]
        assert "global VLANs" in trunk["recommendation"]


@pytest.mark.parametrize("command", ["switchport mode dynamic", "switchport trunk allowed vlan 445,545,2449",
                                      "switchport mode trunk\n switchport trunk allowed vlan add"])
def test_unspecified_dynamic_and_unsupported_allowed_never_pass(command):
    result = audit("vlan 445,545,2449\n!\ninterface Gi0/1\n " + command + "\n!")
    assert row(result, "Gi0/1")["status"] == "unable_to_assess"


def test_port_channel_ownership_and_conflict():
    result = audit("""vlan 445,545,2449
!
interface Port-channel1
 switchport mode trunk
!
interface Gi0/1
 channel-group 1 mode active
!
interface Gi0/2
 switchport mode access
 switchport access vlan 445
 channel-group 1 mode active
!""")
    member = row(result, "Gi0/1")
    assert member["observed"]["configuration_owner"] == "Port-channel1"
    assert member["observed"]["valid_interface_vlans"] == [445, 545, 2449]
    assert "PORT_CHANNEL_CONFIG_CONFLICT" in codes(row(result, "Gi0/2"))
    assert "global VLANs" in member["recommendation"]


def evc_config():
    return """vlan 445,545
!
interface Gi0/1
 service instance 10 ethernet
  encapsulation dot1q 3999 second-dot1q 200
  bridge-domain 445
 service instance 20 ethernet
  encapsulation dot1q 3998
 service instance 30 ethernet
  encapsulation dot1q 3997
  bridge-domain 2449
 service instance 40 ethernet
  encapsulation dot1q 4001
 service instance 50 ethernet
  encapsulation dot1q 2400
  bridge-domain 2400
!
bridge-domain 545
 member GigabitEthernet0/1 service-instance 20
!
bridge-domain 2449
 member Gi0/1 service-instance 30
!
bridge-domain 2440
 member Gi0/1 service-instance 50
!
bridge-domain 2222
 member Gi0/99 service-instance 1
!
bridge-domain 3333
!"""


@pytest.mark.parametrize("family,expected_db", [
    ("ME3600X", [445, 545]), ("ASR920", [445, 545, 2222, 2400, 2440, 2449, 3333]),
])
def test_evc_inline_global_conflict_unresolved_and_database(family, expected_db):
    results = audit(evc_config(), family)
    assert results[0]["observed"]["valid_database_vlans"] == expected_db
    port = row(results, "Gi0/1")
    assert port["observed"]["valid_interface_vlans"] == [445, 545, 2449]
    assert port["missing_vlans"] == [*range(2400, 2445), 4001]
    assert {"CONFLICTING_BRIDGE_DOMAIN_BINDING", "UNRESOLVED_SERVICE_INSTANCE"} <= codes(port)
    assert "INVALID_BRIDGE_DOMAIN_MEMBER_REFERENCE" in codes(results[0])
    assert ("BRIDGE_DOMAIN_MISSING_GLOBAL_VLAN" in codes(results[0])) == (family == "ME3600X")
    assert len(results) == 2  # No synthetic Gi0/99 from global membership.
    service = port["observed"]["numeric_mappings"][0]
    assert service["outer_vlan"] == [3999] and service["inner_vlan"] == [200]
    assert service["audit_vlan"] == [445] and service["service_instance_id"] == "10"
    assert service["binding_status"] == "valid"


@pytest.mark.parametrize("untagged_count,tagged,kind", [(1, False, "access"), (0, True, "evc"),
                                                      (1, True, "hybrid"), (2, True, "review")])
def test_asr_tagged_untagged_classification(untagged_count, tagged, kind):
    config = "interface Gi0/1\n"
    for sid in range(untagged_count):
        config += f" service instance {sid+1} ethernet\n  encapsulation untagged\n  bridge-domain 4001\n"
    if tagged:
        for sid, bd in enumerate((445, 545, 2449), 10):
            config += f" service instance {sid} ethernet\n  encapsulation dot1q 3000\n  bridge-domain {bd}\n"
    port = row(audit(config + "!", "ASR920"), "Gi0/1")
    assert port["observed"]["interface_type"] == kind
    if untagged_count == 2:
        assert "MULTIPLE_UNTAGGED_SERVICE_INSTANCES" in codes(port)
        assert 4001 not in port["observed"]["valid_interface_vlans"]
        assert port["status"] == "non_compliant"  # independently valid tagged subset
    elif not tagged:
        assert port["status"] == "not_applicable"


def test_huawei_vsi_validation_parent_consolidation_and_pvid():
    config = """vlan batch 445 545
#
vsi service4001 static
 vsi-id 4001
#
interface GigabitEthernet0/1
 port link-type trunk
 port trunk pvid vlan 100
 port trunk allow-pass vlan 445 545
#
interface GigabitEthernet0/1.987
 description child service
 control-vid 3333 dot1q-termination
 dot1q termination vid 2449
 l2 binding vsi service4001
#
interface GigabitEthernet0/1.4001
 dot1q termination vid 4001
 ip address 192.0.2.1 255.255.255.0
#
interface GigabitEthernet0/1.2400
 dot1q termination vid 2400
 l2 binding vsi absent
#
interface GigabitEthernet0/2
 l2 binding vsi service4001
#
interface Vlanif2444
 ip address 192.0.2.2 255.255.255.0
#"""
    results = audit(config, "NE05E", observed=("GE0/1",))
    assert results[0]["observed"]["valid_database_vlans"] == [445, 545, 2449]
    port = row(results, "GigabitEthernet0/1")
    assert port["observed"]["interface_type"] == "trunk"
    assert port["observed"]["pvid"] == [100]
    assert port["observed"]["valid_interface_vlans"] == [445, 545, 2449]
    assert len(port["observed"]["child_interfaces"]) == 3
    assert {"VSI_REFERENCE_NOT_FOUND", "L2_TERMINATION_WITHOUT_VALID_VSI"} <= codes(port)
    assert port["observed"]["numeric_mappings"][0]["control_vid"] == [3333]
    routed_detail = next(m for m in port["observed"]["numeric_mappings"] if m["binding_status"] == "routed_detail")
    assert routed_detail["termination_vlan"] == [4001] and routed_detail["audit_vlan"] == []
    assert row(results, "GigabitEthernet0/2")["observed"]["valid_interface_vlans"] == []
    assert len(results) == 4


@pytest.mark.parametrize("parent,code", [("", "PARENT_INTERFACE_NOT_FOUND"),
                                         ("interface GE0/1\n port default vlan 100\n#\n", "PARENT_INTERFACE_NOT_TRUNK")])
def test_huawei_parent_validation(parent, code):
    results = audit("vsi example static\n#\n" + parent + "interface GE0/1.10\n dot1q termination vid 2449\n l2 binding vsi example\n#", "NE05E")
    assert results[0]["observed"]["valid_database_vlans"] == [2449]
    assert any(code in codes(f) for f in results[1:])


def xr_config():
    return """interface TenGigE0/0/0/1
 shutdown
!
interface TenGigE0/0/0/1.100 l2transport
 encapsulation dot1q 445
!
interface TenGigE0/0/0/1.200 l2transport
 encapsulation dot1q 545
!
interface TenGigE0/0/0/1.300 l2transport
 encapsulation dot1q 2449
!
interface TenGigE0/0/0/1.4001 l2transport
 encapsulation dot1q 4001
!
interface TenGigE0/0/0/1.700 l2transport
 encapsulation dot1q 2400
!
interface TenGigE0/0/0/1.800 l2transport
 encapsulation untagged
!
interface TenGigE0/0/0/1.900 l2transport
 description no encapsulation
!
interface TenGigE0/0/0/1.2444
 encapsulation dot1q 2444
!
l2vpn
 bridge group arbitrary
  bridge-domain named
   interface TenGigE0/0/0/1.100
   interface TenGigE0/0/0/1.200
   interface TenGigE0/0/0/1.300
   interface TenGigE0/0/0/1.700
   interface TenGigE0/0/0/1.800
   interface TenGigE0/0/0/1.900
   interface TenGigE0/0/0/1.2444
   interface TenGigE0/0/0/99.4001
  bridge-domain 4001
   interface TenGigE0/0/0/1.700
 bridge group other
  bridge-domain named
!"""


def test_xr_numeric_mapping_conflicts_unbound_routed_and_same_name_groups():
    results = audit(xr_config(), "NCS540", observed=("Te0/0/0/1", "Te0/0/0/1.100"))
    assert results[0]["observed"]["valid_database_vlans"] == [445, 545, 2449]
    assert "MISSING_BRIDGE_DOMAIN_INTERFACE_REFERENCE" in codes(results[0])
    assert len(results) == 2
    parent = results[1]
    assert parent["interface"] == "Te0/0/0/1"
    assert parent["observed"]["shutdown"]
    assert parent["missing_vlans"] == [*range(2400, 2445), 4001]
    assert {"UNBOUND_L2TRANSPORT_SUBINTERFACE", "NON_L2TRANSPORT_BRIDGE_DOMAIN_ATTACHMENT",
            "CONFLICTING_BRIDGE_DOMAIN_ATTACHMENTS", "UNTAGGED_BRIDGE_DOMAIN_MAPPING_UNRESOLVED",
            "L2TRANSPORT_WITHOUT_ENCAPSULATION"} <= codes(parent)
    mappings = parent["observed"]["numeric_mappings"]
    conflict = next(m for m in mappings if m["encapsulation"] == [2400])
    assert conflict["bridge_domains"] == [["arbitrary", "4001"], ["arbitrary", "named"]]
    assert conflict["audit_vlan"] == []
    assert {o["object_id"] for o in results[0]["observed"]["database_inventory"]} == {"arbitrary/named", "arbitrary/4001", "other/named"}
    child = next(c for c in parent["observed"]["child_interfaces"] if c["config_interface_name"].endswith(".100"))
    assert child["interface_match_status"] == "matched"
    assert child["admin_status"] == child["oper_status"] == "up"


@pytest.mark.parametrize("reference", ["3/1", "lag 1"])
def test_edgeswitch_order_database_membership_tagging_pvid_and_lag(reference):
    config = """vlan database
vlan 445,545,2449
vlan name 4001 "annotation-only"
exit
interface lag 1
vlan participation include 445,545,2449,4001
vlan participation exclude 545
vlan participation include 545
vlan tagging 445,545,2449,4001
exit
interface 0/1
addport REFERENCE
exit
interface 0/2
description "Access_port"
vlan pvid 100
vlan participation include 445,545,2449
vlan tagging 445,545,2449
addport REFERENCE
exit
interface 0/3
vlan participation include 445,545,2449
vlan tagging 445,545,2449
vlan participation exclude 545
exit
"""
    config = config.replace("REFERENCE", reference)
    results = audit(config, "EdgeSwitch")
    assert results[0]["observed"]["valid_database_vlans"] == [445, 545, 2449]
    assert {o["domain_id"] for o in results[0]["observed"]["database_inventory"]} == {"445", "545", "2449"}
    lag = row(results, "lag 1")
    assert lag["observed"]["configured_tagged_vlans"] == [445, 545, 2449, 4001]
    assert lag["observed"]["tagged"] == [445, 545, 2449]
    assert {"TAGGED_VLAN_NOT_IN_MEMBERSHIP", "INTERFACE_VLAN_NOT_IN_DATABASE"} <= codes(lag)
    assert row(results, "0/1")["observed"]["configuration_owner"] == "lag 1"
    for name in ("0/1", "0/2"):
        member = row(results, name)
        assert member["observed"]["inherited_from"] == "lag 1"
        assert member["observed"]["valid_interface_vlans"] == lag["observed"]["valid_interface_vlans"]
        assert member["status"] == lag["status"] == "non_compliant"
        assert member["missing_vlans"] == lag["missing_vlans"]
        assert "AGGREGATE_NOT_FOUND" not in codes(member)
        assert any(e["excerpt"] == f"addport {reference}" for e in member["evidence"]["sources"])
    assert "AGGREGATE_CONFIG_CONFLICT" not in codes(row(results, "0/1"))
    assert {"PVID_NOT_IN_MEMBERSHIP", "AGGREGATE_CONFIG_CONFLICT"} <= codes(row(results, "0/2"))
    assert row(results, "0/3")["status"] == "not_applicable"


@pytest.mark.parametrize("reference,owner", [("3/2", "lag 2"), ("lag 2", "lag 2"),
                                               ("1", None), ("0/2", None), ("3/9", None)])
def test_edgeswitch_lag_reference_exact_owner_or_review(reference, owner):
    config = f"""vlan database
vlan 445,545,2449
exit
interface 0/1
addport {reference}
exit
interface lag 1
vlan participation include 445
vlan pvid 445
exit
interface lag 2
vlan participation include 445,545,2449
vlan tagging 445,545,2449
exit
"""
    results = audit(config, "EdgeSwitch")
    member = row(results, "0/1")
    assert len(results) == 4  # Never synthesize an interface from a reference.
    if owner:
        assert member["observed"]["configuration_owner"] == owner
        assert member["observed"]["inherited_from"] == owner
        assert member["observed"]["valid_interface_vlans"] == [445, 545, 2449]
        assert "AGGREGATE_NOT_FOUND" not in codes(member)
    else:
        assert member["observed"]["configuration_owner"] == "0/1"
        assert member["observed"]["valid_interface_vlans"] == []
        assert "AGGREGATE_NOT_FOUND" in codes(member)
        assert member["status"] == "unable_to_assess"


@pytest.mark.parametrize("members,tags,pvid,kind", [
    ("445", "", "445", "access"), ("445,545", "545", "445", "hybrid"),
    ("445,545", "445,545", "", "trunk"), ("445,545", "", "445", "access"),
    ("", "445", "445", "no_membership"),
])
def test_edge_classification_no_membership_from_tagging_or_pvid(members, tags, pvid, kind):
    config = "vlan database\nvlan 445,545\nexit\ninterface 0/1\n"
    for command, value in (("vlan participation include", members), ("vlan tagging", tags), ("vlan pvid", pvid)):
        if value:
            config += f"{command} {value}\n"
    port = row(audit(config + "exit", "EdgeSwitch"), "0/1")
    assert port["observed"]["interface_type"] == kind
    if not members:
        assert port["observed"]["valid_interface_vlans"] == []


def test_configuration_evidence_excludes_secrets_and_banner_fake_facts():
    config = """username hidden secret 9 not-for-export
banner motd ^C
vlan 4001
interface Gi0/99
^C
vlan 445
!
interface Gi0/1
 description token=synthetic-token
 switchport mode access
 switchport access vlan 445
 authentication password hidden-password
!"""
    facts = observe_configuration(config, "cisco_ios", source_filename="example.cfg")
    text = json.dumps([asdict(n) for n in facts])
    assert all(secret not in text for secret in ("not-for-export", "synthetic-token", "hidden-password", "Gi0/99", "4001"))
    results = audit(config)
    assert results[0]["observed"]["valid_database_vlans"] == [445]
    assert all(secret not in json.dumps(results) for secret in ("not-for-export", "synthetic-token", "hidden-password"))


def test_uncertain_family_is_explicit_review():
    ctx = replace(context(), device_family="unknown")
    snapshot = VlanState("switch", "192.0.2.1", "cisco_ios", (), (), NOW, ())
    results = evaluate_vlan_compliance(ctx, [], snapshot, POLICY)
    assert all(f["status"] == "unable_to_assess" for f in results)
    assert all(f["reason"] == "unsupported_or_uncertain_family" for f in results)


@pytest.mark.parametrize("family", FAMILIES)
def test_all_required_vlans_compliant_for_every_family(family):
    required = [445, 545, *range(2400, 2445), 2449, 4001]
    ids = ",".join(map(str, required))
    if family in {"C3750X", "C3850"}:
        config = f"vlan {ids}\n!\ninterface Gi0/1\n switchport trunk encapsulation dot1q\n switchport mode trunk\n!"
    elif family in {"ASR920", "ME3600X"}:
        config = (f"vlan {ids}\n!\n" if family == "ME3600X" else "") + "interface Gi0/1\n"
        for sid, vid in enumerate(required, 100):
            config += f" service instance {sid} ethernet\n  encapsulation dot1q 3000\n  bridge-domain {vid}\n"
        config += "!"
    elif family == "NCS540":
        config = "interface TenGigE0/0/0/1\n!\n"
        for sid, vid in enumerate(required, 100):
            config += f"interface TenGigE0/0/0/1.{sid} l2transport\n encapsulation dot1q {vid}\n!\n"
        config += "l2vpn\n bridge group example\n  bridge-domain not-a-number\n"
        config += "".join(f"   interface TenGigE0/0/0/1.{sid}\n" for sid in range(100, 100 + len(required)))
        config += "!"
    elif family in {"NE05", "NE05E"}:
        config = "vlan batch 445 545 2400 to 2444 2449 4001\n#\ninterface GE0/1\n port link-type trunk\n port trunk allow-pass vlan 445 545 2400 to 2444 2449 4001\n#"
    else:
        config = f"vlan database\nvlan {ids}\nexit\ninterface 0/1\nvlan participation include {ids}\nvlan tagging {ids}\nexit"
    results = audit(config, family)
    assert len(results) == 2
    assert results[0]["observed"]["valid_database_vlans"] == required
    assert results[1]["observed"]["valid_interface_vlans"] == required
    assert all(f["status"] == "compliant" and f["missing_vlans"] == [] for f in results)


def test_repeated_edge_stanzas_preserve_order_and_tag_disable():
    config = "vlan database\nvlan 445\nvlan 545,2449\nexit\ninterface 0/1\nvlan participation include 445,545,2449\nvlan tagging 445,545,2449\nexit\ninterface 0/1\nvlan participation exclude 445\nvlan participation include 445\nvlan tagging 445 disable\nvlan pvid 445\nexit"
    results = audit(config, "EdgeSwitch")
    assert len(results) == 2
    assert results[1]["observed"]["interface_type"] == "hybrid"
    assert results[1]["observed"]["tagged"] == [545, 2449]
    assert results[1]["observed"]["untagged"] == [445]
    headings = [s for s in results[1]["evidence"]["sources"] if s["excerpt"] == "interface 0/1"]
    assert len(headings) == 2


def test_xr_routed_dot1q_detail_is_not_unbound_l2_error():
    results = audit("interface TenGigE0/0/0/1\n!\ninterface TenGigE0/0/0/1.4001\n encapsulation dot1q 4001\n!", "NCS540")
    assert results[0]["observed"]["valid_database_vlans"] == []
    assert results[1]["status"] == "not_applicable"
    assert "UNBOUND_L2TRANSPORT_SUBINTERFACE" not in codes(results[1])


def test_no_fabricated_missing_statement_evidence():
    result = audit("interface Gi0/1\n switchport mode access\n!")
    assert result[0]["missing_vlans"] == [445, 545, *range(2400, 2445), 2449, 4001]
    assert all("vlan " not in s["excerpt"] for s in result[0]["evidence"]["sources"])


def test_normalized_routed_catalyst_port_is_not_unresolved_switchport():
    result = audit("interface Gi0/1\n no switchport\n ip address 192.0.2.1 255.255.255.0\n!")
    assert result[1]["status"] == "not_applicable"
    assert "UNRESOLVED_SWITCHPORT_MODE" not in codes(result[1])


@pytest.mark.parametrize("family", ["ASR920", "ME3600X"])
def test_evc_split_horizon_required_range_preserves_identity_and_evidence(family):
    required = [445, 545, *range(2400, 2445), 2449, 4001]
    config = "interface GigabitEthernet0/1\n"
    for sid, vlan in enumerate(required, 1):
        modifier = " split-horizon group 0" if 2400 <= vlan <= 2444 else ""
        config += (f" service instance {sid} ethernet\n"
                   f"  encapsulation dot1q {vlan}\n  bridge-domain {vlan}{modifier}\n")
    results = audit(config, family, observed=("Gi0/1",))
    port = results[1]
    assert port["status"] == "compliant"
    assert port["missing_vlans"] == []
    assert port["observed"]["valid_interface_vlans"] == required
    assert results[0]["observed"]["valid_database_vlans"] == (required if family == "ASR920" else [])
    assert results[0]["status"] == ("compliant" if family == "ASR920" else "non_compliant")
    sources = port["evidence"]["sources"]
    for vlan in range(2400, 2445):
        excerpt = f"  bridge-domain {vlan} split-horizon group 0"
        source = next(s for s in sources if s["excerpt"] == excerpt)
        assert source["line"] == config.splitlines().index(excerpt) + 1
        assert source["source_filename"] == "synthetic.cfg"


@pytest.mark.parametrize("suffix", ["split-horizon", "split-horizon group", "split-horizon group nope",
                                     "split-horizon group -1", "split-horizon group 0 extra", "arbitrary text"])
def test_evc_bridge_domain_rejects_unknown_or_malformed_modifiers(suffix):
    config = f"interface Gi0/1\n service instance 1 ethernet\n  encapsulation dot1q 2400\n  bridge-domain 2400 {suffix}\n"
    result = audit(config, "ASR920")
    assert result[0]["observed"]["valid_database_vlans"] == []
    assert result[1]["observed"]["valid_interface_vlans"] == []
    assert "UNRESOLVED_SERVICE_INSTANCE" in codes(result[1])
    assert suffix not in json.dumps([asdict(f) for f in observe_configuration(config, "cisco_xe")])


def test_interface_states_and_missing_vlans_survive_spool_export(tmp_path):
    from openpyxl import load_workbook
    from orbitflow.compliance_report import export_compliance_spool
    from orbitflow.execution import DeviceOutcome
    from orbitflow.result_spool import ResultSpool

    config = "vlan 445,545,2449\ninterface GigabitEthernet0/1\n shutdown\n switchport mode trunk\ninterface Gi0/2\n switchport mode access\n"
    ctx = replace(context(), platform="cisco_xe", device_family="C3850")
    snapshot = VlanState("switch", ctx.management_ip, ctx.platform, (), (), NOW,
                         observe_configuration(config, ctx.platform))
    records = [replace(interface("Gi0/1"), admin_status="up", oper_status="down"),
               replace(interface("Gi0/99"), admin_status="down", oper_status="down")]
    findings = evaluate_vlan_compliance(ctx, records, snapshot, POLICY)
    assert findings[1]["observed"]["shutdown"] is True
    assert "Gi0/99" not in json.dumps(findings)
    spool = ResultSpool.create(tmp_path / "runs", "vlan_compliance", 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={"findings": findings, "errors": []})
    path = export_compliance_spool(spool.path, tmp_path / "report.xlsx", cleanup=False)
    workbook = load_workbook(path)
    try:
        values = iter(workbook["Interface Results"].values)
        headers = next(values)
        rows = [dict(zip(headers, cells)) for cells in values]
        for finding, cells in zip(findings[1:], rows):
            assert json.loads(cells["Missing VLANs"]) == finding["missing_vlans"]
        database_values = list(workbook["Database Results"].values)
        assert len(database_values) == 2
        database = dict(zip(*database_values))
        assert database["Status"] == "non_compliant"
        assert json.loads(database["Missing VLANs"]) == findings[0]["missing_vlans"] == [*range(2400, 2445), 4001]
        assert len(rows) == 2
        matched = next(r for r in rows if r["Interface Match"] == "matched")
        assert (matched["Admin Status"], matched["Oper Status"], matched["Shutdown"]) == ("up", "down", "True")
        assert json.loads(matched["Missing VLANs"]) == [*range(2400, 2445), 4001]
        config_only = next(r for r in rows if r["Interface Match"] == "config_only")
        assert config_only["Admin Status"] == config_only["Oper Status"] == "not observed"
    finally:
        workbook.close()


@pytest.mark.parametrize("family", FAMILIES)
def test_unused_configured_interface_is_not_applicable(family):
    name = "0/1" if family == "EdgeSwitch" else "GE0/1" if family == "NE05E" else "Gi0/1"
    config = f"interface {name}\n description unused\n shutdown\n"
    port = row(audit(config, family), name)
    assert port["status"] == "not_applicable"
    assert port["reason"] == "no_vlan_service_configuration"
    assert port["observed"]["admin_status"] == port["observed"]["oper_status"] == "not observed"
    assert port["observed"]["shutdown"] is True


@pytest.mark.parametrize("pvid,tags,valid,untagged,kind", [
    ("445", "545,2449", [445, 545, 2449], [445], "hybrid"),
    ("445", "445,545,2449", [445, 545, 2449], [], "trunk"),
    ("4001", "445,545,2449", [445, 545, 2449], [], "trunk"),
    ("", "445,545,2449", [445, 545, 2449], [], "trunk"),
    ("445", "", [445], [445], "access"),
    ("", "", [], [], "no_membership"),
])
def test_edge_participation_does_not_establish_vlan_role(pvid, tags, valid, untagged, kind):
    config = "vlan database\nvlan 445,545,2400,2449,4001\nexit\ninterface 0/1\nvlan participation include 445,545,2400,2449\n"
    if pvid:
        config += f"vlan pvid {pvid}\n"
    if tags:
        config += f"vlan tagging {tags}\n"
    port = audit(config + "exit", "EdgeSwitch")[1]
    observed = port["observed"]
    assert observed["configured_membership_vlans"] == [445, 545, 2400, 2449]
    assert observed["valid_interface_vlans"] == valid
    assert observed["untagged"] == untagged
    assert observed["interface_type"] == kind
    assert "MULTIPLE_UNTAGGED_MEMBERSHIPS" not in codes(port)
    if kind in {"hybrid", "trunk"}:
        assert port["missing_vlans"] == [*range(2400, 2445), 4001]
    elif kind == "review":
        assert port["status"] == "unable_to_assess"
    else:
        assert port["status"] == "not_applicable"


@pytest.mark.parametrize("family,parent,child", [
    ("C3850", "Gi0/1", "Gi0/1.100"),
    ("NE05E", "GE0/1", "GE0/1.100"),
    ("NCS540", "Te0/0/0/1", "Te0/0/0/1.100"),
])
def test_observed_only_identities_do_not_create_rows_or_child_details(family, parent, child):
    results = audit(f"interface {parent}\n description unused\n", family,
                    observed=(parent, child, "Gi0/99"))
    assert len(results) == 2
    assert results[1]["observed"]["child_interfaces"] == []
    assert child not in json.dumps(results)
    assert "Gi0/99" not in json.dumps(results)
    assert "INTERFACE_NOT_IN_CONFIG" not in json.dumps(results)
    # An observed parent cannot supply a configured parent for an L2 child.
    if family == "NCS540":
        results = audit(f"interface {child} l2transport\n encapsulation dot1q 445\n!", family, observed=(parent,))
        assert len(results) == 2
        assert results[1]["observed"]["interface_match_status"] == "config_only"
        assert "PARENT_INTERFACE_NOT_FOUND" in codes(results[1])

@pytest.mark.parametrize("family", ["C3750X", "C3850", "ME3600X"])
@pytest.mark.parametrize("intent", ["", " switchport mode trunk\n", " switchport mode dynamic\n", " no switchport\n", " switchport trunk allowed vlan 445\n"])
def test_access_vlan_infers_access_unless_switching_intent_conflicts(family, intent):
    result = row(audit("vlan 445\n!\ninterface Gi0/1\n switchport access vlan 445\n" + intent, family), "Gi0/1")
    if intent:
        assert result["status"] == "unable_to_assess"
        assert "CONFLICTING_SWITCHPORT_INTENT" in codes(result)
    else:
        assert result["status"] == "not_applicable"
        assert result["reason"] == "access_excluded"
        assert result["observed"]["interface_type"] == "access"


@pytest.mark.parametrize("name", ["0/1", "lag 1"])
@pytest.mark.parametrize("commands", [
    "vlan participation exclude 1-4094",
    "vlan participation include 445\nvlan participation exclude 445",
    "vlan participation include 445",
    "vlan tagging 445\nvlan pvid 445",
    "vlan participation include 4001\nvlan tagging 4001",
])
def test_edge_empty_effective_membership_is_not_applicable(name, commands):
    config = f"vlan database\nvlan 445\nexit\ninterface {name}\n{commands}\nexit\n"
    if name == "lag 1":
        config += "interface 0/1\naddport 3/1\nexit\n"
    results = audit(config, "EdgeSwitch")
    for port in results[1:]:
        assert port["status"] == "not_applicable"
        assert port["observed"]["valid_interface_vlans"] == []
        assert not port["observed"]["review"]


@pytest.mark.parametrize("family,name", [("ASR920", "BDI445"), ("NCS540", "PW-Ether1"), ("NCS540", "BVI445")])
def test_configured_service_only_rows_retained(family, name):
    results = audit(f"interface {name}\n description service-only\n!\n", family, observed=("Gi0/99",))
    assert len(results) == 2
    port = row(results, name)
    assert port["status"] == "not_applicable"
    assert port["observed"]["interface_match_status"] == "config_only"
    assert "Gi0/99" not in json.dumps(results)


@pytest.mark.parametrize("family", [*FAMILIES, "NE05"])
def test_database_evidence_only_contributing_proof_survives_export(family, tmp_path):
    from openpyxl import load_workbook
    from orbitflow.compliance_report import export_compliance_spool
    from orbitflow.execution import DeviceOutcome
    from orbitflow.result_spool import ResultSpool

    if family in {"C3750X", "C3850", "ME3600X"}:
        config = "vlan 445\n!\ninterface Gi0/1\n description unrelated\n switchport access vlan 545\n service instance 1 ethernet\n  encapsulation dot1q 2449\n  bridge-domain 2449\n"
        expected = ["vlan 445"]
    elif family == "ASR920":
        config = "vlan 4001\n!\ninterface Gi0/1\n description unrelated\n service instance 1 ethernet\n  encapsulation dot1q 3000\n  bridge-domain 445 split-horizon group 0\n service instance 2 ethernet\n  encapsulation dot1q 3999\n!\nbridge-domain 545\n member Gi0/1 service-instance 2\n"
        expected = ["interface Gi0/1", " service instance 1 ethernet", "  encapsulation dot1q 3000", "  bridge-domain 445 split-horizon group 0", " service instance 2 ethernet", "  encapsulation dot1q 3999", "bridge-domain 545", " member Gi0/1 service-instance 2"]
    elif family in {"NE05", "NE05E"}:
        config = "vlan batch 445\n#\nvsi example static\n description unrelated\n#\ninterface GE0/1\n port link-type trunk\n#\ninterface GE0/1.10\n description unrelated\n dot1q termination vid 545\n l2 binding vsi example\n#\ninterface GE0/1.20\n dot1q termination vid 4001\n l2 binding vsi absent\n#\n"
        expected = ["vlan batch 445", "vsi example static", "interface GE0/1.10", " dot1q termination vid 545", " l2 binding vsi example"]
    elif family == "NCS540":
        config = xr_config()
        expected = list(dict.fromkeys(line for line in config.splitlines() if line in {
            "interface TenGigE0/0/0/1.100 l2transport", " encapsulation dot1q 445",
            "interface TenGigE0/0/0/1.200 l2transport", " encapsulation dot1q 545",
            "interface TenGigE0/0/0/1.300 l2transport", " encapsulation dot1q 2449",
            "l2vpn", " bridge group arbitrary", "  bridge-domain named",
            "   interface TenGigE0/0/0/1.100", "   interface TenGigE0/0/0/1.200", "   interface TenGigE0/0/0/1.300"}))
    else:
        config = "vlan database\nvlan 445,545\nvlan name 4001 \"annotation-only\"\nexit\ninterface 0/1\nvlan participation include 445\nvlan tagging 445\nexit\n"
        expected = ["vlan database", "vlan 445,545"]
    findings = audit(config, family)
    proof = findings[0]["evidence"]["sources"]
    assert [e["excerpt"] for e in proof] == expected
    for item in proof:
        assert item["excerpt"] == config.splitlines()[item["line"] - 1]
        assert item["source_filename"] == "synthetic.cfg"
    spool = ResultSpool.create(tmp_path / "runs", "vlan_compliance", 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={"findings": findings, "errors": []})
    path = export_compliance_spool(spool.path, tmp_path / "report.xlsx", cleanup=False)
    workbook = load_workbook(path)
    try:
        database = dict(zip(*list(workbook["Database Results"].values)))
        assert database["Configuration Evidence"] == "\n".join(expected)
        assert database["Evidence Source"] == "synthetic.cfg"
        assert database["Evidence Lines"]
        assert len(list(workbook["Details"].values)) > 1
    finally:
        workbook.close()


@pytest.mark.parametrize("family", ["C3750X", "C3850", "ME3600X"])
@pytest.mark.parametrize("operations,expected", [
    (["445,545,2400-2444,2449", "add 4001"], {445, 545, *range(2400, 2445), 2449, 4001}),
    (["445,545,2449,4001", "remove 4001"], {445, 545, 2449}),
    (["none", "add 445,545,2449,4001"], {445, 545, 2449, 4001}),
    (["all"], {445, 545, *range(2400, 2445), 2449, 4001}),
    (["except 2400-2444,4001"], {445, 545, 2449}),
    (["none", "add 445,545,2449", "remove 2449", "add 4001"], {445, 545, 4001}),
    (["none", "add 445,545,2449", "add 4001", "remove 2449,4001"], {445, 545}),
    (["all", "remove 4001", "add 4001"], {445, 545, *range(2400, 2445), 2449, 4001}),
    (["all", "none", "add 445", "545,2449"], {545, 2449}),
    (["none", "add 4001", "except 4001", "add 4001"], {445, 545, *range(2400, 2445), 2449, 4001}),
])
def test_cisco_ordered_trunk_replay(family, operations, expected):
    statements = [" switchport trunk allowed vlan " + op for op in operations]
    config = ("vlan 445,545,2400-2444,2449,4001\n!\ninterface Gi0/1\n"
              " switchport trunk encapsulation dot1q\n" + "\n".join(statements)
              + "\n switchport mode trunk\n!")
    trunk = row(audit(config, family), "Gi0/1")
    assert trunk["observed"]["valid_interface_vlans"] == sorted(expected)
    assert "UNSUPPORTED_ALLOWED_VLAN_OPERATION" not in codes(trunk)
    trigger = {445, 545} <= expected and bool({2449, 4001} & expected)
    missing = {*range(2400, 2445), 2449, 4001} - expected
    expected_status = ("non_compliant" if missing else "compliant") if trigger else "not_applicable"
    assert trunk["status"] == expected_status
    if trigger:
        assert trunk["missing_vlans"] == sorted(missing)
    sources = [e for e in trunk["evidence"]["sources"] if "allowed vlan" in e["excerpt"]]
    assert [e["excerpt"] for e in sources] == statements
    assert [config.splitlines()[e["line"] - 1] for e in sources] == statements


@pytest.mark.parametrize("family", ["C3750X", "C3850", "ME3600X"])
@pytest.mark.parametrize("invalid", ["add", "remove 4001-2449", "except 4095", "add 445,,545", "add 445 token sensitive-value"])
def test_malformed_ordered_trunk_operation_requires_review(family, invalid):
    config = ("vlan 445,545,2449,4001\n!\ninterface Gi0/1\n switchport mode trunk\n"
              " switchport trunk allowed vlan 445,545,2449\n"
              " switchport trunk allowed vlan " + invalid + "\n!")
    result = audit(config, family)
    trunk = row(result, "Gi0/1")
    assert trunk["status"] == "unable_to_assess"
    assert "UNSUPPORTED_ALLOWED_VLAN_OPERATION" in codes(trunk)
    assert "sensitive-value" not in json.dumps(result)


@pytest.mark.parametrize("family", ["C3750X", "C3850", "ME3600X"])
def test_ordered_replay_preserves_family_database_validation(family):
    config = ("vlan 445,545,2449\n!\ninterface Gi0/1\n switchport mode trunk\n"
              " switchport trunk allowed vlan 445,545,2449\n"
              " switchport trunk allowed vlan add 4001\n!")
    trunk = row(audit(config, family), "Gi0/1")
    assert trunk["observed"]["valid_interface_vlans"] == ([445, 545, 2449, 4001] if family == "ME3600X" else [445, 545, 2449])
    assert ("ALLOWED_VLAN_NOT_IN_DATABASE" in codes(trunk)) == (family != "ME3600X")


@pytest.mark.parametrize("family", ["C3750X", "C3850", "ME3600X"])
def test_ordered_replay_repeated_stanza_evidence_survives_spool_report(family, tmp_path):
    from openpyxl import load_workbook
    from orbitflow.compliance_report import export_compliance_spool
    from orbitflow.execution import DeviceOutcome
    from orbitflow.result_spool import ResultSpool

    config = ("vlan 445,545,2400-2444,2449,4001\n!\ninterface Gi0/1\n"
              " switchport mode trunk\n switchport trunk allowed vlan 445,545,2400-2444,2449\n"
              "!\ninterface Gi0/1\n switchport trunk allowed vlan add 4001\n"
              " switchport trunk allowed vlan remove 2449\n!")
    findings = audit(config, family)
    trunk = row(findings, "Gi0/1")
    assert trunk["missing_vlans"] == [2449]
    spool = ResultSpool.create(tmp_path / "runs", "vlan_compliance", 1)
    with spool.collection():
        spool.append(DeviceOutcome(1), payload={"findings": findings, "errors": []})
    path = export_compliance_spool(spool.path, tmp_path / "report.xlsx", cleanup=False)
    workbook = load_workbook(path)
    try:
        report = dict(zip(*list(workbook["Interface Results"].values)))
        excerpts = report["Configuration Evidence"]
        statements = [line for line in config.splitlines() if "allowed vlan" in line]
        assert [line for line in excerpts.splitlines() if "allowed vlan" in line] == statements
        assert report["Evidence Source"] == "synthetic.cfg"
    finally:
        workbook.close()
