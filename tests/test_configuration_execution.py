"""Approved execution regressions. Synthetic channels only; no network access."""
from contextlib import contextmanager
from dataclasses import replace
from datetime import datetime, timezone
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import json
import socket
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from orbitflow.configuration import execution as ex
from orbitflow.configuration.plans import PlanError, identity
from orbitflow.configuration.v2.catalogue import CATALOGUE
from orbitflow.configuration.v2.planning import plan, request
from orbitflow.configuration.v2.workbook import export_plan, output_columns
from orbitflow.configuration.v2.observation import snapshot
from orbitflow.models import DeviceContext, InterfaceRecord, VlanState
from orbitflow.vendors.configuration_facts import observe_configuration
from orbitflow.vendors.cisco.vlans import parse_ios_running_config
from orbitflow.vendors.ubiquiti.vlans import parse_edgeswitch_config, parse_vlan_lookup
from orbitflow.vendors.cisco.switchport import parse_switchport_output as cisco
from orbitflow.vendors.ubiquiti.switchport import parse_switchport_output as edge
from orbitflow.vendors.common import DeviceCLI, InteractiveCLIError
from orbitflow.transport import DeviceSession, DeviceCredentials
from orbitflow.result_spool import ResultJournal

NOW = datetime.now(timezone.utc)
IP = '192.0.2.117'


def template(family, function):
    return next(t for t in CATALOGUE if (t.family, t.function) == (family, function))


def context(family='C3750X', ip=IP):
    platform, profile = ('ubiquiti_edgeswitch', '') if family == 'EdgeSwitch' else ('cisco_xe', 'c3850_switching') if family == 'C3850X' else ('cisco_ios', 'c3750x_switching')
    return DeviceContext('fixture-' + ip, ip, (ip,), 'fixture', 'fixture', platform,
                         'C3850' if family == 'C3850X' else family, family, profile,
                         (), 'FIXTURE-' + ip, '1', '', NOW, NOW)


def observation(ctx, config):
    ports, objects = parse_edgeswitch_config(config) if ctx.device_family == 'EdgeSwitch' else parse_ios_running_config(config)
    facts = observe_configuration(config, ctx.platform, preserve_full_evidence=True)
    interfaces = tuple(InterfaceRecord('fixture', ctx.management_ip, ctx.platform, n.value,
                       next((c.value for c in n.children if c.kind == 'description'), ''), 'down', 'down', NOW)
                       for n in facts if n.kind == 'interface')
    return snapshot(ctx, interfaces, VlanState('fixture', ctx.management_ip, ctx.platform, ports, objects, NOW, facts))


def make_plan(family='C3750X', function='VLAN Database', config='hostname fixture\nend\n', inputs=None, ip=IP):
    ctx = context(family, ip)
    p = dict(management_ip=ip, **(inputs or dict(vlan_id=445, name='Customer')))
    r = request(template(family, function).id, p)
    original = plan([r], {ip: observation(ctx, config)}, now=NOW)
    return ctx, r, original


@pytest.mark.parametrize('raw, expected', [('445,545,1101-1104,2400-2444,2449,4000,4001', 54), ('ALL', 'ALL'), ('NONE', 'NONE'), ('', [])])
def test_cisco_labels_ranges_optional_fields(raw, expected):
    result = cisco('Name: Gi1/0/1\nSwitchport: Enabled\nAdministrative Mode: static access\nOperational Mode: down\nAccess Mode VLAN: 1033 (Customer)\nTrunking Native Mode VLAN: 1 (default)\nTrunking VLANs Enabled: ' + raw + '\nVepa Enabled: false')
    assert result.facts['access_vlan'] == 1033
    assert result.facts['native_vlan'] == 1
    assert result.fields['Vepa Enabled'] == 'false'
    value = result.facts['allowed_vlans']
    assert len(value) == expected if isinstance(expected, int) else value == expected


def test_wrapped_lists_and_unknown_required_facts():
    result = cisco('Trunking VLANs Enabled: 1,445,\n 545-546\nOperational Mode: trunk (member of bundle Po1)')
    assert result.facts['allowed_vlans'] == [1,445,545,546]
    assert result.facts['enabled'] is None
    assert cisco('Trunking VLANs Enabled: broken').facts['allowed_vlans'] is None
    with pytest.raises(ValueError):
        cisco('Switchport: Enabled\nSwitchport: Disabled')


@pytest.mark.parametrize('pvid, untagged, tagged, name, protection', [
    (1, '1', '', None, 'False'),
    (901, '901', '745,2449,4000-4001', 'UNMS', 'True in Group 0:'),
    (1, '1', '445,545,745-746,2436-2444,2446,2449,4000-4001', None, 'False')])
def test_edge_samples_preserve_overlap_name_and_group_zero(pvid, untagged, tagged, name, protection):
    text = f'General Mode PVID: {pvid}\nGeneral Mode Untagged VLANs: {untagged}\nGeneral Mode Tagged VLANs: {tagged}\nGeneral Mode Forbidden VLANs: {tagged}\nProtected Port: {protection}'
    if name:
        text += '\nName: ' + name
    result = edge(text)
    assert result.facts['tagged_vlans'] == result.facts['forbidden_vlans']
    assert result.facts['pvid'] == pvid and result.facts['name'] == name
    assert result.facts['protection']['group'] == (0 if name else None)
    assert result.facts['untagged_vlans'] == [pvid]


def test_edge_lookup_completeness_independent_current_configured():
    text = 'VLAN ID: 445\nVLAN Name: testing max\nVLAN Type: Static\nPort Current Configured Tagging\n--- ----- ----- -----\n0/1 Exclude Include Tagged\n'
    result = parse_vlan_lookup(text, 445)
    assert result['vlan_type'] == 'Static' and result['ports_complete']
    assert result['ports'][0]['current'] == 'Exclude'
    assert result['ports'][0]['configured'] == 'Include'
    broken = parse_vlan_lookup(text + '0/2 Include Include\n', 445)
    assert broken['status'] == 'present' and not broken['ports_complete']


class Channel:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.sent = []

    def sendall(self, value):
        self.sent.append(value)

    def recv(self, _):
        value = next(self.responses)
        if isinstance(value, BaseException):
            raise value
        return value

    def settimeout(self, value):
        assert value > 0

    def close(self):
        pass


def cli_with(responses, prompt='fixture#'):
    channel = Channel([prompt.encode(), ('terminal length 0\n' + prompt).encode(), *responses])
    session = DeviceSession(SimpleNamespace(invoke_shell=lambda: channel), lambda: None)
    return DeviceCLI(session), channel


@pytest.mark.parametrize('prompt, platform, middle', [
    ('fixture#', 'cisco_ios', []),
    ('(<fixture>) #', 'ubiquiti_edgeswitch', [b'Are you sure you want to ', b'save? (y/n)'])])
@pytest.mark.parametrize('response', ['', '% Error saving configuration'])
def test_save_original_prompt_without_success_marker(prompt, platform, middle, response):
    responses = [b'write memory\n', *middle, (response + '\n' + prompt).encode()]
    cli, channel = cli_with(responses, prompt)
    result = cli.save_configuration(prompt, platform, timeout=20)
    assert result['status'] == 'Saved'
    assert (b'y\n' in channel.sent) == bool(middle)
    assert ex.command_rejected(result['response']) == bool(response)


@pytest.mark.parametrize('last', [socket.timeout(), b''])
def test_save_unknown_preserves_partial_and_never_retries(last):
    cli, channel = cli_with([b'fixture#\n', b'write memory\nBuilding configuration...\n', last])
    result = cli.save_configuration('fixture#', 'cisco_ios', timeout=20)
    assert result['status'] == 'Unknown' and 'Building configuration' in result['response']
    assert channel.sent.count(b'write memory\n') == 1
    with pytest.raises(InteractiveCLIError):
        cli.run_command('write memory')


def test_save_unknown_question_not_answered_and_stale_prompt_not_success():
    cli, channel = cli_with([b'fixture#', b'write memory\nReally erase? (yes/no)', socket.timeout()])
    result = cli.save_configuration('fixture#', 'ubiquiti_edgeswitch', timeout=20)
    assert result['status'] == 'Unknown' and b'y\n' not in channel.sent


def test_save_prompt_newline_and_interrupt_invalidate():
    cli, _ = cli_with([b'write memory\nfixture#\r\n'])
    assert cli.save_configuration('fixture#','cisco_ios',timeout=20)['status'] == 'Saved'
    cli, _ = cli_with([b'write memory\nBuilding configuration...', KeyboardInterrupt()])
    with pytest.raises(KeyboardInterrupt):
        cli.save_configuration('fixture#','cisco_ios',timeout=20)
    assert 'Building configuration' in cli.last_response
    with pytest.raises(InteractiveCLIError):
        cli.run_command('write memory')


def test_send_failure_cannot_retain_prior_sensitive_capture(monkeypatch):
    cli, channel = cli_with([])
    cli.last_response = 'prior secret-bearing configuration capture'
    def failed_send(data):
        raise OSError('send failed')
    monkeypatch.setattr(channel, 'sendall', failed_send)
    with pytest.raises(OSError):
        cli.run_command('configure terminal')
    assert cli.last_response == ''
    with pytest.raises(InteractiveCLIError):
        cli.run_command('configure terminal')


@pytest.mark.parametrize('mode', ['config-vlan', 'config-vl', 'config', 'config-if'])
def test_learned_host_mode_matching(mode):
    cli, _ = cli_with([f'vlan 445\nfixture({mode})#'.encode()])
    cli.execution_mode('fixture#', 'cisco_ios')
    cli.run_command('vlan 445')
    assert cli.prompt == f'fixture({mode})#'
    cli.prompt = 'other(config)#'
    with pytest.raises(InteractiveCLIError):
        cli.execution_mode('fixture#', 'cisco_ios')


@pytest.mark.parametrize('mode', ['Vlan', 'Interface 0/3', 'Config'])
def test_edge_outer_modes(mode):
    from orbitflow.vendors.prompts import is_structural_prompt
    prompt = f'(<fixture>) ({mode})#'
    assert is_structural_prompt(prompt)
    cli, _ = cli_with([('configure\n' + prompt).encode()], '(<fixture>) #')
    cli.execution_mode('(<fixture>) #', 'ubiquiti_edgeswitch')
    assert cli.run_command('configure') == ''
    assert cli.prompt == prompt


def test_prepare_confirmation_digest_and_replay(tmp_path, monkeypatch):
    _, r, original = make_plan()
    service = ex.ExecutionService(tmp_path / 'authority')
    prepared = service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    assert prepared['status'] == 'awaiting_confirmation'
    assert len(prepared['summary']['selected']) == 1
    assert 'partial' in prepared['summary']['save_behavior']
    with pytest.raises(PlanError):
        service.confirm(prepared['reference'], summary_digest='changed', apply=True)
    with pytest.raises(PlanError):
        service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'])
    calls = []
    monkeypatch.setattr(ex, '_run', lambda *a, **kw: calls.append(a) or SimpleNamespace(path=tmp_path))
    service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True)
    with pytest.raises(PlanError, match='already consumed'):
        service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True)
    again = service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    assert again['reference'] == prepared['reference'] and again['status'] == 'completed'
    assert again['run_path'] == str(tmp_path)
    assert len(calls) == 1
    changed = service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest, settings=ex.ApplySettings(save_timeout=301))
    assert changed['summary_digest'] != prepared['summary_digest']


def test_confirmation_concurrent_claim_and_interrupted_claim(tmp_path, monkeypatch):
    _, r, original = make_plan()
    service = ex.ExecutionService(tmp_path)
    prepared = service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    def fail(*args, **kwargs):
        raise KeyboardInterrupt()
    monkeypatch.setattr(ex, '_run', fail)
    def confirm():
        try:
            service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True)
        except BaseException as exc:
            return type(exc).__name__
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: confirm(), range(2)))
    assert sorted(results) == ['KeyboardInterrupt', 'PlanError']
    assert service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)['status'] == 'interrupted'


def test_workbook_protected_and_decisions_preserved(tmp_path):
    _, r, original = make_plan()
    path = tmp_path / 'review.xlsx'
    export_plan(original, path)
    book = load_workbook(path)
    t = template('C3750X', 'VLAN Database')
    book[t.sheet].cell(3, len(t.columns) + len(output_columns(t)), 'Approve')
    book.save(path)
    source = path.read_bytes()
    service = ex.ExecutionService(tmp_path / 'authority')
    prepared = service.prepare_workbook(original, path, expected_digest=original.digest)
    assert len(prepared['summary']['selected']) == 1 and path.read_bytes() == source
    book[t.sheet].cell(3, 1, '192.0.2.2')
    book.save(path)
    with pytest.raises(PlanError):
        service.prepare_workbook(original, path, expected_digest=original.digest)


def test_pending_rejected_unsupported_blocked_rows_skipped(tmp_path):
    _, r, original = make_plan()
    service = ex.ExecutionService(tmp_path)
    for value in ('Pending', 'Reject'):
        prepared = service.prepare(original, {r['row_id']: value}, expected_digest=original.digest)
        assert not prepared['summary']['selected'] and prepared['summary']['skipped'][0]['reason'] == value
    manual = request(template('All', 'Manual CLI').id, dict(management_ip=IP, platform='cisco_ios', mode='configuration', scope='device', commands='show version'))
    original = plan([manual])
    prepared = service.prepare(original, {manual['row_id']: 'Approve'}, expected_digest=original.digest)
    assert not prepared['summary']['selected']


class WorkerCLI:
    instances = []
    rejection = ''
    interrupt = False
    timeout = False

    def __init__(self, session):
        self.edge = getattr(session, 'platform', '') == 'ubiquiti_edgeswitch'
        self.session, self.prompt, self.sent = session, '(<fixture>) #' if self.edge else 'fixture#', []
        self.__class__.instances.append(self)
        self.dead = False

    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.closed = True

    def require_session(self, session):
        if self.dead:
            raise InteractiveCLIError('unsynchronized')
        assert session is self.session

    def execution_mode(self, *args):
        self.require_session(self.session)

    def run_command(self, command, **kwargs):
        self.require_session(self.session)
        self.sent.append(command)
        if command.startswith('name ') and self.interrupt:
            raise KeyboardInterrupt()
        if command.startswith('name ') and self.timeout:
            self.dead = True
            raise TimeoutError()
        if command.startswith('name ') and self.rejection:
            return self.rejection
        if self.edge:
            if command == 'configure':
                self.prompt = '(<fixture>) (Config)#'
            elif command == 'vlan database':
                self.prompt = '(<fixture>) (Vlan)#'
            elif command.startswith('interface '):
                self.prompt = '(<fixture>) (Interface ' + command[len('interface '):] + ')#'
            elif command == 'exit':
                self.prompt = '(<fixture>) (Config)#' if '(Interface ' in self.prompt else '(<fixture>) #'
        elif command == 'configure terminal':
            self.prompt = 'fixture(config)#'
        elif command.startswith('interface '):
            self.prompt = 'fixture(config-if)#'
        elif command.startswith('vlan '):
            self.prompt = 'fixture(config-vl)#'
        elif command == 'end':
            self.prompt = 'fixture#'
        return ''


def fake_runtime(monkeypatch, ctx, obs):
    WorkerCLI.instances = []
    WorkerCLI.rejection, WorkerCLI.interrupt, WorkerCLI.timeout = '', False, False
    @contextmanager
    def connect(*args):
        yield SimpleNamespace(platform=ctx.platform)
    monkeypatch.setattr(ex, 'connect_device', connect)
    monkeypatch.setattr(ex, 'DeviceCLI', WorkerCLI)
    monkeypatch.setattr(ex, 'DeviceInventoryResolver', lambda *args, **kwargs: SimpleNamespace(resolve=lambda *a, **kw: ctx))
    monkeypatch.setattr(ex, '_observe', lambda *args: obs)
    monkeypatch.setattr(ex, 'VlanService', lambda **kw: SimpleNamespace(lookup=lambda *a, **kw: dict(status='present', name='Customer')))
    monkeypatch.setattr(ex, 'ConfigurationService', lambda **kwargs: SimpleNamespace(collect=lambda *args, **kw: 'hostname fixture\ndescription test-user\npassword hidden\n'))
    saves = []
    monkeypatch.setattr(ex, 'save_configuration', lambda *args, **kwargs: saves.append(args) or dict(status='Saved', response=''))
    monkeypatch.setattr(ex, '_verify_observed', lambda *args, **kwargs: dict(status='Verified', evidence='=literal'))
    return saves


def run_prepared(tmp_path, monkeypatch, *, no_op=False):
    ctx, r, original = make_plan(config='hostname fixture\nvlan 445\n name Customer\n!\nend\n' if no_op else 'hostname fixture\nend\n')
    obs = original.to_dict()['observations'][IP]
    saves = fake_runtime(monkeypatch, ctx, obs)
    service = ex.ExecutionService(tmp_path / 'authority')
    prepared = service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    def run():
        return service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True,
                               credentials={IP: DeviceCredentials('test-user', 'hidden')},
                               run_root=tmp_path / 'runs', backups_dir=tmp_path / 'backups', inventory_path=tmp_path / 'inventory.json')
    return run, saves, original


@pytest.mark.parametrize('failure', ['', 'reject', 'timeout'])
def test_execution_exact_commands_partial_save_and_unknown(tmp_path, monkeypatch, failure):
    run, saves, original = run_prepared(tmp_path, monkeypatch)
    WorkerCLI.rejection = '% Invalid input' if failure == 'reject' else ''
    WorkerCLI.timeout = failure == 'timeout'
    spool = run()
    row = next(spool.records())['payload']['rows'][0]
    assert row['status'] == ('Failed' if failure == 'reject' else 'Unknown' if failure == 'timeout' else 'Verified')
    assert len(saves) == (0 if failure == 'timeout' else 1)
    assert WorkerCLI.instances[0].closed
    events = list(ResultJournal(spool.path / 'events.jsonl').records())
    assert any(e['stage'] == 'command_intent' for e in events)
    text = (spool.path / 'events.jsonl').read_text()
    assert 'hidden' not in text
    capture = next((tmp_path / 'backups').rglob('*.txt')).read_text()
    assert 'description test-user' in capture and 'hidden' not in capture
    path = ex.export_results(spool.path, tmp_path / 'results.xlsx')
    book = load_workbook(path)
    assert book[template('C3750X', 'VLAN Database').sheet].max_row == 2
    assert all(c.data_type != 'f' for s in book for row in s for c in row)


def test_noop_fresh_verification_no_writes_or_save(tmp_path, monkeypatch):
    run, saves, _ = run_prepared(tmp_path, monkeypatch, no_op=True)
    spool = run()
    assert next(spool.records())['payload']['rows'][0]['status'] == 'Verified'
    assert not WorkerCLI.instances[0].sent and not saves


def test_interrupted_journal_recovery_no_reapply(tmp_path, monkeypatch):
    run, saves, _ = run_prepared(tmp_path, monkeypatch)
    WorkerCLI.interrupt = True
    with pytest.raises(KeyboardInterrupt):
        run()
    directory = next((tmp_path / 'runs').iterdir())
    events = list(ResultJournal(directory / 'events.jsonl').records())
    assert any(e['stage'] == 'command_intent' and e['command'].startswith('name ') for e in events)
    assert json.loads((directory / 'manifest.json').read_text())['status'] == 'interrupted'
    sent = list(WorkerCLI.instances[0].sent)
    ex.export_results(directory, tmp_path / 'recovered.xlsx')
    assert WorkerCLI.instances[0].sent == sent


@pytest.mark.parametrize('settings', [dict(max_concurrent_devices=0), dict(save_timeout=0), dict(command_timeout=float('nan'))])
def test_invalid_apply_settings(settings):
    with pytest.raises((PlanError, ValueError)):
        ex.ApplySettings(**settings)


@pytest.mark.parametrize('family', ['C3750X', 'C3850X', 'EdgeSwitch'])
def test_standalone_collector_borrows_and_parses(family):
    from orbitflow.capabilities.switchport import SwitchportService, SwitchportCapabilityError
    ctx = context(family)
    target = '0/3' if family == 'EdgeSwitch' else 'GigabitEthernet1/0/1'
    prompt = '(<fixture>) #' if family == 'EdgeSwitch' else 'fixture#'
    command = f'show interfaces switchport {target}' if family == 'EdgeSwitch' else f'show interfaces {target} switchport'
    text = 'Port: 0/3\nGeneral Mode PVID: 901\nGeneral Mode Untagged VLANs: 901\nGeneral Mode Tagged VLANs: 745,2449,4000-4001' if family == 'EdgeSwitch' else 'Name: Gi1/0/1\nSwitchport: Enabled\nAdministrative Mode: static access\nAccess Mode VLAN: 1033 (Customer)'
    responses = [('terminal length 0\n' + prompt).encode()]
    if family != 'EdgeSwitch':
        responses.append(('terminal width 511\n' + prompt).encode())
    responses.append((command + '\n' + text + '\n' + prompt).encode())
    cli, channel = cli_with(responses, prompt)
    result = SwitchportService().collect(cli.session, ctx, target, cli=cli)
    assert result.facts['pvid' if family == 'EdgeSwitch' else 'access_vlan'] == (901 if family == 'EdgeSwitch' else 1033)
    assert cli._channel is channel
    with pytest.raises(SwitchportCapabilityError):
        SwitchportService().collect(cli.session, ctx, target + '\nwrite memory', cli=cli)


@pytest.mark.parametrize('function, before_config, after_config, inputs, facts, expected', [
    ('Access Port', 'interface GigabitEthernet1/0/1\n switchport mode access\n switchport access vlan 545\n shutdown\n!\nend\n',
     'interface GigabitEthernet1/0/1\n switchport mode access\n switchport access vlan 445\n shutdown\n!\nend\n',
     dict(interface='GigabitEthernet1/0/1', vlan_id=445),
     dict(enabled=True, administrative_mode='static access', access_vlan=445), 'Verified'),
    ('Access Port', 'interface GigabitEthernet1/0/1\n switchport mode access\n switchport access vlan 545\n shutdown\n!\nend\n',
     'interface GigabitEthernet1/0/1\n switchport mode access\n switchport access vlan 445\n!\nend\n',
     dict(interface='GigabitEthernet1/0/1', vlan_id=445),
     dict(enabled=True, administrative_mode='static access', access_vlan=445), 'Failed'),
    ('Trunk Port', 'interface GigabitEthernet1/0/1\n switchport mode trunk\n switchport trunk native vlan 1\n switchport trunk allowed vlan 445\n!\nend\n',
     'interface GigabitEthernet1/0/1\n switchport mode trunk\n switchport trunk native vlan 1\n switchport trunk allowed vlan 445,545\n!\nend\n',
     dict(interface='GigabitEthernet1/0/1', action='Add', vlans=[545]),
     dict(enabled=True, administrative_mode='trunk', native_vlan=1, allowed_vlans=[445,545]), 'Verified'),
    ('Trunk Port', 'interface GigabitEthernet1/0/1\n switchport mode trunk\n switchport trunk allowed vlan 445\n!\nend\n',
     'interface GigabitEthernet1/0/1\n switchport mode trunk\n switchport trunk allowed vlan 445,545,700\n!\nend\n',
     dict(interface='GigabitEthernet1/0/1', action='Add', vlans=[545]),
     dict(enabled=True, administrative_mode='trunk', native_vlan=1, allowed_vlans=[445,545,700]), 'Failed'),
    ('Access Port', 'interface GigabitEthernet1/0/1\n switchport mode access\n switchport access vlan 545\n!\nend\n',
     'interface GigabitEthernet1/0/1\n switchport mode access\n switchport access vlan 445\n!\nend\n',
     dict(interface='GigabitEthernet1/0/1', vlan_id=445),
     dict(enabled=None, administrative_mode='static access', access_vlan=445), 'Unknown'),
])
def test_actual_verification_complete_sets_down_and_preservation(monkeypatch, function, before_config, after_config, inputs, facts, expected):
    ctx, r, original = make_plan(function=function, config=before_config, inputs=inputs)
    before = original.to_dict()['observations'][IP]
    approved = ex.evaluate_decisions(original, {r['row_id']: 'Approve'}, expected_digest=original.digest, now=NOW).to_dict()['results'][0]
    after = observation(ctx, after_config)
    after['vlan_lookups'] = {'445': dict(status='absent')}
    monkeypatch.setattr(ex, '_observe', lambda *args: after)
    monkeypatch.setattr(ex, 'SwitchportService', lambda **kwargs: SimpleNamespace(collect=lambda *a, **kw: SimpleNamespace(to_dict=lambda: dict(facts=facts))))
    result = ex.verify_row(object(), ctx, object(), dict(request=r, result=approved), before)
    assert result['status'] == expected


@pytest.mark.parametrize('configured, current, tagged, complete, expected', [
    ('Include', 'Exclude', 'Tagged', True, 'Verified'),
    ('Exclude', 'Include', 'Tagged', True, 'Failed'),
    ('Include', 'Include', 'Untagged', True, 'Failed'),
    ('Include', 'Include', 'Tagged', False, 'Unknown')])
def test_edge_verification_independent_current_configured(monkeypatch, configured, current, tagged, complete, expected):
    config = 'vlan database\nvlan 445,545\nexit\ninterface 0/1\nvlan pvid 445\nvlan participation include 445\nexit\n'
    ctx, r, original = make_plan('EdgeSwitch', 'VLAN Port', config=config,
                               inputs=dict(interface='0/1', action='Add/Change', tagged_vlans='545'))
    before = original.to_dict()['observations'][IP]
    approved = ex.evaluate_decisions(original, {r['row_id']: 'Approve'}, expected_digest=original.digest, now=NOW).to_dict()['results'][0]
    assert approved['deployment_ready']
    after = observation(ctx, config)
    text = f'VLAN ID: 545\nVLAN Name: Customer\nVLAN Type: Static\nInterface Current Configured Tagging\n--------- ------- ---------- --------\n0/1 {current} {configured} {tagged}\n'
    if not complete:
        text += '0/2 Include Include\n'
    from orbitflow.capabilities.vlans import VlanService
    session = object()
    lookup_cli = SimpleNamespace(require_session=lambda s: None, configure=lambda **k: None, run_command=lambda *a, **k: text)
    after['vlan_lookups'] = {'545': VlanService().lookup(session, ctx, 545, cli=lookup_cli)}
    assert after['vlan_lookups']['545']['ports_complete'] == complete
    monkeypatch.setattr(ex, '_observe', lambda *args: after)
    facts = dict(pvid=445, tagged_vlans=[545], untagged_vlans=[445], forbidden_vlans=[545])
    monkeypatch.setattr(ex, 'SwitchportService', lambda **kwargs: SimpleNamespace(collect=lambda *a, **kw: SimpleNamespace(to_dict=lambda: dict(facts=facts))))
    result = ex.verify_row(object(), ctx, object(), dict(request=r, result=approved), before)
    assert result['status'] == expected


@pytest.mark.parametrize('status, name, expected', [('present','Customer','Verified'), ('present','Other','Failed'), ('absent','','Failed'), ('unknown','','Unknown'), ('error','','Unknown')])
def test_database_verification_unknown_vs_absent(monkeypatch, status, name, expected):
    ctx, r, original = make_plan()
    monkeypatch.setattr(ex, 'VlanService', lambda **kw: SimpleNamespace(lookup=lambda *a, **kw: dict(status=status, name=name, operational_status='inactive')))
    value = ex.verify_row(object(), ctx, object(), dict(request=r, result=original.to_dict()['results'][0]), {})
    assert value['status'] == expected
    assert value['warnings'] == ['VLAN_INACTIVE']


def test_safe_continuation_after_failed_row_and_device_isolation(tmp_path, monkeypatch):
    from threading import Barrier
    ctx, r, original = make_plan()
    obs = original.to_dict()['observations'][IP]
    saves = fake_runtime(monkeypatch, ctx, obs)
    requests = [r, request(r['template_id'], dict(r['inputs'], vlan_id=545), row=4)]
    other_ip = '192.0.2.118'
    other_ctx = replace(context(ip=other_ip), hostname='fixture-other')
    other = request(r['template_id'], dict(r['inputs'], management_ip=other_ip), row=5)
    other_obs = observation(other_ctx, 'hostname fixture\nend\n')
    original = plan([*requests, other], {IP: obs, other_ip: other_obs}, now=NOW)
    service = ex.ExecutionService(tmp_path / 'authority')
    prepared = service.prepare(original, {q['row_id']: 'Approve' for q in [*requests, other]}, expected_digest=original.digest, settings=ex.ApplySettings(max_concurrent_devices=2, connection_start_interval=0))
    barrier = Barrier(2)
    @contextmanager
    def connect(ip, *args):
        barrier.wait(timeout=10)
        yield SimpleNamespace(ip=ip)
    monkeypatch.setattr(ex, 'connect_device', connect)
    monkeypatch.setattr(ex, 'DeviceInventoryResolver', lambda *a, **kw: SimpleNamespace(resolve=lambda session, **kw: ctx if session.ip == IP else other_ctx))
    monkeypatch.setattr(ex, '_observe', lambda session, *a: obs if session.ip == IP else other_obs)
    monkeypatch.setattr(ex, '_verify_observed', lambda session, context, cli, entry, *a, **kw: dict(status='Failed' if entry['request']['inputs']['vlan_id'] == 445 and session.ip == IP else 'Verified'))
    progress_events = []
    spool = service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True,
                           credentials={ip: DeviceCredentials('user','hidden') for ip in (IP,other_ip)},
                           run_root=tmp_path / 'runs', backups_dir=tmp_path / 'backups', inventory_path=tmp_path / 'inventory.json', progress=progress_events.append)
    records = {v['target']: v['payload'] for v in spool.records()}
    assert [v['status'] for v in records[IP]['rows']] == ['Failed','Verified']
    assert records[other_ip]['rows'][0]['status'] == 'Verified'
    assert len(saves) == 2 and len(WorkerCLI.instances) == 2
    assert WorkerCLI.instances[0].session is not WorkerCLI.instances[1].session
    assert all(cli.closed for cli in WorkerCLI.instances)

    for ip, entries in ((IP, requests), (other_ip, [other])):
        sources = {f"{template('C3750X','VLAN Database').sheet} row {e.get('row', '?')}" for e in entries}
        selected = [e for e in progress_events if e['management_ip'] == ip and e['stage'] == 'command_intent']
        assert selected and all(e['source'] in sources for e in selected)
        assert all(e['run_id'] == spool.path.name for e in selected)
    assert progress_events[-2]['devices_completed'] == 2


def test_backup_failure_blocks_all_writes(tmp_path, monkeypatch):
    run, saves, _ = run_prepared(tmp_path, monkeypatch)
    def fail(*args):
        raise OSError('secret-bearing unsafe message')
    monkeypatch.setattr(ex.BackupWriter, 'write', fail)
    spool = run()
    value = next(spool.records())['payload']
    assert value['rows'][0]['status'] == 'Not executed' and not saves
    assert not WorkerCLI.instances[0].sent
    assert 'secret-bearing unsafe message' not in (spool.path / 'events.jsonl').read_text()


def test_drift_blocks_without_rerendering(tmp_path, monkeypatch):
    run, saves, _ = run_prepared(tmp_path, monkeypatch)
    # Current device is already correct: do not silently substitute no-op CLI.
    ctx = context()
    obs = observation(ctx, 'hostname fixture\nvlan 445\n name Customer\n!\nend\n')
    monkeypatch.setattr(ex, '_observe', lambda *args: obs)
    spool = run()
    assert next(spool.records())['payload']['rows'][0]['status'] == 'Not executed'
    assert not saves and not WorkerCLI.instances[0].sent


@pytest.mark.parametrize('family', ['C3750X', 'C3850X', 'EdgeSwitch'])
def test_all_approved_family_database_exact_execution(tmp_path, monkeypatch, family):
    ctx, r, original = make_plan(family=family)
    obs = original.to_dict()['observations'][IP]
    saves = fake_runtime(monkeypatch, ctx, obs)
    service = ex.ExecutionService(tmp_path / 'authority')
    prepared = service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    assert len(prepared['summary']['selected']) == 1
    approved = prepared['summary']['selected'][0]['result']['commands']
    spool = service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True,
                           credentials={IP: DeviceCredentials('user','hidden')}, run_root=tmp_path / 'runs',
                           backups_dir=tmp_path / 'backups', inventory_path=tmp_path / 'inventory.json')
    assert WorkerCLI.instances[0].sent == approved
    assert next(spool.records())['payload']['rows'][0]['status'] == 'Verified'
    assert len(saves) == 1


@pytest.mark.parametrize('status', ['absent', 'unknown', 'error', 'present'])
def test_edge_execution_prerequisites(tmp_path, monkeypatch, status):
    config = 'vlan database\nvlan 445\nexit\ninterface 0/1\nvlan pvid 445\nvlan participation include 445\nexit\n'
    ctx, r, original = make_plan('EdgeSwitch', 'VLAN Port', config=config,
                               inputs=dict(interface='0/1', action='Add/Change', tagged_vlans='545'))
    obs = original.to_dict()['observations'][IP]
    # Planning deliberately warns about absence without creating the VLAN.
    saves = fake_runtime(monkeypatch, ctx, obs)
    service = ex.ExecutionService(tmp_path / 'authority')
    prepared = service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    assert len(prepared['summary']['selected']) == 1
    obs['vlan_lookups'] = {'545': dict(status=status, name='VLAN 545' if status == 'present' else '', operational_status='Static')}
    if status == 'present':
        obs['objects'].append(dict(object_type='vlan',object_id='545',name='VLAN 545',domain_id='545',vlan_ids=[545]))
    spool = service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True,
                           credentials={IP: DeviceCredentials('user','hidden')}, run_root=tmp_path / 'runs',
                           backups_dir=tmp_path / 'backups', inventory_path=tmp_path / 'inventory.json')
    row = next(spool.records())['payload']['rows'][0]
    assert row['status'] == ('Verified' if status == 'present' else 'Not executed')
    assert bool(WorkerCLI.instances[0].sent) == (status == 'present')


def test_execution_cli_decline_stays_offline(tmp_path, monkeypatch, capsys):
    from orbitflow.configuration import execution_cli
    from orbitflow.configuration.v2.planning import save
    _, r, original = make_plan()
    json_path, workbook = tmp_path / 'plan.json', tmp_path / 'plan.xlsx'
    save(original, json_path)
    export_plan(original, workbook)
    monkeypatch.setattr('builtins.input', lambda prompt: 'n')
    monkeypatch.setattr(ex, 'connect_device', lambda *args: pytest.fail('decline must not connect'))
    execution_cli.main(['execute', str(json_path), str(workbook), '--expected-digest', original.digest,
                        '--authority-dir', str(tmp_path / 'authority'), '--proxy', 'example', '--cluster', 'example',
                        '--bastion-host', 'example', '--bastion-user', 'example'])
    assert 'Execution declined' in capsys.readouterr().out


def test_report_oversized_literal_and_torn_event(tmp_path, monkeypatch):
    run, _, _ = run_prepared(tmp_path, monkeypatch)
    spool = run()
    journal = ResultJournal(spool.path / 'events.jsonl')
    marker = '=FORMULA\n' + 'x' * 65000
    journal.append(dict(stage='extra_evidence', response=marker))
    with journal.path.open('a') as stream:
        stream.write('{"torn":')
    output = ex.export_results(spool.path, tmp_path / 'oversized.xlsx')
    book = load_workbook(output)
    parts = [str(row[3].value) for row in list(book['Evidence Details'].rows)[1:]]
    assert marker in ''.join(parts).replace('\\n', '\n')
    assert all(c.data_type != 'f' for s in book for row in s for c in row)


def test_noop_verification_failure_is_unknown(tmp_path, monkeypatch):
    run, saves, _ = run_prepared(tmp_path, monkeypatch, no_op=True)
    def unavailable(*args, **kwargs):
        raise ValueError('unavailable')
    monkeypatch.setattr(ex, '_verify_observed', unavailable)
    spool = run()
    assert next(spool.records())['payload']['rows'][0]['status'] == 'Unknown'
    assert not saves


@pytest.mark.parametrize('response', ["Couldn't configure pvid.", 'Failed to set VLAN participation.', 'Failed to set VLAN tagging.', 'first line\n  ^\nInvalid input'])
def test_supplied_edge_rejections_even_when_prompt_returned(response):
    assert ex.command_rejected(response)


@pytest.mark.parametrize('save_result', [dict(status='Unknown'), dict(status='Saved', explicit_error=True)])
def test_save_errors_do_not_change_verified_rows(tmp_path, monkeypatch, save_result):
    run, _, _ = run_prepared(tmp_path, monkeypatch)
    monkeypatch.setattr(ex, 'save_configuration', lambda *a, **kw: save_result)
    spool = run()
    record = next(spool.records())
    assert record['payload']['rows'][0]['status'] == 'Verified'
    assert record['payload']['save'] == save_result
    assert record['status'] == 'failed'


def test_journal_failure_latches_before_any_further_append(tmp_path, monkeypatch):
    import orbitflow.result_spool as spooling
    journal = ResultJournal(tmp_path / 'events.jsonl')
    def fail(*args):
        raise OSError('disk full')
    monkeypatch.setattr(spooling.os, 'fsync', fail)
    with pytest.raises(OSError):
        journal.append(dict(stage='intent'))
    with pytest.raises(RuntimeError, match='stopped'):
        journal.append(dict(stage='another intent'))


def test_cisco_lookup_complete_name_with_status_word():
    from orbitflow.vendors.cisco.vlans import parse_vlan_lookup
    result = parse_vlan_lookup('VLAN Name Status Ports\n445 Customer active Service    inactive    Gi1/0/1\n', 445)
    assert result['name'] == 'Customer active Service'
    assert result['operational_status'] == 'inactive'


def test_failed_dependency_blocks_row_while_independent_row_continues(tmp_path, monkeypatch):
    ctx, database, _ = make_plan()
    config = 'hostname fixture\ninterface GigabitEthernet1/0/1\n switchport mode access\n switchport access vlan 545\n!\nend\n'
    obs = observation(ctx, config)
    access = request(template('C3750X','Access Port').id, dict(management_ip=IP, interface='GigabitEthernet1/0/1', vlan_id=445), sheet_order=1)
    independent = request(database['template_id'], dict(database['inputs'], vlan_id=700), row=4)
    requests = [database, independent, access]
    original = plan(requests, {IP:obs}, now=NOW)
    service = ex.ExecutionService(tmp_path / 'authority')
    prepared = service.prepare(original, {r['row_id']:'Approve' for r in requests}, expected_digest=original.digest)
    assert prepared['summary']['selected'][-1]['result']['dependencies'] == [database['row_id']]
    saves = fake_runtime(monkeypatch, ctx, obs)
    monkeypatch.setattr(ex, '_verify_observed', lambda session,context,cli,entry,*a,**kw: dict(status='Failed' if entry['request']['row_id'] == database['row_id'] else 'Verified'))
    spool = service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True,
                           credentials={IP:DeviceCredentials('user','hidden')}, run_root=tmp_path/'runs',
                           backups_dir=tmp_path/'backups', inventory_path=tmp_path/'inventory.json')
    rows = next(spool.records())['payload']['rows']
    assert [r['status'] for r in rows] == ['Failed','Verified','Not executed']
    assert rows[-1]['blocker'] == 'DEPENDENCY_FAILED'
    assert not any(c.startswith('switchport') for c in WorkerCLI.instances[0].sent)
    assert len(saves) == 1


@pytest.mark.parametrize('family', ['EdgeSwitch', 'C3750X'])
@pytest.mark.parametrize('header,ok', [('correct', True), ('missing', False), ('wrong', False), ('duplicate', False)])
def test_collector_requires_unambiguous_target(family, header, ok):
    from orbitflow.capabilities.switchport import SwitchportService, SwitchportCapabilityError
    ctx = context(family)
    target = '0/3' if family == 'EdgeSwitch' else 'Gi1/0/1'
    label = 'Port' if family == 'EdgeSwitch' else 'Name'
    name = target if header != 'wrong' else ('0/99' if family == 'EdgeSwitch' else 'Gi1/0/99')
    text = '' if header == 'missing' else f'{label}: {name}\n'
    if header == 'duplicate':
        text += f'{label}: {name}\n'
    text += 'General Mode PVID: 1\nGeneral Mode Tagged VLANs: \nGeneral Mode Untagged VLANs: 1' if family == 'EdgeSwitch' else 'Switchport: Enabled\nAdministrative Mode: static access'
    session = object()
    calls = []
    cli = SimpleNamespace(require_session=lambda s: calls.append(s), configure=lambda **k: None,
                          run_command=lambda *a, **k: text)
    if ok:
        SwitchportService().collect(session, ctx, target, cli=cli)
    else:
        with pytest.raises(SwitchportCapabilityError):
            SwitchportService().collect(session, ctx, target, cli=cli)
    assert calls == [session]


@pytest.mark.parametrize('baseline', ['', ' switchport\n'])
@pytest.mark.parametrize('family', ['C3750X', 'C3850X'])
@pytest.mark.parametrize('damage', ['', ' shutdown\n', ' switchport trunk encapsulation isl\n'])
def test_new_conversion_trunk_exact_encapsulation(monkeypatch, baseline, family, damage):
    config = 'interface GigabitEthernet1/0/1\n description Customer\n' + baseline + '!\nend\n'
    ctx, r, original = make_plan(family, 'Trunk Port', config,
                               dict(interface='GigabitEthernet1/0/1', action='Replace', vlans=[445,545]))
    approved = ex.evaluate_decisions(original, {r['row_id']: 'Approve'}, expected_digest=original.digest, now=NOW).to_dict()['results'][0]
    assert approved['deployment_ready']
    selector = 'switchport trunk encapsulation dot1q'
    assert (selector in approved['commands']) == (family == 'C3750X')
    after_config = 'interface GigabitEthernet1/0/1\n description Customer\n switchport mode trunk\n switchport trunk allowed vlan 445,545\n'
    if baseline:
        after_config += baseline
    if family == 'C3750X' and 'encapsulation' not in damage:
        after_config += ' ' + selector + '\n'
    after = observation(ctx, after_config + damage + '!\nend\n')
    monkeypatch.setattr(ex, '_observe', lambda *a: after)
    facts = dict(enabled=True, administrative_mode='trunk', native_vlan=1, allowed_vlans=[445,545])
    monkeypatch.setattr(ex, 'SwitchportService', lambda **k: SimpleNamespace(collect=lambda *a, **k: SimpleNamespace(to_dict=lambda: dict(facts=facts))))
    assert ex.verify_row(object(), ctx, object(), dict(request=r,result=approved), original.to_dict()['observations'][IP])['status'] == ('Failed' if damage else 'Verified')


@pytest.mark.parametrize('diagnostic', ['Command not found / Incomplete command. Use ? to list commands.', 'Value is out of range. The valid range is 1 to 4093.'])
def test_supplied_rejections_application_and_save(tmp_path, monkeypatch, diagnostic):
    from orbitflow.capabilities.save import save_configuration
    run, saves, _ = run_prepared(tmp_path, monkeypatch)
    WorkerCLI.rejection = diagnostic
    spool = run()
    assert next(spool.records())['payload']['rows'][0]['status'] == 'Failed'
    events = list(ResultJournal(spool.path / 'events.jsonl').records())
    assert any(e.get('status') == 'Rejected' and diagnostic in e['response'] for e in events)
    prompt = '(<fixture>) #'
    cli, channel = cli_with([('write memory\n' + diagnostic + '\n' + prompt).encode()], prompt)
    result = save_configuration(cli.session, context('EdgeSwitch'), cli=cli, exec_prompt=prompt, timeout=20)
    assert result['status'] == 'Saved' and result['explicit_error']
    assert diagnostic in result['response']


def test_concurrent_capability_metadata_and_safe_diagnostics(tmp_path):
    from orbitflow.capabilities.switchport import SwitchportService, SwitchportCapabilityError
    from orbitflow.capabilities.save import save_configuration
    from orbitflow.logging import module_logger
    def worker(index):
        ctx = context('EdgeSwitch', f'192.0.2.{index}')
        target = f'0/{index}'
        session = object()
        def command(*a, **k):
            if index == 2:
                raise RuntimeError('password=never-log-this')
            return f'Port: {target}\nGeneral Mode PVID: 1'
        cli = SimpleNamespace(require_session=lambda s: None, configure=lambda **k: None, run_command=command,
                              save_configuration=lambda *a, **k: dict(status='Saved', response=''))
        correlation = dict(run_id=f'run-{index}', row_id=f'row-{index}', digest='abc', password='never-log-this')
        with module_logger('test', log_root=tmp_path):
            try:
                SwitchportService().collect(session, ctx, target, cli=cli, correlation=correlation)
            except SwitchportCapabilityError:
                pass
            save_configuration(session, ctx, cli=cli, exec_prompt='fixture#', timeout=1, correlation=correlation)
    with ThreadPoolExecutor(max_workers=2) as pool:
        list(pool.map(worker, [1,2]))
    records = [json.loads(line) for p in tmp_path.rglob('*.log') for line in p.read_text().splitlines()]
    for index in [1,2]:
        selected = [r for r in records if r.get('run_id') == f'run-{index}']
        assert selected and all(r['management_ip'] == f'192.0.2.{index}' and r['row_id'] == f'row-{index}' for r in selected)
        completed = [r for r in selected if 'elapsed_seconds' in r]
        assert len(completed) == 2 and all(float(r['elapsed_seconds']) >= 0 for r in completed)
        assert any(r.get('interface') == f'0/{index}' and r['command'].endswith(f'0/{index}') for r in selected)
    assert any(r.get('exception_chain') for r in records)
    assert 'never-log-this' not in json.dumps(records)

@pytest.mark.parametrize('explicit, operational, enabled, reliable, expected', [
    ('445', None, True, True, 'Verified'), ('4001', None, True, True, 'Failed'),
    ('', None, True, True, 'Unknown'), ('445', None, None, True, 'Unknown'),
    ('445', None, True, False, 'Unknown'), ('445', 201, True, True, 'Failed'),
    ('445\n switchport access vlan 201', None, True, True, 'Unknown')])
def test_short_access_fallback_boundaries(monkeypatch, explicit, operational, enabled, reliable, expected):
    config = 'interface GigabitEthernet1/1/1\n description OrbitflowTest1\n switchport mode access\n switchport access vlan 545\n shutdown\n!\nend\n'
    ctx, r, original = make_plan(function='Access Port', config=config, inputs=dict(interface='GigabitEthernet1/1/1', vlan_id=445, description='OrbitflowTest1'))
    before = original.to_dict()['observations'][IP]
    after_config = config.replace(' switchport access vlan 545\n', (' switchport access vlan ' + explicit + '\n') if explicit else '')
    after = observation(ctx, after_config)
    after['configuration_reliable'] = reliable
    after['vlan_lookups'] = {}
    monkeypatch.setattr(ex, '_observe', lambda *a: after)
    facts = dict(enabled=enabled, administrative_mode='static access', access_vlan=operational)
    monkeypatch.setattr(ex, 'SwitchportService', lambda **k: SimpleNamespace(collect=lambda *a, **k: SimpleNamespace(to_dict=lambda: dict(facts=facts))))
    value = ex.verify_row(object(), ctx, object(), dict(request=r, result=original.to_dict()['results'][0]), before)
    assert value['status'] == expected
    assert facts['access_vlan'] == operational  # parser observation stays unknown
    if operational is None:
        assert value['evidence']['missing_fields'] == ['Access Mode VLAN']
    if expected == 'Verified':
        assert 'explicit configuration' in value['evidence']['access_vlan_source']


def test_default_backup_relative_movable_and_pre_row_durable(tmp_path, monkeypatch):
    ctx, r, original = make_plan()
    fake_runtime(monkeypatch, ctx, original.to_dict()['observations'][IP])
    service = ex.ExecutionService(tmp_path / 'authority')
    prepared = service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    spool = service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True,
                            credentials={IP: DeviceCredentials('test-user', 'hidden')}, run_root=tmp_path/'runs', inventory_path=tmp_path/'inventory.json')
    events = list(ResultJournal(spool.path/'events.jsonl').records())
    backup = next(e for e in events if e['stage'] == 'backup_complete')
    assert Path(backup['backup_reference']).parts[0] == 'backups'
    assert (spool.path / backup['backup_reference']).is_file()
    assert (spool.path/'backups'/'.gitignore').read_text() == '*\n'
    assert next(i for i,e in enumerate(events) if e['stage']=='pre_row_evidence') < next(i for i,e in enumerate(events) if e['stage']=='command_intent')
    import shutil
    moved = tmp_path/'moved'
    shutil.copytree(spool.path, moved)
    assert (moved/backup['backup_reference']).is_file()
    monkeypatch.setattr(ex, 'connect_device', lambda *a: pytest.fail('offline export'))
    moved_report = ex.export_results(moved, tmp_path/'moved.xlsx')
    moved_book = load_workbook(moved_report)
    assert moved_book['Run Summary']['B1'].value == spool.manifest['run_id']
    moved_sheet = moved_book[template('C3750X', 'VLAN Database').sheet]
    assert moved_sheet['A2'].value == spool.manifest['run_id']


@pytest.mark.parametrize('failure', ['', 'reject', 'timeout'])
def test_readable_report_actual_commands_and_evidence(tmp_path, monkeypatch, failure):
    run, _, original = run_prepared(tmp_path, monkeypatch)
    WorkerCLI.rejection = '% Invalid input' if failure == 'reject' else ''
    WorkerCLI.timeout = failure == 'timeout'
    spool = run()
    path = ex.export_results(spool.path, tmp_path/'readable.xlsx')
    book = load_workbook(path)
    sheet = book[template('C3750X', 'VLAN Database').sheet]
    row = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert 'command:' in row['Config Sent']
    assert row['Apply Status'] == {'': 'Applied', 'reject': 'Rejected', 'timeout': 'Unknown'}[failure]
    assert row['Pre-config'] != 'Not captured'
    if failure:
        assert 'end' not in row['Config Sent']
    assert row['Original Inputs'] and row['IP'] == IP


def test_operator_projection_counts_order_and_no_evidence_dump(tmp_path):
    ctx, r, original = make_plan()
    service = ex.ExecutionService(tmp_path/'authority')
    prepared = service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    projection = ex.operator_summary(prepared)
    assert (projection['changes'], projection['no_ops'], projection['skipped'], projection['device_count']) == (1,0,0,1)
    assert projection['devices'][0]['ip'] == IP
    rendered = ex.render_summary(projection)
    assert 'backup before writes' in rendered and 'partial changes' in rendered
    assert 'observations' not in rendered and 'configure terminal' not in rendered
    assert service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)['summary_digest'] == prepared['summary_digest']
    rejected = service.prepare(original, {r['row_id']: 'Reject'}, expected_digest=original.digest)
    projection = ex.operator_summary(rejected)
    assert projection['skipped'] == 1 and projection['device_count'] == 0
    assert IP == rejected['summary']['skipped'][0]['request']['inputs']['management_ip']


def test_old_report_does_not_fabricate_pre_config(tmp_path):
    _, r, original = make_plan()
    prepared = ex.ExecutionService(tmp_path/'authority').prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    run = tmp_path/'old'; run.mkdir()
    (run/'summary.json').write_text(json.dumps(prepared['summary']))
    ResultJournal(run/'events.jsonl').append(dict(stage='row_result', management_ip=IP, row_id=r['row_id'], status='Verified', verification=dict(status='Verified', evidence='old scoped evidence')))
    path = ex.export_results(run, tmp_path/'old.xlsx')
    sheet = load_workbook(path)[template('C3750X', 'VLAN Database').sheet]
    values = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert values['Pre-config'] == 'Not captured'
    assert values['Config Evidence (post-row)'] == 'old scoped evidence'


def test_summary_and_report_order_noop_skipped_and_row_specific_evidence(tmp_path):
    ctx = context()
    obs = observation(ctx, 'vlan 545\n name Customer\n!\nend\n')
    rows = [request(template('C3750X', 'VLAN Database').id, dict(management_ip=IP, vlan_id=v, name='Customer'), sheet_order=1, row=i+3) for i,v in enumerate((445,545,201))]
    original = plan(rows, {IP: obs}, now=NOW)
    prepared = ex.ExecutionService(tmp_path/'authority').prepare(original, {r['row_id']: 'Reject' if i==2 else 'Approve' for i,r in enumerate(rows)}, expected_digest=original.digest)
    projection = ex.operator_summary(prepared)
    assert (projection['changes'], projection['no_ops'], projection['skipped']) == (1,1,1)
    assert projection['devices'][0]['actions'] == {'VLAN Database: Set': 2}
    run = tmp_path/'run'; run.mkdir()
    (run/'summary.json').write_text(json.dumps(prepared['summary']))
    journal = ResultJournal(run/'events.jsonl')
    for i,r in enumerate(rows[:2]):
        journal.append(dict(stage='pre_row_evidence', row_id=r['row_id'], management_ip=IP, evidence=f'pre row {i}'))
        journal.append(dict(stage='row_result', row_id=r['row_id'], management_ip=IP, status='Verified', verification=dict(status='Verified', evidence=f'post row {i}')))
    output = ex.export_results(run, tmp_path/'rows.xlsx')
    sheet = load_workbook(output)[template('C3750X','VLAN Database').sheet]
    headers = [c.value for c in sheet[1]]
    values = [dict(zip(headers,[c.value for c in row])) for row in list(sheet.rows)[1:]]
    assert [v['Row ID'] for v in values] == [r['row_id'] for r in rows]
    assert [v['Pre-config'] for v in values[:2]] == ['pre row 0','pre row 1']
    assert [v['Config Evidence (post-row)'] for v in values[:2]] == ['post row 0','post row 1']
    assert values[1]['Apply Status'] == 'No-op' and values[1]['Config Sent'] == 'None (no-op)'
    assert values[2]['Apply Status'] == 'Skipped' and values[2]['IP'] == IP and 'Reject' in values[2]['Error']


def test_unavailable_switchport_keeps_scoped_post_configuration(monkeypatch):
    config = 'interface GigabitEthernet1/1/1\n description OrbitflowTest1\n switchport mode access\n switchport access vlan 445\n!\nend\n'
    ctx, r, original = make_plan(function='Access Port', config=config, inputs=dict(interface='GigabitEthernet1/1/1', vlan_id=445))
    obs = original.to_dict()['observations'][IP]
    monkeypatch.setattr(ex, '_observe', lambda *a: obs)
    class Unavailable(Exception):
        pass
    def unavailable(*args, **kwargs):
        raise Unavailable()
    monkeypatch.setattr(ex, 'SwitchportService', lambda **k: SimpleNamespace(collect=unavailable))
    value = ex.verify_row(object(), ctx, object(), dict(request=r, result=original.to_dict()['results'][0]), obs)
    assert value['status'] == 'Unknown'
    assert any('switchport access vlan 445' in e['excerpt'] for e in value['evidence']['configuration'])
    assert 'Unavailable' in value['evidence']['warnings'][0]
    short = cisco('Name: Gi1/1/1\nSwitchport: Enabled\nAdministrative Mode: static access')
    assert short.facts['access_vlan'] is None and short.facts['enabled'] is True


def test_report_warning_and_apply_verification_separation(tmp_path):
    _, r, original = make_plan()
    prepared = ex.ExecutionService(tmp_path/'authority').prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    run=tmp_path/'run';run.mkdir()
    (run/'summary.json').write_text(json.dumps(prepared['summary']))
    journal=ResultJournal(run/'events.jsonl')
    journal.append(dict(stage='command_intent', management_ip=IP, row_id=r['row_id'], command='vlan 445'))
    journal.append(dict(stage='command_result', management_ip=IP, row_id=r['row_id'], status='Returned', response='fixture#'))
    journal.append(dict(stage='row_result', management_ip=IP, row_id=r['row_id'], status='Verified', preflight_findings=['GLOBAL_VLAN_REQUIRED'], verification=dict(status='Verified', warnings=['VLAN_INACTIVE'], evidence=dict(warnings=['ACCESS_VLAN_CONFIG_FALLBACK'], configuration='scoped post config'))))
    output=ex.export_results(run,tmp_path/'warnings.xlsx')
    sheet=load_workbook(output)[template('C3750X','VLAN Database').sheet]
    value=dict(zip([c.value for c in sheet[1]],[c.value for c in sheet[2]]))
    assert value['Apply Status']=='Applied' and 'Verified' in value['Status Verification']
    assert value['Error'] is None
    assert all(w in value['Warning'] for w in ('GLOBAL_VLAN_REQUIRED','VLAN_INACTIVE','ACCESS_VLAN_CONFIG_FALLBACK'))


def test_pre_revision_confirmation_metadata_cannot_bypass_replay(tmp_path):
    import hashlib
    _, r, original = make_plan()
    service=ex.ExecutionService(tmp_path/'authority')
    prepared=service.prepare(original,{r['row_id']:'Reject'},expected_digest=original.digest)
    with service._db() as db:
        raw=db.execute('SELECT payload FROM confirmations').fetchone()[0]
        payload=json.loads(raw)
        for skipped in payload['summary']['skipped']:
            skipped.pop('request');skipped.pop('identity')
        legacy_digest=hashlib.sha256(ex.canonical(payload['summary']).encode()).hexdigest()
        db.execute("UPDATE confirmations SET digest=?,payload=?,state='completed'",(legacy_digest,ex.canonical(payload)))
    repeated=service.prepare(original,{r['row_id']:'Reject'},expected_digest=original.digest)
    assert repeated['reference']==prepared['reference']
    assert repeated['summary_digest']==legacy_digest and repeated['status']=='completed'
    with pytest.raises(PlanError):
        service.confirm(repeated['reference'],summary_digest=legacy_digest,apply=True)


def test_malformed_operational_access_field_does_not_use_missing_field_fallback(monkeypatch):
    config='interface GigabitEthernet1/1/1\n switchport mode access\n switchport access vlan 445\n!\nend\n'
    ctx,r,original=make_plan(function='Access Port',config=config,inputs=dict(interface='GigabitEthernet1/1/1',vlan_id=445))
    obs=original.to_dict()['observations'][IP]
    monkeypatch.setattr(ex,'_observe',lambda *a:obs)
    parsed=cisco('Name: Gi1/1/1\nSwitchport: Enabled\nAdministrative Mode: static access\nAccess Mode VLAN: malformed')
    monkeypatch.setattr(ex,'SwitchportService',lambda **k:SimpleNamespace(collect=lambda *a,**k:parsed))
    value=ex.verify_row(object(),ctx,object(),dict(request=r,result=original.to_dict()['results'][0]),obs)
    assert value['status']=='Unknown'
    assert value['evidence']['invalid_fields']==['Access Mode VLAN']
    assert 'access_vlan_source' not in value['evidence']


def test_external_backup_override_different_drive_reference(tmp_path, monkeypatch):
    run,_,_=run_prepared(tmp_path,monkeypatch)
    def different_drive(*args):
        raise ValueError('different drives')
    monkeypatch.setattr(ex.os.path,'relpath',different_drive)
    spool=run()
    events=list(ResultJournal(spool.path/'events.jsonl').records())
    backup=next(e for e in events if e['stage']=='backup_complete')
    assert Path(backup['backup_reference']).is_absolute()
    assert next(spool.records())['payload']['rows'][0]['status']=='Verified'


def test_real_shaped_findings_grouping_and_overlap_presentation(tmp_path):
    _, r, original = make_plan()
    prepared = ex.ExecutionService(tmp_path/'authority').prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    import copy
    sample = 'EXISTING_VLAN_WARNING: ' + json.dumps(dict(code='PVID_NOT_IN_MEMBERSHIP', vlans=[745,746,747], evidence=[dict(excerpt='opaque raw evidence', line=12)]))
    overlap = 'EXISTING_VLAN_WARNING: ' + json.dumps(dict(code='TAGGED_VLAN_NOT_IN_MEMBERSHIP', vlans=[746], evidence=[dict(excerpt='retained overlap')]))
    entries = []
    for i in range(36):
        entry = copy.deepcopy(prepared['summary']['selected'][0])
        entry['request'].update(row=i+3, row_id=f'opaque-{i}')
        entry['result']['findings'] = [sample, overlap]
        if i > 18:
            entry['request']['template_id'] = template('EdgeSwitch', 'Interface Description').id
        entries.append(entry)
    prepared['summary']['selected'] = entries
    prepared['summary']['skipped'] = [dict(request=dict(entries[0]['request'], row=50+i), row_id=f'skipped-{i}', reason='Reject') for i in range(10)]
    text = ex.render_summary(ex.operator_summary(prepared))
    assert 'Skipped: 10' in text and 'Warnings/blockers' in text
    assert 'selected rows: 36' in text and len(text.splitlines()) == 4
    assert 'opaque-' not in text and 'evidence' not in text and 'TAGGED_VLAN_NOT_IN_MEMBERSHIP' not in text
    assert 'PVID_NOT_IN_MEMBERSHIP' not in text and 'row 3' not in text
    assert entries[0]['result']['findings'] == [sample, overlap]


def test_report_compact_main_cells_complete_linked_detail_and_canonical_run(tmp_path):
    _, r, original = make_plan()
    prepared = ex.ExecutionService(tmp_path/'authority').prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    run = tmp_path/'renamed'; run.mkdir()
    (run/'summary.json').write_text(json.dumps(prepared['summary']))
    (run/'manifest.json').write_text(json.dumps(dict(run_id='canonical-execution-identity')))
    overlap = 'EXISTING_VLAN_WARNING: ' + json.dumps(dict(code='TAGGED_VLAN_NOT_IN_MEMBERSHIP', vlans=[746]))
    evidence = dict(configuration=[dict(excerpt='interface 0/3'), dict(excerpt=' vlan pvid 901')],
                    access_vlan_check=dict(actual=901, requested=901), access_vlan_source='explicit configuration',
                    lookups={'901': dict(ports=[dict(interface=f'0/{i}', current='Include', configured='Include', tagging='Tagged') for i in range(100)])})
    ResultJournal(run/'events.jsonl').append(dict(stage='row_result', management_ip=IP, row_id=r['row_id'], status='Verified',
       preflight_findings=[overlap, 'GLOBAL_VLAN_REQUIRED'], verification=dict(status='Verified', evidence=evidence)))
    book = load_workbook(ex.export_results(run, tmp_path/'report.xlsx'))
    sheet = book[template('C3750X', 'VLAN Database').sheet]
    values = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert values['Run ID'] == 'canonical-execution-identity'
    assert book['Run Summary']['B1'].value == values['Run ID']
    assert values['vlan_id'] == '445' and values['name'] == 'Customer'
    assert values['Config Evidence (post-row)'] == 'interface 0/3\n vlan pvid 901'
    assert 'actual=901' in values['Status Verification'] and 'requested=901' in values['Status Verification']
    assert 'explicit configuration' in values['Status Verification']
    assert '0/99' not in values['Status Verification'] and len(values['Status Verification']) < 500
    assert values['Warning'] == 'GLOBAL_VLAN_REQUIRED'
    assert sheet.cell(2, sheet.max_column).hyperlink.target.startswith("#'Evidence Details'!A")
    detail = '\n'.join(str(c.value) for row in book['Evidence Details'] for c in row)
    assert '0/99' in detail and 'TAGGED_VLAN_NOT_IN_MEMBERSHIP' in detail
    assert sheet.row_dimensions[2].height == 90 and sheet.column_dimensions[__import__('openpyxl').utils.get_column_letter(list(values).index('Config Evidence (post-row)')+1)].width == 48


@pytest.mark.parametrize('failure', ['', 'timeout', 'interrupt'])
def test_progress_order_confirmation_cleanup_and_optional_failure(tmp_path, monkeypatch, failure):
    from orbitflow.progress import Progress
    ctx, r, original = make_plan()
    fake_runtime(monkeypatch, ctx, original.to_dict()['observations'][IP])
    WorkerCLI.timeout = failure == 'timeout'
    WorkerCLI.interrupt = failure == 'interrupt'
    service = ex.ExecutionService(tmp_path/'authority')
    prepared = service.prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    events=[]
    assert not WorkerCLI.instances and not events
    kwargs=dict(summary_digest=prepared['summary_digest'], apply=True, credentials={IP: DeviceCredentials('test-user','hidden')},
                run_root=tmp_path/'runs', inventory_path=tmp_path/'inventory.json', progress=events.append)
    if failure == 'interrupt':
        with pytest.raises(KeyboardInterrupt):
            service.confirm(prepared['reference'], **kwargs)
    else:
        service.confirm(prepared['reference'], **kwargs)
    stages=[e['stage'] for e in events]
    assert stages[0]=='confirmed_start'
    assert stages.index('connect') < stages.index('identify') < stages.index('preflight') < stages.index('backup') < stages.index('backup_complete') < stages.index('command_intent')
    assert 'cleanup' in stages and WorkerCLI.instances[0].closed
    assert all(e['management_ip']==IP for e in events if e['stage'] in {'connect','command_intent','cleanup'})
    if not failure:
        assert stages.index('verification_intent') < stages.index('row_result') < stages.index('save_intent') < stages.index('save_result')
        assert events[-2]['devices_completed']==1
    before=len(events)
    with pytest.raises(PlanError):
        service.confirm(prepared['reference'], **kwargs)
    assert len(events)==before
    assert not any(t.name=='configuration-progress' for t in __import__('threading').enumerate())
    with Progress(lambda e: (_ for _ in ()).throw(ValueError())) as reporter:
        reporter.emit(IP,'command_intent')
        reporter.emit(IP,'command_result')
    assert reporter.subscriber is None


def test_progress_waiting_clock_concurrent_attribution_and_stop():
    from orbitflow.progress import Progress
    clock=[0.0]; events=[]
    with Progress(events.append, clock=lambda:clock[0]) as reporter:
        def start(ip):
            reporter.emit(ip, 'command_intent', hostname='host-'+ip, timeout=10, command_index=1)
        with ThreadPoolExecutor(max_workers=2) as pool:
            list(pool.map(start, ['192.0.2.1','192.0.2.2']))
        clock[0]=5
        reporter.tick()
        waits=[e for e in events if e['stage']=='waiting']
        assert {e['management_ip'] for e in waits}=={'192.0.2.1','192.0.2.2'}
        assert all(e['remaining_seconds']==5 and e['awaited_stage']=='command_intent' for e in waits)
        assert not any(e.get('status')=='Verified' for e in events)
        reporter.emit('192.0.2.1','command_unknown',status='Unknown')
        reporter.emit('192.0.2.1','row_result',status='Unknown')
        reporter.emit('192.0.2.1','device_complete',status='Failed')
        clock[0]=11; reporter.tick()
        assert events[-1]['management_ip']=='192.0.2.2' and events[-1]['remaining_seconds']==0
        assert events[-1]['rows_completed']==1 and events[-1]['devices_completed']==1
    assert not reporter.thread.is_alive() and reporter.active=={}


def test_save_progress_recognized_question_and_original_prompt_only():
    from orbitflow.progress import Progress
    prompt='(<fixture>) #'
    cli, channel=cli_with([b'write memory\nAre you sure you want to save? (y/n)', ('\n'+prompt).encode()],prompt)
    events=[]
    with Progress(events.append) as reporter, reporter.save_scope(IP,'fixture',20):
        result=cli.save_configuration(prompt,'ubiquiti_edgeswitch',timeout=20)
    assert result['status']=='Saved'
    assert [e['stage'] for e in events]==['save_confirmation_recognized','save_answer_y_waiting_original_exec']
    assert channel.sent.count(b'y\n')==1
    assert all(e['management_ip']==IP for e in events)


@pytest.mark.parametrize('explicit, source', [('445', 'fresh exact-interface explicit configuration'), ('', 'unavailable')])
def test_fallback_actual_formatted_log_values(tmp_path, monkeypatch, explicit, source):
    from orbitflow.logging import module_logger
    config='interface GigabitEthernet1/1/1\n switchport mode access\n switchport access vlan 545\n!\nend\n'
    ctx,r,original=make_plan(function='Access Port',config=config,inputs=dict(interface='GigabitEthernet1/1/1',vlan_id=445))
    before=original.to_dict()['observations'][IP]
    after=observation(ctx,config.replace(' switchport access vlan 545\n',f' switchport access vlan {explicit}\n' if explicit else ''))
    after.update(configuration_reliable=True,vlan_lookups={})
    monkeypatch.setattr(ex,'_observe',lambda *a:after)
    monkeypatch.setattr(ex,'SwitchportService',lambda **k:SimpleNamespace(collect=lambda *a,**kw:SimpleNamespace(to_dict=lambda:dict(facts=dict(enabled=True,administrative_mode='static access',access_vlan=None)))))
    with module_logger('test', log_root=tmp_path):
        ex.verify_row(object(),ctx,object(),dict(request=r,result=original.to_dict()['results'][0]),before,correlation=dict(run_id='test-run',row_id=r['row_id']))
    logs=[json.loads(line) for path in tmp_path.rglob('*.log') for line in path.read_text().splitlines()]
    message=next(v for v in logs if 'verification source=' in v['message'])
    assert f'not reported; verification source={source}' in message['message']
    assert '%s' not in message['message'] and 'arguments omitted' not in message['message']
    assert message['management_ip']==IP and message['run_id']=='test-run'


def test_verification_warning_journal_logged_as_useful_message(tmp_path, monkeypatch):
    from orbitflow.logging import module_logger
    run, _, _ = run_prepared(tmp_path, monkeypatch)
    monkeypatch.setattr(ex, '_verify_observed', lambda *a, **kw: dict(status='Verified', evidence=dict(warnings=['ACCESS_VLAN_CONFIG_FALLBACK: Access Mode VLAN not reported'])))
    with module_logger('test', log_root=tmp_path/'logs'):
        spool=run()
    events=list(ResultJournal(spool.path/'events.jsonl').records())
    warning=next(e for e in events if e['stage']=='verification_warning')
    logs=[json.loads(line) for path in (tmp_path/'logs').rglob('*.log') for line in path.read_text().splitlines()]
    message=next(v['message'] for v in logs if 'stage=verification_warning' in v['message'])
    assert warning['warning'] in message and 'Access Mode VLAN not reported' in message


def test_cli_decline_emits_no_progress_or_device_access(tmp_path, monkeypatch, capsys):
    from orbitflow.configuration import execution_cli
    _, _, original=make_plan()
    json_path=tmp_path/'plan.json'
    # Use the same artifact format as the supported planner writer.
    from orbitflow.configuration.v2.planning import save
    save(original, json_path)
    workbook=tmp_path/'plan.xlsx'; export_plan(original,workbook)
    monkeypatch.setattr('builtins.input',lambda *a:'n')
    monkeypatch.setattr(execution_cli,'console_progress',lambda e:pytest.fail('decline progress'))
    monkeypatch.setattr(ex,'connect_device',lambda *a:pytest.fail('decline connect'))
    execution_cli.main(['execute',str(json_path),str(workbook),'--expected-digest',original.digest,
        '--authority-dir',str(tmp_path/'authority'),'--proxy','fixture','--cluster','fixture','--bastion-host','fixture','--bastion-user','fixture'])
    assert 'Execution declined' in capsys.readouterr().out


def test_slow_save_wait_is_visible_but_timeout_is_unknown():
    from orbitflow.progress import Progress
    events=[]; clock=[0.0]
    cli, channel=cli_with([b'write memory\nBuilding configuration...',socket.timeout()])
    original_recv=channel.recv
    with Progress(events.append,clock=lambda:clock[0]) as reporter:
        reporter.emit(IP,'save_intent',timeout=10)
        def recv(size):
            clock[0]+=5
            reporter.tick()
            return original_recv(size)
        channel.recv=recv
        result=cli.save_configuration('fixture#','cisco_ios',timeout=10)
        reporter.emit(IP,'save_result',status=result['status'])
    assert result['status']=='Unknown'
    assert any(e['stage']=='waiting' and e['awaited_stage']=='save_intent' for e in events)
    assert not any(e.get('status')=='Saved' for e in events)
    assert channel.sent.count(b'write memory\n')==1


def test_progress_display_failure_does_not_repeat_writes(tmp_path,monkeypatch):
    ctx,r,original=make_plan()
    fake_runtime(monkeypatch,ctx,original.to_dict()['observations'][IP])
    service=ex.ExecutionService(tmp_path/'authority')
    prepared=service.prepare(original,{r['row_id']:'Approve'},expected_digest=original.digest)
    def failed(event):
        raise OSError('display unavailable')
    spool=service.confirm(prepared['reference'],summary_digest=prepared['summary_digest'],apply=True,
        credentials={IP:DeviceCredentials('user','hidden')},run_root=tmp_path/'runs',inventory_path=tmp_path/'inventory.json',progress=failed)
    assert WorkerCLI.instances[0].sent==prepared['summary']['selected'][0]['result']['commands']
    assert next(spool.records())['payload']['rows'][0]['status']=='Verified'


def test_save_question_audit_precedes_y_and_mandatory_failure_cannot_retry():
    from orbitflow.progress import Progress
    prompt='(<fixture>) #'
    cli,channel=cli_with([b'write memory\nAre you sure you want to save? (y/n)', ('\n'+prompt).encode()],prompt)
    audited=[]
    def audit(stage):
        assert b'y\n' not in channel.sent
        audited.append(stage)
        raise OSError('journal unavailable')
    with Progress() as reporter,reporter.save_scope(IP,'fixture',20,audit=audit):
        result=cli.save_configuration(prompt,'ubiquiti_edgeswitch',timeout=20)
    assert result['status']=='Unknown' and b'y\n' not in channel.sent
    assert audited==['save_confirmation_recognized']
    assert channel.sent.count(b'write memory\n')==1


def test_console_progress_is_concise_and_flushed(monkeypatch):
    from orbitflow.progress import Progress, console_progress
    printed=[]
    monkeypatch.setattr('builtins.print',lambda *args,**kwargs:printed.append((args,kwargs)))
    with Progress(console_progress) as reporter:
        reporter.correlation.update(run_id='opaque-run',digest='opaque-digest',rows_total=3,devices_total=2)
        reporter.emit(IP,'command_intent',hostname='fixture',command='vlan 445',command_index=2,command_total=4)
    args,kwargs=printed[0]
    assert kwargs['flush'] is True
    assert IP in args[0] and 'fixture' in args[0] and '[command_intent]' in args[0]
    assert 'command 2/4' in args[0]
    assert all(value not in args[0] for value in ('rows=', 'devices=', 'elapsed=', 'vlan 445', 'timeout'))
    assert 'opaque-' not in args[0]


@pytest.mark.parametrize('stage', ['row_preflight', 'command_intent', 'waiting', 'device_failure', 'verification_warning'])
def test_compact_progress_retains_context_and_errors(stage, capsys):
    from orbitflow.progress import console_progress
    event = dict(management_ip=IP, hostname='fixture', stage=stage,
                 source='C3750X VLAN Database row 3', command_index=1, command_total=4,
                 command='secret-free raw command', error_category='TimeoutError',
                 elapsed_seconds=10, waiting_seconds=5, remaining_seconds=20,
                 timeout_seconds=30, rows_completed=2, devices_completed=1)
    console_progress(event)
    line = capsys.readouterr().out.strip()
    assert line.startswith(f'{IP} fixture [{stage}] | C3750X VLAN Database row 3 | command 1/4')
    assert 'error_category=TimeoutError' in line
    assert all(value not in line for value in ('raw command', 'seconds', 'elapsed', 'rows=', 'devices='))
    if stage == 'verification_warning':
        assert 'warning: see configuration logs' in line


def test_compact_console_concurrent_delivery(monkeypatch):
    from orbitflow.progress import Progress, console_progress
    from threading import Lock
    lock = Lock()
    active = [False]
    lines = []
    def output(line, *, flush):
        assert flush
        assert lock.acquire(blocking=False), 'overlapping console writes'
        try:
            assert not active[0]
            active[0] = True
            lines.append(line)
            active[0] = False
        finally:
            lock.release()
    monkeypatch.setattr('builtins.print', output)
    with Progress(console_progress) as progress:
        with ThreadPoolExecutor(max_workers=4) as pool:
            list(pool.map(lambda i: progress.emit(str(i), 'connect', hostname=f'host-{i}'), range(40)))
    assert len(lines) == 40
    assert set(lines) == {f'{i} host-{i} [connect]' for i in range(40)}


@pytest.mark.parametrize('statuses,save,record_status,label', [
    (['Verified'], 'Saved', 'success', 'Completed'),
    (['Verified', 'Failed'], 'Saved', 'failed', 'Partial'),
    (['Failed'], 'Saved', 'failed', 'Failed'),
    (['Unknown'], 'Saved', 'failed', 'Unknown'),
    (['Verified'], 'Unknown', 'failed', 'Unknown'),
    (['Verified'], 'Saved', 'failed', 'Partial'),
])
def test_final_outcome_uses_results(statuses, save, record_status, label):
    from orbitflow.configuration.execution_cli import outcome_summary
    records = [dict(status=record_status, payload=dict(rows=[dict(status=s) for s in statuses], save=dict(status=save)))]
    spool = SimpleNamespace(records=lambda: iter(records), manifest=dict(target_count=1, collection_complete=True))
    text = outcome_summary(spool, 12.34, skipped=2)
    assert text.startswith(f'Execution {label} |')
    assert 'total elapsed: 12.3s' in text and 'skipped=2' in text
    for status in set(statuses):
        assert f'{status}={statuses.count(status)}' in text


def test_final_outcome_interrupted_missing_devices():
    from orbitflow.configuration.execution_cli import outcome_summary
    spool = SimpleNamespace(records=lambda: iter([]), manifest=dict(target_count=2, collection_complete=False))
    text = outcome_summary(spool, 4)
    assert text.startswith('Execution Interrupted') and 'Unknown=2' in text


def test_compact_summary_first_appearance_order_and_retained_detail(tmp_path):
    import copy
    _, r, original = make_plan()
    prepared = ex.ExecutionService(tmp_path/'authority').prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    first = prepared['summary']['selected'][0]
    second = copy.deepcopy(first)
    second['request']['inputs']['management_ip'] = '192.0.2.1'
    second['identity']['hostname'] = 'second'
    prepared['summary']['selected'] = [first, second, copy.deepcopy(first)]
    before = copy.deepcopy(prepared)
    text = ex.render_summary(ex.operator_summary(prepared))
    assert text.splitlines()[1].startswith(f'1. {IP} ')
    assert text.splitlines()[1].endswith('selected rows: 2')
    assert text.splitlines()[2] == '2. 192.0.2.1 second | selected rows: 1'
    assert prepared == before

@pytest.mark.parametrize('failure', [None, RuntimeError, KeyboardInterrupt])
def test_cli_final_elapsed_once_before_report(tmp_path, monkeypatch, capsys, failure):
    from orbitflow.configuration import execution_cli as cli
    _, r, original = make_plan()
    prepared = ex.ExecutionService(tmp_path/'authority').prepare(original, {r['row_id']: 'Approve'}, expected_digest=original.digest)
    spool = SimpleNamespace(path=tmp_path, manifest=dict(target_count=1, collection_complete=True),
        records=lambda: iter([dict(status='failed', payload=dict(rows=[dict(status='Verified'), dict(status='Failed')], save=dict(status='Saved')))]))
    def confirm(*args, **kwargs):
        assert kwargs['apply'] is True and callable(kwargs['progress'])
        if failure:
            raise failure()
        return spool
    service = SimpleNamespace(prepare_workbook=lambda *a, **kw: prepared, confirm=confirm)
    monkeypatch.setattr(cli, 'ExecutionService', lambda *a: service)
    monkeypatch.setattr(cli, 'load', lambda *a: original)
    monkeypatch.setattr('orbitflow.targets.load_device_credentials', lambda *a: {})
    monkeypatch.setattr('builtins.input', lambda prompt: 'y')
    times = iter([100., 112.3])
    monkeypatch.setattr(cli, 'monotonic', lambda: next(times))
    monkeypatch.setattr(cli, 'export_results', lambda *a: tmp_path/'execution-results.xlsx' if not failure else pytest.fail('failed run report'))
    args = ['execute', 'plan.json', 'plan.xlsx', '--expected-digest', original.digest,
            '--proxy', 'example', '--cluster', 'example', '--bastion-host', 'example', '--bastion-user', 'example']
    if failure:
        with pytest.raises(KeyboardInterrupt if failure is KeyboardInterrupt else SystemExit):
            cli.main(args)
    else:
        cli.main(args)
    text = capsys.readouterr().out
    assert text.count('total elapsed:') == 1
    assert 'total elapsed: 12.3s' in text
    if failure:
        assert 'Execution interrupted/failed' in text and 'Execution Completed' not in text
        assert 'report_generation' not in text
        assert 'devices: Unknown/incomplete=1' in text and 'rows: Unknown/incomplete=1' in text
    else:
        assert 'Execution Partial' in text and 'Failed=1' in text
        assert text.index('total elapsed:') < text.index('[report_generation]')


def test_multitemplate_ip_dispatch_row_order_one_backup_and_save(tmp_path, monkeypatch):
    ips = [IP, '192.0.2.1']  # Deliberately different from sorted IP order.
    contexts = {ip: replace(context(ip=ip), hostname=f'fixture-{i}') for i, ip in enumerate(ips)}
    config = 'hostname fixture\ninterface GigabitEthernet1/0/1\n description old\n!\nend\n'
    observations = {ip: observation(contexts[ip], config) for ip in ips}
    rows = [request(template('C3750X', 'VLAN Database').id,
                    dict(management_ip=ip, vlan_id=445, name='Customer'), sheet_order=1, row=i+3)
            for i, ip in enumerate(ips)]
    rows += [request(template('C3750X', 'Interface Description').id,
                     dict(management_ip=ip, interface='GigabitEthernet1/0/1', description='new'), sheet_order=4, row=i+3)
             for i, ip in enumerate(ips)]
    original = plan(rows, observations, now=NOW)
    saves = fake_runtime(monkeypatch, contexts[IP], observations[IP])
    connected, collected = [], []
    @contextmanager
    def connect(ip, *args):
        connected.append(ip)
        yield SimpleNamespace(ip=ip)
    monkeypatch.setattr(ex, 'connect_device', connect)
    monkeypatch.setattr(ex, 'DeviceInventoryResolver', lambda *a, **kw: SimpleNamespace(resolve=lambda session, **kw: contexts[session.ip]))
    def observe(session, context, cli, entries, timeout):
        collected.append((session.ip, [e['request']['row_id'] for e in entries]))
        return observations[session.ip]
    monkeypatch.setattr(ex, '_observe', observe)
    service = ex.ExecutionService(tmp_path/'authority')
    prepared = service.prepare(original, {r['row_id']: 'Approve' for r in rows}, expected_digest=original.digest,
                               settings=ex.ApplySettings(connection_start_interval=0))
    assert len(prepared['summary']['selected']) == 4
    events = []
    spool = service.confirm(prepared['reference'], summary_digest=prepared['summary_digest'], apply=True,
        credentials={ip: DeviceCredentials('user', 'hidden') for ip in ips}, run_root=tmp_path/'runs',
        inventory_path=tmp_path/'inventory.json', progress=events.append)
    assert connected == ips
    assert len(saves) == 2
    for ip in ips:
        expected = [r['row_id'] for r in rows if r['inputs']['management_ip'] == ip]
        record = next(r for r in spool.records() if r['target'] == ip)
        assert [r['row_id'] for r in record['payload']['rows']] == expected
        assert all(r['status'] == 'Verified' for r in record['payload']['rows'])
        assert [row_ids for target, row_ids in collected if target == ip] == [expected, [expected[1]], [expected[1]]]
        stages = [e['stage'] for e in events if e['management_ip'] == ip]
        assert stages.count('backup_complete') == stages.count('save_result') == 1
        assert stages.index('backup_complete') < stages.index('command_intent')
