"""Sequential capabilities must work on devices that allow only one shell ever."""
from io import StringIO
import socket

from openpyxl import load_workbook
import pytest

from orbitflow import reporting
from orbitflow.capabilities import (
    InterfaceService, VlanService, InterfaceCapabilityError, VlanCapabilityError,
)
from orbitflow.inventory import DeviceInventoryResolver, JsonInventoryStore, DeviceInventoryError
from orbitflow.transport import DeviceSession, TransportConfig
from orbitflow.vendors.common import DeviceCLI, InteractiveCLIError
from test_capability_context import DEVICES
from test_interface_capability import CASES as INTERFACES
from test_vlan_capability import CASES as VLANS


class Channel:
    """Inject stale prompts before split echoes; track changed prompts live."""

    def __init__(self, prompt, outputs):
        self.prompt, self.outputs = prompt, outputs
        self.sent, self.pending = [], []
        self.close_calls = 0

    def sendall(self, data):
        assert not self.close_calls
        assert not self.pending, "previous command left unread output"
        command = data.decode().strip()
        self.sent.append(command)
        output = self.outputs[command]
        if isinstance(output, BaseException):
            self.pending = [output]
            return
        old_prompt = self.prompt
        # Prompt changes at capability boundaries must be carried to the next echo.
        if command in {"show inventory", "show chassis", "display esn",
                       "show interfaces description", "display interface description",
                       "show interfaces status all"}:
            self.prompt = self.prompt.replace("sw", "sw-next")
        self.pending = [old_prompt.encode(),
                        f"\r\n{old_prompt}{command}\r\n".encode() if command else b"\r\n",
                        f"{output}\r\n{self.prompt}".encode()]
        if not command:
            self.pending = [self.prompt.encode()]

    def recv(self, size):
        assert self.pending, "read beyond final prompt"
        result = self.pending.pop(0)
        if isinstance(result, BaseException):
            raise result
        return result

    def settimeout(self, timeout):
        assert timeout > 0

    def close(self):
        self.close_calls += 1


class OneShellClient:
    def __init__(self, channel):
        self.channel, self.opens, self.closes = channel, 0, 0

    def invoke_shell(self, **kwargs):
        self.opens += 1
        if self.opens > 1:
            raise EOFError("device rejects a second shell even after closing the first")
        assert not self.closes
        return self.channel

    def close(self):
        assert self.channel.close_calls == 1, "shell must close once before transport"
        self.closes += 1


def device_session(device):
    version, platform = device[:2]
    case = INTERFACES[platform]
    _, _, vlan_command, vlan_output = VLANS[platform]
    outputs = {
        "": "", "terminal length 0": "", "screen-length 0 temporary": "",
        "show version": version if platform != "huawei_vrp" else "% Invalid input",
        "display version": version, "show inventory": "", "show chassis": "", "display esn": "",
        case["command"]: case["output"], vlan_command: vlan_output,
    }
    if platform == "huawei_vrp":
        outputs["terminal length 0"] = "Error: Unrecognized command"
    prompt = "RP/0/RP0/CPU0:sw#" if platform == "cisco_xr" else device[4]
    channel = Channel(prompt, outputs)
    client = OneShellClient(channel)
    return DeviceSession(client, client.close), client, channel


@pytest.mark.parametrize("device", DEVICES, ids=[d[2] for d in DEVICES])
def test_shared_shell_inventory_interface_vlan_prompt_continuity(tmp_path, device):
    session, client, channel = device_session(device)
    resolver = DeviceInventoryResolver(JsonInventoryStore(tmp_path / "inventory.json"))
    with session:
        with DeviceCLI(session) as cli:
            context = resolver.resolve(session, management_ip="192.0.2.1", cli=cli)
            assert channel.close_calls == 0
            assert InterfaceService().collect(session, context, cli=cli)
            assert channel.close_calls == 0
            vlans = VlanService().collect(session, context, cli=cli)
            assert (vlans.interfaces or vlans.objects) and vlans.platform == device[1]
            assert cli.prompt == channel.prompt
            assert "sw-next" in cli.prompt
            assert not channel.pending
            assert channel.sent.count("") == 1
            assert client.opens == 1
        cli.close()
        assert channel.close_calls == 1
        assert client.closes == 0
    assert client.closes == 1


@pytest.mark.parametrize("command", ["", "terminal length 0", "show version",
                                    "show interfaces description", "show running-config"])
def test_shared_shell_exception_cleanup(tmp_path, command):
    session, client, channel = device_session(DEVICES[0])
    channel.outputs[command] = socket.timeout("synthetic timeout")
    resolver = DeviceInventoryResolver(JsonInventoryStore(tmp_path / "inventory.json"))
    error = {"": TimeoutError, "terminal length 0": TimeoutError,
             "show version": DeviceInventoryError,
             "show interfaces description": InterfaceCapabilityError,
             "show running-config": VlanCapabilityError}[command]
    with pytest.raises(error):
        with session:
            with DeviceCLI(session) as cli:
                context = resolver.resolve(session, management_ip="192.0.2.1", cli=cli)
                InterfaceService().collect(session, context, cli=cli)
                VlanService().collect(session, context, cli=cli)
    assert (client.opens, client.closes, channel.close_calls) == (1, 1, 1)


def test_unsynchronized_shell_cannot_send_again_or_reopen():
    session, client, channel = device_session(DEVICES[0])
    with session, DeviceCLI(session) as cli:
        channel.outputs["show version"] = socket.timeout()
        with pytest.raises(TimeoutError):
            cli.run_command("show version")
        sent = list(channel.sent)
        with pytest.raises(InteractiveCLIError, match="unsynchronized"):
            cli.run_command("show running-config")
        assert channel.sent == sent
        assert channel.close_calls == 0
    assert client.opens == 1


def test_shared_shell_rejects_wrong_parent_and_closed_context():
    session, client, channel = device_session(DEVICES[0])
    other = DeviceSession(object(), lambda: None)
    with session:
        with DeviceCLI(session) as cli:
            with pytest.raises(ValueError, match="different DeviceSession"):
                cli.require_session(other)
            with pytest.raises(ValueError):
                cli.run_command("show version", timeout=0)
            assert cli.run_command("show version")
        with pytest.raises(InteractiveCLIError, match="closed"):
            cli.run_command("show version")
    assert client.opens == 1


@pytest.mark.parametrize("device", DEVICES, ids=[d[2] for d in DEVICES])
@pytest.mark.parametrize("failure", [None, "inventory", "interfaces", "vlans", "parse"])
def test_report_real_capabilities_one_shell_and_failure_isolation(tmp_path, monkeypatch, device, failure):
    session, client, channel = device_session(device)
    platform = device[1]
    command = {"inventory": "show version", "interfaces": INTERFACES[platform]["command"],
               "vlans": VLANS[platform][2], "parse": INTERFACES[platform]["command"]}.get(failure)
    if command:
        channel.outputs[command] = "not an interface table" if failure == "parse" else socket.timeout()
    monkeypatch.setattr(reporting, "connect_device", lambda *args: session)
    path = reporting.run_report(
        [{"management_ip": "192.0.2.1", "username": "synthetic-user", "password": "synthetic-password"}],
        TransportConfig("proxy", "cluster", "bastion", "operator"),
        inventory_path=tmp_path / "inventory.json", reports_dir=tmp_path,
        log_root=tmp_path / "logs", output=StringIO(),
    )
    with_workbook = load_workbook(path)
    try:
        errors = list(with_workbook["Run_Errors"].values)[1:]
        if failure is None:
            assert not errors
            assert with_workbook["Interfaces"].max_row > 1
        else:
            assert errors[0][2] == ("interfaces" if failure == "parse" else failure)
        if failure == "parse":
            # A parser error leaves a synchronized shell; VLAN can still succeed.
            assert len(errors) == 1
            assert VLANS[platform][2] in channel.sent
        if failure == "interfaces":
            assert VLANS[platform][2] not in channel.sent
    finally:
        with_workbook.close()
    assert (client.opens, client.closes, channel.close_calls) == (1, 1, 1)
