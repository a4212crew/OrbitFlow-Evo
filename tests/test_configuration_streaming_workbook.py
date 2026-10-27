"""XLSX writer bounds, independent of the still-aggregate immutable plan reader."""
from copy import copy, deepcopy
import tracemalloc
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from test_bulk_configuration_v2 import NOW, template, inputs
from orbitflow.configuration.v2.planning import plan, request
from orbitflow.configuration.v2 import workbook


@pytest.mark.parametrize('count', [88, 500, 1500])
def test_writer_streams_complete_evidence_with_bounded_additional_memory(tmp_path, monkeypatch, count, record_property):
    t = template('All', 'Manual CLI')
    original = plan([request(t.id, inputs(t))], now=NOW)
    data = original.to_dict()
    seed_request, seed_result = data['requests'][0], data['results'][0]
    data['requests'], data['results'] = [], []
    for i in range(count):
        r, result = deepcopy(seed_request), deepcopy(seed_result)
        r.update(row=i + 3, row_id=f'synthetic-{i}')
        result['row_id'] = r['row_id']
        data['requests'].append(r)
        data['results'].append(result)

    class AlreadyDecodedPlan:
        # Isolate writer allocations from the known aggregate JSON decode.
        digest = 'a' * 64
        def to_dict(self):
            return data

    def evidence(r, obs):
        return [dict(source_filename='synthetic-config', line=i + 1,
                     excerpt=f'! synthetic {r["row_id"]} evidence line {i + 1}')
                for i in range(50)]
    monkeypatch.setattr(workbook, 'scoped_evidence', evidence)
    output = tmp_path / 'plan.xlsx'
    tracemalloc.start()
    try:
        workbook.export_plan(AlreadyDecodedPlan(), output)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    record_property('xlsx_writer_peak_python_bytes', peak)
    assert peak < 16 * 1024 * 1024
    with ZipFile(output) as archive:
        # Workbook cells are inline strings, including untrusted formula-like text.
        assert 'xl/sharedStrings.xml' not in archive.namelist()
    book = load_workbook(output, read_only=True)
    try:
        details = book['Evidence Details'].iter_rows(values_only=True)
        assert next(details) == ('Reference', 'Source', 'Line', 'Exact sanitized excerpt')
        for position in range(count * 50):
            row = next(details)
            device, line = divmod(position, 50)
            assert row == (f'synthetic-{device}', 'synthetic-config', str(line + 1),
                           f'! synthetic synthetic-{device} evidence line {line + 1}')
        assert next(details, None) is None
        rows = list(book[t.sheet].iter_rows(min_row=3, values_only=True))
        assert len(rows) == count
        assert all(r[-1] == 'Pending' for r in rows)
        assert book.sheetnames == [*data['layout'], 'Evidence Details', 'Summary']
    finally:
        book.close()


def test_streaming_template_retains_presentation_and_validations(tmp_path):
    t = template('All', 'Manual CLI')
    original = plan([request(t.id, inputs(t), row=5)], now=NOW)
    output = tmp_path / 'plan.xlsx'
    workbook.export_plan(original, output)
    book = load_workbook(output)
    base = workbook._base()
    try:
        intro = book['Introduction']
        assert intro['B1'].value == original.digest
        assert str(intro.merged_cells) == str(base['Introduction'].merged_cells)
        for item in workbook.CATALOGUE:
            sheet, source = book[item.sheet], base[item.sheet]
            assert sheet.freeze_panes == source.freeze_panes
            assert sheet['A1'].hyperlink.target == source['A1'].hyperlink.target
            assert sheet['A1'].comment.text == source['A1'].comment.text
            assert copy(sheet['A2'].font) == copy(source['A2'].font)
            assert copy(sheet['A2'].fill) == copy(source['A2'].fill)
            assert sheet['A2'].comment.text == source['A2'].comment.text
            assert len(sheet.data_validations.dataValidation) == len(source.data_validations.dataValidation) + 1
        assert book[t.sheet].row_dimensions[5].height == 90
        assert book[t.sheet].cell(3, 1).value is None
        assert book[t.sheet].cell(4, 1).value is None
    finally:
        book.close()
        base.close()


def test_streaming_failure_cleans_all_owned_xml_files(tmp_path, monkeypatch):
    from pathlib import Path
    t = template('All', 'Manual CLI')
    original = plan([request(t.id, inputs(t))], now=NOW)
    books = []
    constructor = workbook.Workbook
    def capture(*args, **kwargs):
        book = constructor(*args, **kwargs)
        books.append(book)
        return book
    def fail(*args):
        raise OSError('synthetic evidence read failure')
    monkeypatch.setattr(workbook, 'Workbook', capture)
    monkeypatch.setattr(workbook, 'scoped_evidence', fail)
    output = tmp_path / 'plan.xlsx'
    with pytest.raises(OSError, match='synthetic evidence'):
        workbook.export_plan(original, output)
    assert not output.exists()
    streamed = [b for b in books if b.write_only]
    assert len(streamed) == 1
    for sheet in streamed[0]:
        assert sheet.closed
        assert not Path(sheet._writer.out).exists()


def test_invalid_row_replacement_and_valid_row_precedence_match_cell_writer(tmp_path):
    t = template('All', 'Manual CLI')
    r = request(t.id, inputs(t))
    errors = [dict(sheet=t.sheet, row=row, reason='INVALID_INPUT_ROW',
                   source_inputs={'management_ip': ip})
              for row, ip in [(3, '192.0.2.8'), (5, '192.0.2.9'), (5, '192.0.2.10')]]
    original = plan([r], now=NOW, input_errors=errors)
    output = tmp_path / 'plan.xlsx'
    workbook.export_plan(original, output)
    book = load_workbook(output, read_only=True)
    try:
        rows = list(book[t.sheet].iter_rows(min_row=3, values_only=True))
        assert len(rows) == 3
        assert rows[0][0] == r['source_inputs']['management_ip']
        assert all(value is None for value in rows[1])
        assert rows[2][0] == '192.0.2.10'
        assert rows[2][-1] == 'Pending'
        assert len(list(book['Summary'].values)) == 5
    finally:
        book.close()


def test_oversized_evidence_references_keep_actual_sheet_coordinates(tmp_path, monkeypatch):
    t = template('All', 'Manual CLI')
    original = plan([request(t.id, inputs(t))], now=NOW)
    monkeypatch.setattr(workbook, 'scoped_evidence', lambda *args: [
        dict(source_filename='synthetic', line=i, excerpt='x' * 33000)
        for i in (1, 2)])
    output = tmp_path / 'plan.xlsx'
    workbook.export_plan(original, output)
    book = load_workbook(output, read_only=True)
    try:
        rows = list(book['Evidence Details'].values)
        assert 'Field: Evidence Details!D2;' in rows[1][3]
        assert 'Field: Evidence Details!D3;' in rows[2][3]
        assert original.digest in rows[1][3]
        assert all(len(row[3]) < 32767 for row in rows[1:])
    finally:
        book.close()
