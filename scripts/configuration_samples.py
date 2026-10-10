"""Generate synthetic offline C1-C4 operator inputs and a completed example.

Never reads real inventory or device credentials. Uses documentation-only IPs.
"""

import argparse
from datetime import datetime, timezone
from pathlib import Path
import shutil

from openpyxl import load_workbook

from orbitflow.configuration.jobs import MANUAL, REQUEST, prepare
from orbitflow.configuration.workbooks import write
from orbitflow.inventory.store import JsonInventoryStore
from orbitflow.models import DeviceContext


def generate(output):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    now = datetime(2026, 10, 10, tzinfo=timezone.utc)
    contexts = []
    for index, (name, platform, family, profile, capability) in enumerate([
            ("sample-switch", "cisco_xe", "C3850", "c3850_switching", "classic_switchport"),
            ("sample-evc", "cisco_xe", "ASR920", "asr920_evc", "evc"),
            ("sample-huawei", "huawei_vrp", "NE05E", "ne05e", "dot1q_subinterface")], 1):
        ip = f"192.0.2.{index}"
        context = DeviceContext(name, ip, (ip,), name, "Huawei" if index == 3 else "Cisco", platform,
                                family, family, profile, (capability,), f"SYNTHETIC{index}", "example", "", now, now)
        contexts.append(context)
        JsonInventoryStore(output / "inventory.json", id_factory=lambda name=name: name).reconcile(context)
    rows = [
        ["access", "sample-switch", "GigabitEthernet1/0/1", "access port", ""],
        ["trunk", "sample-switch", "GigabitEthernet1/0/2", "trunk VLAN add", ""],
        ["advanced", "sample-switch", "GigabitEthernet1/0/2", "manual CLI", "engineer-cli"],
        ["evc-local", "sample-evc", "GigabitEthernet0/0/1", "EVC", ""],
        ["evc-global", "sample-evc", "GigabitEthernet0/0/1", "EVC", ""],
        ["tagged", "sample-huawei", "GigabitEthernet0/1/1.3500", "tagged subinterface", ""],
    ]
    manual = [["engineer-cli", 1, "interface GigabitEthernet1/0/2"],
              ["engineer-cli", 2, "description Engineer reviewed uplink"]]
    write(output / "requests.xlsx", {"Requests": (REQUEST, rows), "Manual": (MANUAL, manual)})
    prepare(output / "requests.xlsx", contexts, output / "guided.xlsx")
    shutil.copyfile(output / "guided.xlsx", output / "completed.xlsx")
    shutil.copyfile(output / "guided.json", output / "completed.json")
    book = load_workbook(output / "completed.xlsx")
    for row in book["Requests"].iter_rows(min_row=2):
        row_id = row[0].value
        if row_id.startswith("evc-"):
            row[6].value = "cisco.evc." + row_id.removeprefix("evc-")
            row[7].value = "1"
        if row_id == "advanced":
            row[9].value = "trunk"
    for row in book["Parameters"].iter_rows(min_row=2):
        row_id, template, name = [c.value for c in row[:3]]
        if row_id.startswith("evc-") and template != "cisco.evc." + row_id.removeprefix("evc-"):
            continue
        value = {"vlan": 3500, "vlans": "3500,3501", "outer_vlan": 3500 if row_id == "evc-local" else 3501,
                 "service_instance": 7 if row_id == "evc-local" else 8, "bridge_domain": 3500}.get(name)
        if value is not None:
            row[7].value = value
    book.save(output / "completed.xlsx")
    book.close()
    write(output / "conflicts.xlsx", {"Requests": (REQUEST, [
        ["add", "sample-switch", "GigabitEthernet1/0/2", "trunk VLAN add", ""],
        ["replace", "sample-switch", "GigabitEthernet1/0/2", "trunk VLAN replace", ""],
        ["unknown-device", "missing-device", "GigabitEthernet1/0/1", "access port", ""],
        ["unsupported", "sample-huawei", "GigabitEthernet0/1/1", "access port", ""],
    ])})
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("outputs/validation/configuration-guided-sample"))
    args = parser.parse_args()
    try:
        print("Created synthetic inputs: " + str(generate(args.output)))
    except FileExistsError:
        parser.error("output already exists; choose a new sample directory")


if __name__ == "__main__":
    main()
