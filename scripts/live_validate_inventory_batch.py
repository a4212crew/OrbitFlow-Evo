"""Sequential live inventory validation for targets supplied by Excel."""

from __future__ import annotations

import argparse
import json
import platform as host_platform
import sys
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import TextIO

from openpyxl import load_workbook

from orbitflow.inventory import DeviceInventoryResolver, JsonInventoryStore
from orbitflow.models import DeviceContext
from orbitflow.transport import DeviceCredentials, TransportConfig, connect_device

TELEPORT_PROXY = "teleport.lynhamnetworks.au:443"
TELEPORT_CLUSTER = "lynhamcluster"
BASTION_HOST = "bastion-lyn-dc1-vic"
BASTION_USER = "lightningadmin"
INVENTORY_PATH = Path("data/live_validation/inventory.json")
RESULTS_PATH = Path("data/live_validation/inventory_validation_results.json")
REQUIRED_COLUMNS = ("management_ip", "username", "password")


def _linux_teleport_identity_paths() -> tuple[Path | None, Path | None]:
    if host_platform.system().lower() != "linux":
        return None, None
    proxy_host = TELEPORT_PROXY.rsplit(":", 1)[0]
    key_path = Path.home() / ".tsh" / "keys" / proxy_host / BASTION_USER
    return key_path, key_path.with_name(f"{key_path.name}-cert.pub")


def load_targets(path: str | Path) -> list[dict[str, str]]:
    """Load required target fields while ignoring future-compatible columns."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise ValueError("Excel input is empty")
        headers = {
            str(value).strip().casefold(): index
            for index, value in enumerate(header_row)
            if value is not None
        }
        missing = [name for name in REQUIRED_COLUMNS if name not in headers]
        if missing:
            raise ValueError(
                f"Excel input is missing required columns: {', '.join(missing)}"
            )
        targets = []
        for row_number, row in enumerate(rows, start=2):
            values = {
                name: (
                    ""
                    if headers[name] >= len(row) or row[headers[name]] is None
                    else str(row[headers[name]]).strip()
                )
                for name in REQUIRED_COLUMNS
            }
            if not any(values.values()):
                continue
            if not all(values.values()):
                raise ValueError(f"Excel row {row_number} has an empty required field")
            targets.append(values)
        if not targets:
            raise ValueError("Excel input contains no device rows")
        return targets
    finally:
        workbook.close()


def _serialize_context(context: DeviceContext) -> dict[str, object]:
    raw = asdict(context)
    raw["observed_management_ips"] = list(context.observed_management_ips)
    raw["capability_flags"] = list(context.capability_flags)
    raw["last_successful_collection"] = context.last_successful_collection.isoformat()
    raw["last_collection_attempt"] = context.last_collection_attempt.isoformat()
    return raw


def _safe_error(exc: Exception) -> str:
    """Return useful error classification without potentially secret exception text."""
    return f"inventory validation failed ({type(exc).__name__})"


def _write_json(path: str | Path, payload: dict[str, object]) -> None:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(destination.name + ".tmp")
    with temporary.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2)
        handle.write("\n")
    temporary.replace(destination)


def run_batch_validation(
    targets: list[dict[str, str]],
    transport_config: TransportConfig,
    *,
    inventory_path: str | Path = INVENTORY_PATH,
    results_path: str | Path = RESULTS_PATH,
    output: TextIO = sys.stdout,
    clock=lambda: datetime.now(timezone.utc),
) -> dict[str, object]:
    """Validate every target sequentially and retain per-device snapshots."""
    store = JsonInventoryStore(inventory_path)
    results: list[dict[str, object]] = []
    successful = 0
    for target in targets:
        management_ip = target["management_ip"]
        resolver = DeviceInventoryResolver(store)
        try:
            credentials = DeviceCredentials(target["username"], target["password"])
            with connect_device(
                management_ip, credentials, transport_config
            ) as session:
                context = resolver.resolve(session, management_ip=management_ip)
            result = {
                "management_ip": management_ip,
                "success": True,
                "device_context": _serialize_context(context),
                "reconciliation_events": list(resolver.last_events),
                "error": "",
            }
            successful += 1
            print(f"{management_ip}: success", file=output)
        except Exception as exc:
            result = {
                "management_ip": management_ip,
                "success": False,
                "device_context": None,
                "reconciliation_events": [],
                "error": _safe_error(exc),
            }
            print(f"{management_ip}: failed ({type(exc).__name__})", file=output)
        results.append(result)

    payload = {
        "collection_time": clock().isoformat(),
        "total_devices": len(results),
        "successful": successful,
        "failed": len(results) - successful,
        "results": results,
    }
    _write_json(results_path, payload)
    print(f"results_path: {results_path}", file=output)
    return payload


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("excel_path", help="Excel file containing device targets")
    args = parser.parse_args(argv)
    key_path, cert_path = _linux_teleport_identity_paths()
    config = TransportConfig(
        proxy=TELEPORT_PROXY,
        cluster=TELEPORT_CLUSTER,
        bastion_host=BASTION_HOST,
        bastion_user=BASTION_USER,
        teleport_key_path=key_path,
        teleport_cert_path=cert_path,
    )
    run_batch_validation(load_targets(args.excel_path), config)


if __name__ == "__main__":
    main()
